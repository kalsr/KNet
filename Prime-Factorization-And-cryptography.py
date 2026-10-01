


# PrimeShield: Prime Factorization and Cryptography Suite
# Developed by Randy Singh from Kalsnet (KNet) Consulting Group

# Single file edition. Everything (number theory engine, cryptography,
# Quantum simulation, randomness tests, key auditing, report exporters,
# styling and all use case pages) is contained in this one file.

# from __future__ import annotations

import hashlib
import hmac
import io
import math
import random
import re
import secrets
import textwrap
import time
import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from dataclasses import asdict, dataclass, field
from datetime import datetime, timedelta
from fractions import Fraction
from typing import Callable, Dict, List, Optional, Sequence, Tuple


# ============================================================================
# MODULE: core.numtheory
# Number theory engine for PrimeShield.
#
# Contains primality testing, integer factorization (trial division,
# Pollard rho with Brent cycle detection, Pollard p minus 1, Fermat),
# Euclidean algorithm helpers and prime generation.
#
# All functions work with Python arbitrary precision integers.
# ============================================================================

# ---------------------------------------------------------------------------
# Small primes
# ---------------------------------------------------------------------------
def sieve(limit: int) -> List[int]:
    """Sieve of Eratosthenes returning all primes up to and including limit."""
    if limit < 2:
        return []
    flags = bytearray([1]) * (limit + 1)
    flags[0] = flags[1] = 0
    for i in range(2, int(limit ** 0.5) + 1):
        if flags[i]:
            flags[i * i :: i] = bytearray(len(flags[i * i :: i]))
    return [i for i, f in enumerate(flags) if f]


SMALL_PRIMES: List[int] = sieve(10_000)
_SMALL_PRIME_SET = set(SMALL_PRIMES)

# Deterministic Miller Rabin bases, correct for every n below 3.3e24
_DET_BASES = (2, 3, 5, 7, 11, 13, 17, 19, 23, 29, 31, 37, 41)
_DET_LIMIT = 3_317_044_064_679_887_385_961_981


def is_probable_prime(n: int, extra_rounds: int = 16) -> bool:
    """Miller Rabin primality test.

    Deterministic for n < 3.3e24. For larger n it uses the 13 fixed bases
    plus extra random bases, giving an error probability below 4^-(13+extra).
    """
    if n < 2:
        return False
    if n in _SMALL_PRIME_SET:
        return True
    for p in SMALL_PRIMES[:60]:
        if n % p == 0:
            return False
    d, s = n - 1, 0
    while d % 2 == 0:
        d //= 2
        s += 1

    def witness(a: int) -> bool:
        x = pow(a, d, n)
        if x == 1 or x == n - 1:
            return False
        for _ in range(s - 1):
            x = x * x % n
            if x == n - 1:
                return False
        return True

    for a in _DET_BASES:
        if a % n and witness(a):
            return False
    if n < _DET_LIMIT:
        return True
    rng = random.Random(n)
    for _ in range(extra_rounds):
        if witness(rng.randrange(2, n - 1)):
            return False
    return True


def next_prime(n: int) -> int:
    """Smallest prime strictly greater than n."""
    if n < 2:
        return 2
    c = n + 1
    if c % 2 == 0:
        c += 1
    while not is_probable_prime(c):
        c += 2
    return c


def generate_prime(bits: int, rng: Optional[random.Random] = None,
                   congruent_3_mod_4: bool = False) -> int:
    """Random prime with exactly the requested bit length."""
    if bits < 2:
        raise ValueError("A prime needs at least 2 bits")
    rng = rng or random.SystemRandom()
    if bits == 2:
        return 3 if congruent_3_mod_4 else rng.choice([2, 3])
    while True:
        c = rng.getrandbits(bits) | (1 << (bits - 1)) | 1
        if congruent_3_mod_4:
            c |= 3
        if is_probable_prime(c):
            return c


# ---------------------------------------------------------------------------
# GCD, LCM, inverses
# ---------------------------------------------------------------------------
def gcd(a: int, b: int) -> int:
    return math.gcd(a, b)


def lcm(a: int, b: int) -> int:
    if a == 0 or b == 0:
        return 0
    return abs(a // math.gcd(a, b) * b)


def gcd_list(values: List[int]) -> int:
    g = 0
    for v in values:
        g = math.gcd(g, v)
    return g


def lcm_list(values: List[int]) -> int:
    result = 1
    for v in values:
        result = lcm(result, v)
    return result


def egcd(a: int, b: int) -> Tuple[int, int, int]:
    """Extended Euclid: returns (g, x, y) with a*x + b*y = g."""
    old_r, r = a, b
    old_s, s = 1, 0
    old_t, t = 0, 1
    while r:
        q = old_r // r
        old_r, r = r, old_r - q * r
        old_s, s = s, old_s - q * s
        old_t, t = t, old_t - q * t
    return old_r, old_s, old_t


def modinv(a: int, m: int) -> int:
    g, x, _ = egcd(a % m, m)
    if g != 1:
        raise ValueError(f"{a} has no inverse modulo {m}")
    return x % m


def euclid_steps(a: int, b: int) -> List[Dict[str, int]]:
    """Table of Euclidean algorithm steps: a = q*b + r."""
    steps = []
    k = 1
    while b:
        q, r = divmod(a, b)
        steps.append({"Step": k, "a": a, "b": b, "Quotient q": q, "Remainder r": r})
        a, b = b, r
        k += 1
    return steps


# ---------------------------------------------------------------------------
# Factorization algorithms
# ---------------------------------------------------------------------------
def is_perfect_power(n: int) -> Optional[Tuple[int, int]]:
    """Return (base, k) with base**k == n and k >= 2, if one exists."""
    if n < 4:
        return None
    for k in range(n.bit_length(), 1, -1):
        base = integer_nth_root(n, k)
        if base > 1 and base ** k == n:
            return base, k
    return None


def integer_nth_root(n: int, k: int) -> int:
    """Floor of the k-th root of n."""
    if n < 2:
        return n
    x = 1 << ((n.bit_length() + k - 1) // k)
    while True:
        y = ((k - 1) * x + n // x ** (k - 1)) // k
        if y >= x:
            break
        x = y
    while x ** k > n:
        x -= 1
    while (x + 1) ** k <= n:
        x += 1
    return x


def pollard_rho_brent(n: int, deadline: Optional[float] = None,
                      rng: Optional[random.Random] = None) -> Optional[int]:
    """Pollard rho with Brent cycle detection. Returns a non trivial factor."""
    if n % 2 == 0:
        return 2
    rng = rng or random.Random(n ^ 0x5EED)
    while True:
        y, c, m = rng.randrange(1, n), rng.randrange(1, n), 128
        g, r, q = 1, 1, 1
        x = ys = y
        while g == 1:
            x = y
            for _ in range(r):
                y = (y * y + c) % n
            k = 0
            while k < r and g == 1:
                ys = y
                for _ in range(min(m, r - k)):
                    y = (y * y + c) % n
                    q = q * abs(x - y) % n
                g = math.gcd(q, n)
                k += m
            r *= 2
            if deadline is not None and time.perf_counter() > deadline:
                return None
        if g == n:
            g = 1
            while g == 1:
                ys = (ys * ys + c) % n
                g = math.gcd(abs(x - ys), n)
        if g != n:
            return g


def pollard_p_minus_1(n: int, bound: int = 100_000) -> Optional[int]:
    """Pollard p minus 1. Finds p when p minus 1 is bound smooth."""
    a = 2
    for p in sieve(bound):
        pk = p
        while pk * p <= bound:
            pk *= p
        a = pow(a, pk, n)
    g = math.gcd(a - 1, n)
    if 1 < g < n:
        return g
    return None


def fermat_factor(n: int, max_iterations: int = 1_000_000) -> Optional[Tuple[int, int, int]]:
    """Fermat method: n = a^2 - b^2 = (a-b)(a+b). Fast when p and q are close.

    Returns (p, q, iterations) or None.
    """
    if n % 2 == 0:
        return 2, n // 2, 0
    a = math.isqrt(n)
    if a * a < n:
        a += 1
    for i in range(max_iterations):
        b2 = a * a - n
        b = math.isqrt(b2)
        if b * b == b2:
            p, q = a - b, a + b
            if p > 1:
                return p, q, i + 1
            return None
        a += 1
    return None


@dataclass
class FactorResult:
    n: int
    factors: Dict[int, int] = field(default_factory=dict)
    complete: bool = True
    unfactored: List[int] = field(default_factory=list)
    log: List[Dict[str, str]] = field(default_factory=list)
    elapsed: float = 0.0

    @property
    def is_prime(self) -> bool:
        return self.complete and len(self.factors) == 1 and list(self.factors.values())[0] == 1

    def product(self) -> int:
        p = 1
        for q, e in self.factors.items():
            p *= q ** e
        for u in self.unfactored:
            p *= u
        return p

    def plain_text(self) -> str:
        return format_factorization(self.factors, self.unfactored)

    def latex(self) -> str:
        parts = []
        for p in sorted(self.factors):
            e = self.factors[p]
            parts.append(f"{p}^{{{e}}}" if e > 1 else f"{p}")
        for u in self.unfactored:
            parts.append(r"\underbrace{" + str(u) + r"}_{\text{composite}}")
        return r" \times ".join(parts) if parts else str(self.n)


def format_factorization(factors: Dict[int, int], unfactored: Optional[List[int]] = None) -> str:
    parts = [f"{p}^{e}" if e > 1 else str(p) for p, e in sorted(factors.items())]
    parts += [f"[{u} composite, not split]" for u in (unfactored or [])]
    return " x ".join(parts) if parts else "1"


def factorize(n: int, time_limit: float = 10.0) -> FactorResult:
    """Full prime factorization with a time budget.

    Strategy: trial division by primes below 10 000, then Miller Rabin,
    then perfect power check, Pollard p minus 1 and Pollard rho (Brent),
    recursively. Anything not split within the budget is reported in
    `unfactored` and `complete` is False.
    """
    start = time.perf_counter()
    deadline = start + time_limit
    res = FactorResult(n=n)
    if n < 1:
        raise ValueError("Please enter a positive integer")
    if n == 1:
        res.log.append({"Method": "Definition", "Detail": "1 has no prime factors"})
        return res

    def add(p: int, e: int = 1) -> None:
        res.factors[p] = res.factors.get(p, 0) + e

    m = n
    td_found = []
    for p in SMALL_PRIMES:
        if p * p > m:
            break
        if m % p == 0:
            e = 0
            while m % p == 0:
                m //= p
                e += 1
            add(p, e)
            td_found.append(f"{p}^{e}" if e > 1 else str(p))
    res.log.append({"Method": "Trial division (primes up to 10 000)",
                    "Detail": ", ".join(td_found) if td_found else "no small factors"})

    stack = [m] if m > 1 else []
    while stack:
        c = stack.pop()
        if c == 1:
            continue
        if is_probable_prime(c):
            add(c)
            res.log.append({"Method": "Miller Rabin", "Detail": f"{c} is prime"})
            continue
        pp = is_perfect_power(c)
        if pp:
            base, k = pp
            res.log.append({"Method": "Perfect power check", "Detail": f"{c} = {base}^{k}"})
            stack.extend([base] * k)
            continue
        f = pollard_p_minus_1(c, bound=20_000)
        method = "Pollard p minus 1"
        if f is None:
            f = pollard_rho_brent(c, deadline=deadline)
            method = "Pollard rho (Brent)"
        if f is None:
            res.complete = False
            res.unfactored.append(c)
            res.log.append({"Method": "Time budget reached",
                            "Detail": f"{c} ({c.bit_length()} bits) could not be split in time"})
            continue
        res.log.append({"Method": method, "Detail": f"{c} = {f} x {c // f}"})
        stack.extend([f, c // f])

    res.factors = dict(sorted(res.factors.items()))
    res.unfactored.sort()
    res.elapsed = time.perf_counter() - start
    return res


def factor_tree_dot(n: int, result: FactorResult) -> str:
    """Graphviz DOT source of a factor tree: split off the smallest prime each level."""
    primes = []
    for p, e in sorted(result.factors.items()):
        primes += [p] * e
    lines = [
        "digraph G {",
        'graph [bgcolor="transparent", nodesep=0.35, ranksep=0.45];',
        'node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=11, '
        'color="#1F4E8C", fillcolor="#E8F0FB", fontcolor="#0B2545"];',
        'edge [color="#6B7A90", arrowsize=0.6];',
    ]
    extras = list(result.unfactored)
    current = n
    idx = 0
    lines.append(f'n{idx} [label="{_short(current)}", fillcolor="#1F4E8C", fontcolor="white"];')
    node_id = 0
    items = primes + extras
    for i, p in enumerate(items):
        if i == len(items) - 1:
            break
        pid = f"p{i}"
        is_comp = p in extras
        fill = "#FDECEC" if is_comp else "#D6F0E0"
        border = "#B23A3A" if is_comp else "#1E7B4A"
        lines.append(f'{pid} [label="{_short(p)}", fillcolor="{fill}", color="{border}"];')
        lines.append(f"n{node_id} -> {pid};")
        current //= p
        new_id = node_id + 1
        last = i == len(items) - 2
        if last:
            leaf = items[-1]
            is_comp = leaf in extras
            fill = "#FDECEC" if is_comp else "#D6F0E0"
            border = "#B23A3A" if is_comp else "#1E7B4A"
            lines.append(f'n{new_id} [label="{_short(current)}", fillcolor="{fill}", color="{border}"];')
        else:
            lines.append(f'n{new_id} [label="{_short(current)}"];')
        lines.append(f"n{node_id} -> n{new_id};")
        node_id = new_id
    if len(items) <= 1:
        lines[4] = f'n0 [label="{_short(n)}  (prime)", fillcolor="#D6F0E0", color="#1E7B4A"];'
    lines.append("}")
    return "\n".join(lines)


def _short(v: int, limit: int = 28) -> str:
    s = str(v)
    if len(s) <= limit:
        return s
    return f"{s[:10]}...{s[-8:]} ({len(s)} digits)"


# ---------------------------------------------------------------------------
# Arithmetic functions derived from the factorization
# ---------------------------------------------------------------------------
def divisor_count(factors: Dict[int, int]) -> int:
    out = 1
    for e in factors.values():
        out *= e + 1
    return out


def divisor_sum(factors: Dict[int, int]) -> int:
    out = 1
    for p, e in factors.items():
        out *= (p ** (e + 1) - 1) // (p - 1)
    return out


def euler_phi(factors: Dict[int, int]) -> int:
    out = 1
    for p, e in factors.items():
        out *= (p - 1) * p ** (e - 1)
    return out


def random_semiprime(bits: int, rng: random.Random) -> Tuple[int, int, int]:
    """Semiprime n = p*q with p and q each about bits/2 bits."""
    half = max(2, bits // 2)
    while True:
        p = generate_prime(half, rng)
        q = generate_prime(bits - half, rng)
        if p != q:
            return p * q, min(p, q), max(p, q)


def parse_int(text: str) -> int:
    """Parse decimal, hex (0x prefix) or number with separators into an int."""
    s = str(text).strip().replace(",", "").replace("_", "").replace(" ", "")
    if not s:
        raise ValueError("Empty value")
    if s.lower().startswith("0x"):
        return int(s, 16)
    if "e" in s.lower() or "." in s:
        f = float(s)
        if f != int(f):
            raise ValueError(f"{text} is not a whole number")
        return int(f)
    return int(s)


# ============================================================================
# MODULE: core.crypto
# RSA encryption, RSA digital signatures and an RSA key transport banking
# session, implemented from first principles for teaching and analysis.
#
# Note: this is textbook RSA (no OAEP or PSS padding) so every step can be
# shown. Production systems must use a vetted library with padding.
# ============================================================================

# ---------------------------------------------------------------------------
# RSA keys
# ---------------------------------------------------------------------------
@dataclass
class RSAKey:
    p: int
    q: int
    n: int
    e: int
    d: int
    phi: int
    dp: int
    dq: int
    qinv: int

    @property
    def bits(self) -> int:
        return self.n.bit_length()

    def public(self) -> Tuple[int, int]:
        return self.n, self.e

    def as_rows(self) -> List[Dict[str, str]]:
        return [
            {"Field": "p (secret prime 1)", "Value": str(self.p), "Bits": self.p.bit_length()},
            {"Field": "q (secret prime 2)", "Value": str(self.q), "Bits": self.q.bit_length()},
            {"Field": "n = p x q (public modulus)", "Value": str(self.n), "Bits": self.n.bit_length()},
            {"Field": "phi(n) = (p-1)(q-1) (secret)", "Value": str(self.phi), "Bits": self.phi.bit_length()},
            {"Field": "e (public exponent)", "Value": str(self.e), "Bits": self.e.bit_length()},
            {"Field": "d = e^-1 mod phi(n) (private exponent)", "Value": str(self.d), "Bits": self.d.bit_length()},
            {"Field": "dp = d mod (p-1) (CRT)", "Value": str(self.dp), "Bits": self.dp.bit_length()},
            {"Field": "dq = d mod (q-1) (CRT)", "Value": str(self.dq), "Bits": self.dq.bit_length()},
            {"Field": "qinv = q^-1 mod p (CRT)", "Value": str(self.qinv), "Bits": self.qinv.bit_length()},
        ]


def key_from_primes(p: int, q: int, e: int = 65537) -> RSAKey:
    if p == q:
        raise ValueError("p and q must be different primes")
    n = p * q
    phi = (p - 1) * (q - 1)
    if math.gcd(e, phi) != 1:
        raise ValueError(f"e = {e} is not coprime with phi(n) = {phi}; choose another e")
    d = modinv(e, phi)
    return RSAKey(p=p, q=q, n=n, e=e, d=d, phi=phi,
                  dp=d % (p - 1), dq=d % (q - 1), qinv=modinv(q, p))


def generate_rsa_key(bits: int, e: int = 65537, rng: Optional[random.Random] = None) -> RSAKey:
    """Generate an RSA key whose modulus has exactly `bits` bits."""
    if bits < 16:
        raise ValueError("Use at least 16 bits")
    rng = rng or random.SystemRandom()
    while True:
        p = generate_prime(bits // 2, rng)
        q = generate_prime(bits - bits // 2, rng)
        if p == q or (p * q).bit_length() != bits:
            continue
        if math.gcd(e, (p - 1) * (q - 1)) != 1:
            continue
        return key_from_primes(p, q, e)


def encrypt_int(m: int, n: int, e: int) -> int:
    if not 0 <= m < n:
        raise ValueError("Message block must satisfy 0 <= m < n")
    return pow(m, e, n)


def decrypt_int(c: int, key: RSAKey) -> int:
    return pow(c, key.d, key.n)


def decrypt_crt(c: int, key: RSAKey) -> Dict[str, int]:
    """Chinese Remainder Theorem decryption with the intermediate values."""
    m1 = pow(c, key.dp, key.p)
    m2 = pow(c, key.dq, key.q)
    h = (key.qinv * (m1 - m2)) % key.p
    m = m2 + h * key.q
    return {"m1": m1, "m2": m2, "h": h, "m": m}


# ---------------------------------------------------------------------------
# Text to blocks
# ---------------------------------------------------------------------------
def block_size_bytes(n: int) -> int:
    return (n.bit_length() - 1) // 8


def text_to_blocks(text: str, n: int) -> List[int]:
    """UTF-8 encode, pad (ISO 7816-4: 0x80 then zeros) and split into integers below n."""
    k = block_size_bytes(n)
    if k < 1:
        raise ValueError("Modulus too small: n must have at least 9 bits to hold one byte")
    data = text.encode("utf-8") + b"\x80"
    if len(data) % k:
        data += b"\x00" * (k - len(data) % k)
    return [int.from_bytes(data[i:i + k], "big") for i in range(0, len(data), k)]


def blocks_to_text(blocks: List[int], n: int) -> str:
    k = block_size_bytes(n)
    data = b"".join(b.to_bytes(k, "big") for b in blocks).rstrip(b"\x00")
    if data.endswith(b"\x80"):
        data = data[:-1]
    return data.decode("utf-8", errors="replace")


def encrypt_text(text: str, n: int, e: int) -> List[Dict[str, int]]:
    rows = []
    for i, m in enumerate(text_to_blocks(text, n)):
        rows.append({"block": i + 1, "m": m, "c": pow(m, e, n)})
    return rows


def break_rsa_by_factoring(n: int, e: int, ciphertexts: List[int], time_limit: float = 20.0):
    """Attacker view: factor n, rebuild d, decrypt. Returns dict with results."""
    res = factorize(n, time_limit=time_limit)
    primes = [p for p, k in res.factors.items() for _ in range(k)]
    if not res.complete or len(primes) != 2:
        return {"success": False, "factor_result": res}
    p, q = primes
    key = key_from_primes(p, q, e)
    plain = [pow(c, key.d, n) for c in ciphertexts]
    return {"success": True, "factor_result": res, "key": key, "plain_blocks": plain}


# ---------------------------------------------------------------------------
# Digital signatures (hash then sign, textbook RSA)
# ---------------------------------------------------------------------------
def sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def hash_to_int(message: bytes, n: int) -> int:
    """H(M) = SHA-256(M) interpreted as an integer, reduced mod n."""
    return int.from_bytes(hashlib.sha256(message).digest(), "big") % n


def sign(message: bytes, key: RSAKey) -> int:
    return pow(hash_to_int(message, key.n), key.d, key.n)


def verify(message: bytes, signature: int, n: int, e: int) -> Tuple[bool, int, int]:
    """Returns (valid, recovered hash, expected hash)."""
    h = hash_to_int(message, n)
    recovered = pow(signature, e, n)
    return recovered == h, recovered, h


# ---------------------------------------------------------------------------
# Banking session: RSA key transport + HMAC protected transactions
# ---------------------------------------------------------------------------
def derive_session_key(pre_master: int, client_nonce: bytes, server_nonce: bytes) -> bytes:
    pm = pre_master.to_bytes((pre_master.bit_length() + 7) // 8 or 1, "big")
    return hashlib.sha256(pm + client_nonce + server_nonce).digest()


def mac(session_key: bytes, message: str) -> str:
    return hmac.new(session_key, message.encode("utf-8"), hashlib.sha256).hexdigest()


def transaction_message(row: Dict) -> str:
    """Canonical message string that is authenticated for a transaction."""
    return "|".join(str(row[k]) for k in
                    ("transaction_id", "card_last4", "merchant", "amount", "currency", "timestamp"))


def banking_handshake(bank_key: RSAKey, rng: random.Random) -> Dict:
    pre_master_bits = min(128, bank_key.n.bit_length() - 2)
    pre_master = rng.getrandbits(pre_master_bits) | 1
    client_nonce = rng.getrandbits(128).to_bytes(16, "big")
    server_nonce = rng.getrandbits(128).to_bytes(16, "big")
    encrypted = pow(pre_master, bank_key.e, bank_key.n)
    recovered = pow(encrypted, bank_key.d, bank_key.n)
    k_client = derive_session_key(pre_master, client_nonce, server_nonce)
    k_bank = derive_session_key(recovered, client_nonce, server_nonce)
    return {
        "pre_master": pre_master,
        "client_nonce": client_nonce.hex(),
        "server_nonce": server_nonce.hex(),
        "encrypted_pre_master": encrypted,
        "recovered_pre_master": recovered,
        "client_session_key": k_client.hex(),
        "bank_session_key": k_bank.hex(),
        "keys_match": k_client == k_bank,
        "session_key_bytes": k_client,
    }


# ============================================================================
# MODULE: core.quantum
# Classical simulation of Shor's algorithm and quantum threat modelling.
#
# The period finding step is simulated exactly: the measurement probability
# distribution after the quantum Fourier transform is computed with FFTs,
# then continued fractions recover the period r, and gcd(a^(r/2) +- 1, N)
# gives the factors. Only practical for small N (the state vector grows
# like N^2), which is exactly why a real quantum computer is needed.
# ============================================================================

def multiplicative_order(a: int, n: int, limit: int = 5_000_000) -> Optional[int]:
    """Smallest r > 0 with a^r = 1 (mod n), found classically (slow on purpose)."""
    if math.gcd(a, n) != 1:
        return None
    x = a % n
    for r in range(1, limit + 1):
        if x == 1:
            return r
        x = x * a % n
    return None


def register_size(n: int) -> int:
    """Number of qubits q in the input register, with n^2 <= 2^q < 2 n^2."""
    return max(1, math.ceil(math.log2(n * n)))


def period_sequence(a: int, n: int, length: int) -> List[int]:
    out, x = [], 1
    for _ in range(length):
        out.append(x)
        x = x * a % n
    return out


def qft_distribution(n: int, a: int, q_bits: Optional[int] = None) -> Dict:
    """Exact probability of each measurement outcome y after the QFT.

    P(y) = (1/Q^2) * sum over values v of | sum_{x: a^x mod n = v} exp(2 pi i x y / Q) |^2
    """
    q_bits = q_bits or register_size(n)
    Q = 1 << q_bits
    if Q > 1 << 17:
        raise ValueError("Simulation limited to 17 qubits (N up to about 360)")
    f = np.empty(Q, dtype=np.int64)
    x = 1
    for i in range(Q):
        f[i] = x
        x = x * a % n
    probs = np.zeros(Q)
    for v in np.unique(f):
        indicator = (f == v).astype(np.complex128)
        amp = np.fft.fft(indicator)  # sign convention does not change |.|^2
        probs += np.abs(amp) ** 2
    probs /= Q * Q
    return {"Q": Q, "q_bits": q_bits, "probabilities": probs}


def continued_fraction_period(y: int, Q: int, n: int) -> List[Dict]:
    """Convergents of y/Q; denominators below n are period candidates."""
    rows = []
    if y == 0:
        return rows
    frac = Fraction(y, Q)
    a_terms = []
    num, den = frac.numerator, frac.denominator
    while den:
        a_terms.append(num // den)
        num, den = den, num - (num // den) * den
    h_prev, h = 1, a_terms[0]
    k_prev, k = 0, 1
    rows.append({"Convergent": f"{h}/{k}", "Denominator (candidate r)": k})
    for t in a_terms[1:]:
        h_prev, h = h, t * h + h_prev
        k_prev, k = k, t * k + k_prev
        if k >= n:
            break
        rows.append({"Convergent": f"{h}/{k}", "Denominator (candidate r)": k})
    return rows


def shor_factor(n: int, a: Optional[int] = None, rng: Optional[random.Random] = None,
                max_attempts: int = 20) -> Dict:
    """Run the classical skeleton of Shor's algorithm with a simulated period oracle."""
    rng = rng or random.Random(n)
    log: List[Dict[str, str]] = []
    if n < 4:
        return {"success": False, "log": [{"Step": "Input", "Detail": "N must be at least 4"}]}
    if n % 2 == 0:
        log.append({"Step": "Even check", "Detail": f"N is even, factor 2 found"})
        return {"success": True, "factors": (2, n // 2), "log": log, "a": None, "r": None}
    if is_probable_prime(n):
        log.append({"Step": "Primality", "Detail": f"{n} is prime, nothing to factor"})
        return {"success": False, "log": log}
    pp = is_perfect_power(n)
    if pp:
        log.append({"Step": "Perfect power", "Detail": f"N = {pp[0]}^{pp[1]}"})
        return {"success": True, "factors": (pp[0], n // pp[0]), "log": log, "a": None, "r": None}

    tried = set()
    for attempt in range(1, max_attempts + 1):
        if a is None or attempt > 1:
            choices = [c for c in range(2, n - 1) if c not in tried]
            if not choices:
                break
            cand = rng.choice(choices)
        else:
            cand = a
        tried.add(cand)
        g = math.gcd(cand, n)
        if g > 1:
            log.append({"Step": f"Attempt {attempt}: a = {cand}",
                        "Detail": f"gcd(a, N) = {g}, lucky factor found without the quantum step"})
            return {"success": True, "factors": (g, n // g), "log": log, "a": cand, "r": None}
        r = multiplicative_order(cand, n)
        log.append({"Step": f"Attempt {attempt}: a = {cand}",
                    "Detail": f"Quantum period finding gives r = {r}"})
        if r is None or r % 2:
            log.append({"Step": "Check", "Detail": f"r = {r} is odd, choose another a"})
            continue
        half = pow(cand, r // 2, n)
        if half == n - 1:
            log.append({"Step": "Check", "Detail": f"a^(r/2) = -1 mod N, choose another a"})
            continue
        f1, f2 = math.gcd(half - 1, n), math.gcd(half + 1, n)
        log.append({"Step": "Classical post processing",
                    "Detail": f"a^(r/2) mod N = {half}; gcd({half}-1, N) = {f1}; gcd({half}+1, N) = {f2}"})
        for f in (f1, f2):
            if 1 < f < n:
                return {"success": True, "factors": (f, n // f), "log": log, "a": cand, "r": r,
                        "half": half}
    return {"success": False, "log": log}


# ---------------------------------------------------------------------------
# Threat model curves
# ---------------------------------------------------------------------------
def gnfs_log10_ops(bits: int) -> float:
    """log10 of L_n[1/3, (64/9)^(1/3)], the heuristic GNFS running time."""
    ln_n = bits * math.log(2)
    c = (64 / 9) ** (1 / 3)
    return c * ln_n ** (1 / 3) * math.log(ln_n) ** (2 / 3) / math.log(10)


def shor_log10_ops(bits: int) -> float:
    """log10 of an order of magnitude gate count, about bits^3 (schoolbook arithmetic)."""
    return 3 * math.log10(bits) + math.log10(4)


def shor_logical_qubits(bits: int) -> int:
    """Beauregard style circuit: about 2n + 3 logical qubits."""
    return 2 * bits + 3


# ============================================================================
# MODULE: core.randomness
# Blum Blum Shub pseudo random generator, Very Smooth Hash (VSH) and
# NIST SP 800-22 statistical tests.
#
# Both BBS and VSH are provably secure only while factoring n = p q stays
# hard: anyone who factors n can predict BBS output and forge VSH collisions.
# ============================================================================

# ---------------------------------------------------------------------------
# Blum Blum Shub
# ---------------------------------------------------------------------------
def bbs_parameters(bits: int, rng: random.Random) -> Dict[str, int]:
    """Blum integer n = p q with p = q = 3 (mod 4)."""
    while True:
        p = generate_prime(bits // 2, rng, congruent_3_mod_4=True)
        q = generate_prime(bits - bits // 2, rng, congruent_3_mod_4=True)
        if p != q:
            return {"p": p, "q": q, "n": p * q}


def validate_bbs(p: int, q: int, seed: int) -> List[str]:
    issues = []
    if p % 4 != 3:
        issues.append(f"p = {p} is not congruent to 3 mod 4")
    if q % 4 != 3:
        issues.append(f"q = {q} is not congruent to 3 mod 4")
    if p == q:
        issues.append("p and q must differ")
    n = p * q
    if math.gcd(seed, n) != 1:
        issues.append("seed must be coprime with n")
    if seed in (0, 1):
        issues.append("seed must not be 0 or 1")
    return issues


def bbs_generate(n: int, seed: int, count: int) -> Dict[str, List[int]]:
    """x0 = seed^2 mod n; x_{i+1} = x_i^2 mod n; output bit b_i = x_i mod 2 for i >= 1."""
    x = seed * seed % n
    states = [x]
    bits = []
    for _ in range(count):
        x = x * x % n
        states.append(x)
        bits.append(x & 1)
    return {"states": states, "bits": bits}


def bbs_cycle_length(n: int, seed: int, limit: int = 200_000) -> Optional[int]:
    """Length of the cycle the state sequence eventually enters (small n only)."""
    seen = {}
    x = seed * seed % n
    for i in range(limit):
        if x in seen:
            return i - seen[x]
        seen[x] = i
        x = x * x % n
    return None


def lcg_bits(seed: int, count: int, a: int = 5, c: int = 3, m: int = 16) -> List[int]:
    """A deliberately weak linear congruential generator for comparison."""
    x, out = seed % m, []
    for _ in range(count):
        x = (a * x + c) % m
        out.append(x & 1)
    return out


def bits_to_bytes(bits: List[int]) -> bytes:
    out = bytearray()
    for i in range(0, len(bits) - len(bits) % 8, 8):
        v = 0
        for b in bits[i:i + 8]:
            v = (v << 1) | b
        out.append(v)
    return bytes(out)


# ---------------------------------------------------------------------------
# NIST SP 800-22 tests
# ---------------------------------------------------------------------------
def monobit_test(bits: List[int]) -> Dict[str, float]:
    """Frequency (monobit) test. S = sum(2b - 1); s_obs = |S| / sqrt(n); p = erfc(s_obs / sqrt 2)."""
    n = len(bits)
    s = sum(2 * b - 1 for b in bits)
    s_obs = abs(s) / math.sqrt(n)
    p = math.erfc(s_obs / math.sqrt(2))
    return {"n": n, "S": s, "s_obs": s_obs, "p_value": p, "pass": p >= 0.01}


def runs_test(bits: List[int]) -> Dict[str, float]:
    """Runs test. V = number of runs; p = erfc(|V - 2 n pi (1-pi)| / (2 sqrt(2n) pi (1-pi)))."""
    n = len(bits)
    pi = sum(bits) / n
    tau = 2 / math.sqrt(n)
    if abs(pi - 0.5) >= tau:
        return {"n": n, "pi": pi, "V": float("nan"), "p_value": 0.0, "pass": False,
                "note": "Prerequisite frequency test failed"}
    v = 1 + sum(1 for i in range(n - 1) if bits[i] != bits[i + 1])
    num = abs(v - 2 * n * pi * (1 - pi))
    den = 2 * math.sqrt(2 * n) * pi * (1 - pi)
    p = math.erfc(num / den)
    return {"n": n, "pi": pi, "V": v, "p_value": p, "pass": p >= 0.01, "note": ""}


def block_frequency_test(bits: List[int], m: int = 128) -> Dict[str, float]:
    """Block frequency test: chi^2 = 4 M sum (pi_i - 1/2)^2; p = igamc(N/2, chi^2/2)."""
    n_blocks = len(bits) // m
    if n_blocks == 0:
        return {"blocks": 0, "chi2": float("nan"), "p_value": float("nan"), "pass": False}
    chi2 = 0.0
    for i in range(n_blocks):
        pi_i = sum(bits[i * m:(i + 1) * m]) / m
        chi2 += (pi_i - 0.5) ** 2
    chi2 *= 4 * m
    p = _igamc(n_blocks / 2, chi2 / 2)
    return {"blocks": n_blocks, "chi2": chi2, "p_value": p, "pass": p >= 0.01}


def _igamc(a: float, x: float) -> float:
    """Regularized upper incomplete gamma Q(a, x)."""
    try:
        from scipy.special import gammaincc  # type: ignore
        return float(gammaincc(a, x))
    except Exception:  # pragma: no cover - fallback series / continued fraction
        if x < 0 or a <= 0:
            return float("nan")
        if x < a + 1:
            term = total = 1.0 / a
            ap = a
            for _ in range(1000):
                ap += 1
                term *= x / ap
                total += term
                if abs(term) < abs(total) * 1e-15:
                    break
            return 1 - total * math.exp(-x + a * math.log(x) - math.lgamma(a))
        b = x + 1 - a
        c = 1e300
        d = 1 / b
        h = d
        for i in range(1, 1000):
            an = -i * (i - a)
            b += 2
            d = an * d + b
            d = 1e-300 if abs(d) < 1e-300 else d
            c = b + an / c
            c = 1e-300 if abs(c) < 1e-300 else c
            d = 1 / d
            delta = d * c
            h *= delta
            if abs(delta - 1) < 1e-15:
                break
        return math.exp(-x + a * math.log(x) - math.lgamma(a)) * h


# ---------------------------------------------------------------------------
# Very Smooth Hash (Contini, Lenstra, Steinfeld 2006), simplified
# ---------------------------------------------------------------------------
def vsh_block_primes(n: int) -> List[int]:
    """Largest k such that the product of the first k primes is below n."""
    primes = sieve(10_000)
    prod, k = 1, 0
    for p in primes:
        if prod * p >= n:
            break
        prod *= p
        k += 1
    return primes[:k]


def vsh_hash(message: bytes, n: int) -> Dict:
    """x_0 = 1; x_{j+1} = x_j^2 * prod_i p_i^{m_ij} mod n over k-bit message blocks.

    The final block encodes the message bit length, as in the VSH paper.
    """
    primes = vsh_block_primes(n)
    k = len(primes)
    if k < 2:
        raise ValueError("Modulus too small for VSH")
    bits = []
    for byte in message:
        bits.extend((byte >> (7 - i)) & 1 for i in range(8))
    length = len(bits)
    if length % k:
        bits.extend([0] * (k - length % k))
    blocks = [bits[i:i + k] for i in range(0, len(bits), k)]
    len_bits = [(length >> i) & 1 for i in range(k)]
    blocks.append(len_bits)
    x = 1
    trace = []
    for j, block in enumerate(blocks):
        mult = 1
        for p, b in zip(primes, block):
            if b:
                mult *= p
        x = x * x * mult % n
        trace.append(x)
    return {"digest": x, "k": k, "blocks": len(blocks), "primes": primes, "trace": trace}


def hamming_distance(a: int, b: int) -> int:
    return bin(a ^ b).count("1")


# ============================================================================
# MODULE: core.security
# RSA key auditing: the checks researchers use to find weak keys in the wild.
#
#   * Batch GCD (Bernstein product and remainder trees) finds moduli that
#     share a prime factor, the flaw behind the 2012 "Mining your Ps and Qs"
#     and "Ron was wrong, Whit is right" studies.
#   * Fermat factorization finds keys whose primes are too close together.
#   * Pollard p minus 1 finds keys whose p minus 1 is smooth.
#   * Trial division finds keys with a small factor.
#   * Policy checks for modulus size and public exponent.
# ============================================================================

def product_tree(values: List[int]) -> List[List[int]]:
    tree = [values]
    while len(tree[-1]) > 1:
        level = tree[-1]
        tree.append([level[i] * level[i + 1] if i + 1 < len(level) else level[i]
                     for i in range(0, len(level), 2)])
    return tree


def batch_gcd(moduli: List[int]) -> List[int]:
    """For each N_i return gcd(N_i, product of all other N_j), in quasi linear time."""
    if len(moduli) < 2:
        return [1] * len(moduli)
    tree = product_tree(moduli)
    rems = tree.pop()
    while tree:
        level = tree.pop()
        rems = [rems[i // 2] % (level[i] ** 2) for i in range(len(level))]
    return [math.gcd(r // n, n) for r, n in zip(rems, moduli)]


def naive_pairwise_gcd(moduli: List[int]) -> List[int]:
    """Reference O(k^2) implementation used for testing batch_gcd."""
    out = []
    for i, a in enumerate(moduli):
        g = 1
        for j, b in enumerate(moduli):
            if i != j:
                g = g * math.gcd(a, b) // math.gcd(g, math.gcd(a, b))
        out.append(math.gcd(g, a))
    return out


def audit_keys(records: List[Dict], min_bits: int = 2048, fermat_iterations: int = 20_000,
               pm1_bound: int = 20_000, trial_limit: int = 10_000) -> List[Dict]:
    """Audit RSA public keys. records: dicts with device_id, modulus, exponent."""
    moduli = [int(r["modulus"]) for r in records]
    shared = batch_gcd(moduli)
    small_primes = [p for p in SMALL_PRIMES if p <= trial_limit]
    results = []
    for rec, n, g in zip(records, moduli, shared):
        e = int(rec.get("exponent", 65537))
        findings: List[str] = []
        factor_found: Optional[int] = None
        method = ""
        fermat_iters = None

        if n % 2 == 0:
            findings.append("Even modulus")
            factor_found, method = 2, "Even modulus"
        if factor_found is None:
            for p in small_primes:
                if n % p == 0 and n != p:
                    findings.append(f"Small prime factor {p}")
                    factor_found, method = p, "Trial division"
                    break
        if 1 < g < n:
            findings.append("Shares a prime with another key")
            if factor_found is None:
                factor_found, method = g, "Batch GCD"
        elif g == n and n > 1:
            findings.append("Duplicate modulus (identical key reused)")
        if factor_found is None:
            fr = fermat_factor(n, max_iterations=fermat_iterations)
            if fr:
                findings.append(f"Primes too close together (Fermat, {fr[2]} iterations)")
                factor_found, method, fermat_iters = fr[0], "Fermat", fr[2]
        if factor_found is None:
            f = pollard_p_minus_1(n, bound=pm1_bound)
            if f:
                findings.append(f"p minus 1 is smooth (Pollard p minus 1, bound {pm1_bound})")
                factor_found, method = f, "Pollard p minus 1"
        if n.bit_length() < min_bits:
            findings.append(f"Modulus {n.bit_length()} bits is below policy {min_bits} bits")
        if e < 65537:
            findings.append(f"Low public exponent e = {e}")
        if is_probable_prime(n):
            findings.append("Modulus is prime (not a valid RSA modulus)")
            factor_found, method = None, "Invalid"

        factored = factor_found is not None
        if factored:
            risk = "Critical"
        elif any(f.startswith("Low public exponent") or f.startswith("Duplicate") for f in findings):
            risk = "High"
        elif any("below policy" in f for f in findings):
            risk = "Medium"
        else:
            risk = "Low"
        results.append({
            "device_id": rec.get("device_id", ""),
            "bits": n.bit_length(),
            "exponent": e,
            "risk": risk,
            "factored": factored,
            "method": method or ("None" if not factored else ""),
            "factor_p": str(factor_found) if factored else "",
            "factor_q": str(n // factor_found) if factored else "",
            "fermat_iterations": fermat_iters,
            "shared_gcd": str(g) if 1 < g < n else "",
            "findings": "; ".join(findings) if findings else "No weakness detected",
            "modulus": str(n),
        })
    return results


def _smooth_prime(bits: int, rng: random.Random, bound: int = 5000) -> int:
    """Prime p of the given size where p - 1 has only prime factors below bound."""
    small = [p for p in sieve(bound) if p > 2]
    while True:
        pool = small[:]
        rng.shuffle(pool)
        m = 2
        while m.bit_length() < bits - 1:
            m *= pool.pop()  # distinct primes, so every prime power stays below bound
        p = m + 1
        if p.bit_length() == bits and (p >> (bits - 2)) == 3 and is_probable_prime(p):
            return p


def _prime_top2(bits: int, rng: random.Random) -> int:
    """Prime with its two top bits set, so a product of two has exactly 2*bits bits."""
    while True:
        c = rng.getrandbits(bits) | (3 << (bits - 2)) | 1
        if is_probable_prime(c):
            return c


def synthetic_device_keys(rng: random.Random, bits: int = 256) -> List[Dict]:
    """Synthetic fleet of IoT devices, routers and servers with planted weaknesses."""
    half = bits // 2

    def generate_prime(b, r):  # noqa: F811 - local override for exact modulus size
        return _prime_top2(b, r)
    rows: List[Dict] = []

    def add(dev, kind, p, q, e=65537, planted="None (strong key)"):
        rows.append({"device_id": dev, "device_type": kind, "modulus": p * q,
                     "exponent": e, "planted_weakness": planted})

    kinds = ["Router", "VPN gateway", "IP camera", "Web server", "Smart meter",
             "Firewall", "Printer", "NAS storage", "Mail server", "Payment terminal"]
    # Strong keys
    for i in range(10):
        add(f"DEV-{100 + i}", kinds[i % len(kinds)], generate_prime(half, rng), generate_prime(half, rng))
    # Shared primes (low entropy at boot)
    shared1 = generate_prime(half, rng)
    shared2 = generate_prime(half, rng)
    add("DEV-201", "Router", shared1, generate_prime(half, rng), planted="Shared prime (group A)")
    add("DEV-202", "Router", shared1, generate_prime(half, rng), planted="Shared prime (group A)")
    add("DEV-203", "Router", shared1, generate_prime(half, rng), planted="Shared prime (group A)")
    add("DEV-204", "IP camera", shared2, generate_prime(half, rng), planted="Shared prime (group B)")
    add("DEV-205", "IP camera", shared2, generate_prime(half, rng), planted="Shared prime (group B)")
    # Close primes
    for dev in ("DEV-301", "DEV-302"):
        p = generate_prime(half, rng)
        q = p + 2 + rng.randrange(0, 2000) * 2
        while not is_probable_prime(q):
            q += 2
        add(dev, "Smart meter", p, q, planted="Primes too close (Fermat)")
    # Smooth p - 1
    add("DEV-401", "Printer", _smooth_prime(half, rng), generate_prime(half, rng),
        planted="Smooth p minus 1")
    # Small factor
    add("DEV-501", "Legacy terminal", 7919, generate_prime(bits - 13, rng), planted="Small prime factor")
    # Low exponent
    add("DEV-601", "Web server", generate_prime(half, rng), generate_prime(half, rng), e=3,
        planted="Low public exponent e = 3")
    order = list(range(len(rows)))
    rng.shuffle(order)
    return [rows[i] for i in order]


# ============================================================================
# MODULE: core.validation
# Accuracy validation suite. Each check compares the engine against a
# published reference value or an independent implementation.
# Used both by the automated tests and by the in app validation page.
# ============================================================================

def _check(rows: List[Dict], category: str, name: str, expected, actual, reference: str = "") -> None:
    rows.append({"Category": category, "Check": name, "Expected": str(expected)[:80],
                 "Actual": str(actual)[:80], "Result": "Pass" if expected == actual else "Fail",
                 "Reference": reference})


def run_all(quick: bool = False) -> List[Dict]:
    rows: List[Dict] = []
    rng = random.Random(20260930)

    # --- Factorization -----------------------------------------------------
    known = {
        360: {2: 3, 3: 2, 5: 1},
        561: {3: 1, 11: 1, 17: 1},
        600851475143: {71: 1, 839: 1, 1471: 1, 6857: 1},
        2 ** 64 + 1: {274177: 1, 67280421310721: 1},
        2 ** 67 - 1: {193707721: 1, 761838257287: 1},
        2 ** 32 + 1: {641: 1, 6700417: 1},
        10 ** 18: {2: 18, 5: 18},
        3 ** 40: {3: 40},
    }
    refs = {2 ** 64 + 1: "Landry (1880), sixth Fermat number", 2 ** 67 - 1: "Cole (1903)",
            2 ** 32 + 1: "Euler (1732), fifth Fermat number", 561: "Smallest Carmichael number",
            600851475143: "Project Euler problem 3"}
    for n, f in known.items():
        _check(rows, "Factorization", f"factor({n})", f, factorize(n).factors, refs.get(n, "Direct computation"))

    try:
        import sympy
        count = 60 if quick else 300
        mism = 0
        for _ in range(count):
            n = rng.randrange(2, 10 ** rng.randint(3, 18))
            if factorize(n).factors != {int(k): v for k, v in sympy.factorint(n).items()}:
                mism += 1
        _check(rows, "Factorization", f"{count} random n up to 10^18 versus SymPy factorint", 0, mism,
               "SymPy (independent implementation)")
        mism = 0
        for n in range(1, 20001 if not quick else 5001):
            if is_probable_prime(n) != sympy.isprime(n):
                mism += 1
        _check(rows, "Primality", "Miller Rabin versus SymPy for every n up to 20000" if not quick else
               "Miller Rabin versus SymPy for every n up to 5000", 0, mism, "SymPy isprime")
    except ImportError:
        pass
    for n, truth, ref in ((3215031751, False, "Strong pseudoprime to bases 2, 3, 5, 7"),
                          (3825123056546413051, False, "Strong pseudoprime to bases 2 to 23"),
                          (561, False, "Carmichael number"), (2 ** 89 - 1, True, "Mersenne prime M89"),
                          (2 ** 127 - 1, True, "Mersenne prime M127")):
        _check(rows, "Primality", f"is_prime({n})", truth, is_probable_prime(n), ref)

    for bits in (48, 64) if quick else (48, 64, 72):
        n, p, q = random_semiprime(bits, rng)
        r = factorize(n, time_limit=60)
        _check(rows, "Factorization", f"{bits} bit semiprime recovers p and q", {p: 1, q: 1}, r.factors, "Constructed")

    # --- GCD / LCM -----------------------------------------------------------
    bad = 0
    for _ in range(2000):
        a, b = rng.randrange(1, 10 ** 12), rng.randrange(1, 10 ** 12)
        if gcd(a, b) != math.gcd(a, b) or lcm(a, b) != math.lcm(a, b):
            bad += 1
        g, x, y = egcd(a, b)
        if a * x + b * y != g:
            bad += 1
    _check(rows, "GCD and LCM", "2000 random pairs: gcd, lcm and Bezout identity", 0, bad, "Python math module")
    _check(rows, "GCD and LCM", "Euclid gcd(1071, 462) final divisor", 21, euclid_steps(1071, 462)[-1]["b"],
           "Euclid's Elements example")
    _check(rows, "GCD and LCM", "lcm(12, 18, 30, 45)", 180, lcm_list([12, 18, 30, 45]), "Direct computation")

    # --- RSA -----------------------------------------------------------------
    k = key_from_primes(61, 53, 17)
    _check(rows, "RSA", "Textbook key p=61, q=53, e=17: d", 2753, k.d, "Wikipedia RSA example")
    _check(rows, "RSA", "Encrypt m=65", 2790, pow(65, k.e, k.n), "Wikipedia RSA example")
    _check(rows, "RSA", "Decrypt c=2790", 65, decrypt_int(2790, k), "Wikipedia RSA example")
    _check(rows, "RSA", "CRT decrypt c=2790", 65, decrypt_crt(2790, k)["m"], "Wikipedia RSA example")
    ok = 0
    trials = 3 if quick else 6
    for i in range(trials):
        key = generate_rsa_key(1024 if i % 2 else 512, rng=rng)
        msg = f"Test message {i} with unicode text: café, naïve, 東京"
        blocks = text_to_blocks(msg, key.n)
        cts = [pow(m, key.e, key.n) for m in blocks]
        back = blocks_to_text([decrypt_crt(c, key)["m"] for c in cts], key.n)
        ok += back == msg and all(decrypt_crt(c, key)["m"] == decrypt_int(c, key) for c in cts)
    _check(rows, "RSA", f"{trials} random 512 and 1024 bit keys: text round trip and CRT agreement", trials, ok, "Self consistency")
    key = generate_rsa_key(64, rng=random.Random(1))
    att = break_rsa_by_factoring(key.n, key.e, [pow(42, key.e, key.n)], time_limit=30)
    _check(rows, "RSA", "Factoring attack on 64 bit key recovers d", key.d, att["key"].d if att["success"] else None,
           "Attack simulation")

    # --- Signatures -------------------------------------------------------------
    key = generate_rsa_key(512, rng=rng)
    s = sign(b"release 1.0", key)
    _check(rows, "Signatures", "Genuine signature verifies", True, verify(b"release 1.0", s, key.n, key.e)[0], "RSA identity")
    _check(rows, "Signatures", "Tampered message rejected", False, verify(b"release 1.1", s, key.n, key.e)[0], "RSA identity")
    other = generate_rsa_key(512, rng=rng)
    _check(rows, "Signatures", "Signature from another key rejected", False,
           verify(b"release 1.0", sign(b"release 1.0", other), key.n, key.e)[0], "RSA identity")
    _check(rows, "Signatures", "SHA-256 of 'abc'",
           "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad", sha256_hex(b"abc"), "FIPS 180-4 test vector")

    # --- Banking ----------------------------------------------------------------
    hs = banking_handshake(key, random.Random(3))
    _check(rows, "Banking", "Client and bank derive the same session key", True, hs["keys_match"], "Protocol consistency")
    _check(rows, "Banking", "Recovered pre master secret equals original", hs["pre_master"], hs["recovered_pre_master"], "RSA identity")
    _check(rows, "Banking", "HMAC-SHA256 RFC 4231 test case 2",
           "5bdcc146bf60754e6a042426089575c75a003f089d2739839dec58b964ec3843",
           mac(b"Jefe", "what do ya want for nothing?"), "RFC 4231")

    # --- Quantum ---------------------------------------------------------------
    _check(rows, "Quantum", "Order of 7 mod 15", 4, multiplicative_order(7, 15), "Shor (1994) textbook example")
    res = shor_factor(15, a=7)
    _check(rows, "Quantum", "Shor factors 15 with a = 7", (3, 5), tuple(sorted(res["factors"])), "Textbook example")
    res = shor_factor(21, a=2)
    _check(rows, "Quantum", "Shor factors 21 with a = 2", (3, 7), tuple(sorted(res["factors"])), "Textbook example")
    d = qft_distribution(15, 7)
    peaks = sorted(int(i) for i in (d["probabilities"] > 1e-9).nonzero()[0])
    _check(rows, "Quantum", "QFT peaks for N=15, a=7 (Q=256)", [0, 64, 128, 192], peaks, "Exact: multiples of Q/r")
    _check(rows, "Quantum", "Probabilities sum to 1", True, abs(d["probabilities"].sum() - 1) < 1e-9, "Unitarity")
    ok = 0
    ns = [21, 33, 35, 55, 77, 91, 143, 187, 221, 247] if not quick else [21, 35, 77]
    for n in ns:
        rr = shor_factor(n, rng=random.Random(n))
        ok += rr["success"] and rr["factors"][0] * rr["factors"][1] == n and 1 < rr["factors"][0] < n
    _check(rows, "Quantum", f"Shor factors {len(ns)} semiprimes", len(ns), ok, "Self consistency")

    # --- Randomness ------------------------------------------------------------
    _check(rows, "Randomness", "BBS states p=11, q=23, s=3", [9, 81, 236, 36, 31, 202],
           bbs_generate(253, 3, 5)["states"], "Wikipedia Blum Blum Shub example")
    m = monobit_test([int(c) for c in "1011010101"])
    _check(rows, "Randomness", "NIST monobit example p value", 0.527089, round(m["p_value"], 6), "NIST SP 800-22 section 2.1.8")
    r = runs_test([int(c) for c in "1001101011"])
    _check(rows, "Randomness", "NIST runs example p value", 0.147232, round(r["p_value"], 6), "NIST SP 800-22 section 2.3.8")
    bf = block_frequency_test([int(c) for c in "0110011010"], m=3)
    _check(rows, "Randomness", "NIST block frequency example p value", 0.801252, round(bf["p_value"], 6), "NIST SP 800-22 section 2.2.8")
    prm = bbs_parameters(128, random.Random(4))
    bits = bbs_generate(prm["n"], 123457, 20000)["bits"]
    passes = sum(t(bits)["pass"] for t in (monobit_test, runs_test,
                                          lambda b: block_frequency_test(b, 128)))
    _check(rows, "Randomness", "BBS 20000 bits pass 3 NIST tests", 3, passes, "NIST SP 800-22")
    _check(rows, "Randomness", "Weak LCG fails runs test", False, runs_test(lcg_bits(5, 20000))["pass"], "Expected weakness")
    h1 = vsh_hash(b"abc", prm["n"])["digest"]
    h2 = vsh_hash(b"abc", prm["n"])["digest"]
    h3 = vsh_hash(b"abd", prm["n"])["digest"]
    _check(rows, "Randomness", "VSH deterministic and sensitive to input", True, h1 == h2 and h1 != h3, "Hash properties")

    # --- Security testing ---------------------------------------------------------
    moduli = [k["modulus"] for k in synthetic_device_keys(random.Random(8), 256)]
    _check(rows, "Security testing", "Batch GCD equals pairwise GCD", naive_pairwise_gcd(moduli),
           batch_gcd(moduli), "Bernstein product and remainder trees")
    p = generate_prime(128, rng)
    q = next_prime(p + 1000)
    fr = fermat_factor(p * q)
    _check(rows, "Security testing", "Fermat splits primes 1000 apart", (p, q), fr[:2] if fr else None, "Fermat (1643)")
    seeds = range(3) if quick else range(8)
    miss = 0
    for sd in seeds:
        fleet = synthetic_device_keys(random.Random(sd), 256)
        audit = audit_keys(fleet, min_bits=256)
        miss += sum((f["planted_weakness"].startswith("None")) != (a["risk"] == "Low") for f, a in zip(fleet, audit))
    _check(rows, "Security testing", f"All planted weaknesses detected, no false alarms ({len(seeds)} fleets)", 0, miss, "Synthetic ground truth")
    return rows


# ============================================================================
# MODULE: ui.exporters
# Report exporters: every use case can export its results as CSV, plain
# text, Word (.docx) and PDF.
# ============================================================================

APP_NAME = "PrimeShield: Prime Factorization and Cryptography Suite"
AUTHOR_LINE = "Developed by Randy Singh from Kalsnet (KNet) Consulting Group"
BLUE = "#0B4F9C"


@dataclass
class Report:
    use_case: str
    summary: List[Tuple[str, str]] = field(default_factory=list)
    sections: List[Tuple[str, str]] = field(default_factory=list)
    formulas: List[Tuple[str, str]] = field(default_factory=list)
    tables: List[Tuple[str, pd.DataFrame]] = field(default_factory=list)
    chart: Optional[Dict] = None  # {"kind": "bar"|"line", "x": [...], "y": {...}, "title", "xlabel", "ylabel"}
    generated: str = field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d %H:%M"))

    def primary_table(self) -> pd.DataFrame:
        if self.tables:
            return self.tables[0][1]
        return pd.DataFrame(self.summary, columns=["Metric", "Value"])


def _stringify(df: pd.DataFrame) -> pd.DataFrame:
    """Every cell as a Python str, with missing values as empty strings."""
    out = df.astype(object).where(df.notna(), "")
    return out.map(lambda v: v if isinstance(v, str) else str(v))


# ---------------------------------------------------------------------------
def to_csv(report: Report) -> bytes:
    return _stringify(report.primary_table()).to_csv(index=False).encode("utf-8")


def to_txt(report: Report) -> bytes:
    w = 90
    out = io.StringIO()
    out.write("=" * w + "\n")
    out.write(APP_NAME.center(w) + "\n")
    out.write(AUTHOR_LINE.center(w) + "\n")
    out.write("=" * w + "\n")
    out.write(f"Use case : {report.use_case}\nGenerated: {report.generated}\n\n")
    if report.summary:
        out.write("SUMMARY\n" + "-" * w + "\n")
        lw = max(len(k) for k, _ in report.summary) + 2
        for k, v in report.summary:
            out.write(f"{k.ljust(lw)}{v}\n")
        out.write("\n")
    for head, body in report.sections:
        out.write(head.upper() + "\n" + "-" * w + "\n")
        out.write(textwrap.fill(body, w) + "\n\n")
    if report.formulas:
        out.write("FORMULAS\n" + "-" * w + "\n")
        for name, f in report.formulas:
            out.write(f"{name}:\n    {f}\n")
        out.write("\n")
    for name, df in report.tables:
        out.write(name.upper() + "\n" + "-" * w + "\n")
        out.write(_stringify(df).to_string(index=False) + "\n\n")
    return out.getvalue().encode("utf-8")


def chart_png(chart: Dict) -> Optional[bytes]:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return None
    palette = ["#0B4F9C", "#E07A1F", "#2E8B57", "#8E44AD", "#C0392B"]
    fig, ax = plt.subplots(figsize=(8, 3.6), dpi=150)
    x = chart["x"]
    for i, (label, ys) in enumerate(chart["y"].items()):
        if chart.get("kind") == "bar":
            n = len(chart["y"])
            width = 0.8 / n
            pos = [j + (i - (n - 1) / 2) * width for j in range(len(x))]
            ax.bar(pos, ys, width=width, label=label, color=palette[i % len(palette)])
            ax.set_xticks(range(len(x)))
            ax.set_xticklabels([str(v) for v in x], rotation=30, ha="right", fontsize=8)
        else:
            ax.plot(x, ys, marker="o", markersize=3, label=label, color=palette[i % len(palette)])
    ax.set_title(chart.get("title", ""), fontsize=11, color="#0B2545")
    ax.set_xlabel(chart.get("xlabel", ""))
    ax.set_ylabel(chart.get("ylabel", ""))
    if chart.get("logy"):
        ax.set_yscale("log")
    ax.grid(alpha=0.25)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    if len(chart["y"]) > 1:
        ax.legend(frameon=False, fontsize=8)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def to_docx(report: Report, max_rows: int = 300) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Inches, Cm

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.8))
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    p = doc.add_paragraph()
    r = p.add_run(APP_NAME)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x0B, 0x4F, 0x9C)
    p = doc.add_paragraph()
    r = p.add_run(AUTHOR_LINE)
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x44, 0x55, 0x6B)
    h = doc.add_heading(f"Use Case Report: {report.use_case}", level=1)
    doc.add_paragraph(f"Generated {report.generated}")

    if report.summary:
        doc.add_heading("Summary", level=2)
        t = doc.add_table(rows=0, cols=2)
        t.style = "Light Grid Accent 1"
        for k, v in report.summary:
            row = t.add_row().cells
            row[0].text = k
            row[1].text = str(v)
            row[0].paragraphs[0].runs[0].bold = True
    for head, body in report.sections:
        doc.add_heading(head, level=2)
        doc.add_paragraph(body)
    if report.formulas:
        doc.add_heading("Formulas", level=2)
        for name, f in report.formulas:
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(name + ": ").bold = True
            fr = para.add_run(f)
            fr.font.name = "Consolas"
    if report.chart:
        png = chart_png(report.chart)
        if png:
            doc.add_heading("Chart", level=2)
            doc.add_picture(io.BytesIO(png), width=Inches(8.5))
    for name, df in report.tables:
        doc.add_heading(name, level=2)
        d = _stringify(df.head(max_rows))
        t = doc.add_table(rows=1, cols=len(d.columns))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(d.columns):
            t.rows[0].cells[i].text = str(c)
        for _, rowv in d.iterrows():
            cells = t.add_row().cells
            for i, v in enumerate(rowv):
                cells[i].text = v
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(8)
        if len(df) > max_rows:
            doc.add_paragraph(f"Showing first {max_rows} of {len(df)} rows. Full data in the CSV export.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(report: Report, max_rows: int = 300) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle, LongTable)
    from xml.sax.saxutils import escape

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.4 * cm, bottomMargin=1.4 * cm,
                            title=f"{report.use_case} report", author="Randy Singh, KNet Consulting Group")
    ss = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=ss["Title"], textColor=colors.HexColor(BLUE), fontSize=20,
                           alignment=TA_LEFT, spaceAfter=4)
    sub = ParagraphStyle("s", parent=ss["Normal"], textColor=colors.HexColor("#44556B"),
                         fontName="Helvetica-Oblique", fontSize=10, spaceAfter=10)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], textColor=colors.HexColor("#0B2545"), fontSize=15)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor(BLUE), fontSize=12)
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=9.5, leading=13)
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=7.2, leading=9, wordWrap="CJK")
    cellh = ParagraphStyle("ch", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    mono = ParagraphStyle("m", parent=body, fontName="Courier", fontSize=8.5, leading=11)

    width = landscape(A4)[0] - 3 * cm
    story = [Paragraph(escape(APP_NAME), title), Paragraph(escape(AUTHOR_LINE), sub),
             Paragraph(escape(f"Use Case Report: {report.use_case}"), h1),
             Paragraph(f"Generated {report.generated}", body), Spacer(1, 8)]

    def grid(data, header=True, col_widths=None):
        t = LongTable(data, colWidths=col_widths, repeatRows=1 if header else 0)
        st = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C5D3E8")),
              ("VALIGN", (0, 0), (-1, -1), "TOP"),
              ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FC")])]
        if header:
            st.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BLUE)))
        t.setStyle(TableStyle(st))
        return t

    if report.summary:
        story.append(Paragraph("Summary", h2))
        data = [[Paragraph(escape(k), cell), Paragraph(escape(str(v)), cell)] for k, v in report.summary]
        story.append(grid(data, header=False, col_widths=[width * 0.3, width * 0.7]))
    for head, text in report.sections:
        story += [Paragraph(escape(head), h2), Paragraph(escape(text), body)]
    if report.formulas:
        story.append(Paragraph("Formulas", h2))
        for name, f in report.formulas:
            story.append(Paragraph(f"<b>{escape(name)}</b>", body))
            story.append(Paragraph(escape(f), mono))
    if report.chart:
        png = chart_png(report.chart)
        if png:
            story += [Paragraph("Chart", h2), Image(io.BytesIO(png), width=width * 0.75, height=width * 0.75 * 0.45)]
    for name, df in report.tables:
        d = _stringify(df.head(max_rows))
        story.append(Paragraph(escape(name), h2))
        ncol = max(1, len(d.columns))
        data = [[Paragraph(escape(str(c)), cellh) for c in d.columns]]
        for _, rowv in d.iterrows():
            data.append([Paragraph(escape(v), cell) for v in rowv])
        story.append(grid(data, col_widths=[width / ncol] * ncol))
        if len(df) > max_rows:
            story.append(Paragraph(f"Showing first {max_rows} of {len(df)} rows. Full data in the CSV export.", body))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#6B7A90"))
        canvas.drawString(1.5 * cm, 0.8 * cm, f"{APP_NAME}  |  KNet Consulting Group")
        canvas.drawRightString(landscape(A4)[0] - 1.5 * cm, 0.8 * cm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ============================================================================
# MODULE: ui.common
# Shared layout, styling, upload parsing and export widgets.
# ============================================================================

PRIMARY = "#0B4F9C"
NAVY = "#0B2545"
ACCENT = "#E07A1F"
SUCCESS = "#1E7B4A"
DANGER = "#B23A3A"
MUTED = "#6B7A90"
PLOT_COLORS = ["#0B4F9C", "#E07A1F", "#2E8B57", "#8E44AD", "#C0392B", "#16A2B8", "#7F8C8D"]

UPLOAD_TYPES = ["csv", "txt", "docx", "pdf"]

CSS = f"""
<style>
#MainMenu, footer, [data-testid="stDeployButton"], [data-testid="stToolbar"] {{visibility: hidden;}}
header[data-testid="stHeader"] {{background: transparent; height: 0;}}
.block-container {{padding-top: 1.2rem; padding-bottom: 3rem; max-width: 1400px;}}
html, body, [class*="css"] {{font-family: "Segoe UI", "Helvetica Neue", Arial, sans-serif;}}

/* Title bar */
.ps-titlebar {{
  background: linear-gradient(180deg, #FFFFFF 0%, #F3F7FD 100%);
  border: 1px solid #D5E1F2; border-left: 8px solid {PRIMARY};
  border-radius: 10px; padding: 18px 26px 14px 26px; margin-bottom: 18px;
  box-shadow: 0 2px 10px rgba(11,37,69,0.06);
}}
.ps-title {{font-size: 2.45rem; font-weight: 800; color: {PRIMARY}; line-height: 1.15; margin: 0;
  letter-spacing: -0.5px;}}
.ps-subtitle {{font-size: 1.08rem; color: #33475F; margin-top: 6px; font-weight: 500;}}
.ps-subtitle b {{color: {NAVY};}}

/* Page heading */
.ps-page {{display:flex; align-items:baseline; gap:14px; margin: 4px 0 2px 0;}}
.ps-page-num {{background:{PRIMARY}; color:white; font-weight:700; border-radius:6px; padding:3px 10px;
  font-size:0.9rem;}}
.ps-page-title {{font-size:1.65rem; font-weight:700; color:{NAVY};}}
.ps-page-sub {{color:{MUTED}; font-size:1rem; margin-bottom: 10px;}}

/* Cards */
.ps-card {{background:#FFFFFF; border:1px solid #DCE5F1; border-radius:10px; padding:16px 18px;
  height:100%; box-shadow: 0 1px 4px rgba(11,37,69,0.05);}}
.ps-card h4 {{margin:0 0 6px 0; color:{PRIMARY}; font-size:1.02rem;}}
.ps-card p {{margin:0; color:#33475F; font-size:0.93rem; line-height:1.45;}}
.ps-metric {{background:#F5F8FD; border:1px solid #DCE5F1; border-top:4px solid {PRIMARY};
  border-radius:8px; padding:10px 14px;}}
.ps-metric .lbl {{color:{MUTED}; font-size:0.78rem; text-transform:uppercase; letter-spacing:0.6px;}}
.ps-metric .val {{color:{NAVY}; font-size:1.35rem; font-weight:700; word-break: break-all;}}
.ps-metric.ok {{border-top-color:{SUCCESS};}} .ps-metric.bad {{border-top-color:{DANGER};}}
.ps-metric.warn {{border-top-color:{ACCENT};}}
.ps-callout {{background:#EEF4FC; border-left:4px solid {PRIMARY}; padding:12px 16px; border-radius:6px;
  color:#22354D; margin: 8px 0 12px 0; font-size:0.95rem;}}
.ps-callout.warn {{background:#FFF5EB; border-left-color:{ACCENT};}}
.ps-callout.ok {{background:#EBF7F0; border-left-color:{SUCCESS};}}
.ps-callout.bad {{background:#FCEEEE; border-left-color:{DANGER};}}
.ps-formula-name {{font-weight:700; color:{NAVY}; margin-top:10px;}}
.ps-formula-text {{color:#33475F; font-size:0.92rem; margin-bottom:6px;}}
.ps-mono {{font-family: Consolas, "Courier New", monospace; font-size:0.85rem; background:#F5F8FD;
  border:1px solid #DCE5F1; border-radius:6px; padding:8px 10px; word-break: break-all;}}

/* Tabs */
.stTabs [data-baseweb="tab-list"] {{gap: 4px; border-bottom: 2px solid #DCE5F1;}}
.stTabs [data-baseweb="tab"] {{background:#F3F6FB; border-radius:8px 8px 0 0; padding:8px 18px;
  font-weight:600; color:#33475F;}}
.stTabs [aria-selected="true"] {{background:{PRIMARY} !important; color:#FFFFFF !important;}}
.stTabs [aria-selected="true"] p {{color:#FFFFFF !important;}}
.stTabs [data-baseweb="tab-highlight"] {{display:none;}}

/* Sidebar */
section[data-testid="stSidebar"] {{background: linear-gradient(180deg, {NAVY} 0%, #13315C 100%);}}
section[data-testid="stSidebar"] * {{color: #E6EEF8;}}
section[data-testid="stSidebar"] .ps-side-brand {{font-size:1.35rem; font-weight:800; color:#FFFFFF;
  letter-spacing:0.3px; margin-bottom:0;}}
section[data-testid="stSidebar"] .ps-side-tag {{font-size:0.8rem; color:#A9BCD6; margin-bottom:14px;}}
section[data-testid="stSidebar"] .ps-side-head {{font-size:0.72rem; letter-spacing:1.4px; color:#8FA6C4;
  margin: 10px 0 4px 2px; font-weight:700;}}
section[data-testid="stSidebar"] div[role="radiogroup"] {{gap: 3px;}}
section[data-testid="stSidebar"] div[role="radiogroup"] label {{
  padding: 9px 12px; border-radius: 7px; width: 100%; border: 1px solid transparent;
  transition: background 0.15s;}}
section[data-testid="stSidebar"] div[role="radiogroup"] label:hover {{background: rgba(255,255,255,0.08);}}
section[data-testid="stSidebar"] div[role="radiogroup"] label:has(input:checked) {{
  background: #1F63B8; border-color:#4B8BDA;}}
section[data-testid="stSidebar"] div[role="radiogroup"] label > div:first-child {{display:none;}}
section[data-testid="stSidebar"] div[role="radiogroup"] {{width: 100%;}}
section[data-testid="stSidebar"] label[data-testid="stRadioOption"] {{display:flex; width:100% !important;
  box-sizing:border-box;}}
section[data-testid="stSidebar"] label[data-testid="stRadioOption"] > div:not([data-testid="stMarkdownContainer"])
  > div:not([data-testid="stMarkdownContainer"]) {{display:none !important;}}
section[data-testid="stSidebar"] label[data-testid="stRadioOption"][data-selected="true"] {{
  background: #1F63B8; border-color:#4B8BDA;}}
section[data-testid="stSidebar"] div[role="radiogroup"] label p {{font-size:0.95rem; font-weight:600;}}
section[data-testid="stSidebar"] .ps-side-foot {{font-size:0.75rem; color:#8FA6C4; margin-top:18px;
  border-top:1px solid rgba(255,255,255,0.15); padding-top:10px;}}

div[data-testid="stDownloadButton"] button {{width:100%; border:1px solid {PRIMARY}; color:{PRIMARY};
  font-weight:600;}}
div[data-testid="stDownloadButton"] button:hover {{background:{PRIMARY}; color:white;}}
</style>
"""


def inject_css() -> None:
    st.markdown(CSS, unsafe_allow_html=True)


def title_bar() -> None:
    st.markdown(
        f"""<div class="ps-titlebar">
        <div class="ps-title">{APP_NAME}</div>
        <div class="ps-subtitle">Developed by <b>Randy Singh</b> from <b>Kalsnet (KNet) Consulting Group</b></div>
        </div>""",
        unsafe_allow_html=True,
    )


def page_heading(number: str, title: str, subtitle: str) -> None:
    st.markdown(
        f"""<div class="ps-page"><span class="ps-page-num">{number}</span>
        <span class="ps-page-title">{title}</span></div>
        <div class="ps-page-sub">{subtitle}</div>""",
        unsafe_allow_html=True,
    )


def cards(items: Sequence[Tuple[str, str]], cols: int = 3) -> None:
    for i in range(0, len(items), cols):
        row = st.columns(cols)
        for c, (h, p) in zip(row, items[i:i + cols]):
            c.markdown(f'<div class="ps-card"><h4>{h}</h4><p>{p}</p></div>', unsafe_allow_html=True)
        st.write("")


def metrics(items: Sequence[Tuple[str, str]], status: Optional[Sequence[str]] = None) -> None:
    row = st.columns(len(items))
    for i, (c, (lbl, val)) in enumerate(zip(row, items)):
        cls = status[i] if status else ""
        c.markdown(f'<div class="ps-metric {cls}"><div class="lbl">{lbl}</div><div class="val">{val}</div></div>',
                   unsafe_allow_html=True)


def callout(text: str, kind: str = "") -> None:
    st.markdown(f'<div class="ps-callout {kind}">{text}</div>', unsafe_allow_html=True)


def mono(text: str) -> None:
    st.markdown(f'<div class="ps-mono">{text}</div>', unsafe_allow_html=True)


def section(title: str) -> None:
    st.markdown(f"#### {title}")


def formula(name: str, latex: str, explanation: str) -> None:
    st.markdown(f'<div class="ps-formula-name">{name}</div>', unsafe_allow_html=True)
    st.latex(latex)
    st.markdown(f'<div class="ps-formula-text">{explanation}</div>', unsafe_allow_html=True)


def field_table(rows: List[Dict[str, str]]) -> None:
    df = pd.DataFrame(rows, columns=["Field", "Type", "Description", "Example"])
    st.dataframe(df, hide_index=True, width="stretch",
                 column_config={"Description": st.column_config.TextColumn(width="large")})


def show_df(df: pd.DataFrame, height: Optional[int] = None) -> None:
    kwargs = {"hide_index": True, "width": "stretch"}
    if height:
        kwargs["height"] = height
    st.dataframe(safe_df(df), **kwargs)


def safe_df(df: pd.DataFrame) -> pd.DataFrame:
    """Convert columns holding integers beyond 64 bits to strings for display."""
    out = df.copy()
    for c in out.columns:
        if out[c].dtype == object:
            if any(isinstance(v, int) and not isinstance(v, bool) and abs(v) >= 2 ** 53 for v in out[c]):
                out[c] = out[c].astype(str)
    return out


def short(v, limit: int = 24) -> str:
    s = str(v)
    return s if len(s) <= limit else f"{s[:10]}...{s[-8:]}"


def style_fig(fig, height: int = 380, title: Optional[str] = None):
    fig.update_layout(
        template="plotly_white", height=height, margin=dict(l=20, r=20, t=50 if title else 20, b=20),
        font=dict(family="Segoe UI, Helvetica, Arial", size=12, color="#22354D"),
        title=dict(text=title, font=dict(size=15, color=NAVY)) if title else None,
        colorway=PLOT_COLORS, legend=dict(orientation="h", yanchor="bottom", y=1.02, x=0),
    )
    return fig


# ---------------------------------------------------------------------------
# Uploads
# ---------------------------------------------------------------------------
def read_upload(file) -> Dict:
    """Return {'name', 'kind', 'text', 'df'} for CSV, TXT, DOCX or PDF uploads."""
    name = file.name
    ext = name.rsplit(".", 1)[-1].lower()
    raw = file.getvalue()
    text, df = "", None
    if ext == "csv":
        df = pd.read_csv(io.BytesIO(raw), dtype=str, keep_default_na=False)
        df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]
        text = raw.decode("utf-8", errors="replace")
    elif ext == "txt":
        text = raw.decode("utf-8", errors="replace")
    elif ext == "docx":
        from docx import Document
        d = Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            if rows and df is None and len(rows) > 1:
                header = [str(c).strip().lower().replace(" ", "_") for c in rows[0]]
                df = pd.DataFrame(rows[1:], columns=header)
            parts += ["\t".join(r) for r in rows]
        text = "\n".join(parts)
    elif ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    return {"name": name, "kind": ext, "text": text, "df": df}


def extract_integers(text: str, min_digits: int = 1) -> List[int]:
    """All whole numbers in text (decimal, or hex with 0x prefix), in order."""
    out = []
    for tok in re.findall(r"0[xX][0-9a-fA-F]+|\d[\d,_]*\d|\d", text):
        tok = tok.replace(",", "").replace("_", "")
        v = int(tok, 16) if tok.lower().startswith("0x") else int(tok)
        if len(str(v)) >= min_digits:
            out.append(v)
    return out


def data_source(key: str, help_text: str, template_df: Optional[pd.DataFrame] = None):
    """Radio: synthetic or upload. Returns None for synthetic, or the parsed upload dict."""
    c1, c2 = st.columns([1, 2])
    with c1:
        choice = st.radio("Data source", ["Synthetic demonstration data", "Upload real data"],
                          key=f"{key}_src", help="Synthetic data is generated so every step can be checked. "
                                                 "Choose upload to analyse your own file.")
    upload = None
    if choice == "Upload real data":
        with c2:
            st.markdown(f'<div class="ps-callout">{help_text}</div>', unsafe_allow_html=True)
            f = st.file_uploader("Upload file (CSV, TXT, Word DOCX or PDF)", type=UPLOAD_TYPES, key=f"{key}_file")
            if template_df is not None:
                st.download_button("Download CSV template", template_df.to_csv(index=False).encode(),
                                   file_name=f"{key}_template.csv", mime="text/csv", key=f"{key}_tpl")
        if f is not None:
            try:
                upload = read_upload(f)
            except Exception as exc:  # pragma: no cover - user file issues
                st.error(f"Could not read the file: {exc}")
        else:
            st.info("No file uploaded yet. Showing synthetic data until a file is provided.")
    return upload


# ---------------------------------------------------------------------------
# Export panel
# ---------------------------------------------------------------------------
def export_panel(report: Optional[Report], key: str) -> None:
    if report is None:
        st.info("Run the live demonstration first to produce results to export.")
        return
    callout("Download the results of the current demonstration. The PDF and Word reports include the "
            "summary, formulas, a chart and all result tables. The CSV contains the main result table. "
            "The text file is a plain report suitable for logs and email.")
    base = re.sub(r"[^A-Za-z0-9]+", "_", report.use_case).strip("_").lower()
    c = st.columns(4)
    exporters = [("PDF report", to_pdf, "pdf", "application/pdf"),
                 ("Word report", to_docx, "docx",
                  "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                 ("Text report", to_txt, "txt", "text/plain"),
                 ("CSV data", to_csv, "csv", "text/csv")]
    for col, (label, fn, ext, mime) in zip(c, exporters):
        with col:
            try:
                data = fn(report)
                st.download_button(f"Download {label}", data, file_name=f"primeshield_{base}.{ext}",
                                   mime=mime, key=f"{key}_{ext}")
            except Exception as exc:  # pragma: no cover
                st.error(f"{label} failed: {exc}")
    st.markdown("##### Preview of exported summary")
    if report.summary:
        show_df(pd.DataFrame(report.summary, columns=["Metric", "Value"]))
    if report.tables:
        st.markdown(f"##### Preview: {report.tables[0][0]}")
        show_df(report.tables[0][1].head(50))


def use_case_tabs():
    """Standard five tab layout used by every use case page."""
    return st.tabs(["Overview and Benefits", "Fields and Formulas", "Live Demonstration",
                    "Diagrams and Charts", "Export Results"])


# ============================================================================
# PAGE: home
# Page 00: Prime factorization engine and application overview.
# ============================================================================

HOME_KEY = "home"


def home_overview() -> None:
    callout("<b>Prime factorization</b> writes a whole number as a product of primes. By the Fundamental "
            "Theorem of Arithmetic this product is unique, for example 360 = 2<sup>3</sup> x 3<sup>2</sup> x 5. "
            "Multiplying primes is instant, but recovering them from a large product is extremely slow. That "
            "one way asymmetry protects much of the internet, and every use case in this suite builds on it.")
    cards([
        ("Unique building blocks", "Every integer above 1 has exactly one prime factorization, so it acts as "
         "a fingerprint for divisibility, GCD, LCM and modular arithmetic."),
        ("Easy one way, hard the other", "A 2048 bit product of two primes is computed in microseconds, while "
         "the best known classical method (General Number Field Sieve) would need an estimated "
         "10<sup>35</sup> or more operations to reverse it."),
        ("Measurable security", "The time needed to factor grows sub exponentially with key size. This page "
         "measures that growth live on your machine with synthetic semiprimes."),
    ])
    st.markdown("#### Use cases in this suite")
    cards([
        ("01 RSA Encryption", "Encrypt with n and e, decrypt with d. Breaking it means factoring n."),
        ("02 Digital Signatures", "Sign software, contracts and certificates; verify authenticity and integrity."),
        ("03 Online Banking and Payments", "RSA key transport sets up a session key that protects transactions."),
        ("04 Quantum Computing Research", "Simulate Shor's algorithm and compare classical and quantum cost."),
        ("05 Fractions, GCD and LCM", "Simplify fractions, align schedules and design gear trains."),
        ("06 Hashing and Random Numbers", "Blum Blum Shub generator and Very Smooth Hash, tested with NIST methods."),
        ("07 Security Testing", "Find weak keys in a device fleet with batch GCD, Fermat and Pollard methods."),
        ("08 Accuracy Validation", "Live self test of every algorithm against published reference values."),
    ])
    st.markdown("#### How to use each page")
    callout("Select a use case in the left panel. Each page has five tabs: <b>Overview and Benefits</b>, "
            "<b>Fields and Formulas</b>, <b>Live Demonstration</b> (synthetic data or your upload), "
            "<b>Diagrams and Charts</b>, and <b>Export Results</b> (PDF, Word, text and CSV).")


def home_schema() -> None:
    st.markdown("#### Input fields")
    field_table([
        {"Field": "n", "Type": "Whole number", "Description": "The number to factor. Decimal or hexadecimal with a 0x prefix. Commas and spaces are ignored.", "Example": "600851475143"},
        {"Field": "Time budget", "Type": "Seconds", "Description": "Maximum time for the search. Parts not split in time are reported as composite cofactors instead of guessing.", "Example": "10"},
        {"Field": "Upload", "Type": "CSV, TXT, DOCX, PDF", "Description": "CSV with a column named number, or any document; every whole number found in the text is factored.", "Example": "number\n360\n1001"},
    ])
    st.markdown("#### Output fields")
    field_table([
        {"Field": "Factorization", "Type": "Product of prime powers", "Description": "n written as p1^e1 x p2^e2 x ... with every pi proven prime.", "Example": "2^3 x 3^2 x 5"},
        {"Field": "Number of divisors d(n)", "Type": "Integer", "Description": "How many positive whole numbers divide n exactly.", "Example": "24"},
        {"Field": "Sum of divisors sigma(n)", "Type": "Integer", "Description": "Sum of all positive divisors of n.", "Example": "1170"},
        {"Field": "Euler phi(n)", "Type": "Integer", "Description": "Count of numbers from 1 to n sharing no factor with n. The key quantity in RSA.", "Example": "96"},
        {"Field": "Method log", "Type": "Table", "Description": "Which algorithm found each factor, so the result can be audited.", "Example": "Pollard rho: 8051 = 83 x 97"},
    ])
    st.markdown("#### Formulas")
    formula("Fundamental Theorem of Arithmetic", r"n = p_1^{e_1}\,p_2^{e_2}\cdots p_k^{e_k}",
            "Every n greater than 1 is a unique product of primes p<sub>i</sub> raised to exponents e<sub>i</sub>.")
    formula("Number and sum of divisors", r"d(n)=\prod_{i=1}^{k}(e_i+1)\qquad \sigma(n)=\prod_{i=1}^{k}\frac{p_i^{e_i+1}-1}{p_i-1}",
            "Both follow directly from the factorization; without it they are hard to compute for large n.")
    formula("Euler totient", r"\varphi(n)=\prod_{i=1}^{k}p_i^{e_i-1}(p_i-1)",
            "For an RSA modulus n = p q this is (p - 1)(q - 1), which unlocks the private key.")
    formula("Pollard rho iteration", r"x_{i+1} = x_i^2 + c \pmod n,\qquad g=\gcd(|x_i - x_{2i}|,\,n)",
            "The sequence cycles modulo a hidden factor p after about the square root of p steps; the GCD exposes p.")
    formula("Miller Rabin primality", r"n-1 = 2^s d,\quad a^d \equiv 1 \ \text{or}\ a^{2^r d}\equiv -1 \pmod n",
            "If this fails for any base a, n is composite. With the 13 fixed bases used here the test is exact for n below 3.3 x 10<sup>24</sup>.")
    formula("GNFS classical cost", r"L_n\!\left[\tfrac13,\ \sqrt[3]{64/9}\right]=\exp\!\Big(1.923\,(\ln n)^{1/3}(\ln\ln n)^{2/3}\Big)",
            "Heuristic running time of the fastest known classical factoring algorithm for large n.")


def home_demo():
    upload = data_source(HOME_KEY, "CSV with a column named <b>number</b>, or a TXT, Word or PDF file. Every whole "
                              "number found in the file is factored.",
                         pd.DataFrame({"number": ["360", "1001", "600851475143", "18446744073709551617"]}))
    c1, c2 = st.columns([3, 1])
    with c1:
        raw = st.text_input("Number to factor (n)", value="600851475143", key=f"{HOME_KEY}_n",
                            help="Decimal or 0x hexadecimal. Try 18446744073709551617 (2^64 + 1).")
    with c2:
        budget = st.number_input("Time budget (seconds)", 1, 60, 10, key=f"{HOME_KEY}_budget")
    try:
        n = parse_int(raw)
        if n < 1:
            raise ValueError("n must be positive")
    except Exception as exc:
        st.error(f"Invalid number: {exc}")
        return None
    res = factorize(n, time_limit=float(budget))
    st.markdown("##### Result")
    st.latex(f"{short(n, 60)} = {res.latex()}" if len(str(n)) < 60 else res.latex())
    status = "ok" if res.complete else "warn"
    metrics([("Digits", f"{len(str(n))}"), ("Bits", f"{n.bit_length()}"),
             ("Distinct primes", f"{len(res.factors)}"),
             ("Time", f"{res.elapsed * 1000:.1f} ms"),
             ("Status", "Complete" if res.complete else "Partial")],
            status=["", "", "", "", status])
    if res.is_prime:
        callout(f"{short(n, 40)} is prime.", "ok")
    if not res.complete:
        callout("The time budget ran out before every part was split. This is the practical meaning of "
                "factoring hardness: raise the budget, or accept that n is out of reach.", "warn")
    st.markdown("##### Method log")
    show_df(pd.DataFrame(res.log))

    batch_rows = []
    source_numbers = []
    if upload:
        if upload["df"] is not None and "number" in upload["df"].columns:
            source_numbers = [parse_int(v) for v in upload["df"]["number"] if str(v).strip()]
        else:
            source_numbers = extract_integers(upload["text"])
        st.markdown(f"##### Batch results from {upload['name']}")
    else:
        rng = random.Random(2026)
        source_numbers = [360, 1001, 65537, 561, 2 ** 32 + 1, 999_999_000_001, 10 ** 12 + 39 * 10 ** 5]
        source_numbers += [random_semiprime(b, rng)[0] for b in (40, 48, 56, 64)]
        st.markdown("##### Batch results for synthetic numbers")
    t0 = time.perf_counter()
    for v in source_numbers[:200]:
        if v < 1:
            continue
        r = factorize(v, time_limit=3.0)
        batch_rows.append({"number": str(v), "bits": v.bit_length(), "factorization": r.plain_text(),
                           "distinct primes": len(r.factors), "prime": "Yes" if r.is_prime else "No",
                           "divisors d(n)": str(divisor_count(r.factors)) if r.complete else "",
                           "phi(n)": str(euler_phi(r.factors)) if r.complete else "",
                           "complete": "Yes" if r.complete else "No", "ms": round(r.elapsed * 1000, 2)})
    batch = pd.DataFrame(batch_rows)
    show_df(batch)
    return {"n": n, "res": res, "batch": batch}


def home_visuals(ctx) -> None:
    if ctx is None:
        st.info("Enter a valid number in the Live Demonstration tab.")
        return
    n, res = ctx["n"], ctx["res"]
    c1, c2 = st.columns([1, 1])
    with c1:
        st.markdown("##### Factor tree")
        if n > 1:
            st.graphviz_chart(factor_tree_dot(n, res))
        callout("Each level splits off the smallest prime (green). Red nodes, if any, are composite parts "
                "that could not be split within the time budget.")
    with c2:
        st.markdown("##### Prime exponents")
        if res.factors:
            fig = go.Figure(go.Bar(x=[short(p, 14) for p in res.factors], y=list(res.factors.values()),
                                   marker_color="#0B4F9C", text=list(res.factors.values()), textposition="outside"))
            fig.update_xaxes(type="category", title="Prime factor p")
            fig.update_yaxes(title="Exponent e", dtick=1)
            st.plotly_chart(style_fig(fig, 330), key=f"{HOME_KEY}_exp")
        if res.complete and n > 1:
            metrics([("d(n)", str(divisor_count(res.factors))), ("sigma(n)", short(divisor_sum(res.factors), 18)),
                     ("phi(n)", short(euler_phi(res.factors), 18))])

    st.markdown("##### Hardness benchmark: time to factor a semiprime n = p x q")
    callout("Synthetic semiprimes with two primes of equal size are generated and factored live with "
            "Pollard rho. The work grows like the square root of the smaller prime, so each extra 8 bits of "
            "n multiplies the time by about four. Real RSA keys use 2048 bits or more.")
    max_bits = st.select_slider("Largest semiprime size (bits)", options=[40, 48, 56, 64, 72, 80], value=64,
                                key=f"{HOME_KEY}_bench_bits")
    if st.button("Run benchmark", key=f"{HOME_KEY}_bench", type="primary"):
        rng = random.Random(7)
        rows = []
        prog = st.progress(0.0)
        sizes = list(range(24, max_bits + 1, 8))
        for i, b in enumerate(sizes):
            times = []
            for _ in range(3):
                nn, p, q = random_semiprime(b, rng)
                t = time.perf_counter()
                factorize(nn, time_limit=60)
                times.append((time.perf_counter() - t) * 1000)
            rows.append({"bits": b, "median ms": sorted(times)[1], "max ms": max(times)})
            prog.progress((i + 1) / len(sizes))
        st.session_state[f"{HOME_KEY}_bench_df"] = pd.DataFrame(rows)
    bench = st.session_state.get(f"{HOME_KEY}_bench_df")
    if bench is not None:
        fig = go.Figure()
        fig.add_trace(go.Scatter(x=bench["bits"], y=bench["median ms"], mode="lines+markers", name="Median time"))
        fig.add_trace(go.Scatter(x=bench["bits"], y=bench["max ms"], mode="lines+markers", name="Worst of three",
                                 line=dict(dash="dot")))
        fig.update_yaxes(type="log", title="Milliseconds (log scale)")
        fig.update_xaxes(title="Size of n in bits")
        st.plotly_chart(style_fig(fig, 360), key=f"{HOME_KEY}_benchfig")
        show_df(bench.round(2))
    else:
        st.caption("Press Run benchmark to measure factoring time on this server.")


def home_report(ctx):
    if ctx is None:
        return None
    res = ctx["res"]
    rep = Report(use_case="Prime Factorization Engine")
    rep.summary = [("Input n", str(ctx["n"])), ("Bits", str(ctx["n"].bit_length())),
                   ("Factorization", res.plain_text()), ("Complete", "Yes" if res.complete else "No"),
                   ("Time (ms)", f"{res.elapsed * 1000:.2f}")]
    if res.complete and ctx["n"] > 1:
        rep.summary += [("Number of divisors", str(divisor_count(res.factors))),
                        ("Sum of divisors", str(divisor_sum(res.factors))),
                        ("Euler phi", str(euler_phi(res.factors)))]
    rep.formulas = [("Fundamental theorem", "n = p1^e1 x p2^e2 x ... x pk^ek"),
                    ("Divisor count", "d(n) = (e1 + 1)(e2 + 1)...(ek + 1)"),
                    ("Euler phi", "phi(n) = product of p^(e-1) (p - 1)"),
                    ("Pollard rho", "x(i+1) = x(i)^2 + c mod n, g = gcd(|x(i) - x(2i)|, n)")]
    rep.tables = [("Batch factorization results", ctx["batch"]), ("Method log", pd.DataFrame(res.log))]
    b = ctx["batch"]
    if len(b):
        rep.chart = {"kind": "bar", "x": [short(v, 12) for v in b["number"]], "y": {"Bits": list(b["bits"])},
                     "title": "Size of each factored number", "xlabel": "Number", "ylabel": "Bits"}
    return rep


def home_render() -> None:
    page_heading("00", "Prime Factorization Engine", "Factor any whole number, see how it was done, and measure why large numbers resist factoring.")
    t = use_case_tabs()
    with t[2]:
        ctx = home_demo()
    with t[0]:
        home_overview()
    with t[1]:
        home_schema()
    with t[3]:
        home_visuals(ctx)
    with t[4]:
        export_panel(home_report(ctx), HOME_KEY)


# ============================================================================
# PAGE: rsa
# Page 01: RSA encryption.
# ============================================================================

RSA_KEY = "rsa"


@st.cache_data(show_spinner="Generating RSA key...", max_entries=32)
def cached_key(bits: int, e: int, seed: int):
    return generate_rsa_key(bits, e, random.Random(seed))


@st.cache_data(show_spinner="Attacker is trying to factor n...", max_entries=32)
def cached_attack(n: int, e: int, cts: tuple, limit: float):
    return break_rsa_by_factoring(n, e, list(cts), time_limit=limit)

SYNTH_MESSAGES = [
    "Quarterly board minutes: merger talks with Northwind Ltd remain confidential.",
    "VPN credential rotation scheduled for 02:00 UTC Saturday.",
    "Patient record 55821: lab results attached, share only with Dr. Lee.",
    "Wire approval code for invoice INV-2026-0931 is 448213.",
]


def rsa_overview() -> None:
    callout("<b>RSA</b> (Rivest, Shamir, Adleman, 1977) is a public key cryptosystem. Anyone can encrypt with "
            "the public key (n, e); only the holder of the private exponent d can decrypt. Computing d requires "
            "phi(n) = (p - 1)(q - 1), which in turn requires the prime factors of n. <b>Breaking RSA is at most "
            "as hard as factoring n</b>, and no faster general attack is known.")
    cards([
        ("Secure web (HTTPS)", "Server certificates carry RSA public keys, letting browsers authenticate "
         "sites and, in older TLS versions, transport the session key."),
        ("Email and VPN", "S/MIME and PGP encrypt message keys with the recipient's RSA key; IPsec and "
         "OpenVPN authenticate peers with RSA certificates."),
        ("No shared secret needed", "Two parties who have never met can communicate securely: the public "
         "key can be published openly."),
        ("Scalable key management", "One key pair per person instead of one secret per pair of people "
         "(n keys instead of n(n - 1)/2)."),
        ("Tunable security", "Security grows with key size: 2048 bits is the current minimum, 3072 bits "
         "matches 128 bit symmetric strength."),
        ("Proven track record", "Nearly 50 years of public analysis; the largest RSA challenge number "
         "factored publicly is 829 bits (RSA-250, in 2020)."),
    ])
    callout("This demonstration uses textbook RSA so every number is visible. Production systems must add "
            "OAEP padding and use vetted libraries. Keys here are shortened so the factoring attack can be "
            "shown live.", "warn")


def rsa_schema() -> None:
    st.markdown("#### Key fields")
    field_table([
        {"Field": "p, q", "Type": "Secret primes", "Description": "Two large random primes of similar size. Never shared.", "Example": "61, 53"},
        {"Field": "n", "Type": "Public modulus", "Description": "n = p x q. Its bit length is the key size.", "Example": "3233"},
        {"Field": "phi(n)", "Type": "Secret", "Description": "Euler totient (p - 1)(q - 1). Anyone who factors n can compute it.", "Example": "3120"},
        {"Field": "e", "Type": "Public exponent", "Description": "Coprime with phi(n); 65537 is the standard choice.", "Example": "17"},
        {"Field": "d", "Type": "Private exponent", "Description": "Inverse of e modulo phi(n), found with the extended Euclidean algorithm.", "Example": "2753"},
        {"Field": "dp, dq, qinv", "Type": "Secret (CRT)", "Description": "Precomputed values that make decryption about four times faster.", "Example": "53, 49, 38"},
    ])
    st.markdown("#### Message fields")
    field_table([
        {"Field": "Plaintext", "Type": "Text", "Description": "UTF-8 message, padded and cut into blocks smaller than n.", "Example": "HELLO"},
        {"Field": "Block size", "Type": "Bytes", "Description": "floor((bits of n - 1) / 8) bytes per block, so every block m satisfies m < n.", "Example": "63 for 512 bit n"},
        {"Field": "m", "Type": "Integer", "Description": "Plaintext block as a number.", "Example": "65"},
        {"Field": "c", "Type": "Integer", "Description": "Ciphertext block m^e mod n.", "Example": "2790"},
        {"Field": "Upload CSV", "Type": "Columns", "Description": "Column message (one message per row). TXT, Word and PDF files are encrypted as one message.", "Example": "message"},
    ])
    st.markdown("#### Formulas")
    formula("Key generation", r"n = p\,q,\qquad \varphi(n)=(p-1)(q-1),\qquad d \equiv e^{-1} \pmod{\varphi(n)}",
            "Choose e with gcd(e, phi(n)) = 1, then solve e d = 1 + k phi(n) with the extended Euclidean algorithm.")
    formula("Encryption", r"c = m^{e} \bmod n", "Uses only the public key; computed with fast square and multiply.")
    formula("Decryption", r"m = c^{d} \bmod n",
            "Works because m<sup>ed</sup> = m<sup>1 + k phi(n)</sup> = m (mod n) by Euler's theorem.")
    formula("CRT decryption", r"m_1=c^{d_p}\bmod p,\ \ m_2=c^{d_q}\bmod q,\ \ h=q^{-1}(m_1-m_2)\bmod p,\ \ m=m_2+h\,q",
            "Two half size exponentiations instead of one full size one.")
    formula("Attack by factoring", r"n \;\xrightarrow{\text{factor}}\; (p,q) \;\to\; \varphi(n) \;\to\; d = e^{-1} \bmod \varphi(n)",
            "Once n is factored the attacker holds the private key. This is why n must be too large to factor.")
    formula("Security level", r"\text{work} \approx L_n\!\left[\tfrac13,1.923\right]",
            "About 2<sup>112</sup> operations for 2048 bit n and 2<sup>128</sup> for 3072 bit n (NIST SP 800-57).")


def rsa_demo():
    upload = data_source(RSA_KEY, "CSV with a column named <b>message</b>, or a TXT, Word or PDF document "
                              "that is encrypted as one message.",
                         pd.DataFrame({"message": SYNTH_MESSAGES[:2]}))
    c1, c2, c3 = st.columns([1.2, 1, 1])
    with c1:
        mode = st.radio("Key source", ["Generate random key", "Enter my own primes p and q"], key=f"{RSA_KEY}_mode")
    with c2:
        bits = st.select_slider("Key size (bits of n)", options=[32, 48, 64, 80, 96, 128, 256, 512, 1024, 2048],
                                value=64, key=f"{RSA_KEY}_bits",
                                help="Up to about 96 bits the factoring attack succeeds live; above that it times out.")
    with c3:
        e = st.number_input("Public exponent e", min_value=3, value=65537, step=2, key=f"{RSA_KEY}_e")
        seed = st.number_input("Random seed", min_value=0, value=42, key=f"{RSA_KEY}_seed")
    try:
        if mode.startswith("Enter"):
            a, b = st.columns(2)
            p = parse_int(a.text_input("Prime p", "1000003", key=f"{RSA_KEY}_p"))
            q = parse_int(b.text_input("Prime q", "999983", key=f"{RSA_KEY}_q"))
            for name, v in (("p", p), ("q", q)):
                if not is_probable_prime(v):
                    st.error(f"{name} = {v} is not prime.")
                    return None
            key = key_from_primes(p, q, int(e))
        else:
            key = cached_key(int(bits), int(e), int(seed))
    except Exception as exc:
        st.error(str(exc))
        return None
    if key.n.bit_length() < 9:
        st.error("n is too small to hold a byte; choose larger primes.")
        return None

    st.markdown("##### Key pair")
    show_df(pd.DataFrame(key.as_rows()))

    if upload:
        if upload["df"] is not None and "message" in upload["df"].columns:
            messages = [m for m in upload["df"]["message"].astype(str) if m.strip()]
        else:
            messages = [upload["text"].strip()[:4000]]
        st.caption(f"Messages loaded from {upload['name']}: {len(messages)}")
    else:
        messages = SYNTH_MESSAGES
    msg = st.selectbox("Message to walk through", messages, key=f"{RSA_KEY}_msg",
                       format_func=lambda s: s[:90] + ("..." if len(s) > 90 else ""))

    t = time.perf_counter()
    rows = encrypt_text(msg, key.n, key.e)
    enc_ms = (time.perf_counter() - t) * 1000
    t = time.perf_counter()
    crt_rows = [decrypt_crt(r["c"], key) for r in rows]
    dec_ms = (time.perf_counter() - t) * 1000
    recovered = blocks_to_text([r["m"] for r in crt_rows], key.n)
    ok = recovered == msg
    metrics([("Key size", f"{key.bits} bits"), ("Block size", f"{block_size_bytes(key.n)} bytes"),
             ("Blocks", str(len(rows))), ("Encrypt", f"{enc_ms:.2f} ms"), ("Decrypt (CRT)", f"{dec_ms:.2f} ms"),
             ("Round trip", "Exact match" if ok else "Mismatch")], status=["", "", "", "", "", "ok" if ok else "bad"])
    block_df = pd.DataFrame([{"block": r["block"], "plaintext m": str(r["m"]), "ciphertext c = m^e mod n": str(r["c"]),
                              "m1 = c^dp mod p": str(cr["m1"]), "m2 = c^dq mod q": str(cr["m2"]),
                              "recovered m": str(cr["m"]), "match": "Yes" if cr["m"] == r["m"] else "No"}
                             for r, cr in zip(rows, crt_rows)])
    st.markdown("##### Encryption and decryption, block by block")
    show_df(block_df)
    st.markdown("##### Decrypted message")
    mono(recovered.replace("<", "&lt;"))

    # Batch over all messages
    batch = []
    for i, m in enumerate(messages[:100]):
        rr = encrypt_text(m, key.n, key.e)
        back = blocks_to_text([decrypt_crt(x["c"], key)["m"] for x in rr], key.n)
        batch.append({"message_id": i + 1, "characters": len(m), "blocks": len(rr),
                      "first ciphertext block": short(rr[0]["c"], 30), "decrypts correctly": "Yes" if back == m else "No"})
    batch_df = pd.DataFrame(batch)

    st.markdown("##### Attacker view: break the key by factoring n")
    callout("The attacker sees only n, e and the ciphertext. If n can be factored, the private key follows "
            "immediately. Try a key of 96 bits or less to watch it succeed, then 128 bits or more to watch it fail.")
    limit = st.slider("Attacker time budget (seconds)", 1, 30, 5, key=f"{RSA_KEY}_atk_t")
    attack = cached_attack(key.n, key.e, tuple(r["c"] for r in rows), float(limit))
    fr = attack["factor_result"]
    if attack["success"]:
        text = blocks_to_text(attack["plain_blocks"], key.n)
        callout(f"<b>Key broken in {fr.elapsed:.3f} s.</b> n = {short(attack['key'].p, 30)} x "
                f"{short(attack['key'].q, 30)}; recovered d matches the real d: "
                f"<b>{'Yes' if attack['key'].d == key.d else 'No'}</b>. Decrypted text: "
                f"<i>{text[:120].replace('<', '&lt;')}</i>", "bad")
    else:
        callout(f"<b>Attack failed</b> within {limit} s: n ({key.bits} bits) could not be factored. "
                f"Estimated classical effort for this size: about 10<sup>{gnfs_log10_ops(key.bits):.0f}</sup> "
                "operations with the General Number Field Sieve.", "ok")
    return {"key": key, "rows": rows, "block_df": block_df, "batch_df": batch_df, "msg": msg,
            "recovered": recovered, "attack": attack, "enc_ms": enc_ms, "dec_ms": dec_ms}


def rsa_flow_dot() -> str:
    return """digraph R {
rankdir=LR; bgcolor="transparent"; nodesep=0.5;
node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=11, fillcolor="#E8F0FB", color="#1F4E8C"];
edge [fontname="Helvetica", fontsize=9, color="#4A5B72"];
subgraph cluster_k { label="Key owner (bank, website)"; style="rounded,dashed"; color="#1F4E8C"; fontname="Helvetica";
  P [label="Pick primes p, q"]; N [label="n = p x q\\nphi = (p-1)(q-1)"]; D [label="d = e^-1 mod phi", fillcolor="#FDECEC", color="#B23A3A"]; }
PUB [label="Public key (n, e)\\npublished", fillcolor="#D6F0E0", color="#1E7B4A"];
S [label="Sender\\nmessage m"]; C [label="c = m^e mod n"]; R [label="m = c^d mod n"];
A [label="Attacker\\nsees n, e, c", fillcolor="#FFF5EB", color="#E07A1F"]; F [label="Factor n ?", shape=diamond, fillcolor="#FFF5EB", color="#E07A1F"];
P -> N -> D; N -> PUB; PUB -> C [label="encrypt"]; S -> C; C -> R [label="network"]; D -> R [label="decrypt"];
C -> A [style=dashed]; PUB -> A [style=dashed]; A -> F; F -> D [label="if yes: key recovered", style=dashed, color="#B23A3A"];
}"""


def rsa_visuals(ctx) -> None:
    st.markdown("##### RSA data flow")
    st.graphviz_chart(rsa_flow_dot())
    if ctx is None:
        return
    key, rows = ctx["key"], ctx["rows"]
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("##### Plaintext and ciphertext blocks (as fraction of n)")
        fig = go.Figure()
        xs = [r["block"] for r in rows]
        fig.add_trace(go.Bar(x=xs, y=[r["m"] / key.n for r in rows], name="Plaintext m / n"))
        fig.add_trace(go.Bar(x=xs, y=[r["c"] / key.n for r in rows], name="Ciphertext c / n"))
        fig.update_layout(barmode="group")
        fig.update_xaxes(title="Block")
        fig.update_yaxes(title="Value relative to n", range=[0, 1])
        st.plotly_chart(style_fig(fig, 340), key=f"{RSA_KEY}_blocks")
        st.caption("Ciphertext values look uniformly spread over 0 to n, hiding the structure of the text.")
    with c2:
        st.markdown("##### Classical effort to factor n versus key size")
        sizes = [64, 128, 256, 512, 768, 1024, 1536, 2048, 3072, 4096]
        fig = go.Figure(go.Scatter(x=sizes, y=[gnfs_log10_ops(b) for b in sizes], mode="lines+markers",
                                   name="GNFS log10 operations"))
        fig.add_hline(y=gnfs_log10_ops(829), line_dash="dot", annotation_text="RSA-250 (829 bits) factored 2020")
        fig.add_vline(x=key.bits, line_color="#E07A1F", annotation_text=f"Your key: {key.bits} bits")
        fig.update_xaxes(title="Key size (bits)")
        fig.update_yaxes(title="log10 of operations")
        st.plotly_chart(style_fig(fig, 340), key=f"{RSA_KEY}_gnfs")
    st.markdown("##### Security margin by key size")
    ref = pd.DataFrame([
        {"Key size": "1024 bits", "Symmetric equivalent": "80 bits", "Status": "Disallowed (NIST)"},
        {"Key size": "2048 bits", "Symmetric equivalent": "112 bits", "Status": "Minimum today"},
        {"Key size": "3072 bits", "Symmetric equivalent": "128 bits", "Status": "Recommended beyond 2030"},
        {"Key size": "7680 bits", "Symmetric equivalent": "192 bits", "Status": "High security"},
        {"Key size": "15360 bits", "Symmetric equivalent": "256 bits", "Status": "Very high security"},
    ])
    show_df(ref)
    st.caption("Source: NIST SP 800-57 Part 1 comparable strengths.")


def rsa_report(ctx):
    if ctx is None:
        return None
    k = ctx["key"]
    a = ctx["attack"]
    rep = Report(use_case="RSA Encryption")
    rep.summary = [("Key size", f"{k.bits} bits"), ("n", str(k.n)), ("e", str(k.e)),
                   ("Message", ctx["msg"]), ("Blocks", str(len(ctx["rows"]))),
                   ("Round trip correct", "Yes" if ctx["recovered"] == ctx["msg"] else "No"),
                   ("Factoring attack", f"Succeeded in {a['factor_result'].elapsed:.3f} s" if a["success"]
                    else "Failed within time budget")]
    rep.sections = [("Interpretation", "The public key (n, e) encrypts; the private exponent d decrypts. The "
                     "attacker simulation shows that anyone able to factor n rebuilds d. Keys used in practice "
                     "(2048 bits or more) are far beyond reach of classical factoring.")]
    rep.formulas = [("Key generation", "n = p q; phi(n) = (p - 1)(q - 1); d = e^-1 mod phi(n)"),
                    ("Encryption", "c = m^e mod n"), ("Decryption", "m = c^d mod n"),
                    ("CRT", "m1 = c^dp mod p; m2 = c^dq mod q; h = qinv (m1 - m2) mod p; m = m2 + h q")]
    rep.tables = [("Block level results", ctx["block_df"]), ("Key pair", pd.DataFrame(k.as_rows())),
                  ("All messages", ctx["batch_df"])]
    rep.chart = {"kind": "line", "x": [64, 128, 256, 512, 1024, 2048, 3072, 4096],
                 "y": {"GNFS log10 operations": [gnfs_log10_ops(b) for b in [64, 128, 256, 512, 1024, 2048, 3072, 4096]]},
                 "title": "Classical effort to factor n", "xlabel": "Key size (bits)", "ylabel": "log10 operations"}
    return rep


def rsa_render() -> None:
    page_heading("01", "RSA Encryption", "Secures HTTPS websites, email and VPNs. The public key contains n = p x q; only someone who knows p and q can decrypt.")
    t = use_case_tabs()
    with t[2]:
        ctx = rsa_demo()
    with t[0]:
        rsa_overview()
    with t[1]:
        rsa_schema()
    with t[3]:
        rsa_visuals(ctx)
    with t[4]:
        export_panel(rsa_report(ctx), RSA_KEY)


# ============================================================================
# PAGE: signatures
# Page 02: Digital signatures.
# ============================================================================

SIG_KEY = "sig"

SYNTH_DOCS = [
    ("REL-001", "Software update", "AcmeOS 14.2.1 installer build 88213, sha manifest v3"),
    ("REL-002", "Firmware", "Router firmware 3.7.9 for model KX-400, bootloader locked"),
    ("REL-003", "Code signing", "payments-sdk-2.4.0.jar compiled 2026-09-12"),
    ("DOC-101", "Legal e-document", "Master services agreement between KNet and Contoso, value USD 1,250,000"),
    ("DOC-102", "Legal e-document", "Non disclosure agreement, term 36 months, governing law New York"),
    ("CRT-201", "Certificate", "CN=portal.example.com, valid 2026-01-01 to 2026-12-31, key RSA-2048"),
    ("CRT-202", "Certificate", "CN=vpn.example.com, valid 2026-03-01 to 2027-02-28, key RSA-3072"),
    ("REL-004", "Software update", "Mobile banking app 9.1.0 release candidate 2"),
]


@st.cache_data(show_spinner=False)
def publisher_key(bits: int, seed: int):
    return generate_rsa_key(bits, 65537, random.Random(seed))


def signatures_overview() -> None:
    callout("A <b>digital signature</b> proves who created a message and that it has not changed. The signer "
            "hashes the message and raises the hash to the private exponent d. Anyone can verify with the "
            "public key (n, e). Forging a signature without d requires factoring n, the same hard problem "
            "behind RSA encryption.")
    cards([
        ("Software updates", "Operating systems and app stores refuse to install packages whose signature "
         "does not verify, blocking malicious updates."),
        ("Code signing", "Drivers, firmware and scripts are signed so devices only run code from trusted publishers."),
        ("Legal e-documents", "Signed contracts and invoices are tamper evident and support non repudiation "
         "under laws such as the US ESIGN Act and EU eIDAS."),
        ("Certificates (PKI)", "Certificate authorities sign server certificates; browsers verify the chain "
         "up to a trusted root before showing the padlock."),
        ("Integrity", "Changing a single character changes the SHA-256 hash completely, so the old signature fails."),
        ("Non repudiation", "Only the private key holder could have produced a valid signature."),
    ])


def signatures_schema() -> None:
    st.markdown("#### Document fields")
    field_table([
        {"Field": "document_id", "Type": "Text", "Description": "Identifier of the release, contract or certificate.", "Example": "REL-001"},
        {"Field": "category", "Type": "Text", "Description": "Software update, firmware, code signing, legal e-document or certificate.", "Example": "Firmware"},
        {"Field": "content", "Type": "Text or bytes", "Description": "The exact bytes that are signed. Any change invalidates the signature.", "Example": "Router firmware 3.7.9"},
        {"Field": "SHA-256 hash H(M)", "Type": "256 bit digest", "Description": "Fixed length fingerprint of the content.", "Example": "9f86d081..."},
        {"Field": "signature s", "Type": "Integer below n", "Description": "H(M)^d mod n, created with the private key.", "Example": "81236..."},
        {"Field": "verified", "Type": "Yes or No", "Description": "Whether s^e mod n equals H(M) mod n.", "Example": "Yes"},
        {"Field": "Upload CSV", "Type": "Columns", "Description": "document_id, content (category optional). TXT, Word or PDF files are signed as one document.", "Example": "document_id,content"},
    ])
    st.markdown("#### Formulas")
    formula("Hash", r"h = \mathrm{SHA256}(M) \bmod n", "The message of any length is reduced to a 256 bit fingerprint.")
    formula("Sign (private key)", r"s = h^{d} \bmod n", "Only the holder of d can compute this.")
    formula("Verify (public key)", r"s^{e} \bmod n \stackrel{?}{=} h",
            "Because (h<sup>d</sup>)<sup>e</sup> = h<sup>ed</sup> = h (mod n), a genuine signature always verifies.")
    formula("Forgery by factoring", r"n=p\,q \Rightarrow d=e^{-1}\bmod (p-1)(q-1) \Rightarrow s' = h'^{\,d}\bmod n",
            "Anyone who factors n can sign arbitrary content as the publisher.")
    formula("Avalanche effect", r"\Pr[\text{bit } i \text{ of } H(M) \ne \text{bit } i \text{ of } H(M')] \approx \tfrac12",
            "A one character edit flips about 128 of the 256 hash bits.")


def signatures_demo():
    upload = data_source(SIG_KEY, "CSV with columns <b>document_id</b> and <b>content</b>, or a TXT, Word or PDF "
                              "document that is signed as a whole.",
                         pd.DataFrame({"document_id": ["DOC-1"], "category": ["Contract"], "content": ["Agreement text"]}))
    c1, c2 = st.columns(2)
    bits = c1.select_slider("Publisher key size (bits)", options=[64, 96, 128, 256, 512, 1024, 2048], value=512, key=f"{SIG_KEY}_bits")
    seed = c2.number_input("Key seed", 0, 10_000, 7, key=f"{SIG_KEY}_seed")
    key = publisher_key(int(bits), int(seed))

    if upload:
        if upload["df"] is not None and "content" in upload["df"].columns:
            df = upload["df"]
            docs = [(str(r.get("document_id", f"DOC-{i+1}")), str(r.get("category", "Uploaded")), str(r["content"]))
                    for i, r in df.iterrows()]
        else:
            docs = [(upload["name"], "Uploaded document", upload["text"])]
    else:
        docs = SYNTH_DOCS

    # Sign at publisher
    signed = [(d, c, t, sign(t.encode(), key)) for d, c, t in docs]
    # Simulate delivery: tamper one, sign one with an impostor key
    impostor = publisher_key(int(bits), int(seed) + 999)
    delivered = []
    for i, (d, c, t, s) in enumerate(signed):
        scenario, content, sig = "Delivered intact", t, s
        if not upload and i == 1:
            scenario, content = "Content tampered in transit", t.replace("3.7.9", "3.7.9-backdoor")
        if not upload and i == 4:
            scenario, sig = "Signed by impostor key", sign(t.encode(), impostor)
        delivered.append((d, c, content, sig, scenario))

    rows = []
    for d, c, content, sig, scenario in delivered:
        ok, rec, h = verify(content.encode(), sig, key.n, key.e)
        rows.append({"document_id": d, "category": c, "scenario": scenario,
                     "sha256": sha256_hex(content.encode())[:24] + "...", "signature s": short(sig, 26),
                     "s^e mod n": short(rec, 22), "H(M) mod n": short(h, 22),
                     "verified": "Yes" if ok else "No"})
    table = pd.DataFrame(rows)
    n_ok = int((table["verified"] == "Yes").sum())
    metrics([("Documents", str(len(table))), ("Verified", str(n_ok)), ("Rejected", str(len(table) - n_ok)),
             ("Key size", f"{key.bits} bits")], status=["", "ok", "bad" if len(table) - n_ok else "ok", ""])
    st.markdown("##### Verification results at the receiving device")
    show_df(table)

    st.markdown("##### Try it: edit a signed document and verify")
    pick = st.selectbox("Document", [d[0] for d in signed], key=f"{SIG_KEY}_pick")
    orig = next(x for x in signed if x[0] == pick)
    edited = st.text_area("Content (edit to see the signature fail)", orig[2], key=f"{SIG_KEY}_edit_{pick}", height=90)
    ok, rec, h = verify(edited.encode(), orig[3], key.n, key.e)
    h_orig = hashlib.sha256(orig[2].encode()).digest()
    h_new = hashlib.sha256(edited.encode()).digest()
    flipped = sum(bin(a ^ b).count("1") for a, b in zip(h_orig, h_new))
    callout(f"<b>{'Signature valid' if ok else 'Signature invalid'}</b>. Hash bits changed: {flipped} of 256.",
            "ok" if ok else "bad")

    st.markdown("##### Forgery attempt: factor the publisher's n")
    fr = factorize(key.n, time_limit=4.0) if key.bits <= 128 else None
    forged = None
    if fr and fr.complete and len(fr.factors) == 2:
        p, q = list(fr.factors)
        stolen = key_from_primes(p, q, key.e)
        evil = "Malicious update 666 with remote access trojan"
        s_evil = sign(evil.encode(), stolen)
        forged = verify(evil.encode(), s_evil, key.n, key.e)[0]
        callout(f"n factored in {fr.elapsed:.3f} s. The attacker signed <i>{evil}</i> and it "
                f"<b>{'passes' if forged else 'fails'}</b> verification. Small keys make forgery trivial.", "bad")
    else:
        callout(f"With a {key.bits} bit key the attacker cannot factor n, so forged signatures cannot be made. "
                "Choose 128 bits or less to see a successful forgery.", "ok")

    avalanche = []
    for d, c, t in docs[:20]:
        base = hashlib.sha256(t.encode()).digest()
        variant = hashlib.sha256((t + ".").encode()).digest()
        avalanche.append({"document_id": d, "bits changed": sum(bin(a ^ b).count("1") for a, b in zip(base, variant))})
    return {"key": key, "table": table, "avalanche": pd.DataFrame(avalanche), "forged": forged, "edit_ok": ok}


def signatures_flow_dot() -> str:
    return """digraph S {
rankdir=LR; bgcolor="transparent";
node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=11, fillcolor="#E8F0FB", color="#1F4E8C"];
edge [fontname="Helvetica", fontsize=9, color="#4A5B72"];
subgraph cluster_a { label="Publisher (signs)"; style="rounded,dashed"; color="#1F4E8C"; fontname="Helvetica";
  M [label="Document M"]; H [label="h = SHA256(M) mod n"]; S [label="s = h^d mod n", fillcolor="#FDECEC", color="#B23A3A"]; M -> H -> S; }
subgraph cluster_b { label="Receiver (verifies)"; style="rounded,dashed"; color="#1E7B4A"; fontname="Helvetica";
  M2 [label="Received M'"]; H2 [label="h' = SHA256(M') mod n"]; V [label="s^e mod n"]; C [label="Equal ?", shape=diamond, fillcolor="#FFF5EB", color="#E07A1F"];
  OK [label="Accept and install", fillcolor="#D6F0E0", color="#1E7B4A"]; NO [label="Reject", fillcolor="#FDECEC", color="#B23A3A"];
  M2 -> H2 -> C; V -> C; C -> OK [label="yes"]; C -> NO [label="no"]; }
M -> M2 [label="network"]; S -> V [label="signature s"];
}"""


def chain_dot() -> str:
    return """digraph C {
rankdir=TB; bgcolor="transparent";
node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=11, color="#1F4E8C", fillcolor="#E8F0FB"];
R [label="Root CA\\nself signed, in trust store", fillcolor="#0B4F9C", fontcolor="white"];
I [label="Intermediate CA\\nsigned by Root"]; L [label="Code signing certificate\\nsigned by Intermediate"];
F [label="Software package\\nsigned by publisher key", fillcolor="#D6F0E0", color="#1E7B4A"];
R -> I [label=" verify with Root n,e", fontsize=9]; I -> L [label=" verify with Intermediate n,e", fontsize=9];
L -> F [label=" verify with publisher n,e", fontsize=9];
}"""


def signatures_visuals(ctx) -> None:
    c1, c2 = st.columns([1.6, 1])
    with c1:
        st.markdown("##### Sign and verify flow")
        st.graphviz_chart(signatures_flow_dot())
    with c2:
        st.markdown("##### Certificate chain of trust")
        st.graphviz_chart(chain_dot())
    if ctx is None:
        return
    c1, c2 = st.columns(2)
    with c1:
        t = ctx["table"]
        counts = t.groupby(["scenario", "verified"]).size().reset_index(name="count")
        fig = go.Figure()
        for v, color in (("Yes", "#1E7B4A"), ("No", "#B23A3A")):
            sub = counts[counts["verified"] == v]
            fig.add_trace(go.Bar(x=sub["scenario"], y=sub["count"], name=f"Verified {v}", marker_color=color))
        fig.update_layout(barmode="stack")
        st.plotly_chart(style_fig(fig, 340, "Verification outcome by delivery scenario"), key=f"{SIG_KEY}_bar")
    with c2:
        a = ctx["avalanche"]
        fig = go.Figure(go.Bar(x=a["document_id"], y=a["bits changed"], marker_color="#0B4F9C"))
        fig.add_hline(y=128, line_dash="dot", annotation_text="Ideal 128 of 256")
        fig.update_yaxes(range=[0, 256], title="Hash bits changed")
        st.plotly_chart(style_fig(fig, 340, "Avalanche: adding one character to each document"), key=f"{SIG_KEY}_aval")


def signatures_report(ctx):
    if ctx is None:
        return None
    t = ctx["table"]
    rep = Report(use_case="Digital Signatures")
    rep.summary = [("Publisher key size", f"{ctx['key'].bits} bits"), ("Public modulus n", str(ctx["key"].n)),
                   ("Documents checked", str(len(t))), ("Verified", str(int((t['verified'] == 'Yes').sum()))),
                   ("Rejected", str(int((t['verified'] == 'No').sum()))),
                   ("Forgery by factoring", "Succeeded" if ctx["forged"] else "Not possible at this key size")]
    rep.sections = [("Interpretation", "Documents delivered intact verify. A tampered document fails because its "
                     "hash changes, and a signature made with any other key fails because only the publisher's d "
                     "inverts e modulo phi(n).")]
    rep.formulas = [("Hash", "h = SHA256(M) mod n"), ("Sign", "s = h^d mod n"), ("Verify", "s^e mod n == h")]
    rep.tables = [("Verification results", t), ("Avalanche test", ctx["avalanche"])]
    rep.chart = {"kind": "bar", "x": list(ctx["avalanche"]["document_id"]),
                 "y": {"Hash bits changed": list(ctx["avalanche"]["bits changed"])},
                 "title": "Avalanche effect of a one character change", "xlabel": "Document", "ylabel": "Bits of 256"}
    return rep


def signatures_render() -> None:
    page_heading("02", "Digital Signatures", "Used for software updates, code signing, legal e-documents and certificates. RSA signatures prove authenticity with the same factoring principle.")
    t = use_case_tabs()
    with t[2]:
        ctx = signatures_demo()
    with t[0]:
        signatures_overview()
    with t[1]:
        signatures_schema()
    with t[3]:
        signatures_visuals(ctx)
    with t[4]:
        export_panel(signatures_report(ctx), SIG_KEY)


# ============================================================================
# PAGE: banking
# Page 03: Online banking and payments.
# ============================================================================

BANK_KEY = "bank"
MERCHANTS = ["Grocery Mart", "City Transit", "Online Books", "Fuel Station", "Pharmacy Plus",
             "Coffee House", "Electronics Hub", "Utility Power Co", "Airline Travel", "Streaming Service"]


def synthetic_transactions(seed: int, count: int = 16) -> pd.DataFrame:
    rng = random.Random(seed)
    start = datetime(2026, 9, 29, 8, 0)
    rows = []
    for i in range(count):
        rows.append({"transaction_id": f"TXN-{10400 + i}", "card_last4": f"{rng.randint(1000, 9999)}",
                     "merchant": rng.choice(MERCHANTS), "amount": round(rng.uniform(3, 900), 2),
                     "currency": "USD", "timestamp": (start + timedelta(minutes=37 * i)).strftime("%Y-%m-%d %H:%M")})
    return pd.DataFrame(rows)


@st.cache_data(show_spinner=False)
def bank_key(bits: int, seed: int):
    return generate_rsa_key(bits, 65537, random.Random(seed))


def banking_overview() -> None:
    callout("When you open a banking session or pay by card online, the client and the bank must agree on a "
            "secret session key over a public network. In <b>RSA key transport</b> (used by TLS 1.2 and earlier, "
            "and still common in payment hardware security modules) the client picks a random pre master "
            "secret, encrypts it with the bank's RSA public key, and only the bank can decrypt it. Both sides "
            "then derive the same session key, which authenticates every transaction.")
    cards([
        ("Confidential sessions", "Card numbers, balances and credentials travel encrypted under a key only "
         "the client and the bank know."),
        ("Transaction integrity", "Every payment carries a message authentication code (HMAC). Changing the "
         "amount or merchant in transit is detected immediately."),
        ("Replay protection", "Unique transaction IDs and nonces stop an attacker from resubmitting a captured payment."),
        ("Signed receipts", "The bank signs each receipt with its RSA key, giving the customer proof of the transaction."),
        ("Standards compliance", "Supports PCI DSS requirements for strong cryptography over public networks."),
        ("Lesson learned", "If the bank's RSA key is ever factored, recorded sessions can be decrypted. That is "
         "why TLS 1.3 replaced RSA key transport with ephemeral Diffie Hellman (forward secrecy)."),
    ])


def banking_schema() -> None:
    st.markdown("#### Transaction fields")
    field_table([
        {"Field": "transaction_id", "Type": "Text, unique", "Description": "Identifier used for replay detection.", "Example": "TXN-10400"},
        {"Field": "card_last4", "Type": "4 digits", "Description": "Masked card number; the full PAN is never shown.", "Example": "4821"},
        {"Field": "merchant", "Type": "Text", "Description": "Payee name.", "Example": "Grocery Mart"},
        {"Field": "amount", "Type": "Decimal", "Description": "Transaction value.", "Example": "84.20"},
        {"Field": "currency", "Type": "ISO 4217 code", "Description": "Currency of the amount.", "Example": "USD"},
        {"Field": "timestamp", "Type": "Date and time", "Description": "When the payment was authorised.", "Example": "2026-09-29 08:37"},
        {"Field": "hmac", "Type": "64 hex characters", "Description": "HMAC-SHA256 of the canonical message under the session key.", "Example": "3fa1..."},
        {"Field": "receipt signature", "Type": "Integer", "Description": "Bank's RSA signature over the accepted transaction.", "Example": "7712..."},
    ])
    st.markdown("#### Session fields")
    field_table([
        {"Field": "Bank public key (n, e)", "Type": "RSA key", "Description": "From the bank's certificate.", "Example": "512 bit n, e = 65537"},
        {"Field": "Pre master secret PMS", "Type": "Random integer", "Description": "Chosen by the client, below n.", "Example": "128 bits"},
        {"Field": "Client and server nonces", "Type": "128 bit random", "Description": "Fresh per session, prevent reuse of old keys.", "Example": "a91f..."},
        {"Field": "Session key K", "Type": "256 bits", "Description": "Derived from PMS and both nonces.", "Example": "SHA-256 output"},
    ])
    st.markdown("#### Formulas")
    formula("Key transport (client)", r"C = \mathrm{PMS}^{\,e} \bmod n", "Encrypted pre master secret sent to the bank.")
    formula("Key recovery (bank)", r"\mathrm{PMS} = C^{\,d} \bmod n", "Only the bank's private exponent d recovers it.")
    formula("Session key derivation", r"K = \mathrm{SHA256}(\mathrm{PMS} \,\|\, N_c \,\|\, N_s)",
            "Both sides compute the same K from the shared PMS and the public nonces.")
    formula("Transaction authentication", r"\tau = \mathrm{HMAC}_{K}(\text{id}\,|\,\text{card}\,|\,\text{merchant}\,|\,\text{amount}\,|\,\text{currency}\,|\,\text{time})",
            "The bank recomputes the tag; any change to a field produces a different tag.")
    formula("Signed receipt", r"s = \mathrm{SHA256}(\text{receipt})^{d} \bmod n", "Customer verifies with the bank's public key.")
    formula("Retrospective attack", r"n \to (p,q) \to d \to \mathrm{PMS} = C^d \bmod n \to K",
            "Factoring the bank key later exposes every recorded session: the motivation for forward secrecy.")


def banking_demo():
    upload = data_source(BANK_KEY, "CSV with columns <b>transaction_id, card_last4, merchant, amount, currency, "
                              "timestamp</b>. Missing optional columns are filled with defaults.",
                         synthetic_transactions(1, 3))
    c1, c2, c3 = st.columns(3)
    bits = c1.select_slider("Bank RSA key size (bits)", options=[64, 96, 128, 256, 512, 1024, 2048], value=512, key=f"{BANK_KEY}_bits")
    seed = c2.number_input("Session seed", 0, 10_000, 11, key=f"{BANK_KEY}_seed")
    attack_on = c3.checkbox("Simulate attacks in transit", value=True, key=f"{BANK_KEY}_attack",
                            help="Tampers with one amount and replays one transaction.")
    key = bank_key(int(bits), int(seed))
    rng = random.Random(int(seed))
    hs = banking_handshake(key, rng)

    st.markdown("##### Step 1: session handshake (RSA key transport)")
    hs_df = pd.DataFrame([
        {"Step": "Client nonce Nc", "Value": hs["client_nonce"]},
        {"Step": "Server nonce Ns", "Value": hs["server_nonce"]},
        {"Step": "Pre master secret PMS (client secret)", "Value": str(hs["pre_master"])},
        {"Step": "C = PMS^e mod n (sent over network)", "Value": str(hs["encrypted_pre_master"])},
        {"Step": "PMS = C^d mod n (recovered by bank)", "Value": str(hs["recovered_pre_master"])},
        {"Step": "Client session key K", "Value": hs["client_session_key"]},
        {"Step": "Bank session key K", "Value": hs["bank_session_key"]},
    ])
    show_df(hs_df)
    callout(f"Session keys match: <b>{'Yes' if hs['keys_match'] else 'No'}</b>", "ok" if hs["keys_match"] else "bad")

    if upload and upload["df"] is not None and "amount" in upload["df"].columns:
        tx = upload["df"].copy()
        defaults = {"transaction_id": None, "card_last4": "0000", "merchant": "Unknown", "currency": "USD",
                    "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M")}
        for col, dv in defaults.items():
            if col not in tx.columns:
                tx[col] = [f"TXN-U{i+1}" for i in range(len(tx))] if dv is None else dv
        tx["amount"] = pd.to_numeric(tx["amount"].astype(str).str.replace(",", ""), errors="coerce").fillna(0).round(2)
        tx = tx[["transaction_id", "card_last4", "merchant", "amount", "currency", "timestamp"]]
    else:
        if upload:
            st.warning("The uploaded file has no amount column, so synthetic transactions are shown.")
        tx = synthetic_transactions(int(seed))

    K = hs["session_key_bytes"]
    # Client side: compute tags
    sent = []
    for _, r in tx.iterrows():
        row = r.to_dict()
        sent.append({**row, "hmac": mac(K, transaction_message(row))})
    # Network: attacks
    received = [dict(s) for s in sent]
    if attack_on and len(received) >= 3:
        received[2]["amount"] = round(float(received[2]["amount"]) * 10, 2)
        received[2]["_note"] = "Amount altered in transit"
        replay = dict(received[0])
        replay["_note"] = "Replayed copy of an earlier payment"
        received.insert(5 if len(received) > 5 else len(received), replay)
    # Bank side
    seen, results = set(), []
    for r in received:
        tag_ok = mac(K, transaction_message(r)) == r["hmac"]
        dup = r["transaction_id"] in seen
        seen.add(r["transaction_id"])
        if not tag_ok:
            status, reason = "Rejected", "HMAC mismatch: message altered"
        elif dup:
            status, reason = "Rejected", "Replay: duplicate transaction id"
        else:
            status, reason = "Approved", "HMAC valid"
        receipt = f"{transaction_message(r)}|{status}"
        s = sign(receipt.encode(), key)
        rv = verify(receipt.encode(), s, key.n, key.e)[0]
        results.append({"transaction_id": r["transaction_id"], "card_last4": r["card_last4"], "merchant": r["merchant"],
                        "amount": float(r["amount"]), "currency": r["currency"], "timestamp": r["timestamp"],
                        "hmac": r["hmac"][:16] + "...", "status": status, "reason": reason,
                        "network event": r.get("_note", "Normal delivery"),
                        "receipt signature": short(s, 20), "receipt verifies": "Yes" if rv else "No"})
    res = pd.DataFrame(results)
    appr = res[res["status"] == "Approved"]
    metrics([("Transactions received", str(len(res))), ("Approved", str(len(appr))),
             ("Rejected", str(len(res) - len(appr))), ("Approved value", f"{appr['amount'].sum():,.2f}"),
             ("Key size", f"{key.bits} bits")], status=["", "ok", "bad" if len(res) - len(appr) else "ok", "", ""])
    st.markdown("##### Step 2: transactions verified by the bank")
    show_df(res)

    st.markdown("##### Step 3: retrospective attack on a recorded session")
    fr = factorize(key.n, time_limit=4.0) if key.bits <= 128 else None
    attack = None
    if fr and fr.complete and len(fr.factors) == 2:
        p, q = list(fr.factors)
        stolen = key_from_primes(p, q, key.e)
        pms = pow(hs["encrypted_pre_master"], stolen.d, key.n)
        k2 = derive_session_key(pms, bytes.fromhex(hs["client_nonce"]), bytes.fromhex(hs["server_nonce"]))
        attack = k2 == K
        callout(f"The eavesdropper recorded C and both nonces, then factored n in {fr.elapsed:.3f} s. The "
                f"recovered session key {'matches' if attack else 'does not match'} the real one, so every "
                "transaction in the session can be read and forged.", "bad")
    else:
        callout(f"With a {key.bits} bit bank key the recorded session stays protected because n cannot be "
                "factored. Choose 128 bits or less to see the retrospective attack succeed.", "ok")
    return {"key": key, "hs": hs_df, "res": res, "attack": attack}


def sequence_dot() -> str:
    return """digraph B {
rankdir=LR; bgcolor="transparent"; nodesep=0.3; ranksep=0.9;
node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=10.5, fillcolor="#E8F0FB", color="#1F4E8C"];
edge [fontname="Helvetica", fontsize=9, color="#4A5B72"];
C1 [label="1. Client hello\\nnonce Nc"]; B1 [label="2. Server hello\\nnonce Ns + certificate (n, e)"];
C2 [label="3. Verify certificate\\npick random PMS"]; C3 [label="4. Send C = PMS^e mod n"];
B2 [label="5. PMS = C^d mod n", fillcolor="#FDECEC", color="#B23A3A"];
K [label="6. Both derive\\nK = SHA256(PMS|Nc|Ns)", fillcolor="#D6F0E0", color="#1E7B4A"];
T [label="7. Payment + HMAC_K"]; V [label="8. Bank checks HMAC\\nand replay list"]; R [label="9. Signed receipt\\ns = h^d mod n"];
C1 -> B1 -> C2 -> C3 -> B2 -> K -> T -> V -> R;
}"""


def banking_visuals(ctx) -> None:
    st.markdown("##### Session protocol sequence")
    st.graphviz_chart(sequence_dot())
    if ctx is None:
        return
    res = ctx["res"]
    c1, c2 = st.columns(2)
    with c1:
        fig = px.bar(res, x="transaction_id", y="amount", color="status",
                     color_discrete_map={"Approved": "#1E7B4A", "Rejected": "#B23A3A"},
                     hover_data=["merchant", "reason"])
        fig.update_xaxes(title="Transaction", tickangle=-45)
        fig.update_yaxes(title="Amount")
        st.plotly_chart(style_fig(fig, 360, "Transaction amounts and bank decision"), key=f"{BANK_KEY}_bar")
    with c2:
        by_m = res[res["status"] == "Approved"].groupby("merchant")["amount"].sum().reset_index()
        fig = go.Figure(go.Pie(labels=by_m["merchant"], values=by_m["amount"], hole=0.55,
                               marker=dict(colors=PLOT_COLORS * 2)))
        st.plotly_chart(style_fig(fig, 360, "Approved value by merchant"), key=f"{BANK_KEY}_pie")


def banking_report(ctx):
    if ctx is None:
        return None
    res = ctx["res"]
    rep = Report(use_case="Online Banking and Payments")
    appr = res[res["status"] == "Approved"]
    rep.summary = [("Bank key size", f"{ctx['key'].bits} bits"), ("Transactions received", str(len(res))),
                   ("Approved", str(len(appr))), ("Rejected", str(len(res) - len(appr))),
                   ("Approved value", f"{appr['amount'].sum():,.2f}"),
                   ("Retrospective factoring attack", "Succeeded" if ctx["attack"] else "Not possible at this key size")]
    rep.sections = [("Interpretation", "RSA key transport gives both parties the same session key. Tampered and "
                     "replayed payments are rejected. If the bank key could be factored, recorded sessions would "
                     "be exposed, which is why modern TLS uses ephemeral key exchange.")]
    rep.formulas = [("Key transport", "C = PMS^e mod n; PMS = C^d mod n"),
                    ("Session key", "K = SHA256(PMS | Nc | Ns)"), ("Transaction tag", "tag = HMAC-SHA256(K, message)"),
                    ("Receipt signature", "s = SHA256(receipt)^d mod n")]
    rep.tables = [("Transaction decisions", res), ("Handshake", ctx["hs"])]
    rep.chart = {"kind": "bar", "x": list(res["transaction_id"]), "y": {"Amount": list(res["amount"])},
                 "title": "Transaction amounts", "xlabel": "Transaction", "ylabel": "Amount"}
    return rep


def banking_render() -> None:
    page_heading("03", "Online Banking and Payments", "Card transactions and banking sessions rely on protocols that historically used RSA style key exchange.")
    t = use_case_tabs()
    with t[2]:
        ctx = banking_demo()
    with t[0]:
        banking_overview()
    with t[1]:
        banking_schema()
    with t[3]:
        banking_visuals(ctx)
    with t[4]:
        export_panel(banking_report(ctx), BANK_KEY)


# ============================================================================
# PAGE: quantum
# Page 04: Quantum computing research (Shor's algorithm).
# ============================================================================

QC_KEY = "qc"


@st.cache_data(show_spinner="Simulating the quantum register...", max_entries=64)
def cached_qft(n: int, a: int):
    return qft_distribution(n, a)


def quantum_overview() -> None:
    callout("<b>Shor's algorithm</b> (1994) factors n in polynomial time on a large, error corrected quantum "
            "computer. It turns factoring into finding the period r of f(x) = a<sup>x</sup> mod n. A quantum "
            "Fourier transform makes the period visible as sharp peaks, then ordinary arithmetic finishes the "
            "job. This page simulates each step exactly for small n.")
    cards([
        ("Understand the threat", "RSA, and the signatures and banking protocols built on it, would be broken "
         "by a large enough quantum computer."),
        ("Plan the migration", "NIST published post quantum standards in August 2024: FIPS 203 (ML-KEM), "
         "FIPS 204 (ML-DSA) and FIPS 205 (SLH-DSA). Draft NIST IR 8547 proposes deprecating RSA by 2030 and "
         "disallowing it by 2035."),
        ("Harvest now, decrypt later", "Data captured today can be decrypted once quantum machines exist, so "
         "long lived secrets need protection now."),
        ("Resource estimates", "Published estimates (for example Gidney, 2025) put RSA-2048 within reach of "
         "under one million noisy qubits running for about a week; today's machines are far smaller."),
        ("Research and education", "Simulating the algorithm shows exactly which step needs quantum hardware "
         "and which steps are classical."),
        ("Crypto agility", "Inventory RSA usage and design systems that can swap algorithms without redesign."),
    ])


def quantum_schema() -> None:
    st.markdown("#### Fields")
    field_table([
        {"Field": "N", "Type": "Odd composite", "Description": "Number to factor. Simulation supports N up to 255 (16 qubit register).", "Example": "15, 21, 35, 91"},
        {"Field": "a", "Type": "Random base", "Description": "Chosen with 1 < a < N; if gcd(a, N) > 1 a factor is found by luck.", "Example": "7"},
        {"Field": "r", "Type": "Period (order)", "Description": "Smallest r > 0 with a^r = 1 mod N. Finding r is the quantum step.", "Example": "4"},
        {"Field": "q, Q", "Type": "Register size", "Description": "q qubits with N^2 <= Q = 2^q < 2 N^2.", "Example": "q = 8, Q = 256"},
        {"Field": "y", "Type": "Measurement", "Description": "Outcome of measuring the register after the QFT; close to a multiple of Q / r.", "Example": "64, 128, 192"},
        {"Field": "P(y)", "Type": "Probability", "Description": "Chance of observing y, computed exactly by simulation.", "Example": "0.25"},
        {"Field": "Upload", "Type": "CSV, TXT, DOCX, PDF", "Description": "CSV column N, or any document; each number is factored with the Shor procedure.", "Example": "N\\n15\\n77"},
    ])
    st.markdown("#### Formulas")
    formula("Periodic function", r"f(x) = a^{x} \bmod N,\qquad f(x + r) = f(x)",
            "The period r of this function holds the key to the factors.")
    formula("Quantum Fourier transform", r"|x\rangle \;\mapsto\; \frac{1}{\sqrt{Q}}\sum_{y=0}^{Q-1} e^{2\pi i\,x y / Q}\,|y\rangle",
            "Applied to a superposition of all x with the same f(x), it concentrates probability near y = k Q / r.")
    formula("Measurement probability", r"P(y)=\frac{1}{Q^{2}}\sum_{v}\Big|\sum_{x:\,f(x)=v} e^{2\pi i\,x y/Q}\Big|^{2}",
            "Exact formula evaluated with fast Fourier transforms in this simulation.")
    formula("Recover the period", r"\frac{y}{Q} \approx \frac{k}{r}\quad\Rightarrow\quad r = \text{denominator of a continued fraction convergent}",
            "Convergents with denominator below N are candidate periods.")
    formula("Extract the factors", r"p = \gcd\!\big(a^{r/2}-1,\;N\big),\qquad q = \gcd\!\big(a^{r/2}+1,\;N\big)",
            "Works when r is even and a<sup>r/2</sup> is not -1 mod N (probability at least 1/2 per random a).")
    formula("Cost comparison", r"\text{Classical: } e^{1.923\,(\ln N)^{1/3}(\ln\ln N)^{2/3}}\qquad \text{Quantum: } O\big((\log N)^{3}\big)",
            "Sub exponential versus polynomial: the reason quantum computers threaten RSA.")


def quantum_demo():
    upload = data_source(QC_KEY, "CSV with a column named <b>N</b>, or a TXT, Word or PDF file listing numbers. "
                              "Each composite is factored with the simulated Shor procedure (N up to 1,000,000).",
                         pd.DataFrame({"N": [15, 21, 35, 77, 91, 143]}))
    c1, c2 = st.columns(2)
    n = int(c1.selectbox("N to factor (simulated register)", [15, 21, 33, 35, 39, 51, 55, 57, 65, 69, 77, 85, 87,
                                                              91, 95, 111, 115, 119, 133, 143, 155, 161, 187, 203, 221, 247],
                         index=0, key=f"{QC_KEY}_n"))
    coprimes = [a for a in range(2, n - 1) if math.gcd(a, n) == 1]
    a = int(c2.selectbox("Base a (coprime with N)", coprimes, index=coprimes.index(7) if 7 in coprimes else 0,
                         key=f"{QC_KEY}_a"))
    r = multiplicative_order(a, n)
    sim = cached_qft(n, a)
    Q, q = sim["Q"], sim["q_bits"]
    probs = sim["probabilities"]
    peaks = np.argsort(probs)[::-1][:8]
    peak_rows = []
    for y in sorted(int(v) for v in peaks):
        cands = continued_fraction_period(y, Q, n)
        good = [c["Denominator (candidate r)"] for c in cands if c["Denominator (candidate r)"] > 0
                and pow(a, c["Denominator (candidate r)"], n) == 1]
        peak_rows.append({"measured y": y, "P(y)": round(float(probs[y]), 5), "y / Q": f"{y}/{Q}",
                          "k Q / r nearest": round(y * r / Q, 3),
                          "convergents": ", ".join(c["Convergent"] for c in cands) or "none (y = 0)",
                          "period recovered": str(min(good)) if good else "no"})
    peaks_df = pd.DataFrame(peak_rows)
    metrics([("N", str(n)), ("Base a", str(a)), ("Period r", str(r)), ("Qubits q", str(q)), ("Q = 2^q", str(Q)),
             ("Distinct peaks", str(int((probs > 1e-3).sum())))])
    run = shor_factor(n, a=a, rng=random.Random(n * 31 + a))
    if run["success"]:
        f1, f2 = run["factors"]
        callout(f"<b>Factors found: {n} = {min(f1, f2)} x {max(f1, f2)}</b>"
                + (f" using a = {run['a']}, r = {run['r']}" if run.get("r") else ""), "ok")
    else:
        callout("No factor found in the allowed attempts.", "warn")
    st.markdown("##### Algorithm log")
    show_df(pd.DataFrame(run["log"]))
    st.markdown("##### Most likely measurements and period recovery")
    show_df(peaks_df)
    if r % 2 == 1 or pow(a, r // 2, n) == n - 1:
        callout(f"With a = {a} the period r = {r} is {'odd' if r % 2 else 'such that a^(r/2) = -1 mod N'}, so "
                "this base fails and the algorithm retried with another base, as shown in the log.", "warn")

    # Batch
    if upload:
        if upload["df"] is not None and "n" in upload["df"].columns:
            nums = [int(float(v)) for v in upload["df"]["n"] if str(v).strip()]
        else:
            nums = extract_integers(upload["text"])
        title = f"Batch results from {upload['name']}"
    else:
        nums = [15, 21, 35, 77, 91, 143, 221, 323, 437, 667, 899, 1147, 3599, 10403]
        title = "Batch results for synthetic semiprimes"
    batch = []
    for N in nums[:60]:
        if N < 4 or N > 1_000_000:
            batch.append({"N": N, "result": "skipped (outside 4 to 1,000,000)"})
            continue
        rr = shor_factor(N, rng=random.Random(N))
        att = sum(1 for x in rr["log"] if x["Step"].startswith("Attempt"))
        batch.append({"N": N, "qubits needed q": register_size(N),
                      "result": f"{min(rr['factors'])} x {max(rr['factors'])}" if rr["success"] else
                      ("prime" if is_probable_prime(N) else "not found"),
                      "base a": rr.get("a"), "period r": rr.get("r"), "attempts": att})
    batch_df = pd.DataFrame(batch)
    st.markdown(f"##### {title}")
    show_df(batch_df)
    return {"n": n, "a": a, "r": r, "Q": Q, "q": q, "probs": probs, "peaks": peaks_df, "batch": batch_df, "run": run}


def quantum_flow_dot() -> str:
    return """digraph Q {
rankdir=LR; bgcolor="transparent";
node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=10.5, fillcolor="#E8F0FB", color="#1F4E8C"];
edge [color="#4A5B72"];
A [label="Pick random a\\ncheck gcd(a, N)"];
subgraph cluster_q { label="Quantum computer"; style="rounded,filled"; fillcolor="#F3ECFA"; color="#8E44AD"; fontname="Helvetica";
  H [label="Hadamard on q qubits\\nsuperposition of all x"]; U [label="Modular exponentiation\\n|x>|a^x mod N>"];
  F [label="Quantum Fourier\\ntransform"]; M [label="Measure y"]; H -> U -> F -> M; }
CF [label="Continued fractions\\ny/Q ~ k/r"]; G [label="gcd(a^(r/2) +- 1, N)", fillcolor="#D6F0E0", color="#1E7B4A"];
A -> H; M -> CF -> G;
}"""


def quantum_visuals(ctx) -> None:
    st.markdown("##### Shor's algorithm: classical and quantum parts")
    st.graphviz_chart(quantum_flow_dot())
    if ctx is None:
        return
    n, a, r, Q = ctx["n"], ctx["a"], ctx["r"], ctx["Q"]
    c1, c2 = st.columns(2)
    with c1:
        L = min(4 * r + 4, 64)
        seq = period_sequence(a, n, L)
        fig = go.Figure(go.Scatter(x=list(range(L)), y=seq, mode="lines+markers", line=dict(shape="linear")))
        for k in range(0, L, r):
            fig.add_vline(x=k, line_dash="dot", line_color="#E07A1F")
        fig.update_xaxes(title="x")
        fig.update_yaxes(title=f"{a}^x mod {n}")
        st.plotly_chart(style_fig(fig, 340, f"Periodic function, period r = {r}"), key=f"{QC_KEY}_period")
    with c2:
        probs = ctx["probs"]
        fig = go.Figure(go.Bar(x=list(range(Q)), y=probs, marker_color="#8E44AD"))
        fig.update_xaxes(title="Measured value y")
        fig.update_yaxes(title="Probability P(y)")
        st.plotly_chart(style_fig(fig, 340, f"Measurement distribution after QFT (Q = {Q})"), key=f"{QC_KEY}_qft")
    st.caption(f"Peaks sit at multiples of Q / r = {Q / r:.2f}. Each peak reveals the period through continued fractions.")

    st.markdown("##### Classical versus quantum cost to factor an RSA modulus")
    sizes = [256, 512, 768, 1024, 1536, 2048, 3072, 4096]
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=sizes, y=[gnfs_log10_ops(b) for b in sizes], mode="lines+markers",
                             name="Classical GNFS (log10 operations)"))
    fig.add_trace(go.Scatter(x=sizes, y=[shor_log10_ops(b) for b in sizes], mode="lines+markers",
                             name="Shor's algorithm (log10 logical gates, order n^3)"))
    fig.update_xaxes(title="RSA modulus size (bits)")
    fig.update_yaxes(title="log10 of work")
    st.plotly_chart(style_fig(fig, 360), key=f"{QC_KEY}_cost")
    show_df(pd.DataFrame([{"RSA size (bits)": b, "Classical GNFS work": f"10^{gnfs_log10_ops(b):.1f}",
                           "Shor gates (order)": f"10^{shor_log10_ops(b):.1f}",
                           "Logical qubits (2n + 3)": shor_logical_qubits(b)} for b in [1024, 2048, 3072, 4096]]))
    st.caption("Order of magnitude estimates for illustration. Physical qubit counts are much higher once "
               "error correction is included.")


def quantum_report(ctx):
    if ctx is None:
        return None
    rep = Report(use_case="Quantum Computing Research (Shor's Algorithm)")
    run = ctx["run"]
    rep.summary = [("N", str(ctx["n"])), ("Base a", str(ctx["a"])), ("Period r", str(ctx["r"])),
                   ("Qubits q", str(ctx["q"])), ("Q", str(ctx["Q"])),
                   ("Factors", " x ".join(str(f) for f in sorted(run["factors"])) if run["success"] else "not found")]
    rep.sections = [("Interpretation", "The quantum Fourier transform concentrates probability at multiples of "
                     "Q / r. Continued fractions recover r and two GCDs give the factors. The classical cost of "
                     "factoring grows sub exponentially while the quantum cost grows polynomially, which is "
                     "why organisations are migrating to post quantum cryptography.")]
    rep.formulas = [("Periodic function", "f(x) = a^x mod N"), ("QFT peaks", "y ~ k Q / r"),
                    ("Factors", "gcd(a^(r/2) - 1, N) and gcd(a^(r/2) + 1, N)")]
    rep.tables = [("Batch Shor results", ctx["batch"]), ("Measurement peaks", ctx["peaks"])]
    sizes = [256, 512, 1024, 2048, 3072, 4096]
    rep.chart = {"kind": "line", "x": sizes, "y": {"Classical GNFS": [gnfs_log10_ops(b) for b in sizes],
                                                   "Shor (order n^3)": [shor_log10_ops(b) for b in sizes]},
                 "title": "Classical versus quantum factoring cost", "xlabel": "RSA bits", "ylabel": "log10 work"}
    return rep


def quantum_render() -> None:
    page_heading("04", "Quantum Computing Research", "Shor's algorithm could factor large numbers efficiently on a large quantum computer, driving the global shift to post quantum cryptography.")
    t = use_case_tabs()
    with t[2]:
        ctx = quantum_demo()
    with t[0]:
        quantum_overview()
    with t[1]:
        quantum_schema()
    with t[3]:
        quantum_visuals(ctx)
    with t[4]:
        export_panel(quantum_report(ctx), QC_KEY)


# ============================================================================
# PAGE: fractions
# Page 05: Simplifying fractions, GCD and LCM (scheduling and gear design).
# ============================================================================

FRAC_KEY = "frac"

SYNTH_FRACTIONS = [(84, 126), (360, 1024), (1071, 462), (221, 323), (4096, 6144), (97, 89), (1001, 1430), (2310, 4620)]
SYNTH_TASKS = [("Oil change", 12), ("Filter replacement", 18), ("Safety inspection", 30), ("Calibration", 45)]
SYNTH_GEARS = [(12, 36), (17, 43), (20, 48), (24, 60), (31, 47)]


def fac(n: int) -> str:
    r = factorize(n, time_limit=2)
    return format_factorization(r.factors) if n > 1 else "1"


def fractions_overview() -> None:
    callout("The <b>greatest common divisor</b> (GCD) and <b>least common multiple</b> (LCM) of numbers can be "
            "read directly from their prime factorizations: take the lowest power of each shared prime for the "
            "GCD and the highest power of every prime for the LCM. These two quantities solve everyday "
            "problems in mathematics, operations planning and mechanical engineering.")
    cards([
        ("Simplifying fractions", "Divide numerator and denominator by their GCD to get the simplest form, "
         "used in measurements, recipes, ratios and exact computation."),
        ("Scheduling", "Repeating cycles (maintenance, shifts, bus routes, billing periods) all coincide again "
         "after the LCM of their intervals."),
        ("Gear ratio design", "The GCD of tooth counts sets how often the same teeth meet. Coprime counts give a "
         "hunting tooth design that spreads wear evenly and extends gear life."),
        ("Euclid's algorithm", "Computes the GCD without factoring, in a number of steps proportional to the "
         "number of digits, and also yields modular inverses used by RSA."),
        ("Common denominators", "Adding fractions needs the LCM of denominators, the smallest shared unit."),
        ("Data quality", "Reducing ratios to lowest terms makes datasets comparable and avoids overflow."),
    ])


def fractions_schema() -> None:
    st.markdown("#### Fields")
    field_table([
        {"Field": "numerator, denominator", "Type": "Positive integers", "Description": "Fraction a / b to simplify.", "Example": "84, 126"},
        {"Field": "gcd", "Type": "Integer", "Description": "Largest integer dividing both numbers.", "Example": "42"},
        {"Field": "simplified", "Type": "Fraction", "Description": "(a / gcd) / (b / gcd), in lowest terms.", "Example": "2/3"},
        {"Field": "task, interval", "Type": "Text, days", "Description": "Recurring activity and how often it repeats.", "Example": "Oil change, 12"},
        {"Field": "alignment period", "Type": "Days", "Description": "LCM of all intervals: when every task falls on the same day again.", "Example": "180"},
        {"Field": "teeth_a, teeth_b", "Type": "Integers", "Description": "Tooth counts of the driving and driven gear.", "Example": "17, 43"},
        {"Field": "gear ratio", "Type": "Fraction", "Description": "teeth_b / teeth_a in lowest terms: output turns per input turn inverted.", "Example": "43:17"},
        {"Field": "hunting tooth", "Type": "Yes or No", "Description": "Yes when gcd(teeth_a, teeth_b) = 1, so every tooth meets every other tooth.", "Example": "Yes"},
        {"Field": "Upload CSV", "Type": "Columns", "Description": "numerator, denominator; or task, interval; or teeth_a, teeth_b. TXT, Word or PDF numbers are paired as fractions.", "Example": "numerator,denominator"},
    ])
    st.markdown("#### Formulas")
    formula("GCD and LCM from factorizations",
            r"a=\prod p_i^{\alpha_i},\ b=\prod p_i^{\beta_i}\ \Rightarrow\ \gcd(a,b)=\prod p_i^{\min(\alpha_i,\beta_i)},\ \ \mathrm{lcm}(a,b)=\prod p_i^{\max(\alpha_i,\beta_i)}",
            "Lowest shared powers for the GCD, highest powers for the LCM.")
    formula("Product identity", r"\gcd(a,b)\times\mathrm{lcm}(a,b)=a\times b", "Lets the LCM be computed from the GCD.")
    formula("Euclidean algorithm", r"\gcd(a,b)=\gcd(b,\ a \bmod b),\qquad \gcd(a,0)=a",
            "Repeated division with remainder; each row of the table is a = q b + r.")
    formula("Simplified fraction", r"\frac{a}{b}=\frac{a/g}{b/g},\quad g=\gcd(a,b)", "The result is in lowest terms.")
    formula("Schedule alignment", r"T=\mathrm{lcm}(t_1,t_2,\ldots,t_k)", "All tasks coincide every T days.")
    formula("Gear mesh repeat", r"\text{turns of gear A before the same teeth meet} = \frac{\mathrm{lcm}(T_A,T_B)}{T_A}=\frac{T_B}{\gcd(T_A,T_B)}",
            "Larger values spread wear across more tooth pairs; the maximum T<sub>B</sub> occurs when the counts are coprime.")


def fractions_demo():
    upload = data_source(FRAC_KEY, "CSV with columns <b>numerator, denominator</b>, or <b>task, interval</b>, or "
                              "<b>teeth_a, teeth_b</b>. In TXT, Word or PDF files, numbers are read in pairs as fractions.",
                         pd.DataFrame({"numerator": [84, 360], "denominator": [126, 1024]}))
    fractions, tasks, gears = SYNTH_FRACTIONS, SYNTH_TASKS, SYNTH_GEARS
    if upload:
        df = upload["df"]
        if df is not None and {"numerator", "denominator"} <= set(df.columns):
            fractions = [(int(float(a)), int(float(b))) for a, b in zip(df["numerator"], df["denominator"]) if str(a).strip()]
        elif df is not None and {"task", "interval"} <= set(df.columns):
            tasks = [(str(t), int(float(i))) for t, i in zip(df["task"], df["interval"]) if str(i).strip()]
        elif df is not None and {"teeth_a", "teeth_b"} <= set(df.columns):
            gears = [(int(float(a)), int(float(b))) for a, b in zip(df["teeth_a"], df["teeth_b"]) if str(a).strip()]
        else:
            nums = extract_integers(upload["text"])
            fractions = [(nums[i], nums[i + 1]) for i in range(0, len(nums) - 1, 2)] or SYNTH_FRACTIONS

    # A. Fractions
    st.markdown("##### A. Simplify fractions")
    c1, c2 = st.columns(2)
    a = int(c1.number_input("Numerator a", 1, 10 ** 12, 1071, key=f"{FRAC_KEY}_a"))
    b = int(c2.number_input("Denominator b", 1, 10 ** 12, 462, key=f"{FRAC_KEY}_b"))
    g = math.gcd(a, b)
    metrics([("a", f"{a} = {fac(a)}"), ("b", f"{b} = {fac(b)}"), ("gcd(a, b)", str(g)), ("lcm(a, b)", str(lcm(a, b))),
             ("Simplified", f"{a // g} / {b // g}")])
    st.latex(rf"\frac{{{a}}}{{{b}}} = \frac{{{a} \div {g}}}{{{b} \div {g}}} = \frac{{{a // g}}}{{{b // g}}}")
    steps = pd.DataFrame(euclid_steps(a, b))
    st.markdown("Euclidean algorithm steps (a = q b + r)")
    show_df(steps)
    frac_rows = []
    for x, y in fractions[:500]:
        if x <= 0 or y <= 0:
            continue
        gg = math.gcd(x, y)
        frac_rows.append({"numerator": x, "denominator": y, "numerator factors": fac(x), "denominator factors": fac(y),
                          "gcd": gg, "lcm": lcm(x, y), "simplified": f"{x // gg}/{y // gg}",
                          "decimal": round(x / y, 6), "already lowest terms": "Yes" if gg == 1 else "No"})
    frac_df = pd.DataFrame(frac_rows)
    st.markdown("Fraction batch")
    show_df(frac_df)

    # B. Scheduling
    st.markdown("##### B. Scheduling: when do repeating cycles line up?")
    task_df = pd.DataFrame(tasks, columns=["task", "interval"])
    task_df = st.data_editor(task_df, num_rows="dynamic", key=f"{FRAC_KEY}_tasks", hide_index=True, width="stretch")
    task_df = task_df.dropna()
    intervals = [int(v) for v in task_df["interval"] if int(v) > 0]
    T = lcm_list(intervals) if intervals else 0
    metrics([("Tasks", str(len(intervals))), ("Alignment period (LCM)", f"{T} days"),
             ("GCD of intervals", str(gcd_list(intervals)) if intervals else "0"),
             ("LCM factorization", fac(T) if T > 1 else str(T))])
    sched_rows = []
    for t, i in zip(task_df["task"], task_df["interval"]):
        i = int(i)
        if i > 0 and T:
            sched_rows.append({"task": t, "interval (days)": i, "interval factors": fac(i),
                               "occurrences per alignment period": T // i})
    sched_df = pd.DataFrame(sched_rows)
    show_df(sched_df)

    # C. Gears
    st.markdown("##### C. Gear ratio design")
    c1, c2 = st.columns(2)
    ta = int(c1.number_input("Teeth on driving gear A", 5, 500, 17, key=f"{FRAC_KEY}_ta"))
    tb = int(c2.number_input("Teeth on driven gear B", 5, 500, 43, key=f"{FRAC_KEY}_tb"))
    gg = math.gcd(ta, tb)
    turns = tb // gg
    metrics([("Gear ratio", f"{tb // gg} : {ta // gg}"), ("gcd", str(gg)), ("Turns of A to repeat", str(turns)),
             ("Distinct tooth pairings", str(ta * tb // gg)), ("Hunting tooth", "Yes" if gg == 1 else "No")],
            status=["", "", "", "", "ok" if gg == 1 else "warn"])
    if gg > 1:
        callout(f"Because gcd({ta}, {tb}) = {gg}, each tooth on A only ever meets {tb // gg} of the {tb} teeth on B. "
                f"Changing either count by one tooth to make them coprime spreads wear evenly.", "warn")
    gear_rows = []
    for x, y in gears:
        g2 = math.gcd(x, y)
        gear_rows.append({"teeth_a": x, "teeth_b": y, "gcd": g2, "ratio": f"{y // g2}:{x // g2}",
                          "turns of A before repeat": y // g2, "hunting tooth": "Yes" if g2 == 1 else "No",
                          "wear pairs used (percent)": round(100 / g2, 1)})
    gear_df = pd.DataFrame(gear_rows)
    show_df(gear_df)
    return {"a": a, "b": b, "g": g, "steps": steps, "frac": frac_df, "tasks": task_df, "T": T, "sched": sched_df,
            "ta": ta, "tb": tb, "gears": gear_df}


def fractions_visuals(ctx) -> None:
    if ctx is None:
        return
    a, b = ctx["a"], ctx["b"]
    fa, fb = factorize(a).factors, factorize(b).factors
    primes = sorted(set(fa) | set(fb))
    st.markdown("##### Prime exponent comparison: GCD takes the minimum, LCM the maximum")
    fig = go.Figure()
    fig.add_trace(go.Bar(x=[str(p) for p in primes], y=[fa.get(p, 0) for p in primes], name=f"a = {a}"))
    fig.add_trace(go.Bar(x=[str(p) for p in primes], y=[fb.get(p, 0) for p in primes], name=f"b = {b}"))
    fig.add_trace(go.Scatter(x=[str(p) for p in primes], y=[min(fa.get(p, 0), fb.get(p, 0)) for p in primes],
                             mode="markers", marker=dict(size=14, symbol="line-ew-open", color="#1E7B4A", line=dict(width=4)),
                             name="GCD exponent (min)"))
    fig.add_trace(go.Scatter(x=[str(p) for p in primes], y=[max(fa.get(p, 0), fb.get(p, 0)) for p in primes],
                             mode="markers", marker=dict(size=14, symbol="line-ew-open", color="#B23A3A", line=dict(width=4)),
                             name="LCM exponent (max)"))
    fig.update_layout(barmode="group")
    fig.update_xaxes(title="Prime", type="category")
    fig.update_yaxes(title="Exponent", dtick=1)
    st.plotly_chart(style_fig(fig, 340), key=f"{FRAC_KEY}_exp")

    st.markdown("##### Schedule timeline: every task coincides at the LCM")
    T = ctx["T"]
    tasks = ctx["tasks"]
    if T and len(tasks):
        horizon = min(T * 2, 2000)
        fig = go.Figure()
        for idx, (t, i) in enumerate(zip(tasks["task"], tasks["interval"])):
            i = int(i)
            days = list(range(0, horizon + 1, i))
            fig.add_trace(go.Scatter(x=days, y=[t] * len(days), mode="markers", name=t,
                                     marker=dict(size=9, color=PLOT_COLORS[idx % len(PLOT_COLORS)])))
        for k in range(0, horizon + 1, T):
            fig.add_vline(x=k, line_color="#B23A3A", line_dash="dash",
                          annotation_text=f"All align: day {k}" if k else "Start", annotation_position="top")
        fig.update_xaxes(title="Day")
        st.plotly_chart(style_fig(fig, 320), key=f"{FRAC_KEY}_sched")

    st.markdown("##### Gear wear: which teeth on gear B does tooth 1 of gear A meet?")
    ta, tb = ctx["ta"], ctx["tb"]
    contacts = sorted({(k * ta) % tb for k in range(tb)})
    fig = go.Figure()
    theta = [360 * j / tb for j in range(tb)]
    fig.add_trace(go.Scatterpolar(r=[1] * tb, theta=theta, mode="markers", marker=dict(size=10, color="#D0D8E4"),
                                  name="Teeth on gear B"))
    fig.add_trace(go.Scatterpolar(r=[1] * len(contacts), theta=[360 * j / tb for j in contacts], mode="markers",
                                  marker=dict(size=12, color="#0B4F9C"), name="Teeth met by tooth 1 of A"))
    fig.update_layout(polar=dict(radialaxis=dict(visible=False, range=[0, 1.2]), angularaxis=dict(visible=False)))
    c1, c2 = st.columns([1, 1])
    with c1:
        st.plotly_chart(style_fig(fig, 360), key=f"{FRAC_KEY}_gear")
    with c2:
        callout(f"Tooth 1 of gear A meets <b>{len(contacts)} of {tb}</b> teeth on gear B "
                f"(= {tb} / gcd({ta}, {tb})). A hunting tooth design (gcd = 1) uses all of them.",
                "ok" if len(contacts) == tb else "warn")
        gd = ctx["gears"]
        fig2 = go.Figure(go.Bar(x=[f"{x}/{y}" for x, y in zip(gd["teeth_a"], gd["teeth_b"])],
                                y=gd["wear pairs used (percent)"], marker_color="#0B4F9C"))
        fig2.update_yaxes(title="Tooth pairs used (percent)", range=[0, 105])
        fig2.update_xaxes(title="Teeth A / Teeth B")
        st.plotly_chart(style_fig(fig2, 260), key=f"{FRAC_KEY}_gearbar")


def fractions_report(ctx):
    if ctx is None:
        return None
    rep = Report(use_case="Simplifying Fractions, GCD and LCM")
    rep.summary = [("Fraction", f"{ctx['a']}/{ctx['b']}"), ("GCD", str(ctx["g"])),
                   ("Simplified", f"{ctx['a'] // ctx['g']}/{ctx['b'] // ctx['g']}"),
                   ("Schedule alignment period (days)", str(ctx["T"])),
                   ("Gear pair", f"{ctx['ta']} and {ctx['tb']} teeth"),
                   ("Hunting tooth", "Yes" if math.gcd(ctx["ta"], ctx["tb"]) == 1 else "No")]
    rep.formulas = [("GCD", "product of p^min(alpha, beta)"), ("LCM", "product of p^max(alpha, beta)"),
                    ("Identity", "gcd(a, b) x lcm(a, b) = a x b"), ("Euclid", "gcd(a, b) = gcd(b, a mod b)"),
                    ("Gear repeat", "turns = T_B / gcd(T_A, T_B)")]
    rep.tables = [("Fraction batch", ctx["frac"]), ("Euclidean steps", ctx["steps"]),
                  ("Schedule", ctx["sched"]), ("Gear designs", ctx["gears"])]
    s = ctx["sched"]
    if len(s):
        rep.chart = {"kind": "bar", "x": list(s["task"]), "y": {"Occurrences per alignment period": list(s["occurrences per alignment period"])},
                     "title": f"Occurrences within the {ctx['T']} day alignment period", "xlabel": "Task", "ylabel": "Count"}
    return rep


def fractions_render() -> None:
    page_heading("05", "Simplifying Fractions, GCD and LCM", "Used every day in mathematics, scheduling (when repeating cycles line up) and gear ratio design.")
    t = use_case_tabs()
    with t[2]:
        ctx = fractions_demo()
    with t[0]:
        fractions_overview()
    with t[1]:
        fractions_schema()
    with t[3]:
        fractions_visuals(ctx)
    with t[4]:
        export_panel(fractions_report(ctx), FRAC_KEY)


# ============================================================================
# PAGE: randomgen
# Page 06: Hashing and random number generation (Blum Blum Shub, Very Smooth Hash).
# ============================================================================

RNG_KEY = "rng"


@st.cache_data(show_spinner=False)
def params(bits: int, seed: int):
    return bbs_parameters(bits, random.Random(seed))


def randomgen_overview() -> None:
    callout("<b>Blum Blum Shub</b> (1986) generates random bits by repeatedly squaring modulo a Blum integer "
            "n = p q with p and q both congruent to 3 mod 4, outputting the lowest bit of each state. Predicting "
            "the next bit is provably as hard as factoring n. The <b>Very Smooth Hash</b> (VSH, 2006) builds a "
            "hash function whose collision resistance also rests on a factoring related problem.")
    cards([
        ("Provable security", "Unlike ad hoc generators, BBS comes with a proof: breaking it would give an "
         "algorithm to factor n."),
        ("Key and token generation", "Suitable for session tokens, one time passwords, nonces and seeds for "
         "other generators where predictability is unacceptable."),
        ("Reproducible audits", "Given the same seed the sequence repeats exactly, so results can be audited, "
         "yet without p and q an observer cannot predict it."),
        ("Factoring based hashing", "VSH shows how a hash can inherit security from factoring; the digest "
         "changes completely when a single input bit changes."),
        ("Statistical validation", "Output is checked with NIST SP 800-22 frequency, block frequency and runs "
         "tests, the same tests used to certify generators."),
        ("Trade off", "BBS is slow (one modular squaring per bit), so in practice it seeds faster generators "
         "or protects high value secrets."),
    ])


def randomgen_schema() -> None:
    st.markdown("#### Generator fields")
    field_table([
        {"Field": "p, q", "Type": "Secret primes", "Description": "Both congruent to 3 mod 4 (Blum primes).", "Example": "11, 23"},
        {"Field": "n", "Type": "Public modulus", "Description": "n = p x q (a Blum integer).", "Example": "253"},
        {"Field": "seed s", "Type": "Integer", "Description": "Secret starting value, coprime with n and not 0 or 1.", "Example": "3"},
        {"Field": "x_i", "Type": "State", "Description": "x_0 = s^2 mod n, then x_(i+1) = x_i^2 mod n.", "Example": "9, 81, 236"},
        {"Field": "b_i", "Type": "Bit", "Description": "Output bit: lowest bit (parity) of x_i.", "Example": "1, 0, 0"},
        {"Field": "Upload", "Type": "CSV, TXT, DOCX, PDF", "Description": "A bit string (0 and 1 characters) or a list of numbers to test for randomness. CSV column bits or value.", "Example": "0110100111..."},
    ])
    st.markdown("#### Test result fields")
    field_table([
        {"Field": "p value", "Type": "0 to 1", "Description": "Probability a truly random sequence would look at least this unusual.", "Example": "0.53"},
        {"Field": "pass", "Type": "Yes or No", "Description": "Pass when p value >= 0.01 (NIST significance level).", "Example": "Yes"},
    ])
    st.markdown("#### Formulas")
    formula("BBS recurrence", r"x_{0}=s^{2}\bmod n,\qquad x_{i+1}=x_{i}^{2}\bmod n,\qquad b_i = x_i \bmod 2",
            "Squaring is easy; taking square roots modulo n (to run the generator backwards) requires p and q.")
    formula("Jump ahead (with the secret factors)", r"x_i = x_0^{\,2^{i} \bmod \lambda(n)} \bmod n,\qquad \lambda(n)=\mathrm{lcm}(p-1,\,q-1)",
            "Only someone who knows the factorization can compute any state directly.")
    formula("Monobit test", r"S=\sum_{i}(2b_i-1),\qquad s_{obs}=\frac{|S|}{\sqrt{N}},\qquad p=\operatorname{erfc}\!\left(\frac{s_{obs}}{\sqrt2}\right)",
            "Checks that zeros and ones are balanced.")
    formula("Runs test", r"V=1+\sum_{k}[b_k\ne b_{k+1}],\qquad p=\operatorname{erfc}\!\left(\frac{|V-2N\pi(1-\pi)|}{2\sqrt{2N}\,\pi(1-\pi)}\right)",
            "Checks that switches between 0 and 1 happen at the expected rate (pi is the fraction of ones).")
    formula("Block frequency test", r"\chi^{2}=4M\sum_{j=1}^{B}\left(\pi_j-\tfrac12\right)^{2},\qquad p=\mathrm{igamc}\!\left(\tfrac{B}{2},\tfrac{\chi^{2}}{2}\right)",
            "Checks balance inside blocks of M bits.")
    formula("Very Smooth Hash", r"x_0=1,\qquad x_{j+1}=x_j^{2}\prod_{i=1}^{k}p_i^{\,m_{ij}}\bmod n",
            "Message split into k bit blocks m<sub>j</sub>; p<sub>i</sub> are the first k primes. A final block encodes the message length.")


def randomgen_demo():
    upload = data_source(RNG_KEY, "Upload a bit string (characters 0 and 1) or a list of numbers in CSV (column "
                              "<b>bits</b> or <b>value</b>), TXT, Word or PDF. The randomness tests are applied to it.",
                         pd.DataFrame({"value": [12, 7, 255, 3, 90]}))
    c1, c2, c3, c4 = st.columns(4)
    bits = c1.select_slider("Modulus size (bits)", options=[16, 32, 64, 128, 256, 512], value=128, key=f"{RNG_KEY}_bits")
    pseed = c2.number_input("Prime seed", 0, 10_000, 3, key=f"{RNG_KEY}_pseed")
    count = c3.select_slider("Bits to generate", options=[1000, 5000, 10000, 20000, 50000], value=20000, key=f"{RNG_KEY}_count")
    prm = params(int(bits), int(pseed))
    n = prm["n"]
    default_seed = str(random.Random(int(pseed) + 1).randrange(2, n - 1) | 1)
    seed_txt = c4.text_input("Seed s", default_seed, key=f"{RNG_KEY}_seed_{bits}_{pseed}")
    try:
        s = parse_int(seed_txt)
    except Exception:
        st.error("Seed must be a whole number")
        return None
    issues = validate_bbs(prm["p"], prm["q"], s)
    if issues:
        st.error("; ".join(issues))
        return None
    out = bbs_generate(n, s, int(count))
    b = out["bits"]
    show_df(pd.DataFrame([{"Parameter": "p (3 mod 4)", "Value": str(prm["p"])},
                          {"Parameter": "q (3 mod 4)", "Value": str(prm["q"])},
                          {"Parameter": "n = p x q", "Value": str(n)}, {"Parameter": "seed s", "Value": str(s)}]))
    st.markdown("First 20 states and output bits")
    show_df(pd.DataFrame([{"i": i, "x_i": str(x), "b_i = x_i mod 2": x & 1} for i, x in enumerate(out["states"][:21])]))

    def run_tests(label, bitseq):
        m = monobit_test(bitseq)
        r = runs_test(bitseq)
        bf = block_frequency_test(bitseq, 128)
        return [{"sequence": label, "test": "Frequency (monobit)", "statistic": round(m["s_obs"], 4), "p value": round(m["p_value"], 5), "pass": "Yes" if m["pass"] else "No"},
                {"sequence": label, "test": "Runs", "statistic": r["V"], "p value": round(r["p_value"], 5), "pass": "Yes" if r["pass"] else "No"},
                {"sequence": label, "test": "Block frequency (M = 128)", "statistic": round(bf["chi2"], 3), "p value": round(bf["p_value"], 5) if bf["p_value"] == bf["p_value"] else "n/a", "pass": "Yes" if bf["pass"] else "No"}]

    weak = lcg_bits(s, len(b))
    rows = run_tests("Blum Blum Shub", b) + run_tests("Weak LCG (mod 16)", weak)
    uploaded_bits = None
    if upload:
        text = upload["text"]
        if upload["df"] is not None and "bits" in upload["df"].columns:
            text = "".join(upload["df"]["bits"].astype(str))
        digits01 = "".join(ch for ch in text if ch in "01")
        other_digits = any(ch in "23456789" for ch in text)
        if not other_digits and len(digits01) >= 100:
            uploaded_bits = [int(ch) for ch in digits01]
        else:
            vals = extract_integers(text if upload["df"] is None or "value" not in upload["df"].columns
                                    else "\n".join(upload["df"]["value"].astype(str)))
            uploaded_bits = [int(bit) for v in vals for bit in format(v & 0xFF, "08b")]
        if len(uploaded_bits) >= 100:
            rows += run_tests(f"Uploaded: {upload['name']}", uploaded_bits)
        else:
            st.warning("The uploaded data has fewer than 100 bits; at least 100 are needed for the tests.")
            uploaded_bits = None
    tests = pd.DataFrame(rows)
    ones = sum(b)
    cyc = bbs_cycle_length(n, s) if n < 10 ** 9 else None
    metrics([("Bits generated", f"{len(b):,}"), ("Ones", f"{ones / len(b) * 100:.2f} percent"),
             ("BBS tests passed", f"{(tests[tests['sequence'] == 'Blum Blum Shub']['pass'] == 'Yes').sum()} of 3"),
             ("Cycle length", str(cyc) if cyc else "astronomical")], status=["", "", "ok", ""])
    st.markdown("##### NIST SP 800-22 statistical tests")
    show_df(tests)

    st.markdown("##### Random tokens produced from BBS output")
    by = bits_to_bytes(b)
    tokens = pd.DataFrame([{"token": i + 1, "hex (128 bit)": by[i * 16:(i + 1) * 16].hex(),
                            "6 digit one time code": f"{int.from_bytes(by[i * 16:i * 16 + 4], 'big') % 1_000_000:06d}"}
                           for i in range(min(6, len(by) // 16))])
    show_df(tokens)

    st.markdown("##### Very Smooth Hash (factoring based hash)")
    msg = st.text_input("Message to hash", "Transfer 500 USD to account 7788", key=f"{RNG_KEY}_vsh")
    h1 = vsh_hash(msg.encode(), n)
    msg_bytes = bytearray(msg.encode())
    if msg_bytes:
        msg_bytes[0] ^= 1
    h2 = vsh_hash(bytes(msg_bytes), n)
    dist = hamming_distance(h1["digest"], h2["digest"])
    metrics([("Primes per block k", str(h1["k"])), ("Blocks", str(h1["blocks"])),
             ("Digest bits", str(n.bit_length())), ("Bits changed by one bit flip", f"{dist} ({dist / n.bit_length() * 100:.0f} percent)")])
    mono(f"VSH(message) = {h1['digest']:x}<br>VSH(one bit changed) = {h2['digest']:x}")
    aval = []
    rng = random.Random(9)
    base = msg.encode() or b"x"
    for trial in range(40):
        mb = bytearray(base)
        pos = rng.randrange(len(mb) * 8)
        mb[pos // 8] ^= 1 << (7 - pos % 8)
        aval.append({"trial": trial + 1, "flipped bit": pos,
                     "digest bits changed": hamming_distance(vsh_hash(base, n)["digest"], vsh_hash(bytes(mb), n)["digest"])})
    aval_df = pd.DataFrame(aval)
    return {"prm": prm, "s": s, "bits": b, "weak": weak, "tests": tests, "tokens": tokens, "aval": aval_df,
            "n": n, "uploaded": uploaded_bits}


def randomgen_flow_dot() -> str:
    return """digraph R {
rankdir=LR; bgcolor="transparent";
node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=10.5, fillcolor="#E8F0FB", color="#1F4E8C"];
edge [color="#4A5B72", fontname="Helvetica", fontsize=9];
P [label="Secret primes p, q\\n(both 3 mod 4)", fillcolor="#FDECEC", color="#B23A3A"]; N [label="n = p x q\\n(public)"];
S [label="Seed s"]; X0 [label="x0 = s^2 mod n"]; X [label="x(i+1) = x(i)^2 mod n"]; B [label="Output bit\\nx(i) mod 2", fillcolor="#D6F0E0", color="#1E7B4A"];
T [label="Tokens, OTPs,\\nkeys, nonces"];
P -> N -> X0; S -> X0 -> X; X -> X [label="repeat"]; X -> B -> T;
A [label="Predict next bit ?", shape=diamond, fillcolor="#FFF5EB", color="#E07A1F"]; B -> A [style=dashed];
A -> P [label="only by factoring n", style=dashed, color="#B23A3A"];
}"""


def randomgen_visuals(ctx) -> None:
    st.markdown("##### Blum Blum Shub data flow")
    st.graphviz_chart(randomgen_flow_dot())
    if ctx is None:
        return
    b, weak = ctx["bits"], ctx["weak"]
    side = 64
    c1, c2 = st.columns(2)
    for col, seq, title in ((c1, b, "Blum Blum Shub: first 4096 bits"), (c2, weak, "Weak LCG: first 4096 bits")):
        grid = np.array(seq[:side * side]).reshape(side, side)
        fig = go.Figure(go.Heatmap(z=grid, colorscale=[[0, "#F3F6FB"], [1, "#0B4F9C"]], showscale=False))
        fig.update_xaxes(visible=False)
        fig.update_yaxes(visible=False, autorange="reversed")
        with col:
            st.plotly_chart(style_fig(fig, 360, title), key=f"{RNG_KEY}_{title[:4]}")
    st.caption("A good generator looks like static. The weak generator shows stripes because its low bit simply alternates.")
    c1, c2 = st.columns(2)
    with c1:
        walk = np.cumsum([2 * x - 1 for x in b])
        wk = np.cumsum([2 * x - 1 for x in weak])
        fig = go.Figure()
        fig.add_trace(go.Scatter(y=walk, mode="lines", name="BBS"))
        fig.add_trace(go.Scatter(y=wk, mode="lines", name="Weak LCG"))
        lim = 3 * np.sqrt(np.arange(1, len(b) + 1))
        fig.add_trace(go.Scatter(y=lim, mode="lines", line=dict(dash="dot", color="#999"), name="Plus or minus 3 sqrt(N)"))
        fig.add_trace(go.Scatter(y=-lim, mode="lines", line=dict(dash="dot", color="#999"), showlegend=False))
        fig.update_xaxes(title="Bit index")
        fig.update_yaxes(title="Cumulative sum of plus or minus 1")
        st.plotly_chart(style_fig(fig, 340, "Random walk of the bit stream"), key=f"{RNG_KEY}_walk")
    with c2:
        by = bits_to_bytes(b)
        counts = np.bincount(np.frombuffer(by, dtype=np.uint8), minlength=256)
        fig = go.Figure(go.Bar(x=list(range(256)), y=counts, marker_color="#0B4F9C"))
        fig.add_hline(y=len(by) / 256, line_dash="dot", annotation_text="Expected")
        fig.update_xaxes(title="Byte value")
        fig.update_yaxes(title="Count")
        st.plotly_chart(style_fig(fig, 340, "Byte value distribution"), key=f"{RNG_KEY}_bytes")
    a = ctx["aval"]
    fig = go.Figure(go.Bar(x=a["trial"], y=a["digest bits changed"], marker_color="#8E44AD"))
    fig.add_hline(y=ctx["n"].bit_length() / 2, line_dash="dot", annotation_text="Ideal: half the digest bits")
    fig.update_xaxes(title="Trial (one random input bit flipped)")
    fig.update_yaxes(title="Digest bits changed")
    st.plotly_chart(style_fig(fig, 320, "Very Smooth Hash avalanche test"), key=f"{RNG_KEY}_aval")
    tests = ctx["tests"]
    fig = go.Figure()
    for seq in tests["sequence"].unique():
        sub = tests[tests["sequence"] == seq]
        fig.add_trace(go.Bar(x=sub["test"], y=pd.to_numeric(sub["p value"], errors="coerce"), name=seq))
    fig.add_hline(y=0.01, line_color="#B23A3A", annotation_text="Pass threshold 0.01")
    fig.update_yaxes(title="p value", type="log")
    fig.update_layout(barmode="group")
    st.plotly_chart(style_fig(fig, 320, "Test p values by sequence (log scale)"), key=f"{RNG_KEY}_pv")


def randomgen_report(ctx):
    if ctx is None:
        return None
    rep = Report(use_case="Hashing and Random Number Generation")
    rep.summary = [("p", str(ctx["prm"]["p"])), ("q", str(ctx["prm"]["q"])), ("n", str(ctx["n"])),
                   ("Seed", str(ctx["s"])), ("Bits generated", str(len(ctx["bits"]))),
                   ("Fraction of ones", f"{sum(ctx['bits']) / len(ctx['bits']):.4f}"),
                   ("Mean VSH avalanche", f"{ctx['aval']['digest bits changed'].mean():.1f} of {ctx['n'].bit_length()} bits")]
    rep.sections = [("Interpretation", "The BBS output passes the NIST frequency, runs and block frequency tests "
                     "while the weak linear congruential generator fails. Predicting BBS output requires factoring n.")]
    rep.formulas = [("BBS", "x(i+1) = x(i)^2 mod n; b(i) = x(i) mod 2"), ("Monobit", "p = erfc(|S| / sqrt(2N))"),
                    ("Runs", "p = erfc(|V - 2N pi(1 - pi)| / (2 sqrt(2N) pi (1 - pi)))"),
                    ("VSH", "x(j+1) = x(j)^2 x product of p(i)^m(ij) mod n")]
    rep.tables = [("Randomness tests", ctx["tests"]), ("Tokens", ctx["tokens"]), ("VSH avalanche", ctx["aval"])]
    t = ctx["tests"]
    rep.chart = {"kind": "bar", "x": list(t["sequence"] + " " + t["test"]),
                 "y": {"p value": [float(v) if v != "n/a" else 0 for v in t["p value"]]},
                 "title": "Test p values (pass at 0.01 or more)", "xlabel": "", "ylabel": "p value"}
    return rep


def randomgen_render() -> None:
    page_heading("06", "Hashing and Random Number Generation", "Pseudo random generators such as Blum Blum Shub, and hashes such as VSH, rely on the hardness of factoring.")
    t = use_case_tabs()
    with t[2]:
        ctx = randomgen_demo()
    with t[0]:
        randomgen_overview()
    with t[1]:
        randomgen_schema()
    with t[3]:
        randomgen_visuals(ctx)
    with t[4]:
        export_panel(randomgen_report(ctx), RNG_KEY)


# ============================================================================
# PAGE: security
# Page 07: Security testing (weak key discovery).
# ============================================================================

SEC_KEY = "sec"
RISK_COLORS = {"Critical": "#B23A3A", "High": "#E07A1F", "Medium": "#D4AC0D", "Low": "#1E7B4A"}


@st.cache_data(show_spinner="Generating synthetic device fleet...")
def fleet(seed: int, bits: int):
    return synthetic_device_keys(random.Random(seed), bits)


@st.cache_data(show_spinner="Auditing keys...")
def run_audit(records: tuple, min_bits: int, fermat_iter: int, pm1: int):
    return audit_keys([dict(r) for r in records], min_bits=min_bits, fermat_iterations=fermat_iter, pm1_bound=pm1)


def security_overview() -> None:
    callout("Security researchers collect public RSA keys from certificates, SSH servers and devices, then try "
            "cheap factoring attacks on all of them at once. Keys that break reveal <b>faulty random number "
            "generators, bad key generation code and outdated key sizes</b>. In 2012 two research teams found "
            "that tens of thousands of network devices on the internet shared prime factors, letting anyone "
            "compute their private keys with a simple GCD.")
    cards([
        ("Find vulnerable devices", "Batch GCD tests millions of keys against each other in quasi linear time "
         "and exposes any two that share a prime."),
        ("Detect bad key generation", "Primes that are too close (Fermat), have smooth p minus 1 (Pollard) or "
         "are tiny (trial division) indicate broken software."),
        ("Enforce policy", "Flags keys below the organisation's minimum size and risky public exponents."),
        ("Prioritise remediation", "Risk ratings (Critical, High, Medium, Low) direct effort to keys that are "
         "already broken."),
        ("Drive standards", "Findings like these led vendors to fix entropy at boot and pushed standards such "
         "as FIPS 186-5 to require well separated, properly generated primes."),
        ("Continuous monitoring", "Upload your certificate inventory regularly to catch regressions."),
    ])
    callout("Only audit keys you own or are authorised to test. The checks use public keys only.", "warn")


def security_schema() -> None:
    st.markdown("#### Input fields")
    field_table([
        {"Field": "device_id", "Type": "Text", "Description": "Identifier of the device, server or certificate.", "Example": "DEV-201"},
        {"Field": "device_type", "Type": "Text (optional)", "Description": "Category for reporting.", "Example": "Router"},
        {"Field": "modulus", "Type": "Integer (decimal or 0x hex)", "Description": "The public RSA modulus n from the certificate or key.", "Example": "0xC3A1..."},
        {"Field": "exponent", "Type": "Integer (optional)", "Description": "Public exponent e; defaults to 65537.", "Example": "65537"},
        {"Field": "Upload", "Type": "CSV, TXT, DOCX, PDF", "Description": "CSV with the columns above, or any document; numbers with 20 or more digits are read as moduli.", "Example": "device_id,modulus,exponent"},
    ])
    st.markdown("#### Output fields")
    field_table([
        {"Field": "risk", "Type": "Critical, High, Medium, Low", "Description": "Critical: private key recovered. High: risky exponent or duplicate key. Medium: below size policy. Low: nothing found.", "Example": "Critical"},
        {"Field": "method", "Type": "Text", "Description": "Attack that recovered the factor.", "Example": "Batch GCD"},
        {"Field": "factor_p, factor_q", "Type": "Integers", "Description": "Recovered primes (the key is fully compromised).", "Example": "3401..."},
        {"Field": "findings", "Type": "Text", "Description": "All issues found for the key.", "Example": "Shares a prime with another key"},
    ])
    st.markdown("#### Formulas")
    formula("Shared prime", r"\gcd(N_i, N_j) = p > 1 \;\Rightarrow\; N_i = p \cdot \frac{N_i}{p}",
            "Two keys that share a prime are both broken by a single GCD.")
    formula("Batch GCD (Bernstein)", r"P=\prod_j N_j,\qquad g_i=\gcd\!\left(N_i,\ \frac{P \bmod N_i^{2}}{N_i}\right)",
            "A product tree and remainder tree compute every g<sub>i</sub> without comparing all pairs.")
    formula("Fermat factorization", r"N = a^{2}-b^{2} = (a-b)(a+b),\qquad a = \lceil\sqrt N\rceil, \lceil\sqrt N\rceil+1, \ldots",
            "Succeeds in about (p - q)<sup>2</sup> / (8 sqrt N) steps, so instantly when p and q are close.")
    formula("Pollard p minus 1", r"M=\prod_{\ell^{k}\le B}\ell^{k},\qquad g=\gcd\!\left(2^{M}-1,\ N\right)",
            "Finds p whenever every prime power dividing p - 1 is at most B.")
    formula("Risk score", r"\text{Critical if a factor is found; High if } e<65537;\ \text{Medium if } \log_2 N < \text{policy}",
            "Rules used to rank findings.")


def security_demo():
    upload = data_source(SEC_KEY, "CSV with columns <b>device_id, modulus, exponent</b> (device_type optional). In "
                              "TXT, Word or PDF files every number with at least 20 digits is treated as a modulus.",
                         pd.DataFrame({"device_id": ["SRV-1", "SRV-2"], "modulus": ["0xC5", "221"], "exponent": [65537, 65537]}))
    c1, c2, c3, c4 = st.columns(4)
    bits = c1.select_slider("Synthetic key size", options=[128, 256, 512], value=256, key=f"{SEC_KEY}_bits")
    seed = c2.number_input("Fleet seed", 0, 10_000, 5, key=f"{SEC_KEY}_seed")
    policy = c3.select_slider("Minimum size policy (bits)", options=[128, 256, 512, 1024, 2048, 3072],
                              value=256 if not upload else 2048, key=f"{SEC_KEY}_policy_{bool(upload)}")
    effort = c4.select_slider("Attack effort", options=["Quick", "Standard", "Deep"], value="Standard", key=f"{SEC_KEY}_effort")
    fermat_iter, pm1 = {"Quick": (2_000, 5_000), "Standard": (20_000, 20_000), "Deep": (200_000, 100_000)}[effort]

    if upload:
        df = upload["df"]
        if df is not None and "modulus" in df.columns:
            records = []
            for i, r in df.iterrows():
                try:
                    records.append({"device_id": str(r.get("device_id", f"KEY-{i + 1}")),
                                    "device_type": str(r.get("device_type", "Uploaded")),
                                    "modulus": parse_int(r["modulus"]),
                                    "exponent": parse_int(r["exponent"]) if str(r.get("exponent", "")).strip() else 65537})
                except Exception:
                    st.warning(f"Row {i + 1}: could not read the modulus, skipped.")
        else:
            nums = extract_integers(upload["text"], min_digits=20)
            records = [{"device_id": f"KEY-{i + 1}", "device_type": "Uploaded", "modulus": v, "exponent": 65537}
                       for i, v in enumerate(nums)]
        if not records:
            st.error("No moduli found in the file.")
            return None
        truth = None
    else:
        records = fleet(int(seed), int(bits))
        truth = {r["device_id"]: r["planted_weakness"] for r in records}

    t0 = time.perf_counter()
    res = run_audit(tuple(tuple(sorted({k: v for k, v in r.items() if k in ("device_id", "modulus", "exponent")}.items()))
                          for r in records), int(policy), fermat_iter, pm1)
    elapsed = time.perf_counter() - t0
    df = pd.DataFrame(res)
    df.insert(1, "device_type", [r.get("device_type", "") for r in records])
    if truth:
        df.insert(2, "planted weakness (ground truth)", [truth[d] for d in df["device_id"]])
    counts = df["risk"].value_counts()
    metrics([("Keys audited", str(len(df))), ("Critical (broken)", str(counts.get("Critical", 0))),
             ("High", str(counts.get("High", 0))), ("Medium", str(counts.get("Medium", 0))),
             ("Low", str(counts.get("Low", 0))), ("Audit time", f"{elapsed:.2f} s")],
            status=["", "bad", "warn", "warn", "ok", ""])
    if truth:
        planted = df["planted weakness (ground truth)"] != "None (strong key)"
        detected = df["risk"] != "Low"
        tp = int((planted & detected).sum())
        fp = int((~planted & detected & (df["risk"] != "Medium")).sum())
        callout(f"Detection accuracy on synthetic data: <b>{tp} of {int(planted.sum())}</b> planted weaknesses "
                f"found, <b>{fp}</b> false alarms among strong keys.", "ok" if tp == int(planted.sum()) and fp == 0 else "warn")
    risk_order = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
    df = df.sort_values(by="risk", key=lambda s: s.map(risk_order)).reset_index(drop=True)
    view = df.drop(columns=["modulus"]).copy()
    view["factor_p"] = view["factor_p"].map(lambda v: short(v, 22))
    view["factor_q"] = view["factor_q"].map(lambda v: short(v, 22))
    view["shared_gcd"] = view["shared_gcd"].map(lambda v: short(v, 22))
    st.markdown("##### Audit results (sorted by risk)")
    show_df(view)

    st.markdown("##### Fermat method on the closest prime pair")
    close = df[df["method"] == "Fermat"]
    fermat_rows = []
    if len(close):
        n = int(close.iloc[0]["modulus"])
        p, q, it = fermat_factor(n)
        callout(f"Device {close.iloc[0]['device_id']}: p and q differ by {q - p:,}, a relative gap of "
                f"{(q - p) / p:.2e}. Fermat found them after {it} iteration(s).", "bad")
    return {"df": df, "records": records, "policy": policy, "elapsed": elapsed, "truth": truth is not None}


def shared_dot(df: pd.DataFrame) -> str:
    lines = ['graph G {', 'bgcolor="transparent"; layout=neato; overlap=false; splines=true;',
             'node [fontname="Helvetica", fontsize=10, style="filled", shape=box];']
    groups = {}
    for _, r in df.iterrows():
        if r["shared_gcd"]:
            groups.setdefault(r["shared_gcd"], []).append(r["device_id"])
    if not groups:
        lines.append('none [label="No shared primes found", fillcolor="#D6F0E0"];')
    for i, (g, devs) in enumerate(groups.items()):
        lines.append(f'p{i} [label="Shared prime\\n{short(g, 16)}", shape=ellipse, fillcolor="#FDECEC", color="#B23A3A"];')
        for d in devs:
            lines.append(f'"{d}" [fillcolor="#E8F0FB", color="#1F4E8C"];')
            lines.append(f'p{i} -- "{d}" [color="#B23A3A"];')
    lines.append("}")
    return "\n".join(lines)


def pipeline_dot() -> str:
    return """digraph P {
rankdir=LR; bgcolor="transparent";
node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=10.5, fillcolor="#E8F0FB", color="#1F4E8C"];
edge [color="#4A5B72"];
C [label="Collect public keys\\n(certificates, SSH, devices)"]; B [label="Batch GCD\\nshared primes"];
T [label="Trial division\\nsmall factors"]; F [label="Fermat\\nclose primes"]; P [label="Pollard p - 1\\nsmooth primes"];
Y [label="Policy checks\\nsize, exponent"]; R [label="Risk rating\\nand report", fillcolor="#D6F0E0", color="#1E7B4A"];
C -> B -> T -> F -> P -> Y -> R;
}"""


def security_visuals(ctx) -> None:
    st.markdown("##### Audit pipeline")
    st.graphviz_chart(pipeline_dot())
    if ctx is None:
        return
    df = ctx["df"]
    c1, c2 = st.columns(2)
    with c1:
        counts = df["risk"].value_counts().reindex(["Critical", "High", "Medium", "Low"]).fillna(0)
        fig = go.Figure(go.Pie(labels=counts.index, values=counts.values, hole=0.55,
                               marker=dict(colors=[RISK_COLORS[k] for k in counts.index]), sort=False))
        st.plotly_chart(style_fig(fig, 340, "Keys by risk level"), key=f"{SEC_KEY}_pie")
    with c2:
        m = df[df["method"].isin(["Batch GCD", "Fermat", "Pollard p minus 1", "Trial division", "Even modulus"])]
        mc = m["method"].value_counts().reset_index()
        mc.columns = ["method", "keys broken"]
        fig = px.bar(mc, x="method", y="keys broken", color_discrete_sequence=["#B23A3A"])
        st.plotly_chart(style_fig(fig, 340, "Keys broken by attack method"), key=f"{SEC_KEY}_methods")
    st.markdown("##### Shared prime network (keys linked through a common factor)")
    st.graphviz_chart(shared_dot(df))
    if "device_type" in df.columns:
        piv = df.groupby(["device_type", "risk"]).size().reset_index(name="count")
        fig = px.bar(piv, x="device_type", y="count", color="risk", color_discrete_map=RISK_COLORS,
                     category_orders={"risk": ["Critical", "High", "Medium", "Low"]})
        st.plotly_chart(style_fig(fig, 340, "Risk by device type"), key=f"{SEC_KEY}_types")


def security_report(ctx):
    if ctx is None:
        return None
    df = ctx["df"]
    rep = Report(use_case="Security Testing (Weak RSA Key Discovery)")
    counts = df["risk"].value_counts()
    rep.summary = [("Keys audited", str(len(df))), ("Size policy", f"{ctx['policy']} bits"),
                   ("Critical", str(counts.get("Critical", 0))), ("High", str(counts.get("High", 0))),
                   ("Medium", str(counts.get("Medium", 0))), ("Low", str(counts.get("Low", 0))),
                   ("Audit time", f"{ctx['elapsed']:.2f} s")]
    rep.sections = [("Recommendations", "Revoke and reissue every Critical key immediately; their private keys "
                     "can be computed by anyone. Replace keys with e below 65537 and keys below the size policy. "
                     "Investigate the random number generators of device families with shared primes.")]
    rep.formulas = [("Shared prime", "gcd(N_i, N_j) = p > 1"), ("Batch GCD", "g_i = gcd(N_i, (P mod N_i^2) / N_i)"),
                    ("Fermat", "N = a^2 - b^2 = (a - b)(a + b)"), ("Pollard p - 1", "g = gcd(2^M - 1, N)")]
    rep.tables = [("Audit results", df)]
    c = df["risk"].value_counts().reindex(["Critical", "High", "Medium", "Low"]).fillna(0)
    rep.chart = {"kind": "bar", "x": list(c.index), "y": {"Keys": [int(v) for v in c.values]},
                 "title": "Keys by risk level", "xlabel": "Risk", "ylabel": "Keys"}
    return rep


def security_render() -> None:
    page_heading("07", "Security Testing", "Researchers factor weak or badly generated keys to find vulnerable devices and push for stronger standards.")
    t = use_case_tabs()
    with t[2]:
        ctx = security_demo()
    with t[0]:
        security_overview()
    with t[1]:
        security_schema()
    with t[3]:
        security_visuals(ctx)
    with t[4]:
        export_panel(security_report(ctx), SEC_KEY)


# ============================================================================
# PAGE: validation
# Page 08: Accuracy validation.
# ============================================================================

VAL_KEY = "val"


@st.cache_data(show_spinner="Running the validation suite...")
def cached_run(quick: bool):
    t = time.perf_counter()
    rows = run_all(quick=quick)
    return rows, time.perf_counter() - t


def validation_render() -> None:
    page_heading("08", "Accuracy Validation", "Every algorithm in the suite is checked against published reference values and independent implementations.")
    callout("The checks below run live on this server. References include the Wikipedia RSA and Blum Blum Shub "
            "worked examples, NIST SP 800-22 examples, the FIPS 180-4 SHA-256 and RFC 4231 HMAC test vectors, "
            "classical factorizations (Euler, Landry, Cole) and cross checks against the SymPy library.")
    quick = st.toggle("Quick mode (fewer random trials)", value=False, key=f"{VAL_KEY}_quick")
    rows, elapsed = cached_run(quick)
    df = pd.DataFrame(rows)
    passed = int((df["Result"] == "Pass").sum())
    metrics([("Checks", str(len(df))), ("Passed", str(passed)), ("Failed", str(len(df) - passed)),
             ("Pass rate", f"{passed / len(df) * 100:.1f} percent"), ("Run time", f"{elapsed:.1f} s")],
            status=["", "ok", "bad" if len(df) - passed else "ok", "ok" if passed == len(df) else "bad", ""])
    t = st.tabs(["Results", "Summary chart", "Export Results"])
    with t[0]:
        cat = st.multiselect("Filter by category", sorted(df["Category"].unique()), key=f"{VAL_KEY}_cat")
        view = df[df["Category"].isin(cat)] if cat else df
        show_df(view)
    with t[1]:
        g = df.groupby(["Category", "Result"]).size().unstack(fill_value=0)
        fig = go.Figure()
        for res, color in (("Pass", "#1E7B4A"), ("Fail", "#B23A3A")):
            if res in g.columns:
                fig.add_trace(go.Bar(x=g.index, y=g[res], name=res, marker_color=color))
        fig.update_layout(barmode="stack")
        fig.update_yaxes(title="Checks")
        st.plotly_chart(style_fig(fig, 380, "Validation checks by category"), key=f"{VAL_KEY}_chart")
    with t[2]:
        rep = Report(use_case="Accuracy Validation")
        rep.summary = [("Checks", str(len(df))), ("Passed", str(passed)), ("Failed", str(len(df) - passed)),
                       ("Run time (s)", f"{elapsed:.1f}")]
        rep.tables = [("Validation results", df)]
        g2 = df.groupby("Category").size()
        rep.chart = {"kind": "bar", "x": list(g2.index), "y": {"Checks": [int(v) for v in g2.values]},
                     "title": "Checks per category", "xlabel": "Category", "ylabel": "Checks"}
        export_panel(rep, VAL_KEY)


# ============================================================================
# APPLICATION ENTRY POINT
# Page registry, left side navigation and main().
# ============================================================================



st.set_page_config(page_title="PrimeShield | KNet Consulting Group", layout="wide",
                   initial_sidebar_state="expanded")

PAGES = {
    "00  Factorization Engine": home_render,
    "01  RSA Encryption": rsa_render,
    "02  Digital Signatures": signatures_render,
    "03  Online Banking and Payments": banking_render,
    "04  Quantum Computing Research": quantum_render,
    "05  Fractions, GCD and LCM": fractions_render,
    "06  Hashing and Random Numbers": randomgen_render,
    "07  Security Testing": security_render,
    "08  Accuracy Validation": validation_render,
}


def sidebar() -> str:
    with st.sidebar:
        st.markdown('<div class="ps-side-brand">PrimeShield</div>'
                    '<div class="ps-side-tag">Prime Factorization and Cryptography Suite</div>',
                    unsafe_allow_html=True)
        st.markdown('<div class="ps-side-head">USE CASES</div>', unsafe_allow_html=True)
        choice = st.radio("Navigation", list(PAGES), key="nav", label_visibility="collapsed")
        st.markdown('<div class="ps-side-foot">Kalsnet (KNet) Consulting Group<br>'
                    'Version 1.0, September 2026<br>For education, research and authorised testing.</div>',
                    unsafe_allow_html=True)
    return choice


def main() -> None:
    inject_css()
    choice = sidebar()
    title_bar()
    PAGES[choice]()


main()
