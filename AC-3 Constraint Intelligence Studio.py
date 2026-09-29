



# =====================================================================================
#  AC-3 Constraint Intelligence Studio  (single-file Streamlit application)
#  Developed by Randy Singh from Kalsnet (KNet) Consulting
#
#  Run:   pip install -r requirements.txt
#         streamlit run ac3_studio.py
#
#  Contents of this file
#    1. Branding + AC-3 engine (CSP, AC-3, MAC backtracking search)
#    2. Report exporters (PDF, Word, CSV, TXT) + shared charts
#    3. Use case 1  Sudoku & logic puzzles
#    4. Use case 2  Scheduling & timetabling
#    5. Use case 3  Map colouring & radio-frequency assignment
#    6. Use case 4  Product configuration (PC builder / car configurator)
#    7. Use case 5  Planning & resource allocation (job-shop, projects, docks)
#    8. Use case 6  Vision (Waltz line labelling) & NLP (constraint dependency parsing)
#    9. Streamlit user interface
# =====================================================================================
from __future__ import annotations


# =====================================================================================
# SECTION: core/branding.py
# =====================================================================================
APP_NAME = "AC-3 Constraint Intelligence Studio"
APP_TAGLINE = "Arc-consistency powered problem solving for six real-world domains"
DEVELOPER_LINE = "Developed by Randy Singh from Kalsnet (KNet) Consulting"
ORG = "Kalsnet (KNet) Consulting"
VERSION = "1.0.0"

BLUE = "#0B4F9C"
BLUE_DARK = "#07326A"
BLUE_LIGHT = "#E8F1FB"
ACCENT = "#F28C28"
GREEN = "#1E8E5A"
RED = "#C62828"

# qualitative palette used across charts (colour-blind friendly order)
PALETTE = ["#1F77B4", "#FF7F0E", "#2CA02C", "#D62728", "#9467BD", "#8C564B",
           "#E377C2", "#17BECF", "#BCBD22", "#7F7F7F", "#393B79", "#AD494A"]


# =====================================================================================
# SECTION: core/ac3.py
# =====================================================================================
# Generic AC-3 engine and MAC (Maintaining Arc Consistency) backtracking solver.
#
# Terminology
# -----------
# * Variable  X_i            - something we must decide (a Sudoku cell, an exam, a region ...)
# * Domain    D_i            - the finite set of values X_i may still take
# * Constraint C_ij(a, b)    - a binary predicate that is True when X_i=a and X_j=b are compatible
# * Arc       (X_i -> X_j)   - one direction of a binary constraint. Every constraint gives 2 arcs.
#
# An arc (X_i -> X_j) is *consistent* when
#         for every a in D_i there exists b in D_j such that C_ij(a, b) holds.
# AC-3 repeatedly "revises" arcs, deleting unsupported values, until every arc is consistent
# or some domain becomes empty (proof that no solution exists).
#
# Worst-case time complexity: O(e * d^3), e = number of arcs, d = largest domain size.

import time
from collections import deque
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, Hashable, Iterable, List, Optional, Tuple

import pandas as pd

Pred = Callable[[Any, Any], bool]


def _flip(pred: Pred) -> Pred:
    return lambda b, a: pred(a, b)


class CSP:
    """A binary constraint satisfaction problem."""

    def __init__(self, variables: Iterable[Hashable], domains: Dict[Hashable, Iterable[Any]], name: str = "CSP"):
        self.name = name
        self.variables: List[Hashable] = list(variables)
        if len(set(self.variables)) != len(self.variables):
            raise ValueError("Duplicate variable names in CSP.")
        missing = [v for v in self.variables if v not in domains]
        if missing:
            raise ValueError(f"No domain supplied for variables: {missing[:5]}")
        self.domains: Dict[Hashable, List[Any]] = {v: list(domains[v]) for v in self.variables}
        self.original_domains: Dict[Hashable, List[Any]] = {v: list(d) for v, d in self.domains.items()}
        self._preds: Dict[Tuple[Hashable, Hashable], List[Pred]] = {}
        self.neighbors: Dict[Hashable, List[Hashable]] = {v: [] for v in self.variables}
        self.constraint_labels: Dict[Tuple[Hashable, Hashable], List[str]] = {}
        self.unary_log: List[Tuple[Hashable, str, int]] = []  # (var, label, values removed)
        self.value_formatter: Callable[[Hashable, Any], str] = lambda var, val: str(val)

    # ------------------------------------------------------------------ building
    def add_unary(self, var: Hashable, pred: Callable[[Any], bool], label: str = "") -> int:
        """Node consistency: filter a single variable's domain. Returns number of values removed."""
        before = len(self.domains[var])
        self.domains[var] = [v for v in self.domains[var] if pred(v)]
        removed = before - len(self.domains[var])
        self.unary_log.append((var, label, removed))
        return removed

    def add_binary(self, x: Hashable, y: Hashable, pred: Pred, label: str = "") -> None:
        """Add constraint pred(value_of_x, value_of_y). Both arcs x->y and y->x are created."""
        if x == y:
            raise ValueError("Binary constraint needs two different variables.")
        if x not in self.neighbors or y not in self.neighbors:
            raise KeyError(f"Unknown variable in constraint: {x!r}, {y!r}")
        self._preds.setdefault((x, y), []).append(pred)
        self._preds.setdefault((y, x), []).append(_flip(pred))
        if y not in self.neighbors[x]:
            self.neighbors[x].append(y)
            self.neighbors[y].append(x)
        if label:
            self.constraint_labels.setdefault((x, y), []).append(label)
            self.constraint_labels.setdefault((y, x), []).append(label)

    # ------------------------------------------------------------------ queries
    def arcs(self) -> List[Tuple[Hashable, Hashable]]:
        return list(self._preds.keys())

    @property
    def n_constraints(self) -> int:
        return len(self._preds) // 2

    def satisfies(self, x: Hashable, a: Any, y: Hashable, b: Any) -> bool:
        for p in self._preds[(x, y)]:
            if not p(a, b):
                return False
        return True

    def label_of(self, x: Hashable, y: Hashable) -> str:
        return "; ".join(dict.fromkeys(self.constraint_labels.get((x, y), [])))

    def violations(self, assignment: Dict[Hashable, Any]) -> List[Tuple[Hashable, Hashable, str]]:
        """Return every violated binary constraint in a (complete or partial) assignment."""
        out = []
        seen = set()
        for (x, y) in self._preds:
            key = frozenset((x, y))
            if key in seen or x not in assignment or y not in assignment:
                continue
            seen.add(key)
            if not self.satisfies(x, assignment[x], y, assignment[y]):
                out.append((x, y, self.label_of(x, y)))
        return out

    def fmt(self, var: Hashable, val: Any) -> str:
        try:
            return self.value_formatter(var, val)
        except Exception:  # pragma: no cover - formatter is cosmetic
            return str(val)


# ====================================================================== AC-3
@dataclass
class AC3Result:
    consistent: bool
    domains: Dict[Hashable, List[Any]]
    initial_domains: Dict[Hashable, List[Any]]
    original_domains: Dict[Hashable, List[Any]]
    arcs_initial: int
    arcs_processed: int
    revisions: int
    values_removed: int
    constraint_checks: int
    elapsed_ms: float
    trace: List[dict] = field(default_factory=list)
    trace_truncated: bool = False
    wiped_out: Optional[Hashable] = None

    @property
    def solved(self) -> bool:
        return self.consistent and all(len(d) == 1 for d in self.domains.values())

    @property
    def search_space_before(self) -> float:
        return _log10_product(len(d) for d in self.initial_domains.values())

    @property
    def search_space_after(self) -> float:
        return _log10_product(len(d) for d in self.domains.values())

    def domain_table(self, csp: Optional[CSP] = None) -> pd.DataFrame:
        rows = []
        for v in self.domains:
            o, i, a = self.original_domains.get(v, []), self.initial_domains[v], self.domains[v]
            fmt = (lambda val: csp.fmt(v, val)) if csp else str
            rows.append({
                "Variable": str(v),
                "|D| original": len(o),
                "|D| after unary rules": len(i),
                "|D| after AC-3": len(a),
                "Values pruned by AC-3": len(i) - len(a),
                "Remaining values": ", ".join(fmt(x) for x in a[:12]) + (" ..." if len(a) > 12 else ""),
            })
        return pd.DataFrame(rows)

    def trace_table(self, limit: int = 500) -> pd.DataFrame:
        return pd.DataFrame(self.trace[:limit])

    def metrics(self) -> Dict[str, Any]:
        total_init = sum(len(d) for d in self.initial_domains.values())
        total_after = sum(len(d) for d in self.domains.values())
        return {
            "Arc consistent": "Yes" if self.consistent else f"No - domain of {self.wiped_out} wiped out",
            "Arcs in initial queue": self.arcs_initial,
            "Arcs processed": self.arcs_processed,
            "Successful revisions": self.revisions,
            "Values pruned": self.values_removed,
            "Constraint checks": self.constraint_checks,
            "Domain reduction %": round(100.0 * (total_init - total_after) / total_init, 1) if total_init else 0.0,
            "Search space before (log10)": round(self.search_space_before, 2),
            "Search space after (log10)": round(self.search_space_after, 2),
            "Solved by AC-3 alone": "Yes" if self.solved else "No",
            "AC-3 time (ms)": round(self.elapsed_ms, 2),
        }


def _log10_product(sizes: Iterable[int]) -> float:
    import math
    total = 0.0
    for s in sizes:
        if s == 0:
            return float("-inf")
        total += math.log10(s)
    return total


def _revise(csp: CSP, D: Dict[Hashable, List[Any]], x: Hashable, y: Hashable, counter: List[int]) -> List[Any]:
    """Remove values of x that have no support in y. Creates a NEW list (safe with shallow copies)."""
    Dy = D[y]
    preds = csp._preds[(x, y)]
    kept, removed = [], []
    checks = 0
    for a in D[x]:
        supported = False
        for b in Dy:
            checks += 1
            ok = True
            for p in preds:
                if not p(a, b):
                    ok = False
                    break
            if ok:
                supported = True
                break
        (kept if supported else removed).append(a)
    counter[0] += checks
    if removed:
        D[x] = kept
    return removed


def ac3(csp: CSP, domains: Optional[Dict[Hashable, List[Any]]] = None,
        arcs: Optional[Iterable[Tuple[Hashable, Hashable]]] = None,
        record_trace: bool = True, trace_limit: int = 5000) -> AC3Result:
    """Run AC-3 and return the reduced domains plus full statistics and a step trace."""
    t0 = time.perf_counter()
    D = {v: list(vals) for v, vals in (domains if domains is not None else csp.domains).items()}
    initial = {v: list(vals) for v, vals in D.items()}
    queue = deque(arcs if arcs is not None else csp.arcs())
    in_queue = set(queue)
    arcs_initial = len(queue)
    processed = revisions = removed_total = 0
    counter = [0]
    trace: List[dict] = []
    truncated = False
    wiped = None
    consistent = True

    # an empty initial domain is already inconsistent
    for v, d in D.items():
        if not d:
            consistent, wiped = False, v
            break

    while consistent and queue:
        x, y = queue.popleft()
        in_queue.discard((x, y))
        processed += 1
        before = len(D[x])
        removed = _revise(csp, D, x, y, counter)
        if not removed:
            continue
        revisions += 1
        removed_total += len(removed)
        requeued = 0
        if D[x]:
            for z in csp.neighbors[x]:
                if z != y and (z, x) not in in_queue:
                    queue.append((z, x))
                    in_queue.add((z, x))
                    requeued += 1
        if record_trace:
            if len(trace) < trace_limit:
                trace.append({
                    "Step": processed,
                    "Arc (Xi -> Xj)": f"{x} -> {y}",
                    "Constraint": csp.label_of(x, y),
                    "|Di| before": before,
                    "Removed from Di": ", ".join(csp.fmt(x, r) for r in removed[:10]) + (" ..." if len(removed) > 10 else ""),
                    "#Removed": len(removed),
                    "|Di| after": len(D[x]),
                    "Arcs re-queued": requeued,
                    "_x": x, "_y": y, "_removed": removed,
                })
            else:
                truncated = True
        if not D[x]:
            consistent, wiped = False, x
            break

    return AC3Result(
        consistent=consistent, domains=D, initial_domains=initial,
        original_domains={v: list(d) for v, d in csp.original_domains.items()},
        arcs_initial=arcs_initial, arcs_processed=processed, revisions=revisions,
        values_removed=removed_total, constraint_checks=counter[0],
        elapsed_ms=(time.perf_counter() - t0) * 1000, trace=trace,
        trace_truncated=truncated, wiped_out=wiped,
    )


def _ac3_fast(csp: CSP, D: Dict[Hashable, List[Any]], queue: deque, counter: List[int]) -> bool:
    in_queue = set(queue)
    while queue:
        x, y = queue.popleft()
        in_queue.discard((x, y))
        if _revise(csp, D, x, y, counter):
            if not D[x]:
                return False
            for z in csp.neighbors[x]:
                if z != y and (z, x) not in in_queue:
                    queue.append((z, x))
                    in_queue.add((z, x))
    return True


# ====================================================================== search
class _Limit(Exception):
    pass


@dataclass
class SolveResult:
    status: str                      # "solved" | "unsatisfiable" | "limit"
    solutions: List[Dict[Hashable, Any]]
    ac3: AC3Result
    nodes: int = 0
    backtracks: int = 0
    constraint_checks: int = 0
    elapsed_ms: float = 0.0
    exhaustive: bool = True          # False when stopped at max_solutions/limit

    @property
    def solution(self) -> Optional[Dict[Hashable, Any]]:
        return self.solutions[0] if self.solutions else None

    def metrics(self) -> Dict[str, Any]:
        return {
            "Status": self.status,
            "Solutions found": len(self.solutions) if self.exhaustive else f"{len(self.solutions)}+",
            "Search nodes (assignments tried)": self.nodes,
            "Backtracks": self.backtracks,
            "Constraint checks during search": self.constraint_checks,
            "Total time (ms)": round(self.elapsed_ms, 1),
        }


def solve(csp: CSP, max_solutions: int = 1, node_limit: int = 200_000, time_limit: float = 30.0,
          value_order: Optional[Callable[[Hashable, List[Any]], List[Any]]] = None) -> SolveResult:
    """AC-3 pre-processing followed by backtracking search that Maintains Arc Consistency (MAC).

    Variable ordering: MRV (smallest remaining domain) with degree tie-break.
    """
    t0 = time.perf_counter()
    first = ac3(csp)
    if not first.consistent:
        return SolveResult("unsatisfiable", [], first, elapsed_ms=(time.perf_counter() - t0) * 1000)

    solutions: List[Dict[Hashable, Any]] = []
    stats = {"nodes": 0, "backtracks": 0}
    counter = [0]
    deadline = t0 + time_limit
    degree = {v: len(csp.neighbors[v]) for v in csp.variables}

    def select(D):
        best, key = None, None
        for v in csp.variables:
            n = len(D[v])
            if n > 1:
                k = (n, -degree[v])
                if key is None or k < key:
                    best, key = v, k
        return best

    def rec(D) -> bool:
        var = select(D)
        if var is None:
            solutions.append({v: D[v][0] for v in csp.variables})
            return len(solutions) >= max_solutions
        values = value_order(var, D[var]) if value_order else D[var]
        for val in values:
            stats["nodes"] += 1
            if stats["nodes"] > node_limit or time.perf_counter() > deadline:
                raise _Limit()
            D2 = dict(D)
            D2[var] = [val]
            if _ac3_fast(csp, D2, deque((n, var) for n in csp.neighbors[var]), counter):
                if rec(D2):
                    return True
            stats["backtracks"] += 1
        return False

    status, exhaustive = "solved", True
    try:
        stopped_early = rec(first.domains)
        if stopped_early:
            exhaustive = False
    except _Limit:
        exhaustive = False
        status = "limit"
    if solutions:
        status = "solved"
    elif status != "limit":
        status = "unsatisfiable"
    return SolveResult(status, solutions, first, stats["nodes"], stats["backtracks"], counter[0],
                       (time.perf_counter() - t0) * 1000, exhaustive)


# =====================================================================================
# SECTION: core/report.py
# =====================================================================================
# Report model + exporters (PDF, Word, CSV, plain text).

import datetime as _dt
import io
import textwrap
from dataclasses import dataclass, field
from typing import List, Optional, Tuple

import pandas as pd




@dataclass
class Report:
    use_case: str
    dataset: str
    summary: str
    metrics: List[Tuple[str, str]] = field(default_factory=list)
    sections: List[Tuple[str, str]] = field(default_factory=list)       # (heading, paragraph text)
    formulas: List[Tuple[str, str]] = field(default_factory=list)       # (formula as text, meaning)
    tables: List[Tuple[str, pd.DataFrame]] = field(default_factory=list)
    figures: List[Tuple[str, bytes]] = field(default_factory=list)      # (caption, PNG bytes)
    primary_table: int = 0
    generated: str = field(default_factory=lambda: _dt.datetime.now().strftime("%Y-%m-%d %H:%M"))

    @property
    def title(self) -> str:
        return f"{self.use_case} - AC-3 Results Report"

    def primary(self) -> Optional[pd.DataFrame]:
        if not self.tables:
            return None
        return self.tables[min(self.primary_table, len(self.tables) - 1)][1]


def _s(x) -> str:
    """Stringify a cell, keeping it Latin-1 safe for the PDF core fonts."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    s = str(x)
    repl = {"→": "->", "←": "<-", "≤": "<=", "≥": ">=", "≠": "!=", "−": "-",
            "—": "-", "–": "-", "✓": "OK", "✗": "X", "∈": "in", "∧": "and",
            "∀": "for all", "∃": "exists", "×": "x", "…": "..."}
    for k, v in repl.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "replace").decode("latin-1")


# =============================================================== CSV
def to_csv(report: Report) -> bytes:
    df = report.primary()
    if df is None:
        df = pd.DataFrame(report.metrics, columns=["Metric", "Value"])
    return df.to_csv(index=False).encode("utf-8-sig")


# =============================================================== TXT
def _txt_table(df: pd.DataFrame, max_rows: int = 200, width: int = 28) -> str:
    d = df.head(max_rows).copy()
    d = d.astype(str).apply(lambda col: col.map(lambda v: v if len(v) <= width else v[: width - 3] + "..."))
    d.columns = [c if len(str(c)) <= width else str(c)[: width - 3] + "..." for c in d.columns]
    cols = list(d.columns)
    widths = [max(len(str(c)), *(len(v) for v in d[c])) if len(d) else len(str(c)) for c in cols]
    sep = "+" + "+".join("-" * (w + 2) for w in widths) + "+"
    head = "|" + "|".join(f" {str(c):<{w}} " for c, w in zip(cols, widths)) + "|"
    lines = [sep, head, sep.replace("-", "=")]
    for _, r in d.iterrows():
        lines.append("|" + "|".join(f" {str(r[c]):<{w}} " for c, w in zip(cols, widths)) + "|")
    lines.append(sep)
    if len(df) > max_rows:
        lines.append(f"  ... {len(df) - max_rows} more rows (see CSV export for the full table)")
    return "\n".join(lines)


def to_txt(report: Report) -> bytes:
    W = 100
    bar = "=" * W
    out = [bar, APP_NAME.upper().center(W), DEVELOPER_LINE.center(W), bar,
           f"Report      : {report.title}", f"Dataset     : {report.dataset}",
           f"Generated   : {report.generated}", f"Version     : {VERSION}", bar, "",
           "EXECUTIVE SUMMARY", "-" * 17]
    out += textwrap.wrap(report.summary, W)
    out += ["", "KEY METRICS", "-" * 11]
    kw = max((len(k) for k, _ in report.metrics), default=10)
    out += [f"  {k:<{kw}} : {v}" for k, v in report.metrics]
    for h, t in report.sections:
        out += ["", h.upper(), "-" * len(h)]
        for para in str(t).split("\n"):
            out += textwrap.wrap(para, W) or [""]
    if report.formulas:
        out += ["", "FORMULAS USED", "-" * 13]
        for f, m in report.formulas:
            out += [f"  * {f}"] + ["      " + ln for ln in textwrap.wrap(m, W - 6)]
    for name, df in report.tables:
        out += ["", f"TABLE: {name}  ({len(df)} rows)", ""]
        out.append(_txt_table(df))
    out += ["", bar, f"(c) {_dt.date.today().year} {ORG}  -  generated by {APP_NAME}".center(W), bar]
    return "\n".join(out).encode("utf-8")


# =============================================================== PDF
def to_pdf(report: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    blue = colors.HexColor("#0B4F9C")
    light = colors.HexColor("#E8F1FB")
    buf = io.BytesIO()
    page = landscape(letter)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch,
                            title=_s(report.title), author="Randy Singh - Kalsnet (KNet) Consulting")
    ss = getSampleStyleSheet()
    st_title = ParagraphStyle("t", parent=ss["Title"], textColor=blue, fontName="Helvetica-Bold", fontSize=24, leading=28)
    st_dev = ParagraphStyle("d", parent=ss["Title"], textColor=blue, fontName="Helvetica-Bold", fontSize=16, leading=20)
    st_h1 = ParagraphStyle("h1", parent=ss["Heading1"], textColor=blue, fontSize=16, spaceBefore=10)
    st_h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor("#07326A"), fontSize=12.5, spaceBefore=8)
    st_body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10, leading=13.5)
    st_cell = ParagraphStyle("c", parent=ss["BodyText"], fontSize=7.5, leading=9)
    st_hcell = ParagraphStyle("hc", parent=st_cell, textColor=colors.white, fontName="Helvetica-Bold")
    st_meta = ParagraphStyle("m", parent=st_body, alignment=TA_CENTER, textColor=colors.HexColor("#444444"))
    st_cap = ParagraphStyle("cap", parent=st_body, alignment=TA_CENTER, fontSize=9, textColor=colors.HexColor("#555555"))

    def esc(x):
        return _s(x).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = [Paragraph(esc(APP_NAME), st_title), Paragraph(esc(DEVELOPER_LINE), st_dev), Spacer(1, 6),
             Paragraph(f"<b>{esc(report.title)}</b> &nbsp;|&nbsp; Dataset: {esc(report.dataset)} "
                       f"&nbsp;|&nbsp; Generated: {esc(report.generated)}", st_meta), Spacer(1, 10)]

    story += [Paragraph("Executive summary", st_h1), Paragraph(esc(report.summary), st_body)]

    # metrics as a 4-column grid (2 metric pairs per row)
    m = [(Paragraph(f"<b>{esc(k)}</b>", st_cell), Paragraph(esc(v), st_cell)) for k, v in report.metrics]
    rows = []
    for i in range(0, len(m), 2):
        pair = list(m[i]) + (list(m[i + 1]) if i + 1 < len(m) else ["", ""])
        rows.append(pair)
    if rows:
        avail = page[0] - 1.2 * inch
        t = Table(rows, colWidths=[avail * 0.3, avail * 0.2, avail * 0.3, avail * 0.2])
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7CBE3")),
            ("BACKGROUND", (0, 0), (0, -1), light), ("BACKGROUND", (2, 0), (2, -1), light),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story += [Paragraph("Key metrics", st_h1), t]

    for h, txt in report.sections:
        story.append(Paragraph(esc(h), st_h2))
        for para in str(txt).split("\n"):
            if para.strip():
                story.append(Paragraph(esc(para), st_body))

    if report.formulas:
        story.append(Paragraph("Formulas used", st_h1))
        fr = [[Paragraph("<b>Formula</b>", st_hcell), Paragraph("<b>Meaning</b>", st_hcell)]]
        fr += [[Paragraph(f"<font face='Courier'>{esc(f)}</font>", st_cell), Paragraph(esc(mn), st_cell)]
               for f, mn in report.formulas]
        avail = page[0] - 1.2 * inch
        ft = Table(fr, colWidths=[avail * 0.42, avail * 0.58], repeatRows=1)
        ft.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), blue),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B7CBE3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(ft)

    for cap, png in report.figures:
        img = Image(io.BytesIO(png))
        max_w, max_h = page[0] - 1.4 * inch, page[1] - 2.2 * inch
        scale = min(max_w / img.imageWidth, max_h / img.imageHeight, 1.0)
        img.drawWidth, img.drawHeight = img.imageWidth * scale, img.imageHeight * scale
        story += [PageBreak(), KeepTogether([Paragraph(esc(cap), st_h2), img])]

    max_rows = 80
    for name, df in report.tables:
        story += [PageBreak(), Paragraph(esc(name), st_h1)]
        d = df.head(max_rows)
        ncol = max(1, len(d.columns))
        avail = page[0] - 1.2 * inch
        data = [[Paragraph(esc(c), st_hcell) for c in d.columns]]
        data += [[Paragraph(esc(v), st_cell) for v in row] for row in d.itertuples(index=False)]
        t = Table(data, colWidths=[avail / ncol] * ncol, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), blue),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, light]),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B7CBE3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(t)
        if len(df) > max_rows:
            story.append(Paragraph(f"<i>Showing first {max_rows} of {len(df)} rows - the CSV export contains every row.</i>", st_body))

    def on_page(canvas, doc_):
        canvas.saveState()
        canvas.setFillColor(blue)
        canvas.rect(0, page[1] - 0.28 * inch, page[0], 0.28 * inch, fill=1, stroke=0)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(0.6 * inch, 0.35 * inch, _s(f"{APP_NAME} - {ORG}"))
        canvas.drawRightString(page[0] - 0.6 * inch, 0.35 * inch, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# =============================================================== DOCX
def to_docx(report: Report) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    blue = RGBColor(0x0B, 0x4F, 0x9C)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    def para(text, size=10.5, bold=False, color=None, align=None, italic=False):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size, r.bold, r.italic = Pt(size), bold, italic
        if color is not None:
            r.font.color.rgb = color
        if align is not None:
            p.alignment = align
        return p

    def heading(text, size=15):
        p = para(text, size=size, bold=True, color=blue)
        p.paragraph_format.space_before = Pt(10)
        return p

    para(APP_NAME, 26, True, blue, WD_ALIGN_PARAGRAPH.CENTER)
    para(DEVELOPER_LINE, 17, True, blue, WD_ALIGN_PARAGRAPH.CENTER)
    para(f"{report.title}  |  Dataset: {report.dataset}  |  Generated: {report.generated}", 10,
         color=RGBColor(0x44, 0x44, 0x44), align=WD_ALIGN_PARAGRAPH.CENTER)

    heading("Executive summary")
    para(report.summary)

    heading("Key metrics")
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    for i, h in enumerate(("Metric", "Value")):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold, run.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "0B4F9C")
    for i, (k, v) in enumerate(report.metrics):
        cells = t.add_row().cells
        cells[0].text, cells[1].text = str(k), str(v)
        if i % 2:
            shade(cells[0], "E8F1FB"); shade(cells[1], "E8F1FB")

    for h, txt in report.sections:
        heading(h, 13)
        for p in str(txt).split("\n"):
            if p.strip():
                para(p)

    if report.formulas:
        heading("Formulas used")
        ft = doc.add_table(rows=1, cols=2)
        ft.style = "Table Grid"
        for i, h in enumerate(("Formula", "Meaning")):
            c = ft.rows[0].cells[i]
            c.text = ""
            run = c.paragraphs[0].add_run(h)
            run.bold, run.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
            shade(c, "0B4F9C")
        for f, m in report.formulas:
            cells = ft.add_row().cells
            cells[0].text = ""
            r = cells[0].paragraphs[0].add_run(f)
            r.font.name = "Consolas"
            cells[1].text = m

    for cap, png in report.figures:
        heading(cap, 13)
        doc.add_picture(io.BytesIO(png), width=Inches(8.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    max_rows = 150
    for name, df in report.tables:
        heading(f"{name}  ({len(df)} rows)", 13)
        d = df.head(max_rows)
        tb = doc.add_table(rows=1, cols=len(d.columns))
        tb.style = "Table Grid"
        for i, c in enumerate(d.columns):
            cell = tb.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(c))
            run.bold, run.font.size, run.font.color.rgb = True, Pt(8.5), RGBColor(0xFF, 0xFF, 0xFF)
            shade(cell, "0B4F9C")
        for ri, row in enumerate(d.itertuples(index=False)):
            cells = tb.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run("" if v is None else str(v))
                run.font.size = Pt(8.5)
                if ri % 2:
                    shade(cells[i], "E8F1FB")
        if len(df) > max_rows:
            para(f"Showing first {max_rows} of {len(df)} rows - the CSV export contains every row.", 9, italic=True)

    footer = sec.footer.paragraphs[0]
    footer.text = f"{APP_NAME}  |  {ORG}"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def fig_to_png(fig, dpi: int = 130) -> bytes:
    """Matplotlib figure -> PNG bytes (and close it)."""
    import matplotlib.pyplot as plt
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


# =====================================================================================
# SECTION: core/common.py
# =====================================================================================
# Chart + table helpers shared by all six use cases.

from typing import List, Tuple

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
import plotly.graph_objects as go  # noqa: E402





AC3_FORMULAS: List[Tuple[str, str]] = [
    ("Arc (Xi -> Xj) is consistent  <=>  for all a in Di, exists b in Dj : Cij(a, b)",
     "Every value left in the domain of Xi has at least one compatible partner value in Xj."),
    ("REVISE(Xi, Xj):  Di := { a in Di | exists b in Dj : Cij(a, b) }",
     "Deletes the unsupported values of Xi. If Di changes, every arc (Xk -> Xi), k != j, is put back on the queue."),
    ("Domain reduction % = 100 x (sum|Di| before - sum|Di| after) / sum|Di| before",
     "How much of the candidate space AC-3 eliminated without any guessing."),
    ("Search space = product over i of |Di|   (reported as log10)",
     "Number of complete assignments a brute-force search would have to consider."),
    ("Worst-case cost of AC-3 = O(e x d^3)",
     "e = number of arcs (2 per binary constraint), d = largest domain size."),
]


def trace_df(res: AC3Result, limit: int = 500) -> pd.DataFrame:
    df = res.trace_table(limit)
    if df.empty:
        return pd.DataFrame(columns=["Step", "Arc (Xi -> Xj)", "Constraint", "|Di| before",
                                     "Removed from Di", "#Removed", "|Di| after", "Arcs re-queued"])
    return df[[c for c in df.columns if not c.startswith("_")]]


def metrics_pairs(res: AC3Result, search: SolveResult | None = None, extra: dict | None = None) -> List[Tuple[str, str]]:
    pairs = [(k, str(v)) for k, v in (extra or {}).items()]
    pairs += [(k, str(v)) for k, v in res.metrics().items()]
    if search is not None:
        pairs += [(k, str(v)) for k, v in search.metrics().items()]
    return pairs


def domain_reduction_plotly(res: AC3Result, title: str = "Domain size per variable: before vs after AC-3",
                            max_vars: int = 60, label_map=None) -> go.Figure:
    names = list(res.domains.keys())[:max_vars]
    lab = [str(label_map(n)) if label_map else str(n) for n in names]
    orig = [len(res.original_domains.get(n, [])) for n in names]
    before = [len(res.initial_domains[n]) for n in names]
    after = [len(res.domains[n]) for n in names]
    fig = go.Figure()
    if any(o != b for o, b in zip(orig, before)):
        fig.add_bar(x=lab, y=orig, name="Original domain", marker_color="#B8C7DA")
    fig.add_bar(x=lab, y=before, name="After unary rules", marker_color=BLUE)
    fig.add_bar(x=lab, y=after, name="After AC-3", marker_color=ACCENT)
    fig.update_layout(title=title, barmode="group", height=380, margin=dict(l=10, r=10, t=50, b=10),
                      legend=dict(orientation="h", y=1.1, x=0), plot_bgcolor="white",
                      yaxis_title="|D| (values)", xaxis_tickangle=-45)
    fig.update_yaxes(gridcolor="#E5ECF6")
    return fig


def domain_reduction_png(res: AC3Result, title: str = "Domain size before vs after AC-3",
                         max_vars: int = 40, label_map=None) -> bytes:
    names = list(res.domains.keys())[:max_vars]
    lab = [str(label_map(n)) if label_map else str(n) for n in names]
    before = [len(res.initial_domains[n]) for n in names]
    after = [len(res.domains[n]) for n in names]
    fig, ax = plt.subplots(figsize=(11, 4.2))
    x = range(len(names))
    ax.bar([i - 0.2 for i in x], before, width=0.4, label="Before AC-3", color=BLUE)
    ax.bar([i + 0.2 for i in x], after, width=0.4, label="After AC-3", color=ACCENT)
    ax.set_xticks(list(x))
    ax.set_xticklabels(lab, rotation=60, ha="right", fontsize=7)
    ax.set_ylabel("Domain size |D|")
    ax.set_title(title, color=BLUE, fontweight="bold")
    ax.legend()
    ax.grid(axis="y", alpha=0.3)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    return fig_to_png(fig)


def pruning_timeline_plotly(res: AC3Result) -> go.Figure:
    """Cumulative values pruned as AC-3 processes arcs."""
    steps, cum, tot = [0], [0], 0
    for t in res.trace:
        tot += t["#Removed"]
        steps.append(t["Step"])
        cum.append(tot)
    fig = go.Figure(go.Scatter(x=steps, y=cum, mode="lines", line=dict(color=GREEN, width=3),
                               fill="tozeroy", fillcolor="rgba(30,142,90,0.15)"))
    fig.update_layout(title="AC-3 progress: cumulative values pruned vs arcs processed", height=300,
                      xaxis_title="Arcs processed", yaxis_title="Values pruned",
                      margin=dict(l=10, r=10, t=50, b=10), plot_bgcolor="white")
    fig.update_yaxes(gridcolor="#E5ECF6")
    return fig


def constraint_graph_dot(csp: CSP, max_nodes: int = 40, label_map=None) -> str:
    """Graphviz DOT for the constraint graph (variables = nodes, constraints = edges)."""
    vs = csp.variables[:max_nodes]
    keep = set(vs)
    lines = ["graph G {", 'graph [layout=neato, overlap=false, splines=true, bgcolor="transparent"];',
             'node [shape=ellipse, style=filled, fillcolor="#E8F1FB", color="#0B4F9C", fontname="Helvetica", fontsize=10];',
             'edge [color="#7FA7D6"];']
    for v in vs:
        lab = label_map(v) if label_map else v
        lines.append(f'"{v}" [label="{lab}"];')
    seen = set()
    for (x, y) in csp.arcs():
        if x in keep and y in keep and frozenset((x, y)) not in seen:
            seen.add(frozenset((x, y)))
            lines.append(f'"{x}" -- "{y}";')
    lines.append("}")
    return "\n".join(lines)


# =====================================================================================
# USE CASE 1 - SUDOKU & LOGIC PUZZLES
# =====================================================================================
import random as _random

SUDOKU_SAMPLES = {
    "Easy - solved by AC-3 alone (AIMA classic)":
        "..3.2.6..9..3.5..1..18.64....81.29..7.......8..67.82....26.95..8..2.3..9..5.1.3..",
    "Medium - AC-3 + a little search (26 clues)":
        ".8..3.6.....6.2..72...8.9..5.....16.3.......8.97.....5..8.2...46..5.4.....5.1..2.",
    "Hard - 'World's hardest Sudoku' (A. Inkala)":
        "8..........36......7..9.2...5...7.......457.....1...3...1....68..85...1..9....4..",
}

SUDOKU_SCHEMA = {
    "sudoku.csv  (9 x 9 grid, no header)": pd.DataFrame([
        ["Column 1 ... Column 9", "integer 0-9", "Yes", "One CSV row per Sudoku row. 0 (or blank / '.') = empty cell; 1-9 = given clue.", "5,3,0,0,7,0,0,0,0"],
    ], columns=["Field", "Type", "Required", "Description", "Example"]),
    "sudoku.txt  (alternative)": pd.DataFrame([
        ["puzzle", "81-character string", "Yes", "Cells read row by row, left to right. '.' or '0' = empty. One puzzle per line; several lines = several puzzles to choose from.", "..3.2.6..9..3.5..1.."],
    ], columns=["Field", "Type", "Required", "Description", "Example"]),
}


def sudoku_parse_string(s: str) -> List[int]:
    s = "".join(ch for ch in s.strip() if ch in "0123456789.")
    if len(s) != 81:
        raise ValueError(f"A Sudoku needs exactly 81 cells, found {len(s)}.")
    return [0 if ch in ".0" else int(ch) for ch in s]


def sudoku_parse_upload(name: str, raw: bytes) -> List[List[int]]:
    """Return a list of puzzles (each a flat list of 81 ints)."""
    text = raw.decode("utf-8-sig", errors="replace")
    if name.lower().endswith(".csv"):
        rows = [r for r in text.replace(";", ",").splitlines() if r.strip()]
        if len(rows) == 10 and rows[0].replace(" ", "") in ("0,1,2,3,4,5,6,7,8", "C1,C2,C3,C4,C5,C6,C7,C8,C9"):
            rows = rows[1:]                       # tolerate a header row (e.g. the downloadable sample)
        cells = []
        for r in rows:
            for c in r.split(","):
                c = c.strip()
                cells.append(0 if c in ("", ".", "0") else int(c))
        if len(cells) != 81:
            raise ValueError(f"CSV must be 9 rows x 9 columns (81 cells); found {len(cells)} cells.")
        return [cells]
    puzzles = []
    for line in text.splitlines():
        clean = "".join(ch for ch in line if ch in "0123456789.")
        if len(clean) == 81:
            puzzles.append(sudoku_parse_string(clean))
    if not puzzles:
        puzzles = [sudoku_parse_string(text)]
    return puzzles


def sudoku_validate(grid: List[int]) -> List[str]:
    errs = []
    if len(grid) != 81:
        return [f"Grid has {len(grid)} cells, expected 81."]
    if any((not isinstance(v, int)) or v < 0 or v > 9 for v in grid):
        errs.append("Every cell must be an integer 0-9.")
    for i in range(81):
        for j in range(i + 1, 81):
            if grid[i] and grid[i] == grid[j] and j in _SUDOKU_PEERS[i]:
                errs.append(f"Duplicate clue {grid[i]} at R{i//9+1}C{i%9+1} and R{j//9+1}C{j%9+1}.")
    return errs


def _sudoku_peers():
    peers = []
    for i in range(81):
        r, c = divmod(i, 9)
        p = set()
        for k in range(9):
            p.add(r * 9 + k)
            p.add(k * 9 + c)
        br, bc = 3 * (r // 3), 3 * (c // 3)
        for rr in range(br, br + 3):
            for cc in range(bc, bc + 3):
                p.add(rr * 9 + cc)
        p.discard(i)
        peers.append(p)
    return peers


_SUDOKU_PEERS = _sudoku_peers()


def _cell(i: int) -> str:
    return f"R{i//9+1}C{i%9+1}"


def sudoku_build_csp(grid: List[int]) -> CSP:
    vars_ = [_cell(i) for i in range(81)]
    doms = {_cell(i): list(range(1, 10)) for i in range(81)}
    csp = CSP(vars_, doms, "Sudoku")
    for i, v in enumerate(grid):
        if v:
            csp.add_unary(_cell(i), lambda x, v=v: x == v, f"given clue = {v}")
    for i in range(81):
        r, c = divmod(i, 9)
        for j in _SUDOKU_PEERS[i]:
            if j > i:
                r2, c2 = divmod(j, 9)
                why = "same row" if r == r2 else "same column" if c == c2 else "same 3x3 box"
                csp.add_binary(_cell(i), _cell(j), lambda a, b: a != b, f"{why}: values differ")
    return csp


def sudoku_is_valid_solution(sol: List[int], givens: Optional[List[int]] = None) -> bool:
    if len(sol) != 81 or any(v not in range(1, 10) for v in sol):
        return False
    groups = [[r * 9 + c for c in range(9)] for r in range(9)] + [[r * 9 + c for r in range(9)] for c in range(9)]
    groups += [[(br + r) * 9 + bc + c for r in range(3) for c in range(3)] for br in (0, 3, 6) for bc in (0, 3, 6)]
    if any(sorted(sol[i] for i in g) != list(range(1, 10)) for g in groups):
        return False
    return not givens or all(g == 0 or g == s for g, s in zip(givens, sol))


def sudoku_generate(seed: int = 7, clues: int = 32, ac3_only: bool = False) -> Tuple[List[int], List[int]]:
    """Synthetic puzzle: random full grid, then remove clues while the puzzle stays unique
    (or, with ac3_only=True, while AC-3 alone can still solve it)."""
    rng = _random.Random(seed)
    base = [[(3 * (r % 3) + r // 3 + c) % 9 + 1 for c in range(9)] for r in range(9)]
    bands = rng.sample(range(3), 3)
    rows = [b * 3 + r for b in bands for r in rng.sample(range(3), 3)]
    stacks = rng.sample(range(3), 3)
    cols = [s * 3 + c for s in stacks for c in rng.sample(range(3), 3)]
    digits = rng.sample(range(1, 10), 9)
    solution = [digits[base[r][c] - 1] for r in rows for c in cols]
    puzzle = solution[:]
    order = list(range(81))
    rng.shuffle(order)
    for idx in order:
        if sum(1 for v in puzzle if v) <= clues:
            break
        pair = {idx, 80 - idx}
        backup = {k: puzzle[k] for k in pair}
        for k in pair:
            puzzle[k] = 0
        csp = sudoku_build_csp(puzzle)
        if ac3_only:
            ok = ac3(csp, record_trace=False).solved
        else:
            ok = len(solve(csp, max_solutions=2, time_limit=5).solutions) == 1
        if not ok:
            for k, v in backup.items():
                puzzle[k] = v
    return puzzle, solution


def sudoku_run(grid: List[int], dataset: str = "") -> dict:
    csp = sudoku_build_csp(grid)
    res = solve(csp, max_solutions=2, time_limit=30)
    a = res.ac3
    sol = [res.solution[_cell(i)] for i in range(81)] if res.solution else None
    source = []
    for i in range(81):
        if grid[i]:
            source.append("Given")
        elif len(a.domains[_cell(i)]) == 1:
            source.append("AC-3")
        else:
            source.append("Search")
    cells = pd.DataFrame([{
        "Cell": _cell(i), "Row": i // 9 + 1, "Column": i % 9 + 1, "Box": 3 * (i // 27) + (i % 9) // 3 + 1,
        "Given": str(grid[i]) if grid[i] else "",
        "Candidates before AC-3": " ".join(map(str, a.initial_domains[_cell(i)])),
        "Candidates after AC-3": " ".join(map(str, a.domains[_cell(i)])),
        "# after AC-3": len(a.domains[_cell(i)]),
        "Final digit": str(sol[i]) if sol else "",
        "Obtained by": source[i],
    } for i in range(81)])
    grid_df = pd.DataFrame([[sol[r * 9 + c] if sol else "" for c in range(9)] for r in range(9)],
                           columns=[f"C{c+1}" for c in range(9)], index=[f"R{r+1}" for r in range(9)])
    unique = res.exhaustive and len(res.solutions) == 1
    return {"csp": csp, "search": res, "ac3": a, "grid": grid, "solution": sol, "cells": cells,
            "grid_df": grid_df, "source": source, "unique": unique, "dataset": dataset,
            "valid": bool(sol) and sudoku_is_valid_solution(sol, grid)}


def sudoku_plot(grid, values, source=None, candidates=None, title="") -> go.Figure:
    colors_map = {"Given": "#0B4F9C", "AC-3": "#1E8E5A", "Search": "#F28C28", None: "#333"}
    z = [[0] * 9 for _ in range(9)]
    fig = go.Figure()
    for i in range(81):
        r, c = divmod(i, 9)
        if candidates is not None:
            z[r][c] = len(candidates[i])
    if candidates is not None:
        fig.add_trace(go.Heatmap(z=z, colorscale=[[0, "#FFFFFF"], [0.12, "#E8F1FB"], [1, "#F9C98F"]], zmin=1, zmax=9,
                                 showscale=True, colorbar=dict(title="# candidates"), hoverinfo="skip", xgap=1, ygap=1))
    else:
        fig.add_trace(go.Heatmap(z=[[1 if (source and source[r*9+c] == "Given") else 0 for c in range(9)] for r in range(9)],
                                 colorscale=[[0, "#FFFFFF"], [1, "#EAF2FC"]], showscale=False, hoverinfo="skip", xgap=1, ygap=1))
    for i in range(81):
        r, c = divmod(i, 9)
        if candidates is not None and len(candidates[i]) > 1:
            txt = "<br>".join(" ".join(str(d) for d in candidates[i][k:k + 3]) for k in range(0, len(candidates[i]), 3))
            fig.add_annotation(x=c, y=r, text=txt, showarrow=False, font=dict(size=9, color="#8A5A00"))
        else:
            v = values[i] if values else (candidates[i][0] if candidates else "")
            src = source[i] if source else None
            fig.add_annotation(x=c, y=r, text=f"<b>{v}</b>" if v else "", showarrow=False,
                               font=dict(size=20, color=colors_map.get(src, "#333")))
    for k in range(10):
        w = 3 if k % 3 == 0 else 0.6
        fig.add_shape(type="line", x0=k - 0.5, x1=k - 0.5, y0=-0.5, y1=8.5, line=dict(color="#0B4F9C", width=w))
        fig.add_shape(type="line", y0=k - 0.5, y1=k - 0.5, x0=-0.5, x1=8.5, line=dict(color="#0B4F9C", width=w))
    fig.update_layout(title=title, height=520, width=560, margin=dict(l=10, r=10, t=50, b=10), plot_bgcolor="white",
                      xaxis=dict(showticklabels=False, showgrid=False, zeroline=False, range=[-0.5, 8.5]),
                      yaxis=dict(showticklabels=False, showgrid=False, zeroline=False, autorange="reversed",
                                 scaleanchor="x", range=[-0.5, 8.5]))
    return fig


def sudoku_png(grid, values, source, title) -> bytes:
    fig, ax = plt.subplots(figsize=(5.5, 5.8))
    cmap = {"Given": "#0B4F9C", "AC-3": "#1E8E5A", "Search": "#F28C28"}
    for i in range(81):
        r, c = divmod(i, 9)
        if source[i] == "Given":
            ax.add_patch(plt.Rectangle((c, 8 - r), 1, 1, color="#EAF2FC"))
        if values and values[i]:
            ax.text(c + 0.5, 8 - r + 0.5, str(values[i]), ha="center", va="center", fontsize=15,
                    color=cmap[source[i]], fontweight="bold")
    for k in range(10):
        lw = 2.5 if k % 3 == 0 else 0.6
        ax.plot([k, k], [0, 9], color="#0B4F9C", lw=lw)
        ax.plot([0, 9], [k, k], color="#0B4F9C", lw=lw)
    ax.set_xlim(0, 9); ax.set_ylim(0, 9); ax.set_aspect("equal"); ax.axis("off")
    ax.set_title(title, color="#0B4F9C", fontweight="bold")
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color=v, label=k) for k, v in cmap.items()], loc="upper center",
              bbox_to_anchor=(0.5, -0.02), ncol=3, frameon=False)
    return fig_to_png(fig)


def sudoku_report(r: dict) -> Report:
    a, s = r["ac3"], r["search"]
    n_given = sum(1 for v in r["grid"] if v)
    n_ac3 = r["source"].count("AC-3")
    n_search = r["source"].count("Search")
    if r["solution"] is None:
        summary = "The puzzle has no solution: AC-3 or the search proved the clues are contradictory."
    else:
        summary = (f"The puzzle had {n_given} clues. AC-3 alone fixed {n_ac3} of the {81 - n_given} empty cells "
                   f"({'the whole puzzle' if a.solved else 'the rest needed backtracking search'}). "
                   f"{n_search} cells were filled by MAC search using {s.nodes} trial assignments. "
                   f"The solution is {'unique and ' if r['unique'] else ''}verified valid: every row, column and "
                   f"3x3 box contains the digits 1-9 exactly once.")
    return Report(
        use_case="Sudoku & Logic Puzzles", dataset=r["dataset"], summary=summary,
        metrics=metrics_pairs(a, s, {"Clues given": n_given, "Cells fixed by AC-3": n_ac3,
                                     "Cells needing search": n_search, "Unique solution": "Yes" if r["unique"] else "No/unknown",
                                     "Solution verified": "Yes" if r["valid"] else "No"}),
        sections=[("How to read the results",
                   "Blue digits were given. Green digits were deduced by AC-3 purely from the 'all different' constraints "
                   "in rows, columns and boxes. Orange digits needed search (guess + propagate + backtrack). The table "
                   "'Cell details' lists each cell's candidates before and after AC-3.")],
        formulas=[("Xi != Xj  for every pair of peers (same row, column or box)", "Binary 'not-equal' constraint; 810 pairs = 1,620 arcs."),
                  ("Di = {1..9}, or Di = {given} for a clue", "Unary constraint from the clue fixes the domain.")] + AC3_FORMULAS,
        tables=[("Cell details", r["cells"]), ("Solution grid", r["grid_df"].reset_index().rename(columns={"index": "Row"})),
                ("AC-3 trace (first 300 revisions)", trace_df(a, 300))],
        figures=[("Solved grid (blue = given, green = AC-3, orange = search)",
                  sudoku_png(r["grid"], r["solution"], r["source"], "Solved Sudoku")),
                 ("Domain sizes per cell before vs after AC-3", domain_reduction_png(a, max_vars=81))],
    )


# =====================================================================================
# USE CASE 2 - SCHEDULING & TIMETABLING
# =====================================================================================
def _schema(rows):
    return pd.DataFrame(rows, columns=["Field", "Type", "Required", "Description", "Example"])


SCHED_SCHEMA = {
    "events.csv": _schema([
        ["event_id", "text", "Yes", "Unique code of the exam / class / shift / meeting (becomes a CSP variable).", "CS101"],
        ["event_name", "text", "Yes", "Human-readable name.", "Intro to Programming"],
        ["organizer", "text", "Yes", "Instructor / manager who must attend. Two events with the same organizer can't share a timeslot.", "Dr. Patel"],
        ["expected_size", "integer", "No", "Seats needed. If blank, it's counted from attendance.csv. Rooms smaller than this are removed (unary rule).", "85"],
        ["blocked_slots", "text list ';'", "No", "Timeslot ids the event can NOT use (organizer unavailable).", "MON-AM;FRI-PM"],
    ]),
    "attendance.csv": _schema([
        ["person_id", "text", "Yes", "Student / employee id.", "S0042"],
        ["event_id", "text", "Yes", "Event this person attends (must exist in events.csv). Two events that share a person can't share a timeslot.", "CS101"],
    ]),
    "rooms.csv": _schema([
        ["room_id", "text", "Yes", "Unique room code.", "HALL-A"],
        ["capacity", "integer", "Yes", "Number of seats.", "150"],
    ]),
    "timeslots.csv": _schema([
        ["slot_id", "text", "Yes", "Unique timeslot code.", "MON-AM"],
        ["day", "text", "Yes", "Day label used for the timetable display.", "Monday"],
        ["start_time", "text", "Yes", "Start time label.", "09:00"],
    ]),
}


def sched_synthetic(kind: str = "University exam timetabling", seed: int = 11, n_people: int = 240) -> Dict[str, pd.DataFrame]:
    rng = _random.Random(seed)
    if kind.startswith("University"):
        programs = {
            "Computer Science": [("CS101", "Intro to Programming", "Dr. Patel"), ("CS210", "Data Structures", "Dr. Nguyen"),
                                 ("CS305", "Algorithms", "Dr. Patel"), ("CS340", "Artificial Intelligence", "Dr. Okafor")],
            "Mathematics": [("MA101", "Calculus I", "Prof. Chen"), ("MA201", "Linear Algebra", "Prof. Chen"),
                            ("MA250", "Probability", "Dr. Silva"), ("MA330", "Numerical Methods", "Dr. Silva")],
            "Business": [("BU110", "Accounting Basics", "Ms. Garcia"), ("BU220", "Marketing", "Mr. Brown"),
                         ("BU310", "Corporate Finance", "Ms. Garcia"), ("BU330", "Operations Mgmt", "Mr. Brown")],
            "Biology": [("BI101", "Cell Biology", "Dr. Adeyemi"), ("BI205", "Genetics", "Dr. Adeyemi"),
                        ("BI260", "Ecology", "Dr. Kowalski"), ("BI320", "Biochemistry", "Dr. Kowalski")],
        }
        electives = [("GE100", "Academic Writing", "Ms. Larsen"), ("GE150", "Ethics & Society", "Mr. Haddad")]
        slots = [("MON-AM", "Monday", "09:00"), ("MON-PM", "Monday", "14:00"), ("TUE-AM", "Tuesday", "09:00"),
                 ("TUE-PM", "Tuesday", "14:00"), ("WED-AM", "Wednesday", "09:00"), ("WED-PM", "Wednesday", "14:00"),
                 ("THU-AM", "Thursday", "09:00"), ("THU-PM", "Thursday", "14:00")]
        rooms = [("HALL-A", 150), ("HALL-B", 100), ("ROOM-201", 60), ("LAB-3", 40)]
        prefix = "S"
    else:
        programs = {
            "Finance": [("M01", "Quarterly Budget Review", "CFO Rivera"), ("M02", "Audit Prep", "Controller Kim"),
                        ("M03", "Forecast Sync", "Controller Kim")],
            "Product": [("M04", "Roadmap Planning", "VP Shah"), ("M05", "Design Critique", "Lead Moreau"),
                        ("M06", "Sprint Review", "Lead Moreau")],
            "Sales": [("M07", "Pipeline Review", "VP Johnson"), ("M08", "Key Account Strategy", "VP Johnson"),
                      ("M09", "Pricing Committee", "CFO Rivera")],
            "People": [("M10", "Hiring Committee", "HR Dubois"), ("M11", "Training Workshop", "HR Dubois")],
        }
        electives = [("M12", "All-hands Prep", "CEO Tanaka")]
        slots = [(f"{d[:3].upper()}-{t.replace(':', '')}", d, t) for d in ("Monday", "Tuesday", "Wednesday")
                 for t in ("09:00", "11:00", "14:00")]
        rooms = [("BOARDROOM", 20), ("HUDDLE-1", 8), ("HUDDLE-2", 8), ("TRAINING", 40)]
        prefix = "E"
        n_people = min(n_people, 60)

    events = []
    for prog, lst in programs.items():
        for e in lst:
            events.append({"event_id": e[0], "event_name": e[1], "organizer": e[2], "program": prog})
    for e in electives:
        events.append({"event_id": e[0], "event_name": e[1], "organizer": e[2], "program": "Elective"})
    att = []
    prog_names = list(programs)
    for p in range(1, n_people + 1):
        pid = f"{prefix}{p:04d}"
        prog = prog_names[(p - 1) % len(prog_names)]
        courses = [e[0] for e in programs[prog]]
        # everyone takes their programme's first (intro) course + 1-2 others
        chosen = [courses[0]] + rng.sample(courses[1:], rng.randint(1, min(2, len(courses) - 1)))
        if kind.startswith("University") and prog == "Computer Science":
            chosen.append("MA101")            # maths requirement -> MA101 becomes a big exam
        if rng.random() < 0.35:
            chosen.append(rng.choice(electives)[0])
        att += [{"person_id": pid, "event_id": c} for c in chosen]
    att_df = pd.DataFrame(att)
    sizes = att_df.groupby("event_id").size()
    ev = pd.DataFrame(events)
    ev["expected_size"] = ev["event_id"].map(sizes).fillna(0).astype(int)
    # organiser availability: event #2 can't use the first two slots, event #7 not the last one,
    # and the first elective is fixed by the dean to the first slot (all other slots blocked)
    blocked = {ev.iloc[1]["event_id"]: slots[0][0] + ";" + slots[1][0], ev.iloc[6]["event_id"]: slots[-1][0],
               electives[0][0]: ";".join(s_[0] for s_ in slots[1:])}
    ev["blocked_slots"] = ev["event_id"].map(blocked).fillna("")
    ev = ev[["event_id", "event_name", "organizer", "expected_size", "blocked_slots", "program"]]
    return {"events.csv": ev, "attendance.csv": att_df,
            "rooms.csv": pd.DataFrame(rooms, columns=["room_id", "capacity"]),
            "timeslots.csv": pd.DataFrame(slots, columns=["slot_id", "day", "start_time"])}


def _need_cols(df: pd.DataFrame, cols: List[str], fname: str) -> List[str]:
    miss = [c for c in cols if c not in df.columns]
    return [f"{fname}: missing required column(s) {miss}."] if miss else []


def sched_validate(data: Dict[str, pd.DataFrame]) -> List[str]:
    errs = []
    for f in SCHED_SCHEMA:
        if f not in data:
            errs.append(f"Missing file {f}.")
    if errs:
        return errs
    errs += _need_cols(data["events.csv"], ["event_id", "event_name", "organizer"], "events.csv")
    errs += _need_cols(data["attendance.csv"], ["person_id", "event_id"], "attendance.csv")
    errs += _need_cols(data["rooms.csv"], ["room_id", "capacity"], "rooms.csv")
    errs += _need_cols(data["timeslots.csv"], ["slot_id", "day", "start_time"], "timeslots.csv")
    if errs:
        return errs
    ev = data["events.csv"]
    if ev["event_id"].duplicated().any():
        errs.append(f"events.csv: duplicate event_id {ev.loc[ev['event_id'].duplicated(), 'event_id'].tolist()[:5]}.")
    unknown = set(data["attendance.csv"]["event_id"].astype(str)) - set(ev["event_id"].astype(str))
    if unknown:
        errs.append(f"attendance.csv references unknown event_id(s): {sorted(unknown)[:5]}.")
    if not pd.to_numeric(data["rooms.csv"]["capacity"], errors="coerce").notna().all():
        errs.append("rooms.csv: capacity must be numeric.")
    if data["timeslots.csv"]["slot_id"].duplicated().any():
        errs.append("timeslots.csv: duplicate slot_id.")
    return errs


def sched_build(data: Dict[str, pd.DataFrame]):
    ev = data["events.csv"].copy()
    ev["event_id"] = ev["event_id"].astype(str)
    att = data["attendance.csv"].astype(str)
    rooms = data["rooms.csv"].copy()
    rooms["capacity"] = pd.to_numeric(rooms["capacity"]).astype(int)
    slots = data["timeslots.csv"].astype(str)
    counted = att.groupby("event_id").size()
    if "expected_size" not in ev.columns:
        ev["expected_size"] = None
    ev["size"] = [int(s) if pd.notna(s) and str(s).strip() != "" else int(counted.get(e, 0))
                  for e, s in zip(ev["event_id"], ev["expected_size"])]
    cap = dict(zip(rooms["room_id"].astype(str), rooms["capacity"]))
    slot_ids = list(slots["slot_id"])
    room_ids = list(cap)
    # value ordering: best-fit room (smallest adequate) first
    values = sorted([(s, r) for s in slot_ids for r in room_ids], key=lambda v: (cap[v[1]], slot_ids.index(v[0])))
    csp = CSP(list(ev["event_id"]), {e: values for e in ev["event_id"]}, "Timetabling")
    csp.value_formatter = lambda var, v: f"{v[0]}@{v[1]}"
    size = dict(zip(ev["event_id"], ev["size"]))
    for e in ev["event_id"]:
        csp.add_unary(e, lambda v, n=size[e]: cap[v[1]] >= n, f"room capacity >= {size[e]}")
    if "blocked_slots" in ev.columns:
        for e, b in zip(ev["event_id"], ev["blocked_slots"].fillna("").astype(str)):
            bl = {x.strip() for x in b.split(";") if x.strip()}
            if bl:
                csp.add_unary(e, lambda v, bl=bl: v[0] not in bl, f"blocked slots {sorted(bl)}")
    people = att.groupby("event_id")["person_id"].apply(set).to_dict()
    org = dict(zip(ev["event_id"], ev["organizer"].astype(str)))
    ids = list(ev["event_id"])
    shared_pairs = []
    for i, a in enumerate(ids):
        for b in ids[i + 1:]:
            common = people.get(a, set()) & people.get(b, set())
            same_org = org[a] == org[b]
            if common or same_org:
                why = []
                if common:
                    why.append(f"{len(common)} shared attendee(s)")
                if same_org:
                    why.append(f"same organizer {org[a]}")
                csp.add_binary(a, b, lambda x, y: x[0] != y[0], "different timeslot: " + ", ".join(why))
                shared_pairs.append((a, b, len(common), same_org))
            else:
                csp.add_binary(a, b, lambda x, y: x != y, "not same room at same time")
    return csp, ev, rooms, slots, people, shared_pairs


def sched_run(data: Dict[str, pd.DataFrame], dataset: str = "") -> dict:
    csp, ev, rooms, slots, people, pairs = sched_build(data)
    res = solve(csp, time_limit=30)
    sol = res.solution
    cap = dict(zip(rooms["room_id"].astype(str), rooms["capacity"]))
    sl = slots.set_index("slot_id")
    rows = []
    if sol:
        for _, e in ev.iterrows():
            s, r = sol[e["event_id"]]
            rows.append({"Event": e["event_id"], "Name": e["event_name"], "Organizer": e["organizer"],
                         "Attendees": int(e["size"]), "Timeslot": s, "Day": sl.loc[s, "day"],
                         "Start": sl.loc[s, "start_time"], "Room": r, "Capacity": int(cap[r]),
                         "Seat utilisation %": round(100 * e["size"] / cap[r], 1) if cap[r] else 0})
    sched = pd.DataFrame(rows)
    grid = None
    if sol:
        grid = pd.DataFrame("", index=list(slots["slot_id"]), columns=list(cap))
        for _, r in sched.iterrows():
            grid.loc[r["Timeslot"], r["Room"]] = f"{r['Event']} ({r['Attendees']})"
        grid.index = [f"{sl.loc[s, 'day']} {sl.loc[s, 'start_time']}" for s in grid.index]
    # post-hoc verification: personal clashes and room double-booking
    clashes = 0
    if sol:
        by_person = {}
        for e, ps in people.items():
            for p in ps:
                by_person.setdefault(p, []).append(sol[e][0])
        clashes = sum(len(v) - len(set(v)) for v in by_person.values())
    double = (len(sched) - len(sched[["Timeslot", "Room"]].drop_duplicates())) if sol else 0
    over = int((sched["Attendees"] > sched["Capacity"]).sum()) if sol else 0
    conflict_tbl = pd.DataFrame([{"Event A": a, "Event B": b, "Shared attendees": n, "Same organizer": "Yes" if o else "No"}
                                 for a, b, n, o in pairs])
    return {"csp": csp, "search": res, "ac3": res.ac3, "schedule": sched, "grid": grid, "events": ev,
            "rooms": rooms, "slots": slots, "clashes": clashes, "double": double, "over": over,
            "conflicts": conflict_tbl, "dataset": dataset}


def sched_conflict_fig(r: dict) -> go.Figure:
    import networkx as nx
    G = nx.Graph()
    ev = r["events"]
    G.add_nodes_from(ev["event_id"])
    for _, row in r["conflicts"].iterrows():
        G.add_edge(row["Event A"], row["Event B"], w=row["Shared attendees"])
    pos = nx.spring_layout(G, seed=4, k=0.9)
    slot_of = dict(zip(r["schedule"]["Event"], r["schedule"]["Timeslot"])) if len(r["schedule"]) else {}
    slots = list(r["slots"]["slot_id"])
    fig = go.Figure()
    ex, ey = [], []
    for a, b in G.edges():
        ex += [pos[a][0], pos[b][0], None]; ey += [pos[a][1], pos[b][1], None]
    fig.add_trace(go.Scatter(x=ex, y=ey, mode="lines", line=dict(color="#B8C7DA", width=1), hoverinfo="skip", showlegend=False))
    for i, s in enumerate(slots + ["(unscheduled)"]):
        nodes = [n for n in G.nodes if slot_of.get(n, "(unscheduled)") == s]
        if not nodes:
            continue
        fig.add_trace(go.Scatter(x=[pos[n][0] for n in nodes], y=[pos[n][1] for n in nodes], mode="markers+text",
                                 text=nodes, textposition="top center", name=s,
                                 marker=dict(size=22, color=PALETTE[i % len(PALETTE)], line=dict(color="white", width=2))))
    fig.update_layout(title="Conflict graph - edge = events that must NOT share a timeslot; colour = assigned slot",
                      height=480, plot_bgcolor="white", margin=dict(l=10, r=10, t=50, b=10),
                      xaxis=dict(visible=False), yaxis=dict(visible=False))
    return fig


def sched_timetable_png(r: dict) -> bytes:
    g = r["grid"]
    fig, ax = plt.subplots(figsize=(11, 0.55 * len(g) + 1.4))
    ax.axis("off")
    cell_colors = [["#EAF2FC" if v else "white" for v in row] for row in g.values]
    tb = ax.table(cellText=g.values, rowLabels=list(g.index), colLabels=list(g.columns), loc="center",
                  cellColours=cell_colors, colColours=["#0B4F9C"] * len(g.columns), rowColours=["#DDE8F6"] * len(g))
    tb.auto_set_font_size(False); tb.set_fontsize(8.5); tb.scale(1, 1.5)
    for (i, j), c in tb.get_celld().items():
        if i == 0:
            c.get_text().set_color("white"); c.get_text().set_fontweight("bold")
    ax.set_title("Timetable (event id and attendees)", color="#0B4F9C", fontweight="bold")
    return fig_to_png(fig)


def sched_report(r: dict) -> Report:
    a, s = r["ac3"], r["search"]
    if r["schedule"].empty:
        summary = ("No clash-free timetable exists with the given rooms and timeslots. Add timeslots or rooms, "
                   "or relax blocked slots.")
    else:
        summary = (f"{len(r['schedule'])} events were placed into {r['slots'].shape[0]} timeslots and {r['rooms'].shape[0]} rooms. "
                   f"Verification: {r['clashes']} personal clashes, {r['double']} double-booked rooms, {r['over']} over-capacity "
                   f"rooms. AC-3 removed {a.values_removed} impossible (timeslot, room) options before search.")
    return Report(
        use_case="Scheduling & Timetabling", dataset=r["dataset"], summary=summary,
        metrics=metrics_pairs(a, s, {"Events": len(r["events"]), "Conflict pairs": len(r["conflicts"]),
                                     "Personal clashes (verified)": r["clashes"], "Room double-bookings": r["double"],
                                     "Over-capacity rooms": r["over"]}),
        sections=[("How to read the results",
                   "Each row of 'Timetable' is an event with its timeslot, room and seat utilisation. 'Timetable grid' "
                   "shows the same thing as slots x rooms. 'Conflict pairs' lists the events that share people or "
                   "an organizer and therefore got different timeslots.")],
        formulas=[("Xe = (timeslot, room)", "The value chosen for event e."),
                  ("capacity(room(Xe)) >= size(e)", "Unary rule: rooms that are too small are removed from De."),
                  ("slot(Xa) != slot(Xb)  if attendees(a) and attendees(b) overlap, or organizer(a)=organizer(b)", "People-clash rule."),
                  ("Xa != Xb  (not the same room at the same time)", "Room-clash rule for every other pair."),
                  ("Seat utilisation % = 100 x size(e) / capacity(room)", "How full the room is.")] + AC3_FORMULAS,
        tables=[("Timetable", r["schedule"]),
                ("Timetable grid (slot x room)", r["grid"].reset_index().rename(columns={"index": "Timeslot"}) if r["grid"] is not None else pd.DataFrame()),
                ("Conflict pairs", r["conflicts"]), ("AC-3 trace (first 300 revisions)", trace_df(a, 300))],
        figures=([("Timetable grid", sched_timetable_png(r))] if r["grid"] is not None else []) +
                [("Options per event before vs after AC-3", domain_reduction_png(a))],
    )


# =====================================================================================
# USE CASE 3 - MAP COLOURING & RADIO-FREQUENCY ASSIGNMENT
# =====================================================================================
COLOR_NAMES = ["Red", "Green", "Blue", "Yellow", "Purple", "Orange", "Teal", "Pink", "Olive", "Grey", "Navy", "Brown"]
COLOR_HEX = ["#E53935", "#43A047", "#1E88E5", "#FDD835", "#8E24AA", "#FB8C00", "#00897B", "#EC407A",
             "#9E9D24", "#757575", "#283593", "#6D4C41"]

COLOR_SCHEMA = {
    "nodes.csv": _schema([
        ["node_id", "text", "Yes", "Unique id of a region / transmitter (a CSP variable).", "NSW"],
        ["label", "text", "No", "Display name.", "New South Wales"],
        ["x", "number", "No", "Horizontal position for the map (longitude or km). Auto-layout if blank.", "147.0"],
        ["y", "number", "No", "Vertical position (latitude or km).", "-32.5"],
        ["fixed_color", "integer", "No", "Pre-assigned colour/channel number 1..k (unary rule). Blank = free.", "2"],
    ]),
    "edges.csv": _schema([
        ["source", "text", "Yes", "node_id of one region / transmitter.", "NSW"],
        ["target", "text", "Yes", "node_id of the neighbouring region / interfering transmitter.", "V"],
        ["min_separation", "integer", "No", "Required gap |ci - cj| >= s. 1 = just 'different colour' (default); 2+ = adjacent-channel protection for radio.", "2"],
    ]),
}

_AUS = [("WA", "Western Australia", 122.0, -25.5), ("NT", "Northern Territory", 133.5, -19.5),
        ("SA", "South Australia", 135.5, -30.0), ("Q", "Queensland", 144.5, -22.5),
        ("NSW", "New South Wales", 147.0, -32.5), ("V", "Victoria", 144.5, -37.0), ("T", "Tasmania", 146.5, -42.0)]
_AUS_E = [("WA", "NT"), ("WA", "SA"), ("NT", "SA"), ("NT", "Q"), ("SA", "Q"), ("SA", "NSW"), ("SA", "V"),
          ("Q", "NSW"), ("NSW", "V")]
_SAM = [("BR", "Brazil", -52, -10), ("AR", "Argentina", -64, -34), ("CL", "Chile", -71, -30), ("PE", "Peru", -75, -9.5),
        ("CO", "Colombia", -73, 4), ("VE", "Venezuela", -66, 7), ("EC", "Ecuador", -78.5, -1.5), ("BO", "Bolivia", -64.5, -17),
        ("PY", "Paraguay", -58, -23), ("UY", "Uruguay", -56, -33), ("GY", "Guyana", -59, 5), ("SR", "Suriname", -56, 4),
        ("GF", "French Guiana", -53, 4)]
_SAM_E = [("AR", "CL"), ("AR", "BO"), ("AR", "PY"), ("AR", "BR"), ("AR", "UY"), ("BO", "PE"), ("BO", "BR"), ("BO", "PY"),
          ("BO", "CL"), ("BR", "GF"), ("BR", "SR"), ("BR", "GY"), ("BR", "VE"), ("BR", "CO"), ("BR", "PE"), ("BR", "PY"),
          ("BR", "UY"), ("CL", "PE"), ("CO", "VE"), ("CO", "PE"), ("CO", "EC"), ("EC", "PE"), ("GY", "VE"), ("GY", "SR"),
          ("SR", "GF")]


def color_synthetic(kind: str = "Australia - states & territories", seed: int = 5, n_tx: int = 30,
                    radius_km: float = 30.0) -> Dict[str, pd.DataFrame]:
    if kind.startswith("Australia"):
        nodes, edges = _AUS, _AUS_E
    elif kind.startswith("South America"):
        nodes, edges = _SAM, _SAM_E
    else:
        rng = _random.Random(seed)
        pts = [(f"TX{i+1:02d}", f"Transmitter {i+1}", round(rng.uniform(0, 100), 1), round(rng.uniform(0, 100), 1))
               for i in range(n_tx)]
        e = []
        for i in range(len(pts)):
            for j in range(i + 1, len(pts)):
                d = ((pts[i][2] - pts[j][2]) ** 2 + (pts[i][3] - pts[j][3]) ** 2) ** 0.5
                if d < radius_km:
                    e.append((pts[i][0], pts[j][0], 2 if d < radius_km / 2 else 1, round(d, 1)))
        return {"nodes.csv": pd.DataFrame([dict(node_id=a, label=b, x=c, y=d, fixed_color="") for a, b, c, d in pts]),
                "edges.csv": pd.DataFrame(e, columns=["source", "target", "min_separation", "distance_km"])}
    # a couple of pre-coloured regions show AC-3 propagation (WA=Red, NT=Green -> AC-3 deduces the mainland)
    fixed = {"WA": "1", "NT": "2"} if kind.startswith("Australia") else {"BR": "1"}
    return {"nodes.csv": pd.DataFrame([dict(node_id=a, label=b, x=c, y=d, fixed_color=fixed.get(a, "")) for a, b, c, d in nodes]),
            "edges.csv": pd.DataFrame([dict(source=a, target=b, min_separation=1) for a, b in edges])}


def color_validate(data: Dict[str, pd.DataFrame]) -> List[str]:
    if "nodes.csv" not in data or "edges.csv" not in data:
        return ["Both nodes.csv and edges.csv are required."]
    errs = _need_cols(data["nodes.csv"], ["node_id"], "nodes.csv") + _need_cols(data["edges.csv"], ["source", "target"], "edges.csv")
    if errs:
        return errs
    ids = set(data["nodes.csv"]["node_id"].astype(str))
    if data["nodes.csv"]["node_id"].duplicated().any():
        errs.append("nodes.csv: duplicate node_id.")
    bad = (set(data["edges.csv"]["source"].astype(str)) | set(data["edges.csv"]["target"].astype(str))) - ids
    if bad:
        errs.append(f"edges.csv references unknown node(s): {sorted(bad)[:5]}.")
    if (data["edges.csv"]["source"].astype(str) == data["edges.csv"]["target"].astype(str)).any():
        errs.append("edges.csv: a node can't be its own neighbour.")
    return errs


def color_build(data: Dict[str, pd.DataFrame], k: int) -> CSP:
    nodes = data["nodes.csv"].copy()
    nodes["node_id"] = nodes["node_id"].astype(str)
    csp = CSP(list(nodes["node_id"]), {n: list(range(1, k + 1)) for n in nodes["node_id"]}, "Colouring")
    if "fixed_color" in nodes.columns:
        for n, f in zip(nodes["node_id"], nodes["fixed_color"]):
            if pd.notna(f) and str(f).strip() not in ("", "nan"):
                fv = int(float(f))
                csp.add_unary(n, lambda c, fv=fv: c == fv, f"fixed colour {fv}")
    ed = data["edges.csv"]
    for _, e in ed.iterrows():
        a, b = str(e["source"]), str(e["target"])
        sep = int(e["min_separation"]) if "min_separation" in ed.columns and pd.notna(e["min_separation"]) else 1
        sep = max(1, sep)
        lab = "different colour" if sep == 1 else f"|ci - cj| >= {sep}"
        csp.add_binary(a, b, lambda x, y, s=sep: abs(x - y) >= s, lab)
    return csp


def color_run(data: Dict[str, pd.DataFrame], k: int, dataset: str = "", radio: bool = False) -> dict:
    csp = color_build(data, k)
    res = solve(csp, time_limit=20)
    nodes = data["nodes.csv"].copy()
    nodes["node_id"] = nodes["node_id"].astype(str)
    sol = res.solution
    deg = {v: len(csp.neighbors[v]) for v in csp.variables}
    unit = "Channel" if radio else "Colour"
    rows = []
    for _, n in nodes.iterrows():
        v = n["node_id"]
        c = sol[v] if sol else None
        rows.append({"Node": v, "Name": n.get("label", v), "Neighbours": deg[v],
                     f"{unit} #": c if c else None, unit: (f"Ch {c}" if radio else COLOR_NAMES[(c - 1) % len(COLOR_NAMES)]) if c else "",
                     "Options after AC-3": " ".join(map(str, res.ac3.domains[v]))})
    tbl = pd.DataFrame(rows)
    bad = csp.violations(sol) if sol else []
    usage = tbl[f"{unit} #"].value_counts().sort_index().reset_index() if sol else pd.DataFrame()
    if sol:
        usage.columns = [f"{unit} #", "Nodes using it"]
    return {"csp": csp, "search": res, "ac3": res.ac3, "table": tbl, "k": k, "nodes": nodes, "edges": data["edges.csv"],
            "violations": bad, "usage": usage, "dataset": dataset, "radio": radio}


def color_min_k(data: Dict[str, pd.DataFrame], k_max: int = 12) -> Tuple[Optional[int], List[Tuple[int, str, int]]]:
    """Smallest k with a solution. Returns (k, log of (k, status, values pruned by AC-3))."""
    log = []
    for k in range(1, k_max + 1):
        csp = color_build(data, k)
        res = solve(csp, time_limit=10)
        log.append((k, res.status, res.ac3.values_removed))
        if res.status == "solved":
            return k, log
    return None, log


def _positions(r: dict):
    nodes = r["nodes"]
    if "x" in nodes.columns and pd.to_numeric(nodes["x"], errors="coerce").notna().all():
        return {n: (float(x), float(y)) for n, x, y in zip(nodes["node_id"], nodes["x"], nodes["y"])}
    import networkx as nx
    G = nx.Graph()
    G.add_nodes_from(nodes["node_id"])
    G.add_edges_from(zip(r["edges"]["source"].astype(str), r["edges"]["target"].astype(str)))
    return {n: tuple(p) for n, p in nx.spring_layout(G, seed=2).items()}


def color_fig(r: dict) -> go.Figure:
    pos = _positions(r)
    sol = r["search"].solution or {}
    fig = go.Figure()
    for _, e in r["edges"].iterrows():
        a, b = str(e["source"]), str(e["target"])
        sep = int(e.get("min_separation", 1)) if pd.notna(e.get("min_separation", 1)) else 1
        fig.add_trace(go.Scatter(x=[pos[a][0], pos[b][0]], y=[pos[a][1], pos[b][1]], mode="lines", hoverinfo="skip",
                                 showlegend=False, line=dict(color="#D32F2F" if sep > 1 else "#9FB3CC", width=2.5 if sep > 1 else 1.5,
                                                             dash="solid" if sep > 1 else "dot")))
    names = r["table"].set_index("Node")["Name"].to_dict()
    for c in sorted(set(sol.values())) or [None]:
        ns = [n for n in pos if sol.get(n) == c]
        if not ns:
            continue
        lab = (f"Ch {c}" if r["radio"] else COLOR_NAMES[(c - 1) % 12]) if c else "unassigned"
        fig.add_trace(go.Scatter(x=[pos[n][0] for n in ns], y=[pos[n][1] for n in ns], mode="markers+text",
                                 text=ns, textposition="middle center", name=lab,
                                 hovertext=[f"{names.get(n, n)} -> {lab}" for n in ns], hoverinfo="text",
                                 textfont=dict(color="white" if c not in (4,) else "black", size=10),
                                 marker=dict(size=34 if len(pos) < 20 else 26, color=COLOR_HEX[(c - 1) % 12] if c else "#CCC",
                                             line=dict(color="white", width=2))))
    fig.update_layout(title=("Frequency plan - red edges need >=2 channels separation" if r["radio"] else "Coloured map / graph"),
                      height=520, plot_bgcolor="white", margin=dict(l=10, r=10, t=50, b=10),
                      xaxis=dict(visible=False), yaxis=dict(visible=False, scaleanchor="x") if not r["radio"] else dict(visible=False),
                      legend=dict(title="Assigned"))
    return fig


def color_png(r: dict) -> bytes:
    pos = _positions(r)
    sol = r["search"].solution or {}
    fig, ax = plt.subplots(figsize=(9, 6.5))
    for _, e in r["edges"].iterrows():
        a, b = str(e["source"]), str(e["target"])
        sep = int(e.get("min_separation", 1)) if pd.notna(e.get("min_separation", 1)) else 1
        ax.plot([pos[a][0], pos[b][0]], [pos[a][1], pos[b][1]], color="#D32F2F" if sep > 1 else "#9FB3CC",
                lw=2 if sep > 1 else 1, ls="-" if sep > 1 else ":", zorder=1)
    for n, (x, y) in pos.items():
        c = sol.get(n)
        ax.scatter([x], [y], s=650, color=COLOR_HEX[(c - 1) % 12] if c else "#CCC", edgecolor="white", zorder=2)
        ax.text(x, y, n, ha="center", va="center", fontsize=8, color="white" if c != 4 else "black", fontweight="bold", zorder=3)
    ax.axis("off")
    ax.set_title("Colour / channel assignment", color="#0B4F9C", fontweight="bold")
    return fig_to_png(fig)


def color_report(r: dict) -> Report:
    a, s = r["ac3"], r["search"]
    unit = "channels" if r["radio"] else "colours"
    if s.solution:
        summary = (f"All {len(r['table'])} nodes were assigned one of {r['k']} {unit} with {len(r['violations'])} violations "
                   f"across {len(r['edges'])} adjacency/interference constraints. AC-3 pruned {a.values_removed} values "
                   f"before search; search used {s.nodes} assignments and {s.backtracks} backtracks.")
    else:
        summary = f"No valid assignment exists with {r['k']} {unit}. Increase k or use 'Find minimum k'."
    return Report(
        use_case="Map Colouring & Frequency Assignment", dataset=r["dataset"], summary=summary,
        metrics=metrics_pairs(a, s, {f"Number of {unit} (k)": r["k"], "Nodes": len(r["table"]), "Edges": len(r["edges"]),
                                     "Violations (verified)": len(r["violations"])}),
        sections=[("How to read the results",
                   f"'Assignment' gives each region / transmitter its {unit[:-1]}. Neighbours never share one; "
                   "for radio links marked with min_separation >= 2 the channel numbers differ by at least that much.")],
        formulas=[("ci in {1, ..., k}", "Each node chooses a colour / channel number."),
                  ("|ci - cj| >= s_ij  for every edge (i, j)", "s_ij = 1 means simply ci != cj (map colouring); s_ij = 2 adds a guard channel."),
                  ("Chromatic number chi(G) = smallest k with a solution", "Found with 'Find minimum k' (k = 1, 2, ... until solvable).")] + AC3_FORMULAS,
        tables=[("Assignment", r["table"]), (f"{unit.title()} usage", r["usage"]),
                ("AC-3 trace (first 300 revisions)", trace_df(a, 300))],
        figures=[("Coloured graph", color_png(r)), ("Options per node before vs after AC-3", domain_reduction_png(a))],
    )


# =====================================================================================
# USE CASE 4 - PRODUCT CONFIGURATION (PC builder / car configurator)
# =====================================================================================
CONFIG_SCHEMA = {
    "components.csv": _schema([
        ["component", "text", "Yes", "Component slot to fill (a CSP variable), e.g. CPU, Motherboard, Engine.", "CPU"],
        ["option_id", "text", "Yes", "Unique id of this option (a domain value).", "CPU-7600"],
        ["option_name", "text", "Yes", "Display name.", "AMD Ryzen 5 7600"],
        ["price", "number", "Yes", "Price of the option; totals are summed.", "199"],
        ["<any attribute>", "text/number", "No", "Any extra columns (socket, ram_type, tdp_w, length_mm ...). Referenced by rules. Lists use ';'.", "socket = AM5"],
    ]),
    "rules.csv": _schema([
        ["rule_id", "text", "Yes", "Unique rule id.", "R1"],
        ["component_a", "text", "Yes", "Left-hand component.", "CPU"],
        ["attribute_a", "text", "Yes", "Attribute column of component_a.", "socket"],
        ["operator", "==, !=, <=, >=, <, >, in, contains", "Yes",
         "Comparison. 'in' = value of A appears in B's ';' list; 'contains' = B's value appears in A's list.", "=="],
        ["component_b", "text", "Yes", "Right-hand component.", "Motherboard"],
        ["attribute_b", "text", "Yes", "Attribute column of component_b.", "socket"],
        ["offset", "number", "No", "Added to B for numeric rules: A.attr OP B.attr + offset. Default 0.", "250"],
        ["description", "text", "No", "Plain-English meaning shown in the explanation.", "CPU must fit the motherboard socket"],
    ]),
}


def config_synthetic(kind: str = "PC builder") -> Dict[str, pd.DataFrame]:
    if kind.startswith("PC"):
        C = []
        def add(comp, oid, name, price, **attrs):
            C.append(dict(component=comp, option_id=oid, option_name=name, price=price, **attrs))
        for oid, n, p, s, t in [("CPU-7600", "AMD Ryzen 5 7600", 199, "AM5", 65), ("CPU-7800X3D", "AMD Ryzen 7 7800X3D", 399, "AM5", 120),
                                ("CPU-7950X", "AMD Ryzen 9 7950X", 549, "AM5", 170), ("CPU-5600X", "AMD Ryzen 5 5600X", 139, "AM4", 65),
                                ("CPU-5800X3D", "AMD Ryzen 7 5800X3D", 299, "AM4", 105), ("CPU-13400F", "Intel Core i5-13400F", 189, "LGA1700", 65),
                                ("CPU-14700K", "Intel Core i7-14700K", 389, "LGA1700", 125), ("CPU-14900K", "Intel Core i9-14900K", 549, "LGA1700", 253),
                                ("CPU-265K", "Intel Core Ultra 7 265K", 379, "LGA1851", 125)]:
            add("CPU", oid, n, p, socket=s, tdp_w=t)
        for oid, n, p, s, r, f, m in [("MB-B650", "B650 ATX", 179, "AM5", "DDR5", "ATX", 2), ("MB-A620", "A620M mATX", 109, "AM5", "DDR5", "mATX", 1),
                                      ("MB-X670E", "X670E ATX", 329, "AM5", "DDR5", "ATX", 4), ("MB-B550", "B550 ATX", 129, "AM4", "DDR4", "ATX", 2),
                                      ("MB-A520", "A520M mATX", 79, "AM4", "DDR4", "mATX", 1), ("MB-B450I", "B450 Mini-ITX", 99, "AM4", "DDR4", "ITX", 1),
                                      ("MB-B760", "B760 ATX", 149, "LGA1700", "DDR5", "ATX", 3), ("MB-B760M", "B760M mATX DDR4", 119, "LGA1700", "DDR4", "mATX", 2),
                                      ("MB-Z790", "Z790 ATX", 259, "LGA1700", "DDR5", "ATX", 4), ("MB-Z890", "Z890 ATX", 299, "LGA1851", "DDR5", "ATX", 4)]:
            add("Motherboard", oid, n, p, socket=s, ram_type=r, form_factor=f, m2_slots=m)
        for oid, n, p, r in [("RAM-D4-16", "16GB DDR4-3200", 39, "DDR4"), ("RAM-D4-32", "32GB DDR4-3600", 69, "DDR4"),
                             ("RAM-D5-32", "32GB DDR5-6000", 99, "DDR5"), ("RAM-D5-64", "64GB DDR5-6000", 189, "DDR5")]:
            add("RAM", oid, n, p, ram_type=r)
        for oid, n, p, w, l in [("GPU-4060", "GeForce RTX 4060", 299, 115, 240), ("GPU-7600", "Radeon RX 7600", 259, 165, 204),
                                ("GPU-7800XT", "Radeon RX 7800 XT", 499, 263, 267), ("GPU-4070S", "GeForce RTX 4070 Super", 599, 220, 267),
                                ("GPU-4080S", "GeForce RTX 4080 Super", 999, 320, 310), ("GPU-4090", "GeForce RTX 4090", 1799, 450, 336)]:
            add("GPU", oid, n, p, power_w=w, length_mm=l)
        for oid, n, p, w in [("PSU-550", "550 W 80+ Bronze", 59, 550), ("PSU-650", "650 W 80+ Gold", 89, 650),
                             ("PSU-750", "750 W 80+ Gold", 109, 750), ("PSU-850", "850 W 80+ Gold", 129, 850), ("PSU-1000", "1000 W 80+ Platinum", 219, 1000)]:
            add("PSU", oid, n, p, wattage=w)
        for oid, n, p, f, g, c, rad in [("CASE-FULL", "Full tower", 169, "ATX;mATX;ITX", 420, 185, 360), ("CASE-MID", "Mid tower", 99, "ATX;mATX;ITX", 340, 165, 360),
                                        ("CASE-MATX", "Compact mATX", 79, "mATX;ITX", 300, 155, 240), ("CASE-SFF", "Mini-ITX small form factor", 119, "ITX", 280, 70, 240)]:
            add("Case", oid, n, p, supported_form_factors=f, max_gpu_length_mm=g, max_cooler_height_mm=c, max_radiator_mm=rad)
        for oid, n, p, s, h, t, rad in [("COOL-LP", "Low-profile air", 39, "AM4;AM5;LGA1700", 58, 95, 0),
                                        ("COOL-TOWER", "Tower air 155 mm", 49, "AM4;AM5;LGA1700;LGA1851", 155, 220, 0),
                                        ("COOL-AIO240", "240 mm liquid AIO", 99, "AM4;AM5;LGA1700;LGA1851", 55, 250, 240),
                                        ("COOL-AIO360", "360 mm liquid AIO", 149, "AM5;LGA1700;LGA1851", 55, 300, 360)]:
            add("Cooler", oid, n, p, supported_sockets=s, height_mm=h, max_tdp_w=t, radiator_mm=rad)
        for oid, n, p, m in [("SSD-SATA1", "1TB SATA SSD", 59, 0), ("SSD-NVME1", "1TB NVMe Gen4", 79, 1), ("SSD-NVME2", "2TB NVMe Gen4", 139, 1)]:
            add("Storage", oid, n, p, m2_required=m)
        R = [("R1", "CPU", "socket", "==", "Motherboard", "socket", 0, "CPU must match the motherboard socket"),
             ("R2", "RAM", "ram_type", "==", "Motherboard", "ram_type", 0, "Memory type must match the motherboard"),
             ("R3", "Motherboard", "form_factor", "in", "Case", "supported_form_factors", 0, "Case must accept the board size"),
             ("R4", "GPU", "length_mm", "<=", "Case", "max_gpu_length_mm", 0, "Graphics card must fit in the case"),
             ("R5", "CPU", "socket", "in", "Cooler", "supported_sockets", 0, "Cooler must support the CPU socket"),
             ("R6", "Cooler", "height_mm", "<=", "Case", "max_cooler_height_mm", 0, "Air cooler must fit under the side panel"),
             ("R7", "CPU", "tdp_w", "<=", "Cooler", "max_tdp_w", 0, "Cooler must handle the CPU heat"),
             ("R8", "PSU", "wattage", ">=", "GPU", "power_w", 300, "PSU needs GPU power + 300 W headroom"),
             ("R9", "PSU", "wattage", ">=", "CPU", "tdp_w", 350, "PSU needs CPU power + 350 W headroom"),
             ("R10", "Cooler", "radiator_mm", "<=", "Case", "max_radiator_mm", 0, "Radiator must fit in the case"),
             ("R11", "Storage", "m2_required", "<=", "Motherboard", "m2_slots", 0, "NVMe drive needs an M.2 slot")]
    else:
        C = []
        def add(comp, oid, name, price, **attrs):
            C.append(dict(component=comp, option_id=oid, option_name=name, price=price, **attrs))
        for oid, n, p, fuel, hp, tq, trims in [("ENG-20T", "2.0L I4 Turbo", 0, "gas", 250, 370, "Base;Sport;Touring"),
                                               ("ENG-30T", "3.0L I6 Turbo", 4500, "gas", 382, 500, "Sport;Touring;Limited"),
                                               ("ENG-HYB", "2.5L Hybrid", 2800, "hybrid", 219, 300, "Base;Touring;Limited"),
                                               ("ENG-EV", "Dual-motor Electric", 9500, "electric", 480, 700, "Touring;Limited")]:
            add("Engine", oid, n, p, fuel=fuel, hp=hp, torque_nm=tq, allowed_trims=trims)
        for oid, n, p, fuels, mt in [("TR-8AT", "8-speed automatic", 0, "gas", 550), ("TR-6MT", "6-speed manual", -500, "gas", 400),
                                     ("TR-ECVT", "eCVT", 0, "hybrid", 350), ("TR-1SP", "Single-speed reduction", 0, "electric", 900)]:
            add("Transmission", oid, n, p, fuels=fuels, max_torque_nm=mt)
        for oid, n, p, drives, wmin, wmax in [("TRIM-BASE", "Base", 32000, "FWD;AWD", 17, 18), ("TRIM-SPORT", "Sport", 38000, "RWD;AWD", 18, 20),
                                              ("TRIM-TOUR", "Touring", 41000, "FWD;AWD", 18, 19), ("TRIM-LTD", "Limited", 47000, "AWD", 19, 21)]:
            add("Trim", oid, n, p, trim_name=n, drivetrains=drives, min_wheel_in=wmin, max_wheel_in=wmax)
        for oid, n, p, d, need in [("DRV-FWD", "Front-wheel drive", 0, "FWD", 0), ("DRV-RWD", "Rear-wheel drive", 800, "RWD", 250),
                                   ("DRV-AWD", "All-wheel drive", 2200, "AWD", 0)]:
            add("Drivetrain", oid, n, p, drive=d, min_hp=need)
        for oid, n, p, sz in [("WHL-17", "17-inch alloy", 0, 17), ("WHL-18", "18-inch alloy", 600, 18), ("WHL-19", "19-inch machined", 1200, 19),
                              ("WHL-20", "20-inch sport", 1900, 20), ("WHL-21", "21-inch forged", 3200, 21)]:
            add("Wheels", oid, n, p, size_in=sz)
        for oid, n, p, hp, drv in [("TOW-NONE", "No tow package", 0, 0, "FWD;RWD;AWD"), ("TOW-II", "Class II hitch (1,600 kg)", 450, 219, "FWD;RWD;AWD"),
                                   ("TOW-III", "Class III hitch (2,700 kg)", 950, 350, "RWD;AWD")]:
            add("Towing", oid, n, p, min_hp=hp, drives_ok=drv)
        R = [("C1", "Trim", "trim_name", "in", "Engine", "allowed_trims", 0, "Engine offered only on certain trims"),
             ("C2", "Engine", "fuel", "in", "Transmission", "fuels", 0, "Transmission must suit the powertrain type"),
             ("C3", "Engine", "torque_nm", "<=", "Transmission", "max_torque_nm", 0, "Transmission must handle engine torque"),
             ("C4", "Drivetrain", "drive", "in", "Trim", "drivetrains", 0, "Trim decides the available drivetrains"),
             ("C5", "Wheels", "size_in", ">=", "Trim", "min_wheel_in", 0, "Trim sets the smallest wheel"),
             ("C6", "Wheels", "size_in", "<=", "Trim", "max_wheel_in", 0, "Trim sets the largest wheel"),
             ("C7", "Engine", "hp", ">=", "Towing", "min_hp", 0, "Tow package needs enough power"),
             ("C8", "Drivetrain", "drive", "in", "Towing", "drives_ok", 0, "Heavy towing needs RWD/AWD"),
             ("C9", "Engine", "hp", ">=", "Drivetrain", "min_hp", 0, "RWD offered only on powerful engines")]
    comps = pd.DataFrame(C)
    rules = pd.DataFrame(R, columns=["rule_id", "component_a", "attribute_a", "operator", "component_b", "attribute_b", "offset", "description"])
    return {"components.csv": comps, "rules.csv": rules}


_OPS = {"==": lambda a, b: str(a).strip() == str(b).strip(), "!=": lambda a, b: str(a).strip() != str(b).strip(),
        "<=": lambda a, b: float(a) <= float(b), ">=": lambda a, b: float(a) >= float(b),
        "<": lambda a, b: float(a) < float(b), ">": lambda a, b: float(a) > float(b),
        "in": lambda a, b: str(a).strip() in [x.strip() for x in str(b).split(";")],
        "contains": lambda a, b: str(b).strip() in [x.strip() for x in str(a).split(";")]}


def config_validate(data: Dict[str, pd.DataFrame]) -> List[str]:
    if "components.csv" not in data or "rules.csv" not in data:
        return ["Both components.csv and rules.csv are required."]
    comps, rules = data["components.csv"], data["rules.csv"]
    errs = _need_cols(comps, ["component", "option_id", "option_name", "price"], "components.csv")
    errs += _need_cols(rules, ["rule_id", "component_a", "attribute_a", "operator", "component_b", "attribute_b"], "rules.csv")
    if errs:
        return errs
    if comps["option_id"].duplicated().any():
        errs.append("components.csv: option_id must be unique.")
    if not pd.to_numeric(comps["price"], errors="coerce").notna().all():
        errs.append("components.csv: price must be numeric.")
    names = set(comps["component"].astype(str))
    for _, r in rules.iterrows():
        for side in ("a", "b"):
            c, at = str(r[f"component_{side}"]), str(r[f"attribute_{side}"])
            if c not in names:
                errs.append(f"Rule {r['rule_id']}: unknown component '{c}'.")
            elif at not in comps.columns:
                errs.append(f"Rule {r['rule_id']}: attribute '{at}' is not a column of components.csv.")
            elif comps.loc[comps["component"] == c, at].isna().any():
                errs.append(f"Rule {r['rule_id']}: some {c} options have no value for '{at}'.")
        if str(r["operator"]).strip() not in _OPS:
            errs.append(f"Rule {r['rule_id']}: operator '{r['operator']}' not in {list(_OPS)}.")
        if str(r["component_a"]) == str(r["component_b"]):
            errs.append(f"Rule {r['rule_id']}: component_a and component_b must differ (use a pin for single-component rules).")
    return errs


def config_build(data: Dict[str, pd.DataFrame], pins: Optional[Dict[str, str]] = None,
                 max_price: Optional[Dict[str, float]] = None) -> Tuple[CSP, Dict[str, dict]]:
    comps = data["components.csv"].copy()
    comps["price"] = pd.to_numeric(comps["price"])
    comps = comps.sort_values(["component", "price"])  # cheapest first = value ordering
    rec = {r["option_id"]: r for r in comps.to_dict("records")}
    slots = list(dict.fromkeys(data["components.csv"]["component"].astype(str)))
    doms = {c: list(comps.loc[comps["component"].astype(str) == c, "option_id"]) for c in slots}
    csp = CSP(slots, doms, "Configuration")
    csp.value_formatter = lambda var, v: rec[v]["option_name"]
    for c, oid in (pins or {}).items():
        if oid:
            csp.add_unary(c, lambda v, oid=oid: v == oid, f"customer picked {rec[oid]['option_name']}")
    for c, mp in (max_price or {}).items():
        if mp is not None:
            csp.add_unary(c, lambda v, mp=mp: rec[v]["price"] <= mp, f"price <= {mp}")
    for _, r in data["rules.csv"].iterrows():
        op = _OPS[str(r["operator"]).strip()]
        off = float(r["offset"]) if "offset" in r and pd.notna(r.get("offset")) and str(r.get("offset")).strip() != "" else 0.0
        aa, bb = str(r["attribute_a"]), str(r["attribute_b"])
        num = str(r["operator"]).strip() in ("<=", ">=", "<", ">")
        if num:
            pred = lambda x, y, op=op, aa=aa, bb=bb, off=off: op(rec[x][aa], float(rec[y][bb]) + off)
        else:
            pred = lambda x, y, op=op, aa=aa, bb=bb: op(rec[x][aa], rec[y][bb])
        desc = r.get("description", "") if "description" in r else ""
        lab = f"{r['rule_id']}: {r['component_a']}.{aa} {r['operator']} {r['component_b']}.{bb}" + (f" + {off:g}" if off else "")
        if desc and str(desc) != "nan":
            lab += f" ({desc})"
        csp.add_binary(str(r["component_a"]), str(r["component_b"]), pred, lab)
    return csp, rec


def config_optimize(csp: CSP, rec: Dict[str, dict], domains: Dict[str, list], budget: Optional[float] = None,
                    top_n: int = 25, count_cap: int = 200_000) -> Tuple[List[dict], int, int, bool]:
    """Exact search over the AC-3-reduced domains.
    Returns (top_n cheapest valid configs within budget, #valid configs, #valid within budget, complete?)."""
    import heapq
    order = sorted(csp.variables, key=lambda v: len(domains[v]))
    min_rest = [0.0] * (len(order) + 1)
    for i in range(len(order) - 1, -1, -1):
        min_rest[i] = min_rest[i + 1] + (min(rec[o]["price"] for o in domains[order[i]]) if domains[order[i]] else 0)
    heap: List[Tuple[float, int, dict]] = []   # max-heap by -total
    stats = {"valid": 0, "within": 0, "complete": True, "tick": 0}
    assign: Dict[str, str] = {}

    def rec_dfs(i: int, cost: float):
        if stats["valid"] >= count_cap:
            stats["complete"] = False
            return
        if i == len(order):
            stats["valid"] += 1
            if budget is None or cost <= budget:
                stats["within"] += 1
                stats["tick"] += 1
                item = (-cost, stats["tick"], dict(assign))
                if len(heap) < top_n:
                    heapq.heappush(heap, item)
                elif cost < -heap[0][0]:
                    heapq.heapreplace(heap, item)
            return
        var = order[i]
        for val in domains[var]:
            ok = True
            for other, oval in assign.items():
                if (var, other) in csp._preds and not csp.satisfies(var, val, other, oval):
                    ok = False
                    break
            if ok:
                assign[var] = val
                rec_dfs(i + 1, cost + rec[val]["price"])
                del assign[var]

    rec_dfs(0, 0.0)
    best = sorted(({"total": -c, "sol": a} for c, _, a in heap), key=lambda d: d["total"])
    return best, stats["valid"], stats["within"], stats["complete"]


def config_run(data: Dict[str, pd.DataFrame], pins: Optional[Dict[str, str]] = None, budget: Optional[float] = None,
               dataset: str = "") -> dict:
    csp, rec = config_build(data, pins)
    res = solve(csp, max_solutions=1, time_limit=20)          # AC-3 + MAC: is there any valid build?
    if res.ac3.consistent:
        within, n_valid, n_within, complete = config_optimize(csp, rec, res.ac3.domains, budget)
    else:
        within, n_valid, n_within, complete = [], 0, 0, True
    res.exhaustive = complete
    best = within[0] if within else None
    rows = []
    if best:
        for c in csp.variables:
            o = rec[best["sol"][c]]
            attrs = {k: v for k, v in o.items() if k not in ("component", "option_id", "option_name", "price") and pd.notna(v) and str(v) != ""}
            rows.append({"Component": c, "Chosen option": o["option_name"], "Option id": o["option_id"], "Price": o["price"],
                         "Key attributes": ", ".join(f"{k}={v}" for k, v in attrs.items())})
        rows.append({"Component": "TOTAL", "Chosen option": "", "Option id": "", "Price": best["total"], "Key attributes": ""})
    chosen = pd.DataFrame(rows)
    alts = pd.DataFrame([{"Rank": i + 1, "Total price": p["total"], **{c: rec[p["sol"][c]]["option_name"] for c in csp.variables}}
                         for i, p in enumerate(within[:25])])
    # why each option was eliminated (from the AC-3 trace + unary pins)
    why = []
    for t in res.ac3.trace:
        for v in t["_removed"]:
            why.append({"Component": t["_x"], "Eliminated option": rec[v]["option_name"], "Price": rec[v]["price"],
                        "Because of": t["_y"], "Rule": csp.label_of(t["_x"], t["_y"])})
    for c in csp.variables:
        for v in csp.original_domains[c]:
            if v not in csp.domains[c]:
                why.append({"Component": c, "Eliminated option": rec[v]["option_name"], "Price": rec[v]["price"],
                            "Because of": "customer choice", "Rule": "pinned by user"})
    avail = pd.DataFrame([{"Component": c, "Options in catalogue": len(csp.original_domains[c]),
                           "After customer picks": len(res.ac3.initial_domains[c]),
                           "After AC-3 (still compatible)": len(res.ac3.domains[c]),
                           "Compatible options": ", ".join(rec[v]["option_name"] for v in res.ac3.domains[c])}
                          for c in csp.variables])
    violations = csp.violations(best["sol"]) if best else []
    return {"csp": csp, "rec": rec, "search": res, "ac3": res.ac3, "chosen": chosen, "alternatives": alts,
            "eliminated": pd.DataFrame(why), "availability": avail, "n_valid": n_valid, "n_within": n_within,
            "best": best, "budget": budget, "pins": pins or {}, "violations": violations, "dataset": dataset,
            "rules": data["rules.csv"]}


def config_rule_dot(data: Dict[str, pd.DataFrame]) -> str:
    comps = list(dict.fromkeys(data["components.csv"]["component"].astype(str)))
    lines = ['graph G {', 'graph [layout=neato, overlap=false, splines=true, bgcolor="transparent"];',
             'node [shape=box, style="rounded,filled", fillcolor="#E8F1FB", color="#0B4F9C", fontname="Helvetica", fontsize=12];',
             'edge [color="#F28C28", fontname="Helvetica", fontsize=9, fontcolor="#07326A"];']
    lines += [f'"{c}";' for c in comps]
    for _, r in data["rules.csv"].iterrows():
        lines.append(f'"{r["component_a"]}" -- "{r["component_b"]}" [label="{r["rule_id"]}"];')
    return "\n".join(lines + ["}"])


def config_png(r: dict) -> bytes:
    av = r["availability"]
    fig, ax = plt.subplots(figsize=(10, 4.2))
    x = range(len(av))
    ax.bar([i - 0.27 for i in x], av["Options in catalogue"], 0.27, label="Catalogue", color="#B8C7DA")
    ax.bar(list(x), av["After customer picks"], 0.27, label="After customer picks", color="#0B4F9C")
    ax.bar([i + 0.27 for i in x], av["After AC-3 (still compatible)"], 0.27, label="After AC-3", color="#F28C28")
    ax.set_xticks(list(x)); ax.set_xticklabels(av["Component"], rotation=20)
    ax.set_ylabel("Options"); ax.legend(); ax.grid(axis="y", alpha=0.3)
    ax.set_title("How choices and rules narrow each component", color="#0B4F9C", fontweight="bold")
    return fig_to_png(fig)


def config_report(r: dict) -> Report:
    a, s = r["ac3"], r["search"]
    pins = ", ".join(f"{k} = {r['rec'][v]['option_name']}" for k, v in r["pins"].items() if v) or "none"
    if r["best"]:
        summary = (f"Customer choices: {pins}. AC-3 removed {a.values_removed} incompatible options before search. "
                   f"{r['n_valid']} fully compatible configurations exist"
                   f"{'' if s.exhaustive else ' (enumeration capped)'}; {r['n_within']} are within budget. "
                   f"The cheapest valid build costs {r['best']['total']:,.0f} and breaks {len(r['violations'])} rules.")
    else:
        summary = (f"Customer choices: {pins}. No configuration satisfies every rule"
                   f"{' within the budget' if r['n_valid'] else ''}. See 'Eliminated options' to understand why.")
    return Report(
        use_case="Product Configuration", dataset=r["dataset"], summary=summary,
        metrics=metrics_pairs(a, s, {"Components": len(r["csp"].variables), "Rules": len(r["rules"]),
                                     "Valid configurations": r["n_valid"], "Within budget": r["n_within"],
                                     "Budget": "none" if r["budget"] is None else f"{r['budget']:,.0f}",
                                     "Cheapest valid total": f"{r['best']['total']:,.0f}" if r["best"] else "-"}),
        sections=[("How to read the results",
                   "'Recommended configuration' is the cheapest build that satisfies every rule and your picks. "
                   "'Compatibility after AC-3' shows how many options stay possible for each component, and "
                   "'Eliminated options' names the rule that removed each incompatible option.")],
        formulas=[("X_component in {options of that component}", "One variable per component slot."),
                  ("A.attr  OP  B.attr + offset", "Every row of rules.csv becomes one binary constraint between two components."),
                  ("value(A) in list(B)", "'in' rules, e.g. motherboard form factor must be in the case's supported list."),
                  ("Total price = sum of chosen option prices <= budget", "Global rule applied to complete configurations after search.")] + AC3_FORMULAS,
        tables=[("Recommended configuration", r["chosen"]), ("Compatibility after AC-3", r["availability"]),
                ("Eliminated options", r["eliminated"]), ("Top alternatives (cheapest first)", r["alternatives"])],
        figures=[("Options per component", config_png(r))],
    )


# =====================================================================================
# USE CASE 5 - PLANNING & RESOURCE ALLOCATION
# =====================================================================================
PLAN_SCHEMA = {
    "tasks.csv": _schema([
        ["task_id", "text", "Yes", "Unique task id (a CSP variable = its start time).", "A1"],
        ["task_name", "text", "Yes", "Description.", "Cut steel - Job A"],
        ["job", "text", "Yes", "Job / order / project phase the task belongs to (used for colours).", "Job A"],
        ["resource", "text", "Yes", "Machine / crew / dock that performs it. Tasks on the same resource can't overlap.", "M1-Laser"],
        ["duration", "integer", "Yes", "Processing time in time units (hours, days...).", "3"],
        ["release", "integer", "No", "Earliest allowed start (arrival time). Default 0.", "0"],
        ["deadline", "integer", "No", "Latest allowed finish. Default = planning horizon.", "20"],
    ]),
    "precedences.csv": _schema([
        ["before", "text", "Yes", "task_id that must finish first.", "A1"],
        ["after", "text", "Yes", "task_id that starts afterwards.", "A2"],
        ["min_gap", "integer", "No", "Extra waiting time between them (curing, transport). Default 0.", "1"],
    ]),
}


def plan_synthetic(kind: str = "Manufacturing job-shop") -> Dict[str, pd.DataFrame]:
    T, P = [], []
    if kind.startswith("Manufacturing"):
        jobs = {"Job A": [("M1-Laser", 3, "Cut"), ("M2-Press", 2, "Bend"), ("M3-Weld", 2, "Weld")],
                "Job B": [("M1-Laser", 2, "Cut"), ("M3-Weld", 1, "Weld"), ("M2-Press", 4, "Form")],
                "Job C": [("M2-Press", 4, "Stamp"), ("M3-Weld", 3, "Weld"), ("M1-Laser", 1, "Engrave")],
                "Job D": [("M3-Weld", 2, "Tack"), ("M1-Laser", 3, "Cut"), ("M2-Press", 2, "Bend")]}
        for j, ops in jobs.items():
            prev = None
            for k, (m, d, op) in enumerate(ops, 1):
                tid = f"{j[-1]}{k}"
                T.append(dict(task_id=tid, task_name=f"{op} - {j}", job=j, resource=m, duration=d, release=0, deadline=""))
                if prev:
                    P.append(dict(before=prev, after=tid, min_gap=0))
                prev = tid
    elif kind.startswith("Construction"):
        rows = [("SP", "Site preparation", "Groundworks", "Excavation crew", 3), ("FD", "Foundation", "Groundworks", "Concrete crew", 4),
                ("FR", "Framing", "Structure", "Carpentry crew", 6), ("RF", "Roofing", "Structure", "Roofing crew", 3),
                ("PL", "Plumbing rough-in", "Services", "Plumbing crew", 3), ("EL", "Electrical rough-in", "Services", "Electrical crew", 3),
                ("HV", "HVAC install", "Services", "Plumbing crew", 2), ("IN", "Insulation", "Interior", "Carpentry crew", 2),
                ("DW", "Drywall", "Interior", "Drywall crew", 4), ("PT", "Interior painting", "Interior", "Painting crew", 3),
                ("SD", "Exterior siding", "Structure", "Carpentry crew", 3), ("FL", "Flooring", "Interior", "Flooring crew", 3),
                ("LS", "Landscaping", "Exterior", "Excavation crew", 2), ("FI", "Final inspection", "Handover", "Inspector", 1)]
        T = [dict(task_id=a, task_name=b, job=c, resource=d, duration=e, release=0, deadline="") for a, b, c, d, e in rows]
        for a, b, g in [("SP", "FD", 0), ("FD", "FR", 2), ("FR", "RF", 0), ("FR", "PL", 0), ("FR", "EL", 0), ("FR", "HV", 0),
                        ("RF", "SD", 0), ("PL", "IN", 0), ("EL", "IN", 0), ("HV", "IN", 0), ("IN", "DW", 0), ("DW", "PT", 0),
                        ("PT", "FL", 0), ("SD", "LS", 0), ("FL", "FI", 0), ("LS", "FI", 0)]:
            P.append(dict(before=a, after=b, min_gap=g))
    else:  # warehouse dock scheduling
        trucks = [("T01", "Frozen foods", "Reefer dock", 2, 0, 6), ("T02", "Dairy", "Reefer dock", 1, 1, 5),
                  ("T03", "Produce", "Reefer dock", 2, 2, 9), ("T04", "Paper goods", "Dock 1", 3, 0, 8),
                  ("T05", "Beverages", "Dock 1", 2, 1, 8), ("T06", "Hardware", "Dock 2", 3, 0, 7),
                  ("T07", "Electronics", "Dock 2", 2, 3, 10), ("T08", "Apparel", "Dock 1", 1, 4, 10),
                  ("T09", "Furniture", "Dock 2", 3, 5, 12), ("T10", "Returns sort", "Sort team", 2, 0, 12),
                  ("T11", "Cross-dock to store 12", "Sort team", 2, 0, 12), ("T12", "Frozen re-pack", "Sort team", 1, 0, 12)]
        T = [dict(task_id=a, task_name=f"Unload - {b}" if a <= "T09" else b, job="Inbound" if a <= "T09" else "Outbound",
                  resource=c, duration=d, release=r, deadline=dl) for a, b, c, d, r, dl in trucks]
        for a, b in [("T01", "T12"), ("T04", "T11"), ("T05", "T11"), ("T08", "T10")]:
            P.append(dict(before=a, after=b, min_gap=0))
    return {"tasks.csv": pd.DataFrame(T), "precedences.csv": pd.DataFrame(P, columns=["before", "after", "min_gap"])}


def plan_validate(data: Dict[str, pd.DataFrame]) -> List[str]:
    if "tasks.csv" not in data:
        return ["tasks.csv is required (precedences.csv is optional)."]
    t = data["tasks.csv"]
    errs = _need_cols(t, ["task_id", "task_name", "job", "resource", "duration"], "tasks.csv")
    if errs:
        return errs
    if t["task_id"].duplicated().any():
        errs.append("tasks.csv: duplicate task_id.")
    d = pd.to_numeric(t["duration"], errors="coerce")
    if d.isna().any() or (d <= 0).any():
        errs.append("tasks.csv: duration must be a positive number.")
    p = data.get("precedences.csv")
    if p is not None and len(p):
        errs += _need_cols(p, ["before", "after"], "precedences.csv")
        if not errs:
            bad = (set(p["before"].astype(str)) | set(p["after"].astype(str))) - set(t["task_id"].astype(str))
            if bad:
                errs.append(f"precedences.csv references unknown task(s): {sorted(bad)[:5]}.")
            import networkx as nx
            G = nx.DiGraph(list(zip(p["before"].astype(str), p["after"].astype(str))))
            if not nx.is_directed_acyclic_graph(G):
                errs.append("precedences.csv contains a cycle (A before B before ... before A).")
    return errs


def _num(v, default):
    try:
        if v is None or str(v).strip() in ("", "nan"):
            return default
        return int(float(v))
    except ValueError:
        return default


def plan_build(data: Dict[str, pd.DataFrame], horizon: int) -> Tuple[CSP, pd.DataFrame]:
    t = data["tasks.csv"].copy()
    t["task_id"] = t["task_id"].astype(str)
    t["duration"] = pd.to_numeric(t["duration"]).astype(int)
    dur = dict(zip(t["task_id"], t["duration"]))
    doms = {}
    for _, r in t.iterrows():
        rel = _num(r.get("release"), 0)
        dl = min(_num(r.get("deadline"), horizon), horizon)
        doms[r["task_id"]] = list(range(0, horizon - r["duration"] + 1))
    csp = CSP(list(t["task_id"]), doms, "Scheduling")
    csp.value_formatter = lambda var, v: f"t={v}"
    for _, r in t.iterrows():
        rel = _num(r.get("release"), 0)
        dl = _num(r.get("deadline"), horizon)
        tid, d = r["task_id"], r["duration"]
        if rel > 0:
            csp.add_unary(tid, lambda s, rel=rel: s >= rel, f"release: start >= {rel}")
        if dl < horizon:
            csp.add_unary(tid, lambda s, dl=dl, d=d: s + d <= dl, f"deadline: finish <= {dl}")
    p = data.get("precedences.csv")
    prec = set()
    if p is not None and len(p):
        for _, r in p.iterrows():
            a, b = str(r["before"]), str(r["after"])
            gap = _num(r.get("min_gap"), 0)
            csp.add_binary(a, b, lambda sa, sb, da=dur[a], g=gap: sa + da + g <= sb,
                           f"precedence: {a} ends{f' + {gap}' if gap else ''} before {b} starts")
            prec.add(frozenset((a, b)))
    res_of = dict(zip(t["task_id"], t["resource"].astype(str)))
    ids = list(t["task_id"])
    for i, a in enumerate(ids):
        for b in ids[i + 1:]:
            if res_of[a] == res_of[b]:
                csp.add_binary(a, b, lambda sa, sb, da=dur[a], db=dur[b]: sa + da <= sb or sb + db <= sa,
                               f"no overlap on {res_of[a]}")
    return csp, t


def plan_lower_bound(data: Dict[str, pd.DataFrame]) -> int:
    t = data["tasks.csv"].copy()
    t["duration"] = pd.to_numeric(t["duration"]).astype(int)
    lb = int(t.groupby("resource")["duration"].sum().max())
    p = data.get("precedences.csv")
    if p is not None and len(p):
        import networkx as nx
        G = nx.DiGraph()
        dur = dict(zip(t["task_id"].astype(str), t["duration"]))
        for n in dur:
            G.add_node(n)
        for _, r in p.iterrows():
            G.add_edge(str(r["before"]), str(r["after"]), w=dur[str(r["before"])] + _num(r.get("min_gap"), 0))
        longest = {}
        for n in nx.topological_sort(G):
            longest[n] = max([longest[u] + G[u][n]["w"] for u in G.predecessors(n)] + [0])
        lb = max(lb, max(longest[n] + dur[n] for n in dur))
    return lb


def plan_run(data: Dict[str, pd.DataFrame], horizon: Optional[int] = None, optimize: bool = True, dataset: str = "") -> dict:
    t = data["tasks.csv"]
    total = int(pd.to_numeric(t["duration"]).sum())
    max_rel = max([_num(x, 0) for x in t.get("release", pd.Series([0]))] + [0])
    max_dl = max([_num(x, 0) for x in t.get("deadline", pd.Series([0]))] + [0])
    H = horizon or max(total + max_rel, max_dl)
    lb = plan_lower_bound(data)
    log = []
    csp, tt = plan_build(data, H)
    res = solve(csp, time_limit=20)
    log.append({"Horizon tried": H, "Status": res.status, "Makespan found": _makespan(res.solution, tt) if res.solution else "-",
                "AC-3 values pruned": res.ac3.values_removed, "Search nodes": res.nodes})
    best_csp, best_res, best_H = csp, res, H
    if optimize and res.solution:
        cur = _makespan(res.solution, tt)
        while cur > lb:
            H2 = cur - 1
            c2, _ = plan_build(data, H2)
            r2 = solve(c2, time_limit=10, node_limit=50_000)
            log.append({"Horizon tried": H2, "Status": r2.status, "Makespan found": _makespan(r2.solution, tt) if r2.solution else "-",
                        "AC-3 values pruned": r2.ac3.values_removed, "Search nodes": r2.nodes})
            if not r2.solution:
                break
            best_csp, best_res, best_H = c2, r2, H2
            cur = _makespan(r2.solution, tt)
    sol = best_res.solution
    a = best_res.ac3
    if sol:
        # CPM-style windows: AC-3 at horizon = achieved makespan shows the true slack / critical tasks
        ms0 = _makespan(sol, tt)
        c_ms, _ = plan_build(data, ms0)
        a_ms = ac3(c_ms)
        if a_ms.consistent:
            best_csp, a, best_H = c_ms, a_ms, ms0
            best_res.ac3 = a_ms
    rows = []
    for _, r in tt.iterrows():
        tid = r["task_id"]
        es, ls = (min(a.domains[tid]), max(a.domains[tid])) if a.domains[tid] else (None, None)
        s = sol[tid] if sol else None
        rows.append({"Task": tid, "Name": r["task_name"], "Job": r["job"], "Resource": r["resource"], "Duration": int(r["duration"]),
                     "Start": s if sol else "", "End": s + int(r["duration"]) if sol else "",
                     "Earliest start (AC-3)": es, "Latest start (AC-3)": ls,
                     "Window slack": (ls - es) if es is not None else None})
    sched = pd.DataFrame(rows)
    min_slack = sched["Window slack"].min() if sched["Window slack"].notna().any() else None
    sched["Bottleneck?"] = ["Yes" if min_slack is not None and v == min_slack else "" for v in sched["Window slack"]]
    if sol:
        sched = sched.sort_values(["Start", "Resource"]).reset_index(drop=True)
    util = pd.DataFrame()
    ms = _makespan(sol, tt) if sol else None
    if sol:
        util = sched.groupby("Resource")["Duration"].sum().reset_index().rename(columns={"Duration": "Busy time"})
        util["Makespan"] = ms
        util["Utilisation %"] = (100 * util["Busy time"] / ms).round(1)
    violations = best_csp.violations(sol) if sol else []
    return {"csp": best_csp, "search": best_res, "ac3": a, "schedule": sched, "util": util, "makespan": ms, "lower_bound": lb,
            "horizon": best_H, "log": pd.DataFrame(log), "violations": violations, "tasks": tt, "dataset": dataset,
            "optimal": bool(sol) and (ms == lb or (len(log) > 1 and log[-1]["Status"] == "unsatisfiable"))}


def _makespan(sol, tt) -> int:
    dur = dict(zip(tt["task_id"], tt["duration"]))
    return max(sol[k] + int(dur[k]) for k in sol)


def plan_gantt(r: dict) -> go.Figure:
    s = r["schedule"]
    fig = go.Figure()
    jobs = list(dict.fromkeys(r["tasks"]["job"]))
    for i, j in enumerate(jobs):
        d = s[s["Job"] == j]
        if d["Start"].eq("").all():
            continue
        fig.add_trace(go.Bar(y=d["Resource"], x=d["Duration"], base=d["Start"], orientation="h", name=j,
                             marker=dict(color=PALETTE[i % len(PALETTE)], line=dict(color="white", width=1.5)),
                             text=d["Task"], textposition="inside", insidetextanchor="middle",
                             hovertext=[f"{n}<br>{a} -> {b}" for n, a, b in zip(d["Name"], d["Start"], d["End"])], hoverinfo="text"))
    if r["makespan"]:
        fig.add_vline(x=r["makespan"], line=dict(color="#C62828", dash="dash"), annotation_text=f"makespan {r['makespan']}")
    fig.update_layout(title="Gantt chart - optimised schedule", barmode="overlay", height=120 + 45 * s["Resource"].nunique(),
                      xaxis_title="Time", plot_bgcolor="white", margin=dict(l=10, r=10, t=50, b=10),
                      legend=dict(orientation="h", y=1.12, x=0))
    fig.update_xaxes(gridcolor="#E5ECF6", dtick=1 if (r["makespan"] or 0) <= 30 else 5)
    return fig


def plan_window_fig(r: dict) -> go.Figure:
    s = r["schedule"].sort_values("Earliest start (AC-3)")
    fig = go.Figure()
    fig.add_trace(go.Bar(y=s["Task"], x=s["Latest start (AC-3)"] - s["Earliest start (AC-3)"] + s["Duration"],
                         base=s["Earliest start (AC-3)"], orientation="h", name="Feasible window after AC-3",
                         marker_color="rgba(11,79,156,0.18)", marker_line=dict(color="#0B4F9C", width=1)))
    if not s["Start"].eq("").all():
        fig.add_trace(go.Bar(y=s["Task"], x=s["Duration"], base=s["Start"], orientation="h", name="Chosen slot",
                             marker_color="#F28C28"))
    fig.update_layout(title="Start-time windows left by AC-3 (CPM-style) vs chosen slot",
                      barmode="overlay", height=160 + 24 * len(s), plot_bgcolor="white", xaxis_title="Time",
                      margin=dict(l=10, r=10, t=50, b=10), legend=dict(orientation="h", y=1.08, x=0))
    fig.update_yaxes(autorange="reversed")
    return fig


def plan_png(r: dict) -> bytes:
    s = r["schedule"]
    res_list = list(dict.fromkeys(s["Resource"]))
    jobs = list(dict.fromkeys(r["tasks"]["job"]))
    fig, ax = plt.subplots(figsize=(11, 0.6 * len(res_list) + 1.6))
    for _, row in s.iterrows():
        if row["Start"] == "":
            continue
        y = res_list.index(row["Resource"])
        ax.barh(y, row["Duration"], left=row["Start"], color=PALETTE[jobs.index(row["Job"]) % len(PALETTE)], edgecolor="white")
        ax.text(row["Start"] + row["Duration"] / 2, y, row["Task"], ha="center", va="center", color="white", fontsize=8, fontweight="bold")
    ax.set_yticks(range(len(res_list))); ax.set_yticklabels(res_list)
    ax.invert_yaxis(); ax.set_xlabel("Time"); ax.grid(axis="x", alpha=0.3)
    if r["makespan"]:
        ax.axvline(r["makespan"], color="#C62828", ls="--")
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color=PALETTE[i % len(PALETTE)], label=j) for i, j in enumerate(jobs)], loc="upper left",
              bbox_to_anchor=(1.01, 1.0), fontsize=8, frameon=False)
    ax.set_title("Gantt chart", color="#0B4F9C", fontweight="bold")
    return fig_to_png(fig)


def plan_report(r: dict) -> Report:
    a, s = r["ac3"], r["search"]
    if r["makespan"]:
        summary = (f"All {len(r['tasks'])} tasks were scheduled with a makespan of {r['makespan']} time units "
                   f"(lower bound {r['lower_bound']}; {'proven optimal' if r['optimal'] else 'best found within limits'}). "
                   f"{len(r['violations'])} precedence / resource / window rules are violated. AC-3 narrowed every task's "
                   f"start window; {int((r['schedule']['Bottleneck?'] == 'Yes').sum())} bottleneck tasks have the least slack "
                   f"({int(r['schedule']['Window slack'].min())} time units).")
    else:
        summary = "No feasible schedule exists for these release times, deadlines and precedences within the horizon."
    return Report(
        use_case="Planning & Resource Allocation", dataset=r["dataset"], summary=summary,
        metrics=metrics_pairs(a, s, {"Tasks": len(r["tasks"]), "Makespan": r["makespan"] or "-", "Lower bound": r["lower_bound"],
                                     "Optimal": "Yes" if r["optimal"] else "Not proven", "Final horizon": r["horizon"],
                                     "Violations (verified)": len(r["violations"])}),
        sections=[("How to read the results",
                   "'Schedule' gives the start/end of every task. 'Earliest/Latest start (AC-3)' is the window left after arc "
                   "consistency with the horizon set to the optimal makespan; bottleneck tasks have the least slack (slack 0 = strictly critical). 'Makespan search' shows how the horizon "
                   "was tightened until no schedule fits.")],
        formulas=[("S_t in {release_t, ..., H - d_t}", "Start time of task t; H = horizon, d_t = duration."),
                  ("S_a + d_a + gap <= S_b", "Precedence: a finishes (plus waiting gap) before b starts."),
                  ("S_a + d_a <= S_b  OR  S_b + d_b <= S_a", "Same resource: the two tasks don't overlap."),
                  ("S_t + d_t <= deadline_t", "Deadline (unary)."),
                  ("Makespan = max_t (S_t + d_t)", "Finish time of the whole plan (minimised by tightening H)."),
                  ("slack_t = latest start - earliest start (AC-3 at H = makespan)", "Zero slack = critical task: delaying it delays the whole plan."),
                  ("Lower bound = max(busiest resource load, longest precedence chain)", "No schedule can beat it."),
                  ("Utilisation % = 100 x busy time / makespan", "How loaded each resource is.")] + AC3_FORMULAS,
        tables=[("Schedule", r["schedule"]), ("Resource utilisation", r["util"]), ("Makespan search", r["log"]),
                ("AC-3 trace (first 300 revisions)", trace_df(a, 300))],
        figures=([("Gantt chart", plan_png(r))] if r["makespan"] else []) + [("Start-time domain before vs after AC-3", domain_reduction_png(a))],
    )


# =====================================================================================
# USE CASE 6 - VISION (Waltz line labelling) & NLP (constraint dependency parsing)
# =====================================================================================
import math as _math

# ---------------------------------------------------------------- 6A  Waltz filtering
# Huffman-Clowes catalogue for trihedral objects (as in Norvig, PAIP ch.17).
# Labels seen FROM the junction looking along the line:  '+' convex, '-' concave,
# 'R' occluding boundary with the object on the RIGHT, 'L' object on the LEFT.
# Neighbour order:  L: (arm1, arm2) with arm1 -> arm2 CLOCKWISE through the inside (<180 deg) angle;
#                   Y (fork): clockwise;
#                   W (arrow): (barb1, barb2, shaft) with barb1 -> shaft -> barb2 counter-clockwise;
#                   T: (bar1, bar2, stem) with bar1 -> stem -> bar2 counter-clockwise.
WALTZ_CATALOG = {
    "L": [("R", "L"), ("L", "R"), ("+", "R"), ("L", "+"), ("-", "L"), ("R", "-")],
    "Y": [("+", "+", "+"), ("-", "-", "-"), ("L", "R", "-"), ("-", "L", "R"), ("R", "-", "L")],
    "W": [("L", "R", "+"), ("-", "-", "+"), ("+", "+", "-")],
    "T": [("R", "L", "+"), ("R", "L", "-"), ("R", "L", "L"), ("R", "L", "R")],
}
JUNCTION_NAMES = {"L": "L-junction", "Y": "Fork (Y)", "W": "Arrow (W)", "T": "T-junction"}
_REV = {"+": "+", "-": "-", "L": "R", "R": "L"}
LABEL_MEANING = {"+": "convex edge", "-": "concave edge", "R": "occluding boundary (object on right)",
                 "L": "occluding boundary (object on left)"}

WALTZ_SCHEMA = {
    "vertices.csv": _schema([
        ["vertex_id", "text", "Yes", "Unique junction id (a CSP variable).", "A"],
        ["x", "number", "Yes", "Horizontal image coordinate.", "0.87"],
        ["y", "number", "Yes", "Vertical image coordinate (up = larger).", "0.5"],
    ]),
    "lines.csv": _schema([
        ["v1", "text", "Yes", "vertex_id at one end of a line segment.", "A"],
        ["v2", "text", "Yes", "vertex_id at the other end. Each vertex needs 2 or 3 lines (L, fork, arrow or T).", "B"],
    ]),
}


def _iso(X, Y, Z):
    return (round(0.866 * (X - Y), 4), round(-0.5 * (X + Y) + Z, 4))


def waltz_synthetic(kind: str = "Cube") -> Dict[str, pd.DataFrame]:
    if kind.startswith("Cube"):
        pts = {"a": (0, 0)}
        for name, ang in [("b", 90), ("e", 150), ("c", 210), ("f", 270), ("d", 330), ("g", 30)]:
            pts[name] = (round(_math.cos(_math.radians(ang)), 4), round(_math.sin(_math.radians(ang)), 4))
        lines = [("a", "b"), ("a", "c"), ("a", "d"), ("b", "e"), ("e", "c"), ("c", "f"), ("f", "d"), ("d", "g"), ("g", "b")]
    elif kind.startswith("Step"):
        P3 = {"A": (0, 0, 2), "B": (1, 0, 2), "C": (0, 1, 2), "D": (1, 1, 2), "E": (1, 1, 1), "F": (1, 0, 1),
              "G": (2, 0, 1), "H": (2, 1, 1), "I": (2, 1, 0), "J": (2, 0, 0), "K": (0, 1, 0)}
        pts = {k: _iso(*v) for k, v in P3.items()}
        lines = [("A", "B"), ("A", "C"), ("B", "D"), ("C", "D"), ("B", "F"), ("D", "E"), ("E", "F"), ("F", "G"),
                 ("E", "H"), ("G", "H"), ("H", "I"), ("G", "J"), ("J", "I"), ("I", "K"), ("K", "C")]
    else:  # tetrahedron / pyramid seen from above
        pts = {"apex": (1.0, 0.577), "p1": (0.0, 0.0), "p2": (2.0, 0.0), "p3": (1.0, 1.732)}
        lines = [("apex", "p1"), ("apex", "p2"), ("apex", "p3"), ("p1", "p2"), ("p2", "p3"), ("p3", "p1")]
    return {"vertices.csv": pd.DataFrame([dict(vertex_id=k, x=v[0], y=v[1]) for k, v in pts.items()]),
            "lines.csv": pd.DataFrame(lines, columns=["v1", "v2"])}


def _ang(p, q):
    return _math.degrees(_math.atan2(q[1] - p[1], q[0] - p[0])) % 360


def waltz_classify(pts: Dict[str, Tuple[float, float]], adj: Dict[str, List[str]]) -> Dict[str, Tuple[str, List[str]]]:
    """Return {vertex: (junction type, ordered neighbours)} using geometry only."""
    out = {}
    for v, ns in adj.items():
        if len(ns) == 2:
            p, q = ns
            if (_ang(pts[v], pts[q]) - _ang(pts[v], pts[p])) % 360 < 180:
                p, q = q, p          # make p -> q clockwise through the inside angle
            out[v] = ("L", [p, q])
            continue
        if len(ns) != 3:
            raise ValueError(f"Vertex {v} has {len(ns)} lines; Waltz labelling supports 2 or 3 lines per junction.")
        srt = sorted(ns, key=lambda n: _ang(pts[v], pts[n]))
        angs = [_ang(pts[v], pts[n]) for n in srt]
        gaps = [(angs[(i + 1) % 3] - angs[i]) % 360 for i in range(3)]
        big = max(range(3), key=lambda i: gaps[i])
        if abs(gaps[big] - 180) < 2.0:
            kind = "T"
        elif gaps[big] > 180:
            kind = "W"
        else:
            out[v] = ("Y", srt[::-1])   # clockwise
            continue
        # the largest gap is between srt[big] and srt[big+1]; the other three go ccw: srt[big+1] -> srt[big+2] -> srt[big]
        b1, mid, b2 = srt[(big + 1) % 3], srt[(big + 2) % 3], srt[big]
        out[v] = (kind, [b1, b2, mid])
    return out


def waltz_outline(pts, adj) -> List[Tuple[str, str]]:
    """Directed edges of the outer boundary, traversed counter-clockwise (object on the LEFT)."""
    start = min(pts, key=lambda v: (pts[v][0], pts[v][1]))
    back = 180.0
    cur, edges = start, []
    for _ in range(4 * sum(len(n) for n in adj.values())):
        cands = [(((_ang(pts[cur], pts[n]) - back) % 360) or 360, n) for n in adj[cur]]
        nxt = min(cands)[1]
        if (cur, nxt) in edges:
            break
        edges.append((cur, nxt))
        back = _ang(pts[nxt], pts[cur])
        cur = nxt
    return edges


def waltz_build(data: Dict[str, pd.DataFrame], assume_background: bool = True):
    v = data["vertices.csv"]
    pts = {str(r["vertex_id"]): (float(r["x"]), float(r["y"])) for _, r in v.iterrows()}
    adj: Dict[str, List[str]] = {k: [] for k in pts}
    for _, r in data["lines.csv"].iterrows():
        a, b = str(r["v1"]), str(r["v2"])
        adj[a].append(b)
        adj[b].append(a)
    jt = waltz_classify(pts, adj)
    doms = {k: list(WALTZ_CATALOG[jt[k][0]]) for k in pts}
    csp = CSP(list(pts), doms, "Waltz")
    csp.value_formatter = lambda var, lab: "(" + " ".join(lab) + ")"
    outline = waltz_outline(pts, adj) if assume_background else []
    for a, b in outline:
        ia = jt[a][1].index(b)
        csp.add_unary(a, lambda lab, ia=ia: lab[ia] == "L", f"outer boundary {a}-{b}: object inside")
    done = set()
    for a in pts:
        for b in adj[a]:
            if frozenset((a, b)) in done:
                continue
            done.add(frozenset((a, b)))
            ia, ib = jt[a][1].index(b), jt[b][1].index(a)
            csp.add_binary(a, b, lambda la, lb, ia=ia, ib=ib: la[ia] == _REV[lb[ib]], f"line {a}-{b} has one interpretation")
    return csp, pts, adj, jt, outline


def waltz_validate(data: Dict[str, pd.DataFrame]) -> List[str]:
    if "vertices.csv" not in data or "lines.csv" not in data:
        return ["Both vertices.csv and lines.csv are required."]
    errs = _need_cols(data["vertices.csv"], ["vertex_id", "x", "y"], "vertices.csv") + _need_cols(data["lines.csv"], ["v1", "v2"], "lines.csv")
    if errs:
        return errs
    ids = set(data["vertices.csv"]["vertex_id"].astype(str))
    bad = (set(data["lines.csv"]["v1"].astype(str)) | set(data["lines.csv"]["v2"].astype(str))) - ids
    if bad:
        errs.append(f"lines.csv references unknown vertices {sorted(bad)[:5]}.")
        return errs
    deg = pd.concat([data["lines.csv"]["v1"], data["lines.csv"]["v2"]]).astype(str).value_counts()
    wrong = [f"{k} ({deg.get(k, 0)} lines)" for k in ids if deg.get(k, 0) not in (2, 3)]
    if wrong:
        errs.append(f"Every vertex needs 2 or 3 lines; check: {wrong[:5]}.")
    return errs


def waltz_run(data: Dict[str, pd.DataFrame], assume_background: bool = True, dataset: str = "") -> dict:
    csp, pts, adj, jt, outline = waltz_build(data, assume_background)
    res = solve(csp, max_solutions=50, time_limit=10)
    rows = []
    for k in csp.variables:
        rows.append({"Vertex": k, "Junction type": JUNCTION_NAMES[jt[k][0]], "Lines (ordered)": ", ".join(f"{k}-{n}" for n in jt[k][1]),
                     "Catalogue labelings": len(csp.original_domains[k]), "After boundary rule": len(res.ac3.initial_domains[k]),
                     "After AC-3 (Waltz filter)": len(res.ac3.domains[k]),
                     "Remaining labelings": "; ".join(csp.fmt(k, x) for x in res.ac3.domains[k])})
    vt = pd.DataFrame(rows)
    interps = []
    for si, sol in enumerate(res.solutions, 1):
        for a, b in {tuple(sorted(e)) for e in [(a, b) for a in adj for b in adj[a]]}:
            lab = sol[a][jt[a][1].index(b)]
            interps.append({"Interpretation": si, "Line": f"{a}-{b}", "Label (seen from first vertex)": lab,
                            "Meaning": LABEL_MEANING[lab]})
    lt = pd.DataFrame(interps)
    return {"csp": csp, "search": res, "ac3": res.ac3, "pts": pts, "adj": adj, "jt": jt, "outline": outline,
            "vertex_table": vt, "line_table": lt, "dataset": dataset, "assume_background": assume_background,
            "violations": sum(len(csp.violations(s)) for s in res.solutions)}


def _waltz_line_labels(r: dict, sol: dict) -> Dict[Tuple[str, str], str]:
    out = {}
    for a in r["adj"]:
        for b in r["adj"][a]:
            out[(a, b)] = sol[a][r["jt"][a][1].index(b)]
    return out


def waltz_fig(r: dict, idx: int = 0) -> go.Figure:
    pts, adj = r["pts"], r["adj"]
    fig = go.Figure()
    sol = r["search"].solutions[idx] if r["search"].solutions else None
    labs = _waltz_line_labels(r, sol) if sol else {}
    col = {"+": "#1E8E5A", "-": "#C62828", "L": "#0B4F9C", "R": "#0B4F9C"}
    seen = set()
    for a in adj:
        for b in adj[a]:
            if frozenset((a, b)) in seen:
                continue
            seen.add(frozenset((a, b)))
            lab = labs.get((a, b))
            (x0, y0), (x1, y1) = pts[a], pts[b]
            fig.add_trace(go.Scatter(x=[x0, x1], y=[y0, y1], mode="lines", showlegend=False, hoverinfo="skip",
                                     line=dict(color=col.get(lab, "#555"), width=4)))
            mx, my = (x0 + x1) / 2, (y0 + y1) / 2
            if lab in ("+", "-"):
                fig.add_annotation(x=mx, y=my, text="<b>" + ("+" if lab == "+" else "−") + "</b>", showarrow=False,
                                   font=dict(size=22, color=col[lab]), bgcolor="white", borderpad=0)
            elif lab in ("L", "R"):
                # arrow direction: object on the RIGHT of travel.  'R' seen from a => travel a->b
                (sx, sy), (tx, ty) = ((x0, y0), (x1, y1)) if lab == "R" else ((x1, y1), (x0, y0))
                fig.add_annotation(x=mx + (tx - sx) * 0.12, y=my + (ty - sy) * 0.12, ax=mx - (tx - sx) * 0.12,
                                   ay=my - (ty - sy) * 0.12, xref="x", yref="y", axref="x", ayref="y", showarrow=True,
                                   arrowhead=2, arrowsize=1.6, arrowwidth=2.5, arrowcolor="#F28C28", text="")
    fig.add_trace(go.Scatter(x=[p[0] for p in pts.values()], y=[p[1] for p in pts.values()], mode="markers+text",
                             text=[f"{k} ({r['jt'][k][0]})" for k in pts], textposition="top right", showlegend=False,
                             marker=dict(size=11, color="#07326A"), textfont=dict(size=12, color="#07326A")))
    fig.update_layout(title=f"Line labelling - interpretation {idx + 1} of {len(r['search'].solutions)}" if sol else "No consistent labelling",
                      height=500, plot_bgcolor="white", margin=dict(l=10, r=10, t=50, b=10),
                      xaxis=dict(visible=False), yaxis=dict(visible=False, scaleanchor="x"))
    return fig


def waltz_png(r: dict, idx: int = 0) -> bytes:
    pts, adj = r["pts"], r["adj"]
    sol = r["search"].solutions[idx] if r["search"].solutions else None
    labs = _waltz_line_labels(r, sol) if sol else {}
    fig, ax = plt.subplots(figsize=(7, 6))
    col = {"+": "#1E8E5A", "-": "#C62828", "L": "#0B4F9C", "R": "#0B4F9C"}
    seen = set()
    for a in adj:
        for b in adj[a]:
            if frozenset((a, b)) in seen:
                continue
            seen.add(frozenset((a, b)))
            lab = labs.get((a, b))
            (x0, y0), (x1, y1) = pts[a], pts[b]
            ax.plot([x0, x1], [y0, y1], color=col.get(lab, "#555"), lw=3)
            mx, my = (x0 + x1) / 2, (y0 + y1) / 2
            if lab in ("+", "-"):
                ax.text(mx, my, lab, fontsize=18, color=col[lab], ha="center", va="center", fontweight="bold",
                        bbox=dict(facecolor="white", edgecolor="none", pad=0.5))
            elif lab:
                (sx, sy), (tx, ty) = ((x0, y0), (x1, y1)) if lab == "R" else ((x1, y1), (x0, y0))
                ax.annotate("", xy=(mx + (tx - sx) * 0.12, my + (ty - sy) * 0.12), xytext=(mx - (tx - sx) * 0.12, my - (ty - sy) * 0.12),
                            arrowprops=dict(arrowstyle="-|>", color="#F28C28", lw=2.5, mutation_scale=22))
    for k, (x, y) in pts.items():
        ax.plot(x, y, "o", color="#07326A")
        ax.text(x, y, f"  {k}", fontsize=10, color="#07326A")
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("Waltz line labelling (+ convex, - concave, arrow = occluding boundary)", color="#0B4F9C", fontweight="bold", fontsize=10)
    return fig_to_png(fig)


# ---------------------------------------------------------------- 6B  Constraint dependency parsing
NLP_LEXICON = {
    "DET": "the a an this that every some my his her their",
    "ADJ": "old young small big red green fresh long quick happy tall heavy new bright tired lazy wooden",
    "NOUN": ("man woman girl boy dog cat mouse ball telescope park hill book library student teacher chef meal "
             "herbs robot box table letter fork pizza bridge river engineer city car road garden window bird tree "
             "friend school child cake knife office report manager team"),
    "PROPN": "john mary alice bob london paris",
    "PRON": "she he they we i you it",
    "VERB": ("saw chased read wrote cooked moved ate designed barked sat found built bought liked watched opened "
             "painted visited sleeps runs sings reads eats loves sees"),
    "PREP": "with in on near to from under by at for behind over",
    "ADV": "quickly slowly carefully happily yesterday today often",
}

NLP_SENTENCES = [
    "the old man saw a dog with a telescope",
    "a small cat chased the red ball",
    "the student read the book in the library",
    "she quickly wrote a long letter",
    "the chef cooked a meal with fresh herbs",
    "john ate pizza with a fork",
    "the robot moved the box to the table",
    "the engineer designed a bridge near the river",
    "a girl saw the man on the hill with a telescope",
    "the big dog barked",
    "cat the chased mouse the",
]

NLP_SCHEMA = {
    "sentences.txt": _schema([
        ["(one sentence per line)", "text", "Yes", "Lower/upper case both fine; punctuation is ignored. Each sentence is parsed separately.", "the old man saw a dog"],
    ]),
    "lexicon.csv (optional)": _schema([
        ["word", "text", "Yes", "A word form.", "telescope"],
        ["pos", "DET / ADJ / NOUN / PROPN / PRON / VERB / PREP / ADV", "Yes", "Part of speech. It decides which grammatical roles the word may take.", "NOUN"],
    ]),
}

ROLE_MEANING = {"root": "main verb of the sentence", "det": "determiner of a noun", "amod": "adjective modifying a noun",
                "nsubj": "subject of a verb", "obj": "direct object of a verb", "pobj": "object of a preposition",
                "prep": "prepositional phrase attached to a noun or verb", "advmod": "adverb modifying a verb"}


def nlp_default_lexicon() -> pd.DataFrame:
    return pd.DataFrame([(w, p) for p, ws in NLP_LEXICON.items() for w in ws.split()], columns=["word", "pos"])


def nlp_build(words: List[str], pos: List[str]) -> CSP:
    n = len(words)
    N = range(1, n + 1)
    P = {i: pos[i - 1] for i in N}
    nominal = {"NOUN", "PROPN", "PRON"}
    doms = {}
    for i in N:
        p, d = P[i], []
        if p == "VERB":
            d.append(("root", 0))
        for h in N:
            if h == i:
                continue
            ph = P[h]
            if p == "DET" and ph == "NOUN" and h > i:
                d.append(("det", h))
            elif p == "ADJ" and ph == "NOUN" and h > i:
                d.append(("amod", h))
            elif p in nominal:
                if ph == "VERB" and h > i:
                    d.append(("nsubj", h))
                if ph == "VERB" and h < i:
                    d.append(("obj", h))
                if ph == "PREP" and h < i:
                    d.append(("pobj", h))
            elif p == "PREP" and (ph in nominal or ph == "VERB") and h < i:
                d.append(("prep", h))
            elif p == "ADV" and ph == "VERB":
                d.append(("advmod", h))
        doms[i] = d
    csp = CSP(list(N), doms, "Dependency parse")
    csp.value_formatter = lambda var, v: f"{v[0]}->{words[v[1] - 1] if v[1] else 'ROOT'}"

    def cross(i, hi, j, hj):
        a1, b1 = sorted((i, hi))
        a2, b2 = sorted((j, hj))
        return (a1 < a2 < b1 < b2) or (a2 < a1 < b2 < b1)

    unique_roles = {"det", "nsubj", "obj", "pobj", "root"}
    for i in N:
        for j in N:
            if j <= i:
                continue
            def ok(vi, vj, i=i, j=j):
                (ri, hi), (rj, hj) = vi, vj
                if hi == j and hj == i:
                    return False                           # no 2-cycles
                if cross(i, hi, j, hj):
                    return False                           # projective (no crossing arcs)
                if hi == hj and ri == rj and ri in unique_roles:
                    return False                           # one subject/object/determiner per head, one root
                if hi == hj and ri == "amod" and rj == "det":
                    return False                           # determiner comes before adjectives
                return True
            csp.add_binary(i, j, ok, "projective, unique roles, no cycles, det before adj")
    return csp


def nlp_tokenize(sentence: str) -> List[str]:
    import re
    return [w for w in re.findall(r"[A-Za-z']+", sentence.lower())]


def nlp_parse(sentence: str, lexicon: pd.DataFrame, unknown_as_noun: bool = True) -> dict:
    lex = dict(zip(lexicon["word"].astype(str).str.lower(), lexicon["pos"].astype(str).str.upper()))
    words = nlp_tokenize(sentence)
    unknown = [w for w in words if w not in lex]
    if unknown and not unknown_as_noun:
        raise ValueError(f"Unknown word(s): {unknown}")
    pos = [lex.get(w, "NOUN") for w in words]
    csp = nlp_build(words, pos)
    res = solve(csp, max_solutions=50, time_limit=10)
    # global rule (not binary): every preposition must have exactly one object
    good = []
    for s in res.solutions:
        ok = all(sum(1 for j, v in s.items() if v == ("pobj", i)) == 1 for i in s if pos[i - 1] == "PREP")
        ok = ok and sum(1 for v in s.values() if v[0] == "root") == 1
        if ok:
            good.append(s)
    return {"sentence": sentence, "words": words, "pos": pos, "unknown": unknown, "csp": csp, "search": res,
            "ac3": res.ac3, "parses": good}


def nlp_run(sentences: List[str], lexicon: pd.DataFrame, dataset: str = "", unknown_as_noun: bool = True) -> dict:
    parsed = [nlp_parse(s, lexicon, unknown_as_noun) for s in sentences if s.strip()]
    summary = pd.DataFrame([{"#": i + 1, "Sentence": p["sentence"], "Words": len(p["words"]),
                             "Values before AC-3": sum(len(d) for d in p["ac3"].initial_domains.values()),
                             "Values after AC-3": sum(len(d) for d in p["ac3"].domains.values()),
                             "Valid parses": len(p["parses"]),
                             "Verdict": ("ambiguous" if len(p["parses"]) > 1 else "unique parse" if p["parses"] else "ungrammatical"),
                             "Unknown words": ", ".join(p["unknown"])} for i, p in enumerate(parsed)])
    detail = []
    for i, p in enumerate(parsed, 1):
        for k, s in enumerate(p["parses"], 1):
            for w in range(1, len(p["words"]) + 1):
                role, h = s[w]
                detail.append({"Sentence #": i, "Parse #": k, "Position": w, "Word": p["words"][w - 1], "POS": p["pos"][w - 1],
                               "Role": role, "Head word": p["words"][h - 1] if h else "ROOT", "Head position": h,
                               "Meaning": ROLE_MEANING.get(role, "")})
    return {"parsed": parsed, "summary": summary, "detail": pd.DataFrame(detail), "dataset": dataset}


def nlp_fig(p: dict, k: int = 0) -> go.Figure:
    words, pos = p["words"], p["pos"]
    fig = go.Figure()
    n = len(words)
    fig.add_trace(go.Scatter(x=list(range(1, n + 1)), y=[0] * n, mode="text", text=[f"<b>{w}</b>" for w in words],
                             textfont=dict(size=16, color="#07326A"), showlegend=False, hoverinfo="skip"))
    fig.add_trace(go.Scatter(x=list(range(1, n + 1)), y=[-0.35] * n, mode="text", text=pos,
                             textfont=dict(size=11, color="#6B7C93"), showlegend=False, hoverinfo="skip"))
    if p["parses"]:
        s = p["parses"][k]
        roles = sorted({v[0] for v in s.values()})
        for i, (role, h) in s.items():
            c = PALETTE[roles.index(role) % len(PALETTE)]
            if h == 0:
                fig.add_annotation(x=i, y=0.2, ax=i, ay=n * 0.35 + 0.8, xref="x", yref="y", axref="x", ayref="y",
                                   arrowhead=2, arrowcolor=c, arrowwidth=2, showarrow=True, text="")
                fig.add_annotation(x=i, y=n * 0.35 + 1.0, text="ROOT", showarrow=False, font=dict(color=c, size=12))
                continue
            a, b = sorted((i, h))
            height = 0.35 * (b - a) + 0.3
            xs = [a + (b - a) * t / 40 for t in range(41)]
            ys = [0.2 + height * _math.sin(_math.pi * t / 40) for t in range(41)]
            fig.add_trace(go.Scatter(x=xs, y=ys, mode="lines", line=dict(color=c, width=2.5), showlegend=False,
                                     hovertext=f"{words[i-1]} --{role}--> {words[h-1]}", hoverinfo="text"))
            fig.add_annotation(x=i, y=0.22, ax=i + (0.12 if h > i else -0.12), ay=0.45, xref="x", yref="y", axref="x", ayref="y",
                               arrowhead=2, arrowcolor=c, arrowwidth=2, showarrow=True, text="")
            fig.add_annotation(x=(a + b) / 2, y=0.2 + height, text=role, showarrow=False, font=dict(color=c, size=12),
                               bgcolor="white")
    fig.update_layout(title=(f"Dependency parse {k + 1} of {len(p['parses'])}" if p["parses"] else "No valid parse - sentence is ungrammatical for this grammar"),
                      height=300 + 22 * n, plot_bgcolor="white", margin=dict(l=10, r=10, t=50, b=10),
                      xaxis=dict(visible=False, range=[0.3, n + 0.7]), yaxis=dict(visible=False, range=[-0.7, n * 0.35 + 1.4]))
    return fig


def nlp_png(p: dict, k: int = 0) -> bytes:
    words = p["words"]
    n = len(words)
    fig, ax = plt.subplots(figsize=(max(7, 1.1 * n), 3.4))
    for i, w in enumerate(words, 1):
        ax.text(i, 0, w, ha="center", va="center", fontsize=12, fontweight="bold", color="#07326A")
        ax.text(i, -0.35, p["pos"][i - 1], ha="center", va="center", fontsize=8, color="#6B7C93")
    if p["parses"]:
        s = p["parses"][k]
        roles = sorted({v[0] for v in s.values()})
        import numpy as np
        for i, (role, h) in s.items():
            c = PALETTE[roles.index(role) % len(PALETTE)]
            if h == 0:
                ax.annotate("ROOT", xy=(i, 0.2), xytext=(i, n * 0.3 + 0.6), ha="center", color=c,
                            arrowprops=dict(arrowstyle="-|>", color=c))
                continue
            a, b = sorted((i, h))
            t = np.linspace(0, 1, 40)
            hgt = 0.3 * (b - a) + 0.25
            ax.plot(a + (b - a) * t, 0.2 + hgt * np.sin(np.pi * t), color=c, lw=2)
            ax.annotate("", xy=(i, 0.2), xytext=(i + (0.1 if h > i else -0.1), 0.42), arrowprops=dict(arrowstyle="-|>", color=c))
            ax.text((a + b) / 2, 0.22 + hgt, role, ha="center", fontsize=9, color=c,
                    bbox=dict(facecolor="white", edgecolor="none", pad=0.5))
    ax.set_xlim(0.4, n + 0.6); ax.set_ylim(-0.6, n * 0.3 + 1.0); ax.axis("off")
    ax.set_title(f"'{p['sentence']}'  -  parse {k + 1} of {len(p['parses'])}", color="#0B4F9C", fontweight="bold", fontsize=10)
    return fig_to_png(fig)


def perception_report_waltz(r: dict) -> Report:
    a, s = r["ac3"], r["search"]
    n = len(s.solutions)
    summary = (f"The drawing has {len(r['pts'])} junctions and {sum(len(v) for v in r['adj'].values()) // 2} lines. "
               f"Waltz filtering (AC-3) cut the candidate junction labelings from {sum(len(d) for d in a.initial_domains.values())} "
               f"to {sum(len(d) for d in a.domains.values())}. "
               + (f"{n} physically consistent 3-D interpretation{'s' if n != 1 else ''} remain"
                  + (" - AC-3 alone found it without any search." if a.solved else ".") if n else
                  "No consistent interpretation exists - the drawing is an impossible object under trihedral assumptions."))
    return Report(
        use_case="Vision - Waltz Line Labelling", dataset=r["dataset"], summary=summary,
        metrics=metrics_pairs(a, s, {"Junctions": len(r["pts"]), "Interpretations": n,
                                     "Background assumption": "on" if r["assume_background"] else "off",
                                     "Violations (verified)": r["violations"]}),
        sections=[("How to read the results",
                   "'+' = convex edge (sticks out), '-' = concave edge (folds inwards), arrows = occluding boundary with "
                   "the object on the right-hand side of the arrow. Every junction must use a labeling from the "
                   "Huffman-Clowes catalogue, and a line must get the same interpretation at both ends.")],
        formulas=[("X_v in catalogue(type(v))", "Junction v may only use the physically possible labelings of its type (L, fork, arrow, T)."),
                  ("label_v(v,w) = reverse(label_w(w,v))", "+ and - stay the same, L at one end = R at the other end."),
                  ("outer boundary edge -> object inside", "Optional unary rule: the figure is surrounded by background.")] + AC3_FORMULAS,
        tables=[("Junctions", r["vertex_table"]), ("Line interpretations", r["line_table"]),
                ("AC-3 trace", trace_df(a, 300))],
        figures=[("Labelled drawing", waltz_png(r))] if s.solutions else [],
    )


def perception_report_nlp(r: dict) -> Report:
    n_amb = int((r["summary"]["Valid parses"] > 1).sum())
    n_bad = int((r["summary"]["Valid parses"] == 0).sum())
    summary = (f"{len(r['parsed'])} sentences were parsed with a constraint dependency grammar. AC-3 removed "
               f"{int((r['summary']['Values before AC-3'] - r['summary']['Values after AC-3']).sum())} impossible (role, head) "
               f"choices in total. {n_amb} sentences are structurally ambiguous (e.g. prepositional-phrase attachment) and "
               f"{n_bad} were rejected as ungrammatical.")
    figs = [(f"Parse of '{p['sentence']}'", nlp_png(p)) for p in r["parsed"][:6] if p["parses"]]
    return Report(
        use_case="NLP - Constraint Dependency Parsing", dataset=r["dataset"], summary=summary,
        metrics=[("Sentences", str(len(r["parsed"]))), ("Ambiguous", str(n_amb)), ("Ungrammatical", str(n_bad)),
                 ("Total parses", str(int(r["summary"]["Valid parses"].sum())))],
        sections=[("How to read the results",
                   "Each word points to its head word with a grammatical role (subject, object, determiner ...). "
                   "Several parses for one sentence mean real ambiguity: 'saw a dog with a telescope' - who has the telescope?")],
        formulas=[("X_i = (role_i, head_i)", "Word i chooses a grammatical role and the word it depends on (0 = ROOT)."),
                  ("POS(i) restricts role_i and direction of head_i", "Unary rules, e.g. a determiner must point to a noun on its right."),
                  ("arcs (i, h_i) and (j, h_j) must not cross", "Projectivity (binary)."),
                  ("h_i = h_j and role_i = role_j in {det, nsubj, obj, pobj}  =>  violation", "At most one subject/object/determiner per head."),
                  ("not (h_i = j and h_j = i)", "No two-word cycles."),
                  ("every preposition has exactly one pobj", "Global rule checked on complete parses.")] + AC3_FORMULAS,
        tables=[("Sentence summary", r["summary"]), ("Parse details", r["detail"])],
        figures=figs,
    )


# =====================================================================================
# STREAMLIT USER INTERFACE
# =====================================================================================
import hashlib as _hashlib

try:
    import streamlit as st
except ImportError:  # the engine + tests work without Streamlit installed
    st = None

CSS = """
<style>
:root { --knet-blue:#0B4F9C; --knet-dark:#07326A; --knet-light:#E8F1FB; --knet-accent:#F28C28; }
.block-container { padding-top: 1.2rem; max-width: 1400px; }
.app-title { font-size: 2.9rem !important; font-weight: 900 !important; color: #0B4F9C !important; line-height: 1.15 !important;
    margin: 0 !important; letter-spacing: -0.5px; font-family: "Segoe UI", Helvetica, Arial, sans-serif; }
.app-dev   { font-size: 2.05rem !important; font-weight: 800 !important; color: #0B4F9C !important; line-height: 1.25 !important;
    margin: 0.2rem 0 0.25rem 0 !important; font-family: "Segoe UI", Helvetica, Arial, sans-serif; }
[data-testid="stAppDeployButton"], [data-testid="stMainMenu"] { display: none !important; }
header[data-testid="stHeader"] { background: transparent; }
.app-tag   { color:#4A5B72; font-size: 1.02rem; margin-bottom: 0.6rem; }
.hero-rule { height: 5px; border-radius: 3px; background: linear-gradient(90deg,#0B4F9C 0%,#1E88E5 45%,#F28C28 100%); margin: 0.4rem 0 1.0rem 0; }
section[data-testid="stSidebar"] { background: linear-gradient(180deg,#07326A 0%,#0B4F9C 60%,#1565C0 100%); }
section[data-testid="stSidebar"] * { color: #FFFFFF !important; }
section[data-testid="stSidebar"] .stRadio > div { gap: 0.35rem; }
section[data-testid="stSidebar"] .stRadio [role="radiogroup"] { width: 100%; }
section[data-testid="stSidebar"] .stRadio [role="radiogroup"] label { background: rgba(255,255,255,0.08); border: 1px solid rgba(255,255,255,0.18);
    border-radius: 10px; padding: 0.6rem 0.75rem; width: 100% !important; display: flex !important; transition: all .15s; }
section[data-testid="stSidebar"] .stRadio [role="radiogroup"] label p { font-size: 0.98rem !important; font-weight: 600; }
section[data-testid="stSidebar"] .stRadio [role="radiogroup"] label:hover { background: rgba(255,255,255,0.22); }
section[data-testid="stSidebar"] .stRadio [role="radiogroup"] label:has(input:checked) { background: #F28C28; border-color:#FFD2A6; box-shadow: 0 2px 8px rgba(0,0,0,.25); }
.side-brand { font-size: 1.35rem; font-weight: 900; letter-spacing: .5px; margin-bottom: .1rem; }
.side-sub { font-size: .82rem; opacity: .85; margin-bottom: .8rem; }
.uc-banner { border-radius: 14px; padding: 1.0rem 1.3rem; color: white; margin-bottom: 0.8rem;
    background: linear-gradient(120deg, var(--c1) 0%, var(--c2) 100%); box-shadow: 0 3px 12px rgba(11,79,156,.18); }
.uc-banner h2 { color: white !important; margin: 0 0 .2rem 0; font-size: 1.65rem; }
.uc-banner p { margin: 0; font-size: 1.0rem; opacity: .95; }
.kpi { border-radius: 12px; padding: .75rem .9rem; background: white; border: 1px solid #DCE6F2; border-left: 6px solid var(--k);
    box-shadow: 0 1px 4px rgba(7,50,106,.06); height: 100%; }
.kpi .v { font-size: 1.55rem; font-weight: 800; color: var(--k); line-height: 1.1; }
.kpi .l { font-size: .80rem; color: #4A5B72; text-transform: uppercase; letter-spacing: .4px; }
.card { background: #F6F9FE; border: 1px solid #DCE6F2; border-radius: 12px; padding: .9rem 1.1rem; margin-bottom: .7rem; }
.card h4 { color: #0B4F9C; margin-top: 0; }
.pill { display:inline-block; padding:.12rem .55rem; border-radius: 999px; font-size:.8rem; font-weight:700; margin-right:.3rem; }
.ok  { background:#E3F4EA; color:#1E8E5A; } .bad { background:#FDE7E7; color:#C62828; } .info{ background:#E8F1FB; color:#0B4F9C; }
.formula-expl { background:#FFF8EF; border-left: 4px solid #F28C28; padding:.45rem .8rem; border-radius: 6px; margin-bottom: .8rem; font-size:.95rem; }
.stTabs [role="tablist"] { gap: 6px; border-bottom: 3px solid #0B4F9C; }
.stTabs [role="tab"] { background: #EEF4FB; border-radius: 10px 10px 0 0; padding: .55rem 1.1rem !important; margin-right: 4px; }
.stTabs [role="tab"]:hover { background: #D6E6F8; }
.stTabs [role="tab"] p { font-weight: 700 !important; color:#07326A; font-size: 1.0rem !important; }
.stTabs [role="tab"][aria-selected="true"] { background: #0B4F9C !important; border-bottom: none !important; }
.stTabs [role="tab"][aria-selected="true"] p { color: #FFFFFF !important; }
.stTabs [data-baseweb="tab-highlight"], .stTabs [data-baseweb="tab-border"] { display: none; }
.stDownloadButton button { width: 100%; font-weight: 700; border-radius: 10px; padding: .7rem; }
div.stButton > button[kind="primary"] { background: #0B4F9C; border-color: #0B4F9C; font-weight: 700; }
.footer { color:#6B7C93; font-size:.85rem; text-align:center; margin-top: 2rem; padding-top: .8rem; border-top: 1px solid #DCE6F2; }
</style>
"""

USE_CASES = [
    ("home", "🏠  Home - What is AC-3?"),
    ("sudoku", "1️⃣  Sudoku & Logic Puzzles"),
    ("sched", "2️⃣  Scheduling & Timetabling"),
    ("color", "3️⃣  Map Colouring & Frequencies"),
    ("config", "4️⃣  Product Configuration"),
    ("plan", "5️⃣  Planning & Resource Allocation"),
    ("percep", "6️⃣  Vision & Language (NLP)"),
]


# ------------------------------------------------------------------ small UI helpers
def ui_header():
    st.markdown(f'<div class="app-title">{APP_NAME}</div><div class="app-dev">{DEVELOPER_LINE}</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="app-tag">{APP_TAGLINE}</div><div class="hero-rule"></div>', unsafe_allow_html=True)



def ui_tabs(page_key: str, labels: List[str]):
    """Tabs that always open on '1. Understand the use case' when a use case is (re)selected."""
    key = f"tabs_{page_key}"
    if st.session_state.get("_last_page") != page_key:
        st.session_state["_last_page"] = page_key
        st.session_state.pop(key, None)
    try:
        return st.tabs(labels, key=key, default=labels[0], on_change="rerun")
    except TypeError:          # older Streamlit without stateful tabs
        return st.tabs(labels)

def ui_banner(title, text, c1="#0B4F9C", c2="#1E88E5"):
    st.markdown(f'<div class="uc-banner" style="--c1:{c1};--c2:{c2}"><h2>{title}</h2><p>{text}</p></div>', unsafe_allow_html=True)


def ui_kpis(items):
    cols = st.columns(len(items))
    for col, (label, value, color) in zip(cols, items):
        col.markdown(f'<div class="kpi" style="--k:{color}"><div class="v">{value}</div><div class="l">{label}</div></div>',
                     unsafe_allow_html=True)
    st.write("")


def ui_formula(latex, meaning):
    st.latex(latex)
    st.markdown(f'<div class="formula-expl">{meaning}</div>', unsafe_allow_html=True)


def ui_csp_mapping(rows):
    st.markdown("##### How this problem maps onto a CSP")
    st.dataframe(pd.DataFrame(rows, columns=["CSP element", "In this use case", "Example"]), hide_index=True)


def ui_schema(schema: Dict[str, pd.DataFrame], templates: Optional[Dict[str, pd.DataFrame]] = None, key: str = ""):
    st.markdown("##### Data schema - every file and field explained")
    for fname, df in schema.items():
        with st.expander(f"📄 {fname}", expanded=True):
            st.dataframe(df, hide_index=True)
            base = fname.split()[0]
            if templates and base in templates:
                st.download_button(f"⬇️ Download sample {base} (synthetic data, ready to edit)",
                                   templates[base].to_csv(index=False).encode("utf-8-sig"), file_name=base, mime="text/csv",
                                   key=f"tpl_{key}_{base}")


def ui_ac3_generic():
    with st.expander("🔁 How AC-3 itself works (applies to every tab)", expanded=False):
        c1, c2 = st.columns([1, 1])
        with c1:
            st.graphviz_chart(AC3_FLOW_DOT)
        with c2:
            st.code(AC3_PSEUDOCODE, language="text")
        for f, m in AC3_FORMULAS:
            st.markdown(f"- `{f}`  \n  {m}")


def ui_ac3_details(res: AC3Result, csp: CSP, key: str, label_map=None):
    with st.expander("🔍 AC-3 details - domains, progress and step-by-step trace", expanded=False):
        c1, c2 = st.columns([3, 2])
        with c1:
            st.plotly_chart(domain_reduction_plotly(res, label_map=label_map), key=f"dr_{key}")
        with c2:
            st.plotly_chart(pruning_timeline_plotly(res), key=f"pt_{key}")
        st.markdown("**Domains per variable**")
        st.dataframe(res.domain_table(csp), hide_index=True, height=260)
        st.markdown(f"**Trace of successful revisions** ({len(res.trace)} shown"
                    f"{', truncated' if res.trace_truncated else ''}). Each row is one REVISE(Xi, Xj) that deleted values.")
        st.dataframe(trace_df(res, 1000), hide_index=True, height=300)


def ui_upload(key: str, files: List[str], optional: Tuple[str, ...] = ()) -> Optional[Dict[str, pd.DataFrame]]:
    st.info("Upload CSV files that follow the schema shown in **1. Understand the use case**. "
            "Tip: download the sample files there, edit them in Excel, and upload them here.")
    data, cols = {}, st.columns(min(len(files), 4))
    for i, f in enumerate(files):
        up = cols[i % len(cols)].file_uploader(f"{f}{' (optional)' if f in optional else ''}", type=["csv"], key=f"up_{key}_{f}")
        if up is not None:
            try:
                data[f] = pd.read_csv(up)
            except Exception as e:  # noqa: BLE001
                st.error(f"Could not read {f}: {e}")
    missing = [f for f in files if f not in data and f not in optional]
    if missing:
        st.warning(f"Waiting for: {', '.join(missing)}")
        return None
    return data


def _sig(obj) -> str:
    try:
        if isinstance(obj, dict):
            raw = "|".join(f"{k}:{pd.util.hash_pandas_object(v, index=False).sum() if isinstance(v, pd.DataFrame) else v}"
                           for k, v in sorted(obj.items()))
        else:
            raw = str(obj)
    except Exception:  # noqa: BLE001
        raw = str(obj)
    return _hashlib.md5(raw.encode()).hexdigest()


def ui_set_data(key: str, data, label: str):
    sig = _sig(data)
    if st.session_state.get(f"{key}_sig") != sig:
        st.session_state[f"{key}_sig"] = sig
        st.session_state.pop(f"{key}_result", None)
        st.session_state.pop(f"{key}_report", None)
    st.session_state[f"{key}_data"] = data
    st.session_state[f"{key}_label"] = label


def ui_export(key: str, builder: Callable[[dict], Report], slug: str):
    r = st.session_state.get(f"{key}_result")
    if r is None:
        st.info("Run the analysis in tab **3** first - the report is built from its results.")
        return
    if st.session_state.get(f"{key}_report") is None:
        with st.spinner("Building report (PDF, Word, CSV, TXT)..."):
            rep = builder(r)
            st.session_state[f"{key}_report"] = {"rep": rep, "pdf": to_pdf(rep), "docx": to_docx(rep),
                                                "csv": to_csv(rep), "txt": to_txt(rep)}
    b = st.session_state[f"{key}_report"]
    rep = b["rep"]
    st.markdown(f"#### 📤 Download the results - *{rep.title}*")
    st.caption("Each file carries the application title, the developer line, an executive summary, key metrics, "
               "explanations, formulas, charts and every results table.")
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M")
    c1, c2, c3, c4 = st.columns(4)
    c1.download_button("📕 PDF report", b["pdf"], f"{slug}_report_{stamp}.pdf", "application/pdf", key=f"pdf_{key}")
    c2.download_button("📘 Word report", b["docx"], f"{slug}_report_{stamp}.docx",
                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=f"docx_{key}")
    c3.download_button("📗 CSV results", b["csv"], f"{slug}_results_{stamp}.csv", "text/csv", key=f"csv_{key}")
    c4.download_button("📄 Text report", b["txt"], f"{slug}_report_{stamp}.txt", "text/plain", key=f"txt_{key}")
    st.markdown("##### Individual tables as CSV")
    tcols = st.columns(min(4, max(1, len(rep.tables))))
    for i, (name, df) in enumerate(rep.tables):
        tcols[i % len(tcols)].download_button(f"⬇️ {name}", df.to_csv(index=False).encode("utf-8-sig"),
                                              f"{slug}_{re.sub(r'[^a-z0-9]+', '_', name.lower()).strip('_')}.csv",
                                              "text/csv", key=f"tbl_{key}_{i}")
    with st.expander("👁️ Preview of the text report", expanded=False):
        st.code(b["txt"].decode("utf-8")[:12000], language="text")


def ui_status(ok: bool, text_ok: str, text_bad: str):
    st.markdown(f'<span class="pill {"ok" if ok else "bad"}">{"✔ " + text_ok if ok else "✖ " + text_bad}</span>',
                unsafe_allow_html=True)


import re  # noqa: E402

AC3_FLOW_DOT = """digraph {
rankdir=TB; bgcolor="transparent";
node [shape=box, style="rounded,filled", fillcolor="#E8F1FB", color="#0B4F9C", fontname="Helvetica", fontsize=11];
edge [color="#0B4F9C", fontname="Helvetica", fontsize=10];
start [label="Put every arc (Xi -> Xj) in a queue", fillcolor="#0B4F9C", fontcolor="white"];
pop [label="Pop an arc (Xi -> Xj)"];
rev [label="REVISE: delete each value a of Di\\nthat has no partner b in Dj"];
chg [label="Did Di change?", shape=diamond, fillcolor="#FFF3E3", color="#F28C28"];
emp [label="Is Di empty?", shape=diamond, fillcolor="#FFF3E3", color="#F28C28"];
fail [label="STOP: no solution exists", fillcolor="#FDE7E7", color="#C62828"];
add [label="Re-queue every arc (Xk -> Xi), k != j"];
more [label="Queue empty?", shape=diamond, fillcolor="#FFF3E3", color="#F28C28"];
done [label="STOP: network is arc-consistent\\n(solved if every |Di| = 1, else search)", fillcolor="#E3F4EA", color="#1E8E5A"];
start -> pop -> rev -> chg; chg -> emp [label="yes"]; chg -> more [label="no"];
emp -> fail [label="yes"]; emp -> add [label="no"]; add -> more; more -> pop [label="no"]; more -> done [label="yes"];
}"""

AC3_PSEUDOCODE = """function AC-3(csp):
    queue <- all arcs (Xi, Xj) of csp
    while queue is not empty:
        (Xi, Xj) <- queue.pop()
        if REVISE(csp, Xi, Xj):
            if Di is empty: return FAILURE
            for each Xk in neighbours(Xi) - {Xj}:
                queue.add((Xk, Xi))
    return CONSISTENT

function REVISE(csp, Xi, Xj):
    revised <- false
    for each a in Di:
        if no b in Dj satisfies Cij(a, b):
            delete a from Di ; revised <- true
    return revised

function MAC-SEARCH(domains):          # used when AC-3 alone is not enough
    if every |Di| = 1: return solution
    X <- variable with the smallest |Di|  (MRV)
    for each value v in Dx:
        try X = v ; run AC-3 on arcs (Xk, X)
        if still consistent and MAC-SEARCH succeeds: return it
    return FAILURE  (backtrack)"""


# ------------------------------------------------------------------ HOME
def page_home():
    st.session_state["_last_page"] = "home"
    ui_banner("What is AC-3 and why does it matter?",
              "AC-3 (Arc Consistency Algorithm #3, Alan Mackworth 1977) shrinks the possible values of every variable in a "
              "Constraint Satisfaction Problem by deleting values that can't take part in any solution - before and during search.")
    c1, c2 = st.columns([1.05, 1])
    with c1:
        st.markdown("""
<div class="card"><h4>The three ingredients of every CSP</h4>
<b>Variables</b> X<sub>1</sub> ... X<sub>n</sub> - the decisions to make.<br>
<b>Domains</b> D<sub>i</sub> - the values each variable may take.<br>
<b>Constraints</b> C<sub>ij</sub>(a, b) - rules saying which value pairs are compatible.</div>
<div class="card"><h4>What AC-3 guarantees</h4>
After AC-3, every value left in every domain has at least one compatible partner in each neighbouring variable.
If a domain becomes empty the problem is <b>proven impossible</b>; if every domain has one value the problem is
<b>solved without guessing</b>; otherwise the much smaller problem is handed to backtracking search that keeps
running AC-3 after every guess (MAC).</div>
<div class="card"><h4>How to use this studio</h4>
① Pick a use case on the left &nbsp; ② read the explanation, formulas and data schema &nbsp;
③ use the built-in synthetic data or upload your own CSV files &nbsp; ④ run AC-3 + search &nbsp;
⑤ download a formatted PDF, Word, CSV or text report.</div>
""", unsafe_allow_html=True)
    with c2:
        st.graphviz_chart(AC3_FLOW_DOT)
    st.markdown("#### Formulas")
    ui_formula(r"\text{arc } (X_i \rightarrow X_j) \text{ consistent} \iff \forall a \in D_i\ \exists b \in D_j : C_{ij}(a,b)",
               "<b>X<sub>i</sub>, X<sub>j</sub></b> = two variables linked by a constraint; <b>D<sub>i</sub></b> = values still possible for X<sub>i</sub>; "
               "<b>C<sub>ij</sub>(a,b)</b> = true when X<sub>i</sub>=a and X<sub>j</sub>=b are compatible. ∀ = 'for every', ∃ = 'there exists'.")
    ui_formula(r"\text{REVISE}(X_i,X_j):\ D_i \leftarrow \{a \in D_i \mid \exists b \in D_j : C_{ij}(a,b)\}",
               "Keeps only the <b>supported</b> values of X<sub>i</sub>. Whenever D<sub>i</sub> shrinks, every arc pointing <i>into</i> X<sub>i</sub> is re-checked.")
    ui_formula(r"\text{Cost} = O(e \cdot d^3)",
               "<b>e</b> = number of arcs (two per binary constraint), <b>d</b> = size of the largest domain. Each arc is re-queued at most d times and each REVISE costs d².")
    ui_formula(r"\text{Domain reduction \%} = 100 \times \frac{\sum_i |D_i^{before}| - \sum_i |D_i^{after}|}{\sum_i |D_i^{before}|}",
               "Share of candidate values eliminated by pure reasoning. Reported in every tab together with the search space "
               "<b>∏|D<sub>i</sub>|</b> (shown as log<sub>10</sub>).")
    st.markdown("#### Pseudocode used by this application")
    st.code(AC3_PSEUDOCODE, language="text")

    st.markdown("#### Mini example - watch AC-3 work")
    st.caption("Variables X, Y, Z with domains {1, 2, 3}; constraints X < Y and Y < Z.")
    csp = CSP(["X", "Y", "Z"], {v: [1, 2, 3] for v in "XYZ"}, "demo")
    csp.add_binary("X", "Y", lambda a, b: a < b, "X < Y")
    csp.add_binary("Y", "Z", lambda a, b: a < b, "Y < Z")
    res = ac3(csp)
    a, b = st.columns([1.3, 1])
    a.dataframe(trace_df(res), hide_index=True)
    b.dataframe(res.domain_table(csp)[["Variable", "|D| after unary rules", "|D| after AC-3", "Remaining values"]], hide_index=True)
    st.success("AC-3 alone deduced X = 1, Y = 2, Z = 3 - no search needed.")

    st.markdown("#### The six use cases in this studio")
    cards = [("1️⃣ Sudoku", "Cells = variables, digits = values, 'all different' in rows/columns/boxes."),
             ("2️⃣ Timetabling", "Exams/meetings get (timeslot, room) without people or room clashes."),
             ("3️⃣ Colouring", "Neighbouring regions or interfering radio towers get different colours/channels."),
             ("4️⃣ Configuration", "Pick compatible PC parts or car options; every choice prunes the rest."),
             ("5️⃣ Planning", "Start times for jobs on machines / crews / docks with precedences - AC-3 = CPM windows."),
             ("6️⃣ Vision & NLP", "Waltz line labelling of 3-D drawings and constraint dependency parsing of sentences.")]
    cols = st.columns(3)
    for i, (t, d) in enumerate(cards):
        cols[i % 3].markdown(f'<div class="card"><h4>{t}</h4>{d}</div>', unsafe_allow_html=True)


# ------------------------------------------------------------------ 1 SUDOKU
def _peer_fig():
    z = [[0] * 9 for _ in range(9)]
    for j in _SUDOKU_PEERS[40]:
        r, c = divmod(j, 9)
        z[r][c] = 1 if r == 4 else 2 if c == 4 else 3
    z[4][4] = 4
    fig = go.Figure(go.Heatmap(z=z, colorscale=[[0, "#FFFFFF"], [0.25, "#BBDEFB"], [0.5, "#C8E6C9"], [0.75, "#FFE0B2"], [1, "#0B4F9C"]],
                               showscale=False, xgap=2, ygap=2, hoverinfo="skip"))
    fig.add_annotation(x=4, y=4, text="<b>X</b>", showarrow=False, font=dict(color="white", size=18))
    for k in (-0.5, 2.5, 5.5, 8.5):
        fig.add_shape(type="line", x0=k, x1=k, y0=-0.5, y1=8.5, line=dict(color="#07326A", width=3))
        fig.add_shape(type="line", y0=k, y1=k, x0=-0.5, x1=8.5, line=dict(color="#07326A", width=3))
    fig.update_layout(title="The 20 peers of cell R5C5: row (blue), column (green), box (orange)", height=420,
                      margin=dict(l=10, r=10, t=50, b=10), xaxis=dict(visible=False),
                      yaxis=dict(visible=False, autorange="reversed", scaleanchor="x"))
    return fig


def page_sudoku():
    ui_banner("1 · Sudoku & Logic Puzzles", "AC-3 eliminates impossible digits in every cell. Many easy puzzles are solved by AC-3 alone; "
              "harder ones need a few guesses, each followed by more AC-3.", "#0B4F9C", "#42A5F5")
    t1, t2, t3, t4 = ui_tabs("sudoku", ["📘 1. Understand the use case", "🗂️ 2. Data (synthetic / upload)", "⚙️ 3. Run AC-3 & solve", "📤 4. Export report"])
    with t1:
        c1, c2 = st.columns([1.1, 1])
        with c1:
            st.markdown("""
**The problem.** Fill a 9 × 9 grid so that every row, every column and every 3 × 3 box contains the digits 1-9 exactly once.
Some digits (the *clues*) are given.

**Why AC-3 fits.** Each empty cell can hold 9 digits. A clue in a peer cell rules its digit out; once a cell is down to a
single digit, that digit is ruled out in all of *its* peers, and so on. This chain reaction is exactly AC-3's queue.

**Real-world relatives.** Logic puzzles (KenKen, Kakuro, nonograms), Latin-square experiment designs, tournament
schedules and register allocation in compilers all use the same *all-different* constraints.
""")
            ui_csp_mapping([["Variables", "81 cells R1C1 ... R9C9", "R5C5"],
                            ["Domains", "{1..9}; a clue fixes the domain to one digit", "D(R1C3) = {3}"],
                            ["Constraints", "Xi ≠ Xj for the 20 peers of each cell (810 pairs)", "R5C5 ≠ R5C9"],
                            ["Arcs", "2 per constraint = 1,620 arcs", "R5C5 → R5C9"]])
        with c2:
            st.plotly_chart(_peer_fig(), key="sud_peer")
        st.markdown("##### Formulas")
        ui_formula(r"X_{rc} \in \{1,\dots,9\}", "<b>X<sub>rc</sub></b> is the digit in row r, column c. Its domain starts with all nine digits.")
        ui_formula(r"X_{rc} \neq X_{r'c'}\ \ \text{if } r=r' \ \lor\ c=c' \ \lor\ \text{box}(r,c)=\text{box}(r',c')",
                   "Two <b>peer</b> cells (same row, same column or same box) must hold different digits. "
                   "<b>box(r,c)</b> = 3·⌊(r−1)/3⌋ + ⌊(c−1)/3⌋ + 1 numbers the nine 3 × 3 boxes.")
        ui_formula(r"D_{rc} = \{g\}\ \text{ if cell } (r,c) \text{ has clue } g",
                   "Unary (single-variable) rule applied before AC-3 (node consistency).")
        ui_formula(r"\text{Solved by AC-3 alone} \iff |D_{rc}| = 1\ \ \forall r,c",
                   "If every cell is down to one candidate after AC-3, no search is needed. Otherwise MAC search picks the cell "
                   "with the fewest candidates (MRV), tries a digit, and runs AC-3 again.")
        ui_schema(SUDOKU_SCHEMA, {"sudoku.csv": pd.DataFrame([sudoku_parse_string(SUDOKU_SAMPLES[next(iter(SUDOKU_SAMPLES))])[r * 9:(r + 1) * 9] for r in range(9)])}, "sud")
        st.caption("Note: sudoku.csv has no header row - 9 lines of 9 comma-separated digits. The sample file downloads with a "
                   "header of column numbers 0-8, which the uploader ignores automatically.")
        ui_ac3_generic()

    with t2:
        src = st.radio("Data source", ["Built-in sample puzzles", "Synthetic generator", "Upload my own puzzle"], horizontal=True, key="sud_src")
        grid, label = None, ""
        if src == "Built-in sample puzzles":
            name = st.selectbox("Puzzle", list(SUDOKU_SAMPLES), key="sud_sample")
            grid, label = sudoku_parse_string(SUDOKU_SAMPLES[name]), name
        elif src == "Synthetic generator":
            c1, c2, c3 = st.columns(3)
            seed = c1.number_input("Random seed", 1, 99999, 42, key="sud_seed")
            clues = c2.slider("Target number of clues", 22, 50, 30, key="sud_clues")
            mode = c3.radio("Difficulty rule", ["Unique solution (may need search)", "Must be solvable by AC-3 alone"], key="sud_mode")
            grid, _sol = _cached_generate(int(seed), int(clues), mode.startswith("Must"))
            label = f"Synthetic puzzle (seed {seed}, {sum(1 for v in grid if v)} clues)"
        else:
            up = st.file_uploader("sudoku.csv (9 x 9) or sudoku.txt (81 characters per line)", type=["csv", "txt"], key="sud_up")
            if up is not None:
                try:
                    puzzles = sudoku_parse_upload(up.name, up.getvalue())
                    k = st.selectbox("Puzzle in file", range(len(puzzles)), format_func=lambda i: f"Puzzle {i + 1}") if len(puzzles) > 1 else 0
                    grid, label = puzzles[k], f"Uploaded: {up.name}"
                except Exception as e:  # noqa: BLE001
                    st.error(f"Could not read the file: {e}")
        if grid is not None:
            errs = sudoku_validate(grid)
            if errs:
                for e in errs:
                    st.error(e)
            else:
                ui_set_data("sud", grid, label)
                c1, c2 = st.columns([1.2, 1])
                with c1:
                    st.plotly_chart(sudoku_plot(grid, grid, ["Given" if v else None for v in grid], title=label), key="sud_in")
                with c2:
                    n = sum(1 for v in grid if v)
                    ui_kpis([("Clues", n, BLUE), ("Empty cells", 81 - n, ACCENT)])
                    counts = pd.Series([v for v in grid if v]).value_counts().reindex(range(1, 10), fill_value=0)
                    f = go.Figure(go.Bar(x=[str(i) for i in counts.index], y=counts.values, marker_color=BLUE))
                    f.update_layout(title="How often each digit is given", height=300, margin=dict(l=10, r=10, t=40, b=10), plot_bgcolor="white")
                    st.plotly_chart(f, key="sud_cnt")
                    st.success("✔ Puzzle validated - go to tab 3 to run AC-3.")

    with t3:
        grid = st.session_state.get("sud_data")
        if grid is None:
            st.info("Choose or upload a puzzle in tab 2.")
        else:
            if st.button("▶️ Run AC-3 and solve", type="primary", key="sud_run"):
                with st.spinner("Running AC-3 + MAC search..."):
                    st.session_state["sud_result"] = sudoku_run(grid, st.session_state.get("sud_label", ""))
                    st.session_state.pop("sud_report", None)
            r = st.session_state.get("sud_result")
            if r:
                a, s = r["ac3"], r["search"]
                ui_kpis([("Cells fixed by AC-3", r["source"].count("AC-3"), GREEN), ("Cells needing search", r["source"].count("Search"), ACCENT),
                         ("Values pruned", a.values_removed, BLUE), ("Search nodes", s.nodes, "#8E24AA"),
                         ("Solved by AC-3 alone", "Yes" if a.solved else "No", GREEN if a.solved else ACCENT),
                         ("Verified", "Valid" if r["valid"] else "No solution", GREEN if r["valid"] else RED)])
                c1, c2 = st.columns(2)
                with c1:
                    st.plotly_chart(sudoku_plot(grid, None, None, [a.domains[_cell(i)] for i in range(81)],
                                                "After AC-3: remaining candidates per cell"), key="sud_cand")
                with c2:
                    if r["solution"]:
                        st.plotly_chart(sudoku_plot(grid, r["solution"], r["source"], title="Solution: blue = given, green = AC-3, orange = search"), key="sud_sol")
                    else:
                        st.error("This puzzle has no solution.")
                st.markdown(f"**Interpretation.** AC-3 removed **{a.values_removed}** candidate digits by propagation, "
                            f"cutting the search space from 10^{a.search_space_before:.1f} to 10^{a.search_space_after:.1f}. "
                            + ("That was enough to solve the whole puzzle - no guessing." if a.solved else
                               f"The remaining {r['source'].count('Search')} cells were fixed by MAC search in {s.nodes} trial assignments.")
                            + (" The solution is unique." if r["unique"] else ""))
                st.dataframe(r["cells"], hide_index=True, height=300)
                ui_ac3_details(a, r["csp"], "sud")
            with st.expander("📊 Batch benchmark on synthetic puzzles (tests AC-3 on many puzzles at once)"):
                n = st.slider("Number of puzzles", 3, 20, 8, key="sud_bn")
                if st.button("Run benchmark", key="sud_bench"):
                    rows = []
                    prog = st.progress(0.0)
                    for i in range(n):
                        clues = 24 + (i * 3) % 20
                        g, sol = _cached_generate(1000 + i, clues, False)
                        rr = sudoku_run(g, "bench")
                        rows.append({"Puzzle": i + 1, "Clues": sum(1 for v in g if v), "Solved by AC-3 alone": rr["ac3"].solved,
                                     "Values pruned": rr["ac3"].values_removed, "Search nodes": rr["search"].nodes,
                                     "Time ms": round(rr["search"].elapsed_ms, 1), "Matches generator": rr["solution"] == sol})
                        prog.progress((i + 1) / n)
                    df = pd.DataFrame(rows)
                    st.dataframe(df, hide_index=True)
                    f = go.Figure(go.Scatter(x=df["Clues"], y=df["Search nodes"], mode="markers", marker=dict(size=14, color=[GREEN if v else ACCENT for v in df["Solved by AC-3 alone"]])))
                    f.update_layout(title="Fewer clues -> more search (green = AC-3 alone was enough)", xaxis_title="Clues", yaxis_title="Search nodes", height=320, plot_bgcolor="white")
                    st.plotly_chart(f, key="sud_bench_fig")
    with t4:
        ui_export("sud", sudoku_report, "sudoku")


def _cached_generate(seed, clues, ac3_only):
    if st is not None:
        return _gen_cache(seed, clues, ac3_only)
    return sudoku_generate(seed, clues, ac3_only)


if st is not None:
    @st.cache_data(show_spinner="Generating a synthetic puzzle...")
    def _gen_cache(seed, clues, ac3_only):
        return sudoku_generate(seed, clues, ac3_only)


# ------------------------------------------------------------------ 2 SCHEDULING
def page_sched():
    ui_banner("2 · Scheduling & Timetabling", "Exams, classes, shifts and meetings get a (timeslot, room) pair so that no person is in two places "
              "at once, no room is double-booked and every room is big enough.", "#1B5E20", "#43A047")
    t1, t2, t3, t4 = ui_tabs("sched", ["📘 1. Understand the use case", "🗂️ 2. Data (synthetic / upload)", "⚙️ 3. Run AC-3 & solve", "📤 4. Export report"])
    with t1:
        c1, c2 = st.columns([1.1, 1])
        with c1:
            st.markdown("""
**The problem.** A university must place every exam into a timeslot and a room. A student taking two exams can't sit
them at the same time, an instructor can't invigilate two rooms at once, a room holds one exam per slot, and the room must
seat everybody. Companies face the same puzzle with meetings, and hospitals with shifts.

**Why AC-3 fits.** Each event starts with *slots × rooms* options. Capacity and availability rules remove options at once
(node consistency). When one event is fixed to a slot, AC-3 removes that slot from every event that shares people with it,
which can force further choices in a cascade.
""")
            ui_csp_mapping([["Variables", "One per event (exam / class / meeting)", "CS101"],
                            ["Domains", "All (timeslot, room) pairs that are big enough and allowed", "(MON-AM, HALL-A)"],
                            ["Constraints", "Different slot if they share people/organizer; never same slot AND room", "slot(CS101) ≠ slot(MA101)"],
                            ["Arcs", "2 per pair of events", "CS101 → MA101"]])
        with c2:
            st.graphviz_chart("""graph { layout=neato; overlap=false; bgcolor="transparent";
node [shape=circle, style=filled, fontname="Helvetica", fontsize=10, width=0.8, fixedsize=true, color=white, fontcolor=white];
CS101 [fillcolor="#1F77B4"]; MA101 [fillcolor="#FF7F0E"]; CS210 [fillcolor="#2CA02C"]; BU110 [fillcolor="#1F77B4"]; GE100 [fillcolor="#D62728"];
edge [color="#C62828", penwidth=2, fontname="Helvetica", fontsize=9, fontcolor="#C62828"];
CS101 -- MA101 [label="60 shared students"]; CS101 -- CS210 [label="shared"]; MA101 -- CS210 [label="shared"];
CS101 -- GE100 [label="shared"]; BU110 -- GE100 [label="shared"];
edge [color="#9FB3CC", style=dashed, penwidth=1, label=""]; CS101 -- BU110 [label="only: not same room+slot", fontcolor="#6B7C93"];
}""")
            st.caption("Conflict graph: red edge = must use different timeslots; dashed = may share a slot but not a room. Same colour = same slot.")
        st.markdown("##### Formulas")
        ui_formula(r"X_e = (s_e, r_e) \in S \times R", "<b>X<sub>e</sub></b> = assignment of event e; <b>s<sub>e</sub></b> its timeslot from the set of slots <b>S</b>; <b>r<sub>e</sub></b> its room from the set of rooms <b>R</b>.")
        ui_formula(r"\text{cap}(r_e) \ge n_e \quad\text{and}\quad s_e \notin B_e",
                   "Unary rules: <b>cap(r)</b> = seats in room r, <b>n<sub>e</sub></b> = attendees of e (expected_size or counted from attendance.csv), "
                   "<b>B<sub>e</sub></b> = blocked slots of e.")
        ui_formula(r"s_a \ne s_b \quad \text{if } P_a \cap P_b \ne \emptyset \ \lor\ \text{org}(a)=\text{org}(b)",
                   "<b>P<sub>e</sub></b> = set of people attending e; <b>org(e)</b> = organizer. Sharing anyone forces different timeslots.")
        ui_formula(r"(s_a, r_a) \ne (s_b, r_b) \quad \forall a \ne b", "No room is booked twice in the same timeslot.")
        ui_formula(r"\text{Utilisation}_e = 100 \times n_e / \text{cap}(r_e)", "How full the assigned room is (the solver prefers the smallest room that fits).")
        ui_schema(SCHED_SCHEMA, sched_synthetic(), "sched")
        ui_ac3_generic()
    with t2:
        src = st.radio("Data source", ["Synthetic data", "Upload my own data"], horizontal=True, key="sched_src")
        data, label = None, ""
        if src == "Synthetic data":
            c1, c2, c3 = st.columns(3)
            kind = c1.selectbox("Scenario", ["University exam timetabling", "Corporate meeting scheduling"], key="sched_kind")
            seed = c2.number_input("Random seed", 1, 9999, 11, key="sched_seed")
            people = c3.slider("People (students / employees)", 40, 400, 240 if kind.startswith("Uni") else 60, step=20, key="sched_people")
            data, label = sched_synthetic(kind, int(seed), int(people)), f"{kind} (synthetic, seed {seed})"
        else:
            data = ui_upload("sched", list(SCHED_SCHEMA))
            label = "Uploaded data"
        if data is not None:
            errs = sched_validate(data)
            if errs:
                for e in errs:
                    st.error(e)
            else:
                ui_set_data("sched", data, label)
                ev = data["events.csv"]
                ui_kpis([("Events", len(ev), BLUE), ("People", data["attendance.csv"]["person_id"].nunique(), GREEN),
                         ("Attendance rows", len(data["attendance.csv"]), "#8E24AA"), ("Rooms", len(data["rooms.csv"]), ACCENT),
                         ("Timeslots", len(data["timeslots.csv"]), "#00897B")])
                c1, c2 = st.columns([1.3, 1])
                with c1:
                    st.markdown("**events.csv**")
                    st.dataframe(ev, hide_index=True, height=260)
                with c2:
                    sizes = data["attendance.csv"].groupby("event_id").size().reindex(ev["event_id"]).fillna(0)
                    f = go.Figure(go.Bar(x=list(ev["event_id"]), y=sizes.values, marker_color=BLUE, name="attendees"))
                    for _, rr in data["rooms.csv"].iterrows():
                        f.add_hline(y=rr["capacity"], line=dict(dash="dot", color=ACCENT), annotation_text=f"{rr['room_id']} ({rr['capacity']})")
                    f.update_layout(title="Event size vs room capacities", height=300, margin=dict(l=10, r=10, t=40, b=10), plot_bgcolor="white")
                    st.plotly_chart(f, key="sched_sizes")
                c1, c2, c3 = st.columns([1.2, 0.8, 1])
                c1.markdown("**attendance.csv** (first rows)"); c1.dataframe(data["attendance.csv"].head(200), hide_index=True, height=220)
                c2.markdown("**rooms.csv**"); c2.dataframe(data["rooms.csv"], hide_index=True)
                c3.markdown("**timeslots.csv**"); c3.dataframe(data["timeslots.csv"], hide_index=True)
    with t3:
        data = st.session_state.get("sched_data")
        if data is None:
            st.info("Prepare data in tab 2.")
        else:
            if st.button("▶️ Run AC-3 and build the timetable", type="primary", key="sched_run"):
                with st.spinner("Running AC-3 + MAC search..."):
                    st.session_state["sched_result"] = sched_run(data, st.session_state.get("sched_label", ""))
                    st.session_state.pop("sched_report", None)
            r = st.session_state.get("sched_result")
            if r:
                a, s = r["ac3"], r["search"]
                ok = bool(len(r["schedule"]))
                ui_kpis([("Status", "Solved" if ok else "Infeasible", GREEN if ok else RED), ("Personal clashes", r["clashes"], GREEN if not r["clashes"] else RED),
                         ("Room double-bookings", r["double"], GREEN if not r["double"] else RED), ("Options pruned by AC-3", a.values_removed, BLUE),
                         ("Search nodes", s.nodes, "#8E24AA"), ("Avg seat use %", round(r["schedule"]["Seat utilisation %"].mean(), 1) if ok else "-", ACCENT)])
                if ok:
                    st.markdown("##### Timetable (slot × room)")
                    st.dataframe(r["grid"].style.map(lambda v: "background-color:#E8F1FB;font-weight:700;color:#07326A" if v else ""), height=min(420, 40 + 36 * len(r["grid"])))
                    c1, c2 = st.columns([1, 1])
                    with c1:
                        st.plotly_chart(sched_conflict_fig(r), key="sched_graph")
                    with c2:
                        st.dataframe(r["schedule"], hide_index=True, height=470)
                    st.markdown(f"**Interpretation.** Every event has a slot and a room. Post-solve verification finds **{r['clashes']}** personal clashes, "
                                f"**{r['double']}** double-bookings and **{r['over']}** over-capacity rooms. AC-3 alone removed **{a.values_removed}** "
                                f"impossible options (e.g. slots already taken by a fixed event that shares people).")
                else:
                    st.error("No clash-free timetable exists. Add timeslots or rooms, or unblock slots.")
                ui_ac3_details(a, r["csp"], "sched")
    with t4:
        ui_export("sched", sched_report, "timetable")


# ------------------------------------------------------------------ 3 COLOURING
def page_color():
    ui_banner("3 · Map Colouring & Radio-Frequency Assignment", "Neighbouring regions get different colours; transmitters that could interfere get "
              "channels far enough apart. Minimise the number of colours or channels used.", "#6A1B9A", "#AB47BC")
    t1, t2, t3, t4 = ui_tabs("color", ["📘 1. Understand the use case", "🗂️ 2. Data (synthetic / upload)", "⚙️ 3. Run AC-3 & solve", "📤 4. Export report"])
    with t1:
        c1, c2 = st.columns([1.1, 1])
        with c1:
            st.markdown("""
**The problem.** Colour a map so that bordering regions differ - the classic *four-colour* problem. The same mathematics
assigns **radio frequencies**: towers closer than the interference radius need different channels, and very close
towers need channels at least two apart (a *guard band*).

**Why AC-3 fits.** Pre-coloured regions (or channels already licensed) immediately remove their colour from all
neighbours; a region left with one colour propagates further. In the Australia sample, fixing WA = Red and NT = Green
lets AC-3 colour the whole mainland without any guessing.
""")
            ui_csp_mapping([["Variables", "One per region / transmitter", "NSW, TX07"],
                            ["Domains", "Colours or channels 1..k (a fixed colour is a unary rule)", "{Red, Green, Blue}"],
                            ["Constraints", "|ci − cj| ≥ s for every neighbour pair (s = 1 means just ≠)", "c(SA) ≠ c(NSW)"],
                            ["Arcs", "2 per border / interference link", "SA → NSW"]])
        with c2:
            d = color_synthetic("Australia")
            rr = color_run(d, 3, "Australia", False)
            st.plotly_chart(color_fig(rr), key="col_explain")
            st.caption("Australia as a constraint graph, solved with 3 colours.")
        st.markdown("##### Formulas")
        ui_formula(r"c_i \in \{1, 2, \dots, k\}", "<b>c<sub>i</sub></b> = colour / channel of node i; <b>k</b> = number of colours or channels available.")
        ui_formula(r"|c_i - c_j| \ge s_{ij}\quad \forall (i,j) \in E", "<b>E</b> = set of borders / interference links; <b>s<sub>ij</sub></b> = min_separation "
                   "(1 = simply different; 2 = at least one empty channel between them).")
        ui_formula(r"\text{interference}(i,j) \iff \text{dist}(i,j) < \rho,\qquad s_{ij} = 2 \text{ if dist} < \rho/2", "Synthetic radio data: <b>ρ</b> = interference radius in km.")
        ui_formula(r"\chi(G) = \min\{k : \text{a valid colouring with } k \text{ colours exists}\}", "The <b>chromatic number</b>: 'Find minimum k' tries k = 1, 2, 3 ... "
                   "AC-3 often proves small k impossible instantly.")
        ui_schema(COLOR_SCHEMA, color_synthetic("Radio"), "col")
        ui_ac3_generic()
    with t2:
        src = st.radio("Data source", ["Synthetic data", "Upload my own data"], horizontal=True, key="col_src")
        data, label, radio = None, "", False
        if src == "Synthetic data":
            kind = st.selectbox("Dataset", ["Australia - states & territories", "South America - countries",
                                            "Radio frequency assignment - synthetic transmitters"], key="col_kind")
            radio = kind.startswith("Radio")
            if radio:
                c1, c2, c3 = st.columns(3)
                n = c1.slider("Transmitters", 10, 60, 30, key="col_n")
                rad = c2.slider("Interference radius (km)", 10, 45, 30, key="col_rad")
                seed = c3.number_input("Random seed", 1, 9999, 5, key="col_seed")
                data, label = color_synthetic(kind, int(seed), n, float(rad)), f"{n} synthetic transmitters, radius {rad} km, seed {seed}"
            else:
                data, label = color_synthetic(kind), kind
        else:
            data = ui_upload("col", list(COLOR_SCHEMA))
            radio = st.checkbox("Treat as radio channels (show 'Ch 1, Ch 2...' instead of colour names)", key="col_isradio")
            label = "Uploaded graph"
        if data is not None:
            errs = color_validate(data)
            if errs:
                for e in errs:
                    st.error(e)
            else:
                ui_set_data("col", data, label)
                st.session_state["col_radio"] = radio
                ui_kpis([("Nodes", len(data["nodes.csv"]), BLUE), ("Edges", len(data["edges.csv"]), ACCENT),
                         ("Guard-band links (s≥2)", int((pd.to_numeric(data["edges.csv"].get("min_separation", pd.Series([1])), errors="coerce") > 1).sum()), RED),
                         ("Pre-fixed nodes", int(data["nodes.csv"].get("fixed_color", pd.Series(dtype=str)).astype(str).str.strip().replace("nan", "").ne("").sum()), GREEN)])
                c1, c2 = st.columns(2)
                c1.markdown("**nodes.csv**"); c1.dataframe(data["nodes.csv"], hide_index=True, height=300)
                c2.markdown("**edges.csv**"); c2.dataframe(data["edges.csv"], hide_index=True, height=300)
    with t3:
        data = st.session_state.get("col_data")
        if data is None:
            st.info("Prepare data in tab 2.")
        else:
            radio = st.session_state.get("col_radio", False)
            c1, c2, c3 = st.columns([1, 1, 1])
            default_k = 7 if radio else 4 if len(data["nodes.csv"]) > 8 else 3
            k = c1.slider("Number of colours / channels (k)", 1, 12, default_k, key="col_k")
            run = c2.button("▶️ Run AC-3 and colour", type="primary", key="col_run")
            find = c3.button("🔎 Find minimum k", key="col_min")
            if find:
                with st.spinner("Trying k = 1, 2, 3 ..."):
                    mk, log = color_min_k(data)
                st.session_state["col_minlog"] = (mk, log)
                if mk:
                    st.session_state["col_result"] = color_run(data, mk, st.session_state.get("col_label", ""), radio)
                    st.session_state.pop("col_report", None)
            if run:
                with st.spinner("Solving..."):
                    st.session_state["col_result"] = color_run(data, k, st.session_state.get("col_label", ""), radio)
                    st.session_state.pop("col_report", None)
            if st.session_state.get("col_minlog"):
                mk, log = st.session_state["col_minlog"]
                st.info(f"Minimum number of {'channels' if radio else 'colours'}: **{mk}**  ·  " +
                        "  ·  ".join(f"k={a}: {b} (AC-3 pruned {c})" for a, b, c in log))
            r = st.session_state.get("col_result")
            if r:
                a, s = r["ac3"], r["search"]
                ok = s.solution is not None
                ui_kpis([("k used", r["k"], BLUE), ("Status", "Solved" if ok else "Impossible", GREEN if ok else RED),
                         ("Violations", len(r["violations"]), GREEN if not r["violations"] else RED), ("Values pruned by AC-3", a.values_removed, "#8E24AA"),
                         ("Solved by AC-3 alone", "Yes" if a.solved else "No", GREEN if a.solved else ACCENT), ("Search nodes", s.nodes, ACCENT)])
                c1, c2 = st.columns([1.4, 1])
                with c1:
                    st.plotly_chart(color_fig(r), key="col_fig")
                with c2:
                    st.dataframe(r["table"], hide_index=True, height=360)
                    if len(r["usage"]):
                        f = go.Figure(go.Bar(x=r["usage"].iloc[:, 0].astype(str), y=r["usage"].iloc[:, 1],
                                             marker_color=[COLOR_HEX[(int(c) - 1) % 12] for c in r["usage"].iloc[:, 0]]))
                        f.update_layout(title="Nodes per colour / channel", height=240, margin=dict(l=10, r=10, t=40, b=10), plot_bgcolor="white")
                        st.plotly_chart(f, key="col_use")
                if not ok:
                    st.error(f"No valid assignment with k = {r['k']}. AC-3 {'proved it immediately' if not a.consistent else 'plus search proved it'}. Try a larger k.")
                ui_ac3_details(a, r["csp"], "col")
    with t4:
        ui_export("col", color_report, "colouring")


# ------------------------------------------------------------------ 4 CONFIGURATION
def page_config():
    ui_banner("4 · Product Configuration", "Build a PC or configure a car: every component you choose narrows what the remaining "
              "components can be. AC-3 keeps only the options that still fit.", "#E65100", "#FFA726")
    t1, t2, t3, t4 = ui_tabs("config", ["📘 1. Understand the use case", "🗂️ 2. Data (synthetic / upload)", "⚙️ 3. Configure & solve", "📤 4. Export report"])
    with t1:
        c1, c2 = st.columns([1.1, 1])
        with c1:
            st.markdown("""
**The problem.** A configurator must never offer a combination that can't be built: the CPU must fit the motherboard
socket, the memory type must match, the graphics card must fit inside the case, the power supply must be big enough
and so on. Car configurators have the same kind of rules (this engine only with that transmission, this wheel size only on that trim).

**Why AC-3 fits.** Every rule links two components. When a customer picks an option, AC-3 removes every incompatible
option in *all* other components - even ones linked only indirectly - and greys them out instantly. If some component
runs out of options, the customer learns right away that the combination can't be built, and which rule causes it.
""")
            ui_csp_mapping([["Variables", "One per component slot", "CPU, Motherboard, Case"],
                            ["Domains", "The catalogue options of that component (cheapest first)", "{Ryzen 5 7600, i7-14700K, ...}"],
                            ["Constraints", "One per row of rules.csv: A.attr OP B.attr + offset", "CPU.socket == Motherboard.socket"],
                            ["Customer picks", "Unary rules that fix a component", "CPU = Ryzen 7 7800X3D"]])
        with c2:
            st.graphviz_chart(config_rule_dot(config_synthetic("PC builder")))
            st.caption("PC builder rule graph: boxes = components, orange edges = compatibility rules (R1-R11).")
        st.markdown("##### Formulas")
        ui_formula(r"X_c \in O_c", "<b>X<sub>c</sub></b> = option chosen for component c; <b>O<sub>c</sub></b> = catalogue options for c.")
        ui_formula(r"\text{attr}_a(X_A)\ \ \theta\ \ \text{attr}_b(X_B) + \delta",
                   "One rule: <b>θ</b> ∈ {==, !=, ≤, ≥, <, >} compares attribute a of the option chosen for component A with attribute b of the option for "
                   "component B; <b>δ</b> = offset (e.g. +300 W headroom).")
        ui_formula(r"\text{attr}_a(X_A) \in \text{list}(\text{attr}_b(X_B))", "'in' rules, e.g. the board's form factor must be in the case's list 'ATX;mATX;ITX'.")
        ui_formula(r"\text{Total} = \sum_c \text{price}(X_c) \le \text{Budget}",
                   "Budget is a global (non-binary) rule, so it's applied by an exact branch-and-bound search over the AC-3-reduced domains, "
                   "which also returns the cheapest valid build.")
        ui_formula(r"\text{PSU} \ge \text{GPU} + 300 \ \ \wedge\ \ \text{PSU} \ge \text{CPU} + 350",
                   "AC-3 handles <b>binary</b> constraints, so the total-power rule is written as two binary rules with headroom.")
        ui_schema(CONFIG_SCHEMA, config_synthetic("PC builder"), "cfg")
        ui_ac3_generic()
    with t2:
        src = st.radio("Data source", ["Synthetic data", "Upload my own catalogue"], horizontal=True, key="cfg_src")
        if src == "Synthetic data":
            kind = st.selectbox("Catalogue", ["PC builder", "Car configurator"], key="cfg_kind")
            data, label = config_synthetic(kind), f"{kind} (synthetic catalogue)"
        else:
            data, label = ui_upload("cfg", list(CONFIG_SCHEMA)), "Uploaded catalogue"
        if data is not None:
            errs = config_validate(data)
            if errs:
                for e in errs:
                    st.error(e)
            else:
                ui_set_data("cfg", data, label)
                comps = data["components.csv"]
                ui_kpis([("Components", comps["component"].nunique(), BLUE), ("Catalogue options", len(comps), ACCENT),
                         ("Rules", len(data["rules.csv"]), RED),
                         ("Raw combinations", f"{_math.prod(comps.groupby('component').size()):,}", "#8E24AA")])
                c1, c2 = st.columns([1.4, 1])
                c1.markdown("**components.csv**"); c1.dataframe(comps, hide_index=True, height=330)
                c2.markdown("**rules.csv**"); c2.dataframe(data["rules.csv"], hide_index=True, height=330)
    with t3:
        data = st.session_state.get("cfg_data")
        if data is None:
            st.info("Prepare data in tab 2.")
        else:
            comps = data["components.csv"]
            st.markdown("##### Customer choices (leave 'Any' to let the solver decide)")
            slots = list(dict.fromkeys(comps["component"].astype(str)))
            cols = st.columns(4)
            pins = {}
            for i, c in enumerate(slots):
                opts = comps[comps["component"].astype(str) == c]
                labels = {"": "Any"} | {o: f"{n}  (${p:,.0f})" for o, n, p in zip(opts["option_id"], opts["option_name"], opts["price"])}
                default = "CPU-7800X3D" if c == "CPU" and "CPU-7800X3D" in labels else ""
                pins[c] = cols[i % 4].selectbox(c, list(labels), index=list(labels).index(default), format_func=lambda v, labels=labels: labels[v],
                                                key=f"cfg_pin_{st.session_state.get('cfg_sig', '')}_{c}")
            c1, c2 = st.columns([1, 2])
            use_b = c1.checkbox("Apply a budget", value=True, key="cfg_useb")
            budget = c2.number_input("Budget", 0.0, 1e7, 1800.0 if "PC" in st.session_state.get("cfg_label", "") else 60000.0, step=100.0,
                                     key="cfg_budget") if use_b else None
            if st.button("▶️ Run AC-3 and find the best configuration", type="primary", key="cfg_run"):
                with st.spinner("Propagating rules and searching..."):
                    st.session_state["cfg_result"] = config_run(data, {k: v for k, v in pins.items() if v}, budget, st.session_state.get("cfg_label", ""))
                    st.session_state.pop("cfg_report", None)
            r = st.session_state.get("cfg_result")
            if r:
                a = r["ac3"]
                ok = r["best"] is not None
                ui_kpis([("Status", "Buildable" if ok else "Not buildable", GREEN if ok else RED), ("Valid configurations", f"{r['n_valid']:,}", BLUE),
                         ("Within budget", f"{r['n_within']:,}", "#00897B"), ("Options removed by AC-3", a.values_removed, "#8E24AA"),
                         ("Cheapest valid total", f"${r['best']['total']:,.0f}" if ok else "-", ACCENT)])
                if not a.consistent:
                    st.error(f"Impossible combination: after your picks no **{a.wiped_out}** option satisfies the rules. "
                             "See 'Eliminated options' below for the rules responsible.")
                elif not ok:
                    st.warning("Compatible builds exist, but none within the budget.")
                c1, c2 = st.columns([1.1, 1])
                with c1:
                    st.markdown("##### Recommended configuration")
                    if ok:
                        st.dataframe(r["chosen"], hide_index=True)
                with c2:
                    av = r["availability"]
                    f = go.Figure()
                    f.add_bar(x=av["Component"], y=av["Options in catalogue"], name="Catalogue", marker_color="#B8C7DA")
                    f.add_bar(x=av["Component"], y=av["After customer picks"], name="After picks", marker_color=BLUE)
                    f.add_bar(x=av["Component"], y=av["After AC-3 (still compatible)"], name="After AC-3", marker_color=ACCENT)
                    f.update_layout(title="How picks + rules narrow each component", barmode="group", height=380, plot_bgcolor="white",
                                    margin=dict(l=10, r=10, t=80, b=10), legend=dict(orientation="h", y=1.02, yanchor="bottom", x=0))
                    st.plotly_chart(f, key="cfg_av")
                st.markdown("##### Compatibility after AC-3")
                st.dataframe(av, hide_index=True)
                c1, c2 = st.columns(2)
                c1.markdown("##### Eliminated options - and the rule responsible"); c1.dataframe(r["eliminated"], hide_index=True, height=300)
                c2.markdown("##### Top alternatives (cheapest first)"); c2.dataframe(r["alternatives"], hide_index=True, height=300)
                ui_ac3_details(a, r["csp"], "cfg")
    with t4:
        ui_export("cfg", config_report, "configuration")


# ------------------------------------------------------------------ 5 PLANNING
def page_plan():
    ui_banner("5 · Planning & Resource Allocation", "Schedule manufacturing jobs, construction tasks or truck unloading on shared machines, "
              "crews and docks - respecting order, arrival times and deadlines - and minimise the finish time.", "#00695C", "#26A69A")
    t1, t2, t3, t4 = ui_tabs("plan", ["📘 1. Understand the use case", "🗂️ 2. Data (synthetic / upload)", "⚙️ 3. Run AC-3 & optimise", "📤 4. Export report"])
    with t1:
        c1, c2 = st.columns([1.1, 1])
        with c1:
            st.markdown("""
**The problem.** Each task needs one resource (a machine, a crew, a dock door) for a known duration. Some tasks must
wait for others (weld after cut; drywall after insulation), trucks arrive at given times, and some work has deadlines.
Two tasks can't use the same resource at the same time. We want the whole plan to finish as early as possible
(minimum *makespan*).

**Why AC-3 fits.** Each task's start time is a variable with a window of possible values. AC-3 on the precedence rules
pushes earliest starts forward and latest starts backward - exactly the *Critical Path Method (CPM)* - and the
no-overlap rules prune further. Tasks whose window shrinks to a single time are **critical**; the ones with the least
room are the **bottlenecks**. To minimise the makespan,
the app keeps tightening the horizon until AC-3/search proves no schedule fits.
""")
            ui_csp_mapping([["Variables", "Start time S_t of each task", "S(A2)"],
                            ["Domains", "{release, ..., H − duration} within the horizon H", "{0, ..., 26}"],
                            ["Constraints", "Precedence (a before b) and no-overlap on each resource", "S(A1) + 3 ≤ S(A2)"],
                            ["Objective", "Minimise makespan = latest finish", "12 time units"]])
        with c2:
            d = plan_synthetic("Manufacturing job-shop")
            dot = ['digraph { rankdir=LR; bgcolor="transparent"; node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=10, color=white, fontcolor=white];']
            for _, t in d["tasks.csv"].iterrows():
                col = PALETTE[["Job A", "Job B", "Job C", "Job D"].index(t["job"])]
                dot.append(f'"{t["task_id"]}" [label="{t["task_id"]}\\n{t["resource"]}\\n{t["duration"]}h", fillcolor="{col}"];')
            for _, p in d["precedences.csv"].iterrows():
                dot.append(f'"{p["before"]}" -> "{p["after"]}" [color="#07326A"];')
            st.graphviz_chart("\n".join(dot) + "}")
            st.caption("Job-shop precedence graph: each row is a job, arrows = 'must finish before'. Tasks on the same machine also can't overlap.")
        st.markdown("##### Formulas")
        ui_formula(r"S_t \in \{\rho_t, \dots, H - d_t\}", "<b>S<sub>t</sub></b> = start of task t; <b>ρ<sub>t</sub></b> = release (arrival) time; "
                   "<b>d<sub>t</sub></b> = duration; <b>H</b> = planning horizon.")
        ui_formula(r"S_a + d_a + g_{ab} \le S_b", "Precedence a → b; <b>g<sub>ab</sub></b> = min_gap (curing, transport).")
        ui_formula(r"S_a + d_a \le S_b\ \ \lor\ \ S_b + d_b \le S_a \quad \text{if } \text{res}(a)=\text{res}(b)",
                   "No overlap on a shared resource (a disjunctive binary constraint).")
        ui_formula(r"S_t + d_t \le \delta_t", "<b>δ<sub>t</sub></b> = deadline of task t (unary).")
        ui_formula(r"C_{max} = \max_t (S_t + d_t),\qquad LB = \max\big(\max_r \textstyle\sum_{t: res(t)=r} d_t,\ \text{longest chain}\big)",
                   "<b>C<sub>max</sub></b> = makespan. <b>LB</b> = lower bound: no plan can beat the busiest resource's total load or the longest precedence chain.")
        ui_formula(r"\text{slack}_t = \max D_t - \min D_t", "Width of the start-time window after AC-3; <b>slack 0 = critical task</b>.")
        ui_schema(PLAN_SCHEMA, plan_synthetic("Manufacturing job-shop"), "plan")
        ui_ac3_generic()
    with t2:
        src = st.radio("Data source", ["Synthetic data", "Upload my own data"], horizontal=True, key="plan_src")
        if src == "Synthetic data":
            kind = st.selectbox("Scenario", ["Manufacturing job-shop", "Construction project", "Warehouse dock scheduling"], key="plan_kind")
            data, label = plan_synthetic(kind), f"{kind} (synthetic)"
        else:
            data, label = ui_upload("plan", list(PLAN_SCHEMA), optional=("precedences.csv",)), "Uploaded plan"
            if data is not None and "precedences.csv" not in data:
                data["precedences.csv"] = pd.DataFrame(columns=["before", "after", "min_gap"])
        if data is not None:
            errs = plan_validate(data)
            if errs:
                for e in errs:
                    st.error(e)
            else:
                ui_set_data("plan", data, label)
                t = data["tasks.csv"]
                ui_kpis([("Tasks", len(t), BLUE), ("Resources", t["resource"].nunique(), ACCENT), ("Precedences", len(data["precedences.csv"]), "#8E24AA"),
                         ("Total work", int(pd.to_numeric(t["duration"]).sum()), GREEN), ("Lower bound on makespan", plan_lower_bound(data), RED)])
                c1, c2 = st.columns([1.5, 1])
                c1.markdown("**tasks.csv**"); c1.dataframe(t, hide_index=True, height=320)
                c2.markdown("**precedences.csv**"); c2.dataframe(data["precedences.csv"], hide_index=True, height=320)
    with t3:
        data = st.session_state.get("plan_data")
        if data is None:
            st.info("Prepare data in tab 2.")
        else:
            c1, c2, c3 = st.columns(3)
            H = c1.number_input("Horizon H (0 = automatic)", 0, 1000, 0, key="plan_h")
            opt = c2.checkbox("Minimise makespan", True, key="plan_opt")
            if c3.button("▶️ Run AC-3 and schedule", type="primary", key="plan_run"):
                with st.spinner("Propagating time windows and searching..."):
                    st.session_state["plan_result"] = plan_run(data, int(H) or None, opt, st.session_state.get("plan_label", ""))
                    st.session_state.pop("plan_report", None)
            r = st.session_state.get("plan_result")
            if r:
                a = r["ac3"]
                ok = r["makespan"] is not None
                ui_kpis([("Makespan", r["makespan"] if ok else "Infeasible", GREEN if ok else RED), ("Lower bound", r["lower_bound"], BLUE),
                         ("Optimal", "Proven" if r["optimal"] else "Not proven", GREEN if r["optimal"] else ACCENT),
                         ("Bottleneck tasks (least slack)", int((r["schedule"]["Bottleneck?"] == "Yes").sum()), RED), ("Violations", len(r["violations"]), GREEN if not r["violations"] else RED),
                         ("Values pruned by AC-3", a.values_removed, "#8E24AA")])
                if ok:
                    st.plotly_chart(plan_gantt(r), key="plan_gantt")
                c1, c2 = st.columns([1.1, 1])
                with c1:
                    st.plotly_chart(plan_window_fig(r), key="plan_win")
                with c2:
                    st.markdown("##### Makespan search (tightening the horizon)")
                    st.dataframe(r["log"], hide_index=True)
                    if len(r["util"]):
                        f = go.Figure(go.Bar(x=r["util"]["Resource"], y=r["util"]["Utilisation %"], marker_color=PALETTE[:len(r["util"])]))
                        f.update_layout(title="Resource utilisation %", height=280, margin=dict(l=10, r=10, t=40, b=10), plot_bgcolor="white", yaxis_range=[0, 100])
                        st.plotly_chart(f, key="plan_util")
                st.dataframe(r["schedule"], hide_index=True)
                if ok:
                    st.markdown(f"**Interpretation.** The best plan finishes at **t = {r['makespan']}** "
                                f"({'equal to the lower bound' if r['makespan'] == r['lower_bound'] else 'the next shorter horizon was proven infeasible' if r['optimal'] else 'search limit reached'}). "
                                "Orange bars show the chosen start; the shaded bar is the window AC-3 left open at H = makespan. Bottleneck tasks have the "
                                "least slack - watch them first when something slips (slack 0 would mean strictly critical).")
                ui_ac3_details(a, r["csp"], "plan")
    with t4:
        ui_export("plan", plan_report, "planning")


# ------------------------------------------------------------------ 6 PERCEPTION (VISION + NLP)
def page_percep():
    ui_banner("6 · Vision & Natural Language", "AC-3's earliest successes: Waltz filtering labels the lines of a 3-D line drawing, "
              "and constraint dependency grammars parse sentences by pruning impossible word attachments.", "#AD1457", "#EC407A")
    mode = st.radio("Sub-domain", ["👁️ Vision - Waltz line labelling", "🗣️ Language - constraint dependency parsing"], horizontal=True, key="pc_mode")
    vision = mode.startswith("👁️")
    t1, t2, t3, t4 = ui_tabs("percep", ["📘 1. Understand the use case", "🗂️ 2. Data (synthetic / upload)", "⚙️ 3. Run AC-3 & solve", "📤 4. Export report"])
    if vision:
        _percep_vision(t1, t2, t3, t4)
    else:
        _percep_nlp(t1, t2, t3, t4)


def _percep_vision(t1, t2, t3, t4):
    with t1:
        c1, c2 = st.columns([1.1, 1])
        with c1:
            st.markdown("""
**The problem.** Given a line drawing of objects made of flat faces (trihedral: exactly three faces meet at each corner),
decide what every line *means* in 3-D: a **convex** edge (+), a **concave** edge (−), or an **occluding boundary**
(→, with the object on the right of the arrow). David Waltz (1972) showed that constraint propagation solves this almost
without search - the historical birth of arc consistency.

**Why AC-3 fits.** Each junction (L, fork, arrow, T) has only a few physically possible labelings (the Huffman-Clowes
catalogue). A line connects two junctions and must be read the same way from both ends. AC-3 deletes junction labelings
with no matching partner - "Waltz filtering".
""")
            ui_csp_mapping([["Variables", "One per junction (vertex) of the drawing", "a (a fork)"],
                            ["Domains", "Catalogue labelings for its junction type", "Fork: (+,+,+), (−,−,−), (L,R,−) ..."],
                            ["Constraints", "Shared line must have one interpretation at both ends", "label_a(a−b) = reverse(label_b(b−a))"],
                            ["Background rule", "Outer boundary lines have the object inside (optional)", "b−e is occluding"]])
            st.dataframe(pd.DataFrame([{"Junction": JUNCTION_NAMES[k], "Lines": 2 if k == "L" else 3, "# labelings": len(v),
                                        "Catalogue": "  ".join("(" + " ".join(x) + ")" for x in v)} for k, v in WALTZ_CATALOG.items()]), hide_index=True)
        with c2:
            ex = waltz_run(waltz_synthetic("Cube"), True, "Cube")
            st.plotly_chart(waltz_fig(ex), key="wz_ex")
            st.caption("The cube: green '+' = convex edges, orange arrows = occluding boundary (object on the right).")
        st.markdown("##### Formulas")
        ui_formula(r"X_v \in \mathcal{C}(\text{type}(v))", "<b>X<sub>v</sub></b> = labeling of junction v; <b>𝒞</b> = Huffman-Clowes catalogue; "
                   "type(v) ∈ {L, fork, arrow, T} is detected from the angles between the lines (arrow: one gap > 180°, T: one gap = 180°).")
        ui_formula(r"\ell_v(v\!\to\!w) = \text{rev}\big(\ell_w(w\!\to\!v)\big),\quad \text{rev}(+)=+,\ \text{rev}(-)=-,\ \text{rev}(L)=R,\ \text{rev}(R)=L",
                   "<b>ℓ<sub>v</sub>(v→w)</b> = label of line v−w as seen from v. L/R = object on the left/right when walking along the line from v.")
        ui_formula(r"\ell_v(v\!\to\!w) = L \ \ \forall (v,w) \text{ on the outer boundary (counter-clockwise)}",
                   "Background assumption: the figure floats in front of the background, so its silhouette is occluding with the object inside.")
        ui_schema(WALTZ_SCHEMA, waltz_synthetic("Step"), "wz")
        ui_ac3_generic()
    with t2:
        src = st.radio("Data source", ["Synthetic drawings", "Upload my own drawing"], horizontal=True, key="wz_src")
        if src == "Synthetic drawings":
            kind = st.selectbox("Drawing", ["Cube", "Step block (L-shaped)", "Pyramid (tetrahedron)"], key="wz_kind")
            data, label = waltz_synthetic(kind), kind
        else:
            data, label = ui_upload("wz", list(WALTZ_SCHEMA)), "Uploaded drawing"
        if data is not None:
            errs = waltz_validate(data)
            if errs:
                for e in errs:
                    st.error(e)
            else:
                try:
                    _, pts, adj, jt, _ = waltz_build(data)
                    ui_set_data("wz", data, label)
                    ui_kpis([("Junctions", len(pts), BLUE), ("Lines", len(data["lines.csv"]), ACCENT)] +
                            [(JUNCTION_NAMES[k], sum(1 for v in jt.values() if v[0] == k), c) for k, c in zip("LYWT", [GREEN, "#8E24AA", RED, "#00897B"])])
                    c1, c2 = st.columns([1, 1])
                    with c1:
                        blank = {"pts": pts, "adj": adj, "jt": jt, "search": SolveResult("none", [], None)}
                        st.plotly_chart(waltz_fig(blank), key="wz_in")
                    with c2:
                        st.dataframe(data["vertices.csv"], hide_index=True); st.dataframe(data["lines.csv"], hide_index=True, height=240)
                except ValueError as e:
                    st.error(str(e))
    with t3:
        data = st.session_state.get("wz_data")
        if data is None:
            st.info("Prepare a drawing in tab 2.")
        else:
            bg = st.checkbox("Assume the object stands in front of a background (outer boundary = occluding)", True, key="wz_bg")
            if st.button("▶️ Run Waltz filtering (AC-3)", type="primary", key="wz_run"):
                st.session_state["wz_result"] = waltz_run(data, bg, st.session_state.get("wz_label", ""))
                st.session_state.pop("wz_report", None)
            r = st.session_state.get("wz_result")
            if r:
                a, s = r["ac3"], r["search"]
                n = len(s.solutions)
                ui_kpis([("Interpretations", n, GREEN if n == 1 else ACCENT if n else RED), ("Labelings before", sum(len(d) for d in a.initial_domains.values()), BLUE),
                         ("After AC-3", sum(len(d) for d in a.domains.values()), "#8E24AA"), ("Solved by AC-3 alone", "Yes" if a.solved else "No", GREEN if a.solved else ACCENT),
                         ("Violations", r["violations"], GREEN if not r["violations"] else RED)])
                c1, c2 = st.columns([1.1, 1])
                with c1:
                    idx = st.selectbox("Interpretation", range(n), format_func=lambda i: f"Interpretation {i + 1}", key="wz_idx") if n > 1 else 0
                    st.plotly_chart(waltz_fig(r, idx), key="wz_out")
                with c2:
                    st.dataframe(r["vertex_table"], hide_index=True, height=300)
                    if n:
                        lt = r["line_table"]
                        st.dataframe(lt[lt["Interpretation"] == idx + 1], hide_index=True, height=260)
                st.markdown("**Interpretation.** " + ("With the background assumption AC-3 alone leaves exactly one labeling per junction - the object is understood "
                                                      "without any search, as Waltz observed." if a.solved else
                                                      f"{n} consistent 3-D readings exist (e.g. an object floating vs. glued to a wall)." if n else
                                                      "No labeling is consistent: this is an impossible object under the trihedral assumptions."))
                ui_ac3_details(a, r["csp"], "wz")
    with t4:
        ui_export("wz", perception_report_waltz, "waltz")


def _percep_nlp(t1, t2, t3, t4):
    with t1:
        c1, c2 = st.columns([1.1, 1])
        with c1:
            st.markdown("""
**The problem.** Find the grammatical structure of a sentence: which word is the subject, which is the object, and
what each word modifies. In *"the old man saw a dog with a telescope"* the phrase *with a telescope* can attach to *saw*
(the man used a telescope) or to *dog* (the dog had one) - a genuine ambiguity.

**Why AC-3 fits.** In Maruyama's **Constraint Dependency Grammar** (1990) every word is a variable whose value is a pair
(role, head). The part of speech limits the pairs (a determiner points to a noun on its right), and binary rules forbid
crossing arcs, two subjects for one verb, and so on. AC-3 prunes most attachments before any search; the survivors
are exactly the grammatical readings.
""")
            ui_csp_mapping([["Variables", "One per word", "'telescope' (word 9)"],
                            ["Domains", "(role, head position) pairs allowed by its part of speech", "(pobj, 7) or (obj, 4)"],
                            ["Constraints", "No crossing arcs; unique subject/object/determiner; no cycles; det before adj", "arcs 3→7 and 4→6 can't cross"],
                            ["Global check", "Each preposition has exactly one object; exactly one ROOT", "'with' has pobj 'telescope'"]])
        with c2:
            p = nlp_parse(NLP_SENTENCES[0], nlp_default_lexicon())
            st.plotly_chart(nlp_fig(p, 1), key="nlp_ex")
            st.caption("One of the two valid parses: 'with a telescope' attached to 'dog'.")
        st.markdown("##### Formulas")
        ui_formula(r"X_i = (\text{role}_i, h_i),\quad h_i \in \{0, 1, \dots, n\} \setminus \{i\}",
                   "<b>X<sub>i</sub></b> = analysis of word i; <b>role<sub>i</sub></b> ∈ {root, nsubj, obj, det, amod, prep, pobj, advmod}; <b>h<sub>i</sub></b> = position of its head word (0 = ROOT); <b>n</b> = sentence length.")
        ui_formula(r"\neg\big(\min(i,h_i) < \min(j,h_j) < \max(i,h_i) < \max(j,h_j)\big)", "<b>Projectivity:</b> dependency arcs may not cross.")
        ui_formula(r"h_i = h_j \wedge \text{role}_i = \text{role}_j \in \{\text{det, nsubj, obj, pobj, root}\} \Rightarrow \bot", "At most one determiner / subject / object per head and one ROOT (⊥ = forbidden).")
        ui_formula(r"\neg(h_i = j \wedge h_j = i)", "No word may depend on its own dependent (no 2-cycles).")
        ui_formula(r"\text{POS}(i)=\text{DET} \Rightarrow \text{role}_i=\text{det} \wedge \text{POS}(h_i)=\text{NOUN} \wedge h_i > i",
                   "Example unary rule from the lexicon. Similar rules exist for adjectives, nouns, verbs, prepositions and adverbs.")
        ui_schema(NLP_SCHEMA, {"lexicon.csv": nlp_default_lexicon()}, "nlp")
        st.download_button("⬇️ Download sample sentences.txt", "\n".join(NLP_SENTENCES).encode(), "sentences.txt", "text/plain", key="nlp_tpl_s")
        ui_ac3_generic()
    with t2:
        src = st.radio("Data source", ["Synthetic sentences", "Upload my own sentences"], horizontal=True, key="nlp_src")
        lex = nlp_default_lexicon()
        if src == "Synthetic sentences":
            txt = st.text_area("Sentences (one per line) - edit freely", "\n".join(NLP_SENTENCES), height=260, key="nlp_txt")
            label = "Synthetic demo sentences"
        else:
            c1, c2 = st.columns(2)
            up = c1.file_uploader("sentences.txt", type=["txt"], key="nlp_up")
            upl = c2.file_uploader("lexicon.csv (optional - extends the built-in lexicon)", type=["csv"], key="nlp_lex")
            txt = up.getvalue().decode("utf-8-sig", errors="replace") if up else ""
            if upl is not None:
                try:
                    extra = pd.read_csv(upl)
                    if {"word", "pos"} <= set(extra.columns):
                        lex = pd.concat([extra[["word", "pos"]], lex]).drop_duplicates("word")
                        st.success(f"Lexicon extended with {len(extra)} entries.")
                    else:
                        st.error("lexicon.csv needs columns: word, pos")
                except Exception as e:  # noqa: BLE001
                    st.error(str(e))
            label = f"Uploaded: {up.name}" if up else ""
        unk = st.checkbox("Treat unknown words as nouns", True, key="nlp_unk")
        sents = [s.strip() for s in txt.splitlines() if s.strip()]
        if sents:
            ui_set_data("nlp", {"sentences": sents, "lexicon": lex, "unk": unk}, label)
            known = set(lex["word"].str.lower())
            unknown = sorted({w for s in sents for w in nlp_tokenize(s) if w not in known})
            ui_kpis([("Sentences", len(sents), BLUE), ("Words", sum(len(nlp_tokenize(s)) for s in sents), ACCENT),
                     ("Lexicon entries", len(lex), GREEN), ("Unknown words", len(unknown), RED if unknown else GREEN)])
            if unknown:
                st.warning(f"Unknown words: {', '.join(unknown)}" + (" - treated as nouns." if unk else " - add them to lexicon.csv."))
            st.dataframe(lex.groupby("pos")["word"].apply(lambda w: ", ".join(sorted(w))).reset_index().rename(columns={"word": "Words"}), hide_index=True)
    with t3:
        d = st.session_state.get("nlp_data")
        if d is None:
            st.info("Enter or upload sentences in tab 2.")
        else:
            if st.button("▶️ Parse with AC-3", type="primary", key="nlp_run"):
                try:
                    st.session_state["nlp_result"] = nlp_run(d["sentences"], d["lexicon"], st.session_state.get("nlp_label", ""), d["unk"])
                    st.session_state.pop("nlp_report", None)
                except ValueError as e:
                    st.error(str(e))
            r = st.session_state.get("nlp_result")
            if r:
                sm = r["summary"]
                ui_kpis([("Sentences", len(sm), BLUE), ("Unique parse", int((sm["Valid parses"] == 1).sum()), GREEN),
                         ("Ambiguous", int((sm["Valid parses"] > 1).sum()), ACCENT), ("Ungrammatical", int((sm["Valid parses"] == 0).sum()), RED),
                         ("Values pruned by AC-3", int((sm["Values before AC-3"] - sm["Values after AC-3"]).sum()), "#8E24AA")])
                st.dataframe(sm, hide_index=True)
                c1, c2 = st.columns([1, 3])
                si = c1.selectbox("Sentence", range(len(r["parsed"])), format_func=lambda i: f"{i + 1}. {r['parsed'][i]['sentence'][:40]}", key="nlp_si")
                p = r["parsed"][si]
                k = c1.selectbox("Parse", range(max(1, len(p["parses"]))), format_func=lambda i: f"Parse {i + 1}", key=f"nlp_k_{si}")
                with c2:
                    st.plotly_chart(nlp_fig(p, k if p["parses"] else 0), key="nlp_fig")
                det = r["detail"]
                st.dataframe(det[(det["Sentence #"] == si + 1) & (det["Parse #"] == k + 1)], hide_index=True)
                ui_ac3_details(p["ac3"], p["csp"], "nlp", label_map=lambda i, w=p["words"]: f"{i}:{w[i - 1]}")
    with t4:
        ui_export("nlp", perception_report_nlp, "parsing")


# ------------------------------------------------------------------ MAIN
PAGES = {"home": page_home, "sudoku": page_sudoku, "sched": page_sched, "color": page_color,
         "config": page_config, "plan": page_plan, "percep": page_percep}


def _apply_theme():
    """Blue KNet theme without needing a separate .streamlit/config.toml file."""
    try:
        from streamlit import config as _cfg
        for k, v in {"theme.primaryColor": BLUE, "theme.base": "light", "theme.backgroundColor": "#FFFFFF",
                     "theme.secondaryBackgroundColor": "#F2F6FC", "theme.textColor": "#1B2A3D",
                     "client.toolbarMode": "viewer"}.items():
            if _cfg.get_option(k) != v:
                _cfg.set_option(k, v)
    except Exception:  # noqa: BLE001 - cosmetic only
        pass


def main():
    _apply_theme()
    st.set_page_config(page_title=APP_NAME, page_icon="🧩", layout="wide", initial_sidebar_state="expanded")
    st.markdown(CSS, unsafe_allow_html=True)
    with st.sidebar:
        st.markdown('<div class="side-brand">🧩 KNet · AC-3 Studio</div><div class="side-sub">Kalsnet (KNet) Consulting</div>', unsafe_allow_html=True)
        choice = st.radio("Use cases", [lbl for _, lbl in USE_CASES], key="nav", label_visibility="collapsed")
        st.markdown("---")
        st.markdown("**Workflow in every tab**  \n① Understand → ② Data → ③ Run → ④ Export")
        st.markdown("**Exports:** PDF · Word · CSV · TXT")
        st.caption(f"Version {VERSION}")
    ui_header()
    key = dict((lbl, k) for k, lbl in USE_CASES)[choice]
    PAGES[key]()
    st.markdown(f'<div class="footer">{APP_NAME} · {DEVELOPER_LINE} · v{VERSION}</div>', unsafe_allow_html=True)


if __name__ == "__main__" and st is not None:
    main()

