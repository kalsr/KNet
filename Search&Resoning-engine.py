


# Search & Resoning engine app

# Message. For JAWS, turn virtual PC Cursor on if needed.

# Kalsnet Graph Search and Reasoning Engine
# Developed By Randy Singh from Kalsnet (KNet) Consulting Group

# This Streamlit application demonstrates Uninformed Search techniques
# (Depth First Search, Breadth First Search, and Uniform Cost Search)
# applied to a security and enterprise Knowledge Graph.

# The application allows the user to:
# 1. Read an explanation of each search technique and its schema
# 2. Upload real data or generate synthetic data for each case
# 3. Display the entire synthetic dataset on request
# 4. Run the search algorithm and view a graph visualization of the result
# 5. Export the results as PDF, Word, CSV, and plain text files


import io
import random
import string
import heapq
import datetime

import pandas as pd
import networkx as nx
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

import streamlit as st

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.units import inch


# ---------------------------------------------------------------------------
# Static reference data
# ---------------------------------------------------------------------------

NODE_TYPES = [
    "Endpoint", "Employee", "Laptop", "VPN", "API", "Gateway",
    "Application", "Server", "Database", "Sensitive Data",
    "Vendor", "Software", "Library"
]

OWNERS = [
    "IT Security", "Network Team", "Application Team",
    "Database Team", "Cloud Team", "Vendor Management"
]

RELATIONSHIP_TYPES = [
    "connects_to", "accesses", "uses", "depends_on",
    "communicates_with", "hosts"
]

NODE_SCHEMA = pd.DataFrame([
    {"Field": "ID", "Type": "String", "Description": "Unique identifier for the node"},
    {"Field": "Type", "Type": "String", "Description": "Category of the node, for example Endpoint, Server, Database, Application"},
    {"Field": "Risk", "Type": "Integer 0 to 10", "Description": "Risk score assigned to the node"},
    {"Field": "Criticality", "Type": "Integer 0 to 10", "Description": "Business criticality of the node"},
    {"Field": "Owner", "Type": "String", "Description": "Team or individual responsible for the node"},
    {"Field": "Security_Score", "Type": "Integer 0 to 100", "Description": "Overall security posture score of the node"},
])

EDGE_SCHEMA = pd.DataFrame([
    {"Field": "Source", "Type": "String", "Description": "ID of the source node"},
    {"Field": "Target", "Type": "String", "Description": "ID of the target node"},
    {"Field": "Relationship", "Type": "String", "Description": "Type of relationship, for example connects_to, accesses, uses"},
    {"Field": "Cost", "Type": "Float", "Description": "Cost of traversing this relationship, used by Uniform Cost Search"},
    {"Field": "Risk", "Type": "Integer 0 to 10", "Description": "Risk score of this relationship"},
    {"Field": "Probability", "Type": "Float 0 to 1", "Description": "Estimated probability of exploit along this relationship"},
    {"Field": "Time", "Type": "Float", "Description": "Estimated time in hours required to traverse this relationship"},
])

CASE_INFO = {
    "DFS": {
        "full_name": "Depth First Search",
        "purpose": "Attack Chain Exploration",
        "explanation": (
            "Depth First Search follows one path as deeply as possible before backtracking. "
            "Inside Kalsnet Hub this is used to trace a full attack chain from a compromised "
            "endpoint all the way to a sensitive target, for example moving from a compromised "
            "PC to a user account, then to an API, then to an application server, then to a "
            "database, and finally to sensitive data. "
            "The typical space complexity of a stack based Depth First Search is approximately "
            "O(b times m), where b is the branching factor and m is the maximum depth, rather "
            "than O(b to the power m), because only the current path is kept in memory rather "
            "than the entire tree."
        ),
        "use_cases": [
            "Tracing a cyber attack path from a compromised device to a sensitive asset",
            "Following a chain of dependencies as deeply as possible",
            "Root cause analysis that requires exploring one branch fully before trying another"
        ],
    },
    "BFS": {
        "full_name": "Breadth First Search",
        "purpose": "Blast Radius and Impact Analysis",
        "explanation": (
            "Breadth First Search explores a graph level by level, expanding all neighbors of "
            "the current node before moving further out. Inside Kalsnet Hub this is used to "
            "answer questions such as which systems are within three network relationships of "
            "a compromised server. "
            "This is extremely useful for cyber blast radius analysis, dependency analysis, "
            "incident response, network topology mapping, API dependency analysis, application "
            "dependency analysis, and Zero Trust analysis. "
            "The space requirement of Breadth First Search can reach approximately O(b to the "
            "power s), where b is the branching factor and s is the solution depth, because it "
            "must retain an entire frontier level in memory at once."
        ),
        "use_cases": [
            "Finding the shortest number of relationships between two systems",
            "Blast radius analysis after a compromise",
            "Mapping all assets connected to an incident within a limited number of hops"
        ],
    },
    "UCS": {
        "full_name": "Uniform Cost Search",
        "purpose": "Least Cost Remediation Path",
        "explanation": (
            "Uniform Cost Search finds the path with the lowest cumulative cost between a start "
            "node and a goal node, where cost can represent risk, financial cost, time, or "
            "operational impact, or a combination of these factors. Inside Kalsnet Hub this "
            "becomes a basic form of decision optimization, for example comparing patch server, "
            "replace API, and isolate network as three different remediation paths and selecting "
            "the one with the lowest overall combined cost. "
            "Uniform Cost Search always expands the node with the lowest cumulative cost so far, "
            "which guarantees an optimal solution when all edge costs are non negative."
        ),
        "use_cases": [
            "Selecting the lowest risk and lowest cost remediation path",
            "SOAR playbook selection based on operational cost",
            "Recovery sequencing based on cost, time, risk, and business priority"
        ],
    },
}


# ---------------------------------------------------------------------------
# Synthetic data generation
# ---------------------------------------------------------------------------

def generate_synthetic_nodes(num_nodes, seed=42):
    rng = random.Random(seed)
    rows = []
    for i in range(1, num_nodes + 1):
        rows.append({
            "ID": "N" + str(i),
            "Type": rng.choice(NODE_TYPES),
            "Risk": rng.randint(0, 10),
            "Criticality": rng.randint(0, 10),
            "Owner": rng.choice(OWNERS),
            "Security_Score": rng.randint(0, 100),
        })
    return pd.DataFrame(rows)


def generate_synthetic_edges(nodes_df, extra_edge_ratio=0.4, seed=42):
    rng = random.Random(seed)
    node_ids = list(nodes_df["ID"])
    rows = []

    # build a spanning tree first so the graph stays connected
    for i in range(1, len(node_ids)):
        source = rng.choice(node_ids[:i])
        target = node_ids[i]
        rows.append(_make_edge_row(rng, source, target))

    # add extra random edges for additional complexity
    num_extra = int(len(node_ids) * extra_edge_ratio)
    for _ in range(num_extra):
        source, target = rng.sample(node_ids, 2)
        rows.append(_make_edge_row(rng, source, target))

    return pd.DataFrame(rows)


def _make_edge_row(rng, source, target):
    return {
        "Source": source,
        "Target": target,
        "Relationship": rng.choice(RELATIONSHIP_TYPES),
        "Cost": round(rng.uniform(1, 50), 2),
        "Risk": rng.randint(0, 10),
        "Probability": round(rng.uniform(0, 1), 2),
        "Time": round(rng.uniform(1, 20), 2),
    }


# ---------------------------------------------------------------------------
# Graph construction
# ---------------------------------------------------------------------------

def build_graph(nodes_df, edges_df):
    graph = nx.DiGraph()
    for _, row in nodes_df.iterrows():
        graph.add_node(row["ID"], **row.drop(labels=["ID"]).to_dict())
    for _, row in edges_df.iterrows():
        graph.add_edge(row["Source"], row["Target"], **row.drop(labels=["Source", "Target"]).to_dict())
    return graph


# ---------------------------------------------------------------------------
# Search algorithms
# ---------------------------------------------------------------------------

def dfs_search(graph, start, goal):
    visited_order = []
    visited = set()
    stack = [(start, [start])]

    while stack:
        node, path = stack.pop()
        if node in visited:
            continue
        visited.add(node)
        visited_order.append(node)

        if node == goal:
            return path, visited_order

        neighbors = sorted(graph.successors(node), reverse=True)
        for neighbor in neighbors:
            if neighbor not in visited:
                stack.append((neighbor, path + [neighbor]))

    return None, visited_order


def bfs_search(graph, start, goal):
    visited_order = []
    visited = {start}
    queue = [(start, [start], 0)]
    levels = {start: 0}
    head = 0

    while head < len(queue):
        node, path, depth = queue[head]
        head += 1
        visited_order.append(node)

        if node == goal:
            return path, visited_order, levels

        for neighbor in sorted(graph.successors(node)):
            if neighbor not in visited:
                visited.add(neighbor)
                levels[neighbor] = depth + 1
                queue.append((neighbor, path + [neighbor], depth + 1))

    return None, visited_order, levels


def ucs_search(graph, start, goal, weight_field="Cost"):
    visited_order = []
    visited = set()
    counter = 0
    frontier = [(0.0, counter, start, [start])]

    while frontier:
        cost, _, node, path = heapq.heappop(frontier)
        if node in visited:
            continue
        visited.add(node)
        visited_order.append(node)

        if node == goal:
            return path, cost, visited_order

        for neighbor in sorted(graph.successors(node)):
            if neighbor not in visited:
                edge_cost = graph[node][neighbor].get(weight_field, 1.0)
                counter += 1
                heapq.heappush(frontier, (cost + edge_cost, counter, neighbor, path + [neighbor]))

    return None, None, visited_order


# ---------------------------------------------------------------------------
# Visualization
# ---------------------------------------------------------------------------

def plot_graph(graph, path=None, title="Knowledge Graph"):
    fig, ax = plt.subplots(figsize=(9, 6))
    layout = nx.spring_layout(graph, seed=7)

    nx.draw_networkx_nodes(graph, layout, node_color="#A9C6E8", node_size=650, ax=ax)
    nx.draw_networkx_labels(graph, layout, font_size=8, ax=ax)
    nx.draw_networkx_edges(graph, layout, edge_color="#B0B0B0", arrows=True, ax=ax)

    if path and len(path) > 1:
        path_edges = list(zip(path[:-1], path[1:]))
        nx.draw_networkx_nodes(graph, layout, nodelist=path, node_color="#FF9900", node_size=750, ax=ax)
        nx.draw_networkx_edges(graph, layout, edgelist=path_edges, edge_color="#CC0000", width=2.5, arrows=True, ax=ax)
        nx.draw_networkx_nodes(graph, layout, nodelist=[path[0]], node_color="#2E8B57", node_size=800, ax=ax)
        nx.draw_networkx_nodes(graph, layout, nodelist=[path[-1]], node_color="#B22222", node_size=800, ax=ax)

    ax.set_title(title, fontsize=13, fontweight="bold", color="#0000CC")
    ax.axis("off")
    fig.tight_layout()
    return fig


# ---------------------------------------------------------------------------
# Export helpers
# ---------------------------------------------------------------------------

def build_result_summary(case_key, start, goal, path, cost_or_extra, nodes_df, edges_df, visited_order):
    info = CASE_INFO[case_key]
    generated_at = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if path:
        path_text = " then ".join(path)
        status = "Path found"
    else:
        path_text = "No path was found between the selected start and goal nodes"
        status = "No path found"

    summary = {
        "case_key": case_key,
        "case_name": info["full_name"],
        "purpose": info["purpose"],
        "explanation": info["explanation"],
        "generated_at": generated_at,
        "start": start,
        "goal": goal,
        "status": status,
        "path": path,
        "path_text": path_text,
        "extra": cost_or_extra,
        "node_count": len(nodes_df),
        "edge_count": len(edges_df),
        "visited_order": visited_order,
    }
    return summary


def export_csv_bytes(path, nodes_df):
    if path:
        result_df = nodes_df[nodes_df["ID"].isin(path)].copy()
        result_df["Order_In_Path"] = result_df["ID"].apply(lambda x: path.index(x) + 1)
        result_df = result_df.sort_values("Order_In_Path")
    else:
        result_df = pd.DataFrame(columns=list(nodes_df.columns) + ["Order_In_Path"])
    buffer = io.StringIO()
    result_df.to_csv(buffer, index=False)
    return buffer.getvalue().encode("utf-8")


def export_text_bytes(summary):
    lines = []
    lines.append("KALSNET GRAPH SEARCH AND REASONING ENGINE")
    lines.append("Developed By Randy Singh from Kalsnet (KNet) Consulting Group")
    lines.append("")
    lines.append("Case: " + summary["case_name"] + " -- " + summary["purpose"])
    lines.append("Generated: " + summary["generated_at"])
    lines.append("")
    lines.append("Explanation")
    lines.append(summary["explanation"])
    lines.append("")
    lines.append("Search Parameters")
    lines.append("Start node: " + str(summary["start"]))
    lines.append("Goal node: " + str(summary["goal"]))
    lines.append("Total nodes in graph: " + str(summary["node_count"]))
    lines.append("Total edges in graph: " + str(summary["edge_count"]))
    lines.append("")
    lines.append("Result")
    lines.append("Status: " + summary["status"])
    lines.append("Path: " + summary["path_text"])
    if summary["extra"] is not None:
        lines.append("Additional detail: " + str(summary["extra"]))
    lines.append("")
    lines.append("Visited Order")
    lines.append(" then ".join(summary["visited_order"]))
    return "\n".join(lines).encode("utf-8")


def export_word_bytes(summary):
    document = Document()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Kalsnet Graph Search and Reasoning Engine")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0xCC)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run("Developed By Randy Singh from Kalsnet (KNet) Consulting Group")
    run2.bold = True
    run2.font.size = Pt(13)
    run2.font.color.rgb = RGBColor(0x00, 0x00, 0xCC)

    document.add_heading(summary["case_name"] + " -- " + summary["purpose"], level=1)
    document.add_paragraph("Generated: " + summary["generated_at"])

    document.add_heading("Explanation", level=2)
    document.add_paragraph(summary["explanation"])

    document.add_heading("Search Parameters", level=2)
    param_table = document.add_table(rows=4, cols=2)
    param_table.style = "Light Grid Accent 1"
    param_rows = [
        ("Start node", str(summary["start"])),
        ("Goal node", str(summary["goal"])),
        ("Total nodes in graph", str(summary["node_count"])),
        ("Total edges in graph", str(summary["edge_count"])),
    ]
    for i, (label, value) in enumerate(param_rows):
        param_table.cell(i, 0).text = label
        param_table.cell(i, 1).text = value

    document.add_heading("Result", level=2)
    document.add_paragraph("Status: " + summary["status"])
    document.add_paragraph("Path: " + summary["path_text"])
    if summary["extra"] is not None:
        document.add_paragraph("Additional detail: " + str(summary["extra"]))

    document.add_heading("Visited Order", level=2)
    document.add_paragraph(" then ".join(summary["visited_order"]))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_pdf_bytes(summary):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleBlue", parent=styles["Title"], textColor=colors.HexColor("#0000CC"), fontSize=20
    )
    subtitle_style = ParagraphStyle(
        "SubtitleBlue", parent=styles["Normal"], textColor=colors.HexColor("#0000CC"),
        fontSize=12, alignment=1, spaceAfter=14
    )
    heading_style = styles["Heading2"]
    body_style = styles["BodyText"]

    elements = []
    elements.append(Paragraph("Kalsnet Graph Search and Reasoning Engine", title_style))
    elements.append(Paragraph("Developed By Randy Singh from Kalsnet (KNet) Consulting Group", subtitle_style))
    elements.append(Spacer(1, 10))

    elements.append(Paragraph(summary["case_name"] + " -- " + summary["purpose"], heading_style))
    elements.append(Paragraph("Generated: " + summary["generated_at"], body_style))
    elements.append(Spacer(1, 10))

    elements.append(Paragraph("Explanation", heading_style))
    elements.append(Paragraph(summary["explanation"], body_style))
    elements.append(Spacer(1, 10))

    elements.append(Paragraph("Search Parameters", heading_style))
    param_data = [
        ["Field", "Value"],
        ["Start node", str(summary["start"])],
        ["Goal node", str(summary["goal"])],
        ["Total nodes in graph", str(summary["node_count"])],
        ["Total edges in graph", str(summary["edge_count"])],
    ]
    param_table = Table(param_data, colWidths=[2.5 * inch, 3.5 * inch])
    param_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0000CC")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF2FA")]),
    ]))
    elements.append(param_table)
    elements.append(Spacer(1, 10))

    elements.append(Paragraph("Result", heading_style))
    elements.append(Paragraph("Status: " + summary["status"], body_style))
    elements.append(Paragraph("Path: " + summary["path_text"], body_style))
    if summary["extra"] is not None:
        elements.append(Paragraph("Additional detail: " + str(summary["extra"]), body_style))
    elements.append(Spacer(1, 10))

    elements.append(Paragraph("Visited Order", heading_style))
    elements.append(Paragraph(" then ".join(summary["visited_order"]), body_style))

    doc.build(elements)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Streamlit UI helpers
# ---------------------------------------------------------------------------

def apply_page_style():
    st.set_page_config(page_title="Kalsnet Graph Search and Reasoning Engine", layout="wide")
    st.markdown(
        """
        <style>
        .kalsnet-title {
            color: #0000CC;
            font-size: 44px;
            font-weight: 800;
            text-align: center;
            margin-bottom: 0px;
        }
        .kalsnet-subtitle {
            color: #0000CC;
            font-size: 22px;
            font-weight: 700;
            text-align: center;
            margin-top: 4px;
            margin-bottom: 20px;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )
    st.markdown('<div class="kalsnet-title">Kalsnet Graph Search and Reasoning Engine</div>', unsafe_allow_html=True)
    st.markdown('<div class="kalsnet-subtitle">Developed By Randy Singh from Kalsnet (KNet) Consulting Group</div>', unsafe_allow_html=True)


def render_schema_section():
    with st.expander("View Node and Edge Schema"):
        st.subheader("Node Schema")
        st.dataframe(NODE_SCHEMA, use_container_width=True, hide_index=True)
        st.subheader("Edge Schema")
        st.dataframe(EDGE_SCHEMA, use_container_width=True, hide_index=True)


def render_data_section(case_key):
    st.subheader("Step 1: Provide Data")
    source = st.radio(
        "Choose a data source",
        ["Generate Synthetic Data", "Upload Real Data"],
        key=case_key + "_source",
        horizontal=True,
    )

    nodes_df = None
    edges_df = None

    if source == "Generate Synthetic Data":
        col1, col2, col3 = st.columns(3)
        with col1:
            num_nodes = st.slider("Number of nodes", min_value=5, max_value=60, value=15, key=case_key + "_numnodes")
        with col2:
            seed = st.number_input("Random seed", min_value=1, max_value=9999, value=42, key=case_key + "_seed")
        with col3:
            show_all = st.checkbox("Display Entire Synthetic Data", key=case_key + "_showall")

        if st.button("Generate Synthetic Data Now", key=case_key + "_gen_button"):
            nodes_df = generate_synthetic_nodes(num_nodes, seed=seed)
            edges_df = generate_synthetic_edges(nodes_df, seed=seed)
            st.session_state[case_key + "_nodes"] = nodes_df
            st.session_state[case_key + "_edges"] = edges_df

        if case_key + "_nodes" in st.session_state:
            nodes_df = st.session_state[case_key + "_nodes"]
            edges_df = st.session_state[case_key + "_edges"]
            if show_all:
                st.write("Complete synthetic node data, " + str(len(nodes_df)) + " rows")
                st.dataframe(nodes_df, use_container_width=True, hide_index=True)
                st.write("Complete synthetic edge data, " + str(len(edges_df)) + " rows")
                st.dataframe(edges_df, use_container_width=True, hide_index=True)
            else:
                st.write("Preview of synthetic node data, first 5 rows")
                st.dataframe(nodes_df.head(5), use_container_width=True, hide_index=True)
                st.write("Preview of synthetic edge data, first 5 rows")
                st.dataframe(edges_df.head(5), use_container_width=True, hide_index=True)

    else:
        node_file = st.file_uploader("Upload Node CSV File", type=["csv"], key=case_key + "_nodefile")
        edge_file = st.file_uploader("Upload Edge CSV File", type=["csv"], key=case_key + "_edgefile")

        if node_file is not None and edge_file is not None:
            nodes_df = pd.read_csv(node_file)
            edges_df = pd.read_csv(edge_file)
            st.session_state[case_key + "_nodes"] = nodes_df
            st.session_state[case_key + "_edges"] = edges_df

        if case_key + "_nodes" in st.session_state:
            nodes_df = st.session_state[case_key + "_nodes"]
            edges_df = st.session_state[case_key + "_edges"]
            st.write("Uploaded node data")
            st.dataframe(nodes_df, use_container_width=True, hide_index=True)
            st.write("Uploaded edge data")
            st.dataframe(edges_df, use_container_width=True, hide_index=True)

    return nodes_df, edges_df


def render_case(case_key):
    info = CASE_INFO[case_key]
    st.header(info["full_name"] + " -- " + info["purpose"])
    st.write(info["explanation"])

    st.subheader("Typical Use Cases")
    for item in info["use_cases"]:
        st.write("- " + item)

    render_schema_section()

    nodes_df, edges_df = render_data_section(case_key)

    if nodes_df is None or edges_df is None:
        st.info("Provide data above to continue with this case")
        return

    graph = build_graph(nodes_df, edges_df)
    node_ids = list(nodes_df["ID"])

    st.subheader("Step 2: Choose Start and Goal Nodes")
    col1, col2 = st.columns(2)
    with col1:
        start = st.selectbox("Start node", node_ids, key=case_key + "_start")
    with col2:
        goal = st.selectbox("Goal node", node_ids, index=min(len(node_ids) - 1, 1), key=case_key + "_goal")

    weight_field = "Cost"
    if case_key == "UCS":
        weight_field = st.selectbox(
            "Cost field used by Uniform Cost Search",
            ["Cost", "Risk", "Time", "Probability"],
            key=case_key + "_weightfield",
        )

    if st.button("Run " + info["full_name"], key=case_key + "_run"):
        if case_key == "DFS":
            path, visited_order = dfs_search(graph, start, goal)
            extra = None
        elif case_key == "BFS":
            path, visited_order, levels = bfs_search(graph, start, goal)
            extra = "Depth level of goal node: " + str(levels.get(goal, "not reached"))
        else:
            path, cost, visited_order = ucs_search(graph, start, goal, weight_field=weight_field)
            extra = "Total cumulative " + weight_field.lower() + " of path: " + (str(round(cost, 2)) if cost is not None else "not applicable")

        summary = build_result_summary(case_key, start, goal, path, extra, nodes_df, edges_df, visited_order)
        st.session_state[case_key + "_summary"] = summary

    if case_key + "_summary" in st.session_state:
        summary = st.session_state[case_key + "_summary"]

        st.subheader("Step 3: Results")
        st.write("Status: " + summary["status"])
        st.write("Path: " + summary["path_text"])
        if summary["extra"] is not None:
            st.write(summary["extra"])
        st.write("Nodes visited during search, in order:")
        st.write(" then ".join(summary["visited_order"]))

        st.subheader("Graph Visualization")
        fig = plot_graph(graph, path=summary["path"], title=info["full_name"] + " Result")
        st.pyplot(fig)

        st.subheader("Step 4: Export Results")
        col1, col2, col3, col4 = st.columns(4)
        with col1:
            st.download_button(
                "Download PDF",
                data=export_pdf_bytes(summary),
                file_name=case_key + "_result.pdf",
                mime="application/pdf",
                key=case_key + "_pdf",
            )
        with col2:
            st.download_button(
                "Download Word",
                data=export_word_bytes(summary),
                file_name=case_key + "_result.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=case_key + "_word",
            )
        with col3:
            st.download_button(
                "Download CSV",
                data=export_csv_bytes(summary["path"], nodes_df),
                file_name=case_key + "_result.csv",
                mime="text/csv",
                key=case_key + "_csv",
            )
        with col4:
            st.download_button(
                "Download Text",
                data=export_text_bytes(summary),
                file_name=case_key + "_result.txt",
                mime="text/plain",
                key=case_key + "_text",
            )


# ---------------------------------------------------------------------------
# Main application
# ---------------------------------------------------------------------------

def main():
    apply_page_style()

    st.write(
        "This application demonstrates three Uninformed Search techniques, Depth First Search, "
        "Breadth First Search, and Uniform Cost Search, applied to a security and enterprise "
        "Knowledge Graph. Select a tab below to explore each technique."
    )

    tab_dfs, tab_bfs, tab_ucs = st.tabs([
        "DFS -- Attack Chain Exploration",
        "BFS -- Blast Radius and Impact Analysis",
        "UCS -- Least Cost Remediation Path",
    ])

    with tab_dfs:
        render_case("DFS")

    with tab_bfs:
        render_case("BFS")

    with tab_ucs:
        render_case("UCS")

    st.markdown("---")
    st.caption("Kalsnet Graph Search and Reasoning Engine, Developed By Randy Singh from Kalsnet (KNet) Consulting Group")


if __name__ == "__main__":
    main()