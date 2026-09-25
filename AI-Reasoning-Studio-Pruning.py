# Kalsnet AI Reasoning Studio
# ============================
# Three self-contained reasoning modules for the Kalsnet AI + Knowledge Graph
# platform, each demonstrating a distinct class of reasoning used across
# industrial AI decision systems:
#
#   1. Minimax with Alpha-Beta Pruning -> Adversarial reasoning (game-tree search)
#   2. Expectimax                      -> Probabilistic reasoning (decisions under uncertainty)
#   3. 3-SAT                           -> Logical constraint reasoning (satisfiability)
#
# Every module ships with generated synthetic data so it works out of the box,
# and every module also accepts real data (CSV upload or in-app editing) to
# override the synthetic values. Every input field and every result is
# explained in plain language, and all results (summary, tables, charts and
# explanations) can be exported as PDF, Word (.docx), CSV, or plain text.
#
# Developed by Randy Singh - Kalsnet (KNet) Consulting Group

import csv
import inspect
import io
import itertools
import math
import os
import random
import struct
import tempfile
from datetime import datetime

import matplotlib

matplotlib.use("Agg")  # headless backend for Streamlit Cloud servers

import matplotlib.pyplot as plt  # noqa: E402
import networkx as nx  # noqa: E402
import pandas as pd  # noqa: E402
import streamlit as st  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Inches, RGBColor  # noqa: E402

# ============================================================================
# PDF LIBRARY IMPORT (works with fpdf2 AND the legacy "fpdf" package)
# ============================================================================
try:
    from fpdf import FPDF, XPos, YPos  # fpdf2 (preferred)

    FPDF_MODERN = True
except ImportError:  # legacy PyFPDF is installed instead of fpdf2
    from enum import Enum

    from fpdf import FPDF as _LegacyFPDF

    FPDF_MODERN = False

    class XPos(str, Enum):
        LEFT = "LEFT"
        RIGHT = "RIGHT"
        START = "START"
        END = "END"
        WCONT = "WCONT"
        CENTER = "CENTER"
        LMARGIN = "LMARGIN"
        RMARGIN = "RMARGIN"

    class YPos(str, Enum):
        TOP = "TOP"
        LAST = "LAST"
        NEXT = "NEXT"
        TMARGIN = "TMARGIN"
        BMARGIN = "BMARGIN"

    class FPDF(_LegacyFPDF):
        """Lets the legacy library accept fpdf2-style arguments."""

        def cell(self, w=0, h=0, txt="", border=0, ln=0, align="", fill=False,
                 link="", text=None, new_x=None, new_y=None, **kwargs):
            if text is not None:
                txt = text
            if new_y == YPos.NEXT:
                ln = 1 if new_x in (None, XPos.LMARGIN) else 2
            return super().cell(w, h, str(txt), border, ln, align, fill, link)

        def multi_cell(self, w, h=0, txt="", border=0, align="J", fill=False,
                       text=None, new_x=None, new_y=None, **kwargs):
            if text is not None:
                txt = text
            return super().multi_cell(w, h, str(txt), border, align, fill)

        def output(self, name="", dest=""):
            if not name and not dest:
                return super().output(dest="S").encode("latin-1")
            return super().output(name, dest)


# ============================================================================
# STREAMLIT VERSION COMPATIBILITY
# ============================================================================
_PYPLOT_HAS_WIDTH = "width" in inspect.signature(st.pyplot).parameters


def show_fig(fig) -> bytes:
    """Renders a matplotlib figure full-width, returns it as PNG bytes for the
    PDF / Word exports, then frees its memory."""
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=120, bbox_inches="tight", facecolor="white")
    if _PYPLOT_HAS_WIDTH:
        st.pyplot(fig, width="stretch")
    else:
        st.pyplot(fig, use_container_width=True)
    plt.close(fig)
    return buf.getvalue()


def full_width_table(df):
    try:
        st.dataframe(df, width="stretch", hide_index=True)
    except Exception:  # older Streamlit versions
        st.dataframe(df, use_container_width=True, hide_index=True)


def full_width_editor(df, key, column_config=None):
    kwargs = dict(key=key, num_rows="fixed", hide_index=True, column_config=column_config)
    try:
        return st.data_editor(df, width="stretch", **kwargs)
    except Exception:  # older Streamlit versions only accept use_container_width
        return st.data_editor(df, use_container_width=True, **kwargs)


# ============================================================================
# PAGE CONFIG & GLOBAL STYLE
# ============================================================================
st.set_page_config(
    page_title="Kalsnet AI Reasoning Studio",
    layout="wide",
    initial_sidebar_state="expanded",
)

PRIMARY_BLUE = "#1a3d6d"
ACCENT_BLUE = "#1a73e8"
ACCENT_RED = "#d93025"
ACCENT_AMBER = "#f9ab00"
ACCENT_GREEN = "#1e8e3e"
ACCENT_PURPLE = "#8430ce"
PRUNED_GREY = "#c9d1db"

# Plain string (not an f-string) so CSS braces never break Python parsing.
CSS = """
<style>
.stApp {
    background: linear-gradient(180deg, #f4f7fb 0%, #eef3fa 100%);
}
.title-block {
    text-align: center;
    padding: 18px 10px 8px 10px;
}
.title-main {
    font-family: 'Segoe UI', Arial, sans-serif;
    font-weight: 800;
    color: PRIMARY_BLUE;
    font-size: 42px;
    line-height: 1.25;
    margin: 0;
}
.title-sub {
    font-family: 'Segoe UI', Arial, sans-serif;
    font-weight: 800;
    color: PRIMARY_BLUE;
    font-size: 24px;
    line-height: 1.3;
    margin: 4px 0 0 0;
}
.module-banner {
    border-radius: 10px;
    padding: 14px 20px;
    margin-bottom: 14px;
    color: white;
    font-size: 20px;
    font-weight: 700;
    font-family: 'Segoe UI', Arial, sans-serif;
}
.banner-minimax { background: linear-gradient(90deg, ACCENT_BLUE, ACCENT_RED); }
.banner-expectimax { background: linear-gradient(90deg, ACCENT_AMBER, ACCENT_BLUE); }
.banner-3sat { background: linear-gradient(90deg, ACCENT_PURPLE, ACCENT_GREEN); }
.banner-about { background: linear-gradient(90deg, PRIMARY_BLUE, ACCENT_BLUE); }
.metric-card {
    background: white;
    border-radius: 10px;
    padding: 14px 18px;
    box-shadow: 0 1px 4px rgba(0,0,0,0.08);
    border-left: 5px solid ACCENT_BLUE;
    min-height: 118px;
    margin-bottom: 10px;
}
.metric-note {
    font-size: 12.5px;
    color: #5f6b7a;
    line-height: 1.35;
    margin-top: 4px;
}
.stTabs [data-baseweb="tab-list"] { gap: 6px; }
.stTabs [data-baseweb="tab"] {
    background-color: #e4ecf7;
    border-radius: 8px 8px 0 0;
    padding: 10px 16px;
    font-weight: 600;
}
.stTabs [aria-selected="true"] {
    background-color: PRIMARY_BLUE !important;
    color: white !important;
}
</style>
"""
for _name, _value in {
    "PRIMARY_BLUE": PRIMARY_BLUE, "ACCENT_BLUE": ACCENT_BLUE, "ACCENT_RED": ACCENT_RED,
    "ACCENT_AMBER": ACCENT_AMBER, "ACCENT_GREEN": ACCENT_GREEN, "ACCENT_PURPLE": ACCENT_PURPLE,
}.items():
    CSS = CSS.replace(_name, _value)
st.markdown(CSS, unsafe_allow_html=True)

st.markdown(
    """
    <div class="title-block">
        <div class="title-main">Kalsnet AI Reasoning Studio</div>
        <div class="title-sub">Developed by Randy Singh - Kalsnet (KNet) Consulting Group</div>
    </div>
    """,
    unsafe_allow_html=True,
)
st.caption(
    "Adversarial reasoning (Minimax with Alpha-Beta Pruning) | Probabilistic reasoning "
    "(Expectimax) | Logical constraint reasoning (3-SAT) - building blocks for the "
    "Kalsnet AI + Knowledge Graph platform."
)
st.divider()


# ============================================================================
# REPORT SECTIONS (one structure feeds the screen AND all four export formats)
# ============================================================================

def sec_text(heading, body):
    return {"kind": "text", "heading": heading, "body": str(body)}


def sec_table(heading, df, note=""):
    return {"kind": "table", "heading": heading, "df": df, "note": note}


def sec_defs(heading, pairs):
    """pairs: list of (term, plain-language meaning)."""
    return {"kind": "defs", "heading": heading, "pairs": list(pairs)}


def sec_image(heading, png, caption=""):
    return {"kind": "image", "heading": heading, "png": png, "caption": caption}


def defs_df(pairs, left="Field", right="What it means"):
    return pd.DataFrame(pairs, columns=[left, right])


def show_defs(pairs, left="Field", right="What it means"):
    full_width_table(defs_df(pairs, left, right))


# ============================================================================
# EXPORT HELPERS
# ============================================================================
_PDF_REPLACEMENTS = {
    "\u2014": "-", "\u2013": "-", "\u2018": "'", "\u2019": "'", "\u201c": '"',
    "\u201d": '"', "\u2026": "...", "\u2022": "-", "\u00a0": " ", "\u2192": "->",
    "\u2264": "<=", "\u2265": ">=", "\u221e": "inf",
}


def pdf_safe(text) -> str:
    """Built-in PDF fonts only support Latin-1; convert everything else."""
    text = str(text)
    for k, v in _PDF_REPLACEMENTS.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def png_size(png: bytes):
    """Returns (width, height) in pixels from a PNG header."""
    w, h = struct.unpack(">II", png[16:24])
    return w, h


def export_txt_bytes(title: str, sections: list) -> bytes:
    lines = [title, "=" * len(title),
             f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
             "Kalsnet (KNet) Consulting Group - Randy Singh", ""]
    for s in sections:
        lines += [s["heading"], "-" * len(s["heading"])]
        if s["kind"] == "text":
            lines.append(s["body"])
        elif s["kind"] == "table":
            if s["note"]:
                lines.append(s["note"])
            lines.append(s["df"].to_string(index=False))
        elif s["kind"] == "defs":
            for term, meaning in s["pairs"]:
                lines.append(f"* {term}: {meaning}")
        elif s["kind"] == "image":
            lines.append(f"[Chart: {s['caption'] or s['heading']} - "
                         "the chart image is included in the PDF and Word exports.]")
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def export_csv_report_bytes(title: str, sections: list) -> bytes:
    """Full report as one CSV: every section, table and explanation, stacked
    top-to-bottom with a section title row and a blank separator row."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([title])
    w.writerow([f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
    w.writerow([])
    for s in sections:
        w.writerow([f"SECTION: {s['heading']}"])
        if s["kind"] == "text":
            for line in s["body"].split("\n"):
                w.writerow([line])
        elif s["kind"] == "table":
            if s["note"]:
                w.writerow([s["note"]])
            w.writerow(list(s["df"].columns))
            for row in s["df"].itertuples(index=False):
                w.writerow(list(row))
        elif s["kind"] == "defs":
            w.writerow(["Field", "What it means"])
            for term, meaning in s["pairs"]:
                w.writerow([term, meaning])
        elif s["kind"] == "image":
            w.writerow([f"Chart: {s['caption'] or s['heading']} "
                        "(image included in the PDF and Word exports)"])
        w.writerow([])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens it cleanly


def export_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8-sig")


def export_docx_bytes(title: str, sections: list) -> bytes:
    doc = Document()
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x3D, 0x6D)
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph("Kalsnet (KNet) Consulting Group - Randy Singh")
    for s in sections:
        doc.add_heading(s["heading"], level=2)
        if s["kind"] == "text":
            for para in s["body"].split("\n"):
                doc.add_paragraph(para)
        elif s["kind"] == "table":
            if s["note"]:
                doc.add_paragraph().add_run(s["note"]).italic = True
            df = s["df"]
            table = doc.add_table(rows=1, cols=max(len(df.columns), 1))
            try:
                table.style = "Light Grid Accent 1"
            except Exception:
                table.style = "Table Grid"
            for i, col in enumerate(df.columns):
                table.rows[0].cells[i].text = str(col)
            for row in df.itertuples(index=False):
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = str(val)
        elif s["kind"] == "defs":
            for term, meaning in s["pairs"]:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{term}: ").bold = True
                p.add_run(meaning)
        elif s["kind"] == "image":
            doc.add_picture(io.BytesIO(s["png"]), width=Inches(6.3))
            if s["caption"]:
                doc.add_paragraph().add_run(s["caption"]).italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_paragraph(pdf, height, text):
    """Full-width paragraph that always returns the cursor to the left margin."""
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, height, pdf_safe(text), align="L", new_x=XPos.LMARGIN, new_y=YPos.NEXT)


def _pdf_fit(pdf, text, width):
    """Shortens text so it fits inside a table cell of the given width."""
    text = pdf_safe(text)
    if pdf.get_string_width(text) <= width - 2:
        return text
    while text and pdf.get_string_width(text + "..") > width - 2:
        text = text[:-1]
    return text + ".."


def _pdf_table(pdf, df, usable_w):
    cols = list(df.columns)
    if not cols:
        return
    str_rows = [[str(v) for v in row] for row in df.itertuples(index=False)]
    weights = []
    for i, c in enumerate(cols):
        longest = max([len(str(c))] + [len(r[i]) for r in str_rows]) if str_rows else len(str(c))
        weights.append(min(max(longest, 4), 45))
    total = sum(weights)
    widths = [usable_w * wgt / total for wgt in weights]

    def header():
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_fill_color(228, 236, 247)
        for c, wd in zip(cols, widths):
            pdf.cell(wd, 6, _pdf_fit(pdf, c, wd), border=1, fill=True)
        pdf.ln()
        pdf.set_font("Helvetica", "", 8)

    header()
    for r in str_rows:
        if pdf.get_y() + 6 > pdf.h - pdf.b_margin:
            pdf.add_page()
            header()
        pdf.set_x(pdf.l_margin)
        for v, wd in zip(r, widths):
            pdf.cell(wd, 6, _pdf_fit(pdf, v, wd), border=1)
        pdf.ln()


def _pdf_image(pdf, png, usable_w):
    wpx, hpx = png_size(png)
    w_mm = usable_w
    h_mm = w_mm * hpx / max(wpx, 1)
    max_h = pdf.h - pdf.t_margin - pdf.b_margin - 10
    if h_mm > max_h:
        h_mm = max_h
        w_mm = h_mm * wpx / max(hpx, 1)
    if pdf.get_y() + h_mm > pdf.h - pdf.b_margin:
        pdf.add_page()
    fd, path = tempfile.mkstemp(suffix=".png")
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(png)
        y = pdf.get_y()
        pdf.image(path, x=pdf.l_margin, y=y, w=w_mm, h=h_mm)
        pdf.set_y(y + h_mm + 2)
    finally:
        try:
            os.remove(path)
        except OSError:
            pass


def export_pdf_bytes(title: str, sections: list) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    usable_w = pdf.w - pdf.l_margin - pdf.r_margin

    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(26, 61, 109)
    _pdf_paragraph(pdf, 9, title)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(90, 90, 90)
    _pdf_paragraph(pdf, 5, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  "
                           "Kalsnet (KNet) Consulting Group - Randy Singh")
    pdf.ln(3)

    for s in sections:
        needed = 20
        if s["kind"] == "image":  # keep a chart together with its heading
            wpx, hpx = png_size(s["png"])
            needed = min(usable_w * hpx / max(wpx, 1), pdf.h - pdf.t_margin - pdf.b_margin - 10) + 12
        if pdf.get_y() + needed > pdf.h - pdf.b_margin:
            pdf.add_page()
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(26, 61, 109)
        _pdf_paragraph(pdf, 7, s["heading"])
        pdf.set_text_color(30, 30, 30)
        if s["kind"] == "text":
            pdf.set_font("Helvetica", "", 10)
            _pdf_paragraph(pdf, 5.5, s["body"])
        elif s["kind"] == "table":
            if s["note"]:
                pdf.set_font("Helvetica", "I", 9)
                _pdf_paragraph(pdf, 5, s["note"])
            _pdf_table(pdf, s["df"], usable_w)
        elif s["kind"] == "defs":
            for term, meaning in s["pairs"]:
                pdf.set_font("Helvetica", "B", 9.5)
                _pdf_paragraph(pdf, 5, term)
                pdf.set_font("Helvetica", "", 9.5)
                _pdf_paragraph(pdf, 5, meaning)
                pdf.ln(1)
        elif s["kind"] == "image":
            _pdf_image(pdf, s["png"], usable_w)
            if s["caption"]:
                pdf.set_font("Helvetica", "I", 8.5)
                _pdf_paragraph(pdf, 4.5, s["caption"])
        pdf.ln(3)
    return bytes(pdf.output())


def render_export_bar(key_prefix: str, title: str, sections: list, data_df: pd.DataFrame):
    st.markdown("##### Export all results")
    st.caption(
        "PDF and Word contain everything on this tab: settings, results, explanations, "
        "tables and charts. 'CSV full report' contains every section and table in one "
        "spreadsheet-friendly file. 'CSV data only' is the main data table, ready for "
        "analysis. 'Text' is a plain-text copy of the full report."
    )
    c1, c2, c3, c4, c5 = st.columns(5)
    with c1:
        try:
            st.download_button("PDF", data=export_pdf_bytes(title, sections),
                               file_name=f"{key_prefix}.pdf", mime="application/pdf",
                               key=f"{key_prefix}_pdf")
        except Exception as e:  # never let an export failure crash the app
            st.warning(f"PDF export unavailable: {e}")
    with c2:
        try:
            st.download_button("Word (.docx)", data=export_docx_bytes(title, sections),
                               file_name=f"{key_prefix}.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               key=f"{key_prefix}_docx")
        except Exception as e:
            st.warning(f"Word export unavailable: {e}")
    with c3:
        st.download_button("CSV full report", data=export_csv_report_bytes(title, sections),
                           file_name=f"{key_prefix}_full_report.csv", mime="text/csv",
                           key=f"{key_prefix}_csv_full")
    with c4:
        st.download_button("CSV data only", data=export_csv_bytes(data_df),
                           file_name=f"{key_prefix}_data.csv", mime="text/csv",
                           key=f"{key_prefix}_csv_data")
    with c5:
        st.download_button("Text", data=export_txt_bytes(title, sections),
                           file_name=f"{key_prefix}.txt", mime="text/plain",
                           key=f"{key_prefix}_txt")


# ============================================================================
# SMALL UI HELPERS
# ============================================================================

def metric_card(col, label, value, color, note="", size=26):
    col.markdown(
        f'<div class="metric-card"><b>{label}</b><br>'
        f'<span style="font-size:{size}px;color:{color};font-weight:700">{value}</span>'
        f'<div class="metric-note">{note}</div></div>',
        unsafe_allow_html=True,
    )


def fmt(x):
    """Readable number: +inf / -inf for infinities, 2 decimals otherwise."""
    if x is None:
        return ""
    if isinstance(x, bool):
        return str(x)
    if isinstance(x, (int, float)):
        if math.isinf(x):
            return "+inf" if x > 0 else "-inf"
        if float(x).is_integer():
            return f"{x:.0f}"
        return f"{x:.2f}"
    return str(x)


def read_leaf_csv(upload, n_leaves):
    """Reads a leaf_value CSV safely; returns a list of floats or None."""
    try:
        real_df = pd.read_csv(upload)
        real_df.columns = [str(c).strip().lower() for c in real_df.columns]
        if "leaf_value" not in real_df.columns:
            st.error("The CSV must contain a column named 'leaf_value'. "
                     "Download the template below to see the exact format.")
            return None
        values = pd.to_numeric(real_df["leaf_value"], errors="coerce").dropna().tolist()
        if not values:
            st.error("No numeric values were found in the 'leaf_value' column.")
            return None
        if len(values) != n_leaves:
            st.warning(f"The current tree has {n_leaves} leaves but the file has "
                       f"{len(values)} values. Missing values are filled with synthetic "
                       "data and extra values are ignored.")
        else:
            st.success(f"Loaded {len(values)} real leaf values.")
        return values
    except Exception as e:
        st.error(f"Could not read the CSV file: {e}")
        return None


LEAF_UPLOAD_SCHEMA = [
    ("leaf_value", "Required. Number (whole or decimal, negative allowed). One row per "
                   "leaf, in left-to-right order across the bottom of the tree. Example: 7, -3, 4.5"),
]

TREE_INPUT_DEFS = [
    ("Tree depth (plies)", "How many decision levels the tree has below the root. Each level "
                           "is one turn (a 'ply'). Deeper trees look further ahead but grow "
                           "quickly: leaves = branching factor to the power of depth."),
    ("Branching factor", "How many options (child branches) every decision point has."),
    ("Random seed", "A number that fixes the synthetic data. The same seed always produces "
                    "the same tree, so results are repeatable. Change it to get a new scenario."),
    ("Upload leaf-value CSV", "Optional. Replaces the synthetic leaf scores with your own "
                              "real scores (see the file format below)."),
    ("Leaf value table", "Optional. Type directly into the 'value' column to change any "
                         "leaf score. Results update instantly."),
]


# ============================================================================
# TREE UTILITIES (shared by Minimax + Expectimax)
# ============================================================================

def build_tree(depth: int, branching: int, node_types: list, leaf_values=None, seed: int = 42):
    """Builds a synthetic layered tree. node_types[level] gives the node type
    for every node at that level ('MAX', 'MIN', 'CHANCE', or 'LEAF' for the last).
    Node ids describe the route from the root, e.g. n0-2 = first branch, then third."""
    rng = random.Random(int(seed))
    G = nx.DiGraph()
    leaf_iter = itertools.count()

    def add_node(level, path):
        node_id = "n" + ("-".join(map(str, path)) if path else "root")
        ntype = node_types[level]
        G.add_node(node_id, type=ntype, level=level)
        if ntype == "LEAF":
            idx = next(leaf_iter)
            synthetic = float(rng.randint(-10, 10))
            if leaf_values is not None and idx < len(leaf_values):
                val = float(leaf_values[idx])
            else:
                val = synthetic
            G.nodes[node_id]["value"] = val
            return node_id
        children = [add_node(level + 1, path + [i]) for i in range(branching)]
        probs = None
        if ntype == "CHANCE":
            raw = [rng.random() + 0.1 for _ in children]
            total = sum(raw)
            probs = [round(p / total, 3) for p in raw]
            probs[-1] = round(1 - sum(probs[:-1]), 3)  # normalize rounding drift
        for j, child in enumerate(children):
            edge_kwargs = {"prob": probs[j]} if probs else {}
            G.add_edge(node_id, child, **edge_kwargs)
        return node_id

    root = add_node(0, [])
    return G, root


def layered_layout(G):
    levels = {}
    for n, d in G.nodes(data=True):
        levels.setdefault(d["level"], []).append(n)
    pos = {}
    for level, nodes in levels.items():
        count = len(nodes)
        for i, node in enumerate(nodes):
            pos[node] = (i - (count - 1) / 2, -level)
    return pos


def tree_leaves(G):
    return [n for n, d in G.nodes(data=True) if d["type"] == "LEAF"]


def apply_leaf_edits(G, leaves, key):
    """Shows an editable leaf table and writes any edits back into the tree."""
    default_leaf_df = pd.DataFrame({"leaf": leaves, "value": [G.nodes[l]["value"] for l in leaves]})
    cfg = {
        "leaf": st.column_config.TextColumn(
            "leaf", help="Leaf id: the route from the root, e.g. n1-0-1 = 2nd, 1st, then 2nd branch.",
            disabled=True),
        "value": st.column_config.NumberColumn(
            "value", help="Score of this final outcome for MAX. Higher is better for MAX."),
    }
    try:
        edited = full_width_editor(default_leaf_df, key, column_config=cfg)
    except Exception:
        edited = full_width_editor(default_leaf_df, key)
    for l, v in zip(edited["leaf"], edited["value"]):
        try:
            if v is not None and not pd.isna(v):
                G.nodes[l]["value"] = float(v)
        except (TypeError, ValueError):
            pass  # keep the previous value if a cell is blank or invalid


def reconstruct_path(G, root):
    path = [root]
    node = root
    while "chosen" in G.nodes[node]:
        node = G.nodes[node]["chosen"]
        path.append(node)
    return path, list(zip(path[:-1], path[1:]))


def draw_tree(G, pos, path_edges=None, value_fmt="{:.1f}", pruned=None):
    pruned = pruned or set()
    n_nodes = G.number_of_nodes()
    node_size = 900 if n_nodes <= 31 else (520 if n_nodes <= 60 else 260)
    font_size = 8 if n_nodes <= 31 else (7 if n_nodes <= 60 else 5)
    fig, ax = plt.subplots(figsize=(10, 4.8))
    color_map = {"MAX": ACCENT_BLUE, "MIN": ACCENT_RED, "CHANCE": ACCENT_AMBER, "LEAF": ACCENT_GREEN}

    live_edges = [(u, v) for u, v in G.edges if v not in pruned]
    cut_edges = [(u, v) for u, v in G.edges if v in pruned]
    nx.draw_networkx_edges(G, pos, edgelist=live_edges, edge_color="#b0b8c4", width=1.3,
                           arrows=False, ax=ax)
    if cut_edges:
        nx.draw_networkx_edges(G, pos, edgelist=cut_edges, edge_color="#c9d1db", width=1.1,
                               style="dashed", arrows=False, ax=ax)
    if path_edges:
        nx.draw_networkx_edges(G, pos, edgelist=path_edges, edge_color=ACCENT_GREEN, width=3.2,
                               arrows=False, ax=ax)
    nodes = list(G.nodes)
    colors = [PRUNED_GREY if n in pruned else color_map.get(G.nodes[n]["type"], "#888888")
              for n in nodes]
    nx.draw_networkx_nodes(G, pos, nodelist=nodes, node_color=colors, node_size=node_size, ax=ax)

    labels = {}
    for n in nodes:
        if n in pruned:
            labels[n] = "cut"
        elif "value" in G.nodes[n]:
            labels[n] = value_fmt.format(G.nodes[n]["value"])
    nx.draw_networkx_labels(G, pos, labels, font_size=font_size, font_color="white",
                            font_weight="bold", ax=ax)
    edge_labels = nx.get_edge_attributes(G, "prob")
    if edge_labels and n_nodes <= 60:
        nx.draw_networkx_edge_labels(G, pos, edge_labels={k: f"p={v}" for k, v in edge_labels.items()},
                                     font_size=6.5, ax=ax)
    ax.axis("off")
    fig.tight_layout()
    return fig


# ============================================================================
# MODULE 1: MINIMAX WITH ALPHA-BETA PRUNING (Adversarial Reasoning)
# ============================================================================

def minimax_full(G, node, counter):
    """Plain minimax: visits every node. Used as the baseline for comparison."""
    counter["nodes"] += 1
    ntype = G.nodes[node]["type"]
    if ntype == "LEAF":
        counter["leaves"] += 1
        return G.nodes[node]["value"]
    vals = [(c, minimax_full(G, c, counter)) for c in G.successors(node)]
    if ntype == "MAX":
        chosen, val = max(vals, key=lambda cv: cv[1])
    else:
        chosen, val = min(vals, key=lambda cv: cv[1])
    G.nodes[node]["value"] = val
    G.nodes[node]["chosen"] = chosen
    return val


def alphabeta(G, node, alpha, beta, ctx):
    """Minimax with alpha-beta pruning.
    alpha = best score MAX is already guaranteed on the path so far (lower bound).
    beta  = best score MIN is already guaranteed on the path so far (upper bound).
    When alpha >= beta, the remaining children cannot change the decision, so
    they are skipped (pruned)."""
    d = G.nodes[node]
    ctx["visited"].add(node)
    ntype = d["type"]
    if ntype == "LEAF":
        ctx["leaves"] += 1
        ctx["trace"].append((node, ntype, alpha, beta, d["value"], "Leaf evaluated"))
        return d["value"]

    children = ctx["order"](G, node)
    best = -math.inf if ntype == "MAX" else math.inf
    chosen = None
    for i, c in enumerate(children):
        v = alphabeta(G, c, alpha, beta, ctx)
        if ntype == "MAX":
            if v > best:
                best, chosen = v, c
            alpha = max(alpha, best)
        else:
            if v < best:
                best, chosen = v, c
            beta = min(beta, best)
        if alpha >= beta:
            skipped = children[i + 1:]
            if skipped:
                ctx["cutoffs"] += 1
                for s in skipped:
                    ctx["pruned_roots"].append(s)
                kind = "Beta cutoff" if ntype == "MAX" else "Alpha cutoff"
                ctx["trace"].append((node, ntype, alpha, beta, best,
                                     f"{kind}: alpha >= beta, skipped {len(skipped)} branch(es): "
                                     f"{', '.join(skipped)}"))
            break
    d["value"] = best
    d["chosen"] = chosen
    d["alpha"], d["beta"] = alpha, beta
    ctx["trace"].append((node, ntype, alpha, beta, best, "Value returned to parent"))
    return best


def make_order(mode, true_values):
    def as_generated(G, node):
        return list(G.successors(node))

    def reversed_order(G, node):
        return list(G.successors(node))[::-1]

    def ideal(G, node):
        kids = list(G.successors(node))
        desc = G.nodes[node]["type"] == "MAX"
        return sorted(kids, key=lambda c: true_values[c], reverse=desc)

    return {"Left to right (as generated)": as_generated,
            "Right to left (reversed)": reversed_order,
            "Ideal ordering (best move first)": ideal}[mode]


MINIMAX_INPUT_DEFS = TREE_INPUT_DEFS[:3] + [
    ("Child ordering", "The order in which alpha-beta examines the options at each decision "
                       "point. The final answer never changes, but the amount of pruning does: "
                       "examining the best move first prunes the most. 'Ideal ordering' shows "
                       "the best case possible, because it uses the already known values."),
] + TREE_INPUT_DEFS[3:]

MINIMAX_RESULT_DEFS = [
    ("Minimax value", "The score MAX is guaranteed to get if both players play perfectly. "
                      "It is the value at the top (root) of the tree."),
    ("Best first move", "Which branch MAX should take at the root to secure the minimax value."),
    ("Optimal path (principal variation)", "The sequence of moves both players make when each "
                                           "plays perfectly. Drawn in green on the tree."),
    ("Alpha", "The best score MAX can already guarantee on the current path (a lower bound). "
              "Starts at -inf and only rises."),
    ("Beta", "The best score MIN can already guarantee on the current path (an upper bound). "
             "Starts at +inf and only falls."),
    ("Cutoff (pruning)", "When alpha >= beta, the remaining options at that point cannot affect "
                         "the final decision, so they are skipped. A 'beta cutoff' happens at a "
                         "MAX node, an 'alpha cutoff' at a MIN node."),
    ("Nodes visited", "How many tree positions alpha-beta actually examined."),
    ("Leaves evaluated", "How many final outcomes were scored. Plain minimax scores all of them."),
    ("Nodes pruned", "Tree positions alpha-beta proved it did not need to look at. Shown as grey "
                     "'cut' circles with dashed lines on the tree."),
    ("Work saved", "Percentage of leaf evaluations avoided compared with plain minimax."),
    ("Best-case leaves", "The theoretical minimum number of leaves alpha-beta must evaluate with "
                         "perfect move ordering: b^ceil(d/2) + b^floor(d/2) - 1."),
    ("Matches plain minimax", "A built-in correctness check: alpha-beta must always produce the "
                              "same value as plain minimax. 'Yes' confirms the result."),
]

MINIMAX_NODE_COLS = [
    ("node", "Position id: the route from the root, e.g. n0-1 = first branch, then second branch."),
    ("type", "MAX (our decision), MIN (opponent decision) or LEAF (final outcome with a score)."),
    ("level", "Depth in the tree: 0 is the root, higher numbers are further ahead."),
    ("status", "Visited = examined by alpha-beta. Pruned = skipped because it could not change the result."),
    ("value", "Score of the position. For pruned nodes it is blank because it was never computed."),
    ("alpha_final / beta_final", "The alpha and beta bounds when the search finished at this node."),
    ("on_optimal_path", "Yes if the position lies on the green optimal path."),
]

MINIMAX_TRACE_COLS = [
    ("step", "Order in which the search did things (1 = first)."),
    ("node / type", "Which position was handled and whether it was MAX, MIN or LEAF."),
    ("alpha / beta", "The bounds at that moment. Pruning happens as soon as alpha >= beta."),
    ("value", "Score at that moment: a leaf's score, or the best value found so far at a MAX/MIN node."),
    ("event", "What happened: a leaf was scored, a value was returned, or a cutoff pruned branches."),
]


def render_minimax_tab():
    st.markdown('<div class="module-banner banner-minimax">Minimax with Alpha-Beta Pruning - '
                'Adversarial Reasoning</div>', unsafe_allow_html=True)
    st.write(
        "Two opposing agents alternate turns down a game tree. MAX (blue) always picks the "
        "branch with the highest guaranteed score; MIN (red) always picks the branch that is "
        "worst for MAX. This models zero-sum adversarial decisions such as competitive bidding, "
        "negotiation, or an attacker and defender on an industrial network."
    )
    st.write(
        "**Alpha-beta pruning** gives exactly the same answer as plain minimax but skips "
        "branches that provably cannot change the decision, so it examines far fewer "
        "positions. The grey 'cut' nodes on the tree are the branches it skipped."
    )

    left, right = st.columns([1, 2])
    with left:
        st.markdown("**Settings**")
        depth = st.select_slider("Tree depth (plies)", options=[2, 3, 4], value=3, key="mm_depth",
                                 help=MINIMAX_INPUT_DEFS[0][1])
        branching = st.select_slider("Branching factor", options=[2, 3], value=2, key="mm_branch",
                                     help=MINIMAX_INPUT_DEFS[1][1])
        seed = int(st.number_input("Random seed", value=7, step=1, key="mm_seed",
                                   help=MINIMAX_INPUT_DEFS[2][1]))
        order_mode = st.radio("Child ordering", ["Left to right (as generated)",
                                                 "Right to left (reversed)",
                                                 "Ideal ordering (best move first)"],
                              key="mm_order", help=MINIMAX_INPUT_DEFS[3][1])

        node_types = ["MAX" if lvl % 2 == 0 else "MIN" for lvl in range(depth)] + ["LEAF"]
        n_leaves = branching ** depth

        st.markdown("**Use your own data (optional)**")
        upload = st.file_uploader("Upload leaf-value CSV", type=["csv"], key="mm_upload",
                                  help=MINIMAX_INPUT_DEFS[4][1])
        with st.expander("CSV file format"):
            st.write(f"The current tree needs exactly **{n_leaves}** rows.")
            show_defs(LEAF_UPLOAD_SCHEMA, "Column", "Rules and example")
        leaf_values = read_leaf_csv(upload, n_leaves) if upload is not None else None
        template = pd.DataFrame({"leaf_value": [0] * n_leaves})
        st.download_button("Download CSV template", data=export_csv_bytes(template),
                           file_name="minimax_leaf_template.csv", mime="text/csv", key="mm_template")

    G, root = build_tree(depth, branching, node_types, leaf_values=leaf_values, seed=seed)
    leaves = tree_leaves(G)

    st.markdown("**Leaf values** - edit any score in the 'value' column to test a scenario")
    apply_leaf_edits(G, leaves, key=f"mm_editor_{depth}_{branching}_{seed}")

    # --- Baseline: plain minimax on a copy (visits everything) ---
    G_full = G.copy()
    full_counter = {"nodes": 0, "leaves": 0}
    full_val = minimax_full(G_full, root, full_counter)
    true_values = {n: G_full.nodes[n]["value"] for n in G_full.nodes}

    # --- Alpha-beta search ---
    ctx = {"visited": set(), "leaves": 0, "cutoffs": 0, "pruned_roots": [], "trace": [],
           "order": make_order(order_mode, true_values)}
    ab_val = alphabeta(G, root, -math.inf, math.inf, ctx)
    pruned = set()
    for pr in ctx["pruned_roots"]:
        pruned.add(pr)
        pruned |= nx.descendants(G, pr)
    pruned -= ctx["visited"]
    path, path_edges = reconstruct_path(G, root)
    best_first_move = path[1] if len(path) > 1 else root

    total_nodes = G.number_of_nodes()
    visited_n = len(ctx["visited"])
    pruned_leaves = [l for l in leaves if l in pruned]
    saved_pct = 100.0 * (len(leaves) - ctx["leaves"]) / max(len(leaves), 1)
    best_case = branching ** math.ceil(depth / 2) + branching ** math.floor(depth / 2) - 1
    matches = abs(ab_val - full_val) < 1e-9

    with right:
        st.markdown("**Game tree** - blue = MAX, red = MIN, green circles = leaves, "
                    "green lines = optimal path, grey 'cut' = pruned by alpha-beta")
        tree_png = show_fig(draw_tree(G, layered_layout(G), path_edges=path_edges, pruned=pruned))

    r1 = st.columns(4)
    metric_card(r1[0], "Minimax value", fmt(ab_val), ACCENT_BLUE,
                "Score MAX is guaranteed with perfect play by both sides.")
    metric_card(r1[1], "Best first move", best_first_move, ACCENT_BLUE,
                "Branch MAX should choose at the root.")
    metric_card(r1[2], "Leaves evaluated", f"{ctx['leaves']} of {len(leaves)}", ACCENT_BLUE,
                f"Plain minimax scores all {len(leaves)}. Best case possible: {best_case}.")
    metric_card(r1[3], "Work saved by pruning", f"{saved_pct:.0f}%", ACCENT_GREEN,
                "Share of leaf evaluations skipped versus plain minimax.")
    r2 = st.columns(4)
    metric_card(r2[0], "Nodes visited", f"{visited_n} of {total_nodes}", ACCENT_BLUE,
                "Tree positions alpha-beta examined.")
    metric_card(r2[1], "Nodes pruned", f"{len(pruned)}", ACCENT_RED,
                "Positions skipped because they could not change the decision.")
    metric_card(r2[2], "Cutoffs", f"{ctx['cutoffs']}", ACCENT_RED,
                "Times alpha >= beta triggered pruning.")
    metric_card(r2[3], "Matches plain minimax", "Yes" if matches else "No",
                ACCENT_GREEN if matches else ACCENT_RED,
                f"Correctness check. Plain minimax value = {fmt(full_val)}.")

    # --- Comparison chart ---
    st.markdown("**Plain minimax vs. alpha-beta pruning** - same answer, less work")
    fig_c, ax_c = plt.subplots(figsize=(9, 2.6))
    cats = ["Nodes examined", "Leaves evaluated"]
    x = range(len(cats))
    ax_c.bar([i - 0.2 for i in x], [full_counter["nodes"], full_counter["leaves"]], width=0.4,
             color="#a9b8c8", label="Plain minimax")
    ax_c.bar([i + 0.2 for i in x], [visited_n, ctx["leaves"]], width=0.4,
             color=ACCENT_BLUE, label="Alpha-beta")
    ax_c.set_xticks(list(x))
    ax_c.set_xticklabels(cats)
    ax_c.set_ylabel("Count")
    ax_c.legend(fontsize=8)
    fig_c.tight_layout()
    compare_png = show_fig(fig_c)

    # --- Leaf chart ---
    st.markdown("**Leaf values** - green = outcome on the optimal path, blue = evaluated, "
                "grey = pruned (never evaluated)")
    fig2, ax2 = plt.subplots(figsize=(9, 2.8))
    leaf_vals_now = [G.nodes[l]["value"] for l in leaves]
    bar_colors = [ACCENT_GREEN if l in path else (PRUNED_GREY if l in pruned else ACCENT_BLUE)
                  for l in leaves]
    ax2.bar(range(len(leaves)), leaf_vals_now, color=bar_colors)
    ax2.axhline(0, color="#888888", linewidth=0.8)
    ax2.set_xticks(range(len(leaves)))
    ax2.set_xticklabels([l[1:] for l in leaves], rotation=45, fontsize=7)
    ax2.set_ylabel("Score for MAX")
    fig2.tight_layout()
    leaf_png = show_fig(fig2)

    # --- Result tables ---
    node_rows = []
    for n, d in G.nodes(data=True):
        is_pruned = n in pruned
        node_rows.append({
            "node": n,
            "type": d["type"],
            "level": d["level"],
            "status": "Pruned" if is_pruned else "Visited",
            "value": "" if is_pruned else fmt(d.get("value")),
            "alpha_final": fmt(d["alpha"]) if "alpha" in d and not is_pruned else "",
            "beta_final": fmt(d["beta"]) if "beta" in d and not is_pruned else "",
            "on_optimal_path": "Yes" if n in path else "No",
        })
    node_df = pd.DataFrame(node_rows)
    trace_df = pd.DataFrame(
        [{"step": i + 1, "node": n, "type": t, "alpha": fmt(a), "beta": fmt(b),
          "value": fmt(v), "event": e} for i, (n, t, a, b, v, e) in enumerate(ctx["trace"])]
    )

    with st.expander("How to read these results", expanded=False):
        show_defs(MINIMAX_RESULT_DEFS)
    with st.expander("All tree positions (node table)"):
        full_width_table(node_df)
        st.caption("Column meanings:")
        show_defs(MINIMAX_NODE_COLS, "Column")
    with st.expander("Step-by-step alpha-beta trace"):
        full_width_table(trace_df)
        st.caption("Column meanings:")
        show_defs(MINIMAX_TRACE_COLS, "Column")

    st.divider()
    summary = (
        f"Minimax value (alpha-beta): {fmt(ab_val)}\n"
        f"Plain minimax value: {fmt(full_val)}  |  Matches: {'Yes' if matches else 'No'}\n"
        f"Best first move for MAX: {best_first_move}\n"
        f"Optimal path: {' -> '.join(path)}\n"
        f"Nodes visited: {visited_n} of {total_nodes}  |  Nodes pruned: {len(pruned)}  |  "
        f"Cutoffs: {ctx['cutoffs']}\n"
        f"Leaves evaluated: {ctx['leaves']} of {len(leaves)}  |  Work saved: {saved_pct:.0f}%  |  "
        f"Best-case leaves: {best_case}\n"
        f"Pruned leaves: {', '.join(pruned_leaves) if pruned_leaves else 'none'}"
    )
    sections = [
        sec_text("Overview", "Minimax with alpha-beta pruning finds the best move for MAX against "
                             "an opponent (MIN) who always responds with the move worst for MAX. "
                             "Alpha-beta returns exactly the same answer as plain minimax while "
                             "skipping branches that cannot change the decision."),
        sec_text("Settings used", f"Tree depth: {depth}\nBranching factor: {branching}\n"
                                  f"Random seed: {seed}\nChild ordering: {order_mode}\n"
                                  f"Leaf data source: {'Uploaded CSV' if leaf_values else 'Synthetic (plus any manual edits)'}"),
        sec_defs("Input fields explained", MINIMAX_INPUT_DEFS),
        sec_text("Results summary", summary),
        sec_defs("Results explained", MINIMAX_RESULT_DEFS),
        sec_image("Game tree", tree_png, "Blue = MAX, red = MIN, green circles = leaves, green "
                                         "lines = optimal path, grey 'cut' nodes = pruned."),
        sec_image("Plain minimax vs. alpha-beta", compare_png,
                  "Grey = plain minimax, blue = alpha-beta. Same answer, less work."),
        sec_image("Leaf values", leaf_png, "Green = on optimal path, blue = evaluated, grey = pruned."),
        sec_table("Node table", node_df),
        sec_defs("Node table columns", MINIMAX_NODE_COLS),
        sec_table("Alpha-beta trace", trace_df, "Every step of the search, in order."),
        sec_defs("Trace columns", MINIMAX_TRACE_COLS),
    ]
    render_export_bar("minimax_alphabeta_results",
                      "Kalsnet AI Reasoning Studio - Minimax with Alpha-Beta Pruning",
                      sections, node_df)


# ============================================================================
# MODULE 2: EXPECTIMAX (Probabilistic Reasoning)
# ============================================================================

def expectimax(G, node):
    ntype = G.nodes[node]["type"]
    if ntype == "LEAF":
        return G.nodes[node]["value"]
    children = list(G.successors(node))
    if ntype == "MAX":
        vals = [(c, expectimax(G, c)) for c in children]
        chosen, val = max(vals, key=lambda cv: cv[1])
        G.nodes[node]["chosen"] = chosen
    else:  # CHANCE
        val = 0.0
        for c in children:
            p = G.edges[node, c].get("prob", 1 / len(children))
            val += p * expectimax(G, c)
    G.nodes[node]["value"] = val
    return val


def outcome_probabilities(G, root):
    """Probability of reaching each node if MAX follows the recommended strategy."""
    prob = {root: 1.0}
    for n in nx.topological_sort(G):
        if n not in prob:
            continue
        d = G.nodes[n]
        if d["type"] == "MAX":
            chosen = d.get("chosen")
            for c in G.successors(n):
                prob[c] = prob[n] if c == chosen else 0.0
        elif d["type"] == "CHANCE":
            for c in G.successors(n):
                prob[c] = prob[n] * G.edges[n, c].get("prob", 0.0)
    return prob


EXPECTIMAX_INPUT_DEFS = TREE_INPUT_DEFS

EXPECTIMAX_RESULT_DEFS = [
    ("Expected value", "The average score MAX can expect with the best strategy, weighting "
                       "every uncertain outcome by its probability. It is the value at the root."),
    ("Recommended first move", "The branch MAX should take at the root to get the highest "
                               "expected value."),
    ("Recommended decisions", "At every MAX node, the option with the highest expected value. "
                              "Drawn in green on the tree."),
    ("CHANCE node", "A point where the outcome is random (e.g. a machine fails or not). Its value "
                    "is the probability-weighted average of its outcomes."),
    ("p (edge probability)", "The chance that a particular random outcome happens. The "
                             "probabilities leaving each CHANCE node add up to 1."),
    ("Best / worst possible outcome", "The highest and lowest leaf scores that can actually happen "
                                      "when MAX follows the recommended strategy (the risk range)."),
    ("Probability of reaching", "How likely each final outcome is when MAX follows the recommended "
                                "strategy. Outcomes MAX avoids have probability 0."),
    ("Why no alpha-beta pruning here", "Pruning relies on an opponent that always picks the "
                                       "minimum. Chance nodes average all outcomes, so any "
                                       "unexplored outcome could still change the average. "
                                       "Expectimax therefore evaluates every leaf."),
]

EXPECTIMAX_LEAF_COLS = [
    ("leaf", "Final outcome id: the route from the root."),
    ("value", "Score of this outcome for MAX."),
    ("path_probability", "Product of chance probabilities along the route, assuming MAX takes "
                         "that route. Shows how likely the outcome is along its own path."),
    ("prob_under_strategy", "Probability of ending here when MAX follows the recommended "
                            "strategy (0 if the strategy avoids this outcome)."),
    ("contribution", "value x prob_under_strategy. These add up to the expected value."),
]

EXPECTIMAX_NODE_COLS = [
    ("node", "Position id: the route from the root."),
    ("type", "MAX (our decision), CHANCE (random event) or LEAF (final outcome)."),
    ("level", "Depth in the tree: 0 is the root."),
    ("expected_value", "Expected score of this position."),
    ("recommended_choice", "For MAX nodes: the best option to pick. Blank for other nodes."),
]


def render_expectimax_tab():
    st.markdown('<div class="module-banner banner-expectimax">Expectimax - Probabilistic Reasoning</div>',
                unsafe_allow_html=True)
    st.write(
        "MAX nodes (blue) still pick the best option, but the opponent is replaced by CHANCE "
        "nodes (amber) whose outcomes follow a probability distribution, for example equipment "
        "failure rates, demand uncertainty or sensor noise. The value of a chance node is the "
        "probability-weighted average of its outcomes. This models planning under uncertainty "
        "rather than against an adversary."
    )

    left, right = st.columns([1, 2])
    with left:
        st.markdown("**Settings**")
        depth = st.select_slider("Tree depth (plies)", options=[2, 3, 4], value=3, key="em_depth",
                                 help=EXPECTIMAX_INPUT_DEFS[0][1])
        branching = st.select_slider("Branching factor", options=[2, 3], value=2, key="em_branch",
                                     help=EXPECTIMAX_INPUT_DEFS[1][1])
        seed = int(st.number_input("Random seed", value=11, step=1, key="em_seed",
                                   help=EXPECTIMAX_INPUT_DEFS[2][1]))

        node_types = ["MAX" if lvl % 2 == 0 else "CHANCE" for lvl in range(depth)] + ["LEAF"]
        n_leaves = branching ** depth

        st.markdown("**Use your own data (optional)**")
        upload = st.file_uploader("Upload leaf-value CSV", type=["csv"], key="em_upload",
                                  help=EXPECTIMAX_INPUT_DEFS[3][1])
        with st.expander("CSV file format"):
            st.write(f"The current tree needs exactly **{n_leaves}** rows.")
            show_defs(LEAF_UPLOAD_SCHEMA, "Column", "Rules and example")
        leaf_values = read_leaf_csv(upload, n_leaves) if upload is not None else None
        template = pd.DataFrame({"leaf_value": [0] * n_leaves})
        st.download_button("Download CSV template", data=export_csv_bytes(template),
                           file_name="expectimax_leaf_template.csv", mime="text/csv",
                           key="em_template")

    G, root = build_tree(depth, branching, node_types, leaf_values=leaf_values, seed=seed)
    leaves = tree_leaves(G)

    st.markdown("**Leaf values** - edit any score in the 'value' column to test a scenario")
    apply_leaf_edits(G, leaves, key=f"em_editor_{depth}_{branching}_{seed}")

    exp_val = expectimax(G, root)
    path, path_edges = reconstruct_path(G, root)
    # Highlight every recommended MAX decision, not just the first one.
    decision_edges = [(n, d["chosen"]) for n, d in G.nodes(data=True)
                      if d["type"] == "MAX" and "chosen" in d]
    reach = outcome_probabilities(G, root)
    root_children = list(G.successors(root))
    best_first_move = G.nodes[root].get("chosen", root)

    reachable = [l for l in leaves if reach.get(l, 0) > 0]
    best_outcome = max(G.nodes[l]["value"] for l in reachable) if reachable else exp_val
    worst_outcome = min(G.nodes[l]["value"] for l in reachable) if reachable else exp_val

    with right:
        st.markdown("**Decision tree** - blue = MAX, amber = CHANCE (p = probability of each "
                    "outcome), green lines = recommended decisions")
        tree_png = show_fig(draw_tree(G, layered_layout(G), path_edges=decision_edges,
                                      value_fmt="{:.2f}"))

    r1 = st.columns(4)
    metric_card(r1[0], "Expected value", fmt(exp_val), ACCENT_AMBER,
                "Average score with the best strategy, weighted by probability.")
    metric_card(r1[1], "Recommended first move", best_first_move, ACCENT_AMBER,
                "Branch MAX should choose at the root.")
    metric_card(r1[2], "Outcome range", f"{fmt(worst_outcome)} to {fmt(best_outcome)}", ACCENT_AMBER,
                "Worst and best scores that can happen under the recommended strategy.")
    metric_card(r1[3], "Leaves evaluated", f"{len(leaves)}", ACCENT_AMBER,
                "Expectimax must evaluate every leaf (no pruning, see explanation).")

    st.markdown("**Expected value of each first move** - green = recommended")
    fig2, ax2 = plt.subplots(figsize=(9, 2.8))
    child_vals = [G.nodes[c]["value"] for c in root_children]
    colors = [ACCENT_GREEN if c == best_first_move else ACCENT_AMBER for c in root_children]
    ax2.bar(root_children, child_vals, color=colors)
    ax2.axhline(0, color="#888888", linewidth=0.8)
    ax2.set_ylabel("Expected value")
    fig2.tight_layout()
    moves_png = show_fig(fig2)

    # --- Result tables ---
    leaf_rows = []
    for l in leaves:
        route = nx.shortest_path(G, root, l)
        p_path = 1.0
        for u, v in zip(route[:-1], route[1:]):
            p_path *= G.edges[u, v].get("prob", 1.0)
        p_strat = reach.get(l, 0.0)
        leaf_rows.append({
            "leaf": l,
            "value": G.nodes[l]["value"],
            "path_probability": round(p_path, 4),
            "prob_under_strategy": round(p_strat, 4),
            "contribution": round(G.nodes[l]["value"] * p_strat, 4),
        })
    leaf_df = pd.DataFrame(leaf_rows)
    node_df = pd.DataFrame([{
        "node": n, "type": d["type"], "level": d["level"],
        "expected_value": round(d["value"], 4),
        "recommended_choice": d.get("chosen", "") if d["type"] == "MAX" else "",
    } for n, d in G.nodes(data=True)])

    with st.expander("How to read these results", expanded=False):
        show_defs(EXPECTIMAX_RESULT_DEFS)
    with st.expander("Leaf outcomes table"):
        full_width_table(leaf_df)
        st.caption(f"Check: the contributions add up to {sum(leaf_df['contribution']):.4f}, "
                   f"which equals the expected value {exp_val:.4f} (small rounding differences are normal).")
        show_defs(EXPECTIMAX_LEAF_COLS, "Column")
    with st.expander("All tree positions (node table)"):
        full_width_table(node_df)
        show_defs(EXPECTIMAX_NODE_COLS, "Column")

    st.divider()
    strategy_lines = "\n".join(f"At {u}: choose {v} (expected value {fmt(G.nodes[v]['value'])})"
                               for u, v in decision_edges)
    summary = (
        f"Expected value at root: {exp_val:.4f}\n"
        f"Recommended first move: {best_first_move}\n"
        f"Outcome range under recommended strategy: {fmt(worst_outcome)} to {fmt(best_outcome)}\n"
        f"Leaves evaluated: {len(leaves)}"
    )
    sections = [
        sec_text("Overview", "Expectimax finds the decisions that give MAX the highest average "
                             "(expected) score when some outcomes are random rather than chosen "
                             "by an opponent."),
        sec_text("Settings used", f"Tree depth: {depth}\nBranching factor: {branching}\n"
                                  f"Random seed: {seed}\n"
                                  f"Leaf data source: {'Uploaded CSV' if leaf_values else 'Synthetic (plus any manual edits)'}"),
        sec_defs("Input fields explained", EXPECTIMAX_INPUT_DEFS),
        sec_text("Results summary", summary),
        sec_text("Recommended strategy", strategy_lines or "No decisions (tree has no MAX nodes)."),
        sec_defs("Results explained", EXPECTIMAX_RESULT_DEFS),
        sec_image("Decision tree", tree_png, "Blue = MAX, amber = CHANCE, green circles = leaves, "
                                             "green lines = recommended decisions."),
        sec_image("Expected value of each first move", moves_png, "Green = recommended first move."),
        sec_table("Leaf outcomes", leaf_df),
        sec_defs("Leaf outcome columns", EXPECTIMAX_LEAF_COLS),
        sec_table("Node table", node_df),
        sec_defs("Node table columns", EXPECTIMAX_NODE_COLS),
    ]
    render_export_bar("expectimax_results", "Kalsnet AI Reasoning Studio - Expectimax",
                      sections, leaf_df)


# ============================================================================
# MODULE 3: 3-SAT (Logical Constraint Reasoning)
# ============================================================================
BRUTE_FORCE_MAX_VARS = 20  # 2^20 ~ 1M assignments; beyond this use WalkSAT


def generate_3sat(n_vars: int, m_clauses: int, seed: int = 42):
    rng = random.Random(int(seed))
    clauses = []
    for _ in range(m_clauses):
        vs = rng.sample(range(1, n_vars + 1), 3)
        clauses.append([v if rng.random() > 0.5 else -v for v in vs])
    return clauses


def lit_true(lit, assignment):
    return assignment[abs(lit)] if lit > 0 else not assignment[abs(lit)]


def is_satisfied(clause, assignment):
    return any(lit_true(lit, assignment) for lit in clause)


def count_satisfied(clauses, assignment):
    return sum(is_satisfied(c, assignment) for c in clauses)


def lit_text(lit):
    return f"x{lit}" if lit > 0 else f"NOT x{-lit}"


def brute_force_sat(clauses, n_vars):
    var_ids = list(range(1, n_vars + 1))
    best_assignment, best_score = {v: False for v in var_ids}, -1
    tried = 0
    for bits in itertools.product([False, True], repeat=n_vars):
        tried += 1
        assignment = dict(zip(var_ids, bits))
        score = count_satisfied(clauses, assignment)
        if score == len(clauses):
            return assignment, True, tried
        if score > best_score:
            best_assignment, best_score = assignment, score
    return best_assignment, False, tried  # best possible assignment, provably UNSAT


def walksat(clauses, n_vars, max_flips=500, p=0.4, seed=42):
    rng = random.Random(int(seed))
    assignment = {v: rng.choice([True, False]) for v in range(1, n_vars + 1)}
    best_assignment, best_count = dict(assignment), -1
    history = []
    for _ in range(max_flips):
        sat_count = count_satisfied(clauses, assignment)
        history.append(sat_count)
        if sat_count > best_count:
            best_assignment, best_count = dict(assignment), sat_count
        if sat_count == len(clauses):
            return assignment, history, True
        unsatisfied = [c for c in clauses if not is_satisfied(c, assignment)]
        clause = rng.choice(unsatisfied)
        if rng.random() < p:
            var = abs(rng.choice(clause))
        else:
            best_var, best_score = abs(clause[0]), -1
            for lit in clause:
                v = abs(lit)
                assignment[v] = not assignment[v]
                score = count_satisfied(clauses, assignment)
                assignment[v] = not assignment[v]
                if score > best_score:
                    best_score, best_var = score, v
            var = best_var
        assignment[var] = not assignment[var]
    sat_count = count_satisfied(clauses, assignment)
    history.append(sat_count)
    if sat_count >= best_count:
        best_assignment, best_count = dict(assignment), sat_count
    return best_assignment, history, best_count == len(clauses)


def read_clause_csv(upload):
    """Reads a lit1,lit2,lit3 CSV safely; returns (clauses, n_vars) or (None, None)."""
    try:
        real_df = pd.read_csv(upload)
        real_df.columns = [str(c).strip().lower() for c in real_df.columns]
        missing = [c for c in ["lit1", "lit2", "lit3"] if c not in real_df.columns]
        if missing:
            st.error(f"The CSV is missing column(s): {', '.join(missing)}. "
                     "Download the template below to see the exact format.")
            return None, None
        lits = real_df[["lit1", "lit2", "lit3"]].apply(pd.to_numeric, errors="coerce").dropna()
        clauses = [[int(x) for x in row] for row in lits.values.tolist()]
        clauses = [c for c in clauses if all(x != 0 for x in c)]
        if not clauses:
            st.error("No valid clauses found. Literals must be non-zero whole numbers.")
            return None, None
        n_vars = max(abs(x) for c in clauses for x in c)
        st.success(f"Loaded {len(clauses)} real clauses over {n_vars} variables.")
        return clauses, n_vars
    except Exception as e:
        st.error(f"Could not read the CSV file: {e}")
        return None, None


SAT_UPLOAD_SCHEMA = [
    ("lit1", "Required. Non-zero whole number. Positive n means 'xn is True', negative -n means "
             "'xn is False'. Example: 3"),
    ("lit2", "Required. Same rules as lit1. Example: -5"),
    ("lit3", "Required. Same rules as lit1. Example: 7. The row 3, -5, 7 means the rule "
             "(x3 OR NOT x5 OR x7): at least one of the three must hold."),
]

SAT_INPUT_DEFS = [
    ("Number of variables", "How many yes/no decisions (x1, x2, ...) the problem has, e.g. "
                            "'is machine 3 switched on?'. There are 2^n possible combinations."),
    ("Clause/variable ratio", "How many rules (clauses) there are per variable. Low ratios are "
                              "easy to satisfy, high ratios are usually impossible. The hardest "
                              "problems sit near 4.27."),
    ("Random seed", "Fixes the synthetic rules so results are repeatable. Change it for a new problem."),
    ("Upload clause CSV", "Optional. Replaces the synthetic rules with your own (see file format)."),
    ("Solver", "Auto = exact brute force for up to 14 variables, otherwise WalkSAT. Brute force "
               "checks every combination and gives a definite answer. WalkSAT is a fast search "
               "that usually finds a solution but cannot prove that none exists."),
]

SAT_RESULT_DEFS = [
    ("Clause", "One rule made of three conditions joined by OR. It is satisfied if at least one "
               "of its three conditions is true."),
    ("Result: SATISFIABLE", "A setting of all variables was found that satisfies every rule at once."),
    ("Result: UNSATISFIABLE (proven)", "Brute force checked every combination: no setting can "
                                       "satisfy all rules. The best partial setting is shown."),
    ("Result: Best effort", "WalkSAT ran out of steps before satisfying every rule. A solution may "
                            "still exist; try brute force (up to 20 variables) or another seed."),
    ("Clauses satisfied", "How many rules the chosen setting satisfies, out of the total."),
    ("Combinations checked / flips", "Work done: brute force counts combinations tried, WalkSAT "
                                     "counts single-variable changes (flips)."),
    ("WalkSAT search progress", "Chart of how many rules were satisfied after each flip. The "
                                "dashed line is the goal (all rules satisfied)."),
]

SAT_ASSIGN_COLS = [
    ("variable", "Variable name (x1, x2, ...)."),
    ("value", "True or False in the best setting found."),
    ("appears_in_clauses", "How many rules mention this variable."),
]

SAT_CLAUSE_COLS = [
    ("clause", "Rule number (1 = first rule)."),
    ("lit1 / lit2 / lit3", "The three conditions as signed numbers (negative = NOT)."),
    ("rule", "The same rule written in words, e.g. x3 OR NOT x5 OR x7."),
    ("satisfied", "Yes if the best setting found makes this rule true."),
    ("satisfied_by", "Which condition(s) make the rule true. Blank if the rule is violated."),
]


def render_3sat_tab():
    st.markdown('<div class="module-banner banner-3sat">3-SAT - Logical Constraint Reasoning</div>',
                unsafe_allow_html=True)
    st.write(
        "Boolean satisfiability with rules (clauses) of exactly 3 conditions each. It asks: "
        "is there a yes/no setting for every variable that satisfies all rules at once? This "
        "models hard logical-constraint problems in industry such as scheduling, configuration "
        "validation, or checking whether design and process rules can all be met together."
    )

    left, right = st.columns([1, 2])
    with left:
        st.markdown("**Settings**")
        n_vars = st.slider("Number of variables", 4, 16, 8, key="sat_nvars", help=SAT_INPUT_DEFS[0][1])
        ratio = st.slider("Clause/variable ratio", 1.0, 6.0, 4.2, step=0.1, key="sat_ratio",
                          help=SAT_INPUT_DEFS[1][1])
        m_clauses = max(1, int(round(n_vars * ratio)))
        seed = int(st.number_input("Random seed", value=3, step=1, key="sat_seed",
                                   help=SAT_INPUT_DEFS[2][1]))
        st.caption(f"Generating {m_clauses} random 3-condition rules over {n_vars} variables "
                   f"(ratio {ratio:.1f}; the hardest problems are near 4.27).")

        st.markdown("**Use your own data (optional)**")
        upload = st.file_uploader("Upload clause CSV", type=["csv"], key="sat_upload",
                                  help=SAT_INPUT_DEFS[3][1])
        with st.expander("CSV file format"):
            show_defs(SAT_UPLOAD_SCHEMA, "Column", "Rules and example")

        clauses = None
        source = "Synthetic"
        if upload is not None:
            clauses, uploaded_vars = read_clause_csv(upload)
            if clauses is not None:
                n_vars = uploaded_vars
                source = "Uploaded CSV"
        if clauses is None:
            clauses = generate_3sat(n_vars, m_clauses, seed=seed)

        template = pd.DataFrame({"lit1": [1, -2, 3], "lit2": [-3, 4, -5], "lit3": [5, -1, 2]})
        st.download_button("Download CSV template", data=export_csv_bytes(template),
                           file_name="3sat_clause_template.csv", mime="text/csv", key="sat_template")

        solver_choice = st.radio("Solver", ["Auto", "Brute force (exact)", "WalkSAT (local search)"],
                                 key="sat_solver", help=SAT_INPUT_DEFS[4][1])

    use_brute = (solver_choice == "Brute force (exact)" and n_vars <= BRUTE_FORCE_MAX_VARS) or \
                (solver_choice == "Auto" and n_vars <= 14)
    if solver_choice == "Brute force (exact)" and n_vars > BRUTE_FORCE_MAX_VARS:
        st.warning(f"Brute force is limited to {BRUTE_FORCE_MAX_VARS} variables; using WalkSAT instead.")

    history = None
    if use_brute:
        assignment, satisfiable, tried = brute_force_sat(clauses, n_vars)
        method = "Brute force (exact)"
        work = f"{tried:,} of {2 ** n_vars:,} combinations"
    else:
        assignment, history, satisfiable = walksat(clauses, n_vars, max_flips=500, seed=seed)
        method = "WalkSAT (local search)"
        work = f"{len(history) - 1} flips"

    sat_count = count_satisfied(clauses, assignment)
    sat_flags = [is_satisfied(c, assignment) for c in clauses]

    if satisfiable:
        verdict, vcolor = "SATISFIABLE", ACCENT_GREEN
    elif use_brute:
        verdict, vcolor = "UNSATISFIABLE (proven)", ACCENT_RED
    else:
        verdict, vcolor = "Best effort", ACCENT_AMBER

    with right:
        st.markdown("**Clause satisfaction** - green = rule satisfied, red = rule violated "
                    "(using the best setting found)")
        fig, ax = plt.subplots(figsize=(9, 3.0))
        colors = [ACCENT_GREEN if s else ACCENT_RED for s in sat_flags]
        ax.bar(range(1, len(clauses) + 1), [1] * len(clauses), color=colors)
        ax.set_xlabel("Clause number")
        ax.set_yticks([])
        fig.tight_layout()
        clause_png = show_fig(fig)

        progress_png = None
        if history:
            st.markdown("**WalkSAT search progress** - rules satisfied after each flip")
            fig3, ax3 = plt.subplots(figsize=(9, 2.6))
            ax3.plot(history, color=ACCENT_PURPLE, linewidth=2)
            ax3.axhline(len(clauses), color=ACCENT_GREEN, linestyle="--", linewidth=1,
                        label="Goal: all rules satisfied")
            ax3.set_xlabel("Flip number")
            ax3.set_ylabel("Rules satisfied")
            ax3.legend(fontsize=7)
            fig3.tight_layout()
            progress_png = show_fig(fig3)

    r1 = st.columns(4)
    metric_card(r1[0], "Result", verdict, vcolor, "Can every rule be met at the same time?", size=20)
    metric_card(r1[1], "Clauses satisfied", f"{sat_count}/{len(clauses)}", ACCENT_PURPLE,
                "Rules met by the best setting found.")
    metric_card(r1[2], "Method", method, ACCENT_PURPLE,
                "Exact check of every combination, or fast local search.", size=16)
    metric_card(r1[3], "Work done", work, ACCENT_PURPLE,
                "Combinations tried (brute force) or variable flips (WalkSAT).", size=16)

    # --- Result tables ---
    usage = {v: 0 for v in assignment}
    for c in clauses:
        for lit in c:
            usage[abs(lit)] = usage.get(abs(lit), 0) + 1
    assign_df = pd.DataFrame({
        "variable": [f"x{v}" for v in assignment],
        "value": [bool(assignment[v]) for v in assignment],
        "appears_in_clauses": [usage.get(v, 0) for v in assignment],
    })
    clause_df = pd.DataFrame([{
        "clause": i + 1,
        "lit1": c[0], "lit2": c[1], "lit3": c[2],
        "rule": " OR ".join(lit_text(l) for l in c),
        "satisfied": "Yes" if ok else "No",
        "satisfied_by": ", ".join(lit_text(l) for l in c if lit_true(l, assignment)),
    } for i, (c, ok) in enumerate(zip(clauses, sat_flags))])

    with st.expander("How to read these results", expanded=False):
        show_defs(SAT_RESULT_DEFS)
    with st.expander("Variable settings"):
        full_width_table(assign_df)
        show_defs(SAT_ASSIGN_COLS, "Column")
    with st.expander("All rules (clause table)"):
        full_width_table(clause_df)
        show_defs(SAT_CLAUSE_COLS, "Column")

    st.divider()
    violated = [str(i + 1) for i, ok in enumerate(sat_flags) if not ok]
    summary = (
        f"Result: {verdict}\n"
        f"Clauses satisfied: {sat_count} of {len(clauses)}\n"
        f"Violated clauses: {', '.join(violated) if violated else 'none'}\n"
        f"Method: {method}  |  Work done: {work}\n"
        f"True variables: {', '.join(f'x{v}' for v in assignment if assignment[v]) or 'none'}"
    )
    sections = [
        sec_text("Overview", "3-SAT checks whether one yes/no setting of all variables can satisfy "
                             "every rule at the same time, where each rule is three conditions "
                             "joined by OR."),
        sec_text("Settings used", f"Variables: {n_vars}\nClauses: {len(clauses)}\n"
                                  f"Clause/variable ratio: {len(clauses) / n_vars:.2f}\n"
                                  f"Random seed: {seed}\nSolver choice: {solver_choice}\n"
                                  f"Data source: {source}"),
        sec_defs("Input fields explained", SAT_INPUT_DEFS),
        sec_text("Results summary", summary),
        sec_defs("Results explained", SAT_RESULT_DEFS),
        sec_image("Clause satisfaction", clause_png, "Green = satisfied, red = violated."),
    ]
    if progress_png:
        sections.append(sec_image("WalkSAT search progress", progress_png,
                                  "Rules satisfied after each flip; dashed line = goal."))
    sections += [
        sec_table("Variable settings", assign_df),
        sec_defs("Variable settings columns", SAT_ASSIGN_COLS),
        sec_table("Clause table", clause_df),
        sec_defs("Clause table columns", SAT_CLAUSE_COLS),
    ]
    render_export_bar("3sat_results", "Kalsnet AI Reasoning Studio - 3-SAT", sections, clause_df)


# ============================================================================
# ABOUT TAB
# ============================================================================

def render_about_tab():
    st.markdown('<div class="module-banner banner-about">About this Studio</div>',
                unsafe_allow_html=True)
    st.markdown(
        """
**Kalsnet AI Reasoning Studio** packages three foundational reasoning
paradigms as interactive modules, intended as building blocks for
Kalsnet's AI + Knowledge Graph applications:

| Module | Reasoning type | Typical Kalsnet use case |
|---|---|---|
| **Minimax with Alpha-Beta Pruning** | Adversarial reasoning | Competitive negotiation, security red-team/blue-team simulation |
| **Expectimax** | Probabilistic reasoning | Planning under uncertain demand, equipment failure risk |
| **3-SAT** | Logical constraint reasoning | Validating that a set of process/design rules can be jointly satisfied |

Every module:
- Ships with **synthetic data** generated on the fly, so it works immediately.
- Can be **overridden with real data** via CSV upload or in-app editing.
- Explains **every input field, result and table column** in plain language
  (hover over the small help icons, or open the "How to read these results" sections).
- Renders **charts** (game trees, comparison charts, search-progress charts).
- **Exports all results** - summary, explanations, tables and charts - as PDF, Word,
  CSV (full report or data only) and plain text.

---
**Developed by Randy Singh, Kalsnet (KNet) Consulting Group.**
"""
    )


# ============================================================================
# MAIN - TABS
# ============================================================================

tab1, tab2, tab3, tab4 = st.tabs(["Minimax + Alpha-Beta", "Expectimax", "3-SAT", "About"])

with tab1:
    render_minimax_tab()

with tab2:
    render_expectimax_tab()

with tab3:
    render_3sat_tab()

with tab4:
    render_about_tab()