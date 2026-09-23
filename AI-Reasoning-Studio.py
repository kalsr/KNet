# Kalsnet AI Reasoning Studio
# ============================
# Three self-contained reasoning modules for the Kalsnet AI + Knowledge Graph
# platform, each demonstrating a distinct class of reasoning used across
# industrial AI decision systems:
#
#   1. Minimax     -> Adversarial reasoning        (game-tree search)
#   2. Expectimax  -> Probabilistic reasoning      (decisions under uncertainty)
#   3. 3-SAT       -> Logical constraint reasoning (satisfiability)
#
# Every module ships with generated synthetic data so it works out of the box,
# but every module also accepts real data (CSV / typed input) to override the
# synthetic values. Results in every tab can be exported as PDF, Word (.docx),
# CSV, or plain text.
#
# Developed by Randy Singh - Kalsnet (KNet) Consulting Group

import inspect
import io
import itertools
import random
from datetime import datetime

import matplotlib

matplotlib.use("Agg")  # headless backend for Streamlit Cloud servers

import matplotlib.pyplot as plt
import networkx as nx
import numpy as np  # noqa: F401  (kept for future numeric extensions)
import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import RGBColor

# ============================================================================
# PDF LIBRARY IMPORT (works with fpdf2 AND the legacy "fpdf" package)
# ============================================================================
try:
    from fpdf import FPDF, XPos, YPos  # fpdf2 (preferred)

    FPDF_MODERN = True
except ImportError:  # legacy PyFPDF is installed instead of fpdf2
    from enum import Enum

    from fpdf import FPDF as _LegacyFPDF

    FPDF_MODERN = False

    class XPos(str, Enum):
        LEFT = "LEFT"
        RIGHT = "RIGHT"
        START = "START"
        END = "END"
        WCONT = "WCONT"
        CENTER = "CENTER"
        LMARGIN = "LMARGIN"
        RMARGIN = "RMARGIN"

    class YPos(str, Enum):
        TOP = "TOP"
        LAST = "LAST"
        NEXT = "NEXT"
        TMARGIN = "TMARGIN"
        BMARGIN = "BMARGIN"

    class FPDF(_LegacyFPDF):
        """Lets the legacy library accept fpdf2-style arguments."""

        def cell(self, w=0, h=0, txt="", border=0, ln=0, align="", fill=False,
                 link="", text=None, new_x=None, new_y=None, **kwargs):
            if text is not None:
                txt = text
            if new_y == YPos.NEXT:
                ln = 1 if new_x in (None, XPos.LMARGIN) else 2
            return super().cell(w, h, str(txt), border, ln, align, fill, link)

        def multi_cell(self, w, h=0, txt="", border=0, align="J", fill=False,
                       text=None, new_x=None, new_y=None, **kwargs):
            if text is not None:
                txt = text
            return super().multi_cell(w, h, str(txt), border, align, fill)

        def output(self, name="", dest=""):
            if not name and not dest:
                return super().output(dest="S").encode("latin-1")
            return super().output(name, dest)


# ============================================================================
# STREAMLIT VERSION COMPATIBILITY
# ============================================================================
# Newer Streamlit replaced use_container_width=True with width="stretch".
_PYPLOT_HAS_WIDTH = "width" in inspect.signature(st.pyplot).parameters


def show_fig(fig):
    """Render a matplotlib figure full-width, then free its memory."""
    if _PYPLOT_HAS_WIDTH:
        st.pyplot(fig, width="stretch")
    else:
        st.pyplot(fig, use_container_width=True)
    plt.close(fig)


def full_width_editor(df, key):
    try:
        return st.data_editor(df, width="stretch", key=key, num_rows="fixed")
    except Exception:  # older Streamlit versions only accept use_container_width
        return st.data_editor(df, use_container_width=True, key=key, num_rows="fixed")


# ============================================================================
# PAGE CONFIG & GLOBAL STYLE
# ============================================================================
st.set_page_config(
    page_title="Kalsnet AI Reasoning Studio",
    layout="wide",
    initial_sidebar_state="expanded",
)

PRIMARY_BLUE = "#1a3d6d"
ACCENT_BLUE = "#1a73e8"
ACCENT_RED = "#d93025"
ACCENT_AMBER = "#f9ab00"
ACCENT_GREEN = "#1e8e3e"
ACCENT_PURPLE = "#8430ce"

# Plain string (not an f-string) so CSS braces never break Python parsing.
CSS = """
<style>
.stApp {
    background: linear-gradient(180deg, #f4f7fb 0%, #eef3fa 100%);
}
.title-block {
    text-align: center;
    padding: 18px 10px 8px 10px;
}
.title-main {
    font-family: 'Segoe UI', Arial, sans-serif;
    font-weight: 800;
    color: PRIMARY_BLUE;
    font-size: 42px;
    line-height: 1.25;
    margin: 0;
}
.title-sub {
    font-family: 'Segoe UI', Arial, sans-serif;
    font-weight: 800;
    color: PRIMARY_BLUE;
    font-size: 24px;
    line-height: 1.3;
    margin: 4px 0 0 0;
}
.module-banner {
    border-radius: 10px;
    padding: 14px 20px;
    margin-bottom: 14px;
    color: white;
    font-size: 20px;
    font-weight: 700;
    font-family: 'Segoe UI', Arial, sans-serif;
}
.banner-minimax { background: linear-gradient(90deg, ACCENT_BLUE, ACCENT_RED); }
.banner-expectimax { background: linear-gradient(90deg, ACCENT_AMBER, ACCENT_BLUE); }
.banner-3sat { background: linear-gradient(90deg, ACCENT_PURPLE, ACCENT_GREEN); }
.banner-about { background: linear-gradient(90deg, PRIMARY_BLUE, ACCENT_BLUE); }
.metric-card {
    background: white;
    border-radius: 10px;
    padding: 14px 18px;
    box-shadow: 0 1px 4px rgba(0,0,0,0.08);
    border-left: 5px solid ACCENT_BLUE;
}
.stTabs [data-baseweb="tab-list"] { gap: 6px; }
.stTabs [data-baseweb="tab"] {
    background-color: #e4ecf7;
    border-radius: 8px 8px 0 0;
    padding: 10px 16px;
    font-weight: 600;
}
.stTabs [aria-selected="true"] {
    background-color: PRIMARY_BLUE !important;
    color: white !important;
}
</style>
"""
for _name, _value in {
    "PRIMARY_BLUE": PRIMARY_BLUE, "ACCENT_BLUE": ACCENT_BLUE, "ACCENT_RED": ACCENT_RED,
    "ACCENT_AMBER": ACCENT_AMBER, "ACCENT_GREEN": ACCENT_GREEN, "ACCENT_PURPLE": ACCENT_PURPLE,
}.items():
    CSS = CSS.replace(_name, _value)
st.markdown(CSS, unsafe_allow_html=True)

st.markdown(
    """
    <div class="title-block">
        <div class="title-main">Kalsnet AI Reasoning Studio</div>
        <div class="title-sub">Developed by Randy Singh - Kalsnet (KNet) Consulting Group</div>
    </div>
    """,
    unsafe_allow_html=True,
)
st.caption(
    "Adversarial reasoning (Minimax) | Probabilistic reasoning (Expectimax) | "
    "Logical constraint reasoning (3-SAT) - building blocks for the Kalsnet "
    "AI + Knowledge Graph platform."
)
st.divider()

# ============================================================================
# SHARED EXPORT HELPERS
# ============================================================================
_PDF_REPLACEMENTS = {
    "\u2014": "-", "\u2013": "-", "\u2018": "'", "\u2019": "'", "\u201c": '"',
    "\u201d": '"', "\u2026": "...", "\u2022": "-", "\u00a0": " ", "\u2192": "->",
}


def pdf_safe(text) -> str:
    """Built-in PDF fonts only support Latin-1; convert everything else."""
    text = str(text)
    for k, v in _PDF_REPLACEMENTS.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def export_txt(content: str) -> bytes:
    return content.encode("utf-8")


def export_csv_bytes(df: pd.DataFrame) -> bytes:
    buf = io.StringIO()
    df.to_csv(buf, index=False)
    return buf.getvalue().encode("utf-8")


def export_docx_bytes(title: str, sections: list) -> bytes:
    doc = Document()
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x3D, 0x6D)
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph("Kalsnet (KNet) Consulting Group - Randy Singh")
    for heading, body in sections:
        doc.add_heading(heading, level=2)
        if isinstance(body, pd.DataFrame):
            table = doc.add_table(rows=1, cols=max(len(body.columns), 1))
            try:
                table.style = "Light Grid Accent 1"
            except Exception:
                table.style = "Table Grid"
            for i, col in enumerate(body.columns):
                table.rows[0].cells[i].text = str(col)
            for _, row in body.iterrows():
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = str(val)
        else:
            doc.add_paragraph(str(body))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_paragraph(pdf, height, text):
    """Full-width paragraph that always returns the cursor to the left margin."""
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, height, pdf_safe(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)


def export_pdf_bytes(title: str, sections: list) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    usable_w = pdf.w - pdf.l_margin - pdf.r_margin

    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(26, 61, 109)
    _pdf_paragraph(pdf, 10, title)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(90, 90, 90)
    _pdf_paragraph(pdf, 6, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  "
                           "Kalsnet (KNet) Consulting Group")
    pdf.ln(3)

    for heading, body in sections:
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(26, 61, 109)
        _pdf_paragraph(pdf, 8, heading)
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(30, 30, 30)
        if isinstance(body, pd.DataFrame):
            col_w = usable_w / max(len(body.columns), 1)
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", "B", 9)
            for col in body.columns:
                pdf.cell(col_w, 7, pdf_safe(str(col)[:20]), border=1)
            pdf.ln()
            pdf.set_font("Helvetica", "", 9)
            for _, row in body.iterrows():
                pdf.set_x(pdf.l_margin)
                for val in row:
                    pdf.cell(col_w, 7, pdf_safe(str(val)[:20]), border=1)
                pdf.ln()
        else:
            _pdf_paragraph(pdf, 6, body)
        pdf.ln(2)
    return bytes(pdf.output())


def render_export_bar(key_prefix: str, title: str, sections: list, df_for_csv: pd.DataFrame):
    st.markdown("##### Export these results")
    c1, c2, c3, c4 = st.columns(4)
    text_blob = "\n\n".join(
        f"{h}\n{'-' * len(h)}\n{(b.to_string(index=False) if isinstance(b, pd.DataFrame) else b)}"
        for h, b in sections
    )
    with c1:
        try:
            st.download_button("PDF", data=export_pdf_bytes(title, sections),
                               file_name=f"{key_prefix}.pdf", mime="application/pdf",
                               key=f"{key_prefix}_pdf")
        except Exception as e:  # never let an export failure crash the app
            st.warning(f"PDF export unavailable: {e}")
    with c2:
        try:
            st.download_button("Word (.docx)", data=export_docx_bytes(title, sections),
                               file_name=f"{key_prefix}.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               key=f"{key_prefix}_docx")
        except Exception as e:
            st.warning(f"Word export unavailable: {e}")
    with c3:
        st.download_button("CSV", data=export_csv_bytes(df_for_csv),
                           file_name=f"{key_prefix}.csv", mime="text/csv",
                           key=f"{key_prefix}_csv")
    with c4:
        st.download_button("Text", data=export_txt(text_blob),
                           file_name=f"{key_prefix}.txt", mime="text/plain",
                           key=f"{key_prefix}_txt")


def read_leaf_csv(upload, n_leaves):
    """Reads a leaf_value CSV safely; returns a list of floats or None."""
    try:
        real_df = pd.read_csv(upload)
        if "leaf_value" not in real_df.columns:
            st.error("CSV must contain a column named 'leaf_value'.")
            return None
        values = pd.to_numeric(real_df["leaf_value"], errors="coerce").dropna().tolist()
        if not values:
            st.error("No numeric values found in 'leaf_value'.")
            return None
        if len(values) != n_leaves:
            st.warning(f"Expected {n_leaves} values, got {len(values)}. "
                       "Synthetic data fills any missing values; extra values are ignored.")
        return values
    except Exception as e:
        st.error(f"Could not read CSV: {e}")
        return None


# ============================================================================
# TREE UTILITIES (shared by Minimax + Expectimax)
# ============================================================================

def build_tree(depth: int, branching: int, node_types: list, leaf_values=None, seed: int = 42):
    """Builds a synthetic layered tree. node_types[level] gives the node type
    for every node at that level ('MAX', 'MIN', 'CHANCE', or 'LEAF' for the last)."""
    rng = random.Random(int(seed))
    G = nx.DiGraph()
    leaf_iter = itertools.count()

    def add_node(level, path):
        node_id = "n" + ("-".join(map(str, path)) if path else "root")
        ntype = node_types[level]
        G.add_node(node_id, type=ntype, level=level)
        if ntype == "LEAF":
            idx = next(leaf_iter)
            synthetic = float(rng.randint(-10, 10))
            if leaf_values is not None and idx < len(leaf_values):
                val = float(leaf_values[idx])
            else:
                val = synthetic
            G.nodes[node_id]["value"] = val
            return node_id
        children = [add_node(level + 1, path + [i]) for i in range(branching)]
        probs = None
        if ntype == "CHANCE":
            raw = [rng.random() + 0.1 for _ in children]
            total = sum(raw)
            probs = [round(p / total, 3) for p in raw]
            probs[-1] = round(1 - sum(probs[:-1]), 3)  # normalize rounding drift
        for j, child in enumerate(children):
            edge_kwargs = {"prob": probs[j]} if probs else {}
            G.add_edge(node_id, child, **edge_kwargs)
        return node_id

    root = add_node(0, [])
    return G, root


def layered_layout(G):
    levels = {}
    for n, d in G.nodes(data=True):
        levels.setdefault(d["level"], []).append(n)
    pos = {}
    for level, nodes in levels.items():
        count = len(nodes)
        for i, node in enumerate(nodes):
            pos[node] = (i - (count - 1) / 2, -level)
    return pos


def draw_tree(G, pos, path_edges=None, value_fmt="{:.1f}"):
    fig, ax = plt.subplots(figsize=(9, 4.5))
    color_map = {"MAX": ACCENT_BLUE, "MIN": ACCENT_RED, "CHANCE": ACCENT_AMBER, "LEAF": ACCENT_GREEN}
    colors = [color_map.get(G.nodes[n]["type"], "#888888") for n in G.nodes]
    nx.draw(G, pos, ax=ax, node_color=colors, node_size=900, arrows=False,
            edge_color="#bbbbbb", width=1.4)
    labels = {n: value_fmt.format(G.nodes[n]["value"]) if "value" in G.nodes[n] else ""
              for n in G.nodes}
    nx.draw_networkx_labels(G, pos, labels, font_size=8, font_color="white",
                            font_weight="bold", ax=ax)
    edge_labels = nx.get_edge_attributes(G, "prob")
    if edge_labels:
        nx.draw_networkx_edge_labels(G, pos, edge_labels={k: f"p={v}" for k, v in edge_labels.items()},
                                     font_size=7, ax=ax)
    if path_edges:
        nx.draw_networkx_edges(G, pos, edgelist=path_edges, edge_color=ACCENT_GREEN,
                               width=3.2, arrows=False, ax=ax)
    ax.axis("off")
    fig.tight_layout()
    return fig


def reconstruct_path(G, root):
    path = [root]
    node = root
    while "chosen" in G.nodes[node]:
        node = G.nodes[node]["chosen"]
        path.append(node)
    edges = list(zip(path[:-1], path[1:]))
    return path, edges


def apply_leaf_edits(G, leaves, key):
    """Shows an editable leaf table and writes any edits back into the tree."""
    default_leaf_df = pd.DataFrame({"leaf": leaves, "value": [G.nodes[l]["value"] for l in leaves]})
    edited = full_width_editor(default_leaf_df, key)
    for l, v in zip(edited["leaf"], edited["value"]):
        try:
            G.nodes[l]["value"] = float(v)
        except (TypeError, ValueError):
            pass  # keep previous value if the cell is blank / invalid


def metric_card(col, label, value, color, size=26):
    col.markdown(
        f'<div class="metric-card"><b>{label}</b><br>'
        f'<span style="font-size:{size}px;color:{color}">{value}</span></div>',
        unsafe_allow_html=True,
    )


# ============================================================================
# MODULE 1: MINIMAX (Adversarial Reasoning)
# ============================================================================

def minimax(G, node):
    ntype = G.nodes[node]["type"]
    if ntype == "LEAF":
        return G.nodes[node]["value"]
    vals = [(c, minimax(G, c)) for c in G.successors(node)]
    if ntype == "MAX":
        chosen, val = max(vals, key=lambda cv: cv[1])
    else:
        chosen, val = min(vals, key=lambda cv: cv[1])
    G.nodes[node]["value"] = val
    G.nodes[node]["chosen"] = chosen
    return val


def render_minimax_tab():
    st.markdown('<div class="module-banner banner-minimax">Minimax - Adversarial Reasoning</div>',
                unsafe_allow_html=True)
    st.write(
        "Two opposing agents (MAX vs. MIN) alternate turns down a game tree. "
        "MAX always picks the branch with the highest guaranteed value; MIN always "
        "picks the branch that is worst for MAX. This models zero-sum adversarial "
        "decisions - e.g. competitive bidding, negotiation, or an attacker/defender "
        "scenario on an industrial network."
    )

    left, right = st.columns([1, 2])
    with left:
        st.markdown("**Synthetic data controls**")
        depth = st.select_slider("Tree depth (plies)", options=[2, 3, 4], value=3, key="mm_depth")
        branching = st.select_slider("Branching factor", options=[2, 3], value=2, key="mm_branch")
        seed = int(st.number_input("Random seed", value=7, step=1, key="mm_seed"))

        node_types = ["MAX" if lvl % 2 == 0 else "MIN" for lvl in range(depth)] + ["LEAF"]
        n_leaves = branching ** depth

        st.markdown("**Or upload real leaf values (CSV)**")
        st.caption(f"CSV needs one column `leaf_value` with exactly {n_leaves} numeric rows "
                   "for the current tree shape.")
        upload = st.file_uploader("Upload leaf-value CSV", type=["csv"], key="mm_upload")
        leaf_values = read_leaf_csv(upload, n_leaves) if upload is not None else None

        template = pd.DataFrame({"leaf_value": [0] * n_leaves})
        st.download_button("Download CSV template", data=export_csv_bytes(template),
                           file_name="minimax_leaf_template.csv", mime="text/csv", key="mm_template")

    G, root = build_tree(depth, branching, node_types, leaf_values=leaf_values, seed=seed)
    leaves = [n for n, d in G.nodes(data=True) if d["type"] == "LEAF"]

    st.markdown("**Or hand-edit the leaves directly**")
    apply_leaf_edits(G, leaves, key=f"mm_editor_{depth}_{branching}_{seed}")

    best_val = minimax(G, root)
    path, path_edges = reconstruct_path(G, root)

    with right:
        st.markdown("**Game tree** (blue = MAX, red = MIN, green = optimal path)")
        show_fig(draw_tree(G, layered_layout(G), path_edges=path_edges))

    m1, m2, m3 = st.columns(3)
    metric_card(m1, "Minimax value", f"{best_val:.2f}", ACCENT_BLUE)
    metric_card(m2, "Optimal path length", f"{len(path)} nodes", ACCENT_BLUE)
    metric_card(m3, "Leaves evaluated", f"{len(leaves)}", ACCENT_BLUE)

    st.markdown("**Leaf value distribution**")
    fig2, ax2 = plt.subplots(figsize=(9, 2.8))
    leaf_vals_now = [G.nodes[l]["value"] for l in leaves]
    bar_colors = [ACCENT_GREEN if l in path else "#a9b8c8" for l in leaves]
    ax2.bar(range(len(leaves)), leaf_vals_now, color=bar_colors)
    ax2.set_xticks(range(len(leaves)))
    ax2.set_xticklabels([l[1:] for l in leaves], rotation=45, fontsize=7)
    ax2.set_ylabel("Utility")
    fig2.tight_layout()
    show_fig(fig2)

    st.divider()
    result_df = pd.DataFrame({"leaf": leaves, "value": leaf_vals_now,
                              "on_optimal_path": [l in path for l in leaves]})
    sections = [
        ("Configuration", f"Depth={depth}, Branching={branching}, Seed={seed}"),
        ("Minimax Result", f"Optimal (minimax) value = {best_val:.2f}\n"
                           f"Optimal path: {' -> '.join(path)}"),
        ("Leaf Values", result_df),
    ]
    render_export_bar("minimax_results", "Kalsnet AI Reasoning Studio - Minimax Results",
                      sections, result_df)


# ============================================================================
# MODULE 2: EXPECTIMAX (Probabilistic Reasoning)
# ============================================================================

def expectimax(G, node):
    ntype = G.nodes[node]["type"]
    if ntype == "LEAF":
        return G.nodes[node]["value"]
    children = list(G.successors(node))
    if ntype == "MAX":
        vals = [(c, expectimax(G, c)) for c in children]
        chosen, val = max(vals, key=lambda cv: cv[1])
        G.nodes[node]["chosen"] = chosen
    else:  # CHANCE
        val = 0.0
        for c in children:
            p = G.edges[node, c].get("prob", 1 / len(children))
            val += p * expectimax(G, c)
    G.nodes[node]["value"] = val
    return val


def render_expectimax_tab():
    st.markdown('<div class="module-banner banner-expectimax">Expectimax - Probabilistic Reasoning</div>',
                unsafe_allow_html=True)
    st.write(
        "MAX nodes still pick the best option, but MIN is replaced by CHANCE nodes "
        "whose outcomes follow a probability distribution (e.g. equipment failure "
        "rates, demand uncertainty, sensor noise). The value of a chance node is the "
        "probability-weighted expectation of its children - this models planning "
        "under uncertainty rather than against an adversary."
    )

    left, right = st.columns([1, 2])
    with left:
        st.markdown("**Synthetic data controls**")
        depth = st.select_slider("Tree depth (plies)", options=[2, 3, 4], value=3, key="em_depth")
        branching = st.select_slider("Branching factor", options=[2, 3], value=2, key="em_branch")
        seed = int(st.number_input("Random seed", value=11, step=1, key="em_seed"))

        node_types = ["MAX" if lvl % 2 == 0 else "CHANCE" for lvl in range(depth)] + ["LEAF"]
        n_leaves = branching ** depth

        st.markdown("**Or upload real leaf values (CSV)**")
        st.caption(f"CSV needs one column `leaf_value` with exactly {n_leaves} numeric rows "
                   "for the current tree shape.")
        upload = st.file_uploader("Upload leaf-value CSV", type=["csv"], key="em_upload")
        leaf_values = read_leaf_csv(upload, n_leaves) if upload is not None else None

        template = pd.DataFrame({"leaf_value": [0] * n_leaves})
        st.download_button("Download CSV template", data=export_csv_bytes(template),
                           file_name="expectimax_leaf_template.csv", mime="text/csv", key="em_template")

    G, root = build_tree(depth, branching, node_types, leaf_values=leaf_values, seed=seed)
    leaves = [n for n, d in G.nodes(data=True) if d["type"] == "LEAF"]

    st.markdown("**Or hand-edit the leaves directly**")
    apply_leaf_edits(G, leaves, key=f"em_editor_{depth}_{branching}_{seed}")

    exp_val = expectimax(G, root)
    path, path_edges = reconstruct_path(G, root)

    with right:
        st.markdown("**Decision tree** (blue = MAX, amber = CHANCE with edge probabilities, "
                    "green = recommended decisions)")
        show_fig(draw_tree(G, layered_layout(G), path_edges=path_edges, value_fmt="{:.2f}"))

    root_children = list(G.successors(root))
    best_first_move = G.nodes[root].get("chosen", root)

    m1, m2, m3 = st.columns(3)
    metric_card(m1, "Expected value", f"{exp_val:.2f}", ACCENT_AMBER)
    metric_card(m2, "Recommended first move", best_first_move, ACCENT_AMBER)
    metric_card(m3, "Leaves evaluated", f"{len(leaves)}", ACCENT_AMBER)

    st.markdown("**Expected value of each first move**")
    fig2, ax2 = plt.subplots(figsize=(9, 2.8))
    child_vals = [G.nodes[c]["value"] for c in root_children]
    colors = [ACCENT_GREEN if c == best_first_move else ACCENT_AMBER for c in root_children]
    ax2.bar(root_children, child_vals, color=colors)
    ax2.set_ylabel("Expected value")
    fig2.tight_layout()
    show_fig(fig2)

    st.divider()
    leaf_vals_now = [G.nodes[l]["value"] for l in leaves]
    result_df = pd.DataFrame({"leaf": leaves, "value": leaf_vals_now})
    sections = [
        ("Configuration", f"Depth={depth}, Branching={branching}, Seed={seed}"),
        ("Expectimax Result", f"Expected value at root = {exp_val:.2f}\n"
                              f"Recommended first move: {best_first_move}"),
        ("Leaf Values", result_df),
    ]
    render_export_bar("expectimax_results", "Kalsnet AI Reasoning Studio - Expectimax Results",
                      sections, result_df)


# ============================================================================
# MODULE 3: 3-SAT (Logical Constraint Reasoning)
# ============================================================================
BRUTE_FORCE_MAX_VARS = 20  # 2^20 ~ 1M assignments; beyond this use WalkSAT


def generate_3sat(n_vars: int, m_clauses: int, seed: int = 42):
    rng = random.Random(int(seed))
    clauses = []
    for _ in range(m_clauses):
        vs = rng.sample(range(1, n_vars + 1), 3)
        clauses.append([v if rng.random() > 0.5 else -v for v in vs])
    return clauses


def is_satisfied(clause, assignment):
    return any((lit > 0 and assignment[abs(lit)]) or (lit < 0 and not assignment[abs(lit)])
               for lit in clause)


def count_satisfied(clauses, assignment):
    return sum(is_satisfied(c, assignment) for c in clauses)


def brute_force_sat(clauses, n_vars):
    var_ids = list(range(1, n_vars + 1))
    best_assignment, best_score = {v: False for v in var_ids}, -1
    for bits in itertools.product([False, True], repeat=n_vars):
        assignment = dict(zip(var_ids, bits))
        score = count_satisfied(clauses, assignment)
        if score == len(clauses):
            return assignment, True
        if score > best_score:
            best_assignment, best_score = assignment, score
    return best_assignment, False  # best possible assignment, provably UNSAT


def walksat(clauses, n_vars, max_flips=500, p=0.4, seed=42):
    rng = random.Random(int(seed))
    assignment = {v: rng.choice([True, False]) for v in range(1, n_vars + 1)}
    best_assignment, best_count = dict(assignment), -1
    history = []
    for _ in range(max_flips):
        sat_count = count_satisfied(clauses, assignment)
        history.append(sat_count)
        if sat_count > best_count:
            best_assignment, best_count = dict(assignment), sat_count
        if sat_count == len(clauses):
            return assignment, history, True
        unsatisfied = [c for c in clauses if not is_satisfied(c, assignment)]
        clause = rng.choice(unsatisfied)
        if rng.random() < p:
            var = abs(rng.choice(clause))
        else:
            best_var, best_score = abs(clause[0]), -1
            for lit in clause:
                v = abs(lit)
                assignment[v] = not assignment[v]
                score = count_satisfied(clauses, assignment)
                assignment[v] = not assignment[v]
                if score > best_score:
                    best_score, best_var = score, v
            var = best_var
        assignment[var] = not assignment[var]
    sat_count = count_satisfied(clauses, assignment)
    history.append(sat_count)
    if sat_count >= best_count:
        best_assignment, best_count = assignment, sat_count
    return best_assignment, history, best_count == len(clauses)


def read_clause_csv(upload):
    """Reads a lit1,lit2,lit3 CSV safely; returns (clauses, n_vars) or (None, None)."""
    try:
        real_df = pd.read_csv(upload)
        missing = [c for c in ["lit1", "lit2", "lit3"] if c not in real_df.columns]
        if missing:
            st.error(f"CSV is missing column(s): {', '.join(missing)}")
            return None, None
        lits = real_df[["lit1", "lit2", "lit3"]].apply(pd.to_numeric, errors="coerce").dropna()
        clauses = [[int(x) for x in row] for row in lits.values.tolist()]
        clauses = [c for c in clauses if all(x != 0 for x in c)]
        if not clauses:
            st.error("No valid clauses found (literals must be non-zero integers).")
            return None, None
        n_vars = max(abs(x) for c in clauses for x in c)
        st.info(f"Loaded {len(clauses)} real clauses over {n_vars} variables.")
        return clauses, n_vars
    except Exception as e:
        st.error(f"Could not read CSV: {e}")
        return None, None


def render_3sat_tab():
    st.markdown('<div class="module-banner banner-3sat">3-SAT - Logical Constraint Reasoning</div>',
                unsafe_allow_html=True)
    st.write(
        "Boolean satisfiability with clauses of exactly 3 literals each. This models "
        "hard logical-constraint problems in industrial settings - e.g. scheduling, "
        "configuration validation, or checking whether a set of design/process rules "
        "can be jointly satisfied. Solved here with exact brute force (small instances) "
        "or WalkSAT local search (larger instances)."
    )

    left, right = st.columns([1, 2])
    with left:
        st.markdown("**Synthetic data controls**")
        n_vars = st.slider("Number of variables", 4, 16, 8, key="sat_nvars")
        ratio = st.slider("Clause/variable ratio", 1.0, 6.0, 4.2, step=0.1, key="sat_ratio")
        m_clauses = max(1, int(round(n_vars * ratio)))
        seed = int(st.number_input("Random seed", value=3, step=1, key="sat_seed"))
        st.caption(f"Generating {m_clauses} random 3-literal clauses over {n_vars} variables "
                   f"(ratio {ratio:.1f} - the classic hardness peak is around 4.27).")

        st.markdown("**Or upload real clauses (CSV)**")
        st.caption("CSV needs columns `lit1,lit2,lit3` - signed integers, e.g. `3,-5,7` "
                   "means (x3 OR NOT x5 OR x7).")
        upload = st.file_uploader("Upload clause CSV", type=["csv"], key="sat_upload")

        clauses = None
        if upload is not None:
            clauses, uploaded_vars = read_clause_csv(upload)
            if clauses is not None:
                n_vars = uploaded_vars
        if clauses is None:
            clauses = generate_3sat(n_vars, m_clauses, seed=seed)

        template = pd.DataFrame({"lit1": [1, -2, 3], "lit2": [-3, 4, -5], "lit3": [5, -1, 2]})
        st.download_button("Download CSV template", data=export_csv_bytes(template),
                           file_name="3sat_clause_template.csv", mime="text/csv", key="sat_template")

        solver_choice = st.radio("Solver", ["Auto", "Brute force (exact)", "WalkSAT (local search)"],
                                 key="sat_solver")

    use_brute = (solver_choice == "Brute force (exact)" and n_vars <= BRUTE_FORCE_MAX_VARS) or \
                (solver_choice == "Auto" and n_vars <= 14)
    if solver_choice == "Brute force (exact)" and n_vars > BRUTE_FORCE_MAX_VARS:
        st.warning(f"Brute force is limited to {BRUTE_FORCE_MAX_VARS} variables; using WalkSAT instead.")

    history = None
    if use_brute:
        assignment, satisfiable = brute_force_sat(clauses, n_vars)
        method = "Brute force (exact)"
    else:
        assignment, history, satisfiable = walksat(clauses, n_vars, max_flips=500, seed=seed)
        method = "WalkSAT (local search, approximate)"

    sat_count = count_satisfied(clauses, assignment)
    sat_flags = [is_satisfied(c, assignment) for c in clauses]

    with right:
        st.markdown("**Clause satisfaction** (green = satisfied, red = violated, "
                    "under best assignment found)")
        fig, ax = plt.subplots(figsize=(9, 3.2))
        colors = [ACCENT_GREEN if s else ACCENT_RED for s in sat_flags]
        ax.bar(range(len(clauses)), [1] * len(clauses), color=colors)
        ax.set_xlabel("Clause index")
        ax.set_yticks([])
        fig.tight_layout()
        show_fig(fig)

        if history:
            st.markdown("**WalkSAT search progress**")
            fig3, ax3 = plt.subplots(figsize=(9, 2.6))
            ax3.plot(history, color=ACCENT_PURPLE, linewidth=2)
            ax3.axhline(len(clauses), color=ACCENT_GREEN, linestyle="--", linewidth=1,
                        label="All clauses satisfied")
            ax3.set_xlabel("Flip #")
            ax3.set_ylabel("Clauses satisfied")
            ax3.legend(fontsize=7)
            fig3.tight_layout()
            show_fig(fig3)

    if satisfiable:
        verdict = "SATISFIABLE"
    elif use_brute:
        verdict = "UNSATISFIABLE (proven)"
    else:
        verdict = "Best effort (not fully satisfied)"

    m1, m2, m3 = st.columns(3)
    metric_card(m1, "Result", verdict, ACCENT_PURPLE, size=22)
    metric_card(m2, "Clauses satisfied", f"{sat_count}/{len(clauses)}", ACCENT_PURPLE)
    metric_card(m3, "Method", method, ACCENT_PURPLE, size=16)

    st.divider()
    clause_df = pd.DataFrame(clauses, columns=["lit1", "lit2", "lit3"])
    clause_df["satisfied"] = sat_flags
    assign_df = pd.DataFrame({"variable": list(assignment.keys()),
                              "value": list(assignment.values())})

    sections = [
        ("Configuration", f"Variables={n_vars}, Clauses={len(clauses)}, "
                          f"Ratio={len(clauses) / n_vars:.2f}, Method={method}"),
        ("Verdict", f"{verdict} - {sat_count}/{len(clauses)} clauses satisfied"),
        ("Variable Assignment", assign_df),
        ("Clauses", clause_df),
    ]
    render_export_bar("3sat_results", "Kalsnet AI Reasoning Studio - 3-SAT Results",
                      sections, clause_df)


# ============================================================================
# ABOUT TAB
# ============================================================================

def render_about_tab():
    st.markdown('<div class="module-banner banner-about">About this Studio</div>',
                unsafe_allow_html=True)
    st.markdown(
        """
**Kalsnet AI Reasoning Studio** packages three foundational reasoning
paradigms as interactive modules, intended as building blocks for
Kalsnet's AI + Knowledge Graph applications:

| Module | Reasoning type | Typical Kalsnet use case |
|---|---|---|
| **Minimax** | Adversarial reasoning | Competitive negotiation, security red-team/blue-team simulation |
| **Expectimax** | Probabilistic reasoning | Planning under uncertain demand, equipment failure risk |
| **3-SAT** | Logical constraint reasoning | Validating that a set of process/design rules can be jointly satisfied |

Every module:
- Ships with **synthetic data** generated on the fly, so it works immediately.
- Can be **overridden with real data** via CSV upload or in-app editing.
- Renders **graphs** (game trees, bar charts, search-progress charts).
- Lets you **export results** as PDF, Word, CSV, or plain text for reporting.

---
**Developed by Randy Singh, Kalsnet (KNet) Consulting Group.**
"""
    )


# ============================================================================
# MAIN - TABS
# ============================================================================

tab1, tab2, tab3, tab4 = st.tabs(["Minimax", "Expectimax", "3-SAT", "About"])

with tab1:
    render_minimax_tab()

with tab2:
    render_expectimax_tab()

with tab3:
    render_3sat_tab()

with tab4:
    render_about_tab()
