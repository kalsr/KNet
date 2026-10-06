


# KNet Quantum and AI Algorithms Analytics Platform
# Developed By Randy Singh from Kalsnet (KNet) Consulting Group

# Modules: 3-SAT, Logical Equivalence, CNF Conversion, Markov Decision Process,
# Principal Component Analysis, Grover Quantum Search, Knowledge Check.

import io
import re
import sys
import math
import time
import textwrap
import datetime
from collections import Counter
from xml.sax.saxutils import escape as xml_escape

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyBboxPatch
import streamlit as st
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler

sys.setrecursionlimit(20000)

APP_NAME = "KNet Quantum and AI Algorithms Analytics Platform"
DEV_LINE = "Developed By Randy Singh from Kalsnet (KNet) Consulting Group"

BLUE = "#0B4FA8"
ORANGE = "#E07B00"
GREEN = "#2E8B57"
RED = "#C0392B"
GREY = "#8A94A6"
LIGHT = "#DCE6F5"

plt.rcParams.update({
    "figure.dpi": 110, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.alpha": 0.25, "font.size": 10,
    "axes.titleweight": "bold", "axes.titlesize": 11,
})


# =====================================================================
# GENERAL HELPERS
# =====================================================================
def natkey(s):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(s))]


def ui_dataframe(df, height=None):
    kw = {}
    if height:
        kw["height"] = height
    try:
        st.dataframe(df, width="stretch", **kw)
    except TypeError:
        st.dataframe(df, use_container_width=True, **kw)


def ui_image(png, caption=None):
    try:
        st.image(png, caption=caption, width="stretch")
    except (TypeError, Exception):
        st.image(png, caption=caption, use_container_width=True)


def ui_download(col, label, data, fname, mime, key):
    try:
        col.download_button(label, data, file_name=fname, mime=mime, key=key, on_click="ignore")
    except TypeError:
        col.download_button(label, data, file_name=fname, mime=mime, key=key)


def fig_png(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def read_upload_table(upl):
    """Read CSV or Excel upload into a DataFrame."""
    name = upl.name.lower()
    raw = upl.getvalue()
    if name.endswith((".xlsx", ".xls")):
        return pd.read_excel(io.BytesIO(raw))
    return pd.read_csv(io.BytesIO(raw))


# =====================================================================
# REPORT OBJECT AND EXPORTS (PDF, WORD, TEXT, CSV)
# =====================================================================
class Report:
    def __init__(self, title, module):
        self.title = title
        self.module = module
        self.items = []
        self.created = datetime.datetime.now()
        self._cache = {}

    def heading(self, text):
        self.items.append(("h", text))

    def para(self, text):
        self.items.append(("p", text))

    def bullets(self, lines):
        self.items.append(("b", list(lines)))

    def latex(self, tex, plain):
        self.items.append(("x", tex, plain))

    def metrics(self, d):
        self.items.append(("m", {k: str(v) for k, v in d.items()}))

    def table(self, name, df):
        self.items.append(("t", name, df.reset_index(drop=True)))

    def figure(self, name, fig):
        self.items.append(("f", name, fig_png(fig)))

    def tables(self):
        return {it[1]: it[2] for it in self.items if it[0] == "t"}


_ASCII_MAP = {
    "¬": "NOT ", "∧": "AND", "∨": "OR", "⊕": "XOR", "→": "->", "↔": "<->", "≡": "==",
    "γ": "gamma", "θ": "theta", "π": "pi", "√": "sqrt", "≈": "~", "≤": "<=", "≥": ">=",
    "−": "-", "–": "-", "—": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "⟩": ">",
    "⟨": "<", "ψ": "psi", "α": "alpha", "β": "beta", "λ": "lambda", "Σ": "Sum",
    "σ": "sigma", "×": "x", "…": "...", "•": "-", "²": "^2", "ⁿ": "^n",
}


def ascii_safe(s):
    s = str(s)
    for k, v in _ASCII_MAP.items():
        s = s.replace(k, v)
    return s.encode("latin-1", "replace").decode("latin-1")


def _fmt_cell(x):
    if isinstance(x, (bool, np.bool_)):
        return "TRUE" if x else "FALSE"
    if isinstance(x, (float, np.floating)):
        if np.isnan(x):
            return ""
        if np.isinf(x):
            return "inf" if x > 0 else "-inf"
        return f"{x:.4g}"
    return str(x)


def _table_rows(df, max_rows=150, max_cols=10):
    d = df.iloc[:max_rows, :max_cols]
    rows = [[str(c) for c in d.columns]]
    rows += [[_fmt_cell(v) for v in r] for r in d.itertuples(index=False)]
    return rows, len(df) > max_rows, df.shape[1] > max_cols


def report_txt(rep):
    out = [APP_NAME, DEV_LINE, "=" * 78, rep.title,
           f"Module: {rep.module}    Generated: {rep.created:%Y-%m-%d %H:%M}", "=" * 78, ""]
    for it in rep.items:
        k = it[0]
        if k == "h":
            out += ["", it[1].upper(), "-" * len(it[1])]
        elif k == "p":
            out += textwrap.wrap(it[1], 100) + [""]
        elif k == "b":
            out += [f"  - {x}" for x in it[1]] + [""]
        elif k == "x":
            out += [f"    Formula: {it[2]}", ""]
        elif k == "m":
            w = max(len(a) for a in it[1])
            out += [f"  {a.ljust(w)} : {b}" for a, b in it[1].items()] + [""]
        elif k == "t":
            out += [f"[Table] {it[1]}  ({len(it[2])} rows)",
                    it[2].to_string(max_rows=500, max_colwidth=60), ""]
        elif k == "f":
            out += [f"[Figure] {it[1]} (see PDF or Word export for the image)", ""]
    return "\n".join(out).encode("utf-8")


def report_pdf(rep):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, Image as RLImage)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from PIL import Image as PILImage

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=rep.title,
                            author="Randy Singh, Kalsnet (KNet) Consulting Group")
    ss = getSampleStyleSheet()
    blue = colors.HexColor(BLUE)
    s_title = ParagraphStyle("kt", parent=ss["Title"], textColor=blue, fontName="Helvetica-Bold",
                             fontSize=16, leading=20, alignment=0, spaceAfter=2)
    s_h = ParagraphStyle("kh", parent=ss["Heading2"], textColor=blue)
    s_p = ss["BodyText"]
    s_cell = ParagraphStyle("kc", parent=s_p, fontSize=7, leading=8.5)
    s_head = ParagraphStyle("khd", parent=s_cell, textColor=colors.white, fontName="Helvetica-Bold")
    s_cap = ParagraphStyle("kcap", parent=s_p, fontSize=8, textColor=colors.grey)
    s_formula = ParagraphStyle("kf", parent=s_p, fontName="Courier", fontSize=9,
                               backColor=colors.HexColor("#F2F5FB"), borderPadding=4)
    W = letter[0] - 1.2 * inch

    def P(t, s=s_p):
        return Paragraph(xml_escape(ascii_safe(t)).replace("\n", "<br/>"), s)

    def grid_table(rows):
        data = [[P(c, s_head) for c in rows[0]]] + [[P(c, s_cell) for c in r] for r in rows[1:]]
        t = Table(data, colWidths=[W / len(rows[0])] * len(rows[0]), repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), blue),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B8C4D6")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FB")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return t

    story = [P(APP_NAME, s_title), P(DEV_LINE, s_title), Spacer(1, 8), P(rep.title, s_h),
             P(f"Module: {rep.module}    Generated: {rep.created:%Y-%m-%d %H:%M}", s_cap), Spacer(1, 8)]
    for it in rep.items:
        k = it[0]
        if k == "h":
            story.append(P(it[1], s_h))
        elif k == "p":
            story += [P(it[1]), Spacer(1, 4)]
        elif k == "b":
            for x in it[1]:
                story.append(P("- " + x))
            story.append(Spacer(1, 4))
        elif k == "x":
            story += [P(it[2], s_formula), Spacer(1, 6)]
        elif k == "m":
            rows = [["Metric", "Value"]] + [[a, b] for a, b in it[1].items()]
            story += [grid_table(rows), Spacer(1, 8)]
        elif k == "t":
            rows, tr, tc = _table_rows(it[2])
            note = f"Table: {it[1]} ({len(it[2])} rows"
            note += ", first 150 shown" if tr else ""
            note += ", first 10 columns shown" if tc else ""
            story += [P(note + ")", s_cap), grid_table(rows), Spacer(1, 8)]
        elif k == "f":
            im = PILImage.open(io.BytesIO(it[2]))
            w, h = im.size
            dw = min(W, 6.8 * inch)
            dh = dw * h / w
            if dh > 8 * inch:
                dw, dh = dw * 8 * inch / dh, 8 * inch
            story += [RLImage(io.BytesIO(it[2]), width=dw, height=dh), P("Figure: " + it[1], s_cap),
                      Spacer(1, 8)]
    doc.build(story)
    return buf.getvalue()


def report_docx(rep):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    for line in (APP_NAME, DEV_LINE):
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = RGBColor(0x0B, 0x4F, 0xA8)
    doc.add_heading(rep.title, level=1)
    doc.add_paragraph(f"Module: {rep.module}    Generated: {rep.created:%Y-%m-%d %H:%M}")

    def add_table(rows):
        tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
        tbl.style = "Table Grid"
        for i, (r, trow) in enumerate(zip(rows, tbl.rows)):
            cells = trow.cells
            for j, v in enumerate(r):
                run = cells[j].paragraphs[0].add_run(str(v))
                run.font.size = Pt(8)
                if i == 0:
                    run.bold = True
        doc.add_paragraph()

    for it in rep.items:
        k = it[0]
        if k == "h":
            doc.add_heading(it[1], level=2)
        elif k == "p":
            doc.add_paragraph(it[1])
        elif k == "b":
            for x in it[1]:
                doc.add_paragraph(x, style="List Bullet")
        elif k == "x":
            p = doc.add_paragraph()
            r = p.add_run(it[2])
            r.font.name = "Courier New"
            r.font.size = Pt(10)
        elif k == "m":
            add_table([["Metric", "Value"]] + [[a, b] for a, b in it[1].items()])
        elif k == "t":
            rows, tr, tc = _table_rows(it[2])
            note = f"Table: {it[1]} ({len(it[2])} rows" + (", first 150 shown" if tr else "") + \
                   (", first 10 columns shown" if tc else "") + ")"
            doc.add_paragraph(note).runs[0].italic = True
            add_table(rows)
        elif k == "f":
            doc.add_picture(io.BytesIO(it[2]), width=Inches(6.3))
            doc.add_paragraph("Figure: " + it[1]).runs[0].italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_downloads(rep, key):
    st.markdown("##### Export results")
    base = re.sub(r"\W+", "_", rep.module.lower()).strip("_") + "_results"
    c1, c2, c3, c4 = st.columns(4)
    try:
        if "pdf" not in rep._cache:
            rep._cache["pdf"] = report_pdf(rep)
        ui_download(c1, "Download PDF", rep._cache["pdf"], base + ".pdf", "application/pdf", key + "_pdf")
    except Exception as e:
        c1.error(f"PDF export unavailable: {e}")
    try:
        if "docx" not in rep._cache:
            rep._cache["docx"] = report_docx(rep)
        ui_download(c2, "Download Word", rep._cache["docx"], base + ".docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key + "_docx")
    except Exception as e:
        c2.error(f"Word export unavailable: {e}")
    if "txt" not in rep._cache:
        rep._cache["txt"] = report_txt(rep)
    ui_download(c3, "Download Text", rep._cache["txt"], base + ".txt", "text/plain", key + "_txt")
    tables = rep.tables()
    if tables:
        choice = c4.selectbox("Table for CSV", list(tables), key=key + "_csvsel", label_visibility="collapsed")
        ui_download(c4, "Download CSV", tables[choice].to_csv(index=False).encode("utf-8"),
                    base + "_" + re.sub(r"\W+", "_", choice.lower()) + ".csv", "text/csv", key + "_csv")


def show_report(rep, key):
    st.markdown(f"### {rep.title}")
    st.caption(f"Generated {rep.created:%Y-%m-%d %H:%M}")
    for i, it in enumerate(rep.items):
        k = it[0]
        if k == "h":
            st.markdown(f"#### {it[1]}")
        elif k == "p":
            st.markdown(it[1])
        elif k == "b":
            st.markdown("\n".join(f"- {x}" for x in it[1]))
        elif k == "x":
            st.latex(it[1])
        elif k == "m":
            items = list(it[1].items())
            for start in range(0, len(items), 4):
                cols = st.columns(4)
                for c, (a, b) in zip(cols, items[start:start + 4]):
                    c.metric(a, b)
        elif k == "t":
            st.markdown(f"**{it[1]}**")
            ui_dataframe(it[2], height=min(420, 40 + 35 * max(1, len(it[2]))))
        elif k == "f":
            ui_image(it[2], caption=it[1])
    render_downloads(rep, key)


def show_saved_result(state_key, label):
    rep = st.session_state.get(state_key)
    if rep is None:
        st.info(f"No {label} results yet. Run the algorithm in the Synthetic Data Lab or Upload Real Data section.")
    else:
        show_report(rep, state_key + "_res")


def template_button(df, fname, label="Download sample upload template", key=None):
    ui_download(st, label, df.to_csv(index=False).encode("utf-8"), fname, "text/csv", key or fname)


def use_cases(cases):
    for title, body in cases:
        with st.container(border=True):
            st.markdown(f"**{title}**")
            st.markdown(body)


# =====================================================================
# PROPOSITIONAL LOGIC ENGINE (shared by Logical Equivalence and CNF)
# =====================================================================
TOKEN_RE = re.compile(r"\s*(<->|<=>|->|=>|&&|\|\||↔|→|¬|∧|∨|⊕|~|!|&|\||\^|\(|\)|"
                      r"[A-Za-z_][A-Za-z0-9_]*|[01])")
KEYWORDS = {"NOT": "not", "AND": "and", "OR": "or", "XOR": "xor", "IMPLIES": "imp",
            "IFF": "iff", "TRUE": "T", "FALSE": "F"}
SYMBOLS = {"~": "not", "!": "not", "¬": "not", "&": "and", "&&": "and", "∧": "and", "|": "or",
           "||": "or", "∨": "or", "^": "xor", "⊕": "xor", "->": "imp", "=>": "imp", "→": "imp",
           "<->": "iff", "<=>": "iff", "↔": "iff", "(": "(", ")": ")", "1": "T", "0": "F"}
BIN_OPS = ("and", "or", "xor", "imp", "iff")
MAX_TT_VARS = 16


def tokenize(s):
    toks, pos, s = [], 0, s.strip()
    while pos < len(s):
        m = TOKEN_RE.match(s, pos)
        if not m or m.end() == pos:
            raise ValueError(f"Unexpected character near '{s[pos:pos + 12]}'")
        t = m.group(1)
        pos = m.end()
        if t in SYMBOLS:
            toks.append((SYMBOLS[t], None))
        elif t.upper() in KEYWORDS:
            toks.append((KEYWORDS[t.upper()], None))
        else:
            toks.append(("var", t))
        while pos < len(s) and s[pos].isspace():
            pos += 1
    return toks


class _Parser:
    def __init__(self, toks):
        self.t, self.i = toks, 0

    def peek(self):
        return self.t[self.i][0] if self.i < len(self.t) else None

    def parse(self):
        node = self.iff()
        if self.i != len(self.t):
            raise ValueError(f"Unexpected token '{self.t[self.i][1] or self.t[self.i][0]}'")
        return node

    def iff(self):
        n = self.imp()
        while self.peek() == "iff":
            self.i += 1
            n = ("iff", n, self.imp())
        return n

    def imp(self):
        n = self.or_()
        if self.peek() == "imp":
            self.i += 1
            return ("imp", n, self.imp())
        return n

    def or_(self):
        n = self.xor()
        while self.peek() == "or":
            self.i += 1
            n = ("or", n, self.xor())
        return n

    def xor(self):
        n = self.and_()
        while self.peek() == "xor":
            self.i += 1
            n = ("xor", n, self.and_())
        return n

    def and_(self):
        n = self.not_()
        while self.peek() == "and":
            self.i += 1
            n = ("and", n, self.not_())
        return n

    def not_(self):
        if self.peek() == "not":
            self.i += 1
            return ("not", self.not_())
        return self.atom()

    def atom(self):
        k = self.peek()
        if k == "(":
            self.i += 1
            n = self.iff()
            if self.peek() != ")":
                raise ValueError("Missing closing parenthesis")
            self.i += 1
            return n
        if k == "var":
            name = self.t[self.i][1]
            self.i += 1
            return ("var", name)
        if k in ("T", "F"):
            self.i += 1
            return ("const", k == "T")
        raise ValueError("Expression ended unexpectedly" if k is None else f"Unexpected '{k}'")


def parse_expr(s):
    toks = tokenize(str(s))
    if not toks:
        raise ValueError("Empty expression")
    return _Parser(toks).parse()


_OPS = {
    "word": {"and": " AND ", "or": " OR ", "xor": " XOR ", "imp": " -> ", "iff": " <-> "},
    "latex": {"and": r" \land ", "or": r" \lor ", "xor": r" \oplus ", "imp": r" \rightarrow ",
              "iff": r" \leftrightarrow "},
}
_NOT = {"word": "NOT ", "latex": r"\neg "}


def _latex_var(name):
    if len(name) == 1:
        return name
    return r"\mathit{" + name.replace("_", r"\_") + "}"


def to_str(n, style="word"):
    ops = _OPS[style]
    k = n[0]
    if k == "var":
        return n[1] if style == "word" else _latex_var(n[1])
    if k == "const":
        if style == "word":
            return "TRUE" if n[1] else "FALSE"
        return r"\top" if n[1] else r"\bot"
    if k == "not":
        c = n[1]
        s = to_str(c, style)
        return _NOT[style] + (f"({s})" if c[0] in BIN_OPS else s)

    def wrap(c):
        s = to_str(c, style)
        if c[0] in BIN_OPS and not (c[0] == k and k in ("and", "or", "xor")):
            return f"({s})"
        return s

    return wrap(n[1]) + ops[k] + wrap(n[2])


def variables(n):
    out = set()

    def walk(x):
        if x[0] == "var":
            out.add(x[1])
        elif x[0] != "const":
            for c in x[1:]:
                walk(c)

    walk(n)
    return sorted(out, key=natkey)


def ast_size(n):
    if n[0] in ("var", "const"):
        return 1
    return 1 + sum(ast_size(c) for c in n[1:])


def assignment_env(vs):
    nv = len(vs)
    size = 1 << nv
    idx = np.arange(size, dtype=np.int64)
    env = {v: ((idx >> (nv - 1 - i)) & 1) == 0 for i, v in enumerate(vs)}  # TRUE first like textbooks
    return env, size


def ev(n, env, size):
    k = n[0]
    if k == "var":
        return env[n[1]]
    if k == "const":
        return np.full(size, n[1])
    if k == "not":
        return ~ev(n[1], env, size)
    a, b = ev(n[1], env, size), ev(n[2], env, size)
    if k == "and":
        return a & b
    if k == "or":
        return a | b
    if k == "xor":
        return a ^ b
    if k == "imp":
        return (~a) | b
    return ~(a ^ b)


def classify(arr):
    if arr.all():
        return "Tautology"
    if not arr.any():
        return "Contradiction"
    return "Contingent"


def fmt_row(vs, env, r):
    return ", ".join(f"{v}={'T' if env[v][r] else 'F'}" for v in vs)


def subst(n, mapping):
    if n[0] == "var":
        return mapping.get(n[1], n)
    if n[0] == "const":
        return n
    return (n[0],) + tuple(subst(c, mapping) for c in n[1:])


def random_expr(vs, depth, rng, ops=("and", "or", "imp", "iff", "xor", "not")):
    if depth <= 0 or rng.random() < 0.25:
        v = ("var", str(rng.choice(vs)))
        return ("not", v) if rng.random() < 0.3 else v
    op = str(rng.choice(ops))
    if op == "not":
        return ("not", random_expr(vs, depth - 1, rng, ops))
    return (op, random_expr(vs, depth - 1, rng, ops), random_expr(vs, depth - 1, rng, ops))


def ast_dot(n, title=None):
    lines = ['digraph G {', 'node [fontname="Helvetica", fontsize=11];', 'edge [color="#8A94A6"];']
    if title:
        lines.append(f'labelloc="t"; label="{title}";')
    counter = [0]
    names = {"and": "AND", "or": "OR", "xor": "XOR", "imp": "IMPLIES", "iff": "IFF", "not": "NOT"}

    def add(x):
        counter[0] += 1
        nid = f"n{counter[0]}"
        if x[0] == "var":
            lines.append(f'{nid} [label="{x[1]}", shape=circle, style=filled, fillcolor="#E8F5EE", color="#2E8B57"];')
        elif x[0] == "const":
            lines.append(f'{nid} [label="{"TRUE" if x[1] else "FALSE"}", shape=box];')
        else:
            lines.append(f'{nid} [label="{names[x[0]]}", shape=box, style="rounded,filled", '
                         f'fillcolor="#DCE6F5", color="#0B4FA8"];')
            for c in x[1:]:
                lines.append(f"{nid} -> {add(c)};")
        return nid

    add(n)
    lines.append("}")
    return "\n".join(lines)


def ast_figure(n, title):
    """Matplotlib expression tree (for PDF and Word reports)."""
    pos, labels, edges = {}, {}, []
    names = {"and": "AND", "or": "OR", "xor": "XOR", "imp": "->", "iff": "<->", "not": "NOT"}
    counter = [0]
    leaf_x = [0]

    def layout(x, depth):
        counter[0] += 1
        nid = counter[0]
        if x[0] in ("var", "const"):
            pos[nid] = (leaf_x[0], -depth)
            leaf_x[0] += 1
            labels[nid] = x[1] if x[0] == "var" else ("TRUE" if x[1] else "FALSE")
        else:
            kids = [layout(c, depth + 1) for c in x[1:]]
            pos[nid] = (np.mean([pos[k][0] for k in kids]), -depth)
            labels[nid] = names[x[0]]
            edges.extend((nid, k) for k in kids)
        return nid

    layout(n, 0)
    w = max(4, 0.7 * leaf_x[0] + 1)
    h = max(2.5, 0.8 * (1 - min(p[1] for p in pos.values())) + 1)
    fig, ax = plt.subplots(figsize=(min(w, 14), min(h, 10)))
    for a, b in edges:
        ax.plot([pos[a][0], pos[b][0]], [pos[a][1], pos[b][1]], color=GREY, lw=1, zorder=1)
    for nid, (x, y) in pos.items():
        is_leaf = labels[nid] not in names.values()
        ax.text(x, y, labels[nid], ha="center", va="center", fontsize=9, zorder=2,
                bbox=dict(boxstyle="round,pad=0.3", fc="#E8F5EE" if is_leaf else LIGHT,
                          ec=GREEN if is_leaf else BLUE))
    ax.set_title(title)
    ax.axis("off")
    return fig


# =====================================================================
# 3-SAT ENGINE
# =====================================================================
def random_3sat(n, m, rng, planted=False):
    hidden = rng.random(n) < 0.5
    clauses, tries = [], 0
    while len(clauses) < m and tries < m * 200:
        tries += 1
        vs = rng.choice(n, 3, replace=False) + 1
        signs = rng.random(3) < 0.5
        cl = tuple(int(v) if s else -int(v) for v, s in zip(vs, signs))
        if planted and not any((l > 0) == hidden[abs(l) - 1] for l in cl):
            continue
        clauses.append(cl)
    return clauses, hidden


def lit_name(l, names):
    nm = names[abs(l) - 1] if names else f"x{abs(l)}"
    return nm if l > 0 else f"NOT {nm}"


def clause_text(c, names):
    return "(" + " OR ".join(lit_name(l, names) for l in c) + ")"


def clause_latex(c, names):
    def one(l):
        nm = _latex_var(names[abs(l) - 1]) if names else f"x_{{{abs(l)}}}"
        return nm if l > 0 else r"\neg " + nm

    return "(" + r" \lor ".join(one(l) for l in c) + ")"


def dpll(clauses, n, max_decisions=200000):
    stats = {"decisions": 0, "propagations": 0, "backtracks": 0}

    def simplify(cls, lit):
        out = []
        for c in cls:
            if lit in c:
                continue
            if -lit in c:
                nc = tuple(x for x in c if x != -lit)
                if not nc:
                    return None
                out.append(nc)
            else:
                out.append(c)
        return out

    def solve(cls, assign):
        while True:
            unit = next((c[0] for c in cls if len(c) == 1), None)
            if unit is None:
                break
            assign[abs(unit)] = unit > 0
            stats["propagations"] += 1
            cls = simplify(cls, unit)
            if cls is None:
                return None
        if not cls:
            return assign
        lits = set(l for c in cls for l in c)
        pures = [l for l in lits if -l not in lits]
        if pures:
            for l in pures:
                assign[abs(l)] = l > 0
                cls = simplify(cls, l)
            return solve(cls, assign)
        minlen = min(len(c) for c in cls)
        cnt = Counter(l for c in cls if len(c) == minlen for l in c)
        lit = max(cnt, key=cnt.get)
        stats["decisions"] += 1
        if stats["decisions"] > max_decisions:
            raise TimeoutError("decision limit reached")
        for i, choice in enumerate((lit, -lit)):
            a2 = dict(assign)
            a2[abs(choice)] = choice > 0
            c2 = simplify(cls, choice)
            if c2 is not None:
                r = solve(c2, a2)
                if r is not None:
                    return r
            if i == 0:
                stats["backtracks"] += 1
        return None

    clean = [tuple(dict.fromkeys(c)) for c in clauses]
    if any(len(c) == 0 for c in clean):
        return None, stats
    res = solve(clean, {})
    if res is None:
        return None, stats
    return {v: res.get(v, False) for v in range(1, n + 1)}, stats


def _clause_matrix(clauses):
    w = max(len(c) for c in clauses)
    L = np.zeros((len(clauses), w), dtype=np.int64)
    for i, c in enumerate(clauses):
        L[i, :len(c)] = c
    return L


def walksat(clauses, n, max_flips=20000, noise=0.5, seed=0):
    rng = np.random.default_rng(seed)
    L = _clause_matrix(clauses)
    mask, V, pos = L != 0, np.abs(L), L > 0

    def unsat_vec(a):
        return ~(((a[V] == pos) & mask).any(1))

    assign = rng.random(n + 1) < 0.5
    uns = unsat_vec(assign)
    hist = [int(uns.sum())]
    best, best_u = assign.copy(), hist[0]
    for _ in range(max_flips):
        idx = np.nonzero(uns)[0]
        if len(idx) == 0:
            break
        ci = rng.choice(idx)
        cand = V[ci][mask[ci]]
        if rng.random() < noise:
            v = int(rng.choice(cand))
        else:
            scores = []
            for v_ in cand:
                assign[v_] = not assign[v_]
                scores.append(int(unsat_vec(assign).sum()))
                assign[v_] = not assign[v_]
            v = int(cand[int(np.argmin(scores))])
        assign[v] = not assign[v]
        uns = unsat_vec(assign)
        u = int(uns.sum())
        hist.append(u)
        if u < best_u:
            best_u, best = u, assign.copy()
    return {v: bool(best[v]) for v in range(1, n + 1)}, best_u == 0, hist


def model_mask(clauses, n):
    """Vectorised truth evaluation over all 2^n assignments. Variable 1 is the most significant bit."""
    size = 1 << n
    idx = np.arange(size, dtype=np.int64)
    bits = [((idx >> (n - v)) & 1) == 1 for v in range(1, n + 1)]
    ok = np.ones(size, dtype=bool)
    for c in clauses:
        cl = np.zeros(size, dtype=bool)
        for l in c:
            cl |= bits[abs(l) - 1] if l > 0 else ~bits[abs(l) - 1]
        ok &= cl
    return ok


def parse_dimacs(text):
    clauses, cur = [], []
    for line in text.splitlines():
        line = line.strip()
        if not line or line[0] in "cp":
            continue
        if line.startswith("%"):
            break
        for tok in line.split():
            v = int(tok)
            if v == 0:
                if cur:
                    clauses.append(tuple(cur))
                cur = []
            else:
                cur.append(v)
    if cur:
        clauses.append(tuple(cur))
    if not clauses:
        raise ValueError("No clauses found in DIMACS file")
    n = max(abs(l) for c in clauses for l in c)
    return clauses, n, [f"x{i}" for i in range(1, n + 1)]


def parse_sat_table(df):
    """Rows are clauses. Cells are signed integers or names such as A, -B, NOT C, ~D."""
    raw_rows = df.astype(str).values.tolist()
    if raw_rows and all(re.match(r"^(l|lit|literal|clause|c|col|column)[ _]?\d+$", str(x).strip(), re.I)
                        for x in raw_rows[0] if str(x).strip() not in ("", "nan")):
        raw_rows = raw_rows[1:]
    parsed, num_max, name_ids = [], 0, {}
    for row in raw_rows:
        lits = []
        for cell in row:
            t = str(cell).strip()
            if t in ("", "nan", "None"):
                continue
            if re.fullmatch(r"[+-]?\d+(\.0)?", t):
                v = int(float(t))
                if v != 0:
                    lits.append(("num", v))
                    num_max = max(num_max, abs(v))
                continue
            neg = False
            m = re.match(r"^(not\s+|-|~|!|¬)\s*(.+)$", t, re.I)
            if m:
                neg, t = True, m.group(2).strip()
            lits.append(("name", t, neg))
        if lits:
            parsed.append(lits)
    if not parsed:
        raise ValueError("No clauses found in the table")
    names = [f"x{i}" for i in range(1, num_max + 1)]
    for row in parsed:
        for l in row:
            if l[0] == "name" and l[1] not in name_ids:
                names.append(l[1])
                name_ids[l[1]] = len(names)
    clauses = []
    for row in parsed:
        c = []
        for l in row:
            if l[0] == "num":
                c.append(l[1])
            else:
                c.append(-name_ids[l[1]] if l[2] else name_ids[l[1]])
        clauses.append(tuple(c))
    return clauses, len(names), names


CYBER_NAMES = ["Firewall", "MFA", "IDS", "VPN", "Patching", "Encryption", "Backup", "EDR", "SIEM",
               "DLP", "WAF", "PAM", "ZeroTrust", "Segmentation", "CertAuth", "TrustedDevice",
               "Logging", "SSO", "Antivirus", "Sandbox"]


def run_sat(clauses, n, names, solver, source, seed=0, count=True, max_flips=20000, noise=0.5):
    rep = Report("3-SAT Satisfiability Analysis", "3-SAT")
    m = len(clauses)
    widths = Counter(len(c) for c in clauses)
    t0 = time.perf_counter()
    assign, status, extra, hist = None, "", {}, None
    if solver.startswith("DPLL"):
        try:
            assign, stats = dpll(clauses, n)
            status = "SATISFIABLE" if assign else "UNSATISFIABLE"
        except TimeoutError:
            stats = {"decisions": "limit"}
            status = "UNKNOWN (decision limit reached)"
        extra = {"Decisions": stats.get("decisions"), "Unit propagations": stats.get("propagations", "-"),
                 "Backtracks": stats.get("backtracks", "-")}
    elif solver.startswith("WalkSAT"):
        assign, ok, hist = walksat(clauses, n, max_flips, noise, seed)
        status = "SATISFIABLE" if ok else "UNKNOWN (local search found no solution)"
        if not ok:
            extra["Best unsatisfied clauses"] = min(hist)
        extra["Flips used"] = len(hist) - 1
    else:
        ok = model_mask(clauses, n)
        sol = np.nonzero(ok)[0]
        status = "SATISFIABLE" if len(sol) else "UNSATISFIABLE"
        if len(sol):
            x = int(sol[0])
            assign = {v: bool((x >> (n - v)) & 1) for v in range(1, n + 1)}
        extra["Assignments checked"] = f"{1 << n:,}"
    elapsed = (time.perf_counter() - t0) * 1000
    n_models = None
    if count and n <= 20:
        n_models = int(model_mask(clauses, n).sum())

    rep.heading("Problem summary")
    rep.para(f"Data source: {source}. Solver: {solver}.")
    ratio = m / max(n, 1)
    metrics = {"Variables": n, "Clauses": m, "Clause to variable ratio": f"{ratio:.2f}", "Result": status,
               "Solve time (ms)": f"{elapsed:.1f}"}
    metrics.update(extra)
    if n_models is not None:
        metrics["Number of satisfying assignments"] = f"{n_models:,} of {1 << n:,}"
    rep.metrics(metrics)
    if len(clauses) <= 6:
        rep.latex(r" \land ".join(clause_latex(c, names) for c in clauses),
                  " AND ".join(clause_text(c, names) for c in clauses))
    if set(widths) != {3}:
        rep.para("Note: not every clause has exactly three literals (clause widths: "
                 + ", ".join(f"{k} literals x {v}" for k, v in sorted(widths.items()))
                 + "). The solver handles general CNF; use the CNF Conversion module to build a strict 3-CNF.")
    phase = ("below the 4.26 phase transition, so random instances are usually satisfiable and easy"
             if ratio < 3.8 else
             "near the 4.26 phase transition, where random 3-SAT instances are hardest" if ratio < 4.8 else
             "above the 4.26 phase transition, so random instances are usually unsatisfiable")
    rep.para(f"Interpretation: the ratio {ratio:.2f} is {phase}. Result: {status}.")

    if assign:
        rep.heading("Satisfying assignment")
        rep.table("Variable assignment", pd.DataFrame({
            "variable_id": list(range(1, n + 1)), "name": names[:n],
            "value": ["TRUE" if assign[v] else "FALSE" for v in range(1, n + 1)]}))
    rows = []
    for j, c in enumerate(clauses[:5000], 1):
        if assign:
            true_lits = [lit_name(l, names) for l in c if assign[abs(l)] == (l > 0)]
        rows.append({"clause": j, "literals": " ".join(str(l) for l in c), "clause_text": clause_text(c, names),
                     "satisfied": (len(true_lits) > 0) if assign else None,
                     "true_literals": ", ".join(true_lits) if assign else ""})
    rep.table("Clause evaluation", pd.DataFrame(rows))

    pos_cnt = Counter(abs(l) for c in clauses for l in c if l > 0)
    neg_cnt = Counter(abs(l) for c in clauses for l in c if l < 0)
    show = list(range(1, min(n, 40) + 1))
    fig, ax = plt.subplots(figsize=(9, 3.4))
    x = np.arange(len(show))
    ax.bar(x - 0.2, [pos_cnt[v] for v in show], 0.4, color=BLUE, label="Positive literal")
    ax.bar(x + 0.2, [neg_cnt[v] for v in show], 0.4, color=ORANGE, label="Negated literal")
    ax.set_xticks(x)
    ax.set_xticklabels([names[v - 1] for v in show], rotation=60, ha="right", fontsize=7)
    ax.set_ylabel("Occurrences")
    ax.set_title("Variable occurrences in the formula" + (" (first 40 variables)" if n > 40 else ""))
    ax.legend()
    rep.figure("Variable occurrence profile", fig)
    if assign:
        tl = [sum(assign[abs(l)] == (l > 0) for l in c) for c in clauses]
        fig, ax = plt.subplots(figsize=(6, 3.2))
        cnt = Counter(tl)
        ks = sorted(cnt)
        ax.bar([str(k) for k in ks], [cnt[k] for k in ks], color=[RED if k == 0 else BLUE for k in ks])
        ax.set_xlabel("Number of TRUE literals in clause")
        ax.set_ylabel("Clauses")
        ax.set_title("How strongly each clause is satisfied")
        rep.figure("True literal count per clause", fig)
    if hist:
        fig, ax = plt.subplots(figsize=(8, 3.2))
        ax.plot(hist, color=BLUE, lw=1)
        ax.set_xlabel("Flip")
        ax.set_ylabel("Unsatisfied clauses")
        ax.set_title("WalkSAT local search progress")
        rep.figure("WalkSAT convergence", fig)
    dimacs = [f"p cnf {n} {m}"] + [" ".join(str(l) for l in c) + " 0" for c in clauses]
    rep.table("Formula in DIMACS format", pd.DataFrame({"dimacs_line": dimacs}))
    return rep


def sat_phase_transition(n, ratios, trials, seed):
    rng = np.random.default_rng(seed)
    rows = []
    for r in ratios:
        m = int(round(r * n))
        sats, decs = 0, []
        for _ in range(trials):
            cl, _ = random_3sat(n, m, rng)
            a, s = dpll(cl, n)
            sats += a is not None
            decs.append(s["decisions"])
        rows.append({"ratio": r, "clauses": m, "p_satisfiable": sats / trials, "mean_decisions": np.mean(decs)})
    return pd.DataFrame(rows)


# =====================================================================
# CNF ENGINE
# =====================================================================
class CNFBlowup(Exception):
    pass


def elim_iff_xor(n):
    k = n[0]
    if k in ("var", "const"):
        return n
    if k == "not":
        return ("not", elim_iff_xor(n[1]))
    a, b = elim_iff_xor(n[1]), elim_iff_xor(n[2])
    if k == "iff":
        return ("and", ("imp", a, b), ("imp", b, a))
    if k == "xor":
        return ("and", ("or", a, b), ("or", ("not", a), ("not", b)))
    return (k, a, b)


def elim_imp(n):
    k = n[0]
    if k in ("var", "const"):
        return n
    if k == "not":
        return ("not", elim_imp(n[1]))
    a, b = elim_imp(n[1]), elim_imp(n[2])
    if k == "imp":
        return ("or", ("not", a), b)
    return (k, a, b)


def to_nnf(n):
    k = n[0]
    if k == "not":
        c = n[1]
        if c[0] == "not":
            return to_nnf(c[1])
        if c[0] == "and":
            return ("or", to_nnf(("not", c[1])), to_nnf(("not", c[2])))
        if c[0] == "or":
            return ("and", to_nnf(("not", c[1])), to_nnf(("not", c[2])))
        if c[0] == "const":
            return ("const", not c[1])
        return n
    if k in ("and", "or"):
        return (k, to_nnf(n[1]), to_nnf(n[2]))
    return n


def simp_const(n):
    k = n[0]
    if k not in ("and", "or"):
        return n
    a, b = simp_const(n[1]), simp_const(n[2])
    for x, y in ((a, b), (b, a)):
        if x[0] == "const":
            if k == "and":
                return y if x[1] else ("const", False)
            return ("const", True) if x[1] else y
    return (k, a, b)


def clauses_of(n, limit=20000):
    k = n[0]
    if k == "var":
        return [frozenset([(n[1], True)])]
    if k == "not":
        return [frozenset([(n[1][1], False)])]
    if k == "const":
        return [] if n[1] else [frozenset()]
    A, B = clauses_of(n[1], limit), clauses_of(n[2], limit)
    if k == "and":
        out = A + B
        if len(out) > limit:
            raise CNFBlowup()
        return out
    if len(A) * len(B) > limit:
        raise CNFBlowup()
    return [x | y for x in A for y in B]


def simplify_clauses(cls):
    out = []
    seen = set()
    for c in cls:
        names = {}
        taut = False
        for v, p in c:
            if names.get(v, p) != p:
                taut = True
                break
            names[v] = p
        if taut or c in seen:
            continue
        seen.add(c)
        out.append(c)
    if len(out) <= 3000:
        out.sort(key=len)
        kept = []
        for c in out:
            if not any(k <= c for k in kept):
                kept.append(c)
        out = kept
    return out


def lit_str(l):
    return l[0] if l[1] else f"NOT {l[0]}"


def clause_str(c):
    if not c:
        return "FALSE"
    lits = sorted(c, key=lambda l: natkey(l[0]))
    return "(" + " OR ".join(lit_str(l) for l in lits) + ")" if len(lits) > 1 else lit_str(lits[0])


def cnf_str(cls):
    if not cls:
        return "TRUE"
    return " AND ".join(clause_str(c) for c in cls)


def eval_clauses(cls, env, size):
    ok = np.ones(size, dtype=bool)
    for c in cls:
        cl = np.zeros(size, dtype=bool)
        for v, p in c:
            cl |= env[v] if p else ~env[v]
        ok &= cl
    return ok


def _fresh_prefix(used, base):
    p = base
    while any(u.startswith(p) for u in used):
        p += "_"
    return p


def tseitin(n):
    used = set(variables(n))
    prefix = _fresh_prefix(used, "aux")
    clauses, gates, counter = [], [], [0]
    neg = lambda l: (l[0], not l[1])
    sym = {"and": "AND", "or": "OR", "imp": "->", "iff": "<->", "xor": "XOR"}

    def enc(x):
        k = x[0]
        if k == "var":
            return (x[1], True)
        if k == "not":
            return neg(enc(x[1]))
        counter[0] += 1
        t = (f"{prefix}{counter[0]}", True)
        if k == "const":
            clauses.append(frozenset([t if x[1] else neg(t)]))
            gates.append({"aux_variable": t[0], "definition": f"{t[0]} <-> {'TRUE' if x[1] else 'FALSE'}"})
            return t
        a, b = enc(x[1]), enc(x[2])
        if k == "and":
            cl = [[neg(t), a], [neg(t), b], [t, neg(a), neg(b)]]
        elif k == "or":
            cl = [[t, neg(a)], [t, neg(b)], [neg(t), a, b]]
        elif k == "imp":
            cl = [[t, a], [t, neg(b)], [neg(t), neg(a), b]]
        elif k == "iff":
            cl = [[neg(t), neg(a), b], [neg(t), a, neg(b)], [t, a, b], [t, neg(a), neg(b)]]
        else:
            cl = [[neg(t), a, b], [neg(t), neg(a), neg(b)], [t, neg(a), b], [t, a, neg(b)]]
        clauses.extend(frozenset(c) for c in cl)
        gates.append({"aux_variable": t[0], "definition": f"{t[0]} <-> ({lit_str(a)} {sym[k]} {lit_str(b)})",
                      "clauses_added": len(cl)})
        return t

    root = enc(n)
    clauses.append(frozenset([root]))
    return clauses, gates


def to_3cnf(cls, used):
    prefix = _fresh_prefix(set(used), "y")
    counter = [0]

    def fresh():
        counter[0] += 1
        return f"{prefix}{counter[0]}"

    out = []
    for c in cls:
        lits = sorted(c, key=lambda l: natkey(l[0]))
        L = len(lits)
        if L == 0:
            y, z, w = fresh(), fresh(), fresh()
            for s in range(8):
                out.append(((y, bool(s & 1)), (z, bool(s & 2)), (w, bool(s & 4))))
        elif L == 1:
            y, z = fresh(), fresh()
            for py in (True, False):
                for pz in (True, False):
                    out.append((lits[0], (y, py), (z, pz)))
        elif L == 2:
            y = fresh()
            out += [(lits[0], lits[1], (y, True)), (lits[0], lits[1], (y, False))]
        elif L == 3:
            out.append(tuple(lits))
        else:
            y = fresh()
            out.append((lits[0], lits[1], (y, True)))
            for i in range(2, L - 2):
                y2 = fresh()
                out.append(((y, False), lits[i], (y2, True)))
                y = y2
            out.append(((y, False), lits[-2], lits[-1]))
    return out


def to_dimacs(cls):
    names = sorted({v for c in cls for v, _ in c}, key=natkey)
    idx = {v: i + 1 for i, v in enumerate(names)}
    lines = [f"p cnf {len(names)} {len(cls)}"]
    for c in cls:
        lines.append(" ".join(str(idx[v] if p else -idx[v]) for v, p in c) + " 0")
    return lines, names


def cnf_pipeline(expr_text, limit=20000):
    ast = parse_expr(expr_text)
    s1 = elim_iff_xor(ast)
    s2 = elim_imp(s1)
    s3 = simp_const(to_nnf(s2))
    steps = [("1. Original expression", to_str(ast)),
             ("2. Replace biconditional and XOR with implications", to_str(s1)),
             ("3. Replace implications: P -> Q becomes NOT P OR Q", to_str(s2)),
             ("4. Push NOT inward (De Morgan, double negation) to get NNF", to_str(s3))]
    try:
        raw = clauses_of(s3, limit)
        steps.append(("5. Distribute OR over AND", cnf_str(raw)))
        final = simplify_clauses(raw)
        steps.append(("6. Remove duplicate, tautological and subsumed clauses", cnf_str(final)))
        blow = False
    except CNFBlowup:
        raw, final, blow = None, None, True
        steps.append(("5. Distribute OR over AND", f"Stopped: more than {limit:,} clauses (exponential growth). "
                                                    "Use the Tseitin encoding instead."))
    return ast, steps, final, blow


# =====================================================================
# MDP ENGINE
# =====================================================================
GRID_ACTIONS = ["Up", "Down", "Left", "Right", "Wait"]
MOVES = {"Up": (-1, 0), "Down": (1, 0), "Left": (0, -1), "Right": (0, 1), "Wait": (0, 0)}
PERP = {"Up": ["Left", "Right"], "Down": ["Left", "Right"], "Left": ["Up", "Down"],
        "Right": ["Up", "Down"], "Wait": []}


def _grid_path_exists(rows, cols, blocked, start, goal):
    from collections import deque
    q, seen = deque([start]), {start}
    while q:
        r, c = q.popleft()
        if (r, c) == goal:
            return True
        for dr, dc in ((1, 0), (-1, 0), (0, 1), (0, -1)):
            nr, nc = r + dr, c + dc
            if 0 <= nr < rows and 0 <= nc < cols and (nr, nc) not in blocked and (nr, nc) not in seen:
                seen.add((nr, nc))
                q.append((nr, nc))
    return False


def build_grid(rows, cols, obstacle_frac, hazard_frac, slip, step_r, wait_r, goal_r, hazard_r, seed):
    rng = np.random.default_rng(seed)
    cells = [(r, c) for r in range(rows) for c in range(cols)]
    start, goal = (0, 0), (rows - 1, cols - 1)
    free = [x for x in cells if x not in (start, goal)]
    k_obs, k_haz = int(obstacle_frac * len(cells)), int(hazard_frac * len(cells))
    for _ in range(60):
        perm = rng.permutation(len(free))
        obstacles = {free[i] for i in perm[:k_obs]}
        hazards = {free[i] for i in perm[k_obs:k_obs + k_haz]}
        if _grid_path_exists(rows, cols, obstacles | hazards, start, goal):
            break
    else:
        obstacles, hazards = set(), set()
    states = [x for x in cells if x not in obstacles]
    idx = {s: i for i, s in enumerate(states)}
    S, A = len(states), len(GRID_ACTIONS)
    P, R = np.zeros((S, A, S)), np.zeros((S, A, S))
    terminal = np.zeros(S, dtype=bool)
    for s in states:
        i = idx[s]
        if s == goal or s in hazards:
            P[i, :, i] = 1.0
            terminal[i] = True
            continue
        for ai, a in enumerate(GRID_ACTIONS):
            outcomes = [(a, 1 - slip)] + [(p, slip / 2) for p in PERP[a]] if a != "Wait" else [("Wait", 1.0)]
            for mv, pr in outcomes:
                dr, dc = MOVES[mv]
                ns = (s[0] + dr, s[1] + dc)
                if not (0 <= ns[0] < rows and 0 <= ns[1] < cols) or ns in obstacles:
                    ns = s
                j = idx[ns]
                r = wait_r if a == "Wait" else step_r
                if ns == goal:
                    r += goal_r
                elif ns in hazards:
                    r += hazard_r
                P[i, ai, j] += pr
                R[i, ai, j] = r
    names = [f"({r},{c})" for r, c in states]
    return dict(P=P, R=R, states=names, actions=GRID_ACTIONS, avail=np.ones((S, A), dtype=bool),
                terminal=terminal, start=idx[start], grid=dict(rows=rows, cols=cols, cells=states, idx=idx,
                                                                obstacles=obstacles, hazards=hazards, goal=goal,
                                                                start=start))


def random_mdp(nS, nA, branch, seed):
    rng = np.random.default_rng(seed)
    act_pool = ["Monitor", "Patch", "Isolate", "Scan", "Restore", "Escalate", "Wait", "Upgrade"]
    states = [f"S{i}" for i in range(nS - 1)] + ["Goal"]
    actions = act_pool[:nA] if nA <= len(act_pool) else [f"A{i}" for i in range(nA)]
    P, R = np.zeros((nS, nA, nS)), np.zeros((nS, nA, nS))
    avail = np.zeros((nS, nA), dtype=bool)
    for s in range(nS - 1):
        for a in range(nA):
            nxt = rng.choice(nS, size=min(branch, nS), replace=False)
            pr = rng.dirichlet(np.ones(len(nxt)))
            base = rng.normal(-1, 1)
            for j, p in zip(nxt, pr):
                P[s, a, j] = p
                R[s, a, j] = round(base + (20 if j == nS - 1 else 0) + rng.normal(0, 0.5), 2)
            avail[s, a] = True
    terminal = ~avail.any(1)
    return dict(P=P, R=R, states=states, actions=actions, avail=avail, terminal=terminal, start=0, grid=None)


def incident_template():
    rows = [
        ("Normal", "Monitor", "Normal", 0.9, 1), ("Normal", "Monitor", "Suspicious", 0.1, 0),
        ("Normal", "Patch", "Normal", 1.0, -0.5),
        ("Suspicious", "Monitor", "Normal", 0.3, 0), ("Suspicious", "Monitor", "Suspicious", 0.4, -1),
        ("Suspicious", "Monitor", "Compromised", 0.3, -20),
        ("Suspicious", "Scan", "Normal", 0.7, -1), ("Suspicious", "Scan", "Suspicious", 0.2, -1),
        ("Suspicious", "Scan", "Compromised", 0.1, -20),
        ("Suspicious", "Isolate", "Contained", 0.9, -5), ("Suspicious", "Isolate", "Compromised", 0.1, -20),
        ("Compromised", "Isolate", "Contained", 0.8, -10), ("Compromised", "Isolate", "Compromised", 0.2, -30),
        ("Compromised", "Monitor", "Compromised", 1.0, -40),
        ("Contained", "Restore", "Recovered", 0.85, 10), ("Contained", "Restore", "Contained", 0.15, -2),
        ("Contained", "Scan", "Contained", 1.0, -1),
    ]
    return pd.DataFrame(rows, columns=["state", "action", "next_state", "probability", "reward"])


def mdp_from_table(df):
    cols = {c.lower().strip().replace(" ", "_"): c for c in df.columns}
    need = ["state", "action", "next_state", "probability", "reward"]
    miss = [c for c in need if c not in cols]
    if miss:
        raise ValueError("Missing columns: " + ", ".join(miss) + ". Required: " + ", ".join(need))
    d = df[[cols[c] for c in need]].copy()
    d.columns = need
    d = d.dropna()
    d["state"] = d["state"].astype(str)
    d["next_state"] = d["next_state"].astype(str)
    d["action"] = d["action"].astype(str)
    d["probability"] = pd.to_numeric(d["probability"], errors="coerce")
    d["reward"] = pd.to_numeric(d["reward"], errors="coerce")
    d = d.dropna()
    states = list(dict.fromkeys(list(d["state"]) + list(d["next_state"])))
    actions = list(dict.fromkeys(d["action"]))
    si = {s: i for i, s in enumerate(states)}
    ai = {a: i for i, a in enumerate(actions)}
    S, A = len(states), len(actions)
    P, RW = np.zeros((S, A, S)), np.zeros((S, A, S))
    for r in d.itertuples(index=False):
        i, a, j = si[r.state], ai[r.action], si[r.next_state]
        P[i, a, j] += r.probability
        RW[i, a, j] += r.probability * r.reward
    R = np.divide(RW, P, out=np.zeros_like(RW), where=P > 0)
    tot = P.sum(2)
    warnings = []
    bad = (tot > 0) & (np.abs(tot - 1) > 1e-6)
    if bad.any():
        warnings.append(f"{int(bad.sum())} state-action pairs had probabilities not summing to 1 and were normalised.")
    P = np.divide(P, tot[:, :, None], out=np.zeros_like(P), where=tot[:, :, None] > 0)
    avail = tot > 0
    absorbing = np.array([avail[s].any() and all(P[s, a, s] > 0.999999 and abs(R[s, a, s]) < 1e-12
                                                 for a in range(A) if avail[s, a]) for s in range(S)])
    terminal = ~avail.any(1) | absorbing
    return dict(P=P, R=R, states=states, actions=actions, avail=avail, terminal=terminal, start=0,
                grid=None), warnings


def expected_reward(P, R):
    return np.einsum("ijk,ijk->ij", P, R)


def q_values(P, ER, V, gamma, avail):
    Q = ER + gamma * np.einsum("ijk,k->ij", P, V)
    return np.where(avail, Q, -np.inf)


def value_iteration(m, gamma, theta=1e-6, max_iter=5000):
    P, avail = m["P"], m["avail"]
    ER = expected_reward(P, m["R"])
    V = np.zeros(P.shape[0])
    hist = []
    has = avail.any(1)
    for _ in range(max_iter):
        Q = q_values(P, ER, V, gamma, avail)
        Vn = np.where(has, Q.max(1, initial=-np.inf, where=avail), 0.0)
        delta = float(np.max(np.abs(Vn - V)))
        hist.append(delta)
        V = Vn
        if delta < theta:
            break
    Q = q_values(P, ER, V, gamma, avail)
    pi = np.where(has, np.argmax(Q, 1), -1)
    return V, Q, pi, hist


def policy_iteration(m, gamma, max_iter=500):
    P, avail = m["P"], m["avail"]
    S = P.shape[0]
    ER = expected_reward(P, m["R"])
    has = avail.any(1)
    pi = np.where(has, np.argmax(avail, 1), 0)
    hist = []
    for _ in range(max_iter):
        Ppi = P[np.arange(S), pi] * has[:, None]
        Rpi = ER[np.arange(S), pi] * has
        V = np.linalg.solve(np.eye(S) - gamma * Ppi, Rpi)
        Q = q_values(P, ER, V, gamma, avail)
        best = np.argmax(Q, 1)
        improve = has & (Q[np.arange(S), best] > Q[np.arange(S), pi] + 1e-9)
        hist.append(int(improve.sum()))
        if not improve.any():
            break
        pi = np.where(improve, best, pi)
    return V, Q, np.where(has, pi, -1), hist


def q_learning(m, gamma, episodes, alpha, eps_start, eps_end, max_steps, seed):
    rng = np.random.default_rng(seed)
    P, R, avail, term = m["P"], m["R"], m["avail"], m["terminal"]
    S, A = avail.shape
    Q = np.zeros((S, A))
    cumP = P.cumsum(2)
    returns = []
    starts = [s for s in range(S) if not term[s]]
    for ep in range(episodes):
        eps = eps_start + (eps_end - eps_start) * ep / max(1, episodes - 1)
        s = m["start"] if not term[m["start"]] else int(rng.choice(starts))
        G = 0.0
        for _ in range(max_steps):
            if term[s]:
                break
            acts = np.nonzero(avail[s])[0]
            a = int(rng.choice(acts)) if rng.random() < eps else int(acts[np.argmax(Q[s, acts])])
            u = rng.random() * cumP[s, a, -1]
            s2 = min(int(np.searchsorted(cumP[s, a], u, side="right")), S - 1)
            r = R[s, a, s2]
            nxt = 0.0 if term[s2] or not avail[s2].any() else gamma * Q[s2, avail[s2]].max()
            Q[s, a] += alpha * (r + nxt - Q[s, a])
            G += r
            s = s2
        returns.append(G)
    Qm = np.where(avail, Q, -np.inf)
    pi = np.where(avail.any(1), np.argmax(Qm, 1), -1)
    return Q, pi, returns


def rollout(m, pi, max_steps, seed):
    rng = np.random.default_rng(seed)
    s, path, G = m["start"], [m["start"]], 0.0
    for _ in range(max_steps):
        if m["terminal"][s] or pi[s] < 0:
            break
        a = pi[s]
        s2 = int(rng.choice(len(m["states"]), p=m["P"][s, a]))
        G += m["R"][s, a, s2]
        s = s2
        path.append(s)
    return path, G


def grid_figure(m, V, pi, path=None, title="Optimal values and policy"):
    g = m["grid"]
    rows, cols = g["rows"], g["cols"]
    Z = np.full((rows, cols), np.nan)
    for s, i in g["idx"].items():
        Z[s] = V[i]
    fig, ax = plt.subplots(figsize=(min(1.0 * cols + 2, 12), min(0.9 * rows + 1.5, 10)))
    cmap = plt.get_cmap("Blues").copy()
    cmap.set_bad("#4A4A4A")
    im = ax.imshow(np.ma.masked_invalid(Z), cmap=cmap)
    fig.colorbar(im, ax=ax, fraction=0.04, label="State value V(s)")
    for s, i in g["idx"].items():
        r, c = s
        if s == g["goal"]:
            ax.text(c, r, "GOAL", ha="center", va="center", fontsize=8, color="white", weight="bold",
                    bbox=dict(fc=GREEN, ec="none"))
            continue
        if s in g["hazards"]:
            ax.text(c, r, "HAZARD", ha="center", va="center", fontsize=7, color="white", weight="bold",
                    bbox=dict(fc=RED, ec="none"))
            continue
        a = GRID_ACTIONS[pi[i]]
        dr, dc = MOVES[a]
        if a == "Wait":
            ax.plot(c, r, "o", color=ORANGE, ms=6)
        else:
            ax.annotate("", xy=(c + 0.32 * dc, r + 0.32 * dr), xytext=(c - 0.32 * dc, r - 0.32 * dr),
                        arrowprops=dict(arrowstyle="->", color=ORANGE, lw=2))
        if rows * cols <= 100:
            ax.text(c, r + 0.38, f"{V[i]:.1f}", ha="center", va="center", fontsize=6, color="#222")
    if path and len(path) > 1:
        pts = [g["cells"][p] for p in path]
        ax.plot([p[1] for p in pts], [p[0] for p in pts], color=GREEN, lw=2, alpha=0.7, label="Simulated path")
        ax.legend(loc="upper right", fontsize=7)
    sr, sc = g["start"]
    ax.text(sc - 0.42, sr - 0.42, "START", fontsize=6, color=GREEN, weight="bold", va="top")
    ax.set_xticks(range(cols))
    ax.set_yticks(range(rows))
    ax.grid(False)
    ax.set_title(title)
    return fig


def mdp_graph_dot(m, pi=None, max_states=14):
    S = len(m["states"])
    if S > max_states:
        return None
    lines = ['digraph M {', 'rankdir=LR;', 'node [fontname="Helvetica", shape=circle, style=filled, '
             'fillcolor="#DCE6F5", color="#0B4FA8"];', 'edge [fontname="Helvetica", fontsize=9];']
    for i, s in enumerate(m["states"]):
        if m["terminal"][i]:
            lines.append(f'"{s}" [shape=doublecircle, fillcolor="#E8F5EE", color="#2E8B57"];')
    for i in range(S):
        for a in range(len(m["actions"])):
            if not m["avail"][i, a]:
                continue
            for j in np.nonzero(m["P"][i, a] > 1e-9)[0]:
                if m["terminal"][i]:
                    continue
                opt = pi is not None and pi[i] == a
                color = "#E07B00" if opt else "#8A94A6"
                lines.append(f'"{m["states"][i]}" -> "{m["states"][j]}" [label="{m["actions"][a]} '
                             f'{m["P"][i, a, j]:.2f} | {m["R"][i, a, j]:+.1f}", color="{color}", '
                             f'fontcolor="{color}", penwidth={2.2 if opt else 1}];')
    lines.append("}")
    return "\n".join(lines)


def run_mdp(m, gamma, algos, ql, source, warnings=()):
    rep = Report("Markov Decision Process Analysis", "MDP")
    S, A = len(m["states"]), len(m["actions"])
    rep.heading("Model summary")
    rep.para(f"Data source: {source}.")
    for w in warnings:
        rep.para("Warning: " + w)
    met = {"States": S, "Actions": A, "Discount factor gamma": gamma,
           "Terminal states": int(m["terminal"].sum())}
    res = {}
    t0 = time.perf_counter()
    if "Value Iteration" in algos:
        res["VI"] = value_iteration(m, gamma)
        met["Value iteration sweeps"] = len(res["VI"][3])
    if "Policy Iteration" in algos:
        res["PI"] = policy_iteration(m, gamma)
        met["Policy iteration rounds"] = len(res["PI"][3])
    if "Q-Learning" in algos:
        res["QL"] = q_learning(m, gamma, **ql)
    met["Compute time (ms)"] = f"{(time.perf_counter() - t0) * 1000:.0f}"
    ref = res.get("VI") or res.get("PI")
    if ref is not None and "QL" in res:
        nt = ~m["terminal"] & m["avail"].any(1)
        agree = (res["QL"][1][nt] == ref[2][nt]).mean() * 100 if nt.any() else 100
        met["Q-learning agreement with optimal policy"] = f"{agree:.0f}%"
    if ref is not None:
        met["Value of start state"] = f"{ref[0][m['start']]:.2f}"
        met["Start state"] = m["states"][m["start"]]
    rep.metrics(met)
    rep.latex(r"V^*(s)=\max_a \sum_{s'} P(s'|s,a)\,[R(s,a,s')+\gamma V^*(s')]",
              "V*(s) = max_a sum_s' P(s'|s,a) [ R(s,a,s') + gamma V*(s') ]")

    rows = []
    for i, s in enumerate(m["states"]):
        row = {"state": s, "terminal": bool(m["terminal"][i])}
        for key, lab in (("VI", "value_iteration"), ("PI", "policy_iteration")):
            if key in res:
                V, _, pi, _ = res[key]
                row[f"V_{lab}"] = round(float(V[i]), 4)
                row[f"policy_{lab}"] = m["actions"][pi[i]] if pi[i] >= 0 and not m["terminal"][i] else "-"
        if "QL" in res:
            Qq, piq, _ = res["QL"]
            row["policy_q_learning"] = m["actions"][piq[i]] if piq[i] >= 0 and not m["terminal"][i] else "-"
            row["max_Q_q_learning"] = round(float(np.max(np.where(m["avail"][i], Qq[i], -np.inf))), 4) \
                if m["avail"][i].any() else 0.0
        rows.append(row)
    rep.heading("Optimal values and policy")
    rep.table("State values and policies", pd.DataFrame(rows))
    if ref is not None:
        Qd = pd.DataFrame(np.where(np.isinf(ref[1]), np.nan, ref[1]).round(4), columns=m["actions"])
        Qd.insert(0, "state", m["states"])
        rep.table("Action values Q(s,a)", Qd)

    if m["grid"] is not None and ref is not None:
        path, G = rollout(m, ref[2], 200, 1)
        rep.figure("Grid world: optimal state values with policy arrows",
                   grid_figure(m, ref[0], ref[2], path))
        rep.para(f"A simulated episode following the optimal policy took {len(path) - 1} steps and collected a "
                 f"total reward of {G:.1f}. Orange arrows show the best action in each cell; dots mean Wait.")
        rep.table("Simulated episode path", pd.DataFrame({"step": range(len(path)),
                                                           "state": [m["states"][p] for p in path]}))
    elif ref is not None:
        fig, ax = plt.subplots(figsize=(min(0.5 * S + 3, 12), 3.4))
        ax.bar(range(S), ref[0], color=[GREEN if t else BLUE for t in m["terminal"]])
        ax.set_xticks(range(S))
        ax.set_xticklabels(m["states"], rotation=45, ha="right", fontsize=8)
        ax.set_ylabel("V*(s)")
        ax.set_title("Optimal state values (green = terminal)")
        rep.figure("Optimal state values", fig)

    fig, axes = plt.subplots(1, 2 if "QL" in res else 1, figsize=(11 if "QL" in res else 6, 3.3))
    axes = np.atleast_1d(axes)
    ax = axes[0]
    if "VI" in res:
        ax.semilogy(np.maximum(res["VI"][3], 1e-12), color=BLUE, label="Value iteration max change")
    if "PI" in res:
        ax.plot(range(len(res["PI"][3])), np.maximum(res["PI"][3], 1e-12), "o-", color=ORANGE,
                label="Policy iteration states changed")
        ax.set_yscale("log")
    ax.set_xlabel("Iteration")
    ax.set_title("Convergence")
    ax.legend(fontsize=8)
    if "QL" in res:
        rets = np.array(res["QL"][2])
        w = max(1, len(rets) // 30)
        sm = np.convolve(rets, np.ones(w) / w, mode="valid")
        axes[1].plot(rets, color=LIGHT, lw=0.8)
        axes[1].plot(np.arange(len(sm)) + w - 1, sm, color=BLUE, lw=2)
        axes[1].set_xlabel("Episode")
        axes[1].set_ylabel("Total reward")
        axes[1].set_title("Q-learning learning curve")
    rep.figure("Convergence and learning curves", fig)
    if "VI" in res:
        rep.table("Value iteration convergence", pd.DataFrame({"iteration": range(1, len(res["VI"][3]) + 1),
                                                                "max_value_change": res["VI"][3]}))
    tr = [(m["states"][i], m["actions"][a], m["states"][j], round(m["P"][i, a, j], 4), round(m["R"][i, a, j], 3))
          for i in range(S) for a in range(A) if m["avail"][i, a] and not m["terminal"][i]
          for j in np.nonzero(m["P"][i, a] > 1e-12)[0]]
    rep.table("Transition model", pd.DataFrame(tr[:20000], columns=["state", "action", "next_state",
                                                                   "probability", "reward"]))
    return rep, res


# =====================================================================
# PCA ENGINE
# =====================================================================
def synth_traffic(n, anomaly_frac, extra_noise, seed):
    rng = np.random.default_rng(seed)
    size = rng.lognormal(0, 0.5, n)
    attack = rng.normal(0, 1, n)
    df = pd.DataFrame({
        "packets": 150 * size + rng.normal(0, 10, n),
        "data_volume_mb": 15 * size + rng.normal(0, 1.5, n),
        "duration_s": 8 * size + rng.normal(0, 1, n),
        "failed_logins": np.clip(1 + 0.8 * attack + rng.normal(0, 0.5, n), 0, None).round(),
        "unusual_requests": np.clip(2 + 1.2 * attack + rng.normal(0, 0.7, n), 0, None).round(),
    })
    for k in range(extra_noise):
        df[f"noise_{k + 1}"] = rng.normal(0, 1, n)
    label = np.array(["Normal"] * n, dtype=object)
    k = int(anomaly_frac * n)
    if k:
        idx = rng.choice(n, k, replace=False)
        df.loc[idx, "failed_logins"] += rng.integers(6, 15, k)
        df.loc[idx, "unusual_requests"] += rng.integers(8, 25, k)
        df.loc[idx, "packets"] *= rng.uniform(0.2, 0.5, k)
        label[idx] = "Anomaly"
    df = df.round(2)
    df.insert(0, "connection_id", [f"C{i:05d}" for i in range(1, n + 1)])
    df["label"] = label
    return df


def run_pca(df, features, label_col, standardize, k_mode, k_value, var_target, pct, source):
    rep = Report("Principal Component Analysis", "PCA")
    X = df[features].apply(pd.to_numeric, errors="coerce")
    n_missing = int(X.isna().sum().sum())
    X = X.fillna(X.mean())
    keep = X.std(ddof=0) > 0
    dropped = [f for f, k in zip(features, keep) if not k]
    X = X.loc[:, keep]
    feats = list(X.columns)
    if len(feats) < 2:
        raise ValueError("Need at least two numeric, non-constant feature columns.")
    if len(X) < 3:
        raise ValueError("Need at least three rows.")
    Xs = StandardScaler().fit_transform(X) if standardize else (X - X.mean()).values
    full = PCA().fit(Xs)
    evr = full.explained_variance_ratio_
    cum = np.cumsum(evr)
    if k_mode == "fixed":
        k = int(min(k_value, len(feats)))
    else:
        k = int(np.searchsorted(cum, var_target) + 1)
        k = min(k, len(feats))
    comps = full.components_[:k]
    scores = Xs @ comps.T
    recon = scores @ comps
    spe = ((Xs - recon) ** 2).sum(1)
    t2 = (scores ** 2 / full.explained_variance_[:k]).sum(1)
    thr = np.percentile(spe, pct)
    flag = spe > thr

    rep.heading("Summary")
    rep.para(f"Data source: {source}. Features used: {', '.join(feats)}."
             + (f" Constant columns dropped: {', '.join(dropped)}." if dropped else "")
             + (f" {n_missing} missing values were replaced with column means." if n_missing else ""))
    rep.metrics({"Rows": len(X), "Original features": len(feats), "Components kept": k,
                 "Variance retained": f"{cum[k - 1] * 100:.1f}%", "Standardized": "Yes" if standardize else "No",
                 "Anomaly threshold (SPE percentile)": pct, "Rows flagged": int(flag.sum())})
    rep.latex(r"Z = X_{std} W_k,\qquad \hat X = Z W_k^T,\qquad SPE_i=\lVert x_i-\hat x_i\rVert^2",
              "Z = X_std W_k ;  X_hat = Z W_k^T ;  SPE_i = ||x_i - x_hat_i||^2")
    ev = pd.DataFrame({"component": [f"PC{i + 1}" for i in range(len(evr))],
                       "eigenvalue": full.explained_variance_, "variance_explained_pct": evr * 100,
                       "cumulative_pct": cum * 100, "kept": [i < k for i in range(len(evr))]})
    rep.table("Explained variance", ev)
    load = pd.DataFrame(comps.T, columns=[f"PC{i + 1}" for i in range(k)])
    load.insert(0, "feature", feats)
    rep.table("Component loadings (weights)", load)
    for i in range(min(k, 3)):
        top = load.reindex(load[f"PC{i + 1}"].abs().sort_values(ascending=False).index)["feature"].head(3)
        rep.para(f"PC{i + 1} ({evr[i] * 100:.1f}% of variance) is driven mostly by: {', '.join(top)}.")

    fig, ax = plt.subplots(figsize=(8, 3.4))
    xs = np.arange(1, len(evr) + 1)
    ax.bar(xs, evr * 100, color=[BLUE if i < k else LIGHT for i in range(len(evr))], label="Per component")
    ax.plot(xs, cum * 100, "o-", color=ORANGE, label="Cumulative")
    ax.axhline(cum[k - 1] * 100, ls="--", color=GREY, lw=1)
    ax.set_xlabel("Principal component")
    ax.set_ylabel("Variance explained (%)")
    ax.set_title(f"Scree plot: {k} components keep {cum[k - 1] * 100:.1f}% of variance")
    ax.set_xticks(xs)
    ax.legend(fontsize=8)
    rep.figure("Scree and cumulative variance", fig)

    lab = df[label_col].astype(str).values if label_col else np.where(flag, "Flagged", "Not flagged")
    fig, ax = plt.subplots(1, 2, figsize=(12, 4.6))
    pc2 = scores[:, 1] if k > 1 else np.zeros(len(scores))
    palette = [BLUE, ORANGE, GREEN, RED, "#7B4FA8", GREY, "#B08800", "#008B8B"]
    for gi, g in enumerate(pd.unique(lab)[:8]):
        mk = lab == g
        ax[0].scatter(scores[mk, 0], pc2[mk], s=14, alpha=0.75, color=palette[gi % 8], label=str(g))
    ax[0].set_xlabel(f"PC1 ({evr[0] * 100:.1f}%)")
    ax[0].set_ylabel(f"PC2 ({evr[1] * 100:.1f}%)" if k > 1 else "PC2 (not kept)")
    ax[0].set_title("Observations projected on PC1 and PC2")
    ax[0].legend(fontsize=7)
    full2 = full.components_[:2]
    sc = np.abs(scores[:, :1]).max() if k else 1
    for j, f in enumerate(feats[:20]):
        ax[1].arrow(0, 0, full2[0, j], full2[1, j], color=BLUE, head_width=0.02, length_includes_head=True)
        ax[1].text(full2[0, j] * 1.1, full2[1, j] * 1.1, f, fontsize=8, ha="center")
    circ = plt.Circle((0, 0), 1, fill=False, ls="--", color=GREY)
    ax[1].add_patch(circ)
    ax[1].set_xlim(-1.2, 1.2)
    ax[1].set_ylim(-1.2, 1.2)
    ax[1].set_aspect("equal")
    ax[1].set_title("Loading plot (feature directions)")
    ax[1].set_xlabel("PC1 weight")
    ax[1].set_ylabel("PC2 weight")
    rep.figure("Score plot and loading plot", fig)

    fig, ax = plt.subplots(figsize=(min(1 + 0.9 * k, 10), max(2.5, 0.4 * len(feats) + 1)))
    im = ax.imshow(comps.T, cmap="RdBu_r", vmin=-1, vmax=1, aspect="auto")
    ax.set_yticks(range(len(feats)))
    ax.set_yticklabels(feats, fontsize=8)
    ax.set_xticks(range(k))
    ax.set_xticklabels([f"PC{i + 1}" for i in range(k)])
    for (i, j), v in np.ndenumerate(comps.T):
        if len(feats) * k <= 200:
            ax.text(j, i, f"{v:.2f}", ha="center", va="center", fontsize=7)
    ax.grid(False)
    fig.colorbar(im, ax=ax, fraction=0.05)
    ax.set_title("Loadings heatmap")
    rep.figure("Loadings heatmap", fig)

    fig, ax = plt.subplots(figsize=(8, 3.2))
    ax.hist(spe, bins=40, color=BLUE, alpha=0.8)
    ax.axvline(thr, color=RED, ls="--", label=f"{pct}th percentile threshold")
    ax.set_xlabel("Reconstruction error (SPE)")
    ax.set_ylabel("Rows")
    ax.set_title("Reconstruction error distribution (high error = unusual row)")
    ax.legend(fontsize=8)
    rep.figure("Reconstruction error histogram", fig)

    out = pd.DataFrame(scores.round(5), columns=[f"PC{i + 1}" for i in range(k)])
    id_cols = [c for c in df.columns if c not in features and c != label_col][:2]
    for c in reversed(id_cols):
        out.insert(0, c, df[c].values)
    if label_col:
        out["label"] = df[label_col].values
    out["reconstruction_error_SPE"] = spe.round(5)
    out["hotelling_T2"] = t2.round(5)
    out["flagged_anomaly"] = flag
    rep.table("Transformed data (component scores and anomaly scores)", out)
    rep.table("Top 25 most unusual rows", out.sort_values("reconstruction_error_SPE", ascending=False).head(25))
    if label_col:
        ct = pd.crosstab(pd.Series(df[label_col].astype(str).values, name="label"),
                         pd.Series(np.where(flag, "Flagged", "Not flagged"), name="pca_flag"))
        ct.columns.name = None
        rep.table("Label versus flag (cross tabulation)", ct.reset_index())
    rep.para("Reminder: PCA keeps directions of high variance, not necessarily the directions that best predict a "
             "target. A low-variance feature can still matter for fraud or attack detection.")
    return rep


# =====================================================================
# GROVER ENGINE
# =====================================================================
MAX_QUBITS = 14


def grover_sim(n, marked, k_max):
    N = 1 << n
    mask = np.zeros(N, dtype=bool)
    mask[list(marked)] = True
    psi = np.full(N, 1 / np.sqrt(N))
    states, probs = [psi.copy()], [float((psi[mask] ** 2).sum())]
    for _ in range(k_max):
        psi[mask] *= -1
        psi = 2 * psi.mean() - psi
        states.append(psi.copy())
        probs.append(float((psi[mask] ** 2).sum()))
    return states, np.array(probs), mask


def grover_theory(N, M):
    if M == 0 or M >= N:
        return 0.0, 0
    theta = math.asin(math.sqrt(M / N))
    k = max(0, int(round(math.pi / (4 * theta) - 0.5)))
    return theta, k


def grover_circuit_fig(n, k):
    nq = min(n, 6)
    blocks = ["H"] + ["O", "D"] * min(k, 2) + (["R"] if k > 2 else []) + ["M"]
    fig, ax = plt.subplots(figsize=(11, 0.55 * nq + 1.6))
    xs = np.arange(len(blocks)) * 1.5 + 1
    ys = np.arange(nq)[::-1]
    for i, y in enumerate(ys):
        ax.plot([0.2, xs[-1] + 0.8], [y, y], color="#333", lw=1)
        ax.text(0, y, f"q{i}  |0>", ha="right", va="center", fontsize=9)
    if n > 6:
        ax.text(0.6, -0.8, f"... {n} qubits in total (first 6 drawn)", fontsize=8, color=GREY)
    for b, x in zip(blocks, xs):
        if b in ("H", "M"):
            for y in ys:
                fc = LIGHT if b == "H" else "#F4E3CC"
                ax.add_patch(Rectangle((x - 0.25, y - 0.25), 0.5, 0.5, fc=fc, ec=BLUE if b == "H" else ORANGE))
                ax.text(x, y, "H" if b == "H" else "M", ha="center", va="center", fontsize=9, weight="bold")
        elif b == "R":
            ax.text(x, (nq - 1) / 2, f"repeat\n{k} times\nin total", ha="center", va="center", fontsize=8,
                    bbox=dict(fc="white", ec=GREY, boxstyle="round"))
        else:
            fc, ec, lab = ("#FBE4E1", RED, "Oracle\nUf") if b == "O" else ("#E8F5EE", GREEN, "Diffusion\nD")
            ax.add_patch(FancyBboxPatch((x - 0.45, -0.4), 0.9, nq - 0.2, boxstyle="round,pad=0.02", fc=fc, ec=ec))
            ax.text(x, (nq - 1) / 2, lab, ha="center", va="center", fontsize=8, weight="bold")
    ax.set_xlim(-1.2, xs[-1] + 1)
    ax.set_ylim(-1.1, nq)
    ax.axis("off")
    ax.set_title(f"Grover circuit: {n} qubits, {k} iterations of (Oracle then Diffusion)")
    return fig


def grover_geometry_fig(theta, k):
    fig, ax = plt.subplots(figsize=(5.2, 5))
    t = np.linspace(0, np.pi / 2, 100)
    ax.plot(np.cos(t), np.sin(t), color=GREY, lw=1)
    shown = list(range(0, min(k, 8) + 1))
    for j in shown:
        ang = (2 * j + 1) * theta
        col = BLUE if j < k else ORANGE
        ax.annotate("", xy=(math.cos(ang), math.sin(ang)), xytext=(0, 0),
                    arrowprops=dict(arrowstyle="->", color=col, lw=1.8 if j in (0, k) else 1))
        ax.text(1.05 * math.cos(ang), 1.05 * math.sin(ang), f"k={j}", fontsize=8, color=col)
    ax.set_xlabel("Amplitude on unmarked states |alpha>")
    ax.set_ylabel("Amplitude on marked states |beta>")
    ax.set_xlim(-0.1, 1.25)
    ax.set_ylim(-0.1, 1.25)
    ax.set_aspect("equal")
    ax.set_title(f"Each iteration rotates the state by 2 theta\n(theta = {math.degrees(theta):.2f} degrees)")
    return fig


def synth_database(n, seed):
    rng = np.random.default_rng(seed)
    depts = ["Finance", "HR", "IT", "Sales", "Operations", "Legal"]
    return pd.DataFrame({
        "record_id": [f"R{i:05d}" for i in range(1, n + 1)],
        "user": [f"user{rng.integers(100, 999)}" for _ in range(n)],
        "department": rng.choice(depts, n),
        "risk_score": rng.integers(1, 101, n),
        "failed_logins": rng.poisson(1.2, n),
        "status": rng.choice(["Active", "Locked", "Suspended"], n, p=[0.85, 0.1, 0.05]),
    })


def run_grover(n, marked, labels, source, shots, seed, k_choice=None):
    rep = Report("Grover Quantum Search Simulation", "Grover")
    N, M = 1 << n, len(marked)
    theta, k_opt = grover_theory(N, M)
    k_run = k_opt if k_choice is None else k_choice
    k_max = int(min(max(2 * k_opt + 3, k_run + 1, 8), 400))
    states, probs, mask = grover_sim(n, marked, k_max)
    if M >= N:
        theory = np.ones(k_max + 1)
    elif M:
        theory = np.array([math.sin((2 * k + 1) * theta) ** 2 for k in range(k_max + 1)])
    else:
        theory = np.zeros(k_max + 1)
    rng = np.random.default_rng(seed)
    p_final = states[k_run] ** 2
    p_final = p_final / p_final.sum()
    samples = rng.choice(N, size=shots, p=p_final)
    cnt = Counter(samples.tolist())
    hits = sum(v for s, v in cnt.items() if mask[s])

    rep.heading("Search summary")
    rep.para(f"Data source: {source}.")
    classical = (N + 1) / (M + 1) if M else N
    rep.metrics({"Qubits n": n, "Search space N = 2^n": f"{N:,}", "Marked items M": M,
                 "Optimal iterations k": k_opt, "Iterations run": k_run,
                 "Success probability (theory)": f"{theory[k_run] * 100:.2f}%",
                 "Success probability (simulated)": f"{probs[k_run] * 100:.2f}%",
                 "Shots that found a marked item": f"{hits} of {shots}",
                 "Classical expected queries": f"{classical:,.1f}",
                 "Grover oracle queries": k_run,
                 "Query speedup": f"{classical / max(k_run, 1):.1f}x"})
    rep.latex(r"k_{opt}=\left\lfloor \frac{\pi}{4\theta} \right\rceil,\quad \sin\theta=\sqrt{M/N},\quad "
              r"P_k=\sin^2\big((2k+1)\theta\big)",
              "k_opt = round(pi/(4 theta) - 1/2), sin(theta) = sqrt(M/N), P_k = sin^2((2k+1) theta)")
    if M == 0:
        rep.para("No item satisfies the oracle, so Grover's iterations cannot amplify anything; every measurement "
                 "returns a random unmarked state. In practice this is how Grover reports 'not found'.")
    elif M >= N / 2:
        rep.para("At least half the items are marked, so a random guess already succeeds with probability at least "
                 "50 percent and Grover gives no advantage.")
    else:
        rep.para(f"After {k_run} iterations the probability of measuring a marked item is "
                 f"{probs[k_run] * 100:.1f}%. Running more iterations than optimal makes it fall again "
                 "(overshooting), which is visible in the success probability curve.")

    fig, ax = plt.subplots(figsize=(8, 3.4))
    ax.plot(range(k_max + 1), theory, color=GREY, lw=3, alpha=0.5, label="Theory sin^2((2k+1) theta)")
    ax.plot(range(k_max + 1), probs, "o-", ms=3, color=BLUE, label="Statevector simulation")
    ax.axvline(k_opt, ls="--", color=ORANGE, label=f"Optimal k = {k_opt}")
    ax.set_xlabel("Grover iterations k")
    ax.set_ylabel("P(measure a marked item)")
    ax.set_ylim(0, 1.05)
    ax.set_title("Success probability versus iterations")
    ax.legend(fontsize=8)
    rep.figure("Success probability versus iterations", fig)

    fig, axes = plt.subplots(1, 2, figsize=(12, 3.4), sharey=True)
    for ax, kk, ttl in ((axes[0], 0, "Before Grover (uniform superposition)"),
                        (axes[1], k_run, f"After {k_run} iterations")):
        amp = states[kk]
        if N <= 64:
            ax.bar(range(N), amp, color=np.where(mask, ORANGE, BLUE))
        else:
            ax.plot(range(N), amp, color=BLUE, lw=0.6)
            mi = np.nonzero(mask)[0]
            ax.scatter(mi, amp[mi], color=ORANGE, s=14, zorder=3, label="Marked")
        ax.axhline(amp.mean(), color=GREY, ls="--", lw=1)
        ax.set_title(ttl)
        ax.set_xlabel("Basis state index")
    axes[0].set_ylabel("Amplitude (orange = marked)")
    rep.figure("Amplitudes before and after amplification", fig)
    if M:
        rep.figure("Geometric picture of amplitude amplification", grover_geometry_fig(theta, k_run))
    rep.figure("Grover circuit", grover_circuit_fig(n, k_run))

    top = cnt.most_common(25)
    rep.table("Measurement results (top 25 outcomes)", pd.DataFrame({
        "basis_state": [format(s, f"0{n}b") for s, _ in top], "index": [s for s, _ in top],
        "item": [labels(s) for s, _ in top], "marked": [bool(mask[s]) for s, _ in top],
        "count": [c for _, c in top], "frequency": [c / shots for _, c in top]}))
    fig, ax = plt.subplots(figsize=(9, 3.2))
    ax.bar([format(s, f"0{n}b") for s, _ in top[:16]], [c for _, c in top[:16]],
           color=[ORANGE if mask[s] else BLUE for s, _ in top[:16]])
    ax.set_title(f"Measurement histogram ({shots} shots, top 16 outcomes, orange = marked)")
    ax.set_ylabel("Count")
    plt.setp(ax.get_xticklabels(), rotation=60, ha="right", fontsize=7)
    rep.figure("Measurement histogram", fig)
    rep.table("Iteration table", pd.DataFrame({"iteration_k": range(k_max + 1), "p_success_simulated": probs,
                                               "p_success_theory": theory,
                                               "amplitude_marked": [float(s[mask][0]) if M else 0.0 for s in states],
                                               "amplitude_unmarked": [float(s[~mask][0]) if M < N else 0.0
                                                                      for s in states]}))
    mi = np.nonzero(mask)[0]
    rep.table("Marked items", pd.DataFrame({"index": mi, "basis_state": [format(int(i), f"0{n}b") for i in mi],
                                            "item": [labels(int(i)) for i in mi]}).head(5000))
    return rep, states, probs, mask


# =====================================================================
# PAGE: PLATFORM OVERVIEW
# =====================================================================
def page_overview():
    st.markdown("## Platform Overview")
    st.markdown(
        "This platform brings together six foundational ideas from logic, artificial intelligence, machine "
        "learning and quantum computing. Every module has the same five sections in the left navigation: "
        "**Concept and Formulas**, **Use Cases**, **Diagrams**, **Synthetic Data Lab**, **Upload Real Data** "
        "and **Results and Export**. Every result can be downloaded as PDF, Word, Text or CSV.")
    ui_dataframe(pd.DataFrame([
        ("3-SAT Satisfiability", "Can all Boolean constraints be satisfied at once?", "DPLL, WalkSAT, exhaustive check",
         "DIMACS .cnf, CSV of clauses"),
        ("Logical Equivalence", "Do two expressions always have the same truth value?", "Vectorised truth tables",
         "CSV with expr1, expr2"),
        ("CNF Conversion", "How do we rewrite rules as an AND of OR clauses?",
         "Direct conversion, Tseitin, 3-CNF, DIMACS", "CSV or TXT of expressions"),
        ("Markov Decision Process", "Which action maximises expected long-term reward?",
         "Value iteration, policy iteration, Q-learning", "CSV of transitions"),
        ("Principal Component Analysis", "How can we simplify data while keeping its main patterns?",
         "PCA with SPE and T2 anomaly scores", "Any numeric CSV or Excel"),
        ("Grover Quantum Search", "How fast can a quantum computer find a marked item?",
         "Statevector simulation", "Any CSV database"),
    ], columns=["Module", "Question it answers", "Algorithms", "Upload format"]), height=260)
    st.markdown("#### How the concepts connect")
    st.graphviz_chart("""
digraph G { rankdir=LR; node [shape=box, style="rounded,filled", fillcolor="#DCE6F5", color="#0B4FA8",
fontname="Helvetica"]; edge [color="#8A94A6", fontname="Helvetica", fontsize=10];
LE [label="Logical Equivalence"]; CNF [label="CNF Conversion"]; SAT [label="3-SAT"];
GR [label="Grover Search", fillcolor="#FBE4E1", color="#C0392B"]; MDP [label="MDP"]; PCA [label="PCA"];
APP [label="AI and security\\ndecision systems", fillcolor="#E8F5EE", color="#2E8B57"];
LE -> CNF [label="rewrite rules"]; CNF -> SAT [label="3-CNF / DIMACS"]; SAT -> GR [label="oracle marks\\nsolutions"];
PCA -> MDP [label="compressed\\nstate features"]; SAT -> APP [label="policy\\nconsistency"]; MDP -> APP [label="actions"];
PCA -> APP [label="anomaly\\nsignals"]; GR -> APP [label="quadratic\\nspeedup"]; }""")
    st.markdown("#### Suggested learning path")
    st.markdown("1. Logical Equivalence, then CNF, then 3-SAT (closely connected).\n"
                "2. Grover's algorithm, using 3-SAT formulas as the quantum oracle.\n"
                "3. Markov Decision Processes for AI decision-making.\n"
                "4. PCA for machine learning and data analysis.\n"
                "5. Finish with the Knowledge Check.")


# =====================================================================
# PAGE: 3-SAT
# =====================================================================
SAT_USE_CASES = [
    ("Scheduling and timetabling", "Variables represent choices such as 'exam X in room Y at time Z'. Clauses forbid "
     "double-booking and enforce capacity. A satisfying assignment is a valid timetable; UNSAT proves none exists."),
    ("Hardware and software verification", "Chip and software designers encode 'can the system reach a bad state "
     "within k steps?' as SAT (bounded model checking). UNSAT is a proof that the bug cannot happen."),
    ("Security policy consistency", "Firewall, MFA, IDS and VPN requirements become clauses. If no assignment "
     "satisfies them all, the policies contradict each other and cannot be deployed together."),
    ("Planning, puzzles and dependency resolution", "Robot plans, Sudoku, and package managers (choosing compatible "
     "library versions) are all solved by translating them to SAT."),
]


def page_sat(section):
    if section == "Concept and Formulas":
        st.markdown("## 3-SAT: Three Satisfiability Problem")
        st.markdown("**What it means.** 3-SAT asks whether we can assign TRUE or FALSE to Boolean variables so that "
                    "an entire formula becomes TRUE. The formula is an AND of clauses and each clause is an OR of "
                    "exactly three *literals* (a variable or its negation).")
        st.latex(r"F=\bigwedge_{j=1}^{m}\left(\ell_{j1}\lor\ell_{j2}\lor\ell_{j3}\right),\qquad "
                 r"\ell\in\{x_i,\ \neg x_i\}")
        st.markdown("A clause is satisfied when at least one of its literals is TRUE; the formula is satisfied when "
                    "every clause is satisfied. With n variables there are 2ⁿ possible assignments.")
        st.markdown("#### Try the example from the notes")
        st.latex(r"(A\lor B\lor C)\land(\neg A\lor B\lor\neg C)")
        c = st.columns(3)
        A = c[0].toggle("A", True, key="sat_demo_a")
        B = c[1].toggle("B", False, key="sat_demo_b")
        C = c[2].toggle("C", True, key="sat_demo_c")
        c1, c2 = (A or B or C), ((not A) or B or (not C))
        m = st.columns(3)
        m[0].metric("Clause 1", "TRUE" if c1 else "FALSE")
        m[1].metric("Clause 2", "TRUE" if c2 else "FALSE")
        m[2].metric("Overall formula", "TRUE" if c1 and c2 else "FALSE")
        st.markdown("#### Key results")
        st.markdown(
            "- **NP-complete (Cook-Levin theorem):** a fast algorithm for 3-SAT would give a fast algorithm for "
            "every problem in NP.\n"
            "- **Phase transition:** for random 3-SAT the hardest instances sit near a clause-to-variable ratio of "
            "about 4.26. Below it almost all formulas are satisfiable; above it almost none are.")
        st.latex(r"\alpha=\frac{m}{n}\approx 4.26")
        st.markdown("#### Algorithms in this module")
        st.markdown(
            "- **DPLL (complete):** unit propagation (a one-literal clause forces that literal), pure-literal "
            "elimination, then branching with backtracking. It can prove UNSAT.\n"
            "- **WalkSAT (incomplete):** start random, repeatedly pick an unsatisfied clause and flip one of its "
            "variables (randomly with probability p, otherwise greedily). Very fast on satisfiable instances but "
            "cannot prove UNSAT.\n"
            "- **Exhaustive check (n up to 20):** evaluates all 2ⁿ assignments at once and also counts every "
            "solution. This is also the classical baseline for Grover's algorithm.")
        st.latex(r"\text{Schoening random walk: } O\!\left((4/3)^n\right)\ \text{vs brute force } O(2^n)")
        st.markdown("**Real-world example.** Let A = firewall rule enabled, B = MFA enabled, C = intrusion detection "
                    "active. The clause (A OR B OR C) says at least one control must be active. A larger formula "
                    "represents many requirements that must hold simultaneously.")
    elif section == "Use Cases":
        st.markdown("## 3-SAT Use Cases")
        use_cases(SAT_USE_CASES)
    elif section == "Diagrams":
        st.markdown("## 3-SAT Diagrams")
        st.markdown("#### DPLL decision procedure")
        st.graphviz_chart("""
digraph D { node [shape=box, style="rounded,filled", fillcolor="#DCE6F5", color="#0B4FA8", fontname="Helvetica"];
edge [color="#8A94A6", fontname="Helvetica", fontsize=10];
F [label="CNF formula F"]; U [label="Unit propagation\\nforce literals in 1-literal clauses"];
E [label="Empty clause?", shape=diamond, fillcolor="#FBE4E1", color="#C0392B"];
N [label="No clauses left?", shape=diamond, fillcolor="#E8F5EE", color="#2E8B57"];
PL [label="Pure literal elimination"]; B [label="Choose variable x\\n(most frequent in shortest clauses)"];
T [label="Try x = TRUE"]; FA [label="Try x = FALSE"]; BT [label="Backtrack", fillcolor="#FBE4E1", color="#C0392B"];
S [label="SATISFIABLE\\nreturn assignment", fillcolor="#E8F5EE", color="#2E8B57"];
F -> U -> E; E -> BT [label="yes"]; E -> N [label="no"]; N -> S [label="yes"]; N -> PL [label="no"];
PL -> B -> T -> U [label="recurse"]; T -> FA [label="conflict"]; FA -> U [label="recurse"]; FA -> BT [label="conflict"]; }""")
        st.markdown("#### Factor graph of the current formula")
        cur = st.session_state.get("sat_formula")
        if cur is None:
            st.info("Run a solver in the Synthetic Data Lab or Upload section to draw the factor graph of that formula.")
        else:
            clauses, n, names = cur
            if len(clauses) > 25:
                st.info(f"The current formula has {len(clauses)} clauses; the first 25 are drawn.")
            lines = ['graph G { layout=neato; overlap=false; node [fontname="Helvetica", fontsize=10];']
            used = sorted({abs(l) for c in clauses[:25] for l in c})
            for v in used:
                lines.append(f'v{v} [label="{names[v - 1]}", shape=circle, style=filled, fillcolor="#DCE6F5", color="#0B4FA8"];')
            for j, c in enumerate(clauses[:25], 1):
                lines.append(f'c{j} [label="C{j}", shape=box, style=filled, fillcolor="#F4E3CC", color="#E07B00"];')
                for l in c:
                    lines.append(f'c{j} -- v{abs(l)} [style={"solid" if l > 0 else "dashed"}];')
            lines.append("}")
            st.graphviz_chart("\n".join(lines))
            st.caption("Boxes are clauses, circles are variables. Solid edge = positive literal, dashed = negated.")
    elif section == "Synthetic Data Lab":
        st.markdown("## 3-SAT Synthetic Data Lab")
        mode = st.radio("Experiment", ["Solve a random 3-SAT formula", "Phase transition experiment"],
                        horizontal=True, key="sat_mode")
        if mode == "Solve a random 3-SAT formula":
            c = st.columns(3)
            n = c[0].slider("Variables n", 3, 150, 20, key="sat_n")
            ratio = c[1].slider("Clause to variable ratio", 1.0, 8.0, 4.26, 0.01, key="sat_ratio")
            seed = c[2].number_input("Random seed", 0, 10 ** 6, 42, key="sat_seed")
            c = st.columns(3)
            planted = c[0].checkbox("Guarantee satisfiable (planted solution)", key="sat_planted")
            cyber = c[1].checkbox("Use cybersecurity variable names", value=n <= 20, key="sat_cyber")
            solver = c[2].selectbox("Solver", ["DPLL (complete)", "WalkSAT (local search)",
                                               "Exhaustive check (n up to 20)"], key="sat_solver")
            with st.expander("WalkSAT settings"):
                flips = st.slider("Maximum flips", 1000, 100000, 20000, 1000, key="sat_flips")
                noise = st.slider("Noise probability p", 0.0, 1.0, 0.5, 0.05, key="sat_noise")
            m = int(round(ratio * n))
            st.caption(f"The formula will have {m} clauses over {n} variables.")
            if st.button("Generate and solve", type="primary", key="sat_go"):
                if solver.startswith("Exhaustive") and n > 20:
                    st.error("Exhaustive check is limited to 20 variables. Choose DPLL or WalkSAT.")
                else:
                    rng = np.random.default_rng(int(seed))
                    clauses, _ = random_3sat(n, m, rng, planted)
                    names = (CYBER_NAMES[:n] if cyber and n <= len(CYBER_NAMES) else [f"x{i}" for i in range(1, n + 1)])
                    with st.spinner("Solving"):
                        rep = run_sat(clauses, n, names, solver,
                                      f"synthetic random 3-SAT (seed {seed}{', planted' if planted else ''})",
                                      int(seed), True, flips, noise)
                    st.session_state["res_sat"] = rep
                    st.session_state["sat_formula"] = (clauses, n, names)
            cur = st.session_state.get("sat_formula")
            if cur and cur[1] <= 12:
                with st.expander("Test your own assignment on the current formula"):
                    clauses, n, names = cur
                    cols = st.columns(min(n, 6))
                    a = {v: cols[(v - 1) % len(cols)].toggle(names[v - 1], key=f"sat_try_{v}") for v in range(1, n + 1)}
                    ok = [any(a[abs(l)] == (l > 0) for l in c) for c in clauses]
                    st.metric("Clauses satisfied", f"{sum(ok)} of {len(ok)}",
                              "Formula TRUE" if all(ok) else "Formula FALSE")
            if st.session_state.get("res_sat"):
                show_report(st.session_state["res_sat"], "sat_lab")
        else:
            c = st.columns(3)
            n = c[0].slider("Variables n", 10, 40, 20, key="pt_n")
            trials = c[1].slider("Trials per ratio", 5, 60, 20, key="pt_trials")
            seed = c[2].number_input("Random seed", 0, 10 ** 6, 7, key="pt_seed")
            if st.button("Run phase transition experiment", type="primary", key="pt_go"):
                ratios = np.round(np.arange(1.0, 8.01, 0.5), 2)
                with st.spinner("Solving many random formulas"):
                    df = sat_phase_transition(n, ratios, trials, int(seed))
                rep = Report("3-SAT Phase Transition Experiment", "3-SAT")
                rep.para(f"For n = {n} variables, {trials} random formulas were solved with DPLL at each "
                         "clause-to-variable ratio. The probability of satisfiability falls sharply around 4.26 and "
                         "the solver effort (decisions) peaks there.")
                rep.metrics({"Variables": n, "Trials per ratio": trials, "Ratios tested": len(ratios)})
                fig, ax = plt.subplots(figsize=(8, 3.6))
                ax.plot(df.ratio, df.p_satisfiable, "o-", color=BLUE, label="P(satisfiable)")
                ax.axvline(4.26, ls="--", color=GREY)
                ax.set_xlabel("Clause to variable ratio m/n")
                ax.set_ylabel("P(satisfiable)", color=BLUE)
                ax2 = ax.twinx()
                ax2.plot(df.ratio, df.mean_decisions, "s-", color=ORANGE, label="Mean DPLL decisions")
                ax2.set_ylabel("Mean decisions", color=ORANGE)
                ax2.grid(False)
                ax.set_title("Satisfiability and difficulty versus clause density")
                rep.figure("Phase transition curve", fig)
                rep.table("Phase transition results", df)
                st.session_state["res_sat"] = rep
            if st.session_state.get("res_sat"):
                show_report(st.session_state["res_sat"], "sat_lab2")
    elif section == "Upload Real Data":
        st.markdown("## 3-SAT Upload Real Data")
        st.markdown("Upload a **DIMACS CNF** file (standard SAT competition format) or a **CSV/Excel** file where "
                    "each row is a clause. Cells may be signed integers (3, -5) or names (Firewall, NOT MFA, -VPN, ~IDS).")
        tmpl = pd.DataFrame([["Firewall", "MFA", "IDS"], ["NOT Firewall", "MFA", "NOT IDS"],
                             ["VPN", "NOT MFA", "Encryption"], ["NOT VPN", "IDS", "Backup"],
                             ["NOT Encryption", "NOT Backup", "Firewall"]], columns=["l1", "l2", "l3"])
        template_button(tmpl, "sat_template.csv", key="sat_tmpl")
        upl = st.file_uploader("Upload .cnf, .txt, .csv or .xlsx", type=["cnf", "txt", "dimacs", "csv", "xlsx"],
                               key="sat_upl")
        solver = st.selectbox("Solver", ["DPLL (complete)", "WalkSAT (local search)", "Exhaustive check (n up to 20)"],
                              key="sat_upl_solver")
        if upl is not None and st.button("Solve uploaded formula", type="primary", key="sat_upl_go"):
            try:
                if upl.name.lower().endswith((".cnf", ".txt", ".dimacs")):
                    clauses, n, names = parse_dimacs(upl.getvalue().decode("utf-8", "replace"))
                else:
                    raw = upl.getvalue()
                    df = (pd.read_excel(io.BytesIO(raw), header=None) if upl.name.lower().endswith("xlsx")
                          else pd.read_csv(io.BytesIO(raw), header=None, dtype=str))
                    clauses, n, names = parse_sat_table(df)
                if solver.startswith("Exhaustive") and n > 20:
                    st.error(f"The formula has {n} variables; exhaustive check is limited to 20.")
                else:
                    with st.spinner("Solving"):
                        rep = run_sat(clauses, n, names, solver, f"uploaded file {upl.name}")
                    st.session_state["res_sat"] = rep
                    st.session_state["sat_formula"] = (clauses, n, names)
            except Exception as e:
                st.error(f"Could not read the file: {e}")
        if st.session_state.get("res_sat"):
            show_report(st.session_state["res_sat"], "sat_upl")
    else:
        show_saved_result("res_sat", "3-SAT")


# =====================================================================
# PAGE: LOGICAL EQUIVALENCE
# =====================================================================
LAWS = [
    ("De Morgan (AND)", "NOT (P AND Q)", "NOT P OR NOT Q", True),
    ("De Morgan (OR)", "NOT (P OR Q)", "NOT P AND NOT Q", True),
    ("Implication elimination", "P -> Q", "NOT P OR Q", True),
    ("Contrapositive", "P -> Q", "NOT Q -> NOT P", True),
    ("Double negation", "NOT NOT P", "P", True),
    ("Distribution of OR over AND", "P OR (Q AND R)", "(P OR Q) AND (P OR R)", True),
    ("Distribution of AND over OR", "P AND (Q OR R)", "(P AND Q) OR (P AND R)", True),
    ("Biconditional", "P <-> Q", "(P -> Q) AND (Q -> P)", True),
    ("Absorption", "P OR (P AND Q)", "P", True),
    ("Exportation", "(P AND Q) -> R", "P -> (Q -> R)", True),
    ("XOR definition", "P XOR Q", "(P OR Q) AND NOT (P AND Q)", True),
    ("Commutativity", "P AND Q", "Q AND P", True),
    ("Mistake: converse", "P -> Q", "Q -> P", False),
    ("Mistake: inverse", "P -> Q", "NOT P -> NOT Q", False),
    ("Mistake: wrong De Morgan", "NOT (P AND Q)", "NOT P AND NOT Q", False),
    ("Mistake: dropped distribution", "P OR (Q AND R)", "(P OR Q) AND R", False),
    ("Mistake: implication is not associative", "(P -> Q) -> R", "P -> (Q -> R)", False),
]
LE_USE_CASES = [
    ("Digital circuit simplification", "Engineers replace a gate-heavy Boolean expression with an equivalent smaller "
     "one (De Morgan, absorption, distribution). Same behaviour, less chip area, power and delay."),
    ("Compiler and database query optimisation", "Compilers and SQL engines rewrite conditions such as NOT (a AND b) "
     "into equivalent forms that short-circuit sooner or use an index."),
    ("Security policy refactoring audits", "Before deploying a rewritten access-control rule, prove it grants exactly "
     "the same access as the old rule for every combination of attributes. A counterexample is a security bug."),
    ("Theorem proving and AI rule engines", "Knowledge bases are converted to standard forms (NNF, CNF) using "
     "equivalences, so inference engines can reason without changing meaning."),
]


def le_pair_report(e1, e2, source):
    a, b = parse_expr(e1), parse_expr(e2)
    vs = sorted(set(variables(a)) | set(variables(b)), key=natkey)
    if len(vs) > MAX_TT_VARS:
        raise ValueError(f"{len(vs)} variables; the truth-table limit is {MAX_TT_VARS}.")
    env, size = assignment_env(vs)
    ra, rb = ev(a, env, size), ev(b, env, size)
    diff = np.nonzero(ra != rb)[0]
    eq = len(diff) == 0
    rep = Report("Logical Equivalence Check", "Logical Equivalence")
    rep.heading("Result")
    rep.para(f"Data source: {source}.")
    rep.latex(to_str(a, "latex") + (r"\ \equiv\ " if eq else r"\ \not\equiv\ ") + to_str(b, "latex"),
              f"{to_str(a)}  {'==' if eq else 'is NOT equivalent to'}  {to_str(b)}")
    ent1 = not np.any(ra & ~rb)
    ent2 = not np.any(rb & ~ra)
    rel = "Equivalent" if eq else ("Expression 1 implies Expression 2" if ent1 else
                                   "Expression 2 implies Expression 1" if ent2 else "Neither implies the other")
    rep.metrics({"Variables": len(vs), "Truth table rows": size, "Equivalent": "YES" if eq else "NO",
                 "Rows that differ": len(diff), "Expression 1 type": classify(ra),
                 "Expression 2 type": classify(rb), "Relationship": rel})
    if eq:
        rep.para("The two expressions agree on every one of the " + f"{size} assignments, so (Expression 1 <-> "
                 "Expression 2) is a tautology and either one can replace the other anywhere.")
    else:
        rep.para(f"Counterexample: {fmt_row(vs, env, diff[0])} gives Expression 1 = "
                 f"{'TRUE' if ra[diff[0]] else 'FALSE'} but Expression 2 = {'TRUE' if rb[diff[0]] else 'FALSE'}.")
    lim = min(size, 1024)
    tt = pd.DataFrame({v: np.where(env[v][:lim], "T", "F") for v in vs})
    tt["expression_1"] = np.where(ra[:lim], "T", "F")
    tt["expression_2"] = np.where(rb[:lim], "T", "F")
    tt["match"] = np.where(ra[:lim] == rb[:lim], "yes", "NO")
    rep.table("Truth table" + (" (first 1024 rows)" if size > 1024 else ""), tt)
    if len(diff):
        rep.table("Counterexamples", pd.DataFrame({"assignment": [fmt_row(vs, env, r) for r in diff[:500]],
                                                   "expression_1": ra[diff[:500]], "expression_2": rb[diff[:500]]}))
    if size <= 64:
        M = np.column_stack([env[v] for v in vs] + [ra, rb]).astype(int)
        fig, ax = plt.subplots(figsize=(1.0 + 0.8 * M.shape[1], 0.6 + 0.26 * size))
        cmap = matplotlib.colors.ListedColormap(["#F4D6D2", "#D6E6F7"])
        ax.imshow(M, cmap=cmap, aspect="auto")
        for (i, j), v in np.ndenumerate(M):
            mismatch = j >= len(vs) and ra[i] != rb[i]
            ax.text(j, i, "T" if v else "F", ha="center", va="center", fontsize=8,
                    weight="bold" if mismatch else "normal", color=RED if mismatch else "#222")
        ax.set_xticks(range(M.shape[1]))
        ax.set_xticklabels(vs + ["Expr 1", "Expr 2"])
        ax.set_yticks([])
        ax.axvline(len(vs) - 0.5, color=BLUE, lw=2)
        ax.grid(False)
        ax.set_title("Truth table (red bold = rows where the expressions differ)")
        rep.figure("Truth table heatmap", fig)
    rep.figure("Expression 1 tree", ast_figure(a, "Expression 1: " + to_str(a)))
    rep.figure("Expression 2 tree", ast_figure(b, "Expression 2: " + to_str(b)))
    st.session_state["le_last_pair"] = (a, b)
    return rep


def le_batch_report(pairs, source):
    rows = []
    for i, p in enumerate(pairs, 1):
        e1, e2, lab, exp = p
        try:
            a, b = parse_expr(e1), parse_expr(e2)
            vs = sorted(set(variables(a)) | set(variables(b)), key=natkey)
            if len(vs) > MAX_TT_VARS:
                raise ValueError(f"too many variables ({len(vs)})")
            env, size = assignment_env(vs)
            ra, rb = ev(a, env, size), ev(b, env, size)
            diff = np.nonzero(ra != rb)[0]
            rows.append({"pair": i, "label": lab, "expression_1": to_str(a), "expression_2": to_str(b),
                         "variables": len(vs), "equivalent": len(diff) == 0,
                         "expected": exp, "matches_expected": (exp == (len(diff) == 0)) if exp is not None else None,
                         "differing_rows": len(diff), "of_rows": size,
                         "counterexample": fmt_row(vs, env, diff[0]) if len(diff) else "", "error": ""})
        except Exception as e:
            rows.append({"pair": i, "label": lab, "expression_1": e1, "expression_2": e2, "variables": None,
                         "equivalent": None, "expected": exp, "matches_expected": None, "differing_rows": None,
                         "of_rows": None, "counterexample": "", "error": str(e)})
    df = pd.DataFrame(rows)
    rep = Report("Logical Equivalence Batch Analysis", "Logical Equivalence")
    rep.heading("Summary")
    rep.para(f"Data source: {source}.")
    ok = df["error"] == ""
    met = {"Pairs checked": len(df), "Equivalent": int((df["equivalent"] == True).sum()),
           "Not equivalent": int((df["equivalent"] == False).sum()), "Parse errors": int((~ok).sum())}
    if df["expected"].notna().any():
        met["Agreement with expected label"] = f"{(df['matches_expected'] == True).mean() * 100:.0f}%"
    rep.metrics(met)
    rep.table("Equivalence results", df)
    if ok.any():
        g = df[ok].groupby("label")["equivalent"].agg(["sum", "count"]).reset_index()
        g["not_equivalent"] = g["count"] - g["sum"]
        fig, ax = plt.subplots(figsize=(9, max(2.5, 0.35 * len(g) + 1)))
        ax.barh(g["label"], g["sum"], color=BLUE, label="Equivalent")
        ax.barh(g["label"], g["not_equivalent"], left=g["sum"], color=ORANGE, label="Not equivalent")
        ax.set_xlabel("Pairs")
        ax.set_title("Equivalence results by law or label")
        ax.legend(fontsize=8)
        ax.invert_yaxis()
        rep.figure("Results by label", fig)
        rep.table("Summary by label", g.rename(columns={"sum": "equivalent", "count": "pairs"}))
    return rep


def page_le(section):
    if section == "Concept and Formulas":
        st.markdown("## Logical Equivalence")
        st.markdown("**What it means.** Two expressions P and Q are logically equivalent when they produce the same "
                    "truth value for **every** possible assignment of their variables. Equivalently, P ↔ Q is a "
                    "tautology.")
        st.latex(r"P\equiv Q \iff \forall\,\text{assignments}:\ v(P)=v(Q) \iff \models (P\leftrightarrow Q)")
        st.markdown("Checking equivalence by truth table costs 2ⁿ rows for n variables. Non-equivalence is the "
                    "complement of a SAT problem: P and Q differ exactly when (P XOR Q) is satisfiable.")
        st.latex(r"P\not\equiv Q \iff (P\oplus Q)\ \text{is satisfiable}")
        st.markdown("#### Example 1: De Morgan's law (computed live)")
        st.latex(r"\neg(A\land B)\equiv(\neg A\lor\neg B)")
        env, size = assignment_env(["A", "B"])
        a, b = parse_expr("NOT (A AND B)"), parse_expr("NOT A OR NOT B")
        ui_dataframe(pd.DataFrame({"A": np.where(env["A"], "T", "F"), "B": np.where(env["B"], "T", "F"),
                                   "NOT (A AND B)": np.where(ev(a, env, size), "T", "F"),
                                   "NOT A OR NOT B": np.where(ev(b, env, size), "T", "F")}))
        st.markdown("The last two columns are identical, so the expressions are equivalent.")
        st.markdown("#### Example 2: if-then statements")
        st.latex(r"A\rightarrow B\equiv\neg A\lor B")
        st.markdown("With A = 'user has valid credentials' and B = 'system grants access', the rule 'if valid "
                    "credentials then grant access' equals 'credentials are not valid OR access is granted'. This "
                    "rewriting is how rules are prepared for SAT solvers.")
        st.markdown("#### Standard laws")
        ui_dataframe(pd.DataFrame([(n, l, r, "Yes" if e else "No (common mistake)") for n, l, r, e in LAWS],
                                  columns=["Law", "Left side", "Right side", "Equivalent"]), height=420)
        st.markdown("**Expression syntax accepted in this app:** NOT ~ ! · AND & · OR | · XOR ^ · implies -> => · "
                    "biconditional <-> <=> · constants TRUE FALSE 1 0 · parentheses. Precedence from strongest: "
                    "NOT, AND, XOR, OR, ->, <->.")
    elif section == "Use Cases":
        st.markdown("## Logical Equivalence Use Cases")
        use_cases(LE_USE_CASES)
    elif section == "Diagrams":
        st.markdown("## Logical Equivalence Diagrams")
        st.markdown("#### Equivalence checking pipeline")
        st.graphviz_chart("""
digraph G { rankdir=LR; node [shape=box, style="rounded,filled", fillcolor="#DCE6F5", color="#0B4FA8", fontname="Helvetica"];
P [label="Expression P"]; Q [label="Expression Q"]; V [label="Collect variables\\n(union)"];
T [label="Enumerate all 2^n\\nassignments"]; C [label="Compare v(P) and v(Q)\\nrow by row"];
D [label="Same everywhere?", shape=diamond]; EQ [label="EQUIVALENT", fillcolor="#E8F5EE", color="#2E8B57"];
NE [label="NOT EQUIVALENT\\n+ counterexample", fillcolor="#FBE4E1", color="#C0392B"];
P -> V; Q -> V; V -> T -> C -> D; D -> EQ [label="yes"]; D -> NE [label="no"]; }""")
        a, b = st.session_state.get("le_last_pair", (parse_expr("NOT (A AND B)"), parse_expr("NOT A OR NOT B")))
        st.markdown("#### Expression trees of the last checked pair")
        c = st.columns(2)
        with c[0]:
            st.graphviz_chart(ast_dot(a, "Expression 1"))
        with c[1]:
            st.graphviz_chart(ast_dot(b, "Expression 2"))
    elif section == "Synthetic Data Lab":
        st.markdown("## Logical Equivalence Synthetic Data Lab")
        mode = st.radio("Mode", ["Check one pair", "Generate synthetic pairs from known laws"], horizontal=True,
                        key="le_mode")
        if mode == "Check one pair":
            e1 = st.text_input("Expression 1", "NOT (A AND B)", key="le_e1")
            e2 = st.text_input("Expression 2", "NOT A OR NOT B", key="le_e2")
            if st.button("Check equivalence", type="primary", key="le_go"):
                try:
                    st.session_state["res_le"] = le_pair_report(e1, e2, "manual entry")
                except Exception as e:
                    st.error(f"Could not parse: {e}")
        else:
            c = st.columns(4)
            count = c[0].slider("Number of pairs", 5, 300, 40, key="le_count")
            nvars = c[1].slider("Variable pool size", 2, 8, 4, key="le_nvars")
            depth = c[2].slider("Sub-expression depth", 0, 3, 1, key="le_depth")
            seed = c[3].number_input("Random seed", 0, 10 ** 6, 11, key="le_seed")
            rand_frac = st.slider("Share of fully random pairs (no law)", 0.0, 0.5, 0.15, key="le_rand")
            if st.button("Generate and check", type="primary", key="le_gen"):
                rng = np.random.default_rng(int(seed))
                pool = [chr(65 + i) for i in range(nvars)]
                pairs = []
                for _ in range(count):
                    if rng.random() < rand_frac:
                        x, y = random_expr(pool, depth + 1, rng), random_expr(pool, depth + 1, rng)
                        pairs.append((to_str(x), to_str(y), "Random pair", None))
                    else:
                        name, l, r, exp = LAWS[rng.integers(len(LAWS))]
                        mp = {v: random_expr(pool, depth, rng) for v in "PQR"}
                        pairs.append((to_str(subst(parse_expr(l), mp)), to_str(subst(parse_expr(r), mp)), name, exp))
                st.session_state["res_le"] = le_batch_report(pairs, f"synthetic pairs (seed {seed})")
        if st.session_state.get("res_le"):
            show_report(st.session_state["res_le"], "le_lab")
    elif section == "Upload Real Data":
        st.markdown("## Logical Equivalence Upload Real Data")
        st.markdown("Upload a **CSV/Excel** with columns `expr1`, `expr2` and optional `label` and `expected` "
                    "(TRUE/FALSE), or a **TXT** file with one pair per line separated by `;` or `==`.")
        tmpl = pd.DataFrame({"expr1": ["NOT (Admin AND Remote)", "Valid -> Grant", "Valid -> Grant", "MFA OR (VPN AND Device)"],
                             "expr2": ["NOT Admin OR NOT Remote", "NOT Valid OR Grant", "Grant -> Valid",
                                       "(MFA OR VPN) AND (MFA OR Device)"],
                             "label": ["De Morgan", "Implication", "Converse", "Distribution"],
                             "expected": [True, True, False, True]})
        template_button(tmpl, "equivalence_template.csv", key="le_tmpl")
        upl = st.file_uploader("Upload .csv, .xlsx or .txt", type=["csv", "xlsx", "txt"], key="le_upl")
        if upl is not None and st.button("Check uploaded pairs", type="primary", key="le_upl_go"):
            try:
                pairs = []
                if upl.name.lower().endswith(".txt"):
                    for line in upl.getvalue().decode("utf-8", "replace").splitlines():
                        if line.strip():
                            parts = re.split(r"\s*(?:;|==|\t)\s*", line.strip(), maxsplit=1)
                            if len(parts) == 2:
                                pairs.append((parts[0], parts[1], "Uploaded", None))
                else:
                    df = read_upload_table(upl)
                    cols = {c.lower().strip(): c for c in df.columns}
                    c1 = cols.get("expr1", df.columns[0])
                    c2 = cols.get("expr2", df.columns[1])
                    for _, r in df.iterrows():
                        exp = r[cols["expected"]] if "expected" in cols else None
                        if isinstance(exp, str):
                            exp = exp.strip().upper() in ("TRUE", "YES", "1")
                        elif exp is not None and not pd.isna(exp):
                            exp = bool(exp)
                        else:
                            exp = None
                        pairs.append((str(r[c1]), str(r[c2]), str(r[cols["label"]]) if "label" in cols else "Uploaded", exp))
                if not pairs:
                    st.error("No expression pairs found.")
                else:
                    st.session_state["res_le"] = le_batch_report(pairs, f"uploaded file {upl.name}")
            except Exception as e:
                st.error(f"Could not read the file: {e}")
        if st.session_state.get("res_le"):
            show_report(st.session_state["res_le"], "le_upl")
    else:
        show_saved_result("res_le", "logical equivalence")


# =====================================================================
# PAGE: CNF
# =====================================================================
CNF_USE_CASES = [
    ("Input format for SAT solvers", "Every modern SAT solver (MiniSat, CaDiCaL, Kissat) reads CNF in DIMACS format. "
     "CNF conversion is the bridge from real-world rules to industrial solvers."),
    ("Compliance and access-control checklists", "CNF reads naturally as 'requirement 1 AND requirement 2 AND ...', "
     "each requirement listing acceptable options: (MFA OR Certificate) AND (TrustedDevice OR VPN)."),
    ("Automated theorem proving", "The resolution rule, used by many provers and Prolog-style systems, works on "
     "clauses, so knowledge must first be put into CNF."),
    ("Model checking and test generation", "Program paths and circuit behaviour are encoded as CNF (often via Tseitin) "
     "to find inputs that reach a branch or violate an assertion."),
]


def cnf_single_report(text, source):
    ast, steps, final, blow = cnf_pipeline(text)
    vs = variables(ast)
    rep = Report("CNF Conversion", "CNF")
    rep.heading("Conversion steps")
    rep.para(f"Data source: {source}.")
    rep.table("Step by step transformation", pd.DataFrame(steps, columns=["step", "expression"]))
    tse, gates = tseitin(ast)
    met = {"Variables": len(vs), "Expression size (nodes)": ast_size(ast)}
    if not blow:
        met.update({"CNF clauses": len(final), "CNF literals": sum(len(c) for c in final),
                    "Largest clause": max((len(c) for c in final), default=0)})
        if len(vs) <= MAX_TT_VARS:
            env, size = assignment_env(vs)
            same = np.array_equal(ev(ast, env, size), eval_clauses(final, env, size))
            met["Verified equivalent by truth table"] = "YES" if same else "NO"
    met["Tseitin clauses"] = len(tse)
    met["Tseitin auxiliary variables"] = len(gates)
    rep.metrics(met)
    if not blow:
        rep.latex(r"\text{CNF: }\ " + (r" \land ".join(
            "(" + r" \lor ".join((_latex_var(v) if p else r"\neg " + _latex_var(v)) for v, p in
                                 sorted(c, key=lambda l: natkey(l[0]))) + ")" for c in final[:8]) or r"\top")
                  + (r"\land\cdots" if len(final) > 8 else ""), "CNF: " + cnf_str(final))
        rep.table("CNF clauses", pd.DataFrame({"clause": range(1, len(final) + 1),
                                               "clause_text": [clause_str(c) for c in final],
                                               "width": [len(c) for c in final]}))
        base = final
    else:
        rep.para("Direct conversion was stopped because distribution would create too many clauses. The Tseitin "
                 "encoding below is equisatisfiable and grows only linearly.")
        base = tse
    rep.heading("Tseitin encoding (equisatisfiable, linear size)")
    rep.table("Tseitin gate definitions", pd.DataFrame(gates))
    rep.table("Tseitin clauses", pd.DataFrame({"clause": range(1, len(tse) + 1),
                                               "clause_text": [clause_str(c) for c in tse]}))
    three = to_3cnf(base, vs + [g["aux_variable"] for g in gates])
    rep.heading("Strict 3-CNF (ready for the 3-SAT module)")
    rep.para(f"{len(base)} clauses were rewritten so every clause has exactly three literals, giving {len(three)} "
             "clauses. Short clauses are padded with fresh variables; long clauses are split with a chain of fresh "
             "variables. The result is equisatisfiable with the original.")
    three_fs = [frozenset(c) if len(set(c)) == 3 else c for c in three]
    lines, names = to_dimacs([tuple(c) for c in three])
    rep.table("3-CNF clauses", pd.DataFrame({"clause": range(1, len(three) + 1),
                                             "clause_text": ["(" + " OR ".join(lit_str(l) for l in c) + ")"
                                                             for c in three]}))
    rep.table("DIMACS export of 3-CNF", pd.DataFrame({"dimacs_line": lines}))
    rep.table("DIMACS variable map", pd.DataFrame({"id": range(1, len(names) + 1), "variable": names}))
    fig, ax = plt.subplots(figsize=(7, 3))
    labs = (["Direct CNF"] if not blow else []) + ["Tseitin CNF", "3-CNF"]
    vals = ([len(final)] if not blow else []) + [len(tse), len(three)]
    ax.bar(labs, vals, color=[BLUE, ORANGE, GREEN][-len(labs):])
    for i, v in enumerate(vals):
        ax.text(i, v, str(v), ha="center", va="bottom")
    ax.set_ylabel("Clauses")
    ax.set_title("Size of each encoding")
    rep.figure("Encoding sizes", fig)
    rep.figure("Expression tree", ast_figure(ast, "Original: " + to_str(ast)))
    st.session_state["cnf_last_ast"] = ast
    st.session_state["cnf_last_3cnf"] = ([tuple(c) for c in three], names)
    return rep


def cnf_batch_report(exprs, source):
    rows = []
    for i, e in enumerate(exprs, 1):
        try:
            ast, steps, final, blow = cnf_pipeline(e, limit=5000)
            vs = variables(ast)
            tse, gates = tseitin(ast)
            eqv = None
            if not blow and len(vs) <= 14:
                env, size = assignment_env(vs)
                eqv = bool(np.array_equal(ev(ast, env, size), eval_clauses(final, env, size)))
            rows.append({"id": i, "expression": to_str(ast), "variables": len(vs), "size_nodes": ast_size(ast),
                         "cnf": cnf_str(final) if not blow else "(too large)",
                         "cnf_clauses": len(final) if not blow else None,
                         "max_width": max((len(c) for c in final), default=0) if not blow else None,
                         "verified_equivalent": eqv, "tseitin_clauses": len(tse),
                         "three_cnf_clauses": len(to_3cnf(final if not blow else tse, vs)), "error": ""})
        except Exception as ex:
            rows.append({"id": i, "expression": e, "error": str(ex)})
    df = pd.DataFrame(rows)
    rep = Report("CNF Batch Conversion", "CNF")
    rep.heading("Summary")
    rep.para(f"Data source: {source}.")
    good = df[df["error"] == ""]
    rep.metrics({"Expressions": len(df), "Converted": len(good), "Errors": len(df) - len(good),
                 "Mean CNF clauses": f"{good['cnf_clauses'].mean():.1f}" if len(good) else "-",
                 "Mean Tseitin clauses": f"{good['tseitin_clauses'].mean():.1f}" if len(good) else "-",
                 "All verified": "YES" if len(good) and good["verified_equivalent"].dropna().all() else "see table"})
    rep.table("Conversion results", df)
    if len(good):
        fig, ax = plt.subplots(figsize=(8, 3.6))
        ax.scatter(good["size_nodes"], good["cnf_clauses"], color=BLUE, label="Direct CNF", alpha=0.75)
        ax.scatter(good["size_nodes"], good["tseitin_clauses"], color=ORANGE, label="Tseitin", alpha=0.75, marker="s")
        ax.set_yscale("log")
        ax.set_xlabel("Expression size (nodes)")
        ax.set_ylabel("Clauses (log scale)")
        ax.set_title("Direct CNF can grow exponentially; Tseitin grows linearly")
        ax.legend(fontsize=8)
        rep.figure("Clause growth", fig)
    return rep


def page_cnf(section):
    if section == "Concept and Formulas":
        st.markdown("## CNF: Conjunctive Normal Form")
        st.markdown("**What it means.** A formula is in CNF when it is an AND of clauses and every clause is an OR of "
                    "literals. AND joins clauses; OR joins literals inside a clause; a literal is a variable or its "
                    "negation.")
        st.latex(r"F=\bigwedge_{i=1}^{m} C_i,\qquad C_i=\bigvee_{j}\ell_{ij}")
        st.markdown("#### Worked example")
        st.latex(r"A\rightarrow(B\land C)")
        st.latex(r"\equiv\ \neg A\lor(B\land C)\qquad\text{(implication elimination)}")
        st.latex(r"\equiv\ \boxed{(\neg A\lor B)\land(\neg A\lor C)}\qquad\text{(distribute OR over AND)}")
        st.markdown("#### Conversion rules used by this app")
        st.latex(r"P\leftrightarrow Q\equiv(P\rightarrow Q)\land(Q\rightarrow P),\qquad P\oplus Q\equiv(P\lor Q)\land(\neg P\lor\neg Q)")
        st.latex(r"P\rightarrow Q\equiv\neg P\lor Q,\qquad \neg(P\land Q)\equiv\neg P\lor\neg Q,\qquad \neg\neg P\equiv P")
        st.latex(r"P\lor(Q\land R)\equiv(P\lor Q)\land(P\lor R)")
        st.markdown("Distribution can make the CNF exponentially larger. The **Tseitin transformation** avoids this by "
                    "naming each sub-formula with a fresh variable t and adding a few clauses that force t to equal the "
                    "sub-formula. The result is *equisatisfiable* (satisfiable exactly when the original is) and only "
                    "linear in size:")
        st.latex(r"t\leftrightarrow(a\land b)\ \Rightarrow\ (\neg t\lor a)\land(\neg t\lor b)\land(t\lor\neg a\lor\neg b)")
        st.markdown("#### From CNF to 3-CNF")
        st.latex(r"(\ell_1\lor\ell_2\lor\ell_3\lor\ell_4)\ \Rightarrow\ (\ell_1\lor\ell_2\lor y)\land(\neg y\lor\ell_3\lor\ell_4)")
        st.markdown("Every CNF can be rewritten so each clause has exactly three literals, which shows that 3-SAT is "
                    "as hard as general SAT.")
        st.markdown("#### Recognising CNF")
        ui_dataframe(pd.DataFrame([("A AND (B OR C)", "CNF", "AND joins two OR clauses"),
                                   ("(NOT A OR B) AND (NOT A OR C)", "CNF", "AND of OR clauses"),
                                   ("A OR (B AND C)", "Not CNF", "AND appears inside an OR"),
                                   ("NOT (A OR B)", "Not CNF", "NOT applies to a compound expression")],
                                  columns=["Expression", "CNF?", "Reason"]))
        st.markdown("**Cybersecurity example:** (MFA OR Certificate) AND (TrustedDevice OR VPN) requires either MFA "
                    "or a certificate, and either a trusted device or a VPN.")
    elif section == "Use Cases":
        st.markdown("## CNF Use Cases")
        use_cases(CNF_USE_CASES)
    elif section == "Diagrams":
        st.markdown("## CNF Diagrams")
        st.markdown("#### Conversion pipeline")
        st.graphviz_chart("""
digraph G { rankdir=LR; node [shape=box, style="rounded,filled", fillcolor="#DCE6F5", color="#0B4FA8", fontname="Helvetica"];
E [label="Any Boolean\\nexpression"]; I [label="Remove <-> and XOR"]; M [label="Remove ->"];
N [label="Push NOT inward\\n(NNF)"]; D [label="Distribute\\nOR over AND"]; S [label="Simplify\\nclauses"];
C [label="CNF", fillcolor="#E8F5EE", color="#2E8B57"]; T [label="Tseitin\\n(linear size)", fillcolor="#F4E3CC", color="#E07B00"];
K [label="3-CNF / DIMACS", fillcolor="#E8F5EE", color="#2E8B57"]; SAT [label="3-SAT solver"];
E -> I -> M -> N -> D -> S -> C -> K -> SAT; E -> T [style=dashed, label="if large"]; T -> K; }""")
        st.markdown("#### Expression tree")
        txt = st.text_input("Expression to draw", "A -> (B AND C)", key="cnf_diag_expr")
        try:
            ast = parse_expr(txt)
            st.graphviz_chart(ast_dot(ast))
            _, steps, final, blow = cnf_pipeline(txt)
            if not blow:
                st.graphviz_chart(ast_dot(parse_expr(cnf_str(final)) if final else ("const", True), "CNF form"))
        except Exception as e:
            st.error(str(e))
    elif section == "Synthetic Data Lab":
        st.markdown("## CNF Synthetic Data Lab")
        mode = st.radio("Mode", ["Convert one expression", "Generate random expressions"], horizontal=True,
                        key="cnf_mode")
        if mode == "Convert one expression":
            ex = st.selectbox("Example", ["A -> (B AND C)", "(MFA OR Certificate) AND (TrustedDevice OR VPN)",
                                          "(A <-> B) XOR C", "NOT (A OR (B -> C)) OR (D AND E)",
                                          "(Admin -> MFA) AND (Remote -> (VPN AND Device))"], key="cnf_ex")
            txt = st.text_input("Expression (edit freely)", ex, key=f"cnf_txt_{ex}")
            if st.button("Convert to CNF", type="primary", key="cnf_go"):
                try:
                    st.session_state["res_cnf"] = cnf_single_report(txt, "manual entry")
                except Exception as e:
                    st.error(f"Could not parse: {e}")
        else:
            c = st.columns(4)
            count = c[0].slider("Expressions", 5, 200, 30, key="cnf_count")
            nv = c[1].slider("Variables", 2, 8, 4, key="cnf_nv")
            depth = c[2].slider("Depth", 1, 5, 3, key="cnf_depth")
            seed = c[3].number_input("Random seed", 0, 10 ** 6, 3, key="cnf_seed")
            if st.button("Generate and convert", type="primary", key="cnf_gen"):
                rng = np.random.default_rng(int(seed))
                pool = [chr(65 + i) for i in range(nv)]
                exprs = [to_str(random_expr(pool, depth, rng)) for _ in range(count)]
                with st.spinner("Converting"):
                    st.session_state["res_cnf"] = cnf_batch_report(exprs, f"synthetic random expressions (seed {seed})")
        if st.session_state.get("res_cnf"):
            if st.session_state.get("cnf_last_3cnf") and st.button("Send last 3-CNF to the 3-SAT module", key="cnf_send"):
                cls, names = st.session_state["cnf_last_3cnf"]
                idx = {v: i + 1 for i, v in enumerate(names)}
                ints = [tuple(idx[v] if p else -idx[v] for v, p in c) for c in cls]
                st.session_state["res_sat"] = run_sat(ints, len(names), names, "DPLL (complete)", "3-CNF from CNF module")
                st.session_state["sat_formula"] = (ints, len(names), names)
                st.success("Solved in the 3-SAT module. Open 3-SAT Satisfiability, Results and Export to view it.")
            show_report(st.session_state["res_cnf"], "cnf_lab")
    elif section == "Upload Real Data":
        st.markdown("## CNF Upload Real Data")
        st.markdown("Upload a **CSV/Excel** with an `expression` column (or the first column is used), or a **TXT** "
                    "file with one expression per line. One expression gives a full step-by-step report; several give "
                    "a batch report.")
        tmpl = pd.DataFrame({"expression": ["(MFA OR Certificate) AND (TrustedDevice OR VPN)",
                                            "Admin -> (MFA AND Logging)", "Remote <-> VPN",
                                            "NOT (Guest AND (Finance OR HR))"]})
        template_button(tmpl, "cnf_template.csv", key="cnf_tmpl")
        upl = st.file_uploader("Upload .csv, .xlsx or .txt", type=["csv", "xlsx", "txt"], key="cnf_upl")
        if upl is not None and st.button("Convert uploaded expressions", type="primary", key="cnf_upl_go"):
            try:
                if upl.name.lower().endswith(".txt"):
                    exprs = [l.strip() for l in upl.getvalue().decode("utf-8", "replace").splitlines() if l.strip()]
                else:
                    df = read_upload_table(upl)
                    cols = {c.lower().strip(): c for c in df.columns}
                    exprs = df[cols.get("expression", df.columns[0])].dropna().astype(str).tolist()
                if not exprs:
                    st.error("No expressions found.")
                elif len(exprs) == 1:
                    st.session_state["res_cnf"] = cnf_single_report(exprs[0], f"uploaded file {upl.name}")
                else:
                    st.session_state["res_cnf"] = cnf_batch_report(exprs, f"uploaded file {upl.name}")
            except Exception as e:
                st.error(f"Could not read the file: {e}")
        if st.session_state.get("res_cnf"):
            show_report(st.session_state["res_cnf"], "cnf_upl")
    else:
        show_saved_result("res_cnf", "CNF")


# =====================================================================
# PAGE: MDP
# =====================================================================
MDP_USE_CASES = [
    ("Robotics and autonomous navigation", "A delivery robot or drone weighs a fast but risky corridor against a "
     "slower safe route when movement is uncertain (wheels slip, doors close). The policy says what to do in every "
     "location."),
    ("Cybersecurity incident response", "States such as Normal, Suspicious, Compromised and Contained; actions such "
     "as Monitor, Scan, Isolate and Restore. The optimal policy balances the cost of disruption against breach risk."),
    ("Inventory and supply-chain control", "Each week decide how much to reorder given uncertain demand, holding cost "
     "and stock-out penalties. The state is the stock level; the reward is profit."),
    ("Healthcare, finance and networking", "Treatment sequencing, dynamic pricing, portfolio rebalancing and network "
     "traffic routing are sequential decisions with uncertain outcomes, the exact setting MDPs and reinforcement "
     "learning were built for."),
]


def page_mdp(section):
    if section == "Concept and Formulas":
        st.markdown("## Markov Decision Process (MDP)")
        st.markdown("**What it means.** An MDP is a mathematical framework for choosing actions when outcomes are "
                    "uncertain and today's decision changes tomorrow's situation. It is the foundation of "
                    "reinforcement learning.")
        st.latex(r"\boxed{\mathcal{M}=(S,\ A,\ P,\ R,\ \gamma)}")
        ui_dataframe(pd.DataFrame([
            ("S", "States", "Situations the system can be in", "Hallway, near door, at destination"),
            ("A", "Actions", "Choices available", "Move left, move right, wait"),
            ("P(s'|s,a)", "Transition probabilities", "Likelihood of each next state", "90% chance the move succeeds"),
            ("R(s,a,s')", "Rewards", "Benefit or penalty of an outcome", "+100 delivery, -50 collision"),
            ("gamma", "Discount factor", "How much future rewards matter", "0.9 gives future rewards substantial weight"),
        ], columns=["Symbol", "Component", "Meaning", "Delivery robot example"]))
        st.markdown("**Markov property:** the next state depends only on the current state and action, not the history.")
        st.latex(r"P(s_{t+1}\mid s_t,a_t,s_{t-1},\dots)=P(s_{t+1}\mid s_t,a_t)")
        st.markdown("**Goal:** find a policy π(s) that maximises the expected discounted return.")
        st.latex(r"G_t=\sum_{k=0}^{\infty}\gamma^k R_{t+k+1},\qquad V^\pi(s)=\mathbb{E}_\pi[G_t\mid s_t=s]")
        st.markdown("#### Bellman optimality equations")
        st.latex(r"V^*(s)=\max_a\sum_{s'}P(s'\mid s,a)\big[R(s,a,s')+\gamma V^*(s')\big]")
        st.latex(r"Q^*(s,a)=\sum_{s'}P(s'\mid s,a)\big[R(s,a,s')+\gamma\max_{a'}Q^*(s',a')\big],\qquad \pi^*(s)=\arg\max_a Q^*(s,a)")
        st.markdown("#### Algorithms in this module")
        st.markdown("- **Value iteration:** repeat the Bellman update for every state until values stop changing.")
        st.latex(r"V_{k+1}(s)\leftarrow\max_a\sum_{s'}P(s'\mid s,a)[R+\gamma V_k(s')]\quad\text{until}\ \max_s|V_{k+1}-V_k|<\theta")
        st.markdown("- **Policy iteration:** evaluate the current policy exactly (solve a linear system), then improve it "
                    "greedily; stops when no state changes its action.")
        st.latex(r"V^\pi=(I-\gamma P^\pi)^{-1}R^\pi,\qquad \pi'(s)=\arg\max_a Q^\pi(s,a)")
        st.markdown("- **Q-learning (model-free reinforcement learning):** learns from sampled experience without "
                    "knowing P, using epsilon-greedy exploration.")
        st.latex(r"Q(s,a)\leftarrow Q(s,a)+\alpha\big[r+\gamma\max_{a'}Q(s',a')-Q(s,a)\big]")
        st.markdown("#### Delivery robot rewards (grid world in the lab)")
        ui_dataframe(pd.DataFrame([("Package delivered (reach GOAL)", "+100"), ("Each step taken", "-1"),
                                   ("Collision (enter HAZARD)", "-50"), ("Wait one step", "-2")],
                                  columns=["Situation", "Reward"]))
        st.markdown("**Rule-based program vs MDP.** A rule says 'if the corridor is blocked, turn left'. An MDP "
                    "compares turning left, turning right and waiting using probabilities, costs and future outcomes. "
                    "The MDP is the model; reinforcement learning is a family of methods that learn a good policy for it.")
    elif section == "Use Cases":
        st.markdown("## MDP Use Cases")
        use_cases(MDP_USE_CASES)
    elif section == "Diagrams":
        st.markdown("## MDP Diagrams")
        st.markdown("#### Agent and environment loop")
        st.graphviz_chart("""
digraph G { rankdir=LR; node [shape=box, style="rounded,filled", fontname="Helvetica"];
edge [fontname="Helvetica", fontsize=10];
AG [label="Agent\\npolicy pi(s)", fillcolor="#DCE6F5", color="#0B4FA8"];
EN [label="Environment\\nP(s'|s,a), R(s,a,s')", fillcolor="#E8F5EE", color="#2E8B57"];
AG -> EN [label="action a_t", color="#E07B00"]; EN -> AG [label="next state s_t+1\\nreward r_t+1", color="#0B4FA8"]; }""")
        st.markdown("#### Transition graph")
        m, pi = st.session_state.get("mdp_model"), st.session_state.get("mdp_policy")
        if m is None or len(m["states"]) > 14:
            m, _ = mdp_from_table(incident_template())
            pi = value_iteration(m, 0.9)[2]
            st.caption("Showing the sample cybersecurity incident-response MDP (the current model is too large or not "
                       "run yet). Orange edges are the optimal policy; labels show action, probability and reward.")
        else:
            st.caption("Current model. Orange edges are the optimal policy; labels show action, probability and reward.")
        st.graphviz_chart(mdp_graph_dot(m, pi))
    elif section in ("Synthetic Data Lab", "Upload Real Data"):
        if section == "Synthetic Data Lab":
            st.markdown("## MDP Synthetic Data Lab")
            kind = st.radio("Synthetic model", ["Grid world delivery robot", "Random generic MDP",
                                                "Cybersecurity incident response (sample)"], horizontal=True,
                            key="mdp_kind")
            if kind == "Grid world delivery robot":
                c = st.columns(4)
                rows = c[0].slider("Rows", 3, 15, 6, key="g_rows")
                cols = c[1].slider("Columns", 3, 15, 8, key="g_cols")
                obs = c[2].slider("Obstacle share", 0.0, 0.35, 0.12, key="g_obs")
                haz = c[3].slider("Hazard share", 0.0, 0.2, 0.06, key="g_haz")
                c = st.columns(5)
                slip = c[0].slider("Slip probability", 0.0, 0.5, 0.1, key="g_slip")
                step_r = c[1].number_input("Step reward", -20.0, 0.0, -1.0, key="g_step")
                wait_r = c[2].number_input("Wait reward", -20.0, 0.0, -2.0, key="g_wait")
                goal_r = c[3].number_input("Goal reward", 0.0, 1000.0, 100.0, key="g_goal")
                haz_r = c[4].number_input("Hazard reward", -1000.0, 0.0, -50.0, key="g_haz_r")
                seed = st.number_input("Random seed", 0, 10 ** 6, 5, key="g_seed")
            elif kind == "Random generic MDP":
                c = st.columns(4)
                nS = c[0].slider("States", 3, 60, 10, key="r_ns")
                nA = c[1].slider("Actions", 2, 8, 3, key="r_na")
                br = c[2].slider("Next states per action", 1, 6, 3, key="r_br")
                seed = c[3].number_input("Random seed", 0, 10 ** 6, 9, key="r_seed")
            upl_df = None
        else:
            st.markdown("## MDP Upload Real Data")
            st.markdown("Upload a **CSV/Excel** transition table with columns `state`, `action`, `next_state`, "
                        "`probability`, `reward`. States with no outgoing rows are treated as terminal. "
                        "Probabilities for each state-action pair are normalised if they do not sum to 1.")
            template_button(incident_template(), "mdp_template.csv", key="mdp_tmpl")
            upl = st.file_uploader("Upload .csv or .xlsx", type=["csv", "xlsx"], key="mdp_upl")
            upl_df = read_upload_table(upl) if upl is not None else None
            kind = "upload"
        c = st.columns(2)
        gamma = c[0].slider("Discount factor gamma", 0.0, 0.99, 0.9, 0.01, key=f"mdp_gamma_{section}")
        algos = c[1].multiselect("Algorithms", ["Value Iteration", "Policy Iteration", "Q-Learning"],
                                 ["Value Iteration", "Policy Iteration", "Q-Learning"], key=f"mdp_algos_{section}")
        with st.expander("Q-learning settings"):
            c = st.columns(5)
            ql = dict(episodes=c[0].slider("Episodes", 50, 5000, 800, 50, key=f"ql_ep_{section}"),
                      alpha=c[1].slider("Learning rate alpha", 0.01, 1.0, 0.2, key=f"ql_a_{section}"),
                      eps_start=c[2].slider("Epsilon start", 0.0, 1.0, 1.0, key=f"ql_e0_{section}"),
                      eps_end=c[3].slider("Epsilon end", 0.0, 1.0, 0.05, key=f"ql_e1_{section}"),
                      max_steps=c[4].slider("Max steps per episode", 10, 500, 150, key=f"ql_ms_{section}"),
                      seed=0)
        start_choice = None
        if kind == "upload" and upl_df is not None:
            try:
                st_states = list(dict.fromkeys(list(upl_df.iloc[:, 0].astype(str))))
                start_choice = st.selectbox("Start state (for Q-learning)", st_states, key="mdp_start")
            except Exception:
                pass
        if st.button("Solve MDP", type="primary", key=f"mdp_go_{section}"):
            if not algos:
                st.error("Choose at least one algorithm.")
            else:
                try:
                    warnings = []
                    if kind == "Grid world delivery robot":
                        m = build_grid(rows, cols, obs, haz, slip, step_r, wait_r, goal_r, haz_r, int(seed))
                        src = f"synthetic {rows}x{cols} grid world (seed {seed})"
                    elif kind == "Random generic MDP":
                        m = random_mdp(nS, nA, br, int(seed))
                        src = f"synthetic random MDP with {nS} states and {nA} actions (seed {seed})"
                    elif kind.startswith("Cyber"):
                        m, warnings = mdp_from_table(incident_template())
                        src = "sample cybersecurity incident-response MDP"
                    else:
                        if upl_df is None:
                            raise ValueError("Please upload a file first.")
                        m, warnings = mdp_from_table(upl_df)
                        if start_choice in m["states"]:
                            m["start"] = m["states"].index(start_choice)
                        src = f"uploaded file {upl.name}"
                    with st.spinner("Solving"):
                        rep, res = run_mdp(m, gamma, algos, ql, src, warnings)
                    st.session_state["res_mdp"] = rep
                    st.session_state["mdp_model"] = m
                    ref = res.get("VI") or res.get("PI")
                    st.session_state["mdp_policy"] = ref[2] if ref else None
                except Exception as e:
                    st.error(f"Could not solve: {e}")
        if st.session_state.get("res_mdp"):
            show_report(st.session_state["res_mdp"], f"mdp_{section[:4]}")
    else:
        show_saved_result("res_mdp", "MDP")


# =====================================================================
# PAGE: PCA
# =====================================================================
PCA_USE_CASES = [
    ("Network intrusion and anomaly detection", "Compress traffic features (packets, bytes, duration, failed logins) "
     "into a few components. Connections that the components reconstruct poorly (high SPE) are unusual and worth "
     "investigating. The lab demonstrates exactly this."),
    ("Visualising high-dimensional data", "Project dozens of measurements onto PC1 and PC2 to see clusters, outliers "
     "and trends by eye, for example customer segments or gene-expression groups."),
    ("Pre-processing for machine learning", "Remove redundant, correlated features before training. This speeds up "
     "models, reduces overfitting and fixes multicollinearity in regression."),
    ("Compression and noise reduction", "Images, sensor streams and financial factor models keep the top components "
     "and discard minor ones, saving storage and filtering noise (for example eigenfaces, yield-curve factors)."),
]


def page_pca(section):
    if section == "Concept and Formulas":
        st.markdown("## PCA: Principal Component Analysis")
        st.markdown("**What it means.** PCA reduces the number of variables in a dataset while keeping as much of its "
                    "variation as possible. It builds new variables, the principal components, each a weighted "
                    "combination of the original features. This is called dimensionality reduction.")
        st.markdown("#### The five steps and their formulas")
        st.markdown("1. **Prepare the data:** handle missing values and standardise features with different scales.")
        st.latex(r"z_{ij}=\frac{x_{ij}-\bar x_j}{s_j}")
        st.markdown("2. **Calculate relationships:** the covariance (or correlation) matrix.")
        st.latex(r"C=\frac{1}{n-1}Z^{T}Z")
        st.markdown("3. **Find principal components:** eigenvectors of C (or the SVD of Z).")
        st.latex(r"C\,w_k=\lambda_k w_k,\qquad Z=U\Sigma W^{T}")
        st.markdown("4. **Rank components:** PC1 has the largest variance; each later component captures the most "
                    "remaining variance while being orthogonal to the earlier ones.")
        st.latex(r"\text{explained ratio}_k=\frac{\lambda_k}{\sum_j\lambda_j}")
        st.markdown("5. **Reduce dimensions:** keep k components and project the data.")
        st.latex(r"T=Z\,W_k\qquad(n\times p)\rightarrow(n\times k)")
        st.markdown("#### Anomaly scores used in this app")
        st.latex(r"SPE_i=\lVert z_i-\hat z_i\rVert^2,\ \ \hat z_i=W_kW_k^{T}z_i\qquad T^2_i=\sum_{k}\frac{t_{ik}^2}{\lambda_k}")
        st.markdown("SPE (squared prediction error) is large when a row does not follow the usual correlations; "
                    "Hotelling T² is large when a row is extreme along the kept components.")
        st.markdown("#### Explained variance example")
        ui_dataframe(pd.DataFrame({"Component": ["PC1", "PC2", "PC3", "PC4"], "Variance explained": ["60%", "25%", "10%", "5%"],
                                   "Cumulative": ["60%", "85%", "95%", "100%"]}))
        st.markdown("Keeping two components retains 85% of total variance. That may be enough for exploration, but it "
                    "does not guarantee that 85% of the information relevant to a prediction is kept.")
        st.markdown("**Simple intuition:** height in centimetres and height in inches carry the same information. PCA "
                    "merges them into one component that holds nearly all of their variation.")
        st.markdown("#### Python example (scikit-learn)")
        st.code('''from sklearn.preprocessing import StandardScaler
from sklearn.decomposition import PCA
import pandas as pd
data = pd.DataFrame({"packets": [100, 120, 180, 200, 250],
                     "data_volume": [10, 12, 19, 21, 26],
                     "duration": [5, 6, 9, 10, 13],
                     "failed_logins": [0, 1, 2, 1, 4]})
X = StandardScaler().fit_transform(data)
pca = PCA(n_components=2)
reduced = pca.fit_transform(X)
print(reduced, pca.explained_variance_ratio_)''', language="python")
        st.markdown("**Limitation:** PCA is unsupervised and follows variance, not the outcome you care about. A "
                    "low-variance feature can still be critical for detecting fraud or an attack.")
    elif section == "Use Cases":
        st.markdown("## PCA Use Cases")
        use_cases(PCA_USE_CASES)
    elif section == "Diagrams":
        st.markdown("## PCA Diagrams")
        st.graphviz_chart("""
digraph G { rankdir=LR; node [shape=box, style="rounded,filled", fillcolor="#DCE6F5", color="#0B4FA8", fontname="Helvetica"];
D [label="Original data\\n5 features\\npackets, volume, duration,\\nfailed logins, unusual requests"];
S [label="Standardise"]; C [label="Covariance\\nmatrix"]; E [label="Eigenvectors /\\nSVD"];
K [label="Keep top k\\ncomponents"]; R [label="Reduced data\\n2 components", fillcolor="#E8F5EE", color="#2E8B57"];
A [label="Reconstruction error\\nanomaly flags", fillcolor="#FBE4E1", color="#C0392B"];
D -> S -> C -> E -> K -> R; K -> A; }""")
        st.markdown("#### Geometric picture")
        rng = np.random.default_rng(1)
        x = rng.normal(0, 1, 300)
        y = 0.8 * x + rng.normal(0, 0.35, 300)
        Z = np.column_stack([x, y])
        p = PCA().fit(Z)
        fig, ax = plt.subplots(figsize=(6, 4.5))
        ax.scatter(x, y, s=10, alpha=0.5, color=BLUE)
        for vec, var, col, lab in zip(p.components_, p.explained_variance_, [ORANGE, GREEN], ["PC1", "PC2"]):
            v = vec * 2 * np.sqrt(var)
            ax.annotate("", xy=v, xytext=(0, 0), arrowprops=dict(arrowstyle="->", color=col, lw=3))
            ax.text(*(v * 1.12), lab, color=col, weight="bold")
        ax.set_aspect("equal")
        ax.set_title(f"PC1 follows the main spread ({p.explained_variance_ratio_[0] * 100:.0f}% of variance)")
        ax.set_xlabel("Feature 1")
        ax.set_ylabel("Feature 2")
        ui_image(fig_png(fig))
    elif section in ("Synthetic Data Lab", "Upload Real Data"):
        if section == "Synthetic Data Lab":
            st.markdown("## PCA Synthetic Data Lab")
            st.markdown("Synthetic network traffic: packets, data volume and duration share a 'traffic size' factor; "
                        "failed logins and unusual requests share an 'attack' factor. Injected anomalies break the "
                        "normal correlations.")
            c = st.columns(4)
            n = c[0].slider("Connections (rows)", 50, 5000, 500, 50, key="pca_n")
            af = c[1].slider("Anomaly share", 0.0, 0.2, 0.04, 0.01, key="pca_af")
            noise = c[2].slider("Extra pure-noise features", 0, 10, 0, key="pca_noise")
            seed = c[3].number_input("Random seed", 0, 10 ** 6, 21, key="pca_seed")
            if st.button("Generate synthetic data", key="pca_gen") or "pca_df" not in st.session_state:
                st.session_state["pca_df"] = synth_traffic(n, af, noise, int(seed))
                st.session_state["pca_src"] = f"synthetic network traffic (seed {seed})"
            df = st.session_state["pca_df"]
            src = st.session_state["pca_src"]
        else:
            st.markdown("## PCA Upload Real Data")
            st.markdown("Upload any **CSV/Excel** table. Choose the numeric feature columns and, optionally, a label "
                        "column used only for colouring and for comparing with the anomaly flags.")
            template_button(synth_traffic(200, 0.05, 0, 1), "pca_template.csv", key="pca_tmpl")
            upl = st.file_uploader("Upload .csv or .xlsx", type=["csv", "xlsx"], key="pca_upl")
            if upl is None:
                df = None
            else:
                try:
                    df = read_upload_table(upl)
                    src = f"uploaded file {upl.name}"
                except Exception as e:
                    st.error(f"Could not read the file: {e}")
                    df = None
        if df is not None:
            with st.expander(f"Preview data ({len(df)} rows, {df.shape[1]} columns)"):
                ui_dataframe(df.head(200), height=260)
            numeric = [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c])]
            others = [c for c in df.columns if c not in numeric or df[c].nunique() <= 10]
            feats = st.multiselect("Feature columns", numeric, numeric, key=f"pca_feats_{section}")
            c = st.columns(4)
            lab_opts = ["(none)"] + [c_ for c_ in df.columns if c_ not in feats]
            default_lab = lab_opts.index("label") if "label" in lab_opts else 0
            label = c[0].selectbox("Label column (optional)", lab_opts, default_lab, key=f"pca_lab_{section}")
            stdz = c[1].checkbox("Standardise features", True, key=f"pca_std_{section}")
            kmode = c[2].radio("Components", ["By variance target", "Fixed number"], key=f"pca_km_{section}")
            if kmode == "Fixed number":
                kv = c[3].slider("Components k", 1, max(1, len(feats)), min(2, max(1, len(feats))), key=f"pca_k_{section}")
                vt = 0.85
            else:
                vt = c[3].slider("Variance to keep", 0.5, 0.99, 0.85, key=f"pca_vt_{section}")
                kv = 2
            pct = st.slider("Anomaly threshold percentile (SPE)", 80, 99, 95, key=f"pca_pct_{section}")
            if st.button("Run PCA", type="primary", key=f"pca_go_{section}"):
                try:
                    with st.spinner("Computing components"):
                        st.session_state["res_pca"] = run_pca(df, feats, None if label == "(none)" else label, stdz,
                                                              "fixed" if kmode == "Fixed number" else "var", kv, vt,
                                                              pct, src)
                except Exception as e:
                    st.error(f"PCA failed: {e}")
        if st.session_state.get("res_pca"):
            show_report(st.session_state["res_pca"], f"pca_{section[:4]}")
    else:
        show_saved_result("res_pca", "PCA")


# =====================================================================
# PAGE: GROVER
# =====================================================================
GROVER_USE_CASES = [
    ("Unstructured search", "Find a marked record among N unsorted items with about (pi/4) sqrt(N) oracle queries "
     "instead of about N/2 classically. For a million items that is roughly 785 queries instead of 500,000."),
    ("Speeding up SAT and constraint satisfaction", "Let the oracle mark assignments that satisfy a formula. Grover "
     "then finds a solution quadratically faster than brute force over all 2^n assignments. The lab's 3-SAT oracle "
     "mode demonstrates this link."),
    ("Cryptanalysis and post-quantum security", "Brute-force key search falls from 2^n to about 2^(n/2) steps, which "
     "is why AES-256 (128-bit quantum security) is recommended over AES-128 for long-term protection."),
    ("Optimisation and minimum finding", "With a threshold oracle (Durr-Hoyer algorithm), Grover finds the minimum of "
     "a cost function, such as the cheapest route or lowest-risk schedule, in about sqrt(N) evaluations."),
]


def page_grover(section):
    if section == "Concept and Formulas":
        st.markdown("## Grover's Quantum Search Algorithm")
        st.markdown("**What it means.** Grover's algorithm searches N = 2ⁿ unsorted items for one that satisfies a "
                    "condition f(x) = 1. A classical computer needs about N/2 checks on average; Grover needs about "
                    "(π/4)√N. It is a provably optimal quadratic speedup for unstructured search.")
        st.markdown("#### 1. Uniform superposition")
        st.latex(r"|s\rangle=H^{\otimes n}|0\rangle^{\otimes n}=\frac{1}{\sqrt N}\sum_{x=0}^{N-1}|x\rangle")
        st.markdown("#### 2. Oracle: flip the sign of marked items")
        st.latex(r"U_f|x\rangle=(-1)^{f(x)}|x\rangle")
        st.markdown("#### 3. Diffusion: inversion about the mean")
        st.latex(r"D=2|s\rangle\langle s|-I=H^{\otimes n}\big(2|0\rangle\langle0|-I\big)H^{\otimes n},\qquad a_x\mapsto 2\bar a-a_x")
        st.markdown("#### 4. Geometry and iteration count")
        st.markdown("The state stays in the plane spanned by |α⟩ (unmarked) and |β⟩ (marked). Each Grover iteration "
                    "G = D·U_f rotates it by 2θ towards |β⟩.")
        st.latex(r"\sin\theta=\sqrt{M/N},\qquad G^k|s\rangle=\cos\big((2k+1)\theta\big)|\alpha\rangle+\sin\big((2k+1)\theta\big)|\beta\rangle")
        st.latex(r"P_{\text{success}}(k)=\sin^2\big((2k+1)\theta\big),\qquad k_{opt}\approx\frac{\pi}{4}\sqrt{\frac{N}{M}}")
        st.markdown("#### 5. Complexity")
        st.latex(r"\text{Classical: }O(N)\qquad\text{Grover: }O(\sqrt{N/M})\qquad\text{(optimal, Bennett-Bernstein-Brassard-Vazirani)}")
        ns = np.arange(1, 31)
        fig, ax = plt.subplots(figsize=(7, 3.2))
        ax.semilogy(ns, 2.0 ** ns / 2, color=GREY, lw=2, label="Classical about N/2")
        ax.semilogy(ns, np.pi / 4 * np.sqrt(2.0 ** ns), color=BLUE, lw=2, label="Grover about (pi/4) sqrt(N)")
        ax.set_xlabel("Qubits n (N = 2^n)")
        ax.set_ylabel("Oracle queries (log)")
        ax.set_title("Query cost: classical versus Grover")
        ax.legend(fontsize=8)
        ui_image(fig_png(fig))
        st.markdown("#### Practical notes")
        st.markdown("- **Overshooting:** iterating past k_opt reduces the success probability, because the rotation "
                    "continues past the target.\n"
                    "- **Unknown M:** the BBHT algorithm tries randomly growing iteration counts; quantum counting can "
                    "estimate M first.\n"
                    "- **The oracle is the cost:** speedup counts oracle calls; building the oracle circuit (for "
                    "example a 3-SAT checker) has its own cost.\n"
                    "- **Hardware:** this module is an exact classical statevector simulation (up to 14 qubits), "
                    "useful for learning; real devices add noise.")
    elif section == "Use Cases":
        st.markdown("## Grover Use Cases")
        use_cases(GROVER_USE_CASES)
    elif section == "Diagrams":
        st.markdown("## Grover Diagrams")
        st.graphviz_chart("""
digraph G { rankdir=LR; node [shape=box, style="rounded,filled", fillcolor="#DCE6F5", color="#0B4FA8", fontname="Helvetica"];
I [label="n qubits\\n|0...0>"]; H [label="Hadamard on all\\nuniform superposition"];
O [label="Oracle Uf\\nflip sign of marked", fillcolor="#FBE4E1", color="#C0392B"];
D [label="Diffusion D\\ninvert about mean", fillcolor="#E8F5EE", color="#2E8B57"];
L [label="Repeat k_opt\\ntimes?", shape=diamond]; M [label="Measure\\nmarked item with\\nhigh probability", fillcolor="#F4E3CC", color="#E07B00"];
I -> H -> O -> D -> L; L -> O [label="not yet"]; L -> M [label="done"]; }""")
        n = st.slider("Qubits for the diagrams", 2, 10, 4, key="gr_diag_n")
        theta, k = grover_theory(1 << n, 1)
        ui_image(fig_png(grover_circuit_fig(n, k)), f"Circuit for N = {1 << n}, one marked item, k = {k}")
        ui_image(fig_png(grover_geometry_fig(theta, k)), "Rotation towards the marked state")
    elif section in ("Synthetic Data Lab", "Upload Real Data"):
        labels = None
        if section == "Synthetic Data Lab":
            st.markdown("## Grover Synthetic Data Lab")
            kind = st.radio("Oracle", ["Random marked items", "3-SAT oracle (mark satisfying assignments)",
                                       "Synthetic user database"], horizontal=True, key="gr_kind")
            if kind == "Random marked items":
                c = st.columns(3)
                n = c[0].slider("Qubits n", 2, MAX_QUBITS, 6, key="gr_n")
                M = c[1].slider("Marked items M", 0, max(1, (1 << n) // 2), 1, key="gr_m")
                seed = c[2].number_input("Random seed", 0, 10 ** 6, 4, key="gr_seed")
                manual = st.text_input("Or type marked indices (comma separated, overrides M)", "", key="gr_manual")
            elif kind.startswith("3-SAT"):
                use_cur = st.session_state.get("sat_formula") and st.session_state["sat_formula"][1] <= MAX_QUBITS
                src_choice = st.radio("Formula", (["Use formula from 3-SAT module"] if use_cur else []) +
                                      ["New random 3-SAT formula"], horizontal=True, key="gr_satsrc")
                if src_choice == "New random 3-SAT formula":
                    c = st.columns(3)
                    n = c[0].slider("Variables (= qubits)", 3, 12, 8, key="gr_satn")
                    ratio = c[1].slider("Clause to variable ratio", 1.0, 6.0, 3.8, 0.1, key="gr_satr")
                    seed = c[2].number_input("Random seed", 0, 10 ** 6, 2, key="gr_sats")
            else:
                c = st.columns(4)
                n = c[0].slider("Qubits n (database holds 2^n records)", 3, 12, 8, key="gr_dbn")
                seed = c[1].number_input("Random seed", 0, 10 ** 6, 8, key="gr_dbs")
                db = synth_database(1 << n, int(seed))
                col = c[2].selectbox("Search column", list(db.columns), 3, key="gr_dbcol")
                op = c[3].selectbox("Condition", ["equals", "contains", "greater than", "less than"],
                                    2 if pd.api.types.is_numeric_dtype(db[col]) else 0, key="gr_dbop")
                val = st.text_input("Value", "97" if pd.api.types.is_numeric_dtype(db[col]) else str(db[col].iloc[0]),
                                    key="gr_dbval")
        else:
            st.markdown("## Grover Upload Real Data")
            st.markdown("Upload any **CSV/Excel** table as the 'database'. Pick a column and a condition; matching rows "
                        "become the marked items. Rows are padded to the next power of two (maximum "
                        f"{1 << MAX_QUBITS:,} rows).")
            template_button(synth_database(200, 3), "grover_template.csv", key="gr_tmpl")
            upl = st.file_uploader("Upload .csv or .xlsx", type=["csv", "xlsx"], key="gr_upl")
            db = None
            if upl is not None:
                try:
                    db = read_upload_table(upl)
                except Exception as e:
                    st.error(f"Could not read the file: {e}")
            if db is not None:
                c = st.columns(3)
                col = c[0].selectbox("Search column", list(db.columns), key="gr_ucol")
                op = c[1].selectbox("Condition", ["equals", "contains", "greater than", "less than"], key="gr_uop")
                val = c[2].text_input("Value", str(db[col].iloc[0]), key="gr_uval")
            kind = "upload"
        c = st.columns(3)
        shots = c[0].slider("Measurement shots", 100, 10000, 1000, 100, key=f"gr_shots_{section}")
        override = c[1].checkbox("Choose iteration count manually", key=f"gr_ov_{section}")
        k_man = c[2].slider("Iterations k", 0, 200, 3, key=f"gr_k_{section}") if override else None
        if st.button("Run Grover simulation", type="primary", key=f"gr_go_{section}"):
            try:
                if kind == "Random marked items":
                    N = 1 << n
                    if manual.strip():
                        marked = sorted({int(x) for x in re.split(r"[,\s]+", manual.strip()) if x})
                        if any(x < 0 or x >= N for x in marked):
                            raise ValueError(f"Indices must be between 0 and {N - 1}.")
                    else:
                        marked = sorted(np.random.default_rng(int(seed)).choice(N, M, replace=False).tolist())
                    labels = lambda i: f"item {i}"
                    src = f"synthetic search space with {len(marked)} marked items"
                elif kind.startswith("3-SAT"):
                    if src_choice.startswith("Use"):
                        clauses, n, names = st.session_state["sat_formula"]
                    else:
                        clauses, _ = random_3sat(n, int(round(ratio * n)), np.random.default_rng(int(seed)))
                        names = [f"x{i}" for i in range(1, n + 1)]
                    marked = np.nonzero(model_mask(clauses, n))[0].tolist()
                    nm = list(names)
                    labels = lambda i, n=n, nm=nm: ", ".join(f"{nm[v]}={'T' if (i >> (n - 1 - v)) & 1 else 'F'}"
                                                            for v in range(n))
                    src = f"3-SAT oracle over {n} variables and {len(clauses)} clauses"
                else:
                    series = db[col]
                    if len(db) > (1 << MAX_QUBITS):
                        st.warning(f"Only the first {1 << MAX_QUBITS:,} rows are used.")
                        db = db.iloc[:1 << MAX_QUBITS]
                        series = db[col]
                    if op in ("greater than", "less than"):
                        num = pd.to_numeric(series, errors="coerce")
                        v = float(val)
                        hit = (num > v) if op == "greater than" else (num < v)
                    elif op == "equals":
                        hit = series.astype(str).str.strip().str.lower() == str(val).strip().lower()
                    else:
                        hit = series.astype(str).str.contains(str(val), case=False, regex=False)
                    marked = np.nonzero(hit.fillna(False).values)[0].tolist()
                    n = max(1, math.ceil(math.log2(max(len(db), 2))))
                    idcol = db.columns[0]
                    labels = lambda i, db=db, idcol=idcol, col=col: (f"{idcol}={db.iloc[i][idcol]}, {col}={db.iloc[i][col]}"
                                                                     if i < len(db) else "empty padding slot")
                    src = (f"{'uploaded file ' + upl.name if kind == 'upload' else 'synthetic user database'}: "
                           f"{col} {op} {val} ({len(db)} rows)")
                with st.spinner("Simulating quantum state"):
                    rep, states, probs, mask = run_grover(n, marked, labels, src, shots, 0, k_man)
                st.session_state["res_grover"] = rep
                st.session_state["gr_states"] = (states, probs, mask, n)
            except Exception as e:
                st.error(f"Simulation failed: {e}")
        if st.session_state.get("gr_states"):
            states, probs, mask, n = st.session_state["gr_states"]
            with st.expander("Step through the iterations interactively", expanded=False):
                kk = st.slider("Iteration", 0, len(states) - 1, 0, key=f"gr_step_{section}")
                amp = states[kk]
                fig, ax = plt.subplots(figsize=(9, 2.8))
                if len(amp) <= 64:
                    ax.bar(range(len(amp)), amp, color=np.where(mask, ORANGE, BLUE))
                else:
                    ax.plot(amp, color=BLUE, lw=0.6)
                    mi = np.nonzero(mask)[0]
                    ax.scatter(mi, amp[mi], color=ORANGE, s=12, zorder=3)
                ax.axhline(amp.mean(), ls="--", color=GREY)
                ax.set_title(f"Amplitudes after {kk} iterations: P(marked) = {probs[kk] * 100:.1f}%")
                ui_image(fig_png(fig))
        if st.session_state.get("res_grover"):
            show_report(st.session_state["res_grover"], f"gr_{section[:4]}")
    else:
        show_saved_result("res_grover", "Grover")


# =====================================================================
# PAGE: KNOWLEDGE CHECK
# =====================================================================
QUIZ = [
    ("Which topic asks whether Boolean variables can satisfy all clauses?", ["PCA", "3-SAT", "MDP"], 1),
    ("Which expression is in CNF?", ["A AND (B OR C)", "A OR (B AND C)", "NOT (A OR B)"], 0),
    ("What does logical equivalence mean?", ["Both expressions always have the same truth value",
                                             "Both expressions have the same number of variables",
                                             "Both expressions contain AND"], 0),
    ("What does an MDP optimise?", ["The number of variables", "Expected long-term reward", "The variance of data"], 1),
    ("What is PCA primarily used for?", ["Boolean satisfiability", "Dimensionality reduction", "Logical proof only"], 1),
    ("About how many oracle queries does Grover need for N items with one marked?",
     ["About N/2", "About (pi/4) sqrt(N)", "About log2(N)"], 1),
    ("Random 3-SAT is hardest near which clause-to-variable ratio?", ["1.0", "4.26", "10"], 1),
    ("What does the Tseitin transformation produce?", ["An equivalent DNF", "An equisatisfiable CNF of linear size",
                                                       "A truth table"], 1),
]


def page_quiz():
    st.markdown("## Knowledge Check")
    answers = []
    for i, (q, opts, _) in enumerate(QUIZ):
        answers.append(st.radio(f"{i + 1}. {q}", opts, index=None, key=f"quiz_{i}"))
    c = st.columns(2)
    if c[0].button("Check answers", type="primary", key="quiz_check"):
        rows = []
        for (q, opts, ok), a in zip(QUIZ, answers):
            rows.append({"question": q, "your_answer": a or "(no answer)", "correct_answer": opts[ok],
                         "correct": a == opts[ok]})
        df = pd.DataFrame(rows)
        rep = Report("Knowledge Check Results", "Knowledge Check")
        score = int(df["correct"].sum())
        rep.metrics({"Score": f"{score} of {len(QUIZ)}", "Percent": f"{score / len(QUIZ) * 100:.0f}%"})
        rep.table("Answers", df)
        rep.para("Recommended order: Logical Equivalence and CNF first, then 3-SAT, then Grover, MDP and PCA.")
        st.session_state["res_quiz"] = rep
    if c[1].button("Try again", key="quiz_reset"):
        for i in range(len(QUIZ)):
            st.session_state.pop(f"quiz_{i}", None)
        st.session_state.pop("res_quiz", None)
        st.rerun()
    if st.session_state.get("res_quiz"):
        show_report(st.session_state["res_quiz"], "quiz")


# =====================================================================
# MAIN
# =====================================================================
CSS = f"""
<style>
.knet-title {{ color:{BLUE} !important; font-weight:800 !important; font-size:2.4rem !important; line-height:1.25 !important; margin:0 0 0.15rem 0 !important; }}
.knet-header {{ padding:0.2rem 0 0.9rem 0; border-bottom:3px solid {BLUE}; margin-bottom:1.1rem; }}
[data-testid="stSidebar"] {{ background:#F2F5FB; }}
[data-testid="stMetricValue"] {{ font-size:1.25rem; color:{BLUE}; }}
</style>
"""

MODULES = ["Platform Overview", "3-SAT Satisfiability", "Logical Equivalence", "CNF Conversion",
           "Markov Decision Process", "Principal Component Analysis", "Grover Quantum Search", "Knowledge Check"]
SECTIONS = ["Concept and Formulas", "Use Cases", "Diagrams", "Synthetic Data Lab", "Upload Real Data",
            "Results and Export"]


def main():
    st.set_page_config(page_title=APP_NAME, layout="wide", initial_sidebar_state="expanded")
    st.markdown(CSS, unsafe_allow_html=True)
    st.markdown(f'<div class="knet-header"><div class="knet-title">{APP_NAME}</div>'
                f'<div class="knet-title">{DEV_LINE}</div></div>', unsafe_allow_html=True)
    with st.sidebar:
        st.markdown("### Navigation")
        module = st.radio("Module", MODULES, key="nav_module")
        section = None
        if module not in ("Platform Overview", "Knowledge Check"):
            st.markdown("---")
            section = st.radio("Section", SECTIONS, key="nav_section")
        st.markdown("---")
        st.caption("Every result exports to PDF, Word, Text and CSV. Each upload section offers a sample template.")
        st.caption("Kalsnet (KNet) Consulting Group")
    pages = {"3-SAT Satisfiability": page_sat, "Logical Equivalence": page_le, "CNF Conversion": page_cnf,
             "Markov Decision Process": page_mdp, "Principal Component Analysis": page_pca,
             "Grover Quantum Search": page_grover}
    if module == "Platform Overview":
        page_overview()
    elif module == "Knowledge Check":
        page_quiz()
    else:
        pages[module](section)


if __name__ == "__main__":
    main()
