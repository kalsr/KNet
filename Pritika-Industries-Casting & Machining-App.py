

# AI for Auto Ancillary Units — Casting & Machining
# Single-file Streamlit application.
# Developed by Randy Singh | Kalsnet (KNet) Consulting Group
# Modules:
#   1. Vision QC - Casting Defect Detection
#   2. Inspection & NDT - Dimensional & Structural Quality Assurance
#   3. Predictive Maintenance - Machine Health
#   4. Tool Wear Prediction
#   5. Demand & Inventory Forecast
#   6. Energy Optimization
#   7. Traceability / Root Cause
#   8. ROI Summary
# Each module has:
#   - Generate Synthetic Data button
#   - Upload Real Data (CSV/Excel) - real data takes priority over synthetic when present
#   - Charts / tables
#   - Formula / methodology explanation, including a field-by-field schema explanation
#   - Export to CSV / JSON / TXT / Word / PDF
#   - Reset Data button

import io
import logging
import os
import datetime as dt
from datetime import datetime, timedelta
from typing import Dict, Optional

import numpy as np
import pandas as pd
import streamlit as st

try:
    from fpdf import FPDF
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# ----------------------------------------------------------------------------
# Logging
# ----------------------------------------------------------------------------
LOG_DIR = os.environ.get("AAU_LOG_DIR", "./logs")
os.makedirs(LOG_DIR, exist_ok=True)
logger = logging.getLogger("auto_ancillary_units")
if not logger.handlers:
    logger.setLevel(logging.INFO)
    _fh = logging.FileHandler(os.path.join(LOG_DIR, "auto_ancillary_units.log"))
    _fh.setFormatter(logging.Formatter("%(asctime)s | %(levelname)s | %(message)s"))
    logger.addHandler(_fh)

# ----------------------------------------------------------------------------
# Page config
# ----------------------------------------------------------------------------
st.set_page_config(
    page_title="Pritika Industries - Casting & Machining/Auto Ancillary Units",
    layout="wide",
)

RNG_SEED = 42

# ----------------------------------------------------------------------------
# Styling - bold blue title bar
# ----------------------------------------------------------------------------
st.markdown(
    """
    <style>
    .main-title {
        text-align: center;
        color: #0047AB;
        font-weight: 900;
        font-size: 2.6rem;
        margin-bottom: 0.3rem;
        letter-spacing: 0.6px;
        font-family: 'Arial Black', sans-serif;
    }
    .subtitle-line {
        text-align: center;
        color: #0047AB;
        font-weight: 800;
        font-size: 1.15rem;
        margin-top: 0;
        margin-bottom: 1.1rem;
        font-family: 'Arial', sans-serif;
    }
    .section-header {
        color: #0047AB;
        font-weight: 800;
        font-size: 1.6rem;
        margin-top: 0.5rem;
        margin-bottom: 0.4rem;
    }
    .formula-box {
        background: #fff9e6;
        border: 2px solid #ffd700;
        padding: 1rem;
        border-radius: 6px;
        font-size: 0.92rem;
        margin: 0.8rem 0;
    }
    .info-box {
        background: #e8f4f8;
        border-left: 4px solid #0099cc;
        padding: 0.8rem;
        border-radius: 4px;
        font-size: 0.85rem;
        margin: 0.6rem 0;
    }
    .divider-line {
        border-top: 2px solid #0047AB;
        margin: 1rem 0;
        opacity: 0.3;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# Top title bar - always visible above the module navigation, both lines bold blue
st.markdown('<div class="main-title">Pritika Industries - Casting &amp; Machining/Auto Ancillary Units.</div>', unsafe_allow_html=True)
st.markdown('<div class="subtitle-line">Developed by Randy Singh, Kalsnet (KNet) Consulting Group.</div>', unsafe_allow_html=True)
st.markdown('<div class="divider-line"></div>', unsafe_allow_html=True)

# ----------------------------------------------------------------------------
# Export helpers (CSV / JSON / TXT / Word / PDF)
# ----------------------------------------------------------------------------
# FPDF's built-in core fonts only support Latin-1. Emoji, arrows, currency
# glyphs, smart quotes and em-dashes crash pdf.cell()/multi_cell() with a
# Unicode error unless sanitized first. This map is a defensive safety net
# for whatever text ends up in an export (including uploaded real data),
# not a UI element itself, so it is left in place.
_PDF_CHAR_MAP = {
    "•": "- ",
    "Rs.": "Rs.",
    "₹": "Rs. ",
    "–": "-", "—": "-",
    "’": "'", "‘": "'", "“": '"', "”": '"', "…": "...",
}


def _pdf_safe(value) -> str:
    """Make text safe for FPDF's Latin-1-only core fonts."""
    text = str(value)
    for original, replacement in _PDF_CHAR_MAP.items():
        text = text.replace(original, replacement)
    return text.encode("latin-1", "ignore").decode("latin-1").strip()


def export_to_pdf(df: pd.DataFrame, title: str, formulas: Optional[Dict] = None) -> Optional[bytes]:
    """Export dataframe (+ optional formula explanations) to PDF."""
    if not HAS_PDF:
        return None
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", "B", 16)
        pdf.cell(0, 10, _pdf_safe(title), ln=True, align="C")
        pdf.set_font("Arial", "B", 10)
        pdf.cell(0, 5, f"Generated: {dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
        pdf.cell(0, 5, "Pritika Industries | AI for Auto Ancillary Units | Kalsnet (KNet) Consulting Group", ln=True)
        pdf.set_font("Arial", "", 8)
        pdf.ln(5)

        col_width = 190 / max(1, len(df.columns))
        row_height = 7
        page_bottom = pdf.h - pdf.b_margin

        def draw_header():
            pdf.set_font("Arial", "B", 8)
            for col in df.columns:
                pdf.cell(col_width, row_height, _pdf_safe(col)[:20], border=1, align="C")
            pdf.ln()
            pdf.set_font("Arial", "", 8)

        draw_header()
        for _, row in df.head(60).iterrows():
            if pdf.get_y() + row_height > page_bottom:
                pdf.add_page()
                draw_header()
            for val in row:
                pdf.cell(col_width, row_height, _pdf_safe(val)[:20], border=1, align="L")
            pdf.ln()

        if formulas:
            pdf.add_page()
            pdf.set_font("Arial", "B", 12)
            pdf.cell(0, 10, "Schema Field Definitions, Process and Formulas", ln=True)
            pdf.set_font("Arial", "", 9)
            for name, text in formulas.items():
                pdf.cell(0, 5, _pdf_safe(name) + ":", ln=True)
                pdf.multi_cell(0, 4, _pdf_safe(text)[:2000])
                pdf.ln(2)

        raw_output = pdf.output(dest="S")
        if isinstance(raw_output, str):
            return raw_output.encode("latin-1")
        return bytes(raw_output)
    except Exception as e:
        st.error(f"PDF export error: {e}")
        logger.exception("PDF export failed")
        return None


def export_to_docx(df: pd.DataFrame, title: str, formulas: Optional[Dict] = None) -> Optional[bytes]:
    """Export dataframe (+ optional formula explanations) to a Word document."""
    if not HAS_DOCX:
        return None
    try:
        doc = Document()
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(title)
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0, 71, 171)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        meta_para = doc.add_paragraph()
        meta_run = meta_para.add_run(f"Generated: {dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        meta_run.font.size = Pt(10)
        meta_run = meta_para.add_run(
            "Pritika Industries | AI for Auto Ancillary Units | Developed by Randy Singh | "
            "Kalsnet (KNet) Consulting Group"
        )
        meta_run.font.size = Pt(10)
        meta_run.italic = True

        doc.add_paragraph()
        table = doc.add_table(rows=1, cols=len(df.columns))
        table.style = "Light Grid Accent 1"
        hdr_cells = table.rows[0].cells
        for i, col in enumerate(df.columns):
            hdr_cells[i].text = str(col)
        for _, row in df.head(60).iterrows():
            row_cells = table.add_row().cells
            for i, val in enumerate(row):
                row_cells[i].text = str(val)[:100]

        if formulas:
            doc.add_page_break()
            heading_para = doc.add_paragraph()
            heading_run = heading_para.add_run("Schema Field Definitions, Process and Formulas")
            heading_run.font.size = Pt(14)
            heading_run.font.bold = True
            heading_run.font.color.rgb = RGBColor(0, 71, 171)
            for name, text in formulas.items():
                doc.add_paragraph(name, style="Heading 3")
                doc.add_paragraph(text[:2000])

        docx_bytes = io.BytesIO()
        doc.save(docx_bytes)
        return docx_bytes.getvalue()
    except Exception as e:
        st.error(f"Word export error: {e}")
        logger.exception("Word export failed")
        return None


def export_to_text(df: pd.DataFrame, title: str, formulas: Optional[Dict] = None) -> str:
    """Export to plain text format."""
    text = f"\n{'=' * 80}\n{title}\n{'=' * 80}\n\n"
    text += f"Generated: {dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
    text += "Pritika Industries | AI for Auto Ancillary Units | Developed by Randy Singh | Kalsnet (KNet) Consulting Group\n\n"
    text += df.to_string()
    if formulas:
        text += f"\n\n{'=' * 80}\nSchema Field Definitions, Process and Formulas\n{'=' * 80}\n\n"
        for name, val in formulas.items():
            text += f"\n{name}:\n{val}\n"
    return text


def load_uploaded_table(uploaded_file) -> Optional[pd.DataFrame]:
    """Read an uploaded CSV/Excel file into a DataFrame with a friendly error on failure."""
    try:
        if uploaded_file.name.lower().endswith(".csv"):
            return pd.read_csv(uploaded_file)
        return pd.read_excel(uploaded_file)
    except Exception as e:
        st.error(f"Could not read file: {e}")
        logger.exception("Failed to read uploaded file")
        return None


def full_export_buttons(df: pd.DataFrame, title: str, formulas: Optional[Dict], key_prefix: str, file_stub: str):
    """Render a row of CSV / JSON / TXT / Word / PDF export buttons for a dataframe."""
    st.markdown("**Export Results:**")
    col1, col2, col3, col4, col5 = st.columns(5)
    with col1:
        st.download_button("CSV", df.to_csv(index=False).encode(), f"{file_stub}.csv", "text/csv",
                            key=f"csv_{key_prefix}", use_container_width=True)
    with col2:
        st.download_button("JSON", df.to_json(orient="records", indent=2).encode(), f"{file_stub}.json",
                            "application/json", key=f"json_{key_prefix}", use_container_width=True)
    with col3:
        txt = export_to_text(df, title, formulas)
        st.download_button("TXT", txt.encode(), f"{file_stub}.txt", "text/plain",
                            key=f"txt_{key_prefix}", use_container_width=True)
    with col4:
        if HAS_DOCX:
            docx_bytes = export_to_docx(df, title, formulas)
            if docx_bytes:
                st.download_button("WORD", docx_bytes, f"{file_stub}.docx",
                                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                    key=f"docx_{key_prefix}", use_container_width=True)
        else:
            st.caption("Word export unavailable")
    with col5:
        if HAS_PDF:
            pdf_bytes = export_to_pdf(df, title, formulas)
            if pdf_bytes:
                st.download_button("PDF", pdf_bytes, f"{file_stub}.pdf", "application/pdf",
                                    key=f"pdf_{key_prefix}", use_container_width=True)
        else:
            st.caption("PDF export unavailable")


# ----------------------------------------------------------------------------
# Generic helpers
# ----------------------------------------------------------------------------
def df_download_button(df: pd.DataFrame, label: str, filename: str, key: str):
    """Render a CSV export button for a dataframe (kept for backward compatibility)."""
    buf = io.StringIO()
    df.to_csv(buf, index=False)
    st.download_button(
        label=label,
        data=buf.getvalue(),
        file_name=filename,
        mime="text/csv",
        key=key,
        use_container_width=True,
    )


def reset_button(state_key: str, label: str = "Reset Data", key: str = None, also_clear=None):
    """Clear synthetic (and optionally real-upload) state and rerun to regenerate a fresh batch."""
    if st.button(label, key=key or f"reset_{state_key}", use_container_width=True):
        if state_key in st.session_state:
            del st.session_state[state_key]
        if also_clear:
            for k in also_clear:
                if k in st.session_state:
                    del st.session_state[k]
        st.rerun()


def upload_real_data(label: str, real_key: str, expected_cols: str, upload_widget_key: str, clear_key: str):
    """Render an upload widget + clear button for real-world data, storing it in session_state[real_key]."""
    st.markdown(f"**OR Upload Real Data - {label}**")
    st.caption(f"Expected columns (case-insensitive recommended): {expected_cols}")
    col_u1, col_u2 = st.columns([4, 1])
    with col_u1:
        uploaded = st.file_uploader("Upload CSV/Excel file", type=["csv", "xlsx"], key=upload_widget_key,
                                     label_visibility="collapsed")
    with col_u2:
        if st.session_state.get(real_key) is not None:
            if st.button("Clear Upload", key=clear_key, use_container_width=True):
                st.session_state[real_key] = None
                st.rerun()
    if uploaded is not None:
        real_df = load_uploaded_table(uploaded)
        if real_df is not None:
            st.session_state[real_key] = real_df
            st.success(f"Loaded {len(real_df)} real records - now overriding synthetic data below.")


def pick_active_df(real_key: str, synth_key: str):
    """Real uploaded data takes priority over synthetic data when both are present."""
    real_df = st.session_state.get(real_key)
    if real_df is not None:
        return real_df, True
    return st.session_state.get(synth_key), False


def section_title(title: str, subtitle: str):
    st.markdown(f"## {title}")
    st.caption(subtitle)
    st.divider()


def kpi_row(items):
    cols = st.columns(len(items))
    for c, (label, value, delta) in zip(cols, items):
        c.metric(label, value, delta)


def formula_box(markdown_text: str):
    st.markdown('<div class="formula-box">', unsafe_allow_html=True)
    st.markdown(markdown_text)
    st.markdown("</div>", unsafe_allow_html=True)


# ----------------------------------------------------------------------------
# Synthetic data generators
# ----------------------------------------------------------------------------
def gen_casting_defects(n_days=30, seed=RNG_SEED):
    rng = np.random.default_rng(seed)
    dates = pd.date_range(end=datetime.today(), periods=n_days)
    defect_types = ["Blowhole", "Shrinkage", "Cold Shut", "Gas Porosity", "Surface Crack"]
    rows = []
    base_rate = 0.035
    for i, d in enumerate(dates):
        produced = rng.integers(800, 1400)
        trend = base_rate * (1 - 0.5 * i / n_days)
        rejected = int(produced * max(0.01, rng.normal(trend, 0.006)))
        rows.append({
            "Date": d.date(),
            "Produced": produced,
            "Rejected": rejected,
            "Rejection_%": round(100 * rejected / produced, 2),
            "Top_Defect": rng.choice(defect_types, p=[0.30, 0.25, 0.15, 0.20, 0.10]),
            "Melt_Temp_C": round(rng.normal(720, 8), 1),
            "Pour_Temp_C": round(rng.normal(700, 6), 1),
            "Mold_Temp_C": round(rng.normal(180, 10), 1),
            "Humidity_%": round(rng.normal(55, 7), 1),
        })
    return pd.DataFrame(rows)


def gen_inspection_qc(n_records=180, seed=RNG_SEED):
    """
    Dimensional & structural inspection / NDT synthetic data.

    Models a 3-tier inspection strategy:
      Tier 1 - In-line vision + hand gauges, 100% of parts, lowest cost.
      Tier 2 - Statistical sampling per shift/lot with optical 3D scan / arm CMM.
      Tier 3 - Melt-lot certification: CMM + NDT (CT/UT/DPT/MPI) + hardness + XRF,
               with a destructive tensile test run on a subset of events.
    """
    rng = np.random.default_rng(seed + 6)
    tiers = [
        "Tier 1 - In-line Vision/Gauge",
        "Tier 2 - Optical 3D/Arm CMM Sampling",
        "Tier 3 - CMM + NDT + Destructive Certification",
    ]
    tier_p = [0.70, 0.22, 0.08]
    instruments_by_tier = {
        tiers[0]: ["Vision System (2D/3D)", "Digital Caliper", "Height Gauge", "Go/No-Go Gauge"],
        tiers[1]: ["Structured-Light 3D Scanner", "Laser 3D Scanner", "Portable Arm CMM"],
        tiers[2]: ["Bridge CMM", "Industrial CT Scanner", "Ultrasonic Tester (UT)",
                   "Dye Penetrant / MPI", "Portable Hardness Tester",
                   "Tensile Testing Machine", "Portable XRF Analyzer"],
    }
    ndt_methods = ["None", "Dye Penetrant", "MPI", "Ultrasonic Testing", "CT Scan"]
    ndt_p = [0.35, 0.20, 0.15, 0.15, 0.15]
    defect_types = ["None", "Porosity", "Shrinkage", "Inclusion", "Crack"]
    melt_lots = [f"ML-{100 + i}" for i in range(25)]
    dates = pd.date_range(end=datetime.today(), periods=n_records)

    rows = []
    for i in range(n_records):
        tier = rng.choice(tiers, p=tier_p)
        instrument = rng.choice(instruments_by_tier[tier])

        tolerance = round(rng.uniform(0.02, 0.15), 3)
        deviation = round(abs(rng.normal(0, tolerance * 0.55)), 4)
        dim_pass = deviation <= tolerance

        is_tier3 = tier == tiers[2]
        ndt_method = rng.choice(ndt_methods, p=ndt_p) if is_tier3 else "None"
        defect_p = [0.90, 0.04, 0.03, 0.02, 0.01] if ndt_method != "None" else [1.0, 0.0, 0.0, 0.0, 0.0]
        internal_defect = rng.choice(defect_types, p=defect_p)

        hardness = round(rng.normal(220, 15), 1)  # HB - illustrative for a typical casting alloy
        ai_predicted_tensile = round(hardness * 1.55 + rng.normal(0, 8), 1)  # proxy correlation model

        actual_tensile = np.nan
        if is_tier3 and rng.random() < 0.4:
            actual_tensile = round(ai_predicted_tensile + rng.normal(0, max(1.0, ai_predicted_tensile * 0.04)), 1)

        base_cost = {
            tiers[0]: rng.uniform(5, 25),
            tiers[1]: rng.uniform(150, 450),
            tiers[2]: rng.uniform(800, 3500),
        }[tier]
        if not np.isnan(actual_tensile):
            base_cost += rng.uniform(1200, 2500)  # destructive coupon test add-on cost

        overall_result = "Pass" if (dim_pass and internal_defect == "None") else "Fail"

        rows.append({
            "Date": dates[i].date(),
            "Part_ID": f"CST-{2000 + i}",
            "Melt_Lot": rng.choice(melt_lots),
            "Inspection_Tier": tier,
            "Instrument_Used": instrument,
            "Sample_Qty": int(rng.integers(1, 3)) if tier == tiers[0] else int(rng.integers(1, 12)),
            "Dimensional_Deviation_mm": deviation,
            "Tolerance_Limit_mm": tolerance,
            "Dimensional_Result": "Pass" if dim_pass else "Fail",
            "NDT_Method": ndt_method,
            "Internal_Defect_Found": internal_defect,
            "Hardness_HB": hardness,
            "AI_Predicted_Tensile_MPa": ai_predicted_tensile,
            "Actual_Tensile_MPa": actual_tensile,
            "Inspection_Cost_Rs": round(base_cost, 0),
            "Overall_Result": overall_result,
        })
    return pd.DataFrame(rows)


def gen_machine_health(n_hours=240, seed=RNG_SEED):
    rng = np.random.default_rng(seed + 1)
    machines = ["VMC-01", "VMC-02", "HMC-01", "Press-01", "Press-02"]
    rows = []
    start = datetime.today() - timedelta(hours=n_hours)
    for m in machines:
        degradation_start = rng.integers(int(n_hours * 0.6), n_hours)
        for h in range(n_hours):
            t = start + timedelta(hours=h)
            failing = h > degradation_start
            vib = rng.normal(2.5, 0.3) + (3.0 * (h - degradation_start) / n_hours if failing else 0)
            temp = rng.normal(55, 3) + (15 * (h - degradation_start) / n_hours if failing else 0)
            current = rng.normal(12, 1) + (4 * (h - degradation_start) / n_hours if failing else 0)
            risk = min(99, max(1, (vib / 6 + temp / 90 + current / 18) * 33))
            rows.append({
                "Timestamp": t,
                "Machine": m,
                "Vibration_mm_s": round(vib, 2),
                "Temperature_C": round(temp, 1),
                "Current_A": round(current, 2),
                "Failure_Risk_%": round(risk, 1),
                "Predicted_Days_to_Failure": round(max(0.5, 21 - risk / 5), 1),
            })
    return pd.DataFrame(rows)


def gen_tool_wear(n_records=200, seed=RNG_SEED):
    rng = np.random.default_rng(seed + 2)
    tools = ["Insert-CNMG", "Insert-WNMG", "Drill-12mm", "Boring-Bar", "End-Mill-10mm"]
    rows = []
    for i in range(n_records):
        tool = rng.choice(tools)
        cycles = rng.integers(10, 500)
        wear = min(100, round(cycles / rng.uniform(4, 7) + rng.normal(0, 4), 1))
        wear = max(0, wear)
        spindle_load = round(rng.normal(60, 10) + wear * 0.2, 1)
        vib = round(rng.normal(1.8, 0.4) + wear * 0.02, 2)
        recommended_change = "Change Now" if wear > 80 else ("Monitor" if wear > 55 else "OK")
        rows.append({
            "Tool_ID": f"{tool}-{i % 30:02d}",
            "Tool_Type": tool,
            "Cycles_Run": cycles,
            "Wear_%": wear,
            "Spindle_Load_%": spindle_load,
            "Vibration_mm_s": vib,
            "Recommendation": recommended_change,
        })
    return pd.DataFrame(rows)


def gen_demand_forecast(n_weeks=26, seed=RNG_SEED):
    rng = np.random.default_rng(seed + 3)
    parts = ["Bracket-A1", "Housing-B2", "Gear-C3", "Hub-D4"]
    rows = []
    weeks = pd.date_range(end=datetime.today(), periods=n_weeks, freq="W")
    for p in parts:
        base = rng.integers(2000, 6000)
        season_amp = base * 0.15
        for i, w in enumerate(weeks):
            actual = None if i >= n_weeks - 4 else int(
                base + season_amp * np.sin(i / 4) + rng.normal(0, base * 0.05)
            )
            forecast = int(base + season_amp * np.sin(i / 4) + rng.normal(0, base * 0.02))
            rows.append({
                "Week": w.date(),
                "Part": p,
                "Actual_Demand": actual,
                "AI_Forecast": max(0, forecast),
                "Safety_Stock": int(forecast * 0.12),
                "Reorder_Point": int(forecast * 0.35),
            })
    return pd.DataFrame(rows)


def gen_energy(n_days=30, seed=RNG_SEED):
    rng = np.random.default_rng(seed + 4)
    dates = pd.date_range(end=datetime.today(), periods=n_days)
    rows = []
    for d in dates:
        baseline_kwh = rng.normal(4200, 200)
        optimized_kwh = baseline_kwh * rng.uniform(0.82, 0.92)
        rows.append({
            "Date": d.date(),
            "Baseline_kWh": round(baseline_kwh, 0),
            "AI_Optimized_kWh": round(optimized_kwh, 0),
            "Savings_%": round(100 * (1 - optimized_kwh / baseline_kwh), 1),
            "Peak_Load_Shifted_kWh": round(rng.uniform(150, 500), 0),
        })
    return pd.DataFrame(rows)


def gen_traceability(n_events=15, seed=RNG_SEED):
    rng = np.random.default_rng(seed + 5)
    machines = ["VMC-01", "VMC-02", "HMC-01"]
    operators = ["Op-A", "Op-B", "Op-C", "Op-D"]
    rows = []
    for i in range(n_events):
        rows.append({
            "Defect_ID": f"DEF-{1000 + i}",
            # FIX: rng.integers() returns numpy.int64, which Python's
            # timedelta() constructor rejects ("unsupported type for
            # timedelta days component: numpy.int64") - this was the
            # exact TypeError that crashed this module. Casting to a
            # native int fixes it.
            "Date": (datetime.today() - timedelta(days=int(rng.integers(0, 30)))).date(),
            "Melt_Lot": f"ML-{rng.integers(100, 999)}",
            "Machine": rng.choice(machines),
            "Operator": rng.choice(operators),
            "Tool_Used": f"Insert-{rng.integers(1, 9)}",
            "Root_Cause_AI": rng.choice([
                "Die temperature deviation",
                "Raw material batch variance",
                "Tool wear beyond threshold",
                "Coolant pressure drop",
                "Operator parameter override",
            ]),
            "Manual_RCA_Time_Hrs": round(rng.uniform(8, 48), 1),
            "AI_RCA_Time_Min": round(rng.uniform(0.3, 3), 1),
        })
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------
# Sidebar
# ----------------------------------------------------------------------------
st.sidebar.title("AI Ops Console")
st.sidebar.markdown(
    "AI Ops dashboard for **casting + machining auto ancillary units** - "
    "scrap reduction, dimensional/structural inspection, downtime prediction, tool cost, "
    "demand forecasting, energy savings, and root-cause tracing."
)
st.sidebar.divider()

module = st.sidebar.radio(
    "Select Module",
    [
        "Overview",
        "Vision QC - Casting Defects",
        "Inspection & NDT - Dimensional/Structural QA",
        "Predictive Maintenance",
        "Tool Wear Prediction",
        "Demand & Inventory Forecast",
        "Energy Optimization",
        "Traceability & Root Cause",
        "ROI Summary",
    ],
)
st.sidebar.divider()
st.sidebar.caption("Synthetic data is for illustrative purposes. Upload real shop-floor data in any module to override it.")

# ----------------------------------------------------------------------------
# OVERVIEW
# ----------------------------------------------------------------------------
if module == "Overview":
    st.markdown(
        """
This app shows where AI creates measurable value on the shop floor of a
**casting + machining auto ancillary unit**: cutting scrap, verifying dimensional and structural
quality at competitive cost, predicting downtime, optimizing tool cost, sharpening demand
forecasts, trimming energy spend, and collapsing root-cause analysis from days to minutes.

Use the sidebar to open each module. Every module lets you:
- **Generate synthetic data** representing a live shop-floor feed
- **Upload your own real data** (CSV/Excel) - it automatically takes priority over synthetic data
- **Visualize** the key metrics AI would surface, with the underlying schema fields, process
  and formula explained
- **Export** results to CSV, JSON, TXT, Word, or PDF
- **Reset** to clear the generated/uploaded data and start fresh
        """
    )
    st.divider()
    kpi_row([
        ("Casting Rejection", "3.5% to 1.8%", "-48%"),
        ("Machining Tool Cost/Part", "Down 12-18%", None),
        ("Unplanned Downtime", "Down 25%", None),
        ("Inventory Reduction", "15-25%", None),
    ])
    st.divider()
    st.markdown("### Module Map")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.markdown("**Vision QC**\nCatch blowholes, shrinkage, cold shuts, porosity, cracks.")
        st.markdown("**Inspection & NDT**\nCMM/3D scanning + NDT tiers to certify dimensional accuracy and structural strength at low cost.")
        st.markdown("**Predictive Maintenance**\nFlag bearing/spindle/die failures 2-4 weeks early.")
    with c2:
        st.markdown("**Tool Wear**\nReplace inserts on condition, not fixed schedule.")
        st.markdown("**Demand Forecast**\nForecast FG/RM better than Excel for JIT OEM supply.")
    with c3:
        st.markdown("**Energy Optimization**\nShift furnace/machining load to cut bills 8-20%.")
        st.markdown("**Root Cause**\nCorrelate batch, machine, tool, operator in seconds.")

# ----------------------------------------------------------------------------
# 1. VISION QC
# ----------------------------------------------------------------------------
elif module == "Vision QC - Casting Defects":
    section_title("Vision QC - Casting Defect Detection",
                  "Computer-vision style defect tracking: blowholes, shrinkage, cold shuts, porosity, cracks.")

    c1, c2, c3 = st.columns([2.5, 1, 1])
    with c1:
        n_days = st.slider("Days of shop-floor data", 7, 90, 30)
    with c2:
        st.write("")
        gen_clicked = st.button("Generate Synthetic Data", key="gen_casting", use_container_width=True)
    with c3:
        st.write("")
        reset_button("casting_df", key="reset_casting", also_clear=["real_casting_df"])

    if gen_clicked:
        try:
            st.session_state["casting_df"] = gen_casting_defects(n_days)
        except Exception as e:
            st.error(f"Could not generate data: {e}")
            logger.exception("gen_casting_defects failed")

    if "casting_df" not in st.session_state:
        try:
            st.session_state["casting_df"] = gen_casting_defects(n_days)
        except Exception as e:
            st.error(f"Could not generate initial data: {e}")
            logger.exception("gen_casting_defects initial failed")

    upload_real_data("Casting QC Data", "real_casting_df",
                      "Date, Produced, Rejected, Rejection_%, Top_Defect, Melt_Temp_C, Pour_Temp_C, Mold_Temp_C, Humidity_%",
                      "upload_casting", "clear_casting")

    df, is_real = pick_active_df("real_casting_df", "casting_df")

    if df is not None and len(df) > 0:
        if is_real:
            st.info("Showing uploaded real data (synthetic data is hidden while real data is active).")

        has_cols = lambda *cols: all(c in df.columns for c in cols)
        kpi_items = []
        if has_cols("Rejection_%"):
            kpi_items.append(("Avg Rejection %", f"{df['Rejection_%'].mean():.2f}%", None))
        if has_cols("Produced"):
            kpi_items.append(("Total Produced", f"{df['Produced'].sum():,.0f}", None))
        if has_cols("Rejected"):
            kpi_items.append(("Total Rejected", f"{df['Rejected'].sum():,.0f}", None))
        if has_cols("Top_Defect"):
            kpi_items.append(("Most Common Defect", df["Top_Defect"].mode()[0], None))
        if kpi_items:
            kpi_row(kpi_items)

        if has_cols("Date", "Rejection_%"):
            st.markdown("#### Rejection % Trend (AI flags drift before it becomes a batch problem)")
            st.line_chart(df.set_index("Date")[["Rejection_%"]])

        cc1, cc2 = st.columns(2)
        with cc1:
            if has_cols("Top_Defect"):
                st.markdown("#### Defect Type Distribution")
                st.bar_chart(df["Top_Defect"].value_counts())
        with cc2:
            if has_cols("Pour_Temp_C", "Rejection_%"):
                st.markdown("#### Process Parameters vs Rejection")
                st.scatter_chart(df, x="Pour_Temp_C", y="Rejection_%")

        formula_box("""
        **Schema Field Definitions**

        - **Date** - the production day each row summarizes.
        - **Produced** - total castings produced that day.
        - **Rejected** - castings that failed quality inspection that day.
        - **Rejection_%** - Rejected as a percentage of Produced; the core quality KPI for the line.
        - **Top_Defect** - the most frequent defect type recorded that day (Blowhole, Shrinkage,
          Cold Shut, Gas Porosity, Surface Crack).
        - **Melt_Temp_C** - furnace melt temperature in Celsius; too high or too low drives
          porosity and shrinkage defects.
        - **Pour_Temp_C** - temperature of the molten metal at the moment of pouring; affects
          cold shuts and surface finish.
        - **Mold_Temp_C** - mold/die temperature; affects solidification rate and shrinkage.
        - **Humidity_%** - ambient foundry humidity; moisture contributes to gas porosity and
          blowholes.

        **Process**

        A camera-based vision model inspects every casting as it comes off the line and
        classifies defects in real time, instead of relying on sample-based manual inspection
        after the batch is already finished. Because inspection covers 100% of output and happens
        immediately, rising defect rates and their link to process parameters (melt, pour, and
        mold temperature, humidity) are visible the same day, not discovered weeks later on a
        customer return.

        **Formula**

        Rejection % = (Rejected / Produced) x 100

        AI tracks this trend day over day rather than waiting for a monthly quality report, so a
        rising trend is flagged before it becomes a full batch write-off. The synthetic generator
        also models a downward drift over time, simulating the improvement typically seen after
        deploying vision-based defect detection on the line.

        **Result**

        Typical deployments cut rejection rate by about 45-50% (3.5% to 1.8%), directly reducing
        scrap and rework cost.
        """)

        st.markdown("#### Raw Data")
        st.dataframe(df, use_container_width=True, height=280)

        formulas_dict = {
            "Field Definitions": (
                "Date: production day. Produced: castings made that day. Rejected: castings "
                "failed inspection that day. Rejection_%: Rejected / Produced x 100, the core "
                "quality KPI. Top_Defect: most frequent defect type that day. Melt_Temp_C, "
                "Pour_Temp_C, Mold_Temp_C, Humidity_%: process parameters most correlated with "
                "casting defects."
            ),
            "Process": (
                "Vision-based inspection classifies every casting in real time as it leaves the "
                "line, so defect trends and their link to process parameters are visible daily "
                "instead of being discovered weeks later."
            ),
            "Rejection % Formula": "Rejection % = (Rejected / Produced) x 100",
        }
        full_export_buttons(df, "Vision QC - Casting Defects", formulas_dict, "casting", "casting_defects")
    else:
        st.info("No data yet - click **Generate Synthetic Data** or upload a real QC file above.")

# ----------------------------------------------------------------------------
# 2. INSPECTION & NDT - DIMENSIONAL & STRUCTURAL QA
# ----------------------------------------------------------------------------
elif module == "Inspection & NDT - Dimensional/Structural QA":
    section_title("Inspection & NDT - Dimensional & Structural Quality Assurance",
                  "Which scanning instruments to use, what procedure to follow, and how to verify structural "
                  "strength and dimensional accuracy at competitive cost.")

    st.markdown(
        """
### Why this module

Every casting and every machined feature has to be verified two ways: **dimensionally** - does it
match the drawing/GD&T within tolerance - and **structurally** - is the material sound (free of
porosity, shrinkage, cracks, inclusions) and does it meet the required mechanical/tensile strength.
Running the most expensive checks on 100% of parts is not affordable; running only the cheapest
checks risks letting a structurally weak part reach the customer. The tiered strategy below is the
standard low-cost, high-confidence answer used across foundries and machine shops.
        """
    )

    with st.expander("Scanning Instruments Guide - what to use, and why", expanded=False):
        st.markdown(
            """
**Dimensional measurement instruments**

- **Digital calipers / micrometers / height gauges / Go-No-Go gauges** - lowest cost, seconds per
  part, operator-held. Used to spot-check critical dimensions on every single part (100%
  coverage). Typical accuracy: about +/-0.01-0.05 mm.
- **In-line 2D/3D Vision System** (see the Vision QC module) - camera-based, sub-second cycle
  time, very low cost per part once installed. Screens 100% of output for surface defects and
  gross dimensional errors before parts leave the line.
- **Structured-Light / Laser 3D Optical Scanner** - captures a full-surface point cloud in a few
  seconds and overlays it on the CAD model to produce a full GD&T colour-map, catching warpage and
  form errors a caliper cannot see. Typical accuracy: about +/-0.02-0.05 mm. Medium cost; much
  cheaper per part than a CMM at scale because it captures the whole part in one pass.
- **Portable Arm CMM (Faro/Romer-type)** - a flexible articulated arm brought to the part rather
  than the part brought to it, ideal for large or heavy castings on the shop floor. Typical
  accuracy: about +/-0.02-0.04 mm. Medium cost.
- **Bridge/Gantry CMM** - the accuracy reference standard, about +/-0.002-0.005 mm. Highest
  capital cost and slowest cycle time, so it is reserved for first-article approval and periodic
  certification of critical-to-function dimensions, not 100% inspection.

**Structural / internal-integrity instruments (Non-Destructive Testing - NDT)**

- **Dye Penetrant Inspection (DPT/LPI)** - the cheapest NDT method; reveals surface-breaking
  cracks on non-ferrous castings (aluminium, etc.) in minutes.
- **Magnetic Particle Inspection (MPI)** - cheap and fast; reveals surface and near-surface cracks
  on ferrous castings and forgings.
- **Ultrasonic Testing (UT)** - a portable probe sends sound waves through the part and measures
  wall thickness and internal voids/inclusions from the time-of-flight/echo pattern. Moderate
  cost, fully portable, well suited to periodic checks on structural (safety-critical) parts.
- **Industrial CT (Computed Tomography) Scanning** - an X-ray based 3D internal scan that finds
  porosity, shrinkage and inclusions and measures internal wall thickness without cutting the
  part open. The most thorough NDT method and the most expensive, so it is reserved for new-part
  qualification, safety-critical parts, and a small statistical sample per melt lot - not for
  100% inspection.
- **Portable Hardness Tester (Brinell/Rockwell/Leeb)** - a quick, cheap, non-destructive proxy for
  material strength and heat-treatment consistency; correlates statistically with tensile
  strength.
- **Portable XRF / Optical Emission Spectrometer** - verifies alloy chemical composition (which
  drives mechanical strength) in seconds, non-destructively, at moderate cost.
- **Tensile / Compression Testing Machine (destructive)** - the ground-truth measurement of
  actual structural strength; a test coupon (often cast-on to the part, or cut from a sacrificial
  casting in the same melt lot) is pulled to failure. This is the most expensive method per test
  because the sample is destroyed, so it is used at the lowest frequency, and its results are used
  to validate the cheaper proxy methods above.
            """
        )

    with st.expander("Recommended Tiered Inspection Procedure - balancing quality and cost", expanded=False):
        st.markdown(
            """
**Tier 1 - In-line, 100% of parts (lowest cost, every part)**
Vision system plus a go/no-go gauge or caliper check on the 2-3 dimensions that matter most for
fit or function, at the machine, in under 10 seconds per part. This catches gross defects and
out-of-tolerance parts immediately, before further value is added downstream.

**Tier 2 - Statistical sampling per shift/lot (medium cost, a handful of parts)**
Pull a sample sized to an AQL / ISO 2859-1 style sampling plan (for example, general inspection
level II) each shift or lot, and run a full-form check with an optical 3D scanner or portable arm
CMM against the CAD/GD&T model, plus a quick dye-penetrant or MPI surface-crack check. This catches
shape/form drift that a caliper alone would miss, without scanning every part.

**Tier 3 - Melt-lot / batch certification (highest cost, lowest frequency)**
Once per melt lot (or per a defined batch size): one Bridge-CMM certification of the
critical-to-function dimensions, one CT scan or UT scan for internal soundness, one hardness test,
one XRF composition check, and one destructive tensile/compression test on a cast-on or
sacrificial test coupon. This last step is what actually certifies structural strength for the
whole lot.

**AI correlation loop - the competitive-cost lever**
As Tier 3 data accumulates, a regression model learns to predict tensile strength from the
cheaper, non-destructive Tier 1-3 signals (hardness, UT response, dimensional pattern). Once the
model's prediction error stays within a validated confidence band (for example, within +/-5% of
actual, checked against periodic destructive audits), the frequency of the expensive destructive
tensile test can be safely reduced - for example, from one test in every 5 melt lots to one in
every 20 - while the AI model's inference, anchored by those periodic ground-truth checks, keeps
certifying the lots in between. That is the main lever for holding inspection cost down without
giving up structural assurance.
            """
        )

    c1, c2, c3 = st.columns([2.5, 1, 1])
    with c1:
        n_records = st.slider("Number of inspection events", 60, 500, 180, step=20, key="insp_slider")
    with c2:
        st.write("")
        gen_clicked = st.button("Generate Synthetic Data", key="gen_insp", use_container_width=True)
    with c3:
        st.write("")
        reset_button("insp_df", key="reset_insp", also_clear=["real_insp_df"])

    if gen_clicked:
        try:
            st.session_state["insp_df"] = gen_inspection_qc(n_records)
        except Exception as e:
            st.error(f"Could not generate data: {e}")
            logger.exception("gen_inspection_qc failed")

    if "insp_df" not in st.session_state:
        try:
            st.session_state["insp_df"] = gen_inspection_qc(n_records)
        except Exception as e:
            st.error(f"Could not generate initial data: {e}")
            logger.exception("gen_inspection_qc initial failed")

    upload_real_data(
        "Inspection / NDT Data", "real_insp_df",
        "Date, Part_ID, Melt_Lot, Inspection_Tier, Instrument_Used, Sample_Qty, "
        "Dimensional_Deviation_mm, Tolerance_Limit_mm, Dimensional_Result, NDT_Method, "
        "Internal_Defect_Found, Hardness_HB, AI_Predicted_Tensile_MPa, Actual_Tensile_MPa, "
        "Inspection_Cost_Rs, Overall_Result",
        "upload_insp", "clear_insp",
    )

    df, is_real = pick_active_df("real_insp_df", "insp_df")

    if df is not None and len(df) > 0:
        if is_real:
            st.info("Showing uploaded real data (synthetic data is hidden while real data is active).")

        has_cols = lambda *cols: all(c in df.columns for c in cols)
        kpi_items = []
        if "Dimensional_Result" in df.columns:
            pass_rate = 100 * (df["Dimensional_Result"] == "Pass").mean()
            kpi_items.append(("Dimensional Pass Rate", f"{pass_rate:.1f}%", None))
        if "Internal_Defect_Found" in df.columns:
            defect_rate = 100 * (df["Internal_Defect_Found"] != "None").mean()
            kpi_items.append(("Structural Defect Rate", f"{defect_rate:.1f}%", None))
        if has_cols("AI_Predicted_Tensile_MPa", "Actual_Tensile_MPa"):
            valid = df.dropna(subset=["Actual_Tensile_MPa"])
            if len(valid) > 0:
                mape = (abs(valid["Actual_Tensile_MPa"] - valid["AI_Predicted_Tensile_MPa"]) /
                        valid["Actual_Tensile_MPa"]).mean() * 100
                kpi_items.append(("Tensile Prediction Accuracy", f"{100 - mape:.1f}%", None))
                kpi_items.append(("Destructive Tests Run", f"{len(valid)} / {len(df)}", None))
        if "Inspection_Cost_Rs" in df.columns:
            kpi_items.append(("Avg Inspection Cost/Event", f"Rs. {df['Inspection_Cost_Rs'].mean():,.0f}", None))
        if kpi_items:
            kpi_row(kpi_items)

        if has_cols("Date", "Dimensional_Deviation_mm", "Tolerance_Limit_mm"):
            st.markdown("#### Dimensional Deviation vs Tolerance Limit")
            st.line_chart(df.set_index("Date")[["Dimensional_Deviation_mm", "Tolerance_Limit_mm"]])

        cc1, cc2 = st.columns(2)
        with cc1:
            if "Internal_Defect_Found" in df.columns:
                st.markdown("#### Internal Defects Found (excluding 'None')")
                defect_counts = df[df["Internal_Defect_Found"] != "None"]["Internal_Defect_Found"].value_counts()
                if len(defect_counts) > 0:
                    st.bar_chart(defect_counts)
                else:
                    st.caption("No internal defects recorded in this dataset.")
        with cc2:
            if has_cols("AI_Predicted_Tensile_MPa", "Actual_Tensile_MPa"):
                valid = df.dropna(subset=["Actual_Tensile_MPa"])
                if len(valid) > 0:
                    st.markdown("#### AI-Predicted vs Actual Tensile Strength")
                    st.scatter_chart(valid, x="AI_Predicted_Tensile_MPa", y="Actual_Tensile_MPa")
                else:
                    st.caption("No destructive test results recorded yet in this dataset.")

        if has_cols("Inspection_Tier", "Inspection_Cost_Rs"):
            st.markdown("#### Average Inspection Cost by Tier")
            st.bar_chart(df.groupby("Inspection_Tier")["Inspection_Cost_Rs"].mean())

        formula_box("""
        **Schema Field Definitions**

        - **Date** - the day the inspection event took place.
        - **Part_ID** - unique identifier of the inspected part.
        - **Melt_Lot** - the furnace melt/batch the part belongs to; used to group Tier 3
          certification results across all parts from that pour.
        - **Inspection_Tier** - which of the three tiers this event belongs to (Tier 1 in-line,
          Tier 2 statistical sampling, Tier 3 melt-lot certification).
        - **Instrument_Used** - the specific scanning/measurement/NDT instrument used for this
          event.
        - **Sample_Qty** - number of parts covered by this inspection event.
        - **Dimensional_Deviation_mm** - measured deviation from the nominal/CAD dimension.
        - **Tolerance_Limit_mm** - the allowed deviation per the drawing/GD&T callout.
        - **Dimensional_Result** - Pass if Dimensional_Deviation_mm is within Tolerance_Limit_mm.
        - **NDT_Method** - which non-destructive test (if any) was run for structural/internal
          soundness (Dye Penetrant, MPI, Ultrasonic Testing, CT Scan, or None).
        - **Internal_Defect_Found** - the internal/structural defect detected, if any (Porosity,
          Shrinkage, Inclusion, Crack, or None).
        - **Hardness_HB** - Brinell hardness reading; a fast, cheap proxy correlated with tensile
          strength.
        - **AI_Predicted_Tensile_MPa** - tensile strength predicted from the cheaper proxy signals
          (primarily hardness) by the correlation model.
        - **Actual_Tensile_MPa** - measured tensile strength from a destructive coupon test; only
          populated for the subset of Tier 3 events where a destructive test was actually run.
        - **Inspection_Cost_Rs** - the fully-loaded cost of this inspection event (labour plus
          instrument time plus consumables, and the destructive-coupon cost when applicable).
        - **Overall_Result** - Pass only if both the dimensional check and the structural/NDT
          check passed.

        **Process**

        Tier 1 screens 100% of parts at minimal cost using vision and hand gauges. Tier 2 samples a
        statistically justified quantity per shift/lot with an optical 3D scanner or arm CMM for
        full-form verification plus a quick surface NDT check. Tier 3 certifies each melt lot with
        the highest-accuracy dimensional method (CMM), the most thorough NDT method (CT/UT), and a
        destructive tensile test on a sacrificial coupon - the only method that measures actual
        structural strength directly.

        **Formula**

        Tensile Prediction Accuracy = 100% minus the mean of |Actual_Tensile - AI_Predicted_Tensile|
        / Actual_Tensile, x 100 - computed only over events where a destructive test was run.

        As this accuracy stays within a validated confidence band across repeated audits, the
        destructive-test frequency in Tier 3 can be reduced (for example, from 1-in-5 to 1-in-20
        melt lots), because the AI model's prediction, anchored by periodic ground-truth checks,
        can be trusted to certify the lots in between - which is the main lever for holding
        structural assurance and inspection cost together.

        **Result**

        A tiered instrument strategy - cheap 100%-coverage screening plus a shrinking, AI-anchored
        share of expensive destructive/NDT testing - typically holds overall inspection cost per
        part to a fraction of what 100% CMM/CT/tensile coverage would cost, while still certifying
        every melt lot for both dimensional accuracy and structural strength.
        """)

        st.markdown("#### Raw Data")
        st.dataframe(df, use_container_width=True, height=280)

        formulas_dict = {
            "Field Definitions": (
                "Date, Part_ID, Melt_Lot: inspection event identifiers. Inspection_Tier, "
                "Instrument_Used, Sample_Qty: which check ran and on how many parts. "
                "Dimensional_Deviation_mm / Tolerance_Limit_mm / Dimensional_Result: dimensional "
                "accuracy check. NDT_Method / Internal_Defect_Found: structural/internal soundness "
                "check. Hardness_HB, AI_Predicted_Tensile_MPa, Actual_Tensile_MPa: strength proxy "
                "and ground truth. Inspection_Cost_Rs: fully-loaded cost of the event. "
                "Overall_Result: combined pass/fail."
            ),
            "Scanning Instruments Guide": (
                "Dimensional: digital calipers/micrometers/gauges (100% coverage, lowest cost), "
                "in-line vision system (100% coverage), structured-light/laser 3D optical scanner "
                "and portable arm CMM (sampling, medium cost), bridge CMM (certification, highest "
                "accuracy and cost). Structural/NDT: dye penetrant and magnetic particle inspection "
                "(cheap surface-crack checks), ultrasonic testing (portable internal-flaw/wall-"
                "thickness check), industrial CT scanning (thorough internal inspection, highest "
                "cost), portable hardness tester and XRF analyzer (cheap non-destructive strength "
                "and composition proxies), tensile/compression testing machine (destructive ground "
                "truth for structural strength)."
            ),
            "Tiered Inspection Procedure": (
                "Tier 1: 100% in-line vision + gauge check, every part, lowest cost. "
                "Tier 2: AQL/ISO 2859-1 style sample per shift/lot, full-form optical 3D scan or "
                "arm CMM plus surface NDT. Tier 3: per melt lot, CMM certification + CT/UT scan + "
                "hardness + XRF + one destructive tensile test on a sacrificial coupon. An AI "
                "correlation model learns to predict tensile strength from the cheaper proxy "
                "signals so destructive-test frequency can be safely reduced once prediction "
                "accuracy is validated, lowering cost while keeping every lot certified."
            ),
            "Tensile Prediction Accuracy Formula": (
                "Accuracy % = 100 - Mean(|Actual_Tensile - AI_Predicted_Tensile| / Actual_Tensile) x 100"
            ),
        }
        full_export_buttons(df, "Inspection & NDT - Dimensional & Structural QA", formulas_dict, "insp", "inspection_ndt_qc")
    else:
        st.info("No data yet - click **Generate Synthetic Data** or upload a real inspection/NDT file above.")

# ----------------------------------------------------------------------------
# 3. PREDICTIVE MAINTENANCE
# ----------------------------------------------------------------------------
elif module == "Predictive Maintenance":
    section_title("Predictive Maintenance - Machine Health",
                  "Vibration, temperature, current sensors feed a failure-risk score for VMC/HMC/Press machines.")

    c1, c2, c3 = st.columns([2.5, 1, 1])
    with c1:
        n_hours = st.slider("Hours of sensor history", 48, 720, 240, step=24)
    with c2:
        st.write("")
        gen_clicked = st.button("Generate Synthetic Data", key="gen_health", use_container_width=True)
    with c3:
        st.write("")
        reset_button("health_df", key="reset_health", also_clear=["real_health_df"])

    if gen_clicked:
        try:
            st.session_state["health_df"] = gen_machine_health(n_hours)
        except Exception as e:
            st.error(f"Could not generate data: {e}")
            logger.exception("gen_machine_health failed")

    if "health_df" not in st.session_state:
        try:
            st.session_state["health_df"] = gen_machine_health(n_hours)
        except Exception as e:
            st.error(f"Could not generate initial data: {e}")
            logger.exception("gen_machine_health initial failed")

    upload_real_data("Machine Health Sensor Data", "real_health_df",
                      "Timestamp, Machine, Vibration_mm_s, Temperature_C, Current_A, Failure_Risk_%, Predicted_Days_to_Failure",
                      "upload_health", "clear_health")

    df, is_real = pick_active_df("real_health_df", "health_df")

    if df is not None and len(df) > 0 and "Machine" in df.columns:
        if is_real:
            st.info("Showing uploaded real data (synthetic data is hidden while real data is active).")

        machine_sel = st.selectbox("Select machine", sorted(df["Machine"].unique()))
        mdf = df[df["Machine"] == machine_sel]
        if "Timestamp" in mdf.columns:
            mdf = mdf.sort_values("Timestamp")
        latest = mdf.iloc[-1]

        kpi_items = []
        if "Failure_Risk_%" in mdf.columns:
            kpi_items.append(("Current Failure Risk", f"{latest['Failure_Risk_%']:.0f}%", None))
        if "Predicted_Days_to_Failure" in mdf.columns:
            kpi_items.append(("Predicted Days to Failure", f"{latest['Predicted_Days_to_Failure']:.1f} d", None))
        if "Vibration_mm_s" in mdf.columns:
            kpi_items.append(("Vibration", f"{latest['Vibration_mm_s']} mm/s", None))
        if "Temperature_C" in mdf.columns:
            kpi_items.append(("Temperature", f"{latest['Temperature_C']} C", None))
        if kpi_items:
            kpi_row(kpi_items)

        if "Timestamp" in mdf.columns and "Failure_Risk_%" in mdf.columns:
            st.markdown(f"#### Failure Risk Trend - {machine_sel}")
            st.line_chart(mdf.set_index("Timestamp")[["Failure_Risk_%"]])

        cc1, cc2 = st.columns(2)
        with cc1:
            if "Timestamp" in mdf.columns and {"Vibration_mm_s", "Current_A"}.issubset(mdf.columns):
                st.markdown("#### Vibration & Current")
                st.line_chart(mdf.set_index("Timestamp")[["Vibration_mm_s", "Current_A"]])
        with cc2:
            if "Timestamp" in mdf.columns and "Temperature_C" in mdf.columns:
                st.markdown("#### Temperature")
                st.line_chart(mdf.set_index("Timestamp")[["Temperature_C"]])

        if "Failure_Risk_%" in mdf.columns:
            risk_val = latest["Failure_Risk_%"]
            if risk_val > 65:
                st.error(f"{machine_sel} is showing elevated failure risk - schedule maintenance within "
                         f"{latest.get('Predicted_Days_to_Failure', 'a few')} days.")
            elif risk_val > 40:
                st.warning(f"{machine_sel} risk trending up - monitor closely.")
            else:
                st.success(f"{machine_sel} operating within normal parameters.")

        formula_box("""
        **Schema Field Definitions**

        - **Timestamp** - the hour each sensor reading was captured.
        - **Machine** - machine ID (VMC/HMC = vertical/horizontal machining center, Press = stamping
          press).
        - **Vibration_mm_s** - vibration velocity reading; rises as bearings/spindles wear.
        - **Temperature_C** - motor/bearing housing temperature; rises with friction from a
          developing fault.
        - **Current_A** - motor current draw; rises when a mechanical fault forces the motor to
          work harder.
        - **Failure_Risk_%** - composite risk score computed from the three sensor readings above.
        - **Predicted_Days_to_Failure** - estimated runway before intervention is needed, derived
          from Failure_Risk_%.

        **Process**

        Vibration, temperature, and current sensors stream continuously from each machine. As a
        bearing, spindle, or die begins to fail, all three readings tend to drift upward together
        well before the machine actually stops, which is what the risk score is designed to catch
        early.

        **Formula**

        Failure_Risk = min(99, max(1, (Vibration/6 + Temperature/90 + Current/18) x 33))

        This is a simple weighted composite of the three sensor channels, each normalized against
        a rough danger ceiling (6 mm/s vibration, 90 C temperature, 18A current) before being
        averaged and scaled to a 1-99% risk score.

        Action thresholds: above 65% - schedule maintenance now. 40-65% - monitor closely. Below
        40% - normal.

        **Result**

        Predictive maintenance catches degrading bearings, spindles, and dies 2-4 weeks before
        failure, cutting unplanned downtime by about 25%.
        """)

        st.markdown("#### Raw Data")
        st.dataframe(df, use_container_width=True, height=280)

        formulas_dict = {
            "Field Definitions": (
                "Timestamp: hour of reading. Machine: machine ID. Vibration_mm_s, Temperature_C, "
                "Current_A: the three sensor channels that feed the risk score. Failure_Risk_%: "
                "composite 1-99% risk score. Predicted_Days_to_Failure: estimated runway before "
                "intervention is needed."
            ),
            "Process": (
                "Vibration, temperature and current sensors stream continuously; as a fault "
                "develops, all three readings tend to rise together well before the machine "
                "actually fails, which the risk score is designed to catch early."
            ),
            "Failure Risk Formula": "Failure_Risk = min(99, max(1, (Vibration/6 + Temperature/90 + Current/18) x 33))",
        }
        full_export_buttons(df, "Predictive Maintenance - Machine Health", formulas_dict, "health", "machine_health")
    else:
        st.info("No data yet - click **Generate Synthetic Data** or upload a real sensor file above.")

# ----------------------------------------------------------------------------
# 4. TOOL WEAR
# ----------------------------------------------------------------------------
elif module == "Tool Wear Prediction":
    section_title("Tool Wear Prediction",
                  "Spindle load, vibration, and cycle count feed condition-based insert/tool replacement.")

    c1, c2, c3 = st.columns([2.5, 1, 1])
    with c1:
        n_records = st.slider("Number of tools", 50, 500, 200, step=50)
    with c2:
        st.write("")
        gen_clicked = st.button("Generate Synthetic Data", key="gen_tool", use_container_width=True)
    with c3:
        st.write("")
        reset_button("tool_df", key="reset_tool", also_clear=["real_tool_df"])

    if gen_clicked:
        try:
            st.session_state["tool_df"] = gen_tool_wear(n_records)
        except Exception as e:
            st.error(f"Could not generate data: {e}")
            logger.exception("gen_tool_wear failed")

    if "tool_df" not in st.session_state:
        try:
            st.session_state["tool_df"] = gen_tool_wear(n_records)
        except Exception as e:
            st.error(f"Could not generate initial data: {e}")
            logger.exception("gen_tool_wear initial failed")

    upload_real_data("Tool Wear Data", "real_tool_df",
                      "Tool_ID, Tool_Type, Cycles_Run, Wear_%, Spindle_Load_%, Vibration_mm_s, Recommendation",
                      "upload_tool", "clear_tool")

    df, is_real = pick_active_df("real_tool_df", "tool_df")

    if df is not None and len(df) > 0:
        if is_real:
            st.info("Showing uploaded real data (synthetic data is hidden while real data is active).")

        has_cols = lambda *cols: all(c in df.columns for c in cols)
        kpi_items = []
        if "Wear_%" in df.columns:
            kpi_items.append(("Avg Wear %", f"{df['Wear_%'].mean():.1f}%", None))
        if "Recommendation" in df.columns:
            kpi_items.append(("Tools to Change Now", int((df["Recommendation"] == "Change Now").sum()), None))
            kpi_items.append(("Tools to Monitor", int((df["Recommendation"] == "Monitor").sum()), None))
            kpi_items.append(("Tools OK", int((df["Recommendation"] == "OK").sum()), None))
        if kpi_items:
            kpi_row(kpi_items)

        cc1, cc2 = st.columns(2)
        with cc1:
            if has_cols("Tool_Type", "Wear_%"):
                st.markdown("#### Wear % by Tool Type")
                st.bar_chart(df.groupby("Tool_Type")["Wear_%"].mean())
        with cc2:
            if has_cols("Wear_%", "Spindle_Load_%"):
                st.markdown("#### Spindle Load vs Wear")
                st.scatter_chart(df, x="Wear_%", y="Spindle_Load_%")

        formula_box("""
        **Schema Field Definitions**

        - **Tool_ID** - unique identifier for each physical tool/insert in circulation.
        - **Tool_Type** - category of cutting tool (insert grade, drill, boring bar, end mill).
        - **Cycles_Run** - number of machining cycles completed since the tool was installed.
        - **Wear_%** - estimated edge wear as a percentage of usable tool life.
        - **Spindle_Load_%** - how hard the spindle motor is working to cut; rises as a worn edge
          loses sharpness.
        - **Vibration_mm_s** - vibration at the tool/spindle; also rises with wear.
        - **Recommendation** - the AI action guidance derived from Wear_% (OK / Monitor / Change
          Now).

        **Process**

        Rather than replacing tools on a fixed cycle-count schedule, AI estimates real remaining
        life from cycle count plus the two physical signals that respond to a dulling edge -
        spindle load and vibration - and cross-checks them against each other before recommending
        a change.

        **Formula**

        Wear % is approximately Cycles_Run / Expected_Tool_Life, plus sensor noise, capped to the
        0-100% range.

        Recommendation thresholds: Wear 55% or below - OK. Wear above 55% and up to 80% - Monitor.
        Wear above 80% - Change Now.

        **Result**

        Condition-based replacement, instead of fixed-interval changes, typically cuts machining
        tool cost per part by 12-18% while avoiding tool-failure scrap.
        """)

        if "Recommendation" in df.columns:
            st.markdown("#### Tools Needing Attention")
            st.dataframe(
                df[df["Recommendation"] != "OK"].sort_values("Wear_%", ascending=False)
                if "Wear_%" in df.columns else df[df["Recommendation"] != "OK"],
                use_container_width=True, height=240,
            )

        st.markdown("#### Full Tool Data")
        st.dataframe(df, use_container_width=True, height=280)

        formulas_dict = {
            "Field Definitions": (
                "Tool_ID: unique tool identifier. Tool_Type: category of cutting tool. "
                "Cycles_Run: machining cycles completed. Wear_%: estimated edge wear. "
                "Spindle_Load_%, Vibration_mm_s: physical signals that rise with wear. "
                "Recommendation: OK / Monitor / Change Now guidance."
            ),
            "Process": (
                "Instead of a fixed-interval change schedule, AI cross-checks cycle count against "
                "spindle load and vibration - both of which rise as a cutting edge dulls - to "
                "recommend condition-based replacement."
            ),
            "Wear % Logic": "Wear % is approximately Cycles_Run / Expected_Tool_Life, plus sensor noise (0-100%)",
            "Recommendation Thresholds": "55% or below OK | above 55% to 80% Monitor | above 80% Change Now",
        }
        full_export_buttons(df, "Tool Wear Prediction", formulas_dict, "tool", "tool_wear")
    else:
        st.info("No data yet - click **Generate Synthetic Data** or upload a real tool-wear file above.")

# ----------------------------------------------------------------------------
# 5. DEMAND & INVENTORY FORECAST
# ----------------------------------------------------------------------------
elif module == "Demand & Inventory Forecast":
    section_title("Demand & Inventory Forecast",
                  "AI forecast versus actual demand by part, with safety stock and reorder points.")

    c1, c2, c3 = st.columns([2.5, 1, 1])
    with c1:
        n_weeks = st.slider("Weeks of history", 12, 52, 26)
    with c2:
        st.write("")
        gen_clicked = st.button("Generate Synthetic Data", key="gen_demand", use_container_width=True)
    with c3:
        st.write("")
        reset_button("demand_df", key="reset_demand", also_clear=["real_demand_df"])

    if gen_clicked:
        try:
            st.session_state["demand_df"] = gen_demand_forecast(n_weeks)
        except Exception as e:
            st.error(f"Could not generate data: {e}")
            logger.exception("gen_demand_forecast failed")

    if "demand_df" not in st.session_state:
        try:
            st.session_state["demand_df"] = gen_demand_forecast(n_weeks)
        except Exception as e:
            st.error(f"Could not generate initial data: {e}")
            logger.exception("gen_demand_forecast initial failed")

    upload_real_data("Demand/Inventory Data", "real_demand_df",
                      "Week, Part, Actual_Demand, AI_Forecast, Safety_Stock, Reorder_Point",
                      "upload_demand", "clear_demand")

    df, is_real = pick_active_df("real_demand_df", "demand_df")

    if df is not None and len(df) > 0 and "Part" in df.columns:
        if is_real:
            st.info("Showing uploaded real data (synthetic data is hidden while real data is active).")

        part_sel = st.selectbox("Select part", sorted(df["Part"].unique()))
        pdf_ = df[df["Part"] == part_sel]
        if "Week" in pdf_.columns:
            pdf_ = pdf_.sort_values("Week")

        kpi_items = []
        if {"Actual_Demand", "AI_Forecast"}.issubset(pdf_.columns):
            valid = pdf_.dropna(subset=["Actual_Demand"])
            if len(valid) > 0 and (valid["Actual_Demand"] != 0).any():
                mape = (abs(valid["Actual_Demand"] - valid["AI_Forecast"]) / valid["Actual_Demand"]).mean() * 100
                kpi_items.append(("Forecast Accuracy (MAPE)", f"{100 - mape:.1f}%", None))
                kpi_items.append(("Avg Weekly Demand", f"{valid['Actual_Demand'].mean():.0f}", None))
        if "Safety_Stock" in pdf_.columns and len(pdf_) > 0:
            kpi_items.append(("Latest Safety Stock", int(pdf_.iloc[-1]["Safety_Stock"]), None))
        if "Reorder_Point" in pdf_.columns and len(pdf_) > 0:
            kpi_items.append(("Latest Reorder Point", int(pdf_.iloc[-1]["Reorder_Point"]), None))
        if kpi_items:
            kpi_row(kpi_items)

        if "Week" in pdf_.columns and {"Actual_Demand", "AI_Forecast"}.issubset(pdf_.columns):
            st.markdown(f"#### Actual vs AI Forecast - {part_sel}")
            st.line_chart(pdf_.set_index("Week")[["Actual_Demand", "AI_Forecast"]])

        if "Week" in pdf_.columns and {"Safety_Stock", "Reorder_Point"}.issubset(pdf_.columns):
            st.markdown("#### Safety Stock & Reorder Point")
            st.line_chart(pdf_.set_index("Week")[["Safety_Stock", "Reorder_Point"]])

        formula_box("""
        **Schema Field Definitions**

        - **Week** - the week each row covers.
        - **Part** - the part number/SKU being forecast.
        - **Actual_Demand** - units actually ordered/consumed that week (blank for the most recent
          weeks, which have not closed yet).
        - **AI_Forecast** - the model's predicted demand for that week.
        - **Safety_Stock** - buffer inventory held to absorb demand variability and supply
          lead-time risk.
        - **Reorder_Point** - the inventory level that triggers a new purchase/production order.

        **Process**

        The forecast model learns each part's demand pattern, including seasonal swings, from its
        history, and produces a rolling forward forecast. Safety stock and reorder point are then
        derived directly from that forecast so replenishment stays proportional to expected
        demand instead of a flat, one-size-fits-all buffer.

        **Formula**

        MAPE (Mean Absolute Percentage Error) equals the average of |Actual - Forecast| / Actual,
        x 100. Forecast Accuracy is reported as 100% minus MAPE.

        Safety_Stock is approximately AI_Forecast x 12%.
        Reorder_Point is approximately AI_Forecast x 35%.

        **Result**

        Sharper forecasts typically reduce raw-material and finished-goods inventory by 15-25%
        while still protecting JIT OEM supply commitments.
        """)

        st.markdown("#### Raw Data")
        st.dataframe(df, use_container_width=True, height=280)

        formulas_dict = {
            "Field Definitions": (
                "Week: forecast period. Part: SKU forecast. Actual_Demand: units actually "
                "consumed. AI_Forecast: predicted demand. Safety_Stock: buffer inventory. "
                "Reorder_Point: level that triggers a new order."
            ),
            "Process": (
                "The model learns each part's seasonal demand pattern and produces a rolling "
                "forecast; safety stock and reorder point scale directly off that forecast."
            ),
            "MAPE / Accuracy": "Accuracy % = 100 - avg(|Actual - Forecast| / Actual) x 100",
            "Safety Stock": "Safety_Stock is approximately AI_Forecast x 12%",
            "Reorder Point": "Reorder_Point is approximately AI_Forecast x 35%",
        }
        full_export_buttons(df, "Demand & Inventory Forecast", formulas_dict, "demand", "demand_forecast")
    else:
        st.info("No data yet - click **Generate Synthetic Data** or upload a real demand file above.")

# ----------------------------------------------------------------------------
# 6. ENERGY OPTIMIZATION
# ----------------------------------------------------------------------------
elif module == "Energy Optimization":
    section_title("Energy Optimization",
                  "Baseline versus AI-optimized furnace/machining load, with peak-shift savings.")

    c1, c2, c3 = st.columns([2.5, 1, 1])
    with c1:
        n_days = st.slider("Days of energy data", 7, 90, 30, key="energy_days")
    with c2:
        st.write("")
        gen_clicked = st.button("Generate Synthetic Data", key="gen_energy", use_container_width=True)
    with c3:
        st.write("")
        reset_button("energy_df", key="reset_energy", also_clear=["real_energy_df"])

    if gen_clicked:
        try:
            st.session_state["energy_df"] = gen_energy(n_days)
        except Exception as e:
            st.error(f"Could not generate data: {e}")
            logger.exception("gen_energy failed")

    if "energy_df" not in st.session_state:
        try:
            st.session_state["energy_df"] = gen_energy(n_days)
        except Exception as e:
            st.error(f"Could not generate initial data: {e}")
            logger.exception("gen_energy initial failed")

    upload_real_data("Energy Consumption Data", "real_energy_df",
                      "Date, Baseline_kWh, AI_Optimized_kWh, Savings_%, Peak_Load_Shifted_kWh",
                      "upload_energy", "clear_energy")

    df, is_real = pick_active_df("real_energy_df", "energy_df")

    if df is not None and len(df) > 0:
        if is_real:
            st.info("Showing uploaded real data (synthetic data is hidden while real data is active).")

        kpi_items = []
        if {"Baseline_kWh", "AI_Optimized_kWh"}.issubset(df.columns):
            total_baseline = df["Baseline_kWh"].sum()
            total_optimized = df["AI_Optimized_kWh"].sum()
            savings_pct = 100 * (1 - total_optimized / total_baseline) if total_baseline else 0
            kpi_items.append(("Total Baseline kWh", f"{total_baseline:,.0f}", None))
            kpi_items.append(("Total AI-Optimized kWh", f"{total_optimized:,.0f}", None))
            kpi_items.append(("Avg Savings %", f"{savings_pct:.1f}%", None))
        if "Peak_Load_Shifted_kWh" in df.columns:
            kpi_items.append(("Peak Load Shifted (avg/day)", f"{df['Peak_Load_Shifted_kWh'].mean():.0f} kWh", None))
        if kpi_items:
            kpi_row(kpi_items)

        if "Date" in df.columns and {"Baseline_kWh", "AI_Optimized_kWh"}.issubset(df.columns):
            st.markdown("#### Baseline vs AI-Optimized Consumption")
            st.line_chart(df.set_index("Date")[["Baseline_kWh", "AI_Optimized_kWh"]])

        if "Date" in df.columns and "Savings_%" in df.columns:
            st.markdown("#### Daily Savings %")
            st.bar_chart(df.set_index("Date")[["Savings_%"]])

        formula_box("""
        **Schema Field Definitions**

        - **Date** - the day each row covers.
        - **Baseline_kWh** - energy consumption if operations ran without AI-driven scheduling.
        - **AI_Optimized_kWh** - actual/optimized consumption after AI shifts load away from
          peak-tariff windows and trims idle load.
        - **Savings_%** - the percentage reduction achieved that day.
        - **Peak_Load_Shifted_kWh** - how much load was moved out of the most expensive tariff
          band that day.

        **Process**

        AI shifts furnace pre-heat cycles, machining schedules, and compressed-air usage away
        from peak-tariff windows and trims idle-running load, without changing production volume.

        **Formula**

        Savings % = (1 - AI_Optimized_kWh / Baseline_kWh) x 100

        **Result**

        Typical deployments cut the energy bill by 8-20%, depending on tariff structure and
        existing equipment efficiency.
        """)

        st.markdown("#### Raw Data")
        st.dataframe(df, use_container_width=True, height=280)

        formulas_dict = {
            "Field Definitions": (
                "Date: day covered. Baseline_kWh: consumption without AI scheduling. "
                "AI_Optimized_kWh: consumption with AI scheduling. Savings_%: percentage "
                "reduction. Peak_Load_Shifted_kWh: load moved out of peak-tariff windows."
            ),
            "Process": (
                "AI shifts furnace pre-heat cycles, machining schedules, and compressed-air usage "
                "away from peak-tariff windows and trims idle load, without changing volume."
            ),
            "Energy Savings Formula": "Savings % = (1 - AI_Optimized_kWh / Baseline_kWh) x 100",
        }
        full_export_buttons(df, "Energy Optimization", formulas_dict, "energy", "energy_optimization")
    else:
        st.info("No data yet - click **Generate Synthetic Data** or upload a real energy file above.")

# ----------------------------------------------------------------------------
# 7. TRACEABILITY / ROOT CAUSE
# ----------------------------------------------------------------------------
elif module == "Traceability & Root Cause":
    section_title("Traceability & Root Cause Analysis",
                  "AI correlates melt lot, machine, tool, and operator to find root cause in seconds.")

    c1, c2, c3 = st.columns([2.5, 1, 1])
    with c1:
        n_events = st.slider("Number of defect events", 5, 50, 15)
    with c2:
        st.write("")
        gen_clicked = st.button("Generate Synthetic Data", key="gen_trace", use_container_width=True)
    with c3:
        st.write("")
        reset_button("trace_df", key="reset_trace", also_clear=["real_trace_df"])

    if gen_clicked:
        try:
            st.session_state["trace_df"] = gen_traceability(n_events)
        except Exception as e:
            st.error(f"Could not generate data: {e}")
            logger.exception("gen_traceability failed")

    if "trace_df" not in st.session_state:
        try:
            st.session_state["trace_df"] = gen_traceability(n_events)
        except Exception as e:
            st.error(f"Could not generate initial data: {e}")
            logger.exception("gen_traceability initial failed")

    upload_real_data("Traceability / RCA Data", "real_trace_df",
                      "Defect_ID, Date, Melt_Lot, Machine, Operator, Tool_Used, Root_Cause_AI, Manual_RCA_Time_Hrs, AI_RCA_Time_Min",
                      "upload_trace", "clear_trace")

    df, is_real = pick_active_df("real_trace_df", "trace_df")

    if df is not None and len(df) > 0:
        if is_real:
            st.info("Showing uploaded real data (synthetic data is hidden while real data is active).")

        kpi_items = []
        if {"Manual_RCA_Time_Hrs", "AI_RCA_Time_Min"}.issubset(df.columns) and df["AI_RCA_Time_Min"].mean() > 0:
            kpi_items.append(("Avg Manual RCA Time", f"{df['Manual_RCA_Time_Hrs'].mean():.1f} hrs", None))
            kpi_items.append(("Avg AI RCA Time", f"{df['AI_RCA_Time_Min'].mean():.1f} min", None))
            speedup = (df["Manual_RCA_Time_Hrs"].mean() * 60 / df["AI_RCA_Time_Min"].mean())
            kpi_items.append(("Speed-up Factor", f"{speedup:.0f}x", None))
        if "Root_Cause_AI" in df.columns:
            kpi_items.append(("Top Root Cause", df["Root_Cause_AI"].mode()[0], None))
        if kpi_items:
            kpi_row(kpi_items)

        if "Root_Cause_AI" in df.columns:
            st.markdown("#### Root Cause Frequency")
            st.bar_chart(df["Root_Cause_AI"].value_counts())

        if {"Defect_ID", "Manual_RCA_Time_Hrs", "AI_RCA_Time_Min"}.issubset(df.columns):
            st.markdown("#### Manual vs AI RCA Time (per event)")
            comp = df[["Defect_ID", "Manual_RCA_Time_Hrs"]].copy()
            comp["AI_RCA_Time_Hrs"] = df["AI_RCA_Time_Min"] / 60
            st.bar_chart(comp.set_index("Defect_ID")[["Manual_RCA_Time_Hrs", "AI_RCA_Time_Hrs"]])

        formula_box("""
        **Schema Field Definitions**

        - **Defect_ID** - unique identifier for each defect event under investigation.
        - **Date** - date the defect was recorded.
        - **Melt_Lot** - the furnace melt/batch the affected casting came from.
        - **Machine** - the machine that processed the part.
        - **Operator** - the operator on shift at the time.
        - **Tool_Used** - which tool/insert was in use.
        - **Root_Cause_AI** - the most statistically likely cause identified by cross-referencing
          all of the above against historical defect patterns.
        - **Manual_RCA_Time_Hrs** - how long a manual root-cause investigation typically takes, in
          hours.
        - **AI_RCA_Time_Min** - how long the AI-assisted lookup takes, in minutes.

        **Process**

        AI cross-references melt lot, machine, tool, and operator against every historical defect
        event to surface the most statistically likely root cause in seconds, instead of an
        engineer manually pulling logs across multiple systems over hours or days.

        **Formula**

        Speed-up Factor = (Avg_Manual_RCA_Time_Hrs x 60) / Avg_AI_RCA_Time_Min

        **Result**

        Root-cause analysis time typically collapses from about 2 days (manual) to about 30
        seconds (AI-assisted), a reduction of over 99%, freeing engineering time for prevention
        instead of firefighting.
        """)

        st.markdown("#### Event Log")
        st.dataframe(df, use_container_width=True, height=300)

        formulas_dict = {
            "Field Definitions": (
                "Defect_ID: event identifier. Date: date recorded. Melt_Lot: furnace batch. "
                "Machine, Operator, Tool_Used: production context. Root_Cause_AI: most likely "
                "cause identified by AI. Manual_RCA_Time_Hrs / AI_RCA_Time_Min: time to root "
                "cause, manual versus AI-assisted."
            ),
            "Process": (
                "AI cross-references melt lot, machine, tool and operator against historical "
                "defect events to surface the most likely root cause in seconds."
            ),
            "RCA Speed-up Formula": "Speed-up = (Avg_Manual_RCA_Hrs x 60) / Avg_AI_RCA_Min",
        }
        full_export_buttons(df, "Traceability & Root Cause", formulas_dict, "trace", "traceability_rca")
    else:
        st.info("No data yet - click **Generate Synthetic Data** or upload a real RCA log above.")

# ----------------------------------------------------------------------------
# 8. ROI SUMMARY
# ----------------------------------------------------------------------------
elif module == "ROI Summary":
    section_title("ROI Summary",
                  "Typical results reported by Indian auto ancillary units after AI deployment.")

    roi_df = pd.DataFrame([
        {"Area": "Casting Rejection Rate", "Before": "3.5%", "After": "1.8%", "Improvement": "-48%"},
        {"Area": "Machining Tool Cost / Part", "Before": "Baseline", "After": "Down 12-18%", "Improvement": "12-18%"},
        {"Area": "Unplanned Downtime", "Before": "Baseline", "After": "Down 25%", "Improvement": "25%"},
        {"Area": "Inventory / Working Capital", "Before": "Baseline", "After": "Down 15-25%", "Improvement": "15-25%"},
        {"Area": "Energy Bill", "Before": "Baseline", "After": "Down 8-20%", "Improvement": "8-20%"},
        {"Area": "Customer Complaints (Vision QC)", "Before": "Baseline", "After": "Down 50-80%", "Improvement": "50-80%"},
        {"Area": "Root Cause Analysis Time", "Before": "Approx. 2 days", "After": "Approx. 30 sec", "Improvement": "Over 99%"},
    ])
    st.dataframe(roi_df, use_container_width=True, height=300)

    formula_box("""
    **Schema Field Definitions**

    - **Area** - the operational metric being measured.
    - **Before** - baseline performance prior to AI deployment.
    - **After** - typical performance once AI is deployed and adopted.
    - **Improvement** - the resulting percentage change.

    **Process**

    Each row reflects the typical before-versus-after delta reported across the modules in this
    app once deployed on a live line: Vision QC drives the rejection-rate and complaint-reduction
    rows, Inspection & NDT drives dimensional/structural certification at competitive cost,
    Predictive Maintenance drives downtime, Tool Wear drives tool cost, Demand Forecast drives
    inventory/working capital, Energy Optimization drives the energy-bill row, and Traceability
    drives RCA time.

    Treat these as directional industry benchmarks, not a site-specific guarantee - actual results
    depend on baseline process maturity and the quality of sensor/data coverage at rollout.
    """)

    st.markdown("#### Suggested Rollout Sequence (Low Budget First)")
    st.markdown(
        """
1. **Vision QC pilot** - 2-3 cameras plus off-the-shelf vision software on the highest-rejection
   line. Fastest ROI, easiest operator buy-in.
2. **Tier 1 inspection gauges** - hand gauges/calipers plus the vision system above for 100%
   dimensional screening at near-zero incremental cost.
3. **Sensor retrofit** - vibration and current sensors on the top 3 bottleneck machines
   (approximately Rs. 50,000 per machine).
4. **Tier 2/3 inspection instruments** - one optical 3D scanner or portable arm CMM for sampling,
   plus access to CT/UT/tensile testing (in-house or via a local NDT lab) for melt-lot
   certification.
5. **Use existing data** - CMM logs, rejection logs, and tool-change sheets feed the first
   predictive models, including the tensile-strength correlation model.
6. **Scale to scheduling, demand forecasting, and energy** once the shop floor trusts the QC,
   inspection, and maintenance pilots.
        """
    )

    formulas_dict = {
        "Field Definitions": (
            "Area: operational metric measured. Before: baseline performance. After: typical "
            "performance post-AI. Improvement: resulting percentage change."
        ),
        "Methodology": (
            "Before/after deltas aggregated from the Vision QC, Inspection & NDT, Predictive "
            "Maintenance, Tool Wear, Demand Forecast, Energy, and Traceability modules."
        ),
    }
    full_export_buttons(roi_df, "ROI Summary", formulas_dict, "roi", "roi_summary")

st.sidebar.divider()
st.sidebar.caption("All figures are synthetic / illustrative unless real data is uploaded | Not a substitute for a site-specific audit.")